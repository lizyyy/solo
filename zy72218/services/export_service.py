import csv
import os
from typing import List, Dict, Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class ExportService:
    def __init__(self, output_dir: str = "exports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export_csv(self, records: List[Dict[str, Any]], filename: str) -> str:
        filepath = os.path.join(self.output_dir, filename)
        if not records:
            with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
                f.write("")
            return filepath

        fieldnames = list(records[0].keys())
        with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for row in records:
                clean_row = {}
                for k, v in row.items():
                    if isinstance(v, list):
                        clean_row[k] = ";".join(str(i) for i in v)
                    elif isinstance(v, bool):
                        clean_row[k] = "是" if v else "否"
                    else:
                        clean_row[k] = v
                writer.writerow(clean_row)
        return filepath

    def export_excel(self, records: List[Dict[str, Any]], filename: str, sheet_name: str = "发票池融资匹配") -> str:
        filepath = os.path.join(self.output_dir, filename)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        if not records:
            wb.save(filepath)
            return filepath

        headers = list(records[0].keys())
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        split_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        conflict_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")

        for row_idx, record in enumerate(records, 2):
            is_split = record.get("is_split_record", False)
            status = record.get("status", "")

            for col, header in enumerate(headers, 1):
                value = record.get(header)
                if isinstance(value, list):
                    value = ";".join(str(i) for i in value)
                elif isinstance(value, bool):
                    value = "是" if value else "否"

                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                if is_split:
                    cell.fill = split_fill
                if "冲突" in str(status):
                    cell.fill = conflict_fill

        column_widths = {}
        for col, header in enumerate(headers, 1):
            max_len = len(str(header))
            for row_idx, record in enumerate(records, 2):
                val = str(record.get(header, ""))
                max_len = max(max_len, min(len(val), 40))
            column_widths[col] = max_len + 4

        for col, width in column_widths.items():
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

        wb.save(filepath)
        return filepath

    def export_audit_trail(self, logs: List[Dict[str, Any]], filename: str = "audit_trail.xlsx") -> str:
        return self.export_excel(logs, filename, sheet_name="审计日志")

    def export_discrepancies(self, discrepancies: List[Dict[str, Any]], filename: str = "discrepancies.xlsx") -> str:
        return self.export_excel(discrepancies, filename, sheet_name="差异清单")

    def format_records_for_export(self, records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        COLUMN_LABELS = {
            "record_id": "记录ID",
            "business_no": "业务号",
            "record_type": "记录类型",
            "expected_amount": "预期金额",
            "matched_amount": "匹配金额",
            "match_date": "匹配日期",
            "status": "状态",
            "related_record_id": "关联记录ID",
            "is_split_record": "是否拆分行",
            "holiday_extension_applied": "节假日顺延已应用",
            "tail_adjustment_applied": "尾差调整已应用",
            "counterpart_type": "对应行类型",
            "counterpart_amount": "对应行金额",
            "combined_amount": "合并金额",
            "invoice_ids": "关联发票",
            "conflict_pending": "是否有未解决冲突",
            "conflict_pending_description": "冲突状态说明",
            "conflict_resolution_rule": "冲突选择规则",
            "conflict_resolution_status": "冲突解决状态",
            "conflict_resolution_reason": "冲突解决原因",
            "conflict_resolved_by": "冲突决策人",
            "conflict_final_amount": "冲突最终金额",
        }
        result = []
        for record in records:
            mapped = {}
            for eng_key, label in COLUMN_LABELS.items():
                if eng_key in record:
                    mapped[label] = record[eng_key]
            for key in record:
                if key not in COLUMN_LABELS:
                    mapped[key] = record[key]
            result.append(mapped)
        return result
