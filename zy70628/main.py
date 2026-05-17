from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, and_, func
from datetime import datetime, timedelta
from typing import List, Optional
from pydantic import BaseModel, Field
import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from database import (
    get_db, init_db, Parcel, Recipient, ReminderRecord,
    Rejection, ReturnReport, OperationLog
)

app = FastAPI(title="驿站包裹滞留催取拒收退回确认API", version="1.0.0")


@app.on_event("startup")
async def startup_event():
    init_db()


class RecipientCreate(BaseModel):
    name: str = Field(..., description="收件人姓名")
    phone: str = Field(..., description="收件人电话")
    address: Optional[str] = Field(None, description="收件地址")


class ParcelCreate(BaseModel):
    tracking_number: str = Field(..., description="快递单号")
    courier_company: Optional[str] = Field(None, description="快递公司")
    recipient: RecipientCreate
    inbound_time: Optional[datetime] = Field(None, description="入库时间")
    shelf_location: Optional[str] = Field(None, description="货架位置")
    weight: Optional[str] = Field(None, description="重量")
    remarks: Optional[str] = Field(None, description="备注")
    created_by: str = Field(..., description="创建人")


class ParcelResponse(BaseModel):
    id: int
    tracking_number: str
    courier_company: Optional[str]
    recipient_name: str
    recipient_phone: str
    recipient_address: Optional[str]
    inbound_time: datetime
    status: str
    shelf_location: Optional[str]
    weight: Optional[str]
    remarks: Optional[str]
    created_by: str
    created_at: datetime
    stagnation_level: Optional[str]
    stagnation_hours: Optional[float]
    reminder_count: int

    class Config:
        orm_mode = True


class ParcelDetailResponse(ParcelResponse):
    reminder_records: List[dict]
    rejection: Optional[dict]
    return_report: Optional[dict]
    operation_logs: List[dict]


class ReminderCreate(BaseModel):
    parcel_ids: List[int] = Field(..., description="包裹ID列表")
    reminder_type: str = Field(..., description="催取类型：normal/urgent/final")
    reminder_channel: str = Field("sms", description="催取渠道：sms/phone/wechat")
    reminder_content: Optional[str] = Field(None, description="催取内容")
    sent_by: str = Field(..., description="发送人")


class RejectionCreate(BaseModel):
    parcel_id: int = Field(..., description="包裹ID")
    reason: str = Field(..., description="拒收原因")
    rejected_by: str = Field(..., description="拒收操作人")
    contact_result: Optional[str] = Field(None, description="联系结果")
    follow_up_action: Optional[str] = Field(None, description="后续处理措施")
    remarks: Optional[str] = Field(None, description="备注")


class RejectionConfirm(BaseModel):
    rejection_id: int = Field(..., description="拒收记录ID")
    confirmed_by: str = Field(..., description="确认人")


class ReturnCreate(BaseModel):
    parcel_id: int = Field(..., description="包裹ID")
    return_tracking_number: Optional[str] = Field(None, description="退回快递单号")
    return_courier_company: Optional[str] = Field(None, description="退回快递公司")
    return_reason: str = Field(..., description="退回原因")
    return_address: str = Field(..., description="退回地址")
    return_contact: Optional[str] = Field(None, description="退回联系人")
    return_phone: Optional[str] = Field(None, description="退回联系电话")
    reported_by: str = Field(..., description="报告人")
    remarks: Optional[str] = Field(None, description="备注")


class ReturnConfirm(BaseModel):
    return_id: int = Field(..., description="退回记录ID")
    confirmed_by: str = Field(..., description="确认人")
    actual_shipped_at: Optional[datetime] = Field(None, description="实际寄出时间")


class StatusUpdate(BaseModel):
    parcel_id: int = Field(..., description="包裹ID")
    new_status: str = Field(..., description="新状态")
    operator: str = Field(..., description="操作人")
    remarks: Optional[str] = Field(None, description="备注")


class ManualCorrection(BaseModel):
    parcel_id: int = Field(..., description="包裹ID")
    field_name: str = Field(..., description="要修改的字段名")
    field_value: str = Field(..., description="修改后的值")
    operator: str = Field(..., description="操作人")
    reason: str = Field(..., description="修改原因")


class CloseParcel(BaseModel):
    parcel_id: int = Field(..., description="包裹ID")
    close_reason: str = Field(..., description="关闭原因")
    closed_by: str = Field(..., description="关闭人")


def get_stagnation_level(inbound_time: datetime) -> tuple:
    if not inbound_time:
        return None, 0
    now = datetime.utcnow()
    hours_diff = (now - inbound_time).total_seconds() / 3600
    if hours_diff >= 168:
        return "critical", hours_diff
    elif hours_diff >= 72:
        return "urgent", hours_diff
    elif hours_diff >= 24:
        return "warning", hours_diff
    elif hours_diff >= 12:
        return "normal", hours_diff
    return None, hours_diff


def json_serial(obj):
    if isinstance(obj, datetime):
        return obj.isoformat()
    raise TypeError(f"Type {type(obj)} not serializable")

def log_operation(db: Session, parcel_id: int, operation_type: str,
                  original_input: dict, operator: str, conclusion: str, remarks: str = None):
    log = OperationLog(
        parcel_id=parcel_id,
        operation_type=operation_type,
        original_input=json.dumps(original_input, ensure_ascii=False, default=json_serial),
        operator=operator,
        conclusion=conclusion,
        remarks=remarks
    )
    db.add(log)


@app.post("/api/parcels", response_model=ParcelResponse, summary="创建包裹入库")
def create_parcel(parcel_data: ParcelCreate, db: Session = Depends(get_db)):
    existing = db.query(Parcel).filter(
        Parcel.tracking_number == parcel_data.tracking_number,
        Parcel.is_deleted == False
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="该快递单号已存在")

    recipient = db.query(Recipient).filter(
        Recipient.phone == parcel_data.recipient.phone
    ).first()
    if not recipient:
        recipient = Recipient(
            name=parcel_data.recipient.name,
            phone=parcel_data.recipient.phone,
            address=parcel_data.recipient.address
        )
        db.add(recipient)
        db.flush()
    else:
        recipient.name = parcel_data.recipient.name
        if parcel_data.recipient.address:
            recipient.address = parcel_data.recipient.address

    inbound_time = parcel_data.inbound_time or datetime.utcnow()
    parcel = Parcel(
        tracking_number=parcel_data.tracking_number,
        courier_company=parcel_data.courier_company,
        recipient_id=recipient.id,
        inbound_time=inbound_time,
        shelf_location=parcel_data.shelf_location,
        weight=parcel_data.weight,
        remarks=parcel_data.remarks,
        created_by=parcel_data.created_by
    )
    db.add(parcel)
    db.flush()

    log_operation(db, parcel.id, "create_parcel",
                  parcel_data.dict(), parcel_data.created_by,
                  "包裹入库成功")
    db.commit()
    db.refresh(parcel)

    level, hours = get_stagnation_level(parcel.inbound_time)
    return ParcelResponse(
        id=parcel.id,
        tracking_number=parcel.tracking_number,
        courier_company=parcel.courier_company,
        recipient_name=recipient.name,
        recipient_phone=recipient.phone,
        recipient_address=recipient.address,
        inbound_time=parcel.inbound_time,
        status=parcel.status,
        shelf_location=parcel.shelf_location,
        weight=parcel.weight,
        remarks=parcel.remarks,
        created_by=parcel.created_by,
        created_at=parcel.created_at,
        stagnation_level=level,
        stagnation_hours=round(hours, 1),
        reminder_count=0
    )


@app.get("/api/parcels", response_model=List[ParcelResponse], summary="查询包裹列表")
def list_parcels(
    status: Optional[str] = Query(None, description="状态过滤"),
    stagnation_level: Optional[str] = Query(None, description="滞留级别过滤"),
    tracking_number: Optional[str] = Query(None, description="快递单号搜索"),
    recipient_phone: Optional[str] = Query(None, description="收件人电话搜索"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db)
):
    query = db.query(Parcel).filter(Parcel.is_deleted == False)

    if status:
        query = query.filter(Parcel.status == status)
    if tracking_number:
        query = query.filter(Parcel.tracking_number.contains(tracking_number))
    if recipient_phone:
        query = query.join(Recipient).filter(Recipient.phone.contains(recipient_phone))

    parcels = query.order_by(Parcel.inbound_time.desc()).offset(
        (page - 1) * page_size
    ).limit(page_size).all()

    result = []
    for parcel in parcels:
        level, hours = get_stagnation_level(parcel.inbound_time)
        if stagnation_level and level != stagnation_level:
            continue
        reminder_count = db.query(ReminderRecord).filter(
            ReminderRecord.parcel_id == parcel.id
        ).count()
        result.append(ParcelResponse(
            id=parcel.id,
            tracking_number=parcel.tracking_number,
            courier_company=parcel.courier_company,
            recipient_name=parcel.recipient.name,
            recipient_phone=parcel.recipient.phone,
            recipient_address=parcel.recipient.address,
            inbound_time=parcel.inbound_time,
            status=parcel.status,
            shelf_location=parcel.shelf_location,
            weight=parcel.weight,
            remarks=parcel.remarks,
            created_by=parcel.created_by,
            created_at=parcel.created_at,
            stagnation_level=level,
            stagnation_hours=round(hours, 1),
            reminder_count=reminder_count
        ))
    return result


@app.get("/api/parcels/{parcel_id}", response_model=ParcelDetailResponse, summary="查询包裹详情")
def get_parcel_detail(parcel_id: int, db: Session = Depends(get_db)):
    parcel = db.query(Parcel).filter(
        Parcel.id == parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    level, hours = get_stagnation_level(parcel.inbound_time)
    reminder_count = len(parcel.reminder_records)

    return ParcelDetailResponse(
        id=parcel.id,
        tracking_number=parcel.tracking_number,
        courier_company=parcel.courier_company,
        recipient_name=parcel.recipient.name,
        recipient_phone=parcel.recipient.phone,
        recipient_address=parcel.recipient.address,
        inbound_time=parcel.inbound_time,
        status=parcel.status,
        shelf_location=parcel.shelf_location,
        weight=parcel.weight,
        remarks=parcel.remarks,
        created_by=parcel.created_by,
        created_at=parcel.created_at,
        stagnation_level=level,
        stagnation_hours=round(hours, 1),
        reminder_count=reminder_count,
        reminder_records=[{
            "id": r.id, "type": r.reminder_type, "channel": r.reminder_channel,
            "content": r.reminder_content, "sent_at": r.sent_at, "sent_by": r.sent_by
        } for r in parcel.reminder_records],
        rejection={
            "id": parcel.rejection.id, "reason": parcel.rejection.reason,
            "rejected_at": parcel.rejection.rejected_at, "rejected_by": parcel.rejection.rejected_by,
            "is_confirmed": parcel.rejection.is_confirmed
        } if parcel.rejection else None,
        return_report={
            "id": parcel.return_report.id, "return_tracking_number": parcel.return_report.return_tracking_number,
            "return_reason": parcel.return_report.return_reason, "reported_at": parcel.return_report.reported_at,
            "reported_by": parcel.return_report.reported_by, "is_confirmed": parcel.return_report.is_confirmed
        } if parcel.return_report else None,
        operation_logs=[{
            "id": l.id, "type": l.operation_type, "operator": l.operator,
            "conclusion": l.conclusion, "operation_time": l.operation_time
        } for l in parcel.operation_logs]
    )


@app.post("/api/reminders", summary="批量催取（去重）")
def create_reminders(reminder_data: ReminderCreate, db: Session = Depends(get_db)):
    today = datetime.utcnow().date()
    results = []

    for parcel_id in reminder_data.parcel_ids:
        parcel = db.query(Parcel).filter(
            Parcel.id == parcel_id,
            Parcel.is_deleted == False
        ).first()
        if not parcel:
            results.append({"parcel_id": parcel_id, "success": False, "reason": "包裹不存在"})
            continue

        dedup_key = f"{parcel_id}:{reminder_data.reminder_type}:{today}"
        existing = db.query(ReminderRecord).filter(
            ReminderRecord.deduplication_key == dedup_key
        ).first()
        if existing:
            results.append({"parcel_id": parcel_id, "success": False, "reason": "今日已催取过"})
            continue

        content = reminder_data.reminder_content or f"【驿站提醒】您的包裹{parcel.tracking_number}已到店，请及时取件！"
        record = ReminderRecord(
            parcel_id=parcel_id,
            reminder_type=reminder_data.reminder_type,
            reminder_channel=reminder_data.reminder_channel,
            reminder_content=content,
            sent_by=reminder_data.sent_by,
            deduplication_key=dedup_key
        )
        db.add(record)
        parcel.status = "reminded"
        results.append({"parcel_id": parcel_id, "success": True, "tracking_number": parcel.tracking_number})

    log_operation(db, None, "batch_reminder",
                  reminder_data.dict(), reminder_data.sent_by,
                  f"批量催取完成，共{len([r for r in results if r['success']])}个成功")
    db.commit()
    return {"results": results, "total_success": len([r for r in results if r["success"]])}


@app.post("/api/rejections", summary="创建拒收记录")
def create_rejection(rejection_data: RejectionCreate, db: Session = Depends(get_db)):
    parcel = db.query(Parcel).filter(
        Parcel.id == rejection_data.parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    existing = db.query(Rejection).filter(Rejection.parcel_id == rejection_data.parcel_id).first()
    if existing:
        raise HTTPException(status_code=400, detail="该包裹已有拒收记录")

    rejection = Rejection(
        parcel_id=rejection_data.parcel_id,
        reason=rejection_data.reason,
        rejected_by=rejection_data.rejected_by,
        contact_result=rejection_data.contact_result,
        follow_up_action=rejection_data.follow_up_action,
        remarks=rejection_data.remarks
    )
    db.add(rejection)
    parcel.status = "rejected"

    log_operation(db, rejection_data.parcel_id, "create_rejection",
                  rejection_data.dict(), rejection_data.rejected_by,
                  "拒收记录创建成功")
    db.commit()
    db.refresh(rejection)
    return rejection


@app.post("/api/rejections/confirm", summary="确认拒收")
def confirm_rejection(confirm_data: RejectionConfirm, db: Session = Depends(get_db)):
    rejection = db.query(Rejection).filter(Rejection.id == confirm_data.rejection_id).first()
    if not rejection:
        raise HTTPException(status_code=404, detail="拒收记录不存在")
    if rejection.is_confirmed:
        raise HTTPException(status_code=400, detail="该拒收已确认")

    rejection.is_confirmed = True
    rejection.confirmed_by = confirm_data.confirmed_by
    rejection.confirmed_at = datetime.utcnow()

    log_operation(db, rejection.parcel_id, "confirm_rejection",
                  confirm_data.dict(), confirm_data.confirmed_by,
                  "拒收已确认")
    db.commit()
    return {"message": "拒收已确认", "rejection_id": rejection.id}


@app.post("/api/returns", summary="创建退回报告")
def create_return(return_data: ReturnCreate, db: Session = Depends(get_db)):
    parcel = db.query(Parcel).filter(
        Parcel.id == return_data.parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    existing = db.query(ReturnReport).filter(ReturnReport.parcel_id == return_data.parcel_id).first()
    if existing:
        raise HTTPException(status_code=400, detail="该包裹已有退回记录")

    return_report = ReturnReport(
        parcel_id=return_data.parcel_id,
        return_tracking_number=return_data.return_tracking_number,
        return_courier_company=return_data.return_courier_company,
        return_reason=return_data.return_reason,
        return_address=return_data.return_address,
        return_contact=return_data.return_contact,
        return_phone=return_data.return_phone,
        reported_by=return_data.reported_by,
        remarks=return_data.remarks
    )
    db.add(return_report)
    parcel.status = "returning"

    log_operation(db, return_data.parcel_id, "create_return",
                  return_data.dict(), return_data.reported_by,
                  "退回报告创建成功")
    db.commit()
    db.refresh(return_report)
    return return_report


@app.post("/api/returns/confirm", summary="确认退回")
def confirm_return(confirm_data: ReturnConfirm, db: Session = Depends(get_db)):
    return_report = db.query(ReturnReport).filter(ReturnReport.id == confirm_data.return_id).first()
    if not return_report:
        raise HTTPException(status_code=404, detail="退回记录不存在")
    if return_report.is_confirmed:
        raise HTTPException(status_code=400, detail="该退回已确认")

    return_report.is_confirmed = True
    return_report.confirmed_by = confirm_data.confirmed_by
    return_report.confirmed_at = datetime.utcnow()
    if confirm_data.actual_shipped_at:
        return_report.actual_shipped_at = confirm_data.actual_shipped_at

    parcel = db.query(Parcel).filter(Parcel.id == return_report.parcel_id).first()
    if parcel:
        parcel.status = "returned"

    log_operation(db, return_report.parcel_id, "confirm_return",
                  confirm_data.dict(), confirm_data.confirmed_by,
                  "退回已确认")
    db.commit()
    return {"message": "退回已确认", "return_id": return_report.id}


@app.put("/api/parcels/status", summary="状态推进")
def update_parcel_status(status_data: StatusUpdate, db: Session = Depends(get_db)):
    valid_statuses = ["pending", "reminded", "picked_up", "rejected", "returning", "returned", "closed"]
    if status_data.new_status not in valid_statuses:
        raise HTTPException(status_code=400, detail=f"无效状态，有效状态为：{valid_statuses}")

    parcel = db.query(Parcel).filter(
        Parcel.id == status_data.parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    old_status = parcel.status
    parcel.status = status_data.new_status

    log_operation(db, status_data.parcel_id, "status_update",
                  {"old_status": old_status, "new_status": status_data.new_status},
                  status_data.operator,
                  f"状态从{old_status}变更为{status_data.new_status}",
                  status_data.remarks)
    db.commit()
    return {"message": "状态更新成功", "parcel_id": parcel.id, "old_status": old_status, "new_status": status_data.new_status}


@app.put("/api/parcels/correct", summary="人工修正")
def manual_correction(correction_data: ManualCorrection, db: Session = Depends(get_db)):
    parcel = db.query(Parcel).filter(
        Parcel.id == correction_data.parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    valid_fields = ["tracking_number", "courier_company", "shelf_location", "weight", "remarks", "status"]
    if correction_data.field_name not in valid_fields:
        raise HTTPException(status_code=400, detail=f"不允许修改该字段，可修改字段：{valid_fields}")

    old_value = getattr(parcel, correction_data.field_name)
    setattr(parcel, correction_data.field_name, correction_data.field_value)

    log_operation(db, correction_data.parcel_id, "manual_correction",
                  {
                      "field": correction_data.field_name,
                      "old_value": str(old_value),
                      "new_value": correction_data.field_value
                  },
                  correction_data.operator,
                  "人工修正完成",
                  correction_data.reason)
    db.commit()
    return {
        "message": "修正成功",
        "parcel_id": parcel.id,
        "field": correction_data.field_name,
        "old_value": str(old_value),
        "new_value": correction_data.field_value
    }


@app.post("/api/parcels/close", summary="关闭/撤回包裹")
def close_parcel(close_data: CloseParcel, db: Session = Depends(get_db)):
    parcel = db.query(Parcel).filter(
        Parcel.id == close_data.parcel_id,
        Parcel.is_deleted == False
    ).first()
    if not parcel:
        raise HTTPException(status_code=404, detail="包裹不存在")

    parcel.is_deleted = True
    parcel.status = "closed"
    parcel.closed_at = datetime.utcnow()
    parcel.closed_by = close_data.closed_by
    parcel.close_reason = close_data.close_reason

    log_operation(db, close_data.parcel_id, "close_parcel",
                  close_data.dict(), close_data.closed_by,
                  "包裹已关闭/撤回")
    db.commit()
    return {"message": "包裹已关闭", "parcel_id": parcel.id}


@app.get("/api/export", summary="导出报告")
def export_report(
    status: Optional[str] = Query(None),
    start_date: Optional[datetime] = Query(None),
    end_date: Optional[datetime] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(Parcel).filter(Parcel.is_deleted == False)
    if status:
        query = query.filter(Parcel.status == status)
    if start_date:
        query = query.filter(Parcel.inbound_time >= start_date)
    if end_date:
        query = query.filter(Parcel.inbound_time <= end_date)

    parcels = query.order_by(Parcel.inbound_time.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "包裹管理报告"

    headers = ["快递单号", "快递公司", "收件人", "电话", "入库时间", "状态", "滞留级别",
               "催取次数", "拒收原因", "退回单号", "货架位置", "备注"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row, parcel in enumerate(parcels, 2):
        level, _ = get_stagnation_level(parcel.inbound_time)
        reminder_count = db.query(ReminderRecord).filter(
            ReminderRecord.parcel_id == parcel.id
        ).count()

        ws.cell(row=row, column=1, value=parcel.tracking_number)
        ws.cell(row=row, column=2, value=parcel.courier_company or "")
        ws.cell(row=row, column=3, value=parcel.recipient.name)
        ws.cell(row=row, column=4, value=parcel.recipient.phone)
        ws.cell(row=row, column=5, value=parcel.inbound_time.strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=row, column=6, value=parcel.status)
        ws.cell(row=row, column=7, value=level or "")
        ws.cell(row=row, column=8, value=reminder_count)
        ws.cell(row=row, column=9, value=parcel.rejection.reason if parcel.rejection else "")
        ws.cell(row=row, column=10, value=parcel.return_report.return_tracking_number if parcel.return_report else "")
        ws.cell(row=row, column=11, value=parcel.shelf_location or "")
        ws.cell(row=row, column=12, value=parcel.remarks or "")

    for col in range(1, 13):
        ws.column_dimensions[chr(64 + col)].width = 18

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"parcel_report_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@app.get("/api/statistics", summary="统计概览")
def get_statistics(db: Session = Depends(get_db)):
    total = db.query(Parcel).filter(Parcel.is_deleted == False).count()
    status_stats = db.query(Parcel.status, func.count(Parcel.id)).filter(
        Parcel.is_deleted == False
    ).group_by(Parcel.status).all()

    now = datetime.utcnow()
    stagnation_stats = {
        "超过7天": db.query(Parcel).filter(Parcel.inbound_time <= now - timedelta(days=7)).count(),
        "3-7天": db.query(Parcel).filter(
            and_(Parcel.inbound_time > now - timedelta(days=7),
                 Parcel.inbound_time <= now - timedelta(days=3))
        ).count(),
        "1-3天": db.query(Parcel).filter(
            and_(Parcel.inbound_time > now - timedelta(days=3),
                 Parcel.inbound_time <= now - timedelta(days=1))
        ).count(),
    }

    reminder_today = db.query(ReminderRecord).filter(
        func.date(ReminderRecord.sent_at) == now.date()
    ).count()

    return {
        "total_parcels": total,
        "status_distribution": dict(status_stats),
        "stagnation_distribution": stagnation_stats,
        "reminders_today": reminder_today
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
