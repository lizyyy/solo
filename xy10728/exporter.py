import pandas as pd
from datetime import datetime
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ReportExporter:
    def __init__(self):
        self.header_font = Font(bold=True, color='FFFFFF', size=12)
        self.header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        self.error_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        self.warning_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        self.success_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        self.center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def generate_report(self, batch: Dict[str, Any], output_path: str):
        wb = openpyxl.Workbook()
        
        ws_summary = wb.active
        ws_summary.title = '批次概览'
        self._create_summary_sheet(ws_summary, batch)
        
        ws_clean = wb.create_sheet('成功导入数据')
        self._create_clean_data_sheet(ws_clean, batch)
        
        ws_dirty = wb.create_sheet('脏数据详情')
        self._create_dirty_data_sheet(ws_dirty, batch)
        
        ws_fix = wb.create_sheet('修复建议')
        self._create_fix_suggestion_sheet(ws_fix, batch)
        
        for ws in wb.worksheets:
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width
        
        wb.save(output_path)
    
    def _create_summary_sheet(self, ws, batch: Dict[str, Any]):
        validation = batch['validation_result']
        
        summary_data = [
            ['导入批次报告', ''],
            ['', ''],
            ['基本信息', ''],
            ['批次编号', batch['batch_id']],
            ['源文件名', batch['filename']],
            ['负责人', batch.get('owner', '未指定')],
            ['上传时间', self._format_datetime(batch['upload_time'])],
            ['导入状态', self._get_status_text(batch['status'])],
            ['', ''],
            ['数据统计', ''],
            ['总行数', batch['total_rows']],
            ['干净数据行数', validation['clean_rows']],
            ['脏数据行数', validation['dirty_rows']],
            ['导入成功率', f"{(validation['clean_rows']/batch['total_rows']*100):.1f}%"],
            ['', ''],
            ['按负责人分组统计', ''],
        ]
        
        owner_stats = self._get_owner_stats(batch)
        if owner_stats:
            summary_data.append(['负责人', '数据行数', '干净行数', '脏行数'])
            for owner, stats in owner_stats.items():
                summary_data.append([
                    owner,
                    stats['total'],
                    stats['clean'],
                    stats['dirty']
                ])
        
        for row_idx, row_data in enumerate(summary_data, 1):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = self.thin_border
                cell.alignment = self.center_align
                
                if row_idx == 1:
                    cell.font = Font(bold=True, size=16)
                    cell.fill = self.header_fill
                elif value in ['基本信息', '数据统计', '按负责人分组统计']:
                    cell.font = self.header_font
                    cell.fill = self.header_fill
                elif row_idx >= len(summary_data) - len(owner_stats) and owner_stats:
                    if col_idx == 4 and isinstance(value, int) and value > 0:
                        cell.fill = self.error_fill
        
        ws.merge_cells('A1:B1')
    
    def _create_clean_data_sheet(self, ws, batch: Dict[str, Any]):
        if not batch.get('imported_data'):
            ws.cell(row=1, column=1, value='暂无成功导入的数据，请先执行导入操作')
            return
        
        data = batch['imported_data']
        if not data:
            ws.cell(row=1, column=1, value='没有干净的数据可导入')
            return
        
        headers = list(data[0].keys())
        ws.append(['行号'] + headers)
        
        for col in range(1, len(headers) + 2):
            cell = ws.cell(row=1, column=col)
            cell.font = self.header_font
            cell.fill = self.success_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        for idx, row in enumerate(data, 1):
            ws.append([idx] + [row.get(h, '') for h in headers])
            for col in range(1, len(headers) + 2):
                cell = ws.cell(row=idx + 1, column=col)
                cell.border = self.thin_border
                cell.alignment = self.center_align
    
    def _create_dirty_data_sheet(self, ws, batch: Dict[str, Any]):
        validation = batch['validation_result']
        errors = validation['errors']
        
        if not errors:
            ws.cell(row=1, column=1, value='太棒了！没有发现脏数据')
            ws.cell(row=1, column=1).fill = self.success_fill
            return
        
        headers = ['行号', '字段名', '原始值', '错误原因']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = self.header_font
            cell.fill = self.error_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        row_num = 2
        for error_group in errors:
            for error in error_group['errors']:
                ws.append([
                    error_group['row_number'],
                    error['field'],
                    error['value'],
                    error['error']
                ])
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_num, column=col)
                    cell.border = self.thin_border
                    cell.alignment = self.center_align
                    cell.fill = self.error_fill
                row_num += 1
    
    def _create_fix_suggestion_sheet(self, ws, batch: Dict[str, Any]):
        validation = batch['validation_result']
        suggestions = validation.get('fix_suggestions', {})
        
        if not suggestions:
            ws.cell(row=1, column=1, value='暂无修复建议')
            return
        
        headers = ['行号', '字段名', '当前值', '修复建议']
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = self.header_font
            cell.fill = self.warning_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        row_num = 2
        for row_idx, fixes in suggestions.items():
            for field, fix_info in fixes.items():
                ws.append([
                    row_idx + 2,
                    field,
                    fix_info['original_value'],
                    fix_info['suggestion']
                ])
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_num, column=col)
                    cell.border = self.thin_border
                    cell.alignment = self.center_align
                    cell.fill = self.warning_fill
                row_num += 1
    
    def _format_datetime(self, iso_str: str) -> str:
        try:
            dt = datetime.fromisoformat(iso_str)
            return dt.strftime('%Y年%m月%d日 %H:%M:%S')
        except:
            return iso_str
    
    def _get_status_text(self, status: str) -> str:
        status_map = {
            'uploaded': '已上传，未校验',
            'validated': '已校验，可导入',
            'imported': '已完成导入'
        }
        return status_map.get(status, status)
    
    def _get_owner_stats(self, batch: Dict[str, Any]) -> Dict[str, Dict[str, int]]:
        data = batch['raw_data']
        validation = batch['validation_result']
        dirty_indices = set(validation['dirty_row_indices'])
        
        owner_field = None
        if data:
            for key in data[0].keys():
                if '负责人' in key or 'owner' in key.lower():
                    owner_field = key
                    break
        
        if not owner_field:
            return {}
        
        stats = {}
        for idx, row in enumerate(data):
            owner = str(row.get(owner_field, '未指定'))
            if owner not in stats:
                stats[owner] = {'total': 0, 'clean': 0, 'dirty': 0}
            
            stats[owner]['total'] += 1
            if idx in dirty_indices:
                stats[owner]['dirty'] += 1
            else:
                stats[owner]['clean'] += 1
        
        return stats
