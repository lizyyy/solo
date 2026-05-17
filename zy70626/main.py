from datetime import datetime
from enum import Enum
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, HTTPException, Depends, status, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import create_engine, Column, Integer, String, DateTime, ForeignKey, Boolean, Text, Float
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session, relationship
from io import BytesIO
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SQLALCHEMY_DATABASE_URL = "sqlite:///./cleaning.db"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


class TaskStatus(str, Enum):
    CREATED = "created"
    ASSIGNED = "assigned"
    IN_PROGRESS = "in_progress"
    SUBMITTED = "submitted"
    INSPECTING = "inspecting"
    NEEDS_REWORK = "needs_rework"
    REWORKING = "reworking"
    REWORK_SUBMITTED = "rework_submitted"
    APPROVED = "approved"
    COMPLAINT_OPEN = "complaint_open"
    COMPLAINT_UNDER_REVIEW = "complaint_under_review"
    COMPLAINT_RESOLVED = "complaint_resolved"
    CLOSED = "closed"


class ErrorCode(str, Enum):
    MISSING_FIELD = "missing_field"
    INVALID_STATUS = "invalid_status"
    NEEDS_MANUAL_REVIEW = "needs_manual_review"
    ALREADY_PROCESSED = "already_processed"
    NOT_FOUND = "not_found"
    VALIDATION_ERROR = "validation_error"


class ComplaintResolution(str, Enum):
    ACCEPTED = "accepted"
    REJECTED = "rejected"
    NEEDS_FURTHER_INVESTIGATION = "needs_further_investigation"


class Property(Base):
    __tablename__ = "properties"
    id = Column(Integer, primary_key=True, index=True)
    name = Column(String(100), nullable=False)
    address = Column(String(255), nullable=False)
    room_count = Column(Integer, default=1)
    owner_id = Column(Integer, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow)
    is_active = Column(Boolean, default=True)
    cleaning_tasks = relationship("CleaningTask", back_populates="property")


class ChecklistItem(Base):
    __tablename__ = "checklist_items"
    id = Column(Integer, primary_key=True, index=True)
    property_id = Column(Integer, ForeignKey("properties.id"))
    category = Column(String(50), nullable=False)
    name = Column(String(100), nullable=False)
    description = Column(Text)
    is_mandatory = Column(Boolean, default=True)
    sort_order = Column(Integer, default=0)
    is_active = Column(Boolean, default=True)


class CleaningTask(Base):
    __tablename__ = "cleaning_tasks"
    id = Column(Integer, primary_key=True, index=True)
    property_id = Column(Integer, ForeignKey("properties.id"))
    cleaner_id = Column(Integer, nullable=False)
    inspector_id = Column(Integer)
    scheduled_date = Column(DateTime, nullable=False)
    status = Column(String(20), default=TaskStatus.CREATED)
    rework_count = Column(Integer, default=0)
    max_reworks = Column(Integer, default=2)
    notes = Column(Text)
    created_at = Column(DateTime, default=datetime.utcnow)
    updated_at = Column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    completed_at = Column(DateTime)
    property = relationship("Property", back_populates="cleaning_tasks")
    check_results = relationship("CheckResult", back_populates="task", cascade="all, delete-orphan")
    photos = relationship("PhotoEvidence", back_populates="task", cascade="all, delete-orphan")
    complaints = relationship("Complaint", back_populates="task", cascade="all, delete-orphan")
    reports = relationship("InspectionReport", back_populates="task", cascade="all, delete-orphan")


class CheckResult(Base):
    __tablename__ = "check_results"
    id = Column(Integer, primary_key=True, index=True)
    task_id = Column(Integer, ForeignKey("cleaning_tasks.id"))
    checklist_item_id = Column(Integer, ForeignKey("checklist_items.id"))
    checked = Column(Boolean, default=False)
    passed = Column(Boolean)
    notes = Column(Text)
    checked_at = Column(DateTime)
    checked_by = Column(Integer)
    task = relationship("CleaningTask", back_populates="check_results")


class PhotoEvidence(Base):
    __tablename__ = "photo_evidence"
    id = Column(Integer, primary_key=True, index=True)
    task_id = Column(Integer, ForeignKey("cleaning_tasks.id"))
    checklist_item_id = Column(Integer, ForeignKey("checklist_items.id"))
    photo_url = Column(String(255), nullable=False)
    thumbnail_url = Column(String(255))
    category = Column(String(50))
    uploaded_by = Column(Integer, nullable=False)
    uploaded_at = Column(DateTime, default=datetime.utcnow)
    description = Column(Text)
    is_complaint_photo = Column(Boolean, default=False)
    task = relationship("CleaningTask", back_populates="photos")


class Complaint(Base):
    __tablename__ = "complaints"
    id = Column(Integer, primary_key=True, index=True)
    task_id = Column(Integer, ForeignKey("cleaning_tasks.id"))
    guest_id = Column(Integer, nullable=False)
    category = Column(String(50), nullable=False)
    description = Column(Text, nullable=False)
    severity = Column(String(20), default="medium")
    status = Column(String(20), default=TaskStatus.COMPLAINT_OPEN)
    reported_at = Column(DateTime, default=datetime.utcnow)
    resolution = Column(String(50))
    resolution_notes = Column(Text)
    resolved_at = Column(DateTime)
    resolved_by = Column(Integer)
    task = relationship("CleaningTask", back_populates="complaints")


class InspectionReport(Base):
    __tablename__ = "inspection_reports"
    id = Column(Integer, primary_key=True, index=True)
    task_id = Column(Integer, ForeignKey("cleaning_tasks.id"))
    inspector_id = Column(Integer, nullable=False)
    overall_score = Column(Float)
    passed_count = Column(Integer, default=0)
    failed_count = Column(Integer, default=0)
    total_count = Column(Integer, default=0)
    notes = Column(Text)
    generated_at = Column(DateTime, default=datetime.utcnow)
    export_format = Column(String(20))
    task = relationship("CleaningTask", back_populates="reports")


Base.metadata.create_all(bind=engine)


class PropertyCreate(BaseModel):
    name: str
    address: str
    room_count: int = 1
    owner_id: int


class PropertyResponse(BaseModel):
    id: int
    name: str
    address: str
    room_count: int
    owner_id: int
    is_active: bool

    class Config:
        from_attributes = True


class ChecklistItemCreate(BaseModel):
    property_id: int
    category: str
    name: str
    description: Optional[str] = None
    is_mandatory: bool = True
    sort_order: int = 0


class ChecklistItemResponse(BaseModel):
    id: int
    property_id: int
    category: str
    name: str
    description: Optional[str]
    is_mandatory: bool
    sort_order: int

    class Config:
        from_attributes = True


class CleaningTaskCreate(BaseModel):
    property_id: int
    cleaner_id: int
    scheduled_date: datetime
    inspector_id: Optional[int] = None
    max_reworks: int = 2


class CleaningTaskResponse(BaseModel):
    id: int
    property_id: int
    cleaner_id: int
    inspector_id: Optional[int]
    scheduled_date: datetime
    status: str
    rework_count: int
    max_reworks: int
    notes: Optional[str]
    created_at: datetime
    completed_at: Optional[datetime]

    class Config:
        from_attributes = True


class CheckResultSubmit(BaseModel):
    checklist_item_id: int
    checked: bool
    passed: Optional[bool] = None
    notes: Optional[str] = None


class CleaningSubmitRequest(BaseModel):
    task_id: int
    submitted_by: int
    check_results: List[CheckResultSubmit]
    notes: Optional[str] = None


class PhotoUploadRequest(BaseModel):
    task_id: int
    checklist_item_id: Optional[int] = None
    photo_url: str
    thumbnail_url: Optional[str] = None
    category: Optional[str] = None
    uploaded_by: int
    description: Optional[str] = None


class InspectionRequest(BaseModel):
    task_id: int
    inspector_id: int
    check_results: List[CheckResultSubmit]
    overall_score: Optional[float] = None
    notes: Optional[str] = None


class ReworkRequest(BaseModel):
    task_id: int
    assigned_to: int
    reason: str
    checklist_item_ids: List[int]


class ComplaintCreate(BaseModel):
    task_id: int
    guest_id: int
    category: str
    description: str
    severity: str = "medium"


class ComplaintResolveRequest(BaseModel):
    complaint_id: int
    resolved_by: int
    resolution: ComplaintResolution
    resolution_notes: Optional[str] = None


class ErrorResponse(BaseModel):
    error_code: ErrorCode
    message: str
    details: Optional[Dict[str, Any]] = None


app = FastAPI(title="保洁验收返工客诉证据API", version="1.0.0")


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def create_http_exception(error_code: ErrorCode, message: str, details: Optional[Dict] = None, status_code: int = 400):
    return HTTPException(
        status_code=status_code,
        detail={
            "error_code": error_code.value,
            "message": message,
            "details": details or {}
        }
    )


VALID_TRANSITIONS = {
    TaskStatus.CREATED: [TaskStatus.ASSIGNED, TaskStatus.IN_PROGRESS],
    TaskStatus.ASSIGNED: [TaskStatus.IN_PROGRESS],
    TaskStatus.IN_PROGRESS: [TaskStatus.SUBMITTED],
    TaskStatus.SUBMITTED: [TaskStatus.INSPECTING, TaskStatus.APPROVED],
    TaskStatus.INSPECTING: [TaskStatus.APPROVED, TaskStatus.NEEDS_REWORK],
    TaskStatus.NEEDS_REWORK: [TaskStatus.REWORKING],
    TaskStatus.REWORKING: [TaskStatus.REWORK_SUBMITTED],
    TaskStatus.REWORK_SUBMITTED: [TaskStatus.INSPECTING, TaskStatus.APPROVED],
    TaskStatus.APPROVED: [TaskStatus.COMPLAINT_OPEN, TaskStatus.CLOSED],
    TaskStatus.COMPLAINT_OPEN: [TaskStatus.COMPLAINT_UNDER_REVIEW],
    TaskStatus.COMPLAINT_UNDER_REVIEW: [TaskStatus.COMPLAINT_RESOLVED, TaskStatus.NEEDS_REWORK],
    TaskStatus.COMPLAINT_RESOLVED: [TaskStatus.CLOSED],
    TaskStatus.CLOSED: [],
}


def validate_status_transition(current_status: str, new_status: TaskStatus) -> bool:
    allowed_statuses = VALID_TRANSITIONS.get(current_status, [])
    return new_status in allowed_statuses


@app.post("/properties/", response_model=PropertyResponse, status_code=201)
def create_property(property_data: PropertyCreate, db: Session = Depends(get_db)):
    try:
        db_property = Property(**property_data.model_dump())
        db.add(db_property)
        db.commit()
        db.refresh(db_property)
        return db_property
    except ValidationError as e:
        raise create_http_exception(ErrorCode.MISSING_FIELD, "缺少必要字段", {"errors": e.errors()})


@app.get("/properties/", response_model=List[PropertyResponse])
def list_properties(skip: int = 0, limit: int = 100, owner_id: Optional[int] = None, db: Session = Depends(get_db)):
    query = db.query(Property).filter(Property.is_active == True)
    if owner_id:
        query = query.filter(Property.owner_id == owner_id)
    return query.offset(skip).limit(limit).all()


@app.post("/checklist-items/", response_model=ChecklistItemResponse, status_code=201)
def create_checklist_item(item_data: ChecklistItemCreate, db: Session = Depends(get_db)):
    property = db.query(Property).filter(Property.id == item_data.property_id).first()
    if not property:
        raise create_http_exception(ErrorCode.NOT_FOUND, "房源不存在", {"property_id": item_data.property_id}, 404)
    db_item = ChecklistItem(**item_data.model_dump())
    db.add(db_item)
    db.commit()
    db.refresh(db_item)
    return db_item


@app.get("/checklist-items/{property_id}", response_model=List[ChecklistItemResponse])
def get_checklist_for_property(property_id: int, db: Session = Depends(get_db)):
    return db.query(ChecklistItem).filter(
        ChecklistItem.property_id == property_id,
        ChecklistItem.is_active == True
    ).order_by(ChecklistItem.sort_order).all()


@app.post("/tasks/", response_model=CleaningTaskResponse, status_code=201)
def create_task(task_data: CleaningTaskCreate, db: Session = Depends(get_db)):
    property = db.query(Property).filter(Property.id == task_data.property_id).first()
    if not property:
        raise create_http_exception(ErrorCode.NOT_FOUND, "房源不存在", {"property_id": task_data.property_id}, 404)
    db_task = CleaningTask(**task_data.model_dump())
    db.add(db_task)
    db.commit()
    db.refresh(db_task)
    return db_task


@app.get("/tasks/", response_model=List[CleaningTaskResponse])
def list_tasks(
    skip: int = 0,
    limit: int = 100,
    status: Optional[TaskStatus] = None,
    property_id: Optional[int] = None,
    cleaner_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    query = db.query(CleaningTask)
    if status:
        query = query.filter(CleaningTask.status == status)
    if property_id:
        query = query.filter(CleaningTask.property_id == property_id)
    if cleaner_id:
        query = query.filter(CleaningTask.cleaner_id == cleaner_id)
    return query.order_by(CleaningTask.created_at.desc()).offset(skip).limit(limit).all()


@app.get("/tasks/{task_id}", response_model=CleaningTaskResponse)
def get_task(task_id: int, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": task_id}, 404)
    return task


@app.post("/tasks/submit-cleaning")
def submit_cleaning(request: CleaningSubmitRequest, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == request.task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": request.task_id}, 404)
    
    if task.status not in [TaskStatus.IN_PROGRESS, TaskStatus.REWORKING]:
        raise create_http_exception(
            ErrorCode.INVALID_STATUS,
            "当前状态不允许提交保洁",
            {"current_status": task.status, "allowed_statuses": ["in_progress", "reworking"]}
        )
    
    checklist_items = db.query(ChecklistItem).filter(
        ChecklistItem.property_id == task.property_id,
        ChecklistItem.is_active == True
    ).all()
    mandatory_items = [item for item in checklist_items if item.is_mandatory]
    submitted_ids = {result.checklist_item_id for result in request.check_results}
    
    missing_mandatory = [item.id for item in mandatory_items if item.id not in submitted_ids]
    if missing_mandatory:
        raise create_http_exception(
            ErrorCode.MISSING_FIELD,
            "缺少必填检查项",
            {"missing_checklist_item_ids": missing_mandatory}
        )
    
    new_status = TaskStatus.REWORK_SUBMITTED if task.status == TaskStatus.REWORKING else TaskStatus.SUBMITTED
    
    for result in request.check_results:
        existing = db.query(CheckResult).filter(
            CheckResult.task_id == task.id,
            CheckResult.checklist_item_id == result.checklist_item_id
        ).first()
        if existing:
            existing.checked = result.checked
            existing.passed = result.passed
            existing.notes = result.notes
            existing.checked_at = datetime.utcnow()
            existing.checked_by = request.submitted_by
        else:
            db_result = CheckResult(
                task_id=task.id,
                checklist_item_id=result.checklist_item_id,
                checked=result.checked,
                passed=result.passed,
                notes=result.notes,
                checked_at=datetime.utcnow(),
                checked_by=request.submitted_by
            )
            db.add(db_result)
    
    task.status = new_status
    task.notes = request.notes or task.notes
    db.commit()
    
    return {"message": "保洁提交成功", "task_id": task.id, "new_status": new_status}


@app.post("/tasks/inspect")
def inspect_task(request: InspectionRequest, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == request.task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": request.task_id}, 404)
    
    if task.status not in [TaskStatus.SUBMITTED, TaskStatus.REWORK_SUBMITTED]:
        raise create_http_exception(
            ErrorCode.INVALID_STATUS,
            "当前状态不允许验收",
            {"current_status": task.status, "allowed_statuses": ["submitted", "rework_submitted"]}
        )
    
    if task.rework_count >= task.max_reworks:
        raise create_http_exception(
            ErrorCode.NEEDS_MANUAL_REVIEW,
            "已达到最大返工次数，需要人工复核",
            {"rework_count": task.rework_count, "max_reworks": task.max_reworks}
        )
    
    all_passed = all(result.passed for result in request.check_results if result.checked)
    
    for result in request.check_results:
        existing = db.query(CheckResult).filter(
            CheckResult.task_id == task.id,
            CheckResult.checklist_item_id == result.checklist_item_id
        ).first()
        if existing:
            existing.passed = result.passed
            existing.checked = result.checked
            existing.notes = result.notes
            existing.checked_by = request.inspector_id
        else:
            db_result = CheckResult(
                task_id=task.id,
                checklist_item_id=result.checklist_item_id,
                checked=result.checked,
                passed=result.passed,
                notes=result.notes,
                checked_at=datetime.utcnow(),
                checked_by=request.inspector_id
            )
            db.add(db_result)
    
    passed_count = sum(1 for r in request.check_results if r.passed)
    failed_count = sum(1 for r in request.check_results if not r.passed and r.checked)
    total_count = len(request.check_results)
    
    report = InspectionReport(
        task_id=task.id,
        inspector_id=request.inspector_id,
        overall_score=request.overall_score or (passed_count / total_count * 100 if total_count > 0 else 0),
        passed_count=passed_count,
        failed_count=failed_count,
        total_count=total_count,
        notes=request.notes
    )
    db.add(report)
    
    if all_passed:
        task.status = TaskStatus.APPROVED
        task.completed_at = datetime.utcnow()
    else:
        task.status = TaskStatus.NEEDS_REWORK
        task.rework_count += 1
    
    db.commit()
    
    return {
        "message": "验收完成",
        "task_id": task.id,
        "new_status": task.status,
        "rework_count": task.rework_count,
        "report_id": report.id
    }


@app.post("/tasks/start-rework")
def start_rework(request: ReworkRequest, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == request.task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": request.task_id}, 404)
    
    if task.status != TaskStatus.NEEDS_REWORK:
        raise create_http_exception(
            ErrorCode.INVALID_STATUS,
            "当前状态不允许开始返工",
            {"current_status": task.status, "expected_status": "needs_rework"}
        )
    
    task.status = TaskStatus.REWORKING
    task.cleaner_id = request.assigned_to
    task.notes = f"返工原因: {request.reason}\n需要重做的检查项: {request.checklist_item_ids}\n" + (task.notes or "")
    db.commit()
    
    return {"message": "返工已开始", "task_id": task.id, "new_status": task.status}


@app.post("/photos/", status_code=201)
def upload_photo(photo_request: PhotoUploadRequest, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == photo_request.task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": photo_request.task_id}, 404)
    
    photo = PhotoEvidence(**photo_request.model_dump())
    db.add(photo)
    db.commit()
    db.refresh(photo)
    
    return {"message": "照片上传成功", "photo_id": photo.id}


@app.get("/photos/{task_id}")
def get_task_photos(task_id: int, category: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(PhotoEvidence).filter(PhotoEvidence.task_id == task_id)
    if category:
        query = query.filter(PhotoEvidence.category == category)
    return query.order_by(PhotoEvidence.uploaded_at.desc()).all()


@app.post("/complaints/", status_code=201)
def create_complaint(complaint_data: ComplaintCreate, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == complaint_data.task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": complaint_data.task_id}, 404)
    
    if task.status not in [TaskStatus.APPROVED, TaskStatus.COMPLAINT_OPEN, TaskStatus.CLOSED]:
        raise create_http_exception(
            ErrorCode.INVALID_STATUS,
            "只有已验收的任务才能创建客诉",
            {"current_status": task.status}
        )
    
    existing_open = db.query(Complaint).filter(
        Complaint.task_id == complaint_data.task_id,
        Complaint.status.in_([TaskStatus.COMPLAINT_OPEN, TaskStatus.COMPLAINT_UNDER_REVIEW])
    ).first()
    if existing_open:
        raise create_http_exception(
            ErrorCode.ALREADY_PROCESSED,
            "该任务已有未处理的客诉",
            {"existing_complaint_id": existing_open.id}
        )
    
    complaint = Complaint(**complaint_data.model_dump())
    db.add(complaint)
    
    task.status = TaskStatus.COMPLAINT_OPEN
    db.commit()
    db.refresh(complaint)
    
    return {"message": "客诉创建成功", "complaint_id": complaint.id, "task_status": task.status}


@app.get("/complaints/")
def list_complaints(
    skip: int = 0,
    limit: int = 100,
    status: Optional[str] = None,
    severity: Optional[str] = None,
    task_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    query = db.query(Complaint)
    if status:
        query = query.filter(Complaint.status == status)
    if severity:
        query = query.filter(Complaint.severity == severity)
    if task_id:
        query = query.filter(Complaint.task_id == task_id)
    return query.order_by(Complaint.reported_at.desc()).offset(skip).limit(limit).all()


@app.post("/complaints/resolve")
def resolve_complaint(request: ComplaintResolveRequest, db: Session = Depends(get_db)):
    complaint = db.query(Complaint).filter(Complaint.id == request.complaint_id).first()
    if not complaint:
        raise create_http_exception(ErrorCode.NOT_FOUND, "客诉不存在", {"complaint_id": request.complaint_id}, 404)
    
    if complaint.status == TaskStatus.COMPLAINT_RESOLVED:
        raise create_http_exception(
            ErrorCode.ALREADY_PROCESSED,
            "该客诉已处理",
            {"complaint_id": complaint.id}
        )
    
    complaint.resolution = request.resolution.value
    complaint.resolution_notes = request.resolution_notes
    complaint.resolved_at = datetime.utcnow()
    complaint.resolved_by = request.resolved_by
    complaint.status = TaskStatus.COMPLAINT_RESOLVED
    
    task = db.query(CleaningTask).filter(CleaningTask.id == complaint.task_id).first()
    if task:
        task.status = TaskStatus.COMPLAINT_RESOLVED
    
    db.commit()
    
    return {
        "message": "客诉已处理",
        "complaint_id": complaint.id,
        "resolution": complaint.resolution
    }


@app.get("/reports/{task_id}/pdf")
def export_report_pdf(task_id: int, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": task_id}, 404)
    
    report = db.query(InspectionReport).filter(InspectionReport.task_id == task_id).order_by(InspectionReport.generated_at.desc()).first()
    if not report:
        raise create_http_exception(ErrorCode.NOT_FOUND, "该任务暂无验收报告", {"task_id": task_id}, 404)
    
    check_results = db.query(CheckResult).filter(CheckResult.task_id == task_id).all()
    photos = db.query(PhotoEvidence).filter(PhotoEvidence.task_id == task_id).all()
    property = task.property
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []
    
    elements.append(Paragraph("保洁验收报告", styles['Title']))
    elements.append(Spacer(1, 12))
    
    elements.append(Paragraph(f"房源: {property.name}", styles['Heading2']))
    elements.append(Paragraph(f"地址: {property.address}", styles['Normal']))
    elements.append(Paragraph(f"任务ID: {task.id}", styles['Normal']))
    elements.append(Paragraph(f"验收状态: {task.status}", styles['Normal']))
    elements.append(Paragraph(f"验收日期: {report.generated_at.strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Paragraph(f"得分: {report.overall_score:.1f}%", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    data = [["检查项", "状态", "备注"]]
    for result in check_results:
        checklist_item = db.query(ChecklistItem).filter(ChecklistItem.id == result.checklist_item_id).first()
        item_name = checklist_item.name if checklist_item else f"Item {result.checklist_item_id}"
        status = "通过" if result.passed else "不通过" if result.checked else "未检查"
        data.append([item_name, status, result.notes or ""])
    
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(table)
    elements.append(Spacer(1, 12))
    
    elements.append(Paragraph(f"照片凭证数量: {len(photos)}", styles['Normal']))
    for photo in photos:
        elements.append(Paragraph(f"- {photo.category}: {photo.description or '无描述'}", styles['Normal']))
    
    if report.notes:
        elements.append(Spacer(1, 12))
        elements.append(Paragraph("验收备注:", styles['Heading3']))
        elements.append(Paragraph(report.notes, styles['Normal']))
    
    doc.build(elements)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=inspection_report_{task_id}.pdf"}
    )


@app.get("/reports/{task_id}/excel")
def export_report_excel(task_id: int, db: Session = Depends(get_db)):
    task = db.query(CleaningTask).filter(CleaningTask.id == task_id).first()
    if not task:
        raise create_http_exception(ErrorCode.NOT_FOUND, "任务不存在", {"task_id": task_id}, 404)
    
    report = db.query(InspectionReport).filter(InspectionReport.task_id == task_id).order_by(InspectionReport.generated_at.desc()).first()
    if not report:
        raise create_http_exception(ErrorCode.NOT_FOUND, "该任务暂无验收报告", {"task_id": task_id}, 404)
    
    check_results = db.query(CheckResult).filter(CheckResult.task_id == task_id).all()
    property = task.property
    
    buffer = BytesIO()
    wb = Workbook()
    
    ws = wb.active
    ws.title = "验收报告"
    
    headers = ["保洁验收报告", "", "", ""]
    ws.append(headers)
    ws.append(["房源", property.name, "", ""])
    ws.append(["地址", property.address, "", ""])
    ws.append(["任务ID", task.id, "", ""])
    ws.append(["验收状态", task.status, "", ""])
    ws.append(["得分", f"{report.overall_score:.1f}%", "", ""])
    ws.append(["", "", "", ""])
    
    ws.append(["检查项", "分类", "状态", "备注"])
    for result in check_results:
        checklist_item = db.query(ChecklistItem).filter(ChecklistItem.id == result.checklist_item_id).first()
        item_name = checklist_item.name if checklist_item else f"Item {result.checklist_item_id}"
        category = checklist_item.category if checklist_item else ""
        status = "通过" if result.passed else "不通过" if result.checked else "未检查"
        ws.append([item_name, category, status, result.notes or ""])
    
    for row in ws.iter_rows(min_row=8, max_row=8):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    wb.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=inspection_report_{task_id}.xlsx"}
    )


@app.get("/health")
def health_check():
    return {"status": "healthy", "timestamp": datetime.utcnow().isoformat()}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
