import json
import csv
from io import StringIO, BytesIO
from datetime import datetime
from typing import Dict, Any, List, Optional
from collections import Counter
from sqlalchemy.orm import Session
from backend.database import Project, Sample, Annotation, Conflict, ReviewDecision, Version


class ReportGenerator:
    def __init__(self, db: Session):
        self.db = db
    
    def generate_summary(self, project_id: int) -> Dict[str, Any]:
        project = self.db.query(Project).filter(Project.id == project_id).first()
        if not project:
            raise ValueError(f"Project {project_id} not found")
        
        total_samples = self.db.query(Sample).filter(Sample.project_id == project_id).count()
        total_annotations = (
            self.db.query(Annotation)
            .join(Sample)
            .filter(Sample.project_id == project_id)
            .count()
        )
        
        all_conflicts = (
            self.db.query(Conflict)
            .join(Sample)
            .filter(Sample.project_id == project_id)
            .all()
        )
        
        total_conflicts = len(all_conflicts)
        resolved_conflicts = sum(1 for c in all_conflicts if c.status == "resolved")
        pending_conflicts = total_conflicts - resolved_conflicts
        
        conflict_rate = (total_conflicts / total_samples * 100) if total_samples > 0 else 0
        
        annotators = set()
        annotator_accuracy = {}
        
        for sample in project.samples:
            for ann in sample.annotations:
                annotators.add(ann.annotator)
                if ann.annotator not in annotator_accuracy:
                    annotator_accuracy[ann.annotator] = {"total": 0, "accepted": 0}
                annotator_accuracy[ann.annotator]["total"] += 1
                if ann.is_active:
                    annotator_accuracy[ann.annotator]["accepted"] += 1
        
        annotator_stats = []
        for annotator, stats in annotator_accuracy.items():
            accuracy = (stats["accepted"] / stats["total"] * 100) if stats["total"] > 0 else 0
            annotator_stats.append({
                "annotator": annotator,
                "total_annotations": stats["total"],
                "accepted_annotations": stats["accepted"],
                "accuracy_rate": round(accuracy, 2)
            })
        
        severity_dist = Counter(c.severity for c in all_conflicts)
        
        latest_versions = (
            self.db.query(Version)
            .filter(Version.project_id == project_id)
            .order_by(Version.version_number.desc())
            .limit(5)
            .all()
        )
        
        return {
            "project_id": project.id,
            "project_name": project.name,
            "generated_at": datetime.utcnow().isoformat(),
            "summary": {
                "total_samples": total_samples,
                "total_annotations": total_annotations,
                "total_conflicts": total_conflicts,
                "resolved_conflicts": resolved_conflicts,
                "pending_conflicts": pending_conflicts,
                "conflict_rate_percent": round(conflict_rate, 2),
                "resolution_rate": round((resolved_conflicts / total_conflicts * 100) if total_conflicts > 0 else 0, 2)
            },
            "conflict_severity_distribution": dict(severity_dist),
            "annotator_statistics": annotator_stats,
            "total_annotators": len(annotators),
            "recent_versions": [
                {
                    "version_number": v.version_number,
                    "action": v.action,
                    "description": v.description,
                    "created_at": v.created_at.isoformat(),
                    "affected_samples": v.affected_samples
                }
                for v in latest_versions
            ]
        }
    
    def generate_detailed_report(self, project_id: int) -> Dict[str, Any]:
        summary = self.generate_summary(project_id)
        
        conflicts = (
            self.db.query(Conflict)
            .join(Sample)
            .filter(Sample.project_id == project_id)
            .order_by(Conflict.created_at.desc())
            .all()
        )
        
        conflict_details = []
        for conflict in conflicts:
            annotations = [
                {
                    "annotator": a.annotator,
                    "label": a.label,
                    "confidence": a.confidence,
                    "is_active": a.is_active
                }
                for a in conflict.sample.annotations
            ]
            
            latest_decision = None
            for d in conflict.review_decisions:
                if d.is_final:
                    latest_decision = {
                        "reviewer": d.reviewer,
                        "decision": d.decision,
                        "reasoning": d.reasoning,
                        "created_at": d.created_at.isoformat()
                    }
                    break
            
            conflict_details.append({
                "conflict_id": conflict.id,
                "sample_id": conflict.sample_id,
                "sample_content": conflict.sample.content[:200] + ("..." if len(conflict.sample.content) > 200 else ""),
                "status": conflict.status,
                "severity": conflict.severity,
                "detection_method": conflict.detection_method,
                "created_at": conflict.created_at.isoformat(),
                "resolved_at": conflict.resolved_at.isoformat() if conflict.resolved_at else None,
                "resolved_by": conflict.resolved_by,
                "annotations": annotations,
                "latest_decision": latest_decision,
                "decision_count": len(conflict.review_decisions)
            })
        
        summary["conflict_details"] = conflict_details
        return summary
    
    def export_to_json(self, data: Dict[str, Any]) -> bytes:
        return json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8')
    
    def export_to_csv(self, project_id: int, data_type: str = "conflicts") -> bytes:
        output = StringIO()
        writer = csv.writer(output)
        
        if data_type == "conflicts":
            writer.writerow([
                "Conflict ID", "Sample ID", "Status", "Severity", "Detection Method",
                "Created At", "Resolved At", "Resolved By", "Sample Content",
                "Annotations", "Latest Decision", "Decision Reasoning"
            ])
            
            conflicts = (
                self.db.query(Conflict)
                .join(Sample)
                .filter(Sample.project_id == project_id)
                .all()
            )
            
            for conflict in conflicts:
                annotations_str = "; ".join([
                    f"{a.annotator}: {a.label} (conf: {a.confidence})"
                    for a in conflict.sample.annotations
                ])
                
                latest_decision = None
                latest_reasoning = None
                for d in conflict.review_decisions:
                    if d.is_final:
                        latest_decision = d.decision
                        latest_reasoning = d.reasoning
                        break
                
                writer.writerow([
                    conflict.id,
                    conflict.sample_id,
                    conflict.status,
                    conflict.severity,
                    conflict.detection_method or "",
                    conflict.created_at.isoformat(),
                    conflict.resolved_at.isoformat() if conflict.resolved_at else "",
                    conflict.resolved_by or "",
                    conflict.sample.content[:100],
                    annotations_str,
                    latest_decision or "",
                    latest_reasoning or ""
                ])
        
        elif data_type == "samples":
            writer.writerow([
                "Sample ID", "External ID", "Content", "Annotation Count",
                "Conflict Count", "Has Active Conflict"
            ])
            
            samples = self.db.query(Sample).filter(Sample.project_id == project_id).all()
            
            for sample in samples:
                active_conflicts = sum(1 for c in sample.conflicts if c.status == "pending")
                writer.writerow([
                    sample.id,
                    sample.external_id or "",
                    sample.content[:100],
                    len(sample.annotations),
                    len(sample.conflicts),
                    active_conflicts > 0
                ])
        
        elif data_type == "decisions":
            writer.writerow([
                "Decision ID", "Conflict ID", "Reviewer", "Decision",
                "Reasoning", "Created At", "Is Final"
            ])
            
            decisions = (
                self.db.query(ReviewDecision)
                .join(Conflict)
                .join(Sample)
                .filter(Sample.project_id == project_id)
                .all()
            )
            
            for d in decisions:
                writer.writerow([
                    d.id,
                    d.conflict_id,
                    d.reviewer,
                    d.decision,
                    d.reasoning,
                    d.created_at.isoformat(),
                    d.is_final
                ])
        
        return output.getvalue().encode('utf-8')
    
    def export_training_data(self, project_id: int) -> List[Dict[str, Any]]:
        samples = self.db.query(Sample).filter(Sample.project_id == project_id).all()
        
        training_data = []
        for sample in samples:
            active_annotations = [a for a in sample.annotations if a.is_active]
            
            if active_annotations:
                labels = [a.label for a in active_annotations]
                label_counts = Counter(labels)
                final_label = max(label_counts, key=label_counts.get)
                
                training_data.append({
                    "id": sample.external_id or sample.id,
                    "content": sample.content,
                    "label": final_label,
                    "annotation_count": len(active_annotations),
                    "label_distribution": dict(label_counts),
                    "metadata": json.loads(sample.metadata) if sample.metadata else {}
                })
        
        return training_data
    
    def export_to_excel(self, project_id: int) -> bytes:
        try:
            import openpyxl
            from openpyxl import Workbook
        except ImportError:
            raise ImportError("openpyxl is required for Excel export")
        
        wb = Workbook()
        
        ws1 = wb.active
        ws1.title = "Summary"
        summary = self.generate_summary(project_id)
        
        ws1.append(["Project", summary["project_name"]])
        ws1.append(["Generated At", summary["generated_at"]])
        ws1.append([""])
        ws1.append(["Key Metrics"])
        ws1.append(["Total Samples", summary["summary"]["total_samples"]])
        ws1.append(["Total Annotations", summary["summary"]["total_annotations"]])
        ws1.append(["Total Conflicts", summary["summary"]["total_conflicts"]])
        ws1.append(["Resolved Conflicts", summary["summary"]["resolved_conflicts"]])
        ws1.append(["Pending Conflicts", summary["summary"]["pending_conflicts"]])
        ws1.append(["Conflict Rate (%)", summary["summary"]["conflict_rate_percent"]])
        ws1.append(["Resolution Rate (%)", summary["summary"]["resolution_rate"]])
        
        ws2 = wb.create_sheet("Annotators")
        ws2.append(["Annotator", "Total Annotations", "Accepted", "Accuracy Rate (%)"])
        for stat in summary["annotator_statistics"]:
            ws2.append([
                stat["annotator"],
                stat["total_annotations"],
                stat["accepted_annotations"],
                stat["accuracy_rate"]
            ])
        
        ws3 = wb.create_sheet("Conflicts")
        ws3.append([
            "Conflict ID", "Status", "Severity", "Detection Method",
            "Created At", "Resolved By", "Decision"
        ])
        
        conflicts = (
            self.db.query(Conflict)
            .join(Sample)
            .filter(Sample.project_id == project_id)
            .all()
        )
        
        for c in conflicts:
            latest_decision = ""
            for d in c.review_decisions:
                if d.is_final:
                    latest_decision = d.decision
                    break
            
            ws3.append([
                c.id,
                c.status,
                c.severity,
                c.detection_method or "",
                c.created_at.isoformat(),
                c.resolved_by or "",
                latest_decision
            ])
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
