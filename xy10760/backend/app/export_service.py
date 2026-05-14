from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session
from datetime import datetime
from .services import SimulationService
from . import models

class ExportService:
    @staticmethod
    def export_simulation_to_excel(db: Session, simulation_id: int) -> BytesIO:
        detail = SimulationService.get_simulation_detail(db, simulation_id)
        if not detail:
            return None

        wb = Workbook()
        
        ws1 = wb.active
        ws1.title = "模拟结果概览"
        ExportService._fill_overview_sheet(ws1, detail)
        
        ws2 = wb.create_sheet("申请单信息")
        ExportService._fill_application_sheet(ws2, detail["application"])
        
        ws3 = wb.create_sheet("规则版本信息")
        ExportService._fill_rule_version_sheet(ws3, detail["rule_version"])
        
        ws4 = wb.create_sheet("命中条件")
        ExportService._fill_hit_conditions_sheet(ws4, detail["hit_conditions"])
        
        ws5 = wb.create_sheet("审批人信息")
        ExportService._fill_approvers_sheet(ws5, detail["approvers"], detail["simulation"].approver_errors)
        
        ws6 = wb.create_sheet("错误明细")
        ExportService._fill_error_details_sheet(ws6, detail["simulation"])

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _set_header_style(cell):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    @staticmethod
    def _fill_overview_sheet(ws, detail):
        simulation = detail["simulation"]
        
        headers = ["项目", "内容"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        status_map = {
            "success": ("成功", "70AD47"),
            "warning": ("警告", "FFC000"),
            "error": ("失败", "FF0000")
        }
        status_text, status_color = status_map.get(simulation.simulation_status, ("未知", "808080"))

        data = [
            ["模拟编号", simulation.id],
            ["申请单编号", detail["application"].application_no if detail["application"] else "-"],
            ["规则版本", detail["rule_version"].version if detail["rule_version"] else "-"],
            ["模拟状态", status_text],
            ["最终审批结果", simulation.final_approval_result or "待确认"],
            ["是否跳过审批", "是" if simulation.skip_reason_id else "否"],
            ["跳过原因", simulation.skip_reason_text or "-"],
            ["跳过原因已确认", "是" if simulation.skip_manual_confirmed else "否"],
            ["确认人", simulation.skip_confirmed_by or "-"],
            ["确认时间", simulation.skip_confirmed_at.strftime("%Y-%m-%d %H:%M:%S") if simulation.skip_confirmed_at else "-"],
            ["命中条件数量", len(simulation.hit_condition_ids)],
            ["审批人数量", len(simulation.approver_ids)],
            ["是否已发布", "是" if simulation.published else "否"],
            ["发布人", simulation.published_by or "-"],
            ["发布时间", simulation.published_at.strftime("%Y-%m-%d %H:%M:%S") if simulation.published_at else "-"],
            ["创建人", simulation.created_by],
            ["创建时间", simulation.created_at.strftime("%Y-%m-%d %H:%M:%S")]
        ]

        for row, (key, value) in enumerate(data, 2):
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40

    @staticmethod
    def _fill_application_sheet(ws, application):
        if not application:
            return

        headers = ["字段", "值"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        data = [
            ["申请单编号", application.application_no],
            ["申请人", application.applicant],
            ["申请部门", application.department],
            ["申请类型", application.application_type],
            ["申请金额", f"{application.amount:,} 元"],
            ["申请单状态", application.status],
            ["创建时间", application.created_at.strftime("%Y-%m-%d %H:%M:%S")]
        ]

        for row, (key, value) in enumerate(data, 2):
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40

    @staticmethod
    def _fill_rule_version_sheet(ws, rule_version):
        if not rule_version:
            return

        headers = ["字段", "值"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        data = [
            ["规则版本号", rule_version.version],
            ["规则名称", rule_version.name],
            ["规则描述", rule_version.description],
            ["是否启用", "是" if rule_version.is_active else "否"],
            ["创建人", rule_version.created_by],
            ["创建时间", rule_version.created_at.strftime("%Y-%m-%d %H:%M:%S")]
        ]

        for row, (key, value) in enumerate(data, 2):
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40

    @staticmethod
    def _fill_hit_conditions_sheet(ws, hit_conditions):
        headers = ["序号", "条件类型", "条件表达式", "条件值", "运算符", "优先级"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        for row, condition in enumerate(hit_conditions, 2):
            ws.cell(row=row, column=1, value=row - 1)
            ws.cell(row=row, column=2, value=condition.condition_type)
            ws.cell(row=row, column=3, value=condition.condition_expression)
            ws.cell(row=row, column=4, value=condition.condition_value)
            ws.cell(row=row, column=5, value=condition.operator)
            ws.cell(row=row, column=6, value=condition.priority)

        for col in range(1, 7):
            ws.column_dimensions[chr(64 + col)].width = 20

    @staticmethod
    def _fill_approvers_sheet(ws, approvers, approver_errors):
        headers = ["序号", "审批人姓名", "邮箱", "部门", "级别", "状态", "异常说明"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        error_map = {}
        for error in approver_errors:
            error_map[error.get("approver_id")] = error

        for row, approver in enumerate(approvers, 2):
            error = error_map.get(approver.id)
            ws.cell(row=row, column=1, value=row - 1)
            ws.cell(row=row, column=2, value=approver.name)
            ws.cell(row=row, column=3, value=approver.email)
            ws.cell(row=row, column=4, value=approver.department)
            ws.cell(row=row, column=5, value=f"L{approver.level}")
            ws.cell(row=row, column=6, value="正常" if approver.is_active else "停用")
            ws.cell(row=row, column=7, value=error.get("error_message", "") if error else "")

        for col in range(1, 8):
            ws.column_dimensions[chr(64 + col)].width = 18

    @staticmethod
    def _fill_error_details_sheet(ws, simulation):
        headers = ["序号", "异常类型", "异常对象", "异常详情"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            ExportService._set_header_style(cell)

        row = 2
        for error in simulation.approver_errors:
            ws.cell(row=row, column=1, value=row - 1)
            ws.cell(row=row, column=2, value="审批人异常")
            ws.cell(row=row, column=3, value=error.get("approver_name", ""))
            ws.cell(row=row, column=4, value=error.get("error_message", ""))
            row += 1

        if simulation.error_details:
            ws.cell(row=row, column=1, value=row - 1)
            ws.cell(row=row, column=2, value="系统异常")
            ws.cell(row=row, column=3, value="系统")
            ws.cell(row=row, column=4, value=simulation.error_details)

        for col in range(1, 5):
            ws.column_dimensions[chr(64 + col)].width = 25
