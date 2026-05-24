from typing import List, Dict, Any
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
from sqlalchemy.orm import Session
from app.models import BlastPlan, BlastStatus
from app.services.receipt_tracker import receipt_tracker


class ReportExporter:
    def export_plan_to_excel(
        self,
        db: Session,
        plan: BlastPlan,
    ) -> BytesIO:
        output = BytesIO()
        wb = Workbook()

        ws = wb.active
        ws.title = "爆破计划详情"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        def write_header(row: int, col: int, text: str):
            cell = ws.cell(row=row, column=col, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A1:H1")
        title_cell = ws.cell(row=1, column=1, value="采石场爆破计划执行报告")
        title_cell.font = Font(bold=True, size=16)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.cell(row=3, column=1, value="业务编号:").font = Font(bold=True)
        ws.cell(row=3, column=2, value=plan.business_no)
        ws.cell(row=3, column=5, value="采石场名称:").font = Font(bold=True)
        ws.cell(row=3, column=6, value=plan.quarry_name)

        ws.cell(row=4, column=1, value="计划爆破时间:").font = Font(bold=True)
        ws.cell(row=4, column=2, value=plan.blast_time.strftime("%Y-%m-%d %H:%M:%S") if plan.blast_time else "")
        ws.cell(row=4, column=5, value="当前状态:").font = Font(bold=True)
        ws.cell(row=4, column=6, value=plan.status.value)

        ws.cell(row=5, column=1, value="预计爆破量:").font = Font(bold=True)
        ws.cell(row=5, column=2, value=f"{plan.expected_blast_volume} 立方米")
        ws.cell(row=5, column=5, value="创建人:").font = Font(bold=True)
        ws.cell(row=5, column=6, value=plan.created_by or "")

        ws.cell(row=6, column=1, value="风向:").font = Font(bold=True)
        ws.cell(row=6, column=2, value=plan.wind_direction or "")
        ws.cell(row=6, column=5, value="风速:").font = Font(bold=True)
        ws.cell(row=6, column=6, value=f"{plan.wind_speed} m/s" if plan.wind_speed else "")

        ws.cell(row=7, column=1, value="安全措施:").font = Font(bold=True)
        ws.cell(row=7, column=2, value=plan.safety_measures or "")

        ws.row_dimensions[9].height = 20
        headers = ["警戒区名称", "边界描述", "半径(米)", "版本"]
        for col, header in enumerate(headers, 1):
            write_header(9, col, header)

        row = 10
        for zone in plan.zones:
            if zone.is_active:
                ws.cell(row=row, column=1, value=zone.zone_name)
                ws.cell(row=row, column=2, value=zone.boundary_description)
                ws.cell(row=row, column=3, value=zone.radius_meters or "")
                ws.cell(row=row, column=4, value=zone.version)
                row += 1

        row += 2
        ws.row_dimensions[row].height = 20
        notice_headers = ["通知类型", "接收方", "联系电话", "地址", "回执状态"]
        for col, header in enumerate(notice_headers, 1):
            write_header(row, col, header)

        row += 1
        for notice in plan.notices:
            if notice.is_active:
                ws.cell(row=row, column=1, value=notice.notice_type.value)
                ws.cell(row=row, column=2, value=notice.recipient_name)
                ws.cell(row=row, column=3, value=notice.contact_phone or "")
                ws.cell(row=row, column=4, value=notice.address or "")
                has_receipt = any(r.notice_id == notice.id for r in plan.receipts)
                ws.cell(row=row, column=5, value="已确认" if has_receipt else "未确认")
                row += 1

        row += 2
        ws.row_dimensions[row].height = 20
        audit_headers = ["时间", "操作", "操作人", "原状态", "新状态", "详情"]
        for col, header in enumerate(audit_headers, 1):
            write_header(row, col, header)

        row += 1
        for audit in plan.audits:
            ws.cell(row=row, column=1, value=audit.created_at.strftime("%Y-%m-%d %H:%M:%S") if audit.created_at else "")
            ws.cell(row=row, column=2, value=audit.action.value)
            ws.cell(row=row, column=3, value=audit.operator or "")
            ws.cell(row=row, column=4, value=audit.old_status or "")
            ws.cell(row=row, column=5, value=audit.new_status or "")
            ws.cell(row=row, column=6, value=audit.detail or "")
            row += 1

        if plan.report:
            row += 2
            ws.merge_cells(f"A{row}:F{row}")
            report_title = ws.cell(row=row, column=1, value="执行报告")
            report_title.font = Font(bold=True, size=14)
            row += 2

            ws.cell(row=row, column=1, value="实际爆破时间:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=plan.report.actual_blast_time.strftime("%Y-%m-%d %H:%M:%S") if plan.report.actual_blast_time else "")
            row += 1

            ws.cell(row=row, column=1, value="实际爆破量:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=f"{plan.report.actual_blast_volume} 立方米" if plan.report.actual_blast_volume else "")
            row += 1

            ws.cell(row=row, column=1, value="爆破时风向:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=plan.report.wind_direction_at_blast or "")
            row += 1

            ws.cell(row=row, column=1, value="爆破时风速:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=f"{plan.report.wind_speed_at_blast} m/s" if plan.report.wind_speed_at_blast else "")
            row += 1

            ws.cell(row=row, column=1, value="现场监督:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=plan.report.on_site_supervisor or "")
            row += 1

            ws.cell(row=row, column=1, value="安全检查结果:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=plan.report.safety_check_result or "")
            row += 1

            ws.cell(row=row, column=1, value="异常情况:").font = Font(bold=True)
            ws.cell(row=row, column=2, value=plan.report.abnormal_situation or "无")

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 30

        wb.save(output)
        output.seek(0)
        return output

    def export_batch_plans_to_excel(
        self,
        db: Session,
        plans: List[BlastPlan],
    ) -> BytesIO:
        output = BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "爆破计划汇总"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        def write_header(col: int, text: str):
            cell = ws.cell(row=1, column=col, value=text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = [
            "业务编号", "采石场", "爆破时间", "状态", "风向", "风速(m/s)",
            "通知数", "回执数", "创建人", "创建时间"
        ]
        for col, header in enumerate(headers, 1):
            write_header(col, header)

        for row_idx, plan in enumerate(plans, 2):
            receipt_summary = receipt_tracker.get_receipt_summary(db, plan.id)

            ws.cell(row=row_idx, column=1, value=plan.business_no)
            ws.cell(row=row_idx, column=2, value=plan.quarry_name)
            ws.cell(row=row_idx, column=3, value=plan.blast_time.strftime("%Y-%m-%d %H:%M") if plan.blast_time else "")
            ws.cell(row=row_idx, column=4, value=plan.status.value)
            ws.cell(row=row_idx, column=5, value=plan.wind_direction or "")
            ws.cell(row=row_idx, column=6, value=plan.wind_speed or "")
            ws.cell(row=row_idx, column=7, value=receipt_summary["total_notices"])
            ws.cell(row=row_idx, column=8, value=receipt_summary["received_count"])
            ws.cell(row=row_idx, column=9, value=plan.created_by or "")
            ws.cell(row=row_idx, column=10, value=plan.created_at.strftime("%Y-%m-%d %H:%M") if plan.created_at else "")

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        wb.save(output)
        output.seek(0)
        return output


report_exporter = ReportExporter()
