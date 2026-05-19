from typing import List, Dict, Any, Optional
from datetime import datetime
from pathlib import Path
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from exhibition_material.models import RecordStatus, AnomalyType
from exhibition_material.services import QueryService

class ReportExporter:
    def __init__(self, session: Session, export_dir: str = None):
        self.session = session
        self.query_service = QueryService(session)
        self.export_dir = Path(export_dir) if export_dir else Path.cwd() / "exports"
        self.export_dir.mkdir(parents=True, exist_ok=True)
    
    def export_allocations_to_csv(self, file_name: str = None,
                                   responsible_person: str = None,
                                   booth_number: str = None,
                                   status: RecordStatus = None,
                                   has_anomaly: bool = None,
                                   anomaly_type: AnomalyType = None,
                                   start_date: datetime = None,
                                   end_date: datetime = None,
                                   material_id: int = None) -> Dict[str, Any]:
        allocations = self.query_service.query_allocations(
            responsible_person=responsible_person,
            booth_number=booth_number,
            status=status,
            has_anomaly=has_anomaly,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            material_id=material_id
        )
        
        if not file_name:
            file_name = f"allocations_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        
        file_path = self.export_dir / file_name
        
        fieldnames = [
            "调拨单号", "展位号", "物料编码", "物料名称", "物料类型",
            "数量", "已归还数量", "剩余数量", "单位", "负责人", "联系电话",
            "调拨日期", "预计归还日期", "状态", "是否异常", "异常类型",
            "异常备注", "审批人", "审批日期", "备注"
        ]
        
        with open(file_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            
            for alloc in allocations:
                material = alloc.get("material", {})
                writer.writerow({
                    "调拨单号": alloc.get("allocation_no", ""),
                    "展位号": alloc.get("booth_number", ""),
                    "物料编码": material.get("code", ""),
                    "物料名称": material.get("name", ""),
                    "物料类型": material.get("type", ""),
                    "数量": alloc.get("quantity", 0),
                    "已归还数量": alloc.get("returned_quantity", 0),
                    "剩余数量": alloc.get("remaining_quantity", 0),
                    "单位": material.get("unit", ""),
                    "负责人": alloc.get("responsible_person", ""),
                    "联系电话": alloc.get("contact_phone", ""),
                    "调拨日期": alloc.get("allocated_at", ""),
                    "预计归还日期": alloc.get("expected_return_at", ""),
                    "状态": alloc.get("status", ""),
                    "是否异常": "是" if alloc.get("has_anomaly") else "否",
                    "异常类型": alloc.get("anomaly_type", ""),
                    "异常备注": alloc.get("anomaly_remark", ""),
                    "审批人": alloc.get("approved_by", ""),
                    "审批日期": alloc.get("approved_at", ""),
                    "备注": alloc.get("remarks", "")
                })
        
        return {
            "file_path": str(file_path),
            "file_name": file_name,
            "record_count": len(allocations),
            "export_time": datetime.now().isoformat()
        }
    
    def export_returns_to_csv(self, file_name: str = None,
                              received_by: str = None,
                              returned_by: str = None,
                              booth_number: str = None,
                              status: RecordStatus = None,
                              has_anomaly: bool = None,
                              anomaly_type: AnomalyType = None,
                              start_date: datetime = None,
                              end_date: datetime = None,
                              allocation_id: int = None,
                              material_id: int = None) -> Dict[str, Any]:
        records = self.query_service.query_return_records(
            received_by=received_by,
            returned_by=returned_by,
            booth_number=booth_number,
            status=status,
            has_anomaly=has_anomaly,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            allocation_id=allocation_id,
            material_id=material_id
        )
        
        if not file_name:
            file_name = f"returns_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        
        file_path = self.export_dir / file_name
        
        fieldnames = [
            "归还单号", "调拨单号", "展位号", "物料编码", "物料名称",
            "归还数量", "归还日期", "接收人", "归还人", "状态",
            "是否异常", "异常类型", "异常备注", "状况备注", "备注"
        ]
        
        with open(file_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            
            for record in records:
                material = record.get("material", {})
                writer.writerow({
                    "归还单号": record.get("return_no", ""),
                    "调拨单号": record.get("allocation", {}).get("allocation_no", ""),
                    "展位号": record.get("booth_number", ""),
                    "物料编码": material.get("code", ""),
                    "物料名称": material.get("name", ""),
                    "归还数量": record.get("quantity", 0),
                    "归还日期": record.get("returned_at", ""),
                    "接收人": record.get("received_by", ""),
                    "归还人": record.get("returned_by", ""),
                    "状态": record.get("status", ""),
                    "是否异常": "是" if record.get("has_anomaly") else "否",
                    "异常类型": record.get("anomaly_type", ""),
                    "异常备注": record.get("anomaly_remark", ""),
                    "状况备注": record.get("condition_remark", ""),
                    "备注": record.get("remarks", "")
                })
        
        return {
            "file_path": str(file_path),
            "file_name": file_name,
            "record_count": len(records),
            "export_time": datetime.now().isoformat()
        }
    
    def export_allocations_to_excel(self, file_name: str = None,
                                    responsible_person: str = None,
                                    booth_number: str = None,
                                    status: RecordStatus = None,
                                    has_anomaly: bool = None,
                                    anomaly_type: AnomalyType = None,
                                    start_date: datetime = None,
                                    end_date: datetime = None,
                                    material_id: int = None) -> Dict[str, Any]:
        allocations = self.query_service.query_allocations(
            responsible_person=responsible_person,
            booth_number=booth_number,
            status=status,
            has_anomaly=has_anomaly,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            material_id=material_id
        )
        
        if not file_name:
            file_name = f"allocations_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        file_path = self.export_dir / file_name
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "调拨记录"
        
        headers = [
            "调拨单号", "展位号", "物料编码", "物料名称", "物料类型",
            "数量", "已归还数量", "剩余数量", "单位", "负责人", "联系电话",
            "调拨日期", "预计归还日期", "状态", "是否异常", "异常类型",
            "异常备注", "审批人", "审批日期", "备注"
        ]
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        
        for row_idx, alloc in enumerate(allocations, 2):
            material = alloc.get("material", {})
            row_data = [
                alloc.get("allocation_no", ""),
                alloc.get("booth_number", ""),
                material.get("code", ""),
                material.get("name", ""),
                material.get("type", ""),
                alloc.get("quantity", 0),
                alloc.get("returned_quantity", 0),
                alloc.get("remaining_quantity", 0),
                material.get("unit", ""),
                alloc.get("responsible_person", ""),
                alloc.get("contact_phone", ""),
                alloc.get("allocated_at", ""),
                alloc.get("expected_return_at", ""),
                alloc.get("status", ""),
                "是" if alloc.get("has_anomaly") else "否",
                alloc.get("anomaly_type", ""),
                alloc.get("anomaly_remark", ""),
                alloc.get("approved_by", ""),
                alloc.get("approved_at", ""),
                alloc.get("remarks", "")
            ]
            
            anomaly_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = center_align
                cell.border = thin_border
                
                if alloc.get("has_anomaly"):
                    cell.fill = anomaly_fill
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(file_path)
        
        return {
            "file_path": str(file_path),
            "file_name": file_name,
            "record_count": len(allocations),
            "export_time": datetime.now().isoformat()
        }
    
    def export_returns_to_excel(self, file_name: str = None,
                                received_by: str = None,
                                returned_by: str = None,
                                booth_number: str = None,
                                status: RecordStatus = None,
                                has_anomaly: bool = None,
                                anomaly_type: AnomalyType = None,
                                start_date: datetime = None,
                                end_date: datetime = None,
                                allocation_id: int = None,
                                material_id: int = None) -> Dict[str, Any]:
        records = self.query_service.query_return_records(
            received_by=received_by,
            returned_by=returned_by,
            booth_number=booth_number,
            status=status,
            has_anomaly=has_anomaly,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            allocation_id=allocation_id,
            material_id=material_id
        )
        
        if not file_name:
            file_name = f"returns_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        file_path = self.export_dir / file_name
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "归还记录"
        
        headers = [
            "归还单号", "调拨单号", "展位号", "物料编码", "物料名称",
            "归还数量", "归还日期", "接收人", "归还人", "状态",
            "是否异常", "异常类型", "异常备注", "状况备注", "备注"
        ]
        
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        
        for row_idx, record in enumerate(records, 2):
            material = record.get("material", {})
            row_data = [
                record.get("return_no", ""),
                record.get("allocation", {}).get("allocation_no", ""),
                record.get("booth_number", ""),
                material.get("code", ""),
                material.get("name", ""),
                record.get("quantity", 0),
                record.get("returned_at", ""),
                record.get("received_by", ""),
                record.get("returned_by", ""),
                record.get("status", ""),
                "是" if record.get("has_anomaly") else "否",
                record.get("anomaly_type", ""),
                record.get("anomaly_remark", ""),
                record.get("condition_remark", ""),
                record.get("remarks", "")
            ]
            
            anomaly_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = center_align
                cell.border = thin_border
                
                if record.get("has_anomaly"):
                    cell.fill = anomaly_fill
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(file_path)
        
        return {
            "file_path": str(file_path),
            "file_name": file_name,
            "record_count": len(records),
            "export_time": datetime.now().isoformat()
        }
    
    def export_summary_report(self, file_name: str = None) -> Dict[str, Any]:
        anomalies = self.query_service.get_all_anomalies()
        
        if not file_name:
            file_name = f"summary_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        file_path = self.export_dir / file_name
        
        wb = openpyxl.Workbook()
        
        ws1 = wb.active
        ws1.title = "汇总概览"
        
        ws1.merge_cells('A1:D1')
        ws1['A1'] = "会展物料管理汇总报告"
        ws1['A1'].font = Font(bold=True, size=16)
        ws1['A1'].alignment = Alignment(horizontal="center", vertical="center")
        
        ws1.merge_cells('A2:D2')
        ws1['A2'] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws1['A2'].alignment = Alignment(horizontal="center", vertical="center")
        
        summary_data = [
            ["统计项", "数量", "", ""],
            ["调拨异常记录", len(anomalies.get("allocation_anomalies", [])), "", ""],
            ["归还异常记录", len(anomalies.get("return_anomalies", [])), "", ""],
            ["导入错误记录", len(anomalies.get("import_errors", [])), "", ""],
            ["异常总计", anomalies.get("total_anomalies", 0), "", ""],
        ]
        
        for row_idx, row in enumerate(summary_data, 4):
            for col_idx, value in enumerate(row, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 4:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
        
        wb.save(file_path)
        
        return {
            "file_path": str(file_path),
            "file_name": file_name,
            "export_time": datetime.now().isoformat(),
            "summary": {
                "allocation_anomalies": len(anomalies.get("allocation_anomalies", [])),
                "return_anomalies": len(anomalies.get("return_anomalies", [])),
                "import_errors": len(anomalies.get("import_errors", [])),
                "total_anomalies": anomalies.get("total_anomalies", 0)
            }
        }
