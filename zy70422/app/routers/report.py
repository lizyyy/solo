from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional, List
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..database import get_db
from ..models import Store, Rectification, RectificationStatus, RectificationEvent, InspectionRecord, Inspection, InspectionItem
from ..schemas import RegionReport
from ..services import get_region_report

import urllib.parse

router = APIRouter(prefix="/api/reports", tags=["统计报表"])


@router.get("/regions", response_model=List[RegionReport])
def region_reports(
    region: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db)
):
    if region:
        report = get_region_report(db, region, start_date, end_date)
        return [report]

    regions = db.query(Store.region).filter(
        Store.region.isnot(None)
    ).distinct().all()

    reports = []
    for (reg,) in regions:
        report = get_region_report(db, reg, start_date, end_date)
        reports.append(report)

    total_report = get_region_report(db, None, start_date, end_date)
    reports.append(total_report)

    return reports


@router.get("/region-list")
def list_regions(db: Session = Depends(get_db)):
    regions = db.query(Store.region).filter(
        Store.region.isnot(None),
        Store.region != ""
    ).distinct().all()
    return {"regions": [r[0] for r in regions]}


@router.get("/export/rectifications")
def export_rectifications(
    region: Optional[str] = None,
    status: Optional[RectificationStatus] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db)
):
    wb = Workbook()
    ws = wb.active
    ws.title = "整改任务明细"

    headers = [
        "整改编号", "门店", "区域", "城市",
        "巡检项", "问题类型",
        "责任人", "状态",
        "创建时间", "截止时间", "整改提交时间", "复查时间",
        "最终得分", "扣分", "是否逾期",
        "重试次数", "整改说明", "复查说明"
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment

    query = db.query(Rectification, Store, InspectionItem) \
        .join(InspectionRecord, Rectification.record_id == InspectionRecord.id) \
        .join(Inspection, InspectionRecord.inspection_id == Inspection.id) \
        .join(Store, Inspection.store_id == Store.id) \
        .join(InspectionItem, InspectionRecord.item_id == InspectionItem.id)

    if region:
        query = query.filter(Store.region == region)
    if status:
        query = query.filter(Rectification.status == status)
    if start_date:
        query = query.filter(Rectification.created_at >= start_date)
    if end_date:
        query = query.filter(Rectification.created_at <= end_date)

    results = query.order_by(Rectification.created_at.desc()).all()

    for row_idx, (rect, store, item) in enumerate(results, 2):
        is_overdue = rect.status == RectificationStatus.OVERDUE
        if rect.status not in [RectificationStatus.PASSED, RectificationStatus.CANCELLED]:
            is_overdue = datetime.utcnow() > rect.deadline

        ws.cell(row=row_idx, column=1, value=f"R{rect.id:06d}")
        ws.cell(row=row_idx, column=2, value=store.name)
        ws.cell(row=row_idx, column=3, value=store.region or "")
        ws.cell(row=row_idx, column=4, value=store.city or "")
        ws.cell(row=row_idx, column=5, value=item.name)
        ws.cell(row=row_idx, column=6, value=item.category or "")
        ws.cell(row=row_idx, column=7, value=rect.assignee)
        ws.cell(row=row_idx, column=8, value=rect.status.value)
        ws.cell(row=row_idx, column=9, value=rect.created_at.strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=row_idx, column=10, value=rect.deadline.strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=row_idx, column=11, value=rect.rectification_at.strftime("%Y-%m-%d %H:%M") if rect.rectification_at else "")
        ws.cell(row=row_idx, column=12, value=rect.recheck_at.strftime("%Y-%m-%d %H:%M") if rect.recheck_at else "")
        ws.cell(row=row_idx, column=13, value=rect.final_score)
        ws.cell(row=row_idx, column=14, value=rect.final_deduction)
        ws.cell(row=row_idx, column=15, value="是" if is_overdue else "否")
        ws.cell(row=row_idx, column=16, value=rect.retry_count)
        ws.cell(row=row_idx, column=17, value=rect.rectification_description or "")
        ws.cell(row=row_idx, column=18, value=rect.recheck_remark or "")

    column_widths = [12, 20, 10, 10, 25, 15, 12, 10, 18, 18, 18, 18, 10, 8, 8, 8, 30, 30]
    for col_idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    if len(results) > 0:
        summary_ws = wb.create_sheet("数据汇总")
        summary_headers = ["统计项", "数值"]
        for col_idx, header in enumerate(summary_headers, 1):
            cell = summary_ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill

        status_counts = {}
        for rect, _, _ in results:
            status_counts[rect.status.value] = status_counts.get(rect.status.value, 0) + 1

        summary_rows = [
            ["整改任务总数", len(results)],
            ["", ""],
            ["各状态分布：", ""],
        ]

        for status, count in status_counts.items():
            summary_rows.append([f"  - {status}", count])

        total_deduction = sum(r[0].final_deduction or 0 for r in results)
        summary_rows.extend([
            ["", ""],
            ["扣分合计", total_deduction],
        ])

        for row_idx, (label, value) in enumerate(summary_rows, 2):
            summary_ws.cell(row=row_idx, column=1, value=label)
            summary_ws.cell(row=row_idx, column=2, value=value)

        summary_ws.column_dimensions["A"].width = 20
        summary_ws.column_dimensions["B"].width = 12

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"rectification_summary_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )


@router.get("/export/trace/{rect_id}")
def export_rectification_trace(rect_id: int, db: Session = Depends(get_db)):
    rect = db.query(Rectification).filter(Rectification.id == rect_id).first()
    if not rect:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="整改任务不存在")

    from ..services import get_rectification_trace
    trace = get_rectification_trace(db, rect_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "整改追踪"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    rect_info = [
        ["整改基本信息", ""],
        ["整改编号", f"R{rect.id:06d}"],
        ["责任人", rect.assignee],
        ["当前状态", rect.status.value],
        ["创建时间", rect.created_at.strftime("%Y-%m-%d %H:%M:%S")],
        ["截止时间", rect.deadline.strftime("%Y-%m-%d %H:%M:%S")],
        ["重试次数", rect.retry_count],
        ["最终得分", rect.final_score if rect.final_score is not None else "未完成"],
        ["扣分", rect.final_deduction],
    ]

    for row_idx, (label, value) in enumerate(rect_info, 1):
        if label == "整改基本信息":
            cell = ws.cell(row=row_idx, column=1, value=label)
            cell.font = header_font
            cell.fill = header_fill
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
        else:
            ws.cell(row=row_idx, column=1, value=label)
            ws.cell(row=row_idx, column=2, value=value)

    ws.cell(row=len(rect_info) + 2, column=1, value="事件日志")
    event_header_cell = ws.cell(row=len(rect_info) + 2, column=1, value="事件日志")
    event_header_cell.font = header_font
    event_header_cell.fill = header_fill
    ws.merge_cells(start_row=len(rect_info) + 2, start_column=1, end_row=len(rect_info) + 2, end_column=6)

    event_headers = ["时间", "事件类型", "从状态", "到状态", "操作人", "说明"]
    for col_idx, header in enumerate(event_headers, 1):
        cell = ws.cell(row=len(rect_info) + 3, column=col_idx, value=header)
        cell.font = Font(bold=True)

    events = trace["events"]
    for row_idx, event in enumerate(events, len(rect_info) + 4):
        ws.cell(row=row_idx, column=1, value=event.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row_idx, column=2, value=event.event_type)
        ws.cell(row=row_idx, column=3, value=event.from_status.value if event.from_status else "")
        ws.cell(row=row_idx, column=4, value=event.to_status.value)
        ws.cell(row=row_idx, column=5, value=event.actor or "")
        ws.cell(row=row_idx, column=6, value=event.description or "")

    photos_start = len(rect_info) + 5 + len(events)
    ws.cell(row=photos_start, column=1, value="照片证据")
    photo_header_cell = ws.cell(row=photos_start, column=1, value="照片证据")
    photo_header_cell.font = header_font
    photo_header_cell.fill = header_fill
    ws.merge_cells(start_row=photos_start, start_column=1, end_row=photos_start, end_column=5)

    photo_headers = ["照片类型", "文件名", "上传人", "上传时间", "说明"]
    for col_idx, header in enumerate(photo_headers, 1):
        cell = ws.cell(row=photos_start + 1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    photos = trace["photos"]
    for row_idx, photo in enumerate(photos, photos_start + 2):
        ws.cell(row=row_idx, column=1, value=photo.photo_type.value)
        ws.cell(row=row_idx, column=2, value=photo.file_name or photo.file_path)
        ws.cell(row=row_idx, column=3, value=photo.uploaded_by or "")
        ws.cell(row=row_idx, column=4, value=photo.upload_time.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row_idx, column=5, value=photo.description or "")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 40

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"rectification_trace_R{rect.id:06d}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )
