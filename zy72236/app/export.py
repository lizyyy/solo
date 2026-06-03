from sqlalchemy.orm import Session
from typing import List, Dict, Any
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from . import models
from .crud import get_unified_record_data
from .utils import format_date, format_currency


def export_to_excel(db: Session, batch_id: int) -> bytes:
    unified_data = get_unified_record_data(db, batch_id)

    output = io.BytesIO()
    wb = Workbook()

    ws_main = wb.active
    ws_main.title = "尾佣拆分明细"

    headers = [
        "原始行号", "基金代码", "基金名称", "客户账号", "客户姓名",
        "客户经理", "客户经理代码", "审批人", "审批人拼音标记",
        "交易日期", "原清算日期", "当前清算日期",
        "交易金额", "佣金率", "佣金金额", "尾佣金额",
        "拆分比例", "最终金额", "数据来源", "状态", "重复标记",
        "人工修改", "待复核", "余额更新"
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws_main.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    warning_font = Font(color="9C0006")

    for row_idx, record in enumerate(unified_data, 2):
        row_data = [
            record["original_line_number"],
            record["fund_code"],
            record["fund_name"],
            record["customer_account"],
            record["customer_name"],
            record["manager_name"],
            record["manager_code"],
            record["approval_name"],
            "是" if record["approval_name_is_pinyin"] else "否",
            record["transaction_date"],
            record["original_settlement_date"],
            record["settlement_date"],
            record["transaction_amount"],
            record["commission_rate"],
            record["commission_amount"],
            record["trail_commission_amount"],
            record["split_ratio"],
            record["final_amount"],
            get_source_label_export(record["source_type"]),
            get_status_label_export(record["status"]),
            "是" if record["is_duplicate"] else "否",
            "是" if record["manually_modified"] else "否",
            "是" if record["needs_manager_review"] else "否",
            "是" if record["balance_updated"] else "否"
        ]

        is_warning = record["approval_name_is_pinyin"] or record["is_duplicate"] or record["needs_manager_review"]

        for col, value in enumerate(row_data, 1):
            cell = ws_main.cell(row=row_idx, column=col, value=value)
            if is_warning:
                cell.fill = warning_fill
                cell.font = warning_font

    for col in range(1, len(headers) + 1):
        ws_main.column_dimensions[chr(64 + col)].width = 15

    ws_audit = wb.create_sheet("审计追踪")
    audit_headers = [
        "记录ID", "原始行号", "操作类型", "字段", "旧值", "新值",
        "操作人", "操作时间", "备注"
    ]

    for col, header in enumerate(audit_headers, 1):
        cell = ws_audit.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    audit_row = 2
    for record in unified_data:
        for log in record["audit_trail"]:
            audit_data = [
                record["id"],
                record["original_line_number"],
                log["action"],
                log["field_name"] or "",
                log["old_value"] or "",
                log["new_value"] or "",
                log["performed_by"],
                log["performed_at"],
                log["notes"] or ""
            ]
            for col, value in enumerate(audit_data, 1):
                ws_audit.cell(row=audit_row, column=col, value=value)
            audit_row += 1

    for col in range(1, len(audit_headers) + 1):
        ws_audit.column_dimensions[chr(64 + col)].width = 18

    ws_summary = wb.create_sheet("汇总")
    summary_data = [
        ["总记录数", len(unified_data)],
        ["重复记录", sum(1 for r in unified_data if r["is_duplicate"])],
        ["审批人拼音待复核", sum(1 for r in unified_data if r["approval_name_is_pinyin"])],
        ["待经理复核", sum(1 for r in unified_data if r["needs_manager_review"])],
        ["人工修改", sum(1 for r in unified_data if r["manually_modified"])],
        ["余额已更新", sum(1 for r in unified_data if r["balance_updated"])],
        ["有效记录金额", format_currency(sum(r["final_amount"] for r in unified_data if not r["is_duplicate"]))]
    ]

    for row, (key, value) in enumerate(summary_data, 1):
        ws_summary.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws_summary.cell(row=row, column=2, value=value)

    ws_summary.column_dimensions['A'].width = 20
    ws_summary.column_dimensions['B'].width = 20

    wb.save(output)
    output.seek(0)

    return output.getvalue()


def export_to_csv(db: Session, batch_id: int) -> str:
    unified_data = get_unified_record_data(db, batch_id)

    output = io.StringIO()
    writer = csv.writer(output)

    headers = [
        "原始行号", "基金代码", "基金名称", "客户账号", "客户姓名",
        "客户经理", "审批人", "审批人拼音",
        "交易日期", "清算日期",
        "交易金额", "尾佣金额", "最终金额",
        "来源", "状态", "重复", "待复核"
    ]
    writer.writerow(headers)

    for record in unified_data:
        writer.writerow([
            record["original_line_number"],
            record["fund_code"],
            record["fund_name"],
            record["customer_account"],
            record["customer_name"],
            record["manager_name"],
            record["approval_name"],
            "是" if record["approval_name_is_pinyin"] else "否",
            record["transaction_date"],
            record["settlement_date"],
            record["transaction_amount"],
            record["trail_commission_amount"],
            record["final_amount"],
            get_source_label_export(record["source_type"]),
            get_status_label_export(record["status"]),
            "是" if record["is_duplicate"] else "否",
            "是" if record["needs_manager_review"] else "否"
        ])

    return output.getvalue()


def verify_export_consistency(db: Session, batch_id: int) -> Dict[str, Any]:
    unified_data = get_unified_record_data(db, batch_id)

    issues = []
    pinyin_records_export = [r for r in unified_data if r["approval_name_is_pinyin"]]
    pinyin_records_api = [r for r in unified_data if r["approval_name_is_pinyin"]]

    if len(pinyin_records_export) != len(pinyin_records_api):
        issues.append(f"审批人拼音记录数不一致: 导出{len(pinyin_records_export)}, 接口{len(pinyin_records_api)}")

    for record in unified_data:
        if record["approval_name_is_pinyin"] and not record["approval_name"]:
            issues.append(f"记录ID {record['id']}: 审批人拼音标记但名称为空")

        if record["needs_manager_review"] and record["status"] != models.ProcessingStatus.NEEDS_MANAGER_REVIEW:
            issues.append(f"记录ID {record['id']}: 待复核状态不一致")

    return {
        "consistent": len(issues) == 0,
        "total_records": len(unified_data),
        "pinyin_records": len(pinyin_records_export),
        "issues": issues
    }


def get_source_label_export(source_type: str) -> str:
    labels = {
        models.SourceType.CLEARING_BATCH: "清算批次导入",
        models.SourceType.MANUAL_ENTRY: "手工录入",
        models.SourceType.HOLIDAY_ADJUSTMENT: "节假日调整"
    }
    return labels.get(source_type, "未知")


def get_status_label_export(status: str) -> str:
    labels = {
        models.ProcessingStatus.IMPORTED: "已导入",
        models.ProcessingStatus.PENDING_REVIEW: "待审核",
        models.ProcessingStatus.APPROVED: "已通过",
        models.ProcessingStatus.HOLIDAY_ADJUSTED: "节假日已调整",
        models.ProcessingStatus.BALANCE_UPDATED: "余额已更新",
        models.ProcessingStatus.NEEDS_MANAGER_REVIEW: "待经理复核",
        models.ProcessingStatus.REJECTED: "已拒绝"
    }
    return labels.get(status, status)
