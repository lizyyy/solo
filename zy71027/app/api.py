from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc
from typing import List, Optional
from datetime import datetime
import io
import csv
from openpyxl import Workbook

from .database import get_db
from . import models, schemas, services
from .models import HandoverStatus, PrescriptionStatus, DiscrepancyType

router = APIRouter(prefix="/api/v1")

@router.post("/drugs", response_model=schemas.Drug, tags=["基础数据"])
def create_drug(drug: schemas.DrugCreate, db: Session = Depends(get_db)):
    db_drug = db.query(models.Drug).filter(models.Drug.drug_code == drug.drug_code).first()
    if db_drug:
        raise HTTPException(status_code=400, detail="药品编码已存在")
    
    db_drug = models.Drug(**drug.model_dump())
    db.add(db_drug)
    db.commit()
    db.refresh(db_drug)
    return db_drug

@router.get("/drugs", response_model=List[schemas.Drug], tags=["基础数据"])
def list_drugs(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return db.query(models.Drug).offset(skip).limit(limit).all()

@router.post("/batches", response_model=schemas.DrugBatch, tags=["基础数据"])
def create_batch(batch: schemas.DrugBatchCreate, db: Session = Depends(get_db)):
    db_batch = db.query(models.DrugBatch).filter(models.DrugBatch.batch_no == batch.batch_no).first()
    if db_batch:
        raise HTTPException(status_code=400, detail="批号已存在")
    
    db_batch = models.DrugBatch(**batch.model_dump())
    db.add(db_batch)
    db.commit()
    db.refresh(db_batch)
    return db_batch

@router.get("/batches", response_model=List[schemas.DrugBatch], tags=["基础数据"])
def list_batches(drug_id: Optional[int] = None, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(models.DrugBatch)
    if drug_id:
        query = query.filter(models.DrugBatch.drug_id == drug_id)
    return query.offset(skip).limit(limit).all()

@router.post("/prescriptions", response_model=schemas.Prescription, tags=["处方管理"])
def create_prescription(prescription: schemas.PrescriptionCreate, db: Session = Depends(get_db)):
    db_prescription = db.query(models.Prescription).filter(
        models.Prescription.prescription_no == prescription.prescription_no
    ).first()
    if db_prescription:
        raise HTTPException(status_code=400, detail="处方号已存在")
    
    db_prescription = models.Prescription(**prescription.model_dump())
    db.add(db_prescription)
    db.commit()
    db.refresh(db_prescription)
    return db_prescription

@router.post("/prescriptions/{prescription_id}/verify", response_model=schemas.Prescription, tags=["处方管理"])
def verify_prescription(prescription_id: int, verify_data: schemas.PrescriptionVerify, db: Session = Depends(get_db)):
    prescription = db.query(models.Prescription).filter(models.Prescription.id == prescription_id).first()
    if not prescription:
        raise HTTPException(status_code=404, detail="处方不存在")
    
    services.verify_prescription(db, prescription_id, verify_data.verified_by)
    db.refresh(prescription)
    return prescription

@router.get("/prescriptions", response_model=List[schemas.Prescription], tags=["处方管理"])
def list_prescriptions(status: Optional[str] = None, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(models.Prescription)
    if status:
        query = query.filter(models.Prescription.status == status)
    return query.order_by(desc(models.Prescription.created_at)).offset(skip).limit(limit).all()

@router.post("/inventory", response_model=schemas.InventoryRecord, tags=["库存管理"])
def create_inventory_record(record: schemas.InventoryRecordCreate, db: Session = Depends(get_db)):
    db_record = db.query(models.InventoryRecord).filter(models.InventoryRecord.record_no == record.record_no).first()
    if db_record:
        raise HTTPException(status_code=400, detail="库存记录号已存在")
    
    db_record = models.InventoryRecord(**record.model_dump())
    db.add(db_record)
    db.commit()
    db.refresh(db_record)
    return db_record

@router.get("/inventory", response_model=List[schemas.InventoryRecord], tags=["库存管理"])
def list_inventory(batch_id: Optional[int] = None, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(models.InventoryRecord)
    if batch_id:
        query = query.filter(models.InventoryRecord.batch_id == batch_id)
    return query.order_by(desc(models.InventoryRecord.created_at)).offset(skip).limit(limit).all()

@router.post("/handovers", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def create_handover(handover: schemas.HandoverRecordCreate, db: Session = Depends(get_db)):
    db_handover = db.query(models.HandoverRecord).filter(
        models.HandoverRecord.handover_no == handover.handover_no
    ).first()
    if db_handover:
        raise HTTPException(status_code=400, detail="交接单号已存在")
    
    handover_data = handover.model_dump(exclude={"items"})
    db_handover = models.HandoverRecord(**handover_data)
    db.add(db_handover)
    db.flush()
    
    for item in handover.items:
        is_consistent, issues, qty_info = services.check_consistency(db, item)
        
        db_item = models.HandoverItem(
            handover_id=db_handover.id,
            **item.model_dump(),
            is_consistent=is_consistent
        )
        db.add(db_item)
        
        if not is_consistent:
            for issue in issues:
                if "未核验" in issue:
                    disc_type = DiscrepancyType.PRESCRIPTION_NOT_VERIFIED
                elif "不足" in issue or "数量" in issue:
                    disc_type = DiscrepancyType.QUANTITY_MISMATCH
                else:
                    disc_type = DiscrepancyType.BATCH_MISMATCH
                
                services.create_discrepancy_report(
                    db, db_handover.id, item, disc_type, issue, qty_info
                )
    
    db.commit()
    db.refresh(db_handover)
    
    services.log_operation(db, db_handover.id, "创建交接单", handover.from_nurse, 
                          f"创建交接单，包含{len(handover.items)}条明细")
    
    return db_handover

@router.post("/handovers/{handover_id}/submit", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def submit_handover(handover_id: int, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status != HandoverStatus.DRAFT and handover.status != HandoverStatus.REJECTED:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能提交")
    
    has_discrepancies = db.query(models.DiscrepancyReport).filter(
        models.DiscrepancyReport.handover_id == handover_id,
        models.DiscrepancyReport.status == "待处理"
    ).first()
    
    target_status = HandoverStatus.CONFLICT if has_discrepancies else HandoverStatus.SUBMITTED
    
    if not services.transition_status(db, handover, target_status, handover.from_nurse, "提交交接单"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.post("/handovers/{handover_id}/first-sign", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def first_sign_handover(handover_id: int, sign_data: schemas.HandoverFirstSign, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status != HandoverStatus.SUBMITTED and handover.status != HandoverStatus.CONFLICT:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能第一签")
    
    sign_time = datetime.now()
    handover.first_signature = sign_data.signature
    handover.first_signed_at = sign_time
    handover.first_sign_remark = sign_data.sign_remark
    
    time_abnormal = services.check_sign_time_abnormal(handover.handover_time, sign_time)
    if sign_data.is_late_sign or time_abnormal:
        handover.is_late_sign = True
        handover.sign_time_abnormal = True
        services.create_discrepancy_report(
            db, handover.id,
            schemas.HandoverItemCreate(
                drug_name="补签时间异常",
                batch_no="N/A",
                handover_quantity=0
            ),
            DiscrepancyType.TIME_ABNORMAL,
            f"第一签时间异常: {sign_data.sign_remark or '补签'}",
            {}
        )
    
    if not services.transition_status(db, handover, HandoverStatus.FIRST_SIGNED, sign_data.signature, 
                                      f"第一签完成{' (补签)' if sign_data.is_late_sign else ''}"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.post("/handovers/{handover_id}/second-sign", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def second_sign_handover(handover_id: int, sign_data: schemas.HandoverSecondSign, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status != HandoverStatus.FIRST_SIGNED:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能第二签")
    
    if sign_data.signature == handover.first_signature:
        raise HTTPException(status_code=400, detail="双人签名不能为同一人")
    
    sign_time = datetime.now()
    handover.second_signature = sign_data.signature
    handover.second_signed_at = sign_time
    handover.second_sign_remark = sign_data.sign_remark
    
    time_abnormal = services.check_sign_time_abnormal(handover.handover_time, sign_time)
    if sign_data.is_late_sign or time_abnormal:
        handover.is_late_sign = True
        handover.sign_time_abnormal = True
        services.create_discrepancy_report(
            db, handover.id,
            schemas.HandoverItemCreate(
                drug_name="补签时间异常",
                batch_no="N/A",
                handover_quantity=0
            ),
            DiscrepancyType.TIME_ABNORMAL,
            f"第二签时间异常: {sign_data.sign_remark or '补签'}",
            {}
        )
    
    if handover.first_signed_at:
        time_diff = (sign_time - handover.first_signed_at).total_seconds()
        if time_diff < 10:
            handover.sign_time_abnormal = True
            services.create_discrepancy_report(
                db, handover.id,
                schemas.HandoverItemCreate(
                    drug_name="双签间隔异常",
                    batch_no="N/A",
                    handover_quantity=0
                ),
                DiscrepancyType.TIME_ABNORMAL,
                f"双签间隔过短: {time_diff:.1f}秒",
                {}
            )
    
    if not services.transition_status(db, handover, HandoverStatus.SECOND_SIGNED, sign_data.signature, 
                                      f"第二签完成{' (补签)' if sign_data.is_late_sign else ''}"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.post("/handovers/{handover_id}/verify", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def verify_handover(handover_id: int, verify_data: schemas.HandoverVerify, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status not in [HandoverStatus.SECOND_SIGNED, HandoverStatus.MANUAL_FIXED]:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能核验")
    
    handover.reviewer = verify_data.reviewer
    handover.reviewed_at = datetime.now()
    
    for item in handover.items:
        if item.batch_id:
            services.deduct_inventory(db, item.batch_id, item.handover_quantity, 
                                      verify_data.reviewer, handover.handover_no)
        if item.prescription_id:
            prescription = db.query(models.Prescription).filter(
                models.Prescription.id == item.prescription_id
            ).first()
            if prescription:
                prescription.status = PrescriptionStatus.USED
    
    for discrepancy in handover.discrepancies:
        discrepancy.status = "已处理"
        discrepancy.handled_by = verify_data.reviewer
        discrepancy.handled_at = datetime.now()
        discrepancy.handle_result = "核验通过，完成交接"
    
    if not services.transition_status(db, handover, HandoverStatus.VERIFIED, verify_data.reviewer, "核验通过"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    db.commit()
    db.refresh(handover)
    return handover

@router.post("/handovers/{handover_id}/reject", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def reject_handover(handover_id: int, reject_data: schemas.HandoverReject, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status not in [HandoverStatus.SUBMITTED, HandoverStatus.FIRST_SIGNED, HandoverStatus.SECOND_SIGNED, HandoverStatus.CONFLICT]:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能驳回")
    
    handover.reviewer = reject_data.reviewer
    handover.reject_reason = reject_data.reject_reason
    handover.reviewed_at = datetime.now()
    
    if not services.transition_status(db, handover, HandoverStatus.REJECTED, reject_data.reviewer, 
                                      f"驳回原因: {reject_data.reject_reason}"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.post("/handovers/{handover_id}/withdraw", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def withdraw_handover(handover_id: int, withdraw_data: schemas.HandoverWithdraw, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status not in [HandoverStatus.SUBMITTED, HandoverStatus.FIRST_SIGNED, HandoverStatus.SECOND_SIGNED]:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能撤回")
    
    if not services.transition_status(db, handover, HandoverStatus.WITHDRAWN, withdraw_data.operator, 
                                      f"撤回原因: {withdraw_data.reason}"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.post("/handovers/{handover_id}/manual-fix", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def manual_fix_handover(handover_id: int, fix_data: schemas.HandoverManualFix, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status not in [HandoverStatus.DRAFT, HandoverStatus.CONFLICT, HandoverStatus.REJECTED]:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能人工修正")
    
    for item in handover.items:
        db.delete(item)
    
    for discrepancy in handover.discrepancies:
        discrepancy.status = "已处理"
        discrepancy.handled_by = fix_data.operator
        discrepancy.handled_at = datetime.now()
        discrepancy.handle_result = f"人工修正: {fix_data.fix_description}"
    
    for item in fix_data.items:
        is_consistent, issues, qty_info = services.check_consistency(db, item)
        
        db_item = models.HandoverItem(
            handover_id=handover.id,
            **item.model_dump(),
            is_consistent=is_consistent
        )
        db.add(db_item)
        
        if not is_consistent:
            for issue in issues:
                if "未核验" in issue:
                    disc_type = DiscrepancyType.PRESCRIPTION_NOT_VERIFIED
                elif "不足" in issue or "数量" in issue:
                    disc_type = DiscrepancyType.QUANTITY_MISMATCH
                else:
                    disc_type = DiscrepancyType.BATCH_MISMATCH
                
                services.create_discrepancy_report(
                    db, handover.id, item, disc_type, issue, qty_info
                )
    
    if not services.transition_status(db, handover, HandoverStatus.MANUAL_FIXED, fix_data.operator, 
                                      f"人工修正: {fix_data.fix_description}"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    db.commit()
    db.refresh(handover)
    return handover

@router.post("/handovers/{handover_id}/complete", response_model=schemas.HandoverRecord, tags=["交接班管理"])
def complete_handover(handover_id: int, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    if handover.status != HandoverStatus.VERIFIED:
        raise HTTPException(status_code=400, detail=f"当前状态[{handover.status}]不能完成")
    
    if not services.transition_status(db, handover, HandoverStatus.COMPLETED, handover.reviewer, "完成交接"):
        raise HTTPException(status_code=400, detail="状态流转失败")
    
    return handover

@router.get("/handovers", response_model=List[schemas.HandoverRecord], tags=["交接班管理"])
def list_handovers(status: Optional[str] = None, shift_type: Optional[str] = None, 
                   skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(models.HandoverRecord)
    if status:
        query = query.filter(models.HandoverRecord.status == status)
    if shift_type:
        query = query.filter(models.HandoverRecord.shift_type == shift_type)
    return query.order_by(desc(models.HandoverRecord.created_at)).offset(skip).limit(limit).all()

@router.get("/handovers/{handover_id}", response_model=schemas.HandoverDetail, tags=["交接班管理"])
def get_handover_detail(handover_id: int, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    logs = db.query(models.OperationLog).filter(
        models.OperationLog.handover_id == handover_id
    ).order_by(models.OperationLog.created_at).all()
    
    previous_reject_reason = None
    if handover.previous_handover_id:
        previous = db.query(models.HandoverRecord).filter(
            models.HandoverRecord.id == handover.previous_handover_id
        ).first()
        if previous:
            previous_reject_reason = previous.reject_reason
    
    result = schemas.HandoverDetail.model_validate(handover)
    result.operation_logs = [schemas.OperationLog.model_validate(log) for log in logs]
    result.previous_reject_reason = previous_reject_reason
    
    return result

@router.get("/handovers/{handover_id}/trace", tags=["交接班管理"])
def trace_handover(handover_id: int, db: Session = Depends(get_db)):
    handover = db.query(models.HandoverRecord).filter(models.HandoverRecord.id == handover_id).first()
    if not handover:
        raise HTTPException(status_code=404, detail="交接单不存在")
    
    trace_data = {
        "handover_no": handover.handover_no,
        "current_status": handover.status,
        "status_history": [],
        "items": [],
        "discrepancies": []
    }
    
    logs = db.query(models.OperationLog).filter(
        models.OperationLog.handover_id == handover_id
    ).order_by(models.OperationLog.created_at).all()
    
    for log in logs:
        trace_data["status_history"].append({
            "operation": log.operation,
            "operator": log.operator,
            "remark": log.remark,
            "time": log.created_at
        })
    
    for item in handover.items:
        trace_data["items"].append({
            "drug_name": item.drug_name,
            "batch_no": item.batch_no,
            "prescription_no": item.prescription_no,
            "prescription_qty": item.prescription_quantity,
            "inventory_qty": item.inventory_quantity,
            "handover_qty": item.handover_quantity,
            "is_consistent": item.is_consistent
        })
    
    for disc in handover.discrepancies:
        trace_data["discrepancies"].append({
            "type": disc.discrepancy_type,
            "description": disc.description,
            "status": disc.status,
            "handle_result": disc.handle_result
        })
    
    return trace_data

@router.get("/discrepancies", response_model=List[schemas.DiscrepancyReport], tags=["差异管理"])
def list_discrepancies(status: Optional[str] = None, handover_id: Optional[int] = None,
                       skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(models.DiscrepancyReport)
    if status:
        query = query.filter(models.DiscrepancyReport.status == status)
    if handover_id:
        query = query.filter(models.DiscrepancyReport.handover_id == handover_id)
    return query.order_by(desc(models.DiscrepancyReport.created_at)).offset(skip).limit(limit).all()

@router.post("/discrepancies/{discrepancy_id}/handle", response_model=schemas.DiscrepancyReport, tags=["差异管理"])
def handle_discrepancy(discrepancy_id: int, handle_data: schemas.DiscrepancyHandle, db: Session = Depends(get_db)):
    discrepancy = db.query(models.DiscrepancyReport).filter(
        models.DiscrepancyReport.id == discrepancy_id
    ).first()
    if not discrepancy:
        raise HTTPException(status_code=404, detail="差异报告不存在")
    
    discrepancy.status = "已处理"
    discrepancy.handled_by = handle_data.handled_by
    discrepancy.handled_at = datetime.now()
    discrepancy.handle_result = handle_data.handle_result
    db.commit()
    db.refresh(discrepancy)
    
    return discrepancy

@router.get("/export/handovers", tags=["导出下载"])
def export_handovers(format: str = Query("excel", enum=["excel", "csv"]),
                     start_date: Optional[str] = None,
                     end_date: Optional[str] = None,
                     db: Session = Depends(get_db)):
    query = db.query(models.HandoverRecord)
    if start_date:
        query = query.filter(models.HandoverRecord.created_at >= start_date)
    if end_date:
        query = query.filter(models.HandoverRecord.created_at <= end_date)
    
    handovers = query.order_by(desc(models.HandoverRecord.created_at)).all()
    
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["交接单号", "班次", "交班人", "接班人", "状态", "第一签", "第二签", "审核人", "创建时间"])
        
        for h in handovers:
            writer.writerow([
                h.handover_no, h.shift_type, h.from_nurse, h.to_nurse,
                h.status, h.first_signature or "", h.second_signature or "",
                h.reviewer or "", h.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])
        
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=handovers.csv"}
        )
    else:
        output = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "交接班记录"
        
        ws.append(["交接单号", "班次", "交班人", "接班人", "状态", "第一签", "第二签", "审核人", "创建时间"])
        
        for h in handovers:
            ws.append([
                h.handover_no, h.shift_type, h.from_nurse, h.to_nurse,
                h.status, h.first_signature or "", h.second_signature or "",
                h.reviewer or "", h.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])
        
        ws2 = wb.create_sheet("明细")
        ws2.append(["交接单号", "药品名称", "批号", "处方号", "处方数量", "库存数量", "交接数量", "是否一致"])
        
        for h in handovers:
            for item in h.items:
                ws2.append([
                    h.handover_no, item.drug_name, item.batch_no, item.prescription_no or "",
                    item.prescription_quantity, item.inventory_quantity, item.handover_quantity,
                    "是" if item.is_consistent else "否"
                ])
        
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=handovers.xlsx"}
        )

@router.get("/export/discrepancies", tags=["导出下载"])
def export_discrepancies(format: str = Query("excel", enum=["excel", "csv"]),
                         status: Optional[str] = None,
                         db: Session = Depends(get_db)):
    query = db.query(models.DiscrepancyReport)
    if status:
        query = query.filter(models.DiscrepancyReport.status == status)
    
    discrepancies = query.order_by(desc(models.DiscrepancyReport.created_at)).all()
    
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["报告编号", "交接单号", "差异类型", "描述", "处方数量", "库存数量", "交接数量", "差异", "状态", "处理人", "处理结果"])
        
        for d in discrepancies:
            writer.writerow([
                d.report_no, d.handover.handover_no if d.handover else "",
                d.discrepancy_type, d.description,
                d.prescription_quantity or "", d.inventory_quantity or "",
                d.handover_quantity or "", d.difference or "",
                d.status, d.handled_by or "", d.handle_result or ""
            ])
        
        output.seek(0)
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=discrepancies.csv"}
        )
    else:
        output = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "差异报告"
        
        ws.append(["报告编号", "交接单号", "差异类型", "描述", "处方数量", "库存数量", "交接数量", "差异", "状态", "处理人", "处理结果"])
        
        for d in discrepancies:
            ws.append([
                d.report_no, d.handover.handover_no if d.handover else "",
                d.discrepancy_type, d.description,
                d.prescription_quantity or "", d.inventory_quantity or "",
                d.handover_quantity or "", d.difference or "",
                d.status, d.handled_by or "", d.handle_result or ""
            ])
        
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=discrepancies.xlsx"}
        )
