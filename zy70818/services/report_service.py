import csv
from datetime import date, datetime
from typing import List, Dict, Any, Optional
from io import StringIO, BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from models.inventory import InventoryItem
from models.recall import RecallNotice
from models.consumption import ConsumptionItem
from models.reconciliation import ReconciliationResult, Discrepancy, DiscrepancyType
from models.report import ReportData, ReportType, ReportDetailItem, ReportSummaryItem
from utils.storage import store


class ReportGenerator:
    def __init__(self):
        self.data_dir = Path(__file__).parent.parent / "data" / "reports"
        self.data_dir.mkdir(parents=True, exist_ok=True)

    def generate_report(
        self,
        reconciliation_id: str,
        report_type: ReportType,
        report_format: ReportFormat,
        generated_by: Optional[str] = None
    ) -> ReportData:
        reconciliation = store.get('reconciliation', reconciliation_id, ReconciliationResult)
        if not reconciliation:
            raise ValueError("对账任务不存在")

        title = self._generate_title(reconciliation.name, report_type)

        report = ReportData(
            reconciliation_id=reconciliation_id,
            report_type=report_type,
            report_format=report_format,
            title=title,
            generated_by=generated_by
        )

        if report_type in [ReportType.SUMMARY, ReportType.DETAIL]:
            report.summary = self._generate_summary(reconciliation)
            report.details = self._generate_details(reconciliation)

        if report_type in [ReportType.DETAIL, ReportType.DISCREPANCY]:
            report.discrepancies = [d.model_dump() for d in reconciliation.discrepancies]

        if report_type == ReportType.RECALL:
            report.details = self._generate_recall_details(reconciliation)

        if report_type == ReportType.EXPIRY:
            report.details = self._generate_expiry_details(reconciliation)

        report.statistics = self._generate_statistics(reconciliation)

        file_path = self._export_report(report, report_format)
        report.file_path = str(file_path)
        report.file_size = file_path.stat().st_size

        return report

    def _generate_title(self, recon_name: str, report_type: ReportType) -> str:
        type_names = {
            ReportType.SUMMARY: "汇总报告",
            ReportType.DETAIL: "明细报告",
            ReportType.DISCREPANCY: "差异分析报告",
            ReportType.RECALL: "召回专项报告",
            ReportType.EXPIRY: "效期专项报告"
        }
        return f"{recon_name} - {type_names.get(report_type, '报告')}"

    def _generate_summary(self, reconciliation: ReconciliationResult) -> List[ReportSummaryItem]:
        inventory = store.get_all('inventory', InventoryItem)
        discrepancies = reconciliation.discrepancies

        by_store: Dict[str, Dict] = {}
        for item in inventory:
            if item.store_name not in by_store:
                by_store[item.store_name] = {
                    'total': 0, 'normal': 0, 'recalled': 0,
                    'near_expiry': 0, 'expired': 0, 'discrepancy': 0
                }
            by_store[item.store_name]['total'] += item.quantity
            if item.is_recalled:
                by_store[item.store_name]['recalled'] += item.quantity
            elif item.status == '已过期':
                by_store[item.store_name]['expired'] += item.quantity
            elif item.status == '近效期':
                by_store[item.store_name]['near_expiry'] += item.quantity
            else:
                by_store[item.store_name]['normal'] += item.quantity

        for d in discrepancies:
            if d.store_name in by_store:
                by_store[d.store_name]['discrepancy'] += 1

        items = []
        for store_name, data in by_store.items():
            items.append(ReportSummaryItem(
                category=store_name,
                total_quantity=data['total'],
                normal_quantity=data['normal'],
                recalled_quantity=data['recalled'],
                near_expiry_quantity=data['near_expiry'],
                expired_quantity=data['expired'],
                discrepancy_quantity=data['discrepancy']
            ))

        return items

    def _generate_details(self, reconciliation: ReconciliationResult) -> List[ReportDetailItem]:
        inventory = store.get_all('inventory', InventoryItem)
        consumption = store.get_all('consumption', ConsumptionItem)

        details = []
        for inv_item in inventory:
            consumed = sum(c.quantity for c in consumption
                          if c.batch_number == inv_item.batch_number and c.store_name == inv_item.store_name)
            transfer_in = sum(c.quantity for c in consumption
                         if c.batch_number == inv_item.batch_number and c.transfer_to_store == inv_item.store_name)
            transfer_out = sum(c.quantity for c in consumption
                          if c.batch_number == inv_item.batch_number and c.transfer_from_store == inv_item.store_name)

            expected = inv_item.quantity - consumed + transfer_in - transfer_out
            diff = inv_item.quantity - expected

            details.append(ReportDetailItem(
                batch_number=inv_item.batch_number,
                material_name=inv_item.material_name,
                material_type=inv_item.material_type,
                specification=inv_item.specification,
                store_name=inv_item.store_name,
                initial_quantity=inv_item.quantity,
                consumed_quantity=consumed,
                transfer_in_quantity=transfer_in,
                transfer_out_quantity=transfer_out,
                current_quantity=inv_item.quantity,
                expected_quantity=expected,
                diff_quantity=diff,
                status=inv_item.status,
                is_recalled=inv_item.is_recalled,
                expiry_date=inv_item.expiry_date
            ))

        return details

    def _generate_recall_details(self, reconciliation: ReconciliationResult) -> List[ReportDetailItem]:
        all_details = self._generate_details(reconciliation)
        return [d for d in all_details if d.is_recalled]

    def _generate_expiry_details(self, reconciliation: ReconciliationResult) -> List[ReportDetailItem]:
        all_details = self._generate_details(reconciliation)
        return [d for d in all_details if d.status in ['已过期', '近效期']]

    def _generate_statistics(self, reconciliation: ReconciliationResult) -> Dict[str, Any]:
        stats = {
            'reconciliation_name': reconciliation.name,
            'period': f"{reconciliation.start_date} ~ {reconciliation.end_date}",
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'total_inventory': reconciliation.total_inventory_count,
            'total_consumption': reconciliation.total_consumption_count,
            'total_recall': reconciliation.total_recall_count,
            'total_discrepancies': reconciliation.discrepancy_count,
            'unresolved_discrepancies': reconciliation.unresolved_discrepancy_count,
            'by_discrepancy_type': {
                'recall': reconciliation.recalled_batch_count,
                'near_expiry': reconciliation.near_expiry_count,
                'expired': reconciliation.expired_count,
                'transfer': reconciliation.transfer_count,
                'others': reconciliation.discrepancy_count - reconciliation.recalled_batch_count - reconciliation.near_expiry_count - reconciliation.expired_count - reconciliation.transfer_count
            }
        }
        return stats

    def _export_report(self, report: ReportData, report_format: ReportFormat) -> Path:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"{report.report_type.value}_{timestamp}.{report_format.value}"
        file_path = self.data_dir / filename

        if report_format == ReportFormat.EXCEL:
            self._export_excel(report, file_path)
        elif report_format == ReportFormat.CSV:
            self._export_csv(report, file_path)

        return file_path

    def _export_excel(self, report: ReportData, file_path: Path):
        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = "汇总"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        ws_summary['A1'] = report.title
        ws_summary['A1'].font = Font(bold=True, size=14)
        ws_summary.merge_cells('A1:H1')

        ws_summary['A3'] = "统计信息"
        ws_summary['A3'].font = Font(bold=True)

        stats = report.statistics
        row = 4
        for key, value in stats.items():
            if isinstance(value, dict):
                ws_summary[f'A{row}'] = key
                ws_summary[f'A{row}'].font = Font(bold=True)
                row += 1
                for k, v in value.items():
                    ws_summary[f'B{row}'] = k
                    ws_summary[f'C{row}'] = v
                    row += 1
            else:
                ws_summary[f'A{row}'] = key
                ws_summary[f'B{row}'] = value
                row += 1

        if report.summary:
            ws_detail = wb.create_sheet("汇总明细")
            headers = ["门店", "总数量", "正常数量", "召回数量", "近效期数量", "过期数量", "差异数量"]
            for col, header in enumerate(headers, 1):
                cell = ws_detail.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            for row_idx, item in enumerate(report.summary, 2):
                    ws_detail.cell(row=row_idx, column=1, value=item.category)
                    ws_detail.cell(row=row_idx, column=2, value=item.total_quantity)
                    ws_detail.cell(row=row_idx, column=3, value=item.normal_quantity)
                    ws_detail.cell(row=row_idx, column=4, value=item.recalled_quantity)
                    ws_detail.cell(row=row_idx, column=5, value=item.near_expiry_quantity)
                    ws_detail.cell(row=row_idx, column=6, value=item.expired_quantity)
                    ws_detail.cell(row=row_idx, column=7, value=item.discrepancy_quantity)

        if report.details:
            ws_detail = wb.create_sheet("明细")
            headers = ["批号", "物料名称", "物料类型", "规格", "门店", "期初", "消耗", "调入", "调出", "当前", "理论", "差异", "状态", "是否召回", "有效期"]
            for col, header in enumerate(headers, 1):
                cell = ws_detail.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            for row_idx, item in enumerate(report.details, 2):
                    ws_detail.cell(row=row_idx, column=1, value=item.batch_number)
                    ws_detail.cell(row=row_idx, column=2, value=item.material_name)
                    ws_detail.cell(row=row_idx, column=3, value=item.material_type)
                    ws_detail.cell(row=row_idx, column=4, value=item.specification)
                    ws_detail.cell(row=row_idx, column=5, value=item.store_name)
                    ws_detail.cell(row=row_idx, column=6, value=item.initial_quantity)
                    ws_detail.cell(row=row_idx, column=7, value=item.consumed_quantity)
                    ws_detail.cell(row=row_idx, column=8, value=item.transfer_in_quantity)
                    ws_detail.cell(row=row_idx, column=9, value=item.transfer_out_quantity)
                    ws_detail.cell(row=row_idx, column=10, value=item.current_quantity)
                    ws_detail.cell(row=row_idx, column=11, value=item.expected_quantity)
                    ws_detail.cell(row=row_idx, column=12, value=item.diff_quantity)
                    ws_detail.cell(row=row_idx, column=13, value=item.status)
                    ws_detail.cell(row=row_idx, column=14, value="是" if item.is_recalled else "否")
                    ws_detail.cell(row=row_idx, column=15, value=item.expiry_date.strftime('%Y-%m-%d') if item.expiry_date else "")

        if report.discrepancies:
            ws_disc = wb.create_sheet("差异")
            headers = ["类型", "批号", "物料名称", "门店", "描述", "解释", "期望值", "实际值", "数量差异"]
            for col, header in enumerate(headers, 1):
                cell = ws_disc.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            for row_idx, disc in enumerate(report.discrepancies, 2):
                    ws_disc.cell(row=row_idx, column=1, value=disc.get('type', ''))
                    ws_disc.cell(row=row_idx, column=2, value=disc.get('batch_number', ''))
                    ws_disc.cell(row=row_idx, column=3, value=disc.get('material_name', ''))
                    ws_disc.cell(row=row_idx, column=4, value=disc.get('store_name', ''))
                    ws_disc.cell(row=row_idx, column=5, value=disc.get('description', ''))
                    ws_disc.cell(row=row_idx, column=6, value=disc.get('explanation', ''))
                    ws_disc.cell(row=row_idx, column=7, value=str(disc.get('expected_value', '')))
                    ws_disc.cell(row=row_idx, column=8, value=str(disc.get('actual_value', '')))
                    ws_disc.cell(row=row_idx, column=9, value=disc.get('quantity_diff', 0))

        wb.save(file_path)

    def _export_csv(self, report: ReportData, file_path: Path):
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow([report.title])
            writer.writerow([])

            writer.writerow(["统计信息"])
            for key, value in report.statistics.items():
                if isinstance(value, dict):
                    writer.writerow([key])
                    for k, v in value.items():
                        writer.writerow(["", k, v])
                else:
                    writer.writerow([key, value])

            writer.writerow([])

            if report.details:
                writer.writerow(["明细数据"])
                writer.writerow([
                    "批号", "物料名称", "物料类型", "规格", "门店",
                    "期初", "消耗", "调入", "调出", "当前", "理论", "差异",
                    "状态", "是否召回", "有效期"
                ])
                for item in report.details:
                    writer.writerow([
                        item.batch_number,
                        item.material_name,
                        item.material_type,
                        item.specification,
                        item.store_name,
                        item.initial_quantity,
                        item.consumed_quantity,
                        item.transfer_in_quantity,
                        item.transfer_out_quantity,
                        item.current_quantity,
                        item.expected_quantity,
                        item.diff_quantity,
                        item.status,
                        "是" if item.is_recalled else "否",
                        item.expiry_date.strftime('%Y-%m-%d') if item.expiry_date else ""
                    ])

    def get_report_file(self, file_path: str) -> BytesIO:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError("报告文件不存在")

        buffer = BytesIO()
        with open(path, 'rb') as f:
            buffer.write(f.read())
        buffer.seek(0)
        return buffer


report_generator = ReportGenerator()
