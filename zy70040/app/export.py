from io import BytesIO
from datetime import datetime, timedelta
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import (
    DepositOrder, ScanRecord, DamageRecord, RefundOrder,
    OperationHistory, ReconciliationBatch, PackageType
)


PACKAGE_TYPE_NAMES = {
    PackageType.CYCLE_BOX: "循环箱",
    PackageType.THERMAL_BAG: "保温袋",
    PackageType.PALLET: "托盘",
}

DEPOSIT_STATUS_NAMES = {
    "paid": "已缴纳",
    "partial_returned": "部分归还",
    "full_returned": "全部归还",
    "closed": "已关闭",
}

REFUND_STATUS_NAMES = {
    "pending": "待处理",
    "processing": "处理中",
    "success": "已成功",
    "failed": "已失败",
}


def _format_datetime(dt: Optional[datetime]) -> str:
    if not dt:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _set_header_style(ws, headers: List[str]):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border


def _set_data_style(ws, row_count: int, col_count: int):
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, row_count + 1):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = center_alignment


def _auto_fit_columns(ws):
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width


def export_deposit_orders(db: Session, orders: List[DepositOrder]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "押金单汇总"

    headers = [
        "押金单号", "客户ID", "客户名称",
        "循环箱数量", "循环箱押金",
        "保温袋数量", "保温袋押金",
        "托盘数量", "托盘押金",
        "总押金", "已退还", "扣款金额", "可退押金",
        "状态", "创建时间", "最后更新", "备注",
    ]
    _set_header_style(ws, headers)

    for idx, order in enumerate(orders, 2):
        ws.cell(row=idx, column=1, value=order.order_no)
        ws.cell(row=idx, column=2, value=order.customer_id)
        ws.cell(row=idx, column=3, value=order.customer_name)

        ws.cell(row=idx, column=4, value=order.cycle_box_count)
        ws.cell(row=idx, column=5, value=order.cycle_box_deposit)

        ws.cell(row=idx, column=6, value=order.thermal_bag_count)
        ws.cell(row=idx, column=7, value=order.thermal_bag_deposit)

        ws.cell(row=idx, column=8, value=order.pallet_count)
        ws.cell(row=idx, column=9, value=order.pallet_deposit)

        ws.cell(row=idx, column=10, value=order.total_deposit)
        ws.cell(row=idx, column=11, value=order.refunded_amount)
        ws.cell(row=idx, column=12, value=order.total_deduction)
        ws.cell(row=idx, column=13, value=order.refundable_amount)

        ws.cell(row=idx, column=14, value=DEPOSIT_STATUS_NAMES.get(order.status, order.status))
        ws.cell(row=idx, column=15, value=_format_datetime(order.created_at))
        ws.cell(row=idx, column=16, value=_format_datetime(order.updated_at))
        ws.cell(row=idx, column=17, value=order.remark or "")

    _set_data_style(ws, len(orders) + 1, len(headers))
    _auto_fit_columns(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_scan_records(db: Session, records: List[ScanRecord]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "扫码归还记录"

    headers = [
        "扫码编号", "押金单号", "包材类型", "包材编码",
        "数量", "操作人ID", "操作人",
        "扫码时间", "创建时间",
        "是否撤回", "撤回时间", "撤回人", "撤回原因", "备注",
    ]
    _set_header_style(ws, headers)

    for idx, record in enumerate(records, 2):
        ws.cell(row=idx, column=1, value=record.scan_no)
        ws.cell(row=idx, column=2, value=record.deposit_order.order_no)
        ws.cell(row=idx, column=3, value=PACKAGE_TYPE_NAMES.get(record.package_type, record.package_type))
        ws.cell(row=idx, column=4, value=record.package_code)
        ws.cell(row=idx, column=5, value=record.quantity)
        ws.cell(row=idx, column=6, value=record.operator_id)
        ws.cell(row=idx, column=7, value=record.operator_name)
        ws.cell(row=idx, column=8, value=_format_datetime(record.scan_time))
        ws.cell(row=idx, column=9, value=_format_datetime(record.created_at))
        ws.cell(row=idx, column=10, value="是" if record.is_reversed else "否")
        ws.cell(row=idx, column=11, value=_format_datetime(record.reversed_at))
        ws.cell(row=idx, column=12, value=record.reversed_by or "")
        ws.cell(row=idx, column=13, value=record.reverse_reason or "")
        ws.cell(row=idx, column=14, value=record.remark or "")

    _set_data_style(ws, len(records) + 1, len(headers))
    _auto_fit_columns(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_damage_records(db: Session, records: List[DamageRecord]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "损坏扣减记录"

    headers = [
        "损坏编号", "押金单号", "包材类型", "包材编码",
        "数量", "扣减金额", "损坏等级", "损坏描述",
        "操作人ID", "操作人", "创建时间",
        "是否撤回", "撤回时间", "撤回人", "撤回原因", "备注",
    ]
    _set_header_style(ws, headers)

    for idx, record in enumerate(records, 2):
        ws.cell(row=idx, column=1, value=record.damage_no)
        ws.cell(row=idx, column=2, value=record.deposit_order.order_no)
        ws.cell(row=idx, column=3, value=PACKAGE_TYPE_NAMES.get(record.package_type, record.package_type))
        ws.cell(row=idx, column=4, value=record.package_code)
        ws.cell(row=idx, column=5, value=record.quantity)
        ws.cell(row=idx, column=6, value=record.deduction_amount)
        ws.cell(row=idx, column=7, value=record.damage_level or "")
        ws.cell(row=idx, column=8, value=record.damage_description or "")
        ws.cell(row=idx, column=9, value=record.operator_id)
        ws.cell(row=idx, column=10, value=record.operator_name)
        ws.cell(row=idx, column=11, value=_format_datetime(record.created_at))
        ws.cell(row=idx, column=12, value="是" if record.is_reversed else "否")
        ws.cell(row=idx, column=13, value=_format_datetime(record.reversed_at))
        ws.cell(row=idx, column=14, value=record.reversed_by or "")
        ws.cell(row=idx, column=15, value=record.reverse_reason or "")
        ws.cell(row=idx, column=16, value=record.remark or "")

    _set_data_style(ws, len(records) + 1, len(headers))
    _auto_fit_columns(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_refund_orders(db: Session, orders: List[RefundOrder]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "退款订单"

    headers = [
        "退款单号", "押金单号",
        "循环箱退款", "保温袋退款", "托盘退款", "退款总金额",
        "退款方式", "交易号", "状态",
        "操作人ID", "操作人", "创建时间",
        "处理时间", "完成时间", "备注",
    ]
    _set_header_style(ws, headers)

    for idx, order in enumerate(orders, 2):
        ws.cell(row=idx, column=1, value=order.refund_no)
        ws.cell(row=idx, column=2, value=order.deposit_order.order_no)
        ws.cell(row=idx, column=3, value=order.cycle_box_refund)
        ws.cell(row=idx, column=4, value=order.thermal_bag_refund)
        ws.cell(row=idx, column=5, value=order.pallet_refund)
        ws.cell(row=idx, column=6, value=order.refund_amount)
        ws.cell(row=idx, column=7, value=order.refund_method or "")
        ws.cell(row=idx, column=8, value=order.transaction_id or "")
        ws.cell(row=idx, column=9, value=REFUND_STATUS_NAMES.get(order.status, order.status))
        ws.cell(row=idx, column=10, value=order.operator_id)
        ws.cell(row=idx, column=11, value=order.operator_name)
        ws.cell(row=idx, column=12, value=_format_datetime(order.created_at))
        ws.cell(row=idx, column=13, value=_format_datetime(order.processed_at))
        ws.cell(row=idx, column=14, value=_format_datetime(order.completed_at))
        ws.cell(row=idx, column=15, value=order.remark or "")

    _set_data_style(ws, len(orders) + 1, len(headers))
    _auto_fit_columns(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_operation_history(db: Session, records: List[OperationHistory]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "操作历史"

    headers = [
        "操作类型", "操作描述", "实体类型", "实体ID", "实体编号",
        "操作人ID", "操作人", "是否撤回操作", "关联撤回ID",
        "操作时间", "备注",
    ]
    _set_header_style(ws, headers)

    for idx, record in enumerate(records, 2):
        ws.cell(row=idx, column=1, value=record.operation_type)
        ws.cell(row=idx, column=2, value=record.operation_desc)
        ws.cell(row=idx, column=3, value=record.entity_type or "")
        ws.cell(row=idx, column=4, value=record.entity_id or "")
        ws.cell(row=idx, column=5, value=record.entity_no or "")
        ws.cell(row=idx, column=6, value=record.operator_id)
        ws.cell(row=idx, column=7, value=record.operator_name)
        ws.cell(row=idx, column=8, value="是" if record.is_reverse else "否")
        ws.cell(row=idx, column=9, value=record.reverse_related_id or "")
        ws.cell(row=idx, column=10, value=_format_datetime(record.operation_time))
        ws.cell(row=idx, column=11, value=record.remark or "")

    _set_data_style(ws, len(records) + 1, len(headers))
    _auto_fit_columns(ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_reconciliation_batch(
    db: Session,
    batch: ReconciliationBatch,
) -> bytes:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "对账汇总"

    headers1 = [
        "对账批次号", "开始时间", "结束时间",
        "押金单数量", "总押金", "已退款", "已扣款", "剩余押金",
        "扫码匹配数", "扫码不匹配数",
        "退款匹配数", "退款不匹配数",
        "操作人ID", "操作人", "对账时间",
    ]
    _set_header_style(ws1, headers1)

    ws1.cell(row=2, column=1, value=batch.batch_no)
    ws1.cell(row=2, column=2, value=_format_datetime(batch.start_date))
    ws1.cell(row=2, column=3, value=_format_datetime(batch.end_date))
    ws1.cell(row=2, column=4, value=batch.total_orders)
    ws1.cell(row=2, column=5, value=batch.total_deposit)
    ws1.cell(row=2, column=6, value=batch.total_refund)
    ws1.cell(row=2, column=7, value=batch.total_deduction)
    ws1.cell(row=2, column=8, value=batch.remaining_deposit)
    ws1.cell(row=2, column=9, value=batch.scan_matched)
    ws1.cell(row=2, column=10, value=batch.scan_unmatched)
    ws1.cell(row=2, column=11, value=batch.refund_matched)
    ws1.cell(row=2, column=12, value=batch.refund_unmatched)
    ws1.cell(row=2, column=13, value=batch.operator_id)
    ws1.cell(row=2, column=14, value=batch.operator_name)
    ws1.cell(row=2, column=15, value=_format_datetime(batch.created_at))

    _set_data_style(ws1, 2, len(headers1))
    _auto_fit_columns(ws1)

    ws2 = wb.create_sheet("押金单明细")
    headers2 = [
        "押金单号", "客户ID", "客户名称",
        "循环箱数量", "循环箱押金",
        "保温袋数量", "保温袋押金",
        "托盘数量", "托盘押金",
        "总押金", "已退还", "扣款金额", "可退押金",
        "状态", "创建时间",
    ]
    _set_header_style(ws2, headers2)

    orders = db.query(DepositOrder).filter(
        DepositOrder.is_deleted == False,
        DepositOrder.created_at >= batch.start_date,
        DepositOrder.created_at < batch.end_date,
    ).all()

    for idx, order in enumerate(orders, 2):
        ws2.cell(row=idx, column=1, value=order.order_no)
        ws2.cell(row=idx, column=2, value=order.customer_id)
        ws2.cell(row=idx, column=3, value=order.customer_name)
        ws2.cell(row=idx, column=4, value=order.cycle_box_count)
        ws2.cell(row=idx, column=5, value=order.cycle_box_deposit)
        ws2.cell(row=idx, column=6, value=order.thermal_bag_count)
        ws2.cell(row=idx, column=7, value=order.thermal_bag_deposit)
        ws2.cell(row=idx, column=8, value=order.pallet_count)
        ws2.cell(row=idx, column=9, value=order.pallet_deposit)
        ws2.cell(row=idx, column=10, value=order.total_deposit)
        ws2.cell(row=idx, column=11, value=order.refunded_amount)
        ws2.cell(row=idx, column=12, value=order.total_deduction)
        ws2.cell(row=idx, column=13, value=order.refundable_amount)
        ws2.cell(row=idx, column=14, value=DEPOSIT_STATUS_NAMES.get(order.status, order.status))
        ws2.cell(row=idx, column=15, value=_format_datetime(order.created_at))

    _set_data_style(ws2, len(orders) + 1, len(headers2))
    _auto_fit_columns(ws2)

    order_ids = [o.id for o in orders]

    ws3 = wb.create_sheet("扫码记录")
    headers3 = [
        "扫码编号", "押金单号", "包材类型", "包材编码",
        "数量", "操作人", "扫码时间", "是否撤回",
    ]
    _set_header_style(ws3, headers3)

    scans = db.query(ScanRecord).filter(
        ScanRecord.deposit_order_id.in_(order_ids)
    ).all()

    for idx, scan in enumerate(scans, 2):
        ws3.cell(row=idx, column=1, value=scan.scan_no)
        ws3.cell(row=idx, column=2, value=scan.deposit_order.order_no)
        ws3.cell(row=idx, column=3, value=PACKAGE_TYPE_NAMES.get(scan.package_type, scan.package_type))
        ws3.cell(row=idx, column=4, value=scan.package_code)
        ws3.cell(row=idx, column=5, value=scan.quantity)
        ws3.cell(row=idx, column=6, value=scan.operator_name)
        ws3.cell(row=idx, column=7, value=_format_datetime(scan.scan_time))
        ws3.cell(row=idx, column=8, value="是" if scan.is_reversed else "否")

    _set_data_style(ws3, len(scans) + 1, len(headers3))
    _auto_fit_columns(ws3)

    ws4 = wb.create_sheet("损坏记录")
    headers4 = [
        "损坏编号", "押金单号", "包材类型", "包材编码",
        "数量", "扣减金额", "操作人", "创建时间", "是否撤回",
    ]
    _set_header_style(ws4, headers4)

    damages = db.query(DamageRecord).filter(
        DamageRecord.deposit_order_id.in_(order_ids)
    ).all()

    for idx, damage in enumerate(damages, 2):
        ws4.cell(row=idx, column=1, value=damage.damage_no)
        ws4.cell(row=idx, column=2, value=damage.deposit_order.order_no)
        ws4.cell(row=idx, column=3, value=PACKAGE_TYPE_NAMES.get(damage.package_type, damage.package_type))
        ws4.cell(row=idx, column=4, value=damage.package_code)
        ws4.cell(row=idx, column=5, value=damage.quantity)
        ws4.cell(row=idx, column=6, value=damage.deduction_amount)
        ws4.cell(row=idx, column=7, value=damage.operator_name)
        ws4.cell(row=idx, column=8, value=_format_datetime(damage.created_at))
        ws4.cell(row=idx, column=9, value="是" if damage.is_reversed else "否")

    _set_data_style(ws4, len(damages) + 1, len(headers4))
    _auto_fit_columns(ws4)

    ws5 = wb.create_sheet("退款记录")
    headers5 = [
        "退款单号", "押金单号",
        "循环箱退款", "保温袋退款", "托盘退款", "退款总金额",
        "状态", "操作人", "创建时间", "完成时间",
    ]
    _set_header_style(ws5, headers5)

    refunds = db.query(RefundOrder).filter(
        RefundOrder.deposit_order_id.in_(order_ids)
    ).all()

    for idx, refund in enumerate(refunds, 2):
        ws5.cell(row=idx, column=1, value=refund.refund_no)
        ws5.cell(row=idx, column=2, value=refund.deposit_order.order_no)
        ws5.cell(row=idx, column=3, value=refund.cycle_box_refund)
        ws5.cell(row=idx, column=4, value=refund.thermal_bag_refund)
        ws5.cell(row=idx, column=5, value=refund.pallet_refund)
        ws5.cell(row=idx, column=6, value=refund.refund_amount)
        ws5.cell(row=idx, column=7, value=REFUND_STATUS_NAMES.get(refund.status, refund.status))
        ws5.cell(row=idx, column=8, value=refund.operator_name)
        ws5.cell(row=idx, column=9, value=_format_datetime(refund.created_at))
        ws5.cell(row=idx, column=10, value=_format_datetime(refund.completed_at))

    _set_data_style(ws5, len(refunds) + 1, len(headers5))
    _auto_fit_columns(ws5)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
