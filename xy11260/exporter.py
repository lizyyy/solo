from datetime import datetime
from typing import Optional, List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import json
from database import Hazard, HazardStatus, OperationLog, OperationType
from utils import mask_sensitive_data, logger


class DataExporter:
    def __init__(self, db_session):
        self.db = db_session

    def _format_datetime(self, dt: Optional[datetime]) -> str:
        if not dt:
            return ""
        return dt.strftime("%Y-%m-%d %H:%M:%S")

    def _get_status_text(self, status: Optional[HazardStatus]) -> str:
        status_map = {
            HazardStatus.REGISTERED: "已登记",
            HazardStatus.ASSIGNED: "已派发",
            HazardStatus.RECTIFIED: "已整改",
            HazardStatus.RECHECKED: "已复查",
            HazardStatus.ARCHIVED: "已归档",
            HazardStatus.REJECTED: "已驳回"
        }
        return status_map.get(status, str(status) if status else "")

    def _get_operation_type_text(self, op_type: Optional[OperationType]) -> str:
        type_map = {
            OperationType.REGISTER: "登记",
            OperationType.ASSIGN: "派发",
            OperationType.RECTIFY: "整改",
            OperationType.RECHECK: "复查",
            OperationType.ARCHIVE: "归档",
            OperationType.REJECT: "驳回",
            OperationType.UPDATE: "更新",
            OperationType.DELETE: "删除"
        }
        return type_map.get(op_type, str(op_type) if op_type else "")

    def _apply_header_style(self, worksheet, row: int, col_count: int):
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, col_count + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

    def export_hazards_to_excel(self, file_path: str, status: Optional[HazardStatus] = None,
                                level: Optional[str] = None, start_date: Optional[datetime] = None,
                                end_date: Optional[datetime] = None) -> Dict:
        query = self.db.query(Hazard).filter(Hazard.is_deleted == False)

        if status:
            query = query.filter(Hazard.status == status)
        if level:
            query = query.filter(Hazard.level == level)
        if start_date:
            query = query.filter(Hazard.registered_time >= start_date)
        if end_date:
            query = query.filter(Hazard.registered_time <= end_date)

        hazards = query.order_by(Hazard.registered_time.desc()).all()

        if not hazards:
            return {"success": False, "message": "没有数据可导出"}

        wb = Workbook()

        ws_main = wb.active
        ws_main.title = "隐患清单"

        headers = [
            "隐患编号", "标题", "描述", "位置", "等级", "状态",
            "整改人", "整改部门", "截止时间",
            "整改描述", "整改时间",
            "复查结果", "复查意见", "复查时间",
            "登记人", "登记时间",
            "派发人", "派发时间",
            "归档人", "归档时间"
        ]

        for col, header in enumerate(headers, 1):
            ws_main.cell(row=1, column=col, value=header)

        self._apply_header_style(ws_main, 1, len(headers))

        for row_idx, hazard in enumerate(hazards, 2):
            data = mask_sensitive_data({
                "hazard_no": hazard.hazard_no,
                "title": hazard.title,
                "description": hazard.description,
                "location": hazard.location,
                "level": hazard.level,
                "status": self._get_status_text(hazard.status),
                "rectifier_name": hazard.rectifier_name or "",
                "rectifier_dept": hazard.rectifier_dept or "",
                "deadline": self._format_datetime(hazard.deadline),
                "rectification_desc": hazard.rectification_desc or "",
                "rectification_time": self._format_datetime(hazard.rectification_time),
                "recheck_result": "通过" if hazard.recheck_result else "不通过" if hazard.recheck_result is not None else "",
                "recheck_opinion": hazard.recheck_opinion or "",
                "recheck_time": self._format_datetime(hazard.recheck_time),
                "registered_by_name": hazard.registered_by_name,
                "registered_time": self._format_datetime(hazard.registered_time),
                "assigned_by_name": hazard.assigned_by_name or "",
                "assigned_time": self._format_datetime(hazard.assigned_time),
                "archived_by_name": hazard.archived_by_name or "",
                "archived_time": self._format_datetime(hazard.archived_time)
            })

            values = [
                data["hazard_no"], data["title"], data["description"],
                data["location"], data["level"], data["status"],
                data["rectifier_name"], data["rectifier_dept"], data["deadline"],
                data["rectification_desc"], data["rectification_time"],
                data["recheck_result"], data["recheck_opinion"], data["recheck_time"],
                data["registered_by_name"], data["registered_time"],
                data["assigned_by_name"], data["assigned_time"],
                data["archived_by_name"], data["archived_time"]
            ]

            for col, value in enumerate(values, 1):
                cell = ws_main.cell(row=row_idx, column=col, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

                if data["status"] == "已归档":
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif data["status"] == "已驳回":
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        for col in range(1, len(headers) + 1):
            ws_main.column_dimensions[get_column_letter(col)].width = 15

        ws_logs = wb.create_sheet(title="操作日志")
        self._export_operation_logs_to_sheet(ws_logs, hazards)

        ws_stats = wb.create_sheet(title="统计汇总")
        self._export_statistics_to_sheet(ws_stats, hazards)

        wb.save(file_path)

        logger.info(f"隐患数据导出成功: {file_path}, 共 {len(hazards)} 条记录")

        return {
            "success": True,
            "message": "导出成功",
            "data": {
                "file_path": file_path,
                "record_count": len(hazards)
            }
        }

    def _export_operation_logs_to_sheet(self, worksheet, hazards: List[Hazard]):
        hazard_ids = [h.id for h in hazards]
        logs = self.db.query(OperationLog).filter(
            OperationLog.hazard_id.in_(hazard_ids)
        ).order_by(OperationLog.operation_time.desc()).all()

        headers = [
            "隐患编号", "操作类型", "操作人", "角色",
            "操作时间", "备注", "IP地址"
        ]

        for col, header in enumerate(headers, 1):
            worksheet.cell(row=1, column=col, value=header)

        self._apply_header_style(worksheet, 1, len(headers))

        hazard_map = {h.id: h.hazard_no for h in hazards}

        for row_idx, log in enumerate(logs, 2):
            data = mask_sensitive_data({
                "hazard_no": hazard_map.get(log.hazard_id, ""),
                "operation_type": self._get_operation_type_text(log.operation_type),
                "operator_name": log.operator_name,
                "operator_role": log.operator_role.value if hasattr(log.operator_role, 'value') else str(log.operator_role),
                "operation_time": self._format_datetime(log.operation_time),
                "remark": log.remark or "",
                "ip_address": log.ip_address or ""
            })

            values = [
                data["hazard_no"], data["operation_type"], data["operator_name"],
                data["operator_role"], data["operation_time"], data["remark"],
                data["ip_address"]
            ]

            for col, value in enumerate(values, 1):
                worksheet.cell(row=row_idx, column=col, value=value)

        for col in range(1, len(headers) + 1):
            worksheet.column_dimensions[get_column_letter(col)].width = 18

    def _export_statistics_to_sheet(self, worksheet, hazards: List[Hazard]):
        total = len(hazards)

        status_counts = {}
        for h in hazards:
            status_text = self._get_status_text(h.status)
            status_counts[status_text] = status_counts.get(status_text, 0) + 1

        level_counts = {}
        for h in hazards:
            level_counts[h.level] = level_counts.get(h.level, 0) + 1

        dept_counts = {}
        for h in hazards:
            if h.rectifier_dept:
                dept_counts[h.rectifier_dept] = dept_counts.get(h.rectifier_dept, 0) + 1

        row = 1

        worksheet.cell(row=row, column=1, value="隐患闭环管理统计汇总")
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        worksheet.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 2

        worksheet.cell(row=row, column=1, value="总隐患数").font = Font(bold=True)
        worksheet.cell(row=row, column=2, value=total)
        row += 1

        closed = status_counts.get("已归档", 0)
        worksheet.cell(row=row, column=1, value="已闭环数").font = Font(bold=True)
        worksheet.cell(row=row, column=2, value=closed)
        row += 1

        worksheet.cell(row=row, column=1, value="闭环率").font = Font(bold=True)
        worksheet.cell(row=row, column=2, value=f"{round(closed / total * 100, 2) if total > 0 else 0}%")
        row += 2

        worksheet.cell(row=row, column=1, value="按状态统计").font = Font(bold=True, size=12)
        row += 1
        for status, count in status_counts.items():
            worksheet.cell(row=row, column=1, value=status)
            worksheet.cell(row=row, column=2, value=count)
            row += 1
        row += 1

        worksheet.cell(row=row, column=1, value="按等级统计").font = Font(bold=True, size=12)
        row += 1
        for level, count in level_counts.items():
            worksheet.cell(row=row, column=1, value=level)
            worksheet.cell(row=row, column=2, value=count)
            row += 1
        row += 1

        worksheet.cell(row=row, column=1, value="按部门统计").font = Font(bold=True, size=12)
        row += 1
        for dept, count in dept_counts.items():
            worksheet.cell(row=row, column=1, value=dept)
            worksheet.cell(row=row, column=2, value=count)
            row += 1

        worksheet.column_dimensions['A'].width = 20
        worksheet.column_dimensions['B'].width = 15

    def export_hazards_to_json(self, file_path: str, status: Optional[HazardStatus] = None,
                               level: Optional[str] = None, start_date: Optional[datetime] = None,
                               end_date: Optional[datetime] = None) -> Dict:
        query = self.db.query(Hazard).filter(Hazard.is_deleted == False)

        if status:
            query = query.filter(Hazard.status == status)
        if level:
            query = query.filter(Hazard.level == level)
        if start_date:
            query = query.filter(Hazard.registered_time >= start_date)
        if end_date:
            query = query.filter(Hazard.registered_time <= end_date)

        hazards = query.order_by(Hazard.registered_time.desc()).all()

        if not hazards:
            return {"success": False, "message": "没有数据可导出"}

        from utils import hazard_to_dict

        export_data = {
            "export_time": datetime.utcnow().isoformat(),
            "record_count": len(hazards),
            "hazards": [mask_sensitive_data(hazard_to_dict(h)) for h in hazards]
        }

        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, ensure_ascii=False, indent=2)

        logger.info(f"隐患数据JSON导出成功: {file_path}, 共 {len(hazards)} 条记录")

        return {
            "success": True,
            "message": "导出成功",
            "data": {
                "file_path": file_path,
                "record_count": len(hazards)
            }
        }
