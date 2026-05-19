import json
import csv
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from ..engine import ValidationResult, DecisionType, ValidationLevel
from ..tracker import AuditTrail


class ReportGenerator:
    def __init__(self, audit_trail: AuditTrail):
        self.audit_trail = audit_trail
        self.stats = audit_trail.get_statistics()
        self.audit_report = audit_trail.generate_audit_report()
        self.bad_rows_report = audit_trail.get_bad_rows_report()

    def generate_console_summary(self) -> str:
        lines = []
        lines.append("=" * 70)
        lines.append("多仓库依赖升级许可延期管理排查报告")
        lines.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("=" * 70)
        lines.append("")

        lines.append("【统计摘要】")
        lines.append(f"  仓库总数: {self.stats['total_repos']}")
        lines.append(f"  依赖包总数: {self.stats['total_dependencies']}")
        lines.append(f"  负责人意见数: {self.stats['total_opinions']}")
        lines.append(f"  延期申请数: {self.stats['total_extensions']}")
        lines.append(f"  升级批次数: {self.stats['total_batches']}")
        lines.append("")

        lines.append("【验证结果】")
        lines.append(f"  已批准: {self.stats['validation']['approved']}")
        lines.append(f"  已拒绝: {self.stats['validation']['rejected']}")
        lines.append(f"  需要延期: {self.stats['validation']['need_extension']}")
        lines.append(f"  待定: {self.stats['validation']['pending']}")
        lines.append("")

        lines.append("【问题统计】")
        lines.append(f"  错误: {self.stats['issues_by_level']['error']}")
        lines.append(f"  警告: {self.stats['issues_by_level']['warning']}")
        lines.append(f"  信息: {self.stats['issues_by_level']['info']}")
        lines.append(f"  坏行数: {self.stats['bad_rows_count']}")
        lines.append("")

        lines.append("【仓库详情】")
        for repo in self.audit_report:
            status_icon = self._get_status_icon(repo["decision"])
            lines.append(f"{status_icon} {repo['repo_name']}: {repo['decision']}")
            if repo["approved_packages"]:
                lines.append(f"    可升级包: {', '.join(repo['approved_packages'])}")
            errors = [i for i in repo["issues"] if i["level"] == "error"]
            warnings = [i for i in repo["issues"] if i["level"] == "warning"]
            if errors:
                lines.append(f"    错误 ({len(errors)}):")
                for e in errors[:3]:
                    lines.append(f"      - {e['message']}")
                if len(errors) > 3:
                    lines.append(f"      ... 还有 {len(errors) - 3} 个错误")
            if warnings:
                lines.append(f"    警告 ({len(warnings)}):")
                for w in warnings[:2]:
                    lines.append(f"      - {w['message']}")
                if len(warnings) > 2:
                    lines.append(f"      ... 还有 {len(warnings) - 2} 个警告")

        if self.bad_rows_report:
            lines.append("")
            lines.append("【坏行记录】")
            for row in self.bad_rows_report[:5]:
                lines.append(f"  - {row['source']}: {row['error_message']}")
            if len(self.bad_rows_report) > 5:
                lines.append(f"  ... 还有 {len(self.bad_rows_report) - 5} 个坏行")

        lines.append("")
        lines.append("=" * 70)

        return "\n".join(lines)

    def _get_status_icon(self, decision: DecisionType) -> str:
        return {
            DecisionType.APPROVED: "✅",
            DecisionType.REJECTED: "❌",
            DecisionType.NEED_EXTENSION: "⏳",
            DecisionType.PENDING: "❓",
            DecisionType.CONFLICT: "⚠️",
            DecisionType.VERSION_MISMATCH: "📦",
        }.get(decision, "❓")

    def export_license_list(self, output_path: str) -> str:
        approved_repos = [r for r in self.audit_report if r["decision"] == DecisionType.APPROVED]
        data = {
            "export_time": datetime.now().isoformat(),
            "total_approved": len(approved_repos),
            "repositories": [
                {
                    "repo_name": r["repo_name"],
                    "approved_packages": r["approved_packages"],
                    "can_upgrade": r["can_upgrade"],
                }
                for r in approved_repos
            ],
        }
        self._write_json(output_path, data)
        return output_path

    def export_full_report(self, output_path: str) -> str:
        data = {
            "export_time": datetime.now().isoformat(),
            "statistics": self.stats,
            "repositories": self.audit_report,
            "bad_rows": self.bad_rows_report,
        }
        self._write_json(output_path, data)
        return output_path

    def export_csv_report(self, output_path: str) -> str:
        rows = []
        for repo in self.audit_report:
            errors = [i["message"] for i in repo["issues"] if i["level"] == "error"]
            warnings = [i["message"] for i in repo["issues"] if i["level"] == "warning"]
            rows.append({
                "仓库名称": repo["repo_name"],
                "决策结果": repo["decision"],
                "可升级": "是" if repo["can_upgrade"] else "否",
                "批准的包": ", ".join(repo["approved_packages"]),
                "错误数": len(errors),
                "警告数": len(warnings),
                "错误详情": " | ".join(errors),
                "警告详情": " | ".join(warnings),
            })

        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=rows[0].keys() if rows else [])
            writer.writeheader()
            writer.writerows(rows)
        return output_path

    def export_excel_report(self, output_path: str) -> str:
        if not HAS_OPENPYXL:
            raise ImportError("需要安装 openpyxl 才能导出 Excel 报告")

        wb = Workbook()

        ws_summary = wb.active
        ws_summary.title = "摘要"
        self._write_summary_sheet(ws_summary)

        ws_repos = wb.create_sheet("仓库详情")
        self._write_repos_sheet(ws_repos)

        ws_bad = wb.create_sheet("坏行记录")
        self._write_bad_rows_sheet(ws_bad)

        wb.save(output_path)
        return output_path

    def _write_summary_sheet(self, ws):
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        ws["A1"] = "多仓库依赖升级许可延期管理排查报告"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        row = 4
        ws.cell(row=row, column=1, value="统计项目").fill = header_fill
        ws.cell(row=row, column=2, value="数值").fill = header_fill
        for cell in ws[row]:
            cell.font = header_font

        row = 5
        items = [
            ("仓库总数", self.stats["total_repos"]),
            ("依赖包总数", self.stats["total_dependencies"]),
            ("负责人意见数", self.stats["total_opinions"]),
            ("延期申请数", self.stats["total_extensions"]),
            ("升级批次数", self.stats["total_batches"]),
            ("", ""),
            ("已批准", self.stats["validation"]["approved"]),
            ("已拒绝", self.stats["validation"]["rejected"]),
            ("需要延期", self.stats["validation"]["need_extension"]),
            ("待定", self.stats["validation"]["pending"]),
            ("", ""),
            ("错误数", self.stats["issues_by_level"]["error"]),
            ("警告数", self.stats["issues_by_level"]["warning"]),
            ("信息数", self.stats["issues_by_level"]["info"]),
            ("坏行数", self.stats["bad_rows_count"]),
        ]
        for label, value in items:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 15

    def _write_repos_sheet(self, ws):
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        headers = ["仓库名称", "决策结果", "可升级", "批准的包", "错误数", "警告数", "错误详情", "警告详情"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, repo in enumerate(self.audit_report, 2):
            errors = [i["message"] for i in repo["issues"] if i["level"] == "error"]
            warnings = [i["message"] for i in repo["issues"] if i["level"] == "warning"]

            ws.cell(row=row_idx, column=1, value=repo["repo_name"])
            ws.cell(row=row_idx, column=2, value=repo["decision"])
            ws.cell(row=row_idx, column=3, value="是" if repo["can_upgrade"] else "否")
            ws.cell(row=row_idx, column=4, value=", ".join(repo["approved_packages"]))
            ws.cell(row=row_idx, column=5, value=len(errors))
            ws.cell(row=row_idx, column=6, value=len(warnings))
            ws.cell(row=row_idx, column=7, value=" | ".join(errors))
            ws.cell(row=row_idx, column=8, value=" | ".join(warnings))

            if repo["decision"] == DecisionType.REJECTED:
                fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif repo["decision"] == DecisionType.APPROVED:
                fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif repo["decision"] == DecisionType.NEED_EXTENSION:
                fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            else:
                fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

            for col in range(1, 9):
                ws.cell(row=row_idx, column=col).fill = fill

        for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            ws.column_dimensions[col].width = 20

    def _write_bad_rows_sheet(self, ws):
        header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        headers = ["来源", "错误类型", "错误信息", "原始数据"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, row in enumerate(self.bad_rows_report, 2):
            ws.cell(row=row_idx, column=1, value=row["source"])
            ws.cell(row=row_idx, column=2, value=row["error_type"])
            ws.cell(row=row_idx, column=3, value=row["error_message"])
            ws.cell(row=row_idx, column=4, value=str(row["raw_data"]))

        for col in ["A", "B", "C", "D"]:
            ws.column_dimensions[col].width = 30

    def _write_json(self, path: str, data: Dict[str, Any]):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def export_all(self, output_dir: str, prefix: str = "dep_report") -> List[str]:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        files = []

        json_path = out_dir / f"{prefix}_{timestamp}.json"
        self.export_full_report(str(json_path))
        files.append(str(json_path))

        csv_path = out_dir / f"{prefix}_{timestamp}.csv"
        self.export_csv_report(str(csv_path))
        files.append(str(csv_path))

        if HAS_OPENPYXL:
            xlsx_path = out_dir / f"{prefix}_{timestamp}.xlsx"
            self.export_excel_report(str(xlsx_path))
            files.append(str(xlsx_path))

        license_path = out_dir / f"{prefix}_license_{timestamp}.json"
        self.export_license_list(str(license_path))
        files.append(str(license_path))

        return files
