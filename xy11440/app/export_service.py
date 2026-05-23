from datetime import datetime
from typing import List, Optional
from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import Batch, Record, StatusTrail, BatchStatus, DirtyType
from .schemas import ExportRow, BatchExportResponse


class ExportService:
    @staticmethod
    def get_batch_export_data(db: Session, batch_id: Optional[int] = None) -> BatchExportResponse:
        query = db.query(Batch)
        if batch_id:
            query = query.filter(Batch.id == batch_id)
        batches = query.order_by(Batch.created_at.desc()).all()
        
        export_rows = []
        for batch in batches:
            status_before_freeze = ExportService._get_status_before_freeze(db, batch.id)
            manual_reason = ExportService._get_manual_reason(batch)
            duplicate_warning = ExportService._check_duplicate_calculation(batch)
            
            records = db.query(Record).filter(Record.batch_id == batch.id).all()
            dirty_records = [r for r in records if r.is_dirty]
            
            export_rows.append(ExportRow(
                batch_no=batch.batch_no,
                supplier_name=batch.supplier_name,
                delivery_date=batch.delivery_date.strftime("%Y-%m-%d") if batch.delivery_date else "",
                status_before_freeze=status_before_freeze,
                status_after_freeze=batch.status.value,
                bad_fruit_deduction=batch.bad_fruit_deduction,
                secondary_sorting_loss=batch.secondary_sorting_loss,
                duplicate_calculation_warning=duplicate_warning,
                final_settlement=batch.final_settlement,
                manual_reason=manual_reason,
                record_count=len(records),
                dirty_record_count=len(dirty_records),
                created_at=batch.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ))
        
        return BatchExportResponse(
            data=export_rows,
            export_time=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            total_count=len(export_rows)
        )

    @staticmethod
    def _get_status_before_freeze(db: Session, batch_id: int) -> str:
        trails = db.query(StatusTrail).filter(
            StatusTrail.batch_id == batch_id,
            StatusTrail.to_status == BatchStatus.FROZEN
        ).order_by(StatusTrail.changed_at.desc()).first()
        
        if trails:
            return trails.from_status.value if trails.from_status else "N/A"
        return "未冻结"

    @staticmethod
    def _get_manual_reason(batch: Batch) -> str:
        reasons = []
        if batch.review_comment:
            reasons.append(f"复核意见: {batch.review_comment}")
        if batch.freeze_reason:
            reasons.append(f"冻结原因: {batch.freeze_reason}")
        if batch.archive_reason:
            reasons.append(f"归档原因: {batch.archive_reason}")
        return "; ".join(reasons) if reasons else "无人工备注"

    @staticmethod
    def _check_duplicate_calculation(batch: Batch) -> str:
        if batch.bad_fruit_deduction > 0 and batch.secondary_sorting_loss > 0:
            if batch.bad_fruit_deduction + batch.secondary_sorting_loss > batch.total_weighing_amount * 0.1:
                return "警告: 坏果扣款和二次分拣损耗总和超过10%，请确认是否存在重复计算"
            return "注意: 同时存在坏果扣款和二次分拣损耗，请核对明细"
        return "正常"

    @staticmethod
    def export_to_excel(db: Session, batch_id: Optional[int] = None) -> BytesIO:
        export_data = ExportService.get_batch_export_data(db, batch_id)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "生鲜分拣损耗汇总"
        
        headers = [
            "批次号", "供应商", "送货日期", "冻结前状态", "当前状态",
            "坏果扣款", "二次分拣损耗", "重复计算警告", "最终结算金额",
            "人工备注", "记录总数", "脏记录数", "创建时间"
        ]
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        
        for row_idx, row_data in enumerate(export_data.data, 2):
            ws.cell(row=row_idx, column=1, value=row_data.batch_no)
            ws.cell(row=row_idx, column=2, value=row_data.supplier_name)
            ws.cell(row=row_idx, column=3, value=row_data.delivery_date)
            ws.cell(row=row_idx, column=4, value=row_data.status_before_freeze)
            ws.cell(row=row_idx, column=5, value=row_data.status_after_freeze)
            ws.cell(row=row_idx, column=6, value=row_data.bad_fruit_deduction)
            ws.cell(row=row_idx, column=7, value=row_data.secondary_sorting_loss)
            
            warning_cell = ws.cell(row=row_idx, column=8, value=row_data.duplicate_calculation_warning)
            if "警告" in row_data.duplicate_calculation_warning:
                warning_cell.fill = warning_fill
            
            ws.cell(row=row_idx, column=9, value=row_data.final_settlement)
            ws.cell(row=row_idx, column=10, value=row_data.manual_reason)
            ws.cell(row=row_idx, column=11, value=row_data.record_count)
            ws.cell(row=row_idx, column=12, value=row_data.dirty_record_count)
            ws.cell(row=row_idx, column=13, value=row_data.created_at)
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 18
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def get_manager_dashboard(db: Session) -> dict:
        batches = db.query(Batch).all()
        
        frozen_batches = [b for b in batches if b.status == BatchStatus.FROZEN]
        pending_review = [b for b in batches if b.status == BatchStatus.PENDING_REVIEW]
        with_dirty_records = []
        
        for batch in batches:
            dirty_count = db.query(Record).filter(
                Record.batch_id == batch.id,
                Record.is_dirty == True
            ).count()
            if dirty_count > 0:
                with_dirty_records.append({
                    "batch_no": batch.batch_no,
                    "supplier_name": batch.supplier_name,
                    "dirty_count": dirty_count
                })
        
        total_bad_fruit = sum(b.bad_fruit_deduction for b in batches)
        total_secondary_loss = sum(b.secondary_sorting_loss for b in batches)
        total_settlement = sum(b.final_settlement for b in batches)
        
        return {
            "summary": {
                "total_batches": len(batches),
                "frozen_batches": len(frozen_batches),
                "pending_review": len(pending_review),
                "batches_with_dirty": len(with_dirty_records),
                "total_bad_fruit_deduction": total_bad_fruit,
                "total_secondary_sorting_loss": total_secondary_loss,
                "total_final_settlement": total_settlement
            },
            "dirty_record_batches": with_dirty_records,
            "recent_frozen": [
                {
                    "batch_no": b.batch_no,
                    "supplier_name": b.supplier_name,
                    "freeze_reason": b.freeze_reason,
                    "final_settlement": b.final_settlement
                }
                for b in sorted(frozen_batches, key=lambda x: x.updated_at, reverse=True)[:10]
            ]
        }
