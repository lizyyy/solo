from typing import List, Optional
from datetime import datetime
import json
import uuid
from io import BytesIO
from sqlalchemy.orm import Session

from app.models import SwitchReport, SwitchRequest, SwitchStatus
from app.services.supplier_service import SupplierService
from app.services.qualification_service import QualificationService
from app.services.price_service import PriceService
from app.services.exception_service import ExceptionService
from app.services.delivery_service import DeliveryService


class ReportService:
    def __init__(self, db: Session):
        self.db = db
        self.supplier_service = SupplierService(db)
        self.qualification_service = QualificationService(db)
        self.price_service = PriceService(db)
        self.exception_service = ExceptionService(db)
        self.delivery_service = DeliveryService(db)

    def generate_report_no(self) -> str:
        return f"RPT-{datetime.utcnow().strftime('%Y%m%d')}-{uuid.uuid4().hex[:8].upper()}"

    def generate_switch_report(
        self,
        switch_request_id: int,
        generated_by: Optional[str] = None
    ) -> Optional[SwitchReport]:
        switch_request = self.db.query(SwitchRequest).filter(
            SwitchRequest.id == switch_request_id
        ).first()

        if not switch_request:
            return None

        primary_supplier = self.supplier_service.get_supplier(switch_request.primary_supplier_id)
        alternative_supplier = self.supplier_service.get_supplier(switch_request.alternative_supplier_id)

        qualification_result = self.qualification_service.check_qualification(
            switch_request.alternative_supplier_id
        )

        price_comparison = self.price_service.compare_prices(
            switch_request.primary_supplier_id,
            switch_request.alternative_supplier_id,
            switch_request.product_code
        )

        exceptions = self.exception_service.get_exceptions_by_request(switch_request_id)
        delivery_impacts = self.delivery_service.get_delivery_impacts(switch_request_id)

        report_content = self._build_report_content(
            switch_request=switch_request,
            primary_supplier=primary_supplier,
            alternative_supplier=alternative_supplier,
            qualification_result=qualification_result,
            price_comparison=price_comparison,
            exceptions=exceptions,
            delivery_impacts=delivery_impacts
        )

        report = SwitchReport(
            switch_request_id=switch_request_id,
            report_no=self.generate_report_no(),
            content=json.dumps(report_content, ensure_ascii=False, default=str),
            generated_by=generated_by
        )
        self.db.add(report)
        self.db.commit()
        self.db.refresh(report)

        return report

    def _build_report_content(
        self,
        switch_request: SwitchRequest,
        primary_supplier,
        alternative_supplier,
        qualification_result,
        price_comparison,
        exceptions,
        delivery_impacts
    ) -> dict:
        return {
            "report_header": {
                "report_title": "供应商切换评估报告",
                "request_no": switch_request.request_no,
                "generated_at": datetime.utcnow().isoformat(),
                "request_status": switch_request.status
            },
            "switch_basic_info": {
                "product_code": switch_request.product_code,
                "product_name": switch_request.product_name,
                "quantity": switch_request.quantity,
                "reason": switch_request.reason,
                "requester": switch_request.requester,
                "requester_department": switch_request.requester_department,
                "request_date": switch_request.created_at.isoformat()
            },
            "supplier_information": {
                "primary": {
                    "name": primary_supplier.name if primary_supplier else "未知",
                    "code": primary_supplier.code if primary_supplier else "未知",
                    "status": primary_supplier.status if primary_supplier else "未知"
                },
                "alternative": {
                    "name": alternative_supplier.name if alternative_supplier else "未知",
                    "code": alternative_supplier.code if alternative_supplier else "未知",
                    "status": alternative_supplier.status if alternative_supplier else "未知"
                }
            },
            "qualification_check": {
                "passed": qualification_result.passed,
                "message": qualification_result.message,
                "failed_items": qualification_result.failed_qualifications
            },
            "price_analysis": {
                "comparison": price_comparison.model_dump() if price_comparison else None,
                "assessment": self._assess_price_impact(price_comparison)
            },
            "delivery_impact": {
                "records": [
                    {
                        "original_delivery": di.original_delivery_date.isoformat(),
                        "new_delivery": di.new_delivery_date.isoformat(),
                        "delay_days": di.delay_days,
                        "impact": di.impact_description,
                        "mitigation": di.mitigation_measures
                    }
                    for di in delivery_impacts
                ]
            },
            "exception_records": {
                "count": len(exceptions),
                "items": [
                    {
                        "type": exc.exception_type,
                        "description": exc.description,
                        "detail": exc.detail,
                        "is_resolved": exc.is_resolved,
                        "created_at": exc.created_at.isoformat()
                    }
                    for exc in exceptions
                ]
            },
            "recommendation": self._generate_recommendation(
                qualification_result,
                price_comparison,
                exceptions,
                switch_request.status
            )
        }

    def _assess_price_impact(self, price_comparison) -> str:
        if not price_comparison:
            return "价格信息不完整，无法评估"
        
        diff_pct = price_comparison.price_difference_percent
        if diff_pct < 0:
            return f"价格下降{abs(diff_pct)}%，成本节省"
        elif diff_pct == 0:
            return "价格持平"
        elif diff_pct <= 5:
            return f"价格上涨{diff_pct}%，影响较小"
        elif diff_pct <= 15:
            return f"价格上涨{diff_pct}%，需要关注"
        else:
            return f"价格上涨{diff_pct}%，影响较大"

    def _generate_recommendation(
        self,
        qualification_result,
        price_comparison,
        exceptions,
        request_status
    ) -> str:
        unresolved_exceptions = [e for e in exceptions if not e.is_resolved]

        if not qualification_result.passed:
            return "建议：备选供应商资质校验不通过，不建议切换。请寻找其他合格供应商。"
        
        if unresolved_exceptions:
            return f"建议：存在{len(unresolved_exceptions)}条待处理异常记录，请先处理异常后再评估是否执行切换。"
        
        if request_status == SwitchStatus.APPROVED:
            return "建议：所有校验通过，审批已完成，可以执行供应商切换。"
        elif request_status == SwitchStatus.PENDING:
            return "建议：校验通过，待完成审批流程后执行切换。"
        elif request_status == SwitchStatus.EXECUTED:
            return "建议：切换已执行，请注意监控新供应商的交付质量和服务水平。"
        
        return "建议：请完成所有必要的校验和审批流程。"

    def get_report(self, report_id: int) -> Optional[SwitchReport]:
        return self.db.query(SwitchReport).filter(SwitchReport.id == report_id).first()

    def get_reports_by_request(self, switch_request_id: int) -> List[SwitchReport]:
        return self.db.query(SwitchReport).filter(
            SwitchReport.switch_request_id == switch_request_id
        ).order_by(SwitchReport.created_at.desc()).all()

    def export_report_to_excel(self, report_id: int) -> Optional[BytesIO]:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            return None

        report = self.get_report(report_id)
        if not report:
            return None

        content = json.loads(report.content)

        wb = Workbook()
        ws = wb.active
        ws.title = "供应商切换报告"

        header_font = Font(bold=True, size=14)
        sub_header_font = Font(bold=True, size=11)
        normal_font = Font(size=10)
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")

        ws['A1'] = content['report_header']['report_title']
        ws['A1'].font = header_font
        ws.merge_cells('A1:F1')
        ws['A1'].alignment = center_align

        row = 3
        ws.cell(row=row, column=1, value='申请单号').font = sub_header_font
        ws.cell(row=row, column=2, value=content['report_header']['request_no'])
        row += 1
        ws.cell(row=row, column=1, value='生成时间').font = sub_header_font
        ws.cell(row=row, column=2, value=content['report_header']['generated_at'])
        row += 1
        ws.cell(row=row, column=1, value='申请状态').font = sub_header_font
        ws.cell(row=row, column=2, value=content['report_header']['request_status'])

        row += 2
        ws.cell(row=row, column=1, value='一、切换基本信息').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        basic_info = content['switch_basic_info']
        basic_fields = [
            ('产品编码', basic_info.get('product_code')),
            ('产品名称', basic_info.get('product_name')),
            ('数量', basic_info.get('quantity')),
            ('切换原因', basic_info.get('reason')),
            ('申请人', basic_info.get('requester')),
            ('申请部门', basic_info.get('requester_department')),
            ('申请日期', basic_info.get('request_date'))
        ]
        for field, value in basic_fields:
            ws.cell(row=row, column=1, value=field).font = sub_header_font
            ws.cell(row=row, column=2, value=value).font = normal_font
            ws.merge_cells(f'B{row}:F{row}')
            row += 1

        row += 1
        ws.cell(row=row, column=1, value='二、供应商信息').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        supplier_info = content['supplier_information']
        headers = ['供应商类型', '名称', '编码', '状态']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        row += 1

        for supplier_type in ['primary', 'alternative']:
            supp = supplier_info.get(supplier_type, {})
            ws.cell(row=row, column=1, value='主供应商' if supplier_type == 'primary' else '备选供应商')
            ws.cell(row=row, column=2, value=supp.get('name'))
            ws.cell(row=row, column=3, value=supp.get('code'))
            ws.cell(row=row, column=4, value=supp.get('status'))
            for col in range(1, 5):
                ws.cell(row=row, column=col).border = thin_border
            row += 1

        row += 1
        ws.cell(row=row, column=1, value='三、资质校验结果').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        qual = content['qualification_check']
        ws.cell(row=row, column=1, value='校验结果').font = sub_header_font
        ws.cell(row=row, column=2, value='通过' if qual.get('passed') else '不通过')
        row += 1
        ws.cell(row=row, column=1, value='校验说明').font = sub_header_font
        ws.cell(row=row, column=2, value=qual.get('message'))
        ws.merge_cells(f'B{row}:F{row}')
        row += 1

        if qual.get('failed_items'):
            ws.cell(row=row, column=1, value='不合格项').font = sub_header_font
            for i, item in enumerate(qual.get('failed_items', [])):
                ws.cell(row=row + i, column=2, value=f"{i+1}. {item}")
                ws.merge_cells(f'B{row+i}:F{row+i}')
            row += len(qual.get('failed_items', []))

        row += 1
        ws.cell(row=row, column=1, value='四、价格分析').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        price = content.get('price_analysis', {})
        price_comp = price.get('comparison')
        if price_comp:
            price_headers = ['产品编码', '产品名称', '主供应商价格', '备选供应商价格', '价格差异', '差异百分比']
            for col, header in enumerate(price_headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            row += 1

            ws.cell(row=row, column=1, value=price_comp.get('product_code')).border = thin_border
            ws.cell(row=row, column=2, value=price_comp.get('product_name')).border = thin_border
            ws.cell(row=row, column=3, value=price_comp.get('primary_price')).border = thin_border
            ws.cell(row=row, column=4, value=price_comp.get('alternative_price')).border = thin_border
            ws.cell(row=row, column=5, value=price_comp.get('price_difference')).border = thin_border
            ws.cell(row=row, column=6, value=f"{price_comp.get('price_difference_percent')}%").border = thin_border
            row += 1

        ws.cell(row=row, column=1, value='价格评估').font = sub_header_font
        ws.cell(row=row, column=2, value=price.get('assessment'))
        ws.merge_cells(f'B{row}:F{row}')

        row += 2
        ws.cell(row=row, column=1, value='五、交期影响').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        delivery = content.get('delivery_impact', {})
        delivery_records = delivery.get('records', [])
        if delivery_records:
            delivery_headers = ['原交期', '新交期', '延误天数', '影响描述', '缓解措施']
            for col, header in enumerate(delivery_headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            row += 1

            for record in delivery_records:
                ws.cell(row=row, column=1, value=record.get('original_delivery')).border = thin_border
                ws.cell(row=row, column=2, value=record.get('new_delivery')).border = thin_border
                ws.cell(row=row, column=3, value=record.get('delay_days')).border = thin_border
                ws.cell(row=row, column=4, value=record.get('impact')).border = thin_border
                ws.cell(row=row, column=5, value=record.get('mitigation')).border = thin_border
                row += 1
        else:
            ws.cell(row=row, column=1, value='暂无交期影响记录')
            row += 1

        row += 1
        ws.cell(row=row, column=1, value='六、异常记录').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        exc = content.get('exception_records', {})
        ws.cell(row=row, column=1, value='异常数量').font = sub_header_font
        ws.cell(row=row, column=2, value=exc.get('count', 0))
        row += 1

        if exc.get('items'):
            exc_headers = ['类型', '描述', '详情', '是否已解决', '创建时间']
            for col, header in enumerate(exc_headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            row += 1

            for item in exc.get('items', []):
                ws.cell(row=row, column=1, value=item.get('type')).border = thin_border
                ws.cell(row=row, column=2, value=item.get('description')).border = thin_border
                ws.cell(row=row, column=3, value=item.get('detail')).border = thin_border
                ws.cell(row=row, column=4, value='是' if item.get('is_resolved') else '否').border = thin_border
                ws.cell(row=row, column=5, value=item.get('created_at')).border = thin_border
                row += 1

        row += 1
        ws.cell(row=row, column=1, value='七、建议').font = header_font
        ws.merge_cells(f'A{row}:F{row}')
        row += 1
        ws.cell(row=row, column=1, value=content.get('recommendation'))
        ws.merge_cells(f'A{row}:F{row}')

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 15

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def get_report_statistics(self) -> dict:
        total = self.db.query(SwitchReport).count()
        total_requests = self.db.query(SwitchRequest).count()
        executed_requests = self.db.query(SwitchRequest).filter(
            SwitchRequest.status == SwitchStatus.EXECUTED
        ).count()
        failed_requests = self.db.query(SwitchRequest).filter(
            SwitchRequest.status == SwitchStatus.FAILED
        ).count()

        return {
            "total_reports": total,
            "total_switch_requests": total_requests,
            "executed_requests": executed_requests,
            "failed_requests": failed_requests,
            "success_rate": round(executed_requests / total_requests * 100, 2) if total_requests > 0 else 0
        }
