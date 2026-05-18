import pandas as pd
import os
from datetime import datetime
from typing import Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows


class ReportGenerator:
    def __init__(self, output_dir: str = None):
        if output_dir is None:
            output_dir = os.path.join(os.getcwd(), 'reports')
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
        
        self.timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    def generate_full_report(self, data: Dict[str, Any], 
                             filename: str = None) -> str:
        if filename is None:
            filename = f'social_work_report_{self.timestamp}.xlsx'
        
        filepath = os.path.join(self.output_dir, filename)
        
        wb = Workbook()
        
        # 先删除默认创建的Sheet工作表
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        self._add_summary_sheet(wb, data)
        self._add_risk_sheet(wb, data.get('risk_results', pd.DataFrame()))
        self._add_schedule_sheet(wb, data.get('schedule_results', pd.DataFrame()))
        self._add_material_sheet(wb, data.get('material_results', pd.DataFrame()))
        self._add_duplicates_sheet(wb, data.get('duplicate_results', pd.DataFrame()))
        self._add_source_sheet(wb, data.get('source_results', pd.DataFrame()))
        self._add_changes_sheet(wb, data.get('change_results', pd.DataFrame()))
        self._add_bad_rows_sheet(wb, data.get('bad_rows_results', pd.DataFrame()))
        
        wb.save(filepath)
        return filepath

    def _add_summary_sheet(self, wb: Workbook, data: Dict[str, Any]):
        ws = wb.create_sheet('总体概览', 0)
        
        ws['A1'] = '社工走访风险分级物资核销排查报告'
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.merge_cells('A1:D1')
        
        ws['A3'] = '生成时间'
        ws['B3'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        stats = data.get('statistics', {})
        row = 5
        ws[f'A{row}'] = '数据统计'
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        ws[f'A{row}'] = '居民总数'
        ws[f'B{row}'] = stats.get('total_residents', 0)
        row += 1
        
        ws[f'A{row}'] = '走访记录总数'
        ws[f'B{row}'] = stats.get('total_visits', 0)
        row += 1
        
        ws[f'A{row}'] = '物资发放记录总数'
        ws[f'B{row}'] = stats.get('total_materials', 0)
        row += 1
        
        ws[f'A{row}'] = '坏行记录总数'
        ws[f'B{row}'] = stats.get('total_bad_rows', 0)
        row += 2
        
        ws[f'A{row}'] = '风险分级统计'
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        risk_results = data.get('risk_results', pd.DataFrame())
        if not risk_results.empty:
            high_count = len(risk_results[risk_results['风险等级'] == 'high'])
            medium_count = len(risk_results[risk_results['风险等级'] == 'medium'])
            low_count = len(risk_results[risk_results['风险等级'] == 'low'])
            
            ws[f'A{row}'] = '高风险'
            ws[f'B{row}'] = high_count
            ws[f'B{row}'].fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            row += 1
            
            ws[f'A{row}'] = '中风险'
            ws[f'B{row}'] = medium_count
            ws[f'B{row}'].fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            row += 1
            
            ws[f'A{row}'] = '低风险'
            ws[f'B{row}'] = low_count
            ws[f'B{row}'].fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            row += 2
        
        ws[f'A{row}'] = '来源文件'
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        for f in stats.get('source_files', []):
            ws[f'A{row}'] = f
            row += 1

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20

    def _add_risk_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
        
        ws = wb.create_sheet('风险分级结果')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            risk_level = row[2].value
            if risk_level == 'high':
                fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
            elif risk_level == 'medium':
                fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            else:
                fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            
            for cell in row:
                cell.fill = fill

        self._auto_width(ws)

    def _add_schedule_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('回访计划安排')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            is_overdue = row[7].value
            if is_overdue == '是':
                fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                for cell in row:
                    cell.fill = fill

        self._auto_width(ws)

    def _add_material_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('物资核销结果')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            status = row[7].value
            if status == '异常':
                fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                for cell in row:
                    cell.fill = fill

        self._auto_width(ws)

    def _add_duplicates_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('重复走访记录')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            is_kept = row[6].value
            if is_kept == '是':
                fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            else:
                fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
            
            for cell in row:
                cell.fill = fill

        self._auto_width(ws)

    def _add_source_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('数据来源追踪')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        self._auto_width(ws)

    def _add_changes_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('数据修改记录')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        self._auto_width(ws)

    def _add_bad_rows_sheet(self, wb: Workbook, df: pd.DataFrame):
        if df.empty:
            return
            
        ws = wb.create_sheet('坏行记录')
        
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        self._style_header(ws)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

        self._auto_width(ws)

    def _style_header(self, ws):
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center')

    def _auto_width(self, ws):
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def export_csv(self, df: pd.DataFrame, filename: str) -> str:
        filepath = os.path.join(self.output_dir, f'{filename}_{self.timestamp}.csv')
        df.to_csv(filepath, index=False, encoding='utf-8-sig')
        return filepath

    def generate_summary_text(self, data: Dict[str, Any]) -> str:
        lines = []
        lines.append('=' * 60)
        lines.append('社工走访风险分级物资核销排查报告')
        lines.append('=' * 60)
        lines.append(f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        lines.append('')
        
        stats = data.get('statistics', {})
        lines.append('【数据统计】')
        lines.append(f'  居民总数: {stats.get("total_residents", 0)}')
        lines.append(f'  走访记录总数: {stats.get("total_visits", 0)}')
        lines.append(f'  物资发放记录总数: {stats.get("total_materials", 0)}')
        lines.append(f'  坏行记录总数: {stats.get("total_bad_rows", 0)}')
        lines.append('')
        
        risk_results = data.get('risk_results', pd.DataFrame())
        if not risk_results.empty:
            high_count = len(risk_results[risk_results['风险等级'] == 'high'])
            medium_count = len(risk_results[risk_results['风险等级'] == 'medium'])
            low_count = len(risk_results[risk_results['风险等级'] == 'low'])
            
            lines.append('【风险分级统计】')
            lines.append(f'  高风险: {high_count} 人')
            lines.append(f'  中风险: {medium_count} 人')
            lines.append(f'  低风险: {low_count} 人')
            lines.append('')
            
            if high_count > 0:
                lines.append('【高风险居民清单】')
                high_risk = risk_results[risk_results['风险等级'] == 'high']
                for _, row in high_risk.iterrows():
                    lines.append(f'  {row["姓名"]} ({row["身份证号"]}): {row["风险因素"]}')
                lines.append('')
        
        schedule = data.get('schedule_results', pd.DataFrame())
        if not schedule.empty:
            overdue = schedule[schedule['是否逾期'] == '是']
            lines.append(f'【逾期回访】共 {len(overdue)} 人')
            for _, row in overdue.iterrows():
                lines.append(f'  {row["姓名"]}: 逾期 {row["逾期天数"]} 天, 负责社工: {row["负责社工"]}')
            lines.append('')
        
        materials = data.get('material_results', pd.DataFrame())
        if not materials.empty:
            abnormal = materials[materials['核销状态'] == '异常']
            lines.append(f'【物资异常】共 {len(abnormal)} 条记录')
            lines.append('')
        
        lines.append('【来源文件】')
        for f in stats.get('source_files', []):
            lines.append(f'  {f}')
        
        lines.append('=' * 60)
        
        return '\n'.join(lines)
