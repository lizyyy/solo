from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import List, Dict, Any, Tuple
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import (
    Package, DisposalRecord, SummaryReport,
    DisposalType, ReconciliationStatus, AuditLog
)


class ReportExporter:
    @staticmethod
    def build_summary(packages: List[Package],
                     disposals: List[DisposalRecord]) -> SummaryReport:
        total = len(packages)
        pkg_map = {p.package_id: p for p in packages}

        picked_on_time = 0
        overdue_count = 0
        returned_count = 0
        manual_review_count = 0
        duplicate_reminder_count = 0
        pending_supplement_count = 0
        total_pickup_days = 0
        picked_count = 0

        for d in disposals:
            if d.disposal_type == DisposalType.RETURN:
                returned_count += 1
            if d.status == ReconciliationStatus.PENDING_REVIEW:
                manual_review_count += 1
            if d.disposal_type == DisposalType.SUPPLEMENT:
                pending_supplement_count += 1
            if any(dt.value == "duplicate_reminder" for dt in d.difference_types):
                duplicate_reminder_count += 1

            pkg = pkg_map.get(d.package_id)
            if pkg:
                from .models import PackageStatus, DifferenceType
                if pkg.status == PackageStatus.PICKED and pkg.pickup_date:
                    days = (pkg.pickup_date - pkg.arrival_date).days
                    if days <= 7:
                        picked_on_time += 1
                    total_pickup_days += days
                    picked_count += 1
                if any(dt == DifferenceType.OVERDUE_PICKUP for dt in d.difference_types):
                    overdue_count += 1

        avg_days = total_pickup_days / picked_count if picked_count > 0 else 0.0

        return SummaryReport(
            total_packages=total,
            picked_on_time=picked_on_time,
            overdue_count=overdue_count,
            returned_count=returned_count,
            manual_review_count=manual_review_count,
            duplicate_reminder_count=duplicate_reminder_count,
            pending_supplement_count=pending_supplement_count,
            average_pickup_days=round(avg_days, 2),
            reconciliation_date=datetime.now()
        )

    @staticmethod
    def export_details_csv(packages: List[Package],
                          disposals: List[DisposalRecord]) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        pkg_map = {p.package_id: p for p in packages}

        writer.writerow([
            "包裹ID", "运单号", "收件人", "取件码", "到件日期", "状态",
            "处置类型", "差异原因", "证据", "复核人", "复核备注", "状态"
        ])

        type_map = {
            DisposalType.RELEASE: "放行",
            DisposalType.RETURN: "退回",
            DisposalType.SUPPLEMENT: "补材料",
            DisposalType.MANUAL_REVIEW: "待复核"
        }

        status_map = {
            ReconciliationStatus.AUTO_MATCHED: "自动匹配",
            ReconciliationStatus.PENDING_REVIEW: "待复核",
            ReconciliationStatus.REVIEWED: "已复核",
            ReconciliationStatus.EXPORTED: "已导出"
        }

        for d in disposals:
            pkg = pkg_map.get(d.package_id)
            if not pkg:
                continue
            writer.writerow([
                pkg.package_id,
                pkg.tracking_no,
                pkg.recipient_name,
                pkg.pickup_code,
                pkg.arrival_date.isoformat(),
                pkg.status.value,
                type_map.get(d.disposal_type, d.disposal_type.value),
                d.reason,
                " | ".join(d.evidence),
                d.reviewed_by or "",
                d.review_note or "",
                status_map.get(d.status, d.status.value)
            ])

        return output.getvalue()

    @staticmethod
    def export_summary_json(summary: SummaryReport) -> str:
        return json.dumps(summary.model_dump(), ensure_ascii=False, indent=2,
                         default=str)

    @staticmethod
    def export_excel(packages: List[Package],
                    disposals: List[DisposalRecord],
                    summary: SummaryReport,
                    audit_logs: List[AuditLog] = None) -> bytes:
        output = io.BytesIO()
        wb = Workbook()

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        type_map = {
            DisposalType.RELEASE: "放行",
            DisposalType.RETURN: "退回",
            DisposalType.SUPPLEMENT: "补材料",
            DisposalType.MANUAL_REVIEW: "待复核"
        }
        status_map = {
            ReconciliationStatus.AUTO_MATCHED: "自动匹配",
            ReconciliationStatus.PENDING_REVIEW: "待复核",
            ReconciliationStatus.REVIEWED: "已复核",
            ReconciliationStatus.EXPORTED: "已导出"
        }

        ws1 = wb.active
        ws1.title = "汇总报告"
        summary_data = [
            ["指标", "数值"],
            ["总包裹数", summary.total_packages],
            ["按时取件数", summary.picked_on_time],
            ["超期包裹数", summary.overdue_count],
            ["退回包裹数", summary.returned_count],
            ["待人工复核数", summary.manual_review_count],
            ["重复催取数", summary.duplicate_reminder_count],
            ["待补材料数", summary.pending_supplement_count],
            ["平均取件天数", summary.average_pickup_days],
            ["对账时间", summary.reconciliation_date.strftime("%Y-%m-%d %H:%M:%S")],
        ]
        for row_idx, row in enumerate(summary_data, start=1):
            for col_idx, value in enumerate(row, start=1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                cell.alignment = center_align
        for col in range(1, 3):
            ws1.column_dimensions[get_column_letter(col)].width = 20

        ws2 = wb.create_sheet("明细数据")
        headers = ["包裹ID", "运单号", "收件人", "取件码", "到件日期", "状态",
                  "处置类型", "差异原因", "证据", "复核人", "复核备注", "对账状态"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        pkg_map = {p.package_id: p for p in packages}
        for row_idx, d in enumerate(disposals, start=2):
            pkg = pkg_map.get(d.package_id)
            if not pkg:
                continue
            row_data = [
                pkg.package_id, pkg.tracking_no, pkg.recipient_name,
                pkg.pickup_code, pkg.arrival_date.isoformat(), pkg.status.value,
                type_map.get(d.disposal_type, d.disposal_type.value),
                d.reason, " | ".join(d.evidence),
                d.reviewed_by or "", d.review_note or "",
                status_map.get(d.status, d.status.value)
            ]
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws2.cell(row=row_idx, column=col_idx, value=value)
                if col_idx == 7 and d.disposal_type == DisposalType.RETURN:
                    cell.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
                if col_idx == 7 and d.disposal_type == DisposalType.MANUAL_REVIEW:
                    cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
        widths = [12, 18, 10, 10, 12, 8, 10, 30, 40, 10, 20, 10]
        for col_idx, w in enumerate(widths, start=1):
            ws2.column_dimensions[get_column_letter(col_idx)].width = w

        if audit_logs:
            ws3 = wb.create_sheet("审计日志")
            log_headers = ["时间", "包裹ID", "操作", "操作员", "旧值", "新值", "备注"]
            for col_idx, h in enumerate(log_headers, start=1):
                cell = ws3.cell(row=1, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
            for row_idx, log in enumerate(audit_logs, start=2):
                ws3.cell(row=row_idx, column=1, value=log.timestamp.strftime("%Y-%m-%d %H:%M:%S"))
                ws3.cell(row=row_idx, column=2, value=log.package_id)
                ws3.cell(row=row_idx, column=3, value=log.action)
                ws3.cell(row=row_idx, column=4, value=log.operator)
                ws3.cell(row=row_idx, column=5, value=json.dumps(log.old_value, ensure_ascii=False) if log.old_value else "")
                ws3.cell(row=row_idx, column=6, value=json.dumps(log.new_value, ensure_ascii=False) if log.new_value else "")
                ws3.cell(row=row_idx, column=7, value=log.note or "")
            log_widths = [20, 12, 15, 10, 30, 30, 20]
            for col_idx, w in enumerate(log_widths, start=1):
                ws3.column_dimensions[get_column_letter(col_idx)].width = w

        wb.save(output)
        return output.getvalue()

    @staticmethod
    def export_evidence_report(package_id: str, package: Package,
                              disposal: DisposalRecord,
                              audit_logs: List[AuditLog]) -> Dict[str, Any]:
        type_map = {
            DisposalType.RELEASE: "放行",
            DisposalType.RETURN: "退回",
            DisposalType.SUPPLEMENT: "补充材料",
            DisposalType.MANUAL_REVIEW: "待人工复核"
        }
        return {
            "package_id": package_id,
            "recipient": package.recipient_name,
            "pickup_code": package.pickup_code,
            "arrival_date": package.arrival_date.isoformat(),
            "status": package.status.value,
            "disposal_type": type_map.get(disposal.disposal_type, disposal.disposal_type.value),
            "reasons": disposal.reason.split("; "),
            "evidence_chain": disposal.evidence,
            "reviewer": disposal.reviewed_by,
            "review_note": disposal.review_note,
            "audit_trail": [
                {
                    "time": log.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                    "action": log.action,
                    "operator": log.operator,
                    "change": f"{log.old_value} -> {log.new_value}" if log.old_value and log.new_value else log.action
                }
                for log in audit_logs
            ]
        }
