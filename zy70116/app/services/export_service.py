from datetime import datetime
from decimal import Decimal
from typing import List, Dict
import io

from openpyxl import Workbook
from sqlalchemy.orm import Session

from ..models.journal import AccountJournal
from ..models.settlement import StoreSettlement, SettlementDetail
from ..models.store import Store
from ..models.account import PrepaidAccount


class ExportService:
    @staticmethod
    def export_account_journals(db: Session, account_id: int) -> bytes:
        account = db.query(PrepaidAccount).filter(PrepaidAccount.id == account_id).first()
        if not account:
            raise ValueError("账户不存在")
        
        journals = db.query(AccountJournal).filter(
            AccountJournal.account_id == account_id
        ).order_by(AccountJournal.id.asc()).all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "账务流水"
        
        ws.append(["账户信息"])
        ws.append(["账号", account.account_no])
        ws.append(["会员", account.member_name or "-"])
        ws.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws.append([])
        
        ws.append(["流水号", "业务类型", "业务单号", "方向", 
                   "本金变动", "赠金变动", "合计变动",
                   "本金余额", "赠金余额", "总余额",
                   "操作时间", "操作人", "备注", "是否作废"])
        
        for j in journals:
            ws.append([
                j.journal_no,
                j.biz_type,
                j.biz_order_no,
                j.direction,
                float(j.principal_delta),
                float(j.bonus_delta),
                float(j.total_delta),
                float(j.principal_balance_after),
                float(j.bonus_balance_after),
                float(j.total_balance_after),
                j.journal_time.strftime("%Y-%m-%d %H:%M:%S") if j.journal_time else "",
                j.operator or "-",
                j.remark or "-",
                "是" if j.is_void else "否"
            ])
        
        ws.append([])
        ws.append(["统计核对"])
        total_principal_in = sum(j.principal_delta for j in journals if j.principal_delta > 0 and not j.is_void)
        total_principal_out = sum(-j.principal_delta for j in journals if j.principal_delta < 0 and not j.is_void)
        total_bonus_in = sum(j.bonus_delta for j in journals if j.bonus_delta > 0 and not j.is_void)
        total_bonus_out = sum(-j.bonus_delta for j in journals if j.bonus_delta < 0 and not j.is_void)
        
        ws.append(["本金收入", float(total_principal_in)])
        ws.append(["本金支出", float(total_principal_out)])
        ws.append(["赠金收入", float(total_bonus_in)])
        ws.append(["赠金支出", float(total_bonus_out)])
        ws.append(["当前本金余额", float(account.principal_balance)])
        ws.append(["当前赠金余额", float(account.bonus_balance)])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_settlement(db: Session, settlement_id: int) -> bytes:
        settlement = db.query(StoreSettlement).filter(StoreSettlement.id == settlement_id).first()
        if not settlement:
            raise ValueError("结算单不存在")
        
        store = db.query(Store).filter(Store.id == settlement.store_id).first()
        details = db.query(SettlementDetail).filter(
            SettlementDetail.settlement_id == settlement_id
        ).order_by(SettlementDetail.id.asc()).all()
        
        wb = Workbook()
        
        ws_summary = wb.active
        ws_summary.title = "结算汇总"
        
        ws_summary.append(["门店结算单"])
        ws_summary.append(["结算单号", settlement.settlement_no])
        ws_summary.append(["门店", store.store_name if store else "-"])
        ws_summary.append(["结算周期", settlement.settlement_period])
        ws_summary.append(["结算日期", settlement.settlement_date])
        ws_summary.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws_summary.append([])
        
        ws_summary.append(["业务类型", "笔数", "本金金额", "赠金金额"])
        ws_summary.append(["充值", settlement.deposit_count, float(settlement.deposit_amount), "-"])
        ws_summary.append(["消费", settlement.consume_count, float(settlement.consume_amount), "-"])
        ws_summary.append(["退款", settlement.refund_count, float(settlement.refund_amount), "-"])
        ws_summary.append([])
        ws_summary.append(["净结算金额", float(settlement.net_amount)])
        
        ws_details = wb.create_sheet("明细")
        ws_details.append(["序号", "业务类型", "单号", "本金金额", "赠金金额", "总金额", "备注"])
        
        for idx, d in enumerate(details, 1):
            ws_details.append([
                idx,
                d.biz_type,
                d.order_no,
                float(d.principal_amount),
                float(d.bonus_amount),
                float(d.total_amount),
                d.remark or "-"
            ])
        
        ws_details.append([])
        ws_details.append(["核对汇总"])
        total_principal = sum(d.principal_amount for d in details)
        total_bonus = sum(d.bonus_amount for d in details)
        total = sum(d.total_amount for d in details)
        ws_details.append(["本金合计", float(total_principal)])
        ws_details.append(["赠金合计", float(total_bonus)])
        ws_details.append(["总计", float(total)])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_refund_check(db: Session, start_date: str, end_date: str) -> bytes:
        from ..models.refund import RefundOrder
        
        refunds = db.query(RefundOrder).filter(
            RefundOrder.is_void == False,
            RefundOrder.refund_time >= datetime.strptime(start_date, "%Y-%m-%d"),
            RefundOrder.refund_time < datetime.strptime(end_date, "%Y-%m-%d")
        ).order_by(RefundOrder.id.asc()).all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "退款核对表"
        
        ws.append(["退款核对表"])
        ws.append(["统计周期", f"{start_date} 至 {end_date}"])
        ws.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws.append([])
        
        ws.append([
            "退款单号", "申请单号", "账号", 
            "退款总额", "本金退款", "赠金退款", "赠金没收",
            "退款时间", "状态", "备注"
        ])
        
        total_refund = Decimal("0.00")
        total_principal = Decimal("0.00")
        total_bonus = Decimal("0.00")
        total_forfeit = Decimal("0.00")
        
        for r in refunds:
            ws.append([
                r.order_no,
                r.request_id or "-",
                r.account_id,
                float(r.total_refund),
                float(r.principal_refund),
                float(r.bonus_refund),
                float(r.bonus_forfeit),
                r.refund_time.strftime("%Y-%m-%d %H:%M:%S") if r.refund_time else "",
                r.status,
                r.remark or "-"
            ])
            total_refund += r.total_refund
            total_principal += r.principal_refund
            total_bonus += r.bonus_refund
            total_forfeit += r.bonus_forfeit
        
        ws.append([])
        ws.append(["合计"])
        ws.append(["退款总额", float(total_refund)])
        ws.append(["本金退款", float(total_principal)])
        ws.append(["赠金退款", float(total_bonus)])
        ws.append(["赠金没收", float(total_forfeit)])
        ws.append([])
        ws.append(["核对公式: 退款总额 = 本金退款 + 赠金退款"])
        ws.append(["验证", "通过" if total_refund == total_principal + total_bonus else "不通过"])
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
