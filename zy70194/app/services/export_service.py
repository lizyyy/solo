from datetime import datetime
import os
from typing import Dict, Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app import db
from app.models import (
    Inquiry, Quote, QuoteItem, ComparisonResult,
    InquiryStatus, QuoteStatus
)
from config import BASE_DIR

class ExportService:
    EXPORT_DIR = os.path.join(BASE_DIR, 'exports')
    
    @staticmethod
    def _ensure_export_dir():
        os.makedirs(ExportService.EXPORT_DIR, exist_ok=True)
    
    @staticmethod
    def _get_style(style_type: str = None):
        header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        title_font = Font(name='微软雅黑', size=14, bold=True)
        content_font = Font(name='微软雅黑', size=11)
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        center_alignment = Alignment(horizontal='center', vertical='center')
        left_alignment = Alignment(horizontal='left', vertical='center')
        
        return {
            'header_font': header_font,
            'header_fill': header_fill,
            'header_alignment': header_alignment,
            'title_font': title_font,
            'content_font': content_font,
            'border': thin_border,
            'center_alignment': center_alignment,
            'left_alignment': left_alignment
        }
    
    @staticmethod
    def _format_number(num: float) -> str:
        if num is None:
            return ''
        return f'{num:,.2f}'
    
    @staticmethod
    def export_comparison_to_excel(inquiry_id: int, export_type: str = 'all') -> Dict[str, Any]:
        inquiry = Inquiry.query.get(inquiry_id)
        if not inquiry:
            return {
                'success': False,
                'error': '询价单不存在',
                'error_code': 'inquiry_not_found'
            }
        
        comparison = ComparisonResult.query.filter_by(
            inquiry_id=inquiry_id
        ).order_by(ComparisonResult.generated_at.desc()).first()
        
        if not comparison:
            return {
                'success': False,
                'error': '比价结果不存在，请先生成比价',
                'error_code': 'comparison_not_found'
            }
        
        try:
            ExportService._ensure_export_dir()
            styles = ExportService._get_style()
            
            wb = Workbook()
            
            ws_summary = wb.active
            ws_summary.title = '比价汇总'
            
            ws_summary['A1'] = '采购询价比价表'
            ws_summary['A1'].font = styles['title_font']
            ws_summary.merge_cells('A1:J1')
            ws_summary['A1'].alignment = styles['center_alignment']
            
            ws_summary['A3'] = '询价单编号:'
            ws_summary['B3'] = inquiry.inquiry_no
            ws_summary['A4'] = '询价标题:'
            ws_summary['B4'] = inquiry.title
            ws_summary['A5'] = '部门:'
            ws_summary['C5'] = inquiry.department or ''
            ws_summary['A6'] = '项目:'
            ws_summary['C6'] = inquiry.project or ''
            ws_summary['A7'] = '比价日期:'
            ws_summary['B7'] = comparison.generated_at.strftime('%Y-%m-%d %H:%M:%S')
            
            if comparison.recommended_vendor_name:
                ws_summary['A8'] = '推荐供应商:'
                ws_summary['B8'] = comparison.recommended_vendor_name
                ws_summary['A9'] = '推荐理由:'
                ws_summary['B9'] = comparison.recommendation_reason or ''
                ws_summary.merge_cells('B9:J9')
            
            row = 11
            headers = [
                '排名', '供应商', '报价编号', '货币', '税率',
                '不含税总价', '含税总价', '运费', '总金额', '有效期至'
            ]
            
            for col_idx, header in enumerate(headers, 1):
                cell = ws_summary.cell(row=row, column=col_idx, value=header)
                cell.font = styles['header_font']
                cell.fill = styles['header_fill']
                cell.alignment = styles['header_alignment']
                cell.border = styles['border']
            
            row += 1
            quotes_data = comparison.comparison_data.get('quotes', [])
            
            for quote in quotes_data:
                ws_summary.cell(row=row, column=1, value=quote.get('rank', '')).alignment = styles['center_alignment']
                ws_summary.cell(row=row, column=2, value=quote.get('vendor_name', '')).alignment = styles['left_alignment']
                ws_summary.cell(row=row, column=3, value=quote.get('quote_no', '')).alignment = styles['center_alignment']
                ws_summary.cell(row=row, column=4, value=quote.get('currency', '')).alignment = styles['center_alignment']
                ws_summary.cell(row=row, column=5, value=f"{quote.get('tax_rate', 0) * 100:.1f}%").alignment = styles['center_alignment']
                ws_summary.cell(row=row, column=6, value=ExportService._format_number(quote.get('total_price_excl_tax', 0)))
                ws_summary.cell(row=row, column=7, value=ExportService._format_number(quote.get('total_price_incl_tax', 0)))
                ws_summary.cell(row=row, column=8, value=ExportService._format_number(quote.get('total_freight', 0)))
                ws_summary.cell(row=row, column=9, value=ExportService._format_number(quote.get('total_amount', 0)))
                ws_summary.cell(row=row, column=10, value=quote.get('valid_until', '').split('T')[0] if quote.get('valid_until') else '')
                
                for col in range(1, 11):
                    cell = ws_summary.cell(row=row, column=col)
                    cell.font = styles['content_font']
                    cell.border = styles['border']
                
                row += 1
            
            for col in range(1, 11):
                ws_summary.column_dimensions[get_column_letter(col)].width = 15
            ws_summary.column_dimensions['B'].width = 25
            ws_summary.column_dimensions['C'].width = 20
            
            ws_items = wb.create_sheet('分项对比')
            item_comparison = comparison.comparison_data.get('item_comparison', [])
            
            if item_comparison:
                first_item = item_comparison[0]
                vendor_count = len(first_item.get('quotes', []))
                
                ws_items['A1'] = '分项价格对比'
                ws_items['A1'].font = styles['title_font']
                ws_items.merge_cells('A1:F1')
                ws_items['A1'].alignment = styles['center_alignment']
                
                headers = ['物料编码', '物料名称', '规格型号', '单位', '需求数量']
                for quote in first_item.get('quotes', []):
                    headers.extend([
                        f"{quote.get('vendor_name', '')} 单价",
                        f"{quote.get('vendor_name', '')} 小计"
                    ])
                
                for col_idx, header in enumerate(headers, 1):
                    cell = ws_items.cell(row=3, column=col_idx, value=header)
                    cell.font = styles['header_font']
                    cell.fill = styles['header_fill']
                    cell.alignment = styles['header_alignment']
                    cell.border = styles['border']
                
                row = 4
                for item in item_comparison:
                    ws_items.cell(row=row, column=1, value=item.get('item_no', '')).alignment = styles['center_alignment']
                    ws_items.cell(row=row, column=2, value=item.get('item_name', '')).alignment = styles['left_alignment']
                    ws_items.cell(row=row, column=3, value='').alignment = styles['left_alignment']
                    ws_items.cell(row=row, column=4, value='').alignment = styles['center_alignment']
                    ws_items.cell(row=row, column=5, value=item.get('required_quantity', '')).alignment = styles['center_alignment']
                    
                    col = 6
                    for quote_item in item.get('quotes', []):
                        ws_items.cell(row=row, column=col, value=ExportService._format_number(quote_item.get('unit_price_incl_tax', 0)))
                        ws_items.cell(row=row, column=col+1, value=ExportService._format_number(quote_item.get('line_total_incl_tax', 0)))
                        col += 2
                    
                    for col_idx in range(1, len(headers) + 1):
                        cell = ws_items.cell(row=row, column=col_idx)
                        cell.font = styles['content_font']
                        cell.border = styles['border']
                    
                    row += 1
                
                for col_idx in range(1, len(headers) + 1):
                    ws_items.column_dimensions[get_column_letter(col_idx)].width = 15
            
            ws_details = wb.create_sheet('报价详情')
            
            detail_headers = [
                '供应商', '报价编号', '物料编码', '物料名称', '规格型号',
                '单位', '数量', '不含税单价', '含税单价', '税率', '税额',
                '单位运费', '运费总额', '不含税小计', '含税小计'
            ]
            
            for col_idx, header in enumerate(detail_headers, 1):
                cell = ws_details.cell(row=1, column=col_idx, value=header)
                cell.font = styles['header_font']
                cell.fill = styles['header_fill']
                cell.alignment = styles['header_alignment']
                cell.border = styles['border']
            
            row = 2
            for quote in quotes_data:
                quote_items = QuoteItem.query.filter_by(quote_id=quote['quote_id']).all()
                
                for item in quote_items:
                    ws_details.cell(row=row, column=1, value=quote.get('vendor_name', '')).alignment = styles['left_alignment']
                    ws_details.cell(row=row, column=2, value=quote.get('quote_no', '')).alignment = styles['center_alignment']
                    ws_details.cell(row=row, column=3, value=item.item_no).alignment = styles['center_alignment']
                    ws_details.cell(row=row, column=4, value=item.item_name).alignment = styles['left_alignment']
                    ws_details.cell(row=row, column=5, value=item.specification or '').alignment = styles['left_alignment']
                    ws_details.cell(row=row, column=6, value=item.unit).alignment = styles['center_alignment']
                    ws_details.cell(row=row, column=7, value=item.quantity).alignment = styles['center_alignment']
                    ws_details.cell(row=row, column=8, value=ExportService._format_number(item.unit_price_excl_tax))
                    ws_details.cell(row=row, column=9, value=ExportService._format_number(item.unit_price_incl_tax))
                    ws_details.cell(row=row, column=10, value=f'{item.tax_rate * 100:.1f}%').alignment = styles['center_alignment']
                    ws_details.cell(row=row, column=11, value=ExportService._format_number(item.tax_amount))
                    ws_details.cell(row=row, column=12, value=ExportService._format_number(item.freight_per_unit))
                    ws_details.cell(row=row, column=13, value=ExportService._format_number(item.freight_total))
                    ws_details.cell(row=row, column=14, value=ExportService._format_number(item.line_total_excl_tax))
                    ws_details.cell(row=row, column=15, value=ExportService._format_number(item.line_total_incl_tax))
                    
                    for col in range(1, 16):
                        cell = ws_details.cell(row=row, column=col)
                        cell.font = styles['content_font']
                        cell.border = styles['border']
                    
                    row += 1
            
            for col in range(1, 16):
                ws_details.column_dimensions[get_column_letter(col)].width = 14
            ws_details.column_dimensions['B'].width = 20
            ws_details.column_dimensions['D'].width = 20
            
            filename = f'比价表_{inquiry.inquiry_no}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            file_path = os.path.join(ExportService.EXPORT_DIR, filename)
            
            wb.save(file_path)
            
            return {
                'success': True,
                'data': {
                    'file_path': file_path,
                    'filename': filename,
                    'inquiry_id': inquiry_id,
                    'inquiry_no': inquiry.inquiry_no
                }
            }
        except Exception as e:
            return {
                'success': False,
                'error': f'导出失败: {str(e)}',
                'error_code': 'export_failed'
            }
    
    @staticmethod
    def export_quote_to_excel(quote_id: int) -> Dict[str, Any]:
        quote = Quote.query.get(quote_id)
        if not quote:
            return {
                'success': False,
                'error': '报价单不存在',
                'error_code': 'quote_not_found'
            }
        
        try:
            ExportService._ensure_export_dir()
            styles = ExportService._get_style()
            
            wb = Workbook()
            ws = wb.active
            ws.title = '报价单'
            
            ws['A1'] = '报价单'
            ws['A1'].font = styles['title_font']
            ws.merge_cells('A1:H1')
            ws['A1'].alignment = styles['center_alignment']
            
            ws['A3'] = '报价编号:'
            ws['B3'] = quote.quote_no
            ws['A4'] = '供应商:'
            ws['B4'] = quote.vendor_name
            ws['A5'] = '联系人:'
            ws['B5'] = quote.vendor_contact or ''
            ws['C5'] = '联系电话:'
            ws['D5'] = quote.vendor_phone or ''
            ws['A6'] = '报价有效期:'
            ws['B6'] = f"{quote.valid_from.strftime('%Y-%m-%d')} 至 {quote.valid_until.strftime('%Y-%m-%d')}"
            ws['A7'] = '货币:'
            ws['B7'] = quote.currency
            ws['A8'] = '税率:'
            ws['B8'] = f'{quote.tax_rate * 100:.1f}%'
            ws['A9'] = '付款条件:'
            ws['B9'] = quote.payment_terms or ''
            ws['A10'] = '交货地点:'
            ws['B10'] = quote.delivery_location or ''
            ws['A11'] = '交货期:'
            ws['B11'] = quote.delivery_time or ''
            ws['A12'] = '备注:'
            ws['B12'] = quote.remarks or ''
            ws.merge_cells('B12:H12')
            
            headers = [
                '序号', '物料编码', '物料名称', '规格型号', '单位',
                '数量', '不含税单价', '含税单价', '税额',
                '单位运费', '运费总额', '不含税小计', '含税小计'
            ]
            
            row = 14
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx, value=header)
                cell.font = styles['header_font']
                cell.fill = styles['header_fill']
                cell.alignment = styles['header_alignment']
                cell.border = styles['border']
            
            row += 1
            items = quote.items.all()
            
            for idx, item in enumerate(items, 1):
                ws.cell(row=row, column=1, value=idx).alignment = styles['center_alignment']
                ws.cell(row=row, column=2, value=item.item_no).alignment = styles['center_alignment']
                ws.cell(row=row, column=3, value=item.item_name).alignment = styles['left_alignment']
                ws.cell(row=row, column=4, value=item.specification or '').alignment = styles['left_alignment']
                ws.cell(row=row, column=5, value=item.unit).alignment = styles['center_alignment']
                ws.cell(row=row, column=6, value=item.quantity).alignment = styles['center_alignment']
                ws.cell(row=row, column=7, value=ExportService._format_number(item.unit_price_excl_tax))
                ws.cell(row=row, column=8, value=ExportService._format_number(item.unit_price_incl_tax))
                ws.cell(row=row, column=9, value=ExportService._format_number(item.tax_amount))
                ws.cell(row=row, column=10, value=ExportService._format_number(item.freight_per_unit))
                ws.cell(row=row, column=11, value=ExportService._format_number(item.freight_total))
                ws.cell(row=row, column=12, value=ExportService._format_number(item.line_total_excl_tax))
                ws.cell(row=row, column=13, value=ExportService._format_number(item.line_total_incl_tax))
                
                for col in range(1, 14):
                    cell = ws.cell(row=row, column=col)
                    cell.font = styles['content_font']
                    cell.border = styles['border']
                
                row += 1
            
            total_row = row + 1
            ws.cell(row=total_row, column=12, value='总价(不含税):').font = styles['header_font']
            ws.cell(row=total_row, column=13, value=ExportService._format_number(quote.total_price_excl_tax)).font = styles['header_font']
            
            total_row += 1
            ws.cell(row=total_row, column=12, value='总价(含税):').font = styles['header_font']
            ws.cell(row=total_row, column=13, value=ExportService._format_number(quote.total_price_incl_tax)).font = styles['header_font']
            
            total_row += 1
            ws.cell(row=total_row, column=12, value='运费:').font = styles['header_font']
            ws.cell(row=total_row, column=13, value=ExportService._format_number(quote.total_freight)).font = styles['header_font']
            
            total_row += 1
            ws.cell(row=total_row, column=12, value='总计:').font = styles['header_font']
            ws.cell(row=total_row, column=13, value=ExportService._format_number(quote.total_amount)).font = styles['header_font']
            
            for col in range(1, 14):
                ws.column_dimensions[get_column_letter(col)].width = 12
            ws.column_dimensions['C'].width = 20
            
            filename = f'报价单_{quote.quote_no}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
            file_path = os.path.join(ExportService.EXPORT_DIR, filename)
            
            wb.save(file_path)
            
            return {
                'success': True,
                'data': {
                    'file_path': file_path,
                    'filename': filename,
                    'quote_id': quote_id,
                    'quote_no': quote.quote_no
                }
            }
        except Exception as e:
            return {
                'success': False,
                'error': f'导出失败: {str(e)}',
                'error_code': 'export_failed'
            }
