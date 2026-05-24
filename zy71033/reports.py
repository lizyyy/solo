from datetime import datetime
from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import models


def generate_shipping_excel(db: Session, order_no: str) -> BytesIO:
    order = db.query(models.SalesOrder).filter(
        models.SalesOrder.order_no == order_no
    ).first()

    if not order:
        raise ValueError("订单不存在")

    wb = Workbook()
    ws = wb.active
    ws.title = "发货报告"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    ws["A1"] = "苗圃检疫发货报告"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = center_align

    ws["A3"] = "订单编号:"
    ws["B3"] = order.order_no
    ws["A4"] = "客户:"
    ws["B4"] = order.customer
    ws["A5"] = "苗木品种:"
    ws["B5"] = order.batch.species
    ws["A6"] = "数量:"
    ws["B6"] = order.quantity
    ws["A7"] = "目的地:"
    ws["B7"] = f"{order.destination.region_name} ({order.destination.region_code})"
    ws["A8"] = "检疫证号:"
    ws["B8"] = order.certificate.cert_no if order.certificate else "无"
    ws["A9"] = "订单状态:"
    ws["B9"] = order.status.value

    headers = ["时间", "操作", "操作员", "详情", "结果"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=11, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    traces = db.query(models.ProcessingTrace).filter(
        models.ProcessingTrace.order_id == order.id
    ).order_by(models.ProcessingTrace.created_at.asc()).all()

    for row, trace in enumerate(traces, 12):
        ws.cell(row=row, column=1, value=trace.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row, column=2, value=trace.action.value)
        ws.cell(row=row, column=3, value=trace.operator)
        ws.cell(row=row, column=4, value=trace.detail)
        ws.cell(row=row, column=5, value=trace.result)

    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 25

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_batch_report(db: Session, start_date: datetime = None,
                          end_date: datetime = None) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "批次检疫汇总"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = ["批次号", "品种", "数量", "检疫证号", "有效期", "订单数", "已发货", "状态"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    query = db.query(models.SeedlingBatch)
    if start_date:
        query = query.filter(models.SeedlingBatch.created_at >= start_date)
    if end_date:
        query = query.filter(models.SeedlingBatch.created_at <= end_date)

    batches = query.all()

    for row, batch in enumerate(batches, 2):
        cert = batch.certificates[0] if batch.certificates else None
        order_count = len(batch.sales_orders)
        shipped_count = sum(1 for o in batch.sales_orders
                            if o.status == models.SalesOrderStatus.SHIPPED)

        ws.cell(row=row, column=1, value=batch.batch_no)
        ws.cell(row=row, column=2, value=batch.species)
        ws.cell(row=row, column=3, value=batch.quantity)
        ws.cell(row=row, column=4, value=cert.cert_no if cert else "无")
        ws.cell(row=row, column=5,
                value=cert.expiry_date.strftime("%Y-%m-%d") if cert else "")
        ws.cell(row=row, column=6, value=order_count)
        ws.cell(row=row, column=7, value=shipped_count)
        ws.cell(row=row, column=8,
                value="有效" if cert and cert.status == models.CertificateStatus.VALID else "无效")

    for col in range(1, 9):
        ws.column_dimensions[chr(64 + col)].width = 18

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
