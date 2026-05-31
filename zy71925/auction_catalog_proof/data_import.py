"""
数据导入模块
============

支持从 Excel 和 CSV 文件导入作品清单和展墙图数据
"""

import os
import csv
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime

from .models import Artwork, WallLayout, HistoryEntry, CatalogProofSession
from .messages import format_import_error

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ImportError(Exception):
    def __init__(self, message: str, error_type: str, details: str = ""):
        super().__init__(message)
        self.error_type = error_type
        self.details = details
        self.user_message = message


def _read_csv(file_path: str) -> List[Dict[str, Any]]:
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb2312"]
    last_error = None

    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                reader = csv.DictReader(f)
                rows = [dict(row) for row in reader]
                if rows:
                    return rows
        except UnicodeDecodeError as e:
            last_error = e
            continue
        except Exception as e:
            last_error = e
            continue

    raise ImportError(
        format_import_error(file_path, "parse_error", str(last_error)),
        "parse_error",
        str(last_error),
    )


def _read_excel(file_path: str) -> List[Dict[str, Any]]:
    if not HAS_OPENPYXL:
        raise ImportError(
            format_import_error(
                file_path,
                "format_unsupported",
                "需要安装 openpyxl 才能读取 Excel 文件：pip install openpyxl",
            ),
            "format_unsupported",
            "openpyxl not installed",
        )

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        headers = []
        for cell in ws[1]:
            if cell.value is not None:
                headers.append(str(cell.value).strip())
            else:
                headers.append("")

        rows = []
        for row_idx in range(2, ws.max_row + 1):
            row_data = {}
            for col_idx, header in enumerate(headers):
                if not header:
                    continue
                cell = ws.cell(row=row_idx, column=col_idx + 1)
                value = cell.value
                if value is not None:
                    row_data[header] = str(value).strip()
                else:
                    row_data[header] = ""

            if any(v.strip() for v in row_data.values()):
                rows.append(row_data)

        return rows
    except Exception as e:
        raise ImportError(
            format_import_error(file_path, "parse_error", str(e)),
            "parse_error",
            str(e),
        )


def _read_file(file_path: str) -> List[Dict[str, Any]]:
    if not os.path.exists(file_path):
        raise ImportError(
            format_import_error(file_path, "file_not_found", ""),
            "file_not_found",
            "",
        )

    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".xlsx", ".xls"]:
        rows = _read_excel(file_path)
    elif ext == ".csv":
        rows = _read_csv(file_path)
    else:
        raise ImportError(
            format_import_error(file_path, "format_unsupported", f"不支持的格式：{ext}"),
            "format_unsupported",
            ext,
        )

    if not rows:
        raise ImportError(
            format_import_error(file_path, "empty_file", ""),
            "empty_file",
            "",
        )

    return rows


ARTWORK_COLUMN_MAPPING = {
    "拍品编号": "lot_number",
    "Lot": "lot_number",
    "LOT": "lot_number",
    "lot": "lot_number",
    "作品名称": "title_cn",
    "中文名称": "title_cn",
    "Title": "title_cn",
    "英文名称": "title_en",
    "Title EN": "title_en",
    "Artist EN": "artist_en",
    "艺术家": "artist",
    "作者": "artist",
    "Artist": "artist",
    "年代": "year",
    "创作年代": "year",
    "Year": "year",
    "材质": "medium",
    "材质工艺": "medium",
    "Medium": "medium",
    "尺寸": "dimensions",
    "Dimensions": "dimensions",
    "估价": "estimate",
    "Estimate": "estimate",
    "来源": "provenance",
    "Provenance": "provenance",
    "出版": "literature",
    "文献": "literature",
    "Literature": "literature",
    "展览": "exhibition",
    "展览历史": "exhibition",
    "Exhibition": "exhibition",
    "说明": "description",
    "作品说明": "description",
    "Description": "description",
    "备注": "notes",
    "Notes": "notes",
    "灯光": "lighting_scheme",
    "灯光方案": "lighting_scheme",
    "Lighting": "lighting_scheme",
}

WALL_LAYOUT_COLUMN_MAPPING = {
    "展墙编号": "wall_id",
    "Wall ID": "wall_id",
    "展墙名称": "wall_name",
    "Wall Name": "wall_name",
    "拍品编号": "lot_number",
    "Lot": "lot_number",
    "LOT": "lot_number",
    "lot": "lot_number",
    "X坐标": "position_x",
    "X": "position_x",
    "Y坐标": "position_y",
    "Y": "position_y",
    "宽度": "width",
    "Width": "width",
    "高度": "height",
    "Height": "height",
    "灯光": "lighting_scheme",
    "灯光方案": "lighting_scheme",
    "Lighting": "lighting_scheme",
    "顺序": "display_sequence",
    "展示顺序": "display_sequence",
    "Sequence": "display_sequence",
    "备注": "notes",
    "Notes": "notes",
}


def _map_columns(row: Dict[str, Any], mapping: Dict[str, str]) -> Dict[str, Any]:
    mapped = {}
    for header, value in row.items():
        header_clean = str(header).strip()
        if header_clean in mapping:
            field_name = mapping[header_clean]
            if value is not None and str(value).strip() != "":
                mapped[field_name] = str(value).strip()
    return mapped


def _check_required_columns(headers: List[str], required: List[str], mapping: Dict[str, str]) -> Tuple[bool, List[str]]:
    mapped_headers = set()
    for h in headers:
        if h in mapping:
            mapped_headers.add(mapping[h])

    missing = []
    for req in required:
        if req not in mapped_headers:
            original_names = [k for k, v in mapping.items() if v == req]
            missing.append(original_names[0] if original_names else req)

    return len(missing) == 0, missing


def import_works_list(file_path: str, session: Optional[CatalogProofSession] = None) -> Tuple[Dict[str, Artwork], List[HistoryEntry]]:
    rows = _read_file(file_path)
    headers = list(rows[0].keys())

    ok, missing = _check_required_columns(
        headers, ["lot_number", "title_cn"], ARTWORK_COLUMN_MAPPING
    )
    if not ok:
        raise ImportError(
            format_import_error(
                file_path,
                "missing_column",
                f"缺少必要的列：{', '.join(missing)}",
            ),
            "missing_column",
            ", ".join(missing),
        )

    artworks: Dict[str, Artwork] = {}
    history: List[HistoryEntry] = []

    for row in rows:
        mapped = _map_columns(row, ARTWORK_COLUMN_MAPPING)
        if "lot_number" not in mapped or "title_cn" not in mapped:
            continue

        lot = mapped["lot_number"]
        artwork = Artwork(
            lot_number=lot,
            title_cn=mapped.get("title_cn", ""),
            title_en=mapped.get("title_en"),
            artist=mapped.get("artist"),
            artist_en=mapped.get("artist_en"),
            year=mapped.get("year"),
            medium=mapped.get("medium"),
            dimensions=mapped.get("dimensions"),
            estimate=mapped.get("estimate"),
            provenance=mapped.get("provenance"),
            literature=mapped.get("literature"),
            exhibition=mapped.get("exhibition"),
            description=mapped.get("description"),
            notes=mapped.get("notes"),
            source="works_list",
        )
        artworks[lot] = artwork

        history.append(
            HistoryEntry(
                lot_number=lot,
                action="import",
                notes=f"从作品清单导入：{artwork.title_cn}",
                operator="system",
            )
        )

    if session is not None:
        for lot, artwork in artworks.items():
            session.artworks[lot] = artwork
        session.history.extend(history)
        session.works_list_source = file_path

    return artworks, history


def import_wall_layout(file_path: str, session: Optional[CatalogProofSession] = None) -> Tuple[Dict[str, WallLayout], List[HistoryEntry]]:
    rows = _read_file(file_path)
    headers = list(rows[0].keys())

    ok, missing = _check_required_columns(
        headers, ["wall_id", "lot_number"], WALL_LAYOUT_COLUMN_MAPPING
    )
    if not ok:
        raise ImportError(
            format_import_error(
                file_path,
                "missing_column",
                f"缺少必要的列：{', '.join(missing)}",
            ),
            "missing_column",
            ", ".join(missing),
        )

    wall_layouts: Dict[str, WallLayout] = {}
    history: List[HistoryEntry] = []

    for row in rows:
        mapped = _map_columns(row, WALL_LAYOUT_COLUMN_MAPPING)
        if "wall_id" not in mapped or "lot_number" not in mapped:
            continue

        key = f"{mapped['wall_id']}_{mapped['lot_number']}"

        def _parse_float(value: Optional[str]) -> Optional[float]:
            if value is None or value == "":
                return None
            try:
                return float(value)
            except ValueError:
                return None

        def _parse_int(value: Optional[str]) -> Optional[int]:
            if value is None or value == "":
                return None
            try:
                return int(float(value))
            except ValueError:
                return None

        wall_layout = WallLayout(
            wall_id=mapped["wall_id"],
            wall_name=mapped.get("wall_name", mapped["wall_id"]),
            lot_number=mapped["lot_number"],
            position_x=_parse_float(mapped.get("position_x")),
            position_y=_parse_float(mapped.get("position_y")),
            width=_parse_float(mapped.get("width")),
            height=_parse_float(mapped.get("height")),
            lighting_scheme=mapped.get("lighting_scheme"),
            lighting_source="wall_layout",
            display_sequence=_parse_int(mapped.get("display_sequence")),
            notes=mapped.get("notes"),
        )
        wall_layouts[key] = wall_layout

        history.append(
            HistoryEntry(
                lot_number=wall_layout.lot_number,
                action="import",
                notes=f"从展墙图导入：{wall_layout.wall_name}",
                operator="system",
            )
        )

    if session is not None:
        for key, layout in wall_layouts.items():
            session.wall_layouts[key] = layout
        session.history.extend(history)
        session.wall_layout_source = file_path

    return wall_layouts, history


def create_sample_works_list() -> List[Dict[str, str]]:
    return [
        {
            "拍品编号": "LOT001",
            "作品名称": "山水清音",
            "艺术家": "张大千",
            "年代": "1947年",
            "材质": "设色纸本 立轴",
            "尺寸": "136×68cm",
            "估价": "RMB 800,000-1,200,000",
            "来源": "北京文物商店旧藏",
            "灯光": "暖白光 3000K，重点照明",
        },
        {
            "拍品编号": "LOT002",
            "作品名称": "行书七言联",
            "艺术家": "齐白石",
            "年代": "1938年",
            "材质": "水墨纸本 对联",
            "尺寸": "132×32cm×2",
            "估价": "RMB 500,000-800,000",
            "展览历史": "2018年齐白石书法展，中国美术馆",
            "灯光": "冷白光 4000K，漫射照明",
        },
        {
            "拍品编号": "LOT003",
            "作品名称": "秋山策杖图",
            "艺术家": "黄宾虹",
            "年代": "1952年",
            "材质": "设色纸本 镜框",
            "尺寸": "68×45cm",
            "估价": "RMB 300,000-500,000",
            "文献": "《黄宾虹全集》卷五，2012年",
        },
        {
            "拍品编号": "LOT004",
            "作品名称": "幽兰图",
            "艺术家": "潘天寿",
            "年代": "1961年",
            "材质": "水墨纸本 立轴",
            "尺寸": "95×42cm",
            "估价": "RMB 280,000-400,000",
            "灯光": "暖白光 3000K，标准照度",
        },
        {
            "拍品编号": "LOT005",
            "作品名称": "青绿山水",
            "艺术家": "吴湖帆",
            "年代": "1935年",
            "材质": "设色绢本 手卷",
            "尺寸": "28×186cm",
            "估价": "RMB 1,500,000-2,000,000",
            "来源": "沪上名家旧藏",
            "灯光": "低照度 UV防护，漫射光",
        },
    ]


def create_sample_wall_layout() -> List[Dict[str, str]]:
    return [
        {
            "展墙编号": "W1",
            "展墙名称": "主入口展墙",
            "拍品编号": "LOT001",
            "X坐标": "1.5",
            "Y坐标": "1.8",
            "宽度": "1.8",
            "高度": "0.9",
            "灯光": "暖白光 3000K，重点照明",
            "顺序": "1",
        },
        {
            "展墙编号": "W1",
            "展墙名称": "主入口展墙",
            "拍品编号": "LOT002",
            "X坐标": "4.0",
            "Y坐标": "1.8",
            "宽度": "1.5",
            "高度": "1.8",
            "灯光": "冷白光 4000K，重点照明",
            "顺序": "2",
        },
        {
            "展墙编号": "W2",
            "展墙名称": "东侧展墙",
            "拍品编号": "LOT003",
            "X坐标": "1.2",
            "Y坐标": "1.6",
            "宽度": "0.9",
            "高度": "0.7",
            "灯光": "暖白光 3000K，标准照度",
            "顺序": "1",
        },
        {
            "展墙编号": "W2",
            "展墙名称": "东侧展墙",
            "拍品编号": "LOT004",
            "X坐标": "2.8",
            "Y坐标": "1.7",
            "宽度": "0.7",
            "高度": "1.2",
            "灯光": "暖白光 3000K，标准照度",
            "顺序": "2",
        },
        {
            "展墙编号": "W3",
            "展墙名称": "西侧展墙",
            "拍品编号": "LOT005",
            "X坐标": "1.0",
            "Y坐标": "1.2",
            "宽度": "0.5",
            "高度": "2.5",
            "灯光": "标准照明",
            "顺序": "1",
        },
    ]


def write_sample_data(output_dir: str) -> None:
    os.makedirs(output_dir, exist_ok=True)

    works_data = create_sample_works_list()
    works_fieldnames = []
    for row in works_data:
        for key in row.keys():
            if key not in works_fieldnames:
                works_fieldnames.append(key)

    works_path = os.path.join(output_dir, "sample_works_list.csv")
    with open(works_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=works_fieldnames)
        writer.writeheader()
        writer.writerows(works_data)

    wall_data = create_sample_wall_layout()
    wall_fieldnames = []
    for row in wall_data:
        for key in row.keys():
            if key not in wall_fieldnames:
                wall_fieldnames.append(key)

    wall_path = os.path.join(output_dir, "sample_wall_layout.csv")
    with open(wall_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=wall_fieldnames)
        writer.writeheader()
        writer.writerows(wall_data)

    print(f"示例数据已生成：")
    print(f"  作品清单：{works_path}")
    print(f"  展墙图：{wall_path}")
