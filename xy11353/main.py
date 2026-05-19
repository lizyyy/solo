from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime, timedelta
from typing import List, Optional
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import models, schemas, crud
from database import engine, get_db

models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="园区安保系统", description="访客预约、临时车牌、黑名单核验系统")


@app.get("/")
def root():
    return {"message": "园区安保系统 API 服务运行中", "docs": "/docs"}


@app.post("/visitors/", response_model=schemas.Visitor)
def create_visitor(visitor: schemas.VisitorCreate, db: Session = Depends(get_db)):
    db_visitor = crud.get_visitor_by_id_card(db, id_card=visitor.id_card)
    if db_visitor:
        raise HTTPException(status_code=400, detail="该身份证已存在预约记录")
    return crud.create_visitor(db=db, visitor=visitor)


@app.get("/visitors/", response_model=List[schemas.Visitor])
def read_visitors(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    visitors = crud.get_visitors(db, skip=skip, limit=limit)
    return visitors


@app.get("/visitors/{visitor_id}", response_model=schemas.Visitor)
def read_visitor(visitor_id: int, db: Session = Depends(get_db)):
    db_visitor = crud.get_visitor(db, visitor_id=visitor_id)
    if db_visitor is None:
        raise HTTPException(status_code=404, detail="访客不存在")
    return db_visitor


@app.put("/visitors/{visitor_id}", response_model=schemas.Visitor)
def update_visitor(visitor_id: int, visitor_update: schemas.VisitorUpdate, db: Session = Depends(get_db)):
    db_visitor = crud.update_visitor(db, visitor_id=visitor_id, visitor_update=visitor_update)
    if db_visitor is None:
        raise HTTPException(status_code=404, detail="访客不存在")
    return db_visitor


@app.delete("/visitors/{visitor_id}")
def delete_visitor(visitor_id: int, db: Session = Depends(get_db)):
    db_visitor = crud.delete_visitor(db, visitor_id=visitor_id)
    if db_visitor is None:
        raise HTTPException(status_code=404, detail="访客不存在")
    return {"message": "删除成功"}


@app.post("/visitors/batch/", response_model=schemas.BatchResult)
def create_visitors_batch(visitors: List[schemas.VisitorCreate], db: Session = Depends(get_db)):
    successful = []
    failed = []
    for visitor in visitors:
        try:
            db_visitor = crud.get_visitor_by_id_card(db, id_card=visitor.id_card)
            if db_visitor:
                failed.append({"id_card": visitor.id_card, "name": visitor.name, "error": "身份证已存在"})
                continue
            created = crud.create_visitor(db=db, visitor=visitor)
            successful.append({"id": created.id, "id_card": created.id_card, "name": created.name})
        except Exception as e:
            failed.append({"id_card": visitor.id_card, "name": visitor.name, "error": str(e)})
    return schemas.BatchResult(
        total=len(visitors),
        success_count=len(successful),
        failure_count=len(failed),
        successful=successful,
        failed=failed
    )


@app.post("/temporary-plates/", response_model=schemas.TemporaryPlate)
def create_temporary_plate(plate: schemas.TemporaryPlateCreate, db: Session = Depends(get_db)):
    return crud.create_temporary_plate(db=db, plate=plate)


@app.get("/temporary-plates/{plate_number}", response_model=schemas.TemporaryPlate)
def read_temporary_plate(plate_number: str, db: Session = Depends(get_db)):
    db_plate = crud.get_temporary_plate_by_number(db, plate_number=plate_number)
    if db_plate is None:
        raise HTTPException(status_code=404, detail="临时车牌不存在")
    return db_plate


@app.post("/blacklist/", response_model=schemas.Blacklist)
def create_blacklist_entry(blacklist: schemas.BlacklistCreate, db: Session = Depends(get_db)):
    return crud.create_blacklist_entry(db=db, blacklist=blacklist)


@app.get("/blacklist/check/{identifier}")
def check_blacklist(identifier: str, db: Session = Depends(get_db)):
    is_blocked = crud.is_blacklisted(db, identifier=identifier)
    entry = crud.get_blacklist_by_identifier(db, identifier=identifier)
    return {"is_blacklisted": is_blocked, "entry": entry}


@app.post("/blacklist/batch/", response_model=schemas.BatchResult)
def create_blacklist_batch(entries: List[schemas.BlacklistCreate], db: Session = Depends(get_db)):
    successful = []
    failed = []
    for entry in entries:
        try:
            existing = crud.get_blacklist_by_identifier(db, identifier=entry.identifier)
            if existing:
                failed.append({"identifier": entry.identifier, "error": "已在黑名单中"})
                continue
            created = crud.create_blacklist_entry(db=db, blacklist=entry)
            successful.append({"id": created.id, "identifier": created.identifier})
        except Exception as e:
            failed.append({"identifier": entry.identifier, "error": str(e)})
    return schemas.BatchResult(
        total=len(entries),
        success_count=len(successful),
        failure_count=len(failed),
        successful=successful,
        failed=failed
    )


@app.post("/verify/", response_model=schemas.VerificationResponse)
def verify_visitor(request: schemas.VerificationRequest, db: Session = Depends(get_db)):
    visitor = None
    exception_type = schemas.ExceptionType.NO_EXCEPTION
    status = schemas.VerificationStatus.APPROVED
    message = "核验通过"

    if request.id_card:
        if crud.is_blacklisted(db, identifier=request.id_card):
            exception_type = schemas.ExceptionType.BLACKLISTED
            status = schemas.VerificationStatus.REJECTED
            message = "身份证在黑名单中，拒绝放行"
            visitor = crud.get_visitor_by_id_card(db, id_card=request.id_card)
        else:
            visitor = crud.get_visitor_by_id_card(db, id_card=request.id_card)

    if request.plate_number and not visitor and exception_type == schemas.ExceptionType.NO_EXCEPTION:
        if crud.is_blacklisted(db, identifier=request.plate_number):
            exception_type = schemas.ExceptionType.BLACKLISTED
            status = schemas.VerificationStatus.REJECTED
            message = "车牌在黑名单中，拒绝放行"
        else:
            visitor = crud.get_visitor_by_plate(db, plate_number=request.plate_number)
            temp_plate = crud.get_temporary_plate_by_number(db, plate_number=request.plate_number)
            if temp_plate and not visitor:
                from sqlalchemy.orm import Session
                visitor = db.query(models.Visitor).filter(models.Visitor.id == temp_plate.visitor_id).first()

    if not visitor and exception_type == schemas.ExceptionType.NO_EXCEPTION:
        exception_type = schemas.ExceptionType.INVALID_PLATE
        status = schemas.VerificationStatus.REJECTED
        message = "未找到有效预约记录"

    if visitor and exception_type == schemas.ExceptionType.NO_EXCEPTION:
        now = datetime.now()
        if visitor.expected_end < now:
            exception_type = schemas.ExceptionType.EXPIRED_VISIT
            status = schemas.VerificationStatus.REJECTED
            message = "预约已过期"
        elif visitor.expected_start > now:
            status = schemas.VerificationStatus.PENDING
            message = "预约尚未生效"

        if request.gate_number and visitor.gate_number and request.gate_number != visitor.gate_number:
            exception_type = schemas.ExceptionType.WRONG_GATE
            status = schemas.VerificationStatus.REJECTED
            message = f"应从{visitor.gate_number}进入，当前在{request.gate_number}"

    record = schemas.VerificationRecordCreate(
        visitor_id=visitor.id if visitor else None,
        plate_number=request.plate_number,
        id_card=request.id_card,
        gate_number=request.gate_number,
        verified_by=request.verified_by,
        status=status,
        exception_type=exception_type,
        exception_details=message
    )
    db_record = crud.create_verification_record(db=db, record=record)

    return schemas.VerificationResponse(
        success=(status == schemas.VerificationStatus.APPROVED),
        status=status,
        exception_type=exception_type,
        message=message,
        visitor=visitor,
        record_id=db_record.id
    )


@app.get("/verification-records/", response_model=List[schemas.VerificationRecord])
def read_verification_records(
    responsible_person: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    status: Optional[schemas.VerificationStatus] = None,
    exception_type: Optional[schemas.ExceptionType] = None,
    gate_number: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    filters = schemas.QueryFilters(
        responsible_person=responsible_person,
        start_date=start_date,
        end_date=end_date,
        status=status,
        exception_type=exception_type,
        gate_number=gate_number
    )
    return crud.query_verification_records(db, filters=filters, skip=skip, limit=limit)


@app.get("/verification-records/export/")
def export_verification_records(
    responsible_person: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    status: Optional[schemas.VerificationStatus] = None,
    exception_type: Optional[schemas.ExceptionType] = None,
    gate_number: Optional[str] = None,
    db: Session = Depends(get_db)
):
    filters = schemas.QueryFilters(
        responsible_person=responsible_person,
        start_date=start_date,
        end_date=end_date,
        status=status,
        exception_type=exception_type,
        gate_number=gate_number
    )
    records = crud.query_verification_records(db, filters=filters, skip=0, limit=10000)

    wb = Workbook()
    ws = wb.active
    ws.title = "核验记录"

    headers = ["ID", "访客ID", "车牌号", "身份证号", "门岗编号", "核验时间", "核验人", "状态", "异常类型", "异常详情", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    for row_idx, record in enumerate(records, 2):
        ws.cell(row=row_idx, column=1, value=record.id)
        ws.cell(row=row_idx, column=2, value=record.visitor_id)
        ws.cell(row=row_idx, column=3, value=record.plate_number)
        ws.cell(row=row_idx, column=4, value=record.id_card)
        ws.cell(row=row_idx, column=5, value=record.gate_number)
        ws.cell(row=row_idx, column=6, value=record.verified_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row_idx, column=7, value=record.verified_by)
        ws.cell(row=row_idx, column=8, value=record.status.value if record.status else "")
        ws.cell(row=row_idx, column=9, value=record.exception_type.value if record.exception_type else "")
        ws.cell(row=row_idx, column=10, value=record.exception_details)
        ws.cell(row=row_idx, column=11, value=record.notes)

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    from urllib.parse import quote
    filename = f"verification_records_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"}
    )


@app.get("/verification-records/{record_id}", response_model=schemas.VerificationRecord)
def read_verification_record(record_id: int, db: Session = Depends(get_db)):
    db_record = crud.get_verification_record(db, record_id=record_id)
    if db_record is None:
        raise HTTPException(status_code=404, detail="核验记录不存在")
    return db_record


@app.get("/stats/summary")
def get_stats_summary(db: Session = Depends(get_db)):
    today = datetime.now().date()
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())

    total_visitors = db.query(models.Visitor).count()
    total_records = db.query(models.VerificationRecord).count()
    today_records = db.query(models.VerificationRecord).filter(
        models.VerificationRecord.verified_at >= today_start,
        models.VerificationRecord.verified_at <= today_end
    ).count()
    approved_today = db.query(models.VerificationRecord).filter(
        models.VerificationRecord.verified_at >= today_start,
        models.VerificationRecord.verified_at <= today_end,
        models.VerificationRecord.status == models.VerificationStatus.APPROVED
    ).count()
    rejected_today = db.query(models.VerificationRecord).filter(
        models.VerificationRecord.verified_at >= today_start,
        models.VerificationRecord.verified_at <= today_end,
        models.VerificationRecord.status == models.VerificationStatus.REJECTED
    ).count()
    blacklist_count = db.query(models.Blacklist).filter(models.Blacklist.is_active == True).count()

    return {
        "total_visitors": total_visitors,
        "total_verification_records": total_records,
        "today_verifications": today_records,
        "today_approved": approved_today,
        "today_rejected": rejected_today,
        "active_blacklist_entries": blacklist_count
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
