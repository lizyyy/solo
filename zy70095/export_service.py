from datetime import datetime, date
from typing import List, Optional
from sqlalchemy.orm import Session
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

import models
import schemas

class ExportService:
    @staticmethod
    def export_plan_to_csv(
        db: Session,
        plan_id: int
    ) -> bytes:
        plan = db.query(models.ChargePlan).filter(
            models.ChargePlan.id == plan_id
        ).first()
        if not plan:
            raise ValueError("计划不存在")
        
        output = io.StringIO()
        writer = csv.writer(output)
        
        writer.writerow(["储能充放电计划 - 业务复核报表"])
        writer.writerow(["生成时间:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        writer.writerow([])
        
        writer.writerow(["【计划基本信息】"])
        writer.writerow(["计划ID", plan.id])
        writer.writerow(["电站ID", plan.station_id])
        writer.writerow(["计划日期", plan.plan_date])
        writer.writerow(["计划名称", plan.plan_name])
        writer.writerow(["描述", plan.description or ""])
        writer.writerow(["版本", plan.version])
        writer.writerow(["状态", plan.status])
        writer.writerow(["创建人", plan.created_by or "系统"])
        writer.writerow(["创建时间", plan.created_at.strftime("%Y-%m-%d %H:%M:%S")])
        writer.writerow([])
        
        soc_constraint = plan.soc_constraint
        writer.writerow(["【SOC 约束信息】"])
        writer.writerow(["最小SOC(%)", soc_constraint.min_soc])
        writer.writerow(["最大SOC(%)", soc_constraint.max_soc])
        writer.writerow(["初始SOC(%)", soc_constraint.initial_soc])
        writer.writerow(["目标SOC(%)", soc_constraint.target_soc or ""])
        writer.writerow([])
        
        writer.writerow(["【计划时段明细】"])
        writer.writerow([
            "时段序号", "开始时间", "结束时间", "操作类型",
            "功率(kW)", "电量(kWh)", "预期SOC(%)", "关联电价窗口"
        ])
        
        for i, segment in enumerate(sorted(plan.segments, key=lambda s: s.start_time), 1):
            op_type_name = {
                "charge": "充电",
                "discharge": "放电",
                "idle": "待机"
            }.get(segment.operation_type, segment.operation_type)
            
            window_info = ""
            if segment.price_window:
                window_info = f"{segment.price_window.window_type} ({segment.price_window.price_per_kwh}元/kWh)"
            
            writer.writerow([
                i,
                segment.start_time.strftime("%Y-%m-%d %H:%M:%S"),
                segment.end_time.strftime("%Y-%m-%d %H:%M:%S"),
                op_type_name,
                segment.power_kw,
                segment.energy_kwh,
                segment.expected_soc,
                window_info
            ])
        writer.writerow([])
        
        if plan.execution_receipts:
            writer.writerow(["【执行回执明细】"])
            writer.writerow([
                "回执ID", "时段ID", "实际开始时间", "实际结束时间",
                "实际功率(kW)", "实际电量(kWh)", "实际SOC(%)",
                "设备状态", "备注"
            ])
            
            for receipt in plan.execution_receipts:
                status_name = {
                    "available": "可用",
                    "maintenance": "维护",
                    "fault": "故障",
                    "limited": "受限"
                }.get(receipt.equipment_status, receipt.equipment_status)
                
                writer.writerow([
                    receipt.id,
                    receipt.segment_id,
                    receipt.actual_start_time.strftime("%Y-%m-%d %H:%M:%S"),
                    receipt.actual_end_time.strftime("%Y-%m-%d %H:%M:%S"),
                    receipt.actual_power_kw,
                    receipt.actual_energy_kwh,
                    receipt.actual_soc,
                    status_name,
                    receipt.remarks or ""
                ])
            writer.writerow([])
        
        if plan.alerts:
            writer.writerow(["【偏差告警记录】"])
            writer.writerow([
                "告警ID", "告警类型", "告警级别", "消息",
                "偏差值", "阈值", "是否已处理", "处理人", "创建时间"
            ])
            
            for alert in plan.alerts:
                level_name = {
                    "info": "信息",
                    "warning": "警告",
                    "critical": "严重"
                }.get(alert.alert_level, alert.alert_level)
                
                writer.writerow([
                    alert.id,
                    alert.alert_type,
                    level_name,
                    alert.message,
                    alert.deviation_value if alert.deviation_value is not None else "",
                    alert.threshold_value if alert.threshold_value is not None else "",
                    "是" if alert.is_resolved else "否",
                    alert.resolved_by or "",
                    alert.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ])
            writer.writerow([])
        
        if plan.revenue_reports:
            report = plan.revenue_reports[0]
            writer.writerow(["【收益报表】"])
            writer.writerow(["充电总量(kWh)", report.total_charge_kwh])
            writer.writerow(["放电总量(kWh)", report.total_discharge_kwh])
            writer.writerow(["充电成本(元)", round(report.charge_cost, 2)])
            writer.writerow(["放电收益(元)", round(report.discharge_revenue, 2)])
            writer.writerow(["净利润(元)", round(report.net_profit, 2)])
            writer.writerow(["效率(%)", round(report.efficiency_rate, 2)])
            writer.writerow(["偏差率(%)", round(report.deviation_rate, 2)])
            writer.writerow([])
        
        if plan.history:
            writer.writerow(["【状态变更历史】"])
            writer.writerow([
                "记录ID", "操作", "原状态", "新状态", "操作人", "原因", "时间"
            ])
            
            for history in plan.history:
                writer.writerow([
                    history.id,
                    history.action,
                    history.old_status or "",
                    history.new_status,
                    history.operator or "",
                    history.reason or "",
                    history.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ])
            writer.writerow([])
        
        if plan.versions:
            writer.writerow(["【版本快照】"])
            writer.writerow(["版本号", "状态", "变更原因", "变更人", "创建时间"])
            
            for version in plan.versions:
                writer.writerow([
                    version.version_number,
                    version.status,
                    version.change_reason or "",
                    version.changed_by or "",
                    version.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ])
        
        writer.writerow([])
        writer.writerow(["报表说明:"])
        writer.writerow(["- 此报表用于业务复核，包含计划完整生命周期信息"])
        writer.writerow(["- 状态变更历史和版本快照可追溯所有人工干预"])
        writer.writerow(["- 收益报表已考虑偏差告警的影响"])
        writer.writerow(["- 如需验证数据，请核对执行回执与计划时段的对应关系"])
        
        return output.getvalue().encode('utf-8-sig')

    @staticmethod
    def export_plan_to_excel(
        db: Session,
        plan_id: int
    ) -> bytes:
        plan = db.query(models.ChargePlan).filter(
            models.ChargePlan.id == plan_id
        ).first()
        if not plan:
            raise ValueError("计划不存在")
        
        wb = Workbook()
        
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        ws_info = wb.active
        ws_info.title = "计划信息"
        
        ws_info.append(["储能充放电计划 - 业务复核报表"])
        ws_info.merge_cells('A1:F1')
        ws_info['A1'].font = Font(bold=True, size=14)
        ws_info['A1'].alignment = center_align
        
        ws_info.append(["生成时间:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws_info.append([])
        ws_info.append(["【计划基本信息】"])
        
        info_data = [
            ["计划ID", plan.id],
            ["电站ID", plan.station_id],
            ["计划日期", str(plan.plan_date)],
            ["计划名称", plan.plan_name],
            ["描述", plan.description or ""],
            ["版本", plan.version],
            ["状态", plan.status],
            ["创建人", plan.created_by or "系统"],
            ["创建时间", plan.created_at.strftime("%Y-%m-%d %H:%M:%S")]
        ]
        
        for row_data in info_data:
            ws_info.append(row_data)
        
        ws_info.append([])
        ws_info.append(["【SOC 约束信息】"])
        
        soc = plan.soc_constraint
        soc_data = [
            ["最小SOC(%)", soc.min_soc],
            ["最大SOC(%)", soc.max_soc],
            ["初始SOC(%)", soc.initial_soc],
            ["目标SOC(%)", soc.target_soc or ""]
        ]
        
        for row_data in soc_data:
            ws_info.append(row_data)
        
        ws_info.column_dimensions['A'].width = 15
        ws_info.column_dimensions['B'].width = 40
        
        ws_segments = wb.create_sheet("计划时段")
        headers = ["时段序号", "开始时间", "结束时间", "操作类型",
                  "功率(kW)", "电量(kWh)", "预期SOC(%)", "关联电价窗口"]
        ws_segments.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws_segments.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        
        for i, segment in enumerate(sorted(plan.segments, key=lambda s: s.start_time), 1):
            op_type_name = {
                "charge": "充电",
                "discharge": "放电",
                "idle": "待机"
            }.get(segment.operation_type, segment.operation_type)
            
            window_info = ""
            if segment.price_window:
                window_info = f"{segment.price_window.window_type} ({segment.price_window.price_per_kwh}元/kWh)"
            
            row_data = [
                i,
                segment.start_time.strftime("%Y-%m-%d %H:%M:%S"),
                segment.end_time.strftime("%Y-%m-%d %H:%M:%S"),
                op_type_name,
                segment.power_kw,
                segment.energy_kwh,
                segment.expected_soc,
                window_info
            ]
            ws_segments.append(row_data)
            
            for col in range(1, len(headers) + 1):
                cell = ws_segments.cell(row=ws_segments.max_row, column=col)
                cell.border = thin_border
        
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']:
            ws_segments.column_dimensions[col_letter].width = 20
        
        if plan.execution_receipts:
            ws_receipts = wb.create_sheet("执行回执")
            headers = ["回执ID", "时段ID", "实际开始时间", "实际结束时间",
                      "实际功率(kW)", "实际电量(kWh)", "实际SOC(%)",
                      "设备状态", "备注"]
            ws_receipts.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws_receipts.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            
            for receipt in plan.execution_receipts:
                status_name = {
                    "available": "可用",
                    "maintenance": "维护",
                    "fault": "故障",
                    "limited": "受限"
                }.get(receipt.equipment_status, receipt.equipment_status)
                
                row_data = [
                    receipt.id,
                    receipt.segment_id,
                    receipt.actual_start_time.strftime("%Y-%m-%d %H:%M:%S"),
                    receipt.actual_end_time.strftime("%Y-%m-%d %H:%M:%S"),
                    receipt.actual_power_kw,
                    receipt.actual_energy_kwh,
                    receipt.actual_soc,
                    status_name,
                    receipt.remarks or ""
                ]
                ws_receipts.append(row_data)
                
                for col in range(1, len(headers) + 1):
                    cell = ws_receipts.cell(row=ws_receipts.max_row, column=col)
                    cell.border = thin_border
            
            for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
                ws_receipts.column_dimensions[col_letter].width = 18
        
        if plan.alerts:
            ws_alerts = wb.create_sheet("偏差告警")
            headers = ["告警ID", "告警类型", "告警级别", "消息",
                      "偏差值", "阈值", "是否已处理", "处理人", "创建时间"]
            ws_alerts.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws_alerts.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            
            for alert in plan.alerts:
                level_name = {
                    "info": "信息",
                    "warning": "警告",
                    "critical": "严重"
                }.get(alert.alert_level, alert.alert_level)
                
                row_data = [
                    alert.id,
                    alert.alert_type,
                    level_name,
                    alert.message,
                    alert.deviation_value if alert.deviation_value is not None else "",
                    alert.threshold_value if alert.threshold_value is not None else "",
                    "是" if alert.is_resolved else "否",
                    alert.resolved_by or "",
                    alert.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ]
                ws_alerts.append(row_data)
                
                for col in range(1, len(headers) + 1):
                    cell = ws_alerts.cell(row=ws_alerts.max_row, column=col)
                    cell.border = thin_border
                
                if alert.alert_level == "critical" and not alert.is_resolved:
                    for col in range(1, len(headers) + 1):
                        cell = ws_alerts.cell(row=ws_alerts.max_row, column=col)
                        cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            
            for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
                ws_alerts.column_dimensions[col_letter].width = 15
        
        if plan.revenue_reports:
            ws_revenue = wb.create_sheet("收益报表")
            report = plan.revenue_reports[0]
            
            ws_revenue.append(["【收益报表】"])
            ws_revenue.merge_cells('A1:B1')
            ws_revenue['A1'].font = Font(bold=True, size=12)
            
            revenue_data = [
                ["充电总量(kWh)", report.total_charge_kwh],
                ["放电总量(kWh)", report.total_discharge_kwh],
                ["充电成本(元)", round(report.charge_cost, 2)],
                ["放电收益(元)", round(report.discharge_revenue, 2)],
                ["净利润(元)", round(report.net_profit, 2)],
                ["效率(%)", round(report.efficiency_rate, 2)],
                ["偏差率(%)", round(report.deviation_rate, 2)]
            ]
            
            for row_data in revenue_data:
                ws_revenue.append(row_data)
            
            ws_revenue.column_dimensions['A'].width = 15
            ws_revenue.column_dimensions['B'].width = 20
            
            if report.net_profit < 0:
                profit_cell = ws_revenue.cell(row=6, column=2)
                profit_cell.font = Font(bold=True, color="FF0000")
        
        if plan.history:
            ws_history = wb.create_sheet("状态历史")
            headers = ["记录ID", "操作", "原状态", "新状态", "操作人", "原因", "时间"]
            ws_history.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws_history.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            
            for history in plan.history:
                row_data = [
                    history.id,
                    history.action,
                    history.old_status or "",
                    history.new_status,
                    history.operator or "",
                    history.reason or "",
                    history.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ]
                ws_history.append(row_data)
                
                for col in range(1, len(headers) + 1):
                    cell = ws_history.cell(row=ws_history.max_row, column=col)
                    cell.border = thin_border
                
                if history.action == "人工修正":
                    for col in range(1, len(headers) + 1):
                        cell = ws_history.cell(row=ws_history.max_row, column=col)
                        cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            
            for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
                ws_history.column_dimensions[col_letter].width = 18
        
        if plan.versions:
            ws_versions = wb.create_sheet("版本快照")
            headers = ["版本号", "状态", "变更原因", "变更人", "创建时间"]
            ws_versions.append(headers)
            
            for col in range(1, len(headers) + 1):
                cell = ws_versions.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            
            for version in plan.versions:
                row_data = [
                    version.version_number,
                    version.status,
                    version.change_reason or "",
                    version.changed_by or "",
                    version.created_at.strftime("%Y-%m-%d %H:%M:%S")
                ]
                ws_versions.append(row_data)
                
                for col in range(1, len(headers) + 1):
                    cell = ws_versions.cell(row=ws_versions.max_row, column=col)
                    cell.border = thin_border
            
            for col_letter in ['A', 'B', 'C', 'D', 'E']:
                ws_versions.column_dimensions[col_letter].width = 20
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
