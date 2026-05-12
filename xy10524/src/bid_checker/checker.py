import os
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional, Tuple
from .models import (
    CheckStatus, CheckResult, MaterialItem, MaterialType,
    MaterialRequirement, VersionInfo, ReportSummary
)


class BidChecker:
    def __init__(self, project_dir: str, config: Dict[str, Any]):
        self.project_dir = Path(project_dir)
        self.config = config
        self.requirements = [
            MaterialRequirement(**r) for r in config.get("requirements", [])
        ]
        self.rules = config.get("rules", {})

    def discover_materials(self) -> List[MaterialItem]:
        materials = []
        processed_paths = set()

        for req in self.requirements:
            if req.pattern:
                for file_path in self.project_dir.rglob(req.pattern):
                    if file_path.is_file() and str(file_path) not in processed_paths:
                        processed_paths.add(str(file_path))
                        material = self._create_material_item(file_path, req.type)
                        if material:
                            materials.append(material)

        for file_path in self.project_dir.rglob("*"):
            if file_path.is_file() and str(file_path) not in processed_paths:
                if not str(file_path).startswith(str(self.project_dir / ".bid-checker")):
                    material = self._create_material_item(file_path, MaterialType.OTHER)
                    if material:
                        materials.append(material)

        return materials

    def _create_material_item(self, file_path: Path, mat_type: MaterialType) -> Optional[MaterialItem]:
        name = file_path.name
        version = self._extract_version(name)
        is_current = self._is_current_version(name)
        has_seal = self._check_has_seal(file_path)
        amount = None
        amount_source = None

        if mat_type in [MaterialType.QUOTATION]:
            amount, amount_source = self._extract_amount(file_path)

        return MaterialItem(
            id=f"mat_{len(name)}_{abs(hash(str(file_path))) % 10000}",
            name=name,
            type=mat_type,
            path=str(file_path.relative_to(self.project_dir)),
            version=version,
            is_current=is_current,
            has_seal=has_seal,
            amount=amount,
            amount_source=amount_source
        )

    def _extract_version(self, filename: str) -> Optional[str]:
        patterns = [
            r'[vV](\d+(?:\.\d+)*)',
            r'版本\s*(\d+(?:\.\d+)*)',
            r'(\d+(?:\.\d+)*)\s*版',
            r'_(\d{8})_'
        ]
        for pattern in patterns:
            match = re.search(pattern, filename)
            if match:
                return match.group(1)
        return None

    def _is_current_version(self, filename: str) -> bool:
        return '旧' not in filename and 'old' not in filename.lower()

    def _check_has_seal(self, file_path: Path) -> bool:
        name_lower = file_path.name.lower()
        return '盖章' in file_path.name or 'seal' in name_lower or 'signed' in name_lower

    def _extract_amount(self, file_path: Path) -> Tuple[Optional[float], Optional[str]]:
        if file_path.suffix.lower() in ['.xlsx', '.xls']:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(file_path, read_only=True)
                for sheet in wb:
                    for row in sheet.iter_rows(values_only=True):
                        for cell in row:
                            if isinstance(cell, str):
                                amount = self._parse_amount_string(cell)
                                if amount:
                                    return amount, f"{file_path.name}:{sheet.title}"
                            elif isinstance(cell, (int, float)):
                                if cell > 1000:
                                    return float(cell), f"{file_path.name}:{sheet.title}"
            except Exception:
                pass
        return None, None

    def _parse_amount_string(self, text: str) -> Optional[float]:
        patterns = [
            r'[￥¥]?\s*([\d,]+\.?\d*)\s*元',
            r'合计\s*[:：]?\s*[￥¥]?\s*([\d,]+\.?\d*)',
            r'总计\s*[:：]?\s*[￥¥]?\s*([\d,]+\.?\d*)',
            r'报价\s*[:：]?\s*[￥¥]?\s*([\d,]+\.?\d*)'
        ]
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                try:
                    return float(match.group(1).replace(',', ''))
                except ValueError:
                    pass
        return None

    def run_checks(self, materials: List[MaterialItem], versions: List[VersionInfo]) -> List[CheckResult]:
        results = []
        results.extend(self._check_required_materials(materials))
        results.extend(self._check_old_versions(materials))
        results.extend(self._check_seal_pages(materials))
        results.extend(self._check_amount_consistency(materials))
        results.extend(self._check_version_documents(materials, versions))
        return results

    def _check_required_materials(self, materials: List[MaterialItem]) -> List[CheckResult]:
        results = []
        found_types = set(m.type for m in materials if m.is_current)

        for req in self.requirements:
            if req.required:
                if req.type not in found_types:
                    results.append(CheckResult(
                        rule_id="REQUIRED_MISSING",
                        status=CheckStatus.BLOCKED,
                        message=f"缺少必需材料: {req.name}",
                        details={"requirement": req.name, "type": req.type},
                    ))
                else:
                    results.append(CheckResult(
                        rule_id="REQUIRED_PRESENT",
                        status=CheckStatus.PASSED,
                        message=f"必需材料齐全: {req.name}",
                        details={"requirement": req.name},
                    ))
        return results

    def _check_old_versions(self, materials: List[MaterialItem]) -> List[CheckResult]:
        results = []
        for material in materials:
            if not material.is_current:
                results.append(CheckResult(
                    rule_id="OLD_VERSION_DETECTED",
                    status=CheckStatus.WARNING,
                    message=f"检测到旧版本文件: {material.name}",
                    details={"file": material.name, "version": material.version},
                    material_id=material.id
                ))
        if not any(r.rule_id == "OLD_VERSION_DETECTED" for r in results):
            results.append(CheckResult(
                rule_id="NO_OLD_VERSIONS",
                status=CheckStatus.PASSED,
                message="未检测到旧版本文件"
            ))
        return results

    def _check_seal_pages(self, materials: List[MaterialItem]) -> List[CheckResult]:
        results = []
        seal_materials = [m for m in materials if m.type == MaterialType.SEAL_PAGE]

        if self.rules.get("must_have_seal", True):
            if not seal_materials:
                results.append(CheckResult(
                    rule_id="SEAL_MISSING",
                    status=CheckStatus.BLOCKED,
                    message="缺少盖章页文件",
                    details={"required": True}
                ))
            else:
                for material in seal_materials:
                    if not material.has_seal:
                        results.append(CheckResult(
                            rule_id="SEAL_NOT_VERIFIED",
                            status=CheckStatus.WARNING,
                            message=f"盖章页未确认: {material.name}",
                            details={"file": material.name},
                            material_id=material.id
                        ))
                    else:
                        results.append(CheckResult(
                            rule_id="SEAL_VERIFIED",
                            status=CheckStatus.PASSED,
                            message=f"盖章页已确认: {material.name}",
                            material_id=material.id
                        ))
        return results

    def _check_amount_consistency(self, materials: List[MaterialItem]) -> List[CheckResult]:
        results = []
        if not self.rules.get("check_amount_consistency", True):
            return results

        quotations = [m for m in materials if m.type == MaterialType.QUOTATION and m.amount is not None]

        if len(quotations) < 2:
            results.append(CheckResult(
                rule_id="INSUFFICIENT_QUOTATIONS",
                status=CheckStatus.WARNING,
                message=f"报价文件数量不足，无法验证一致性（找到 {len(quotations)} 个）"
            ))
            return results

        amounts = [m.amount for m in quotations if m.amount is not None]
        if amounts:
            first_amount = amounts[0]
            for i, amount in enumerate(amounts[1:], 1):
                if abs(amount - first_amount) > 0.01:
                    results.append(CheckResult(
                        rule_id="AMOUNT_MISMATCH",
                        status=CheckStatus.FAILED,
                        message=f"报价金额不一致: {quotations[0].name} ({first_amount}) vs {quotations[i].name} ({amount})",
                        details={
                            "file1": quotations[0].name,
                            "amount1": first_amount,
                            "file2": quotations[i].name,
                            "amount2": amount
                        },
                        material_id=quotations[i].id
                    ))

        if not any(r.rule_id == "AMOUNT_MISMATCH" for r in results):
            results.append(CheckResult(
                rule_id="AMOUNT_CONSISTENT",
                status=CheckStatus.PASSED,
                message=f"报价金额一致（共 {len(quotations)} 个文件）"
            ))

        return results

    def _check_version_documents(self, materials: List[MaterialItem], versions: List[VersionInfo]) -> List[CheckResult]:
        results = []
        version_docs = [m for m in materials if m.type == MaterialType.VERSION_DOC]

        if not version_docs:
            results.append(CheckResult(
                rule_id="VERSION_DOC_MISSING",
                status=CheckStatus.WARNING,
                message="缺少版本说明文档"
            ))
        else:
            results.append(CheckResult(
                rule_id="VERSION_DOC_PRESENT",
                status=CheckStatus.PASSED,
                message=f"版本说明文档已存在: {', '.join(m.name for m in version_docs)}"
            ))

        if not versions:
            results.append(CheckResult(
                rule_id="NO_VERSION_INFO",
                status=CheckStatus.WARNING,
                message="未导入版本变更历史"
            ))
        else:
            results.append(CheckResult(
                rule_id="VERSION_INFO_PRESENT",
                status=CheckStatus.PASSED,
                message=f"已导入 {len(versions)} 条版本记录"
            ))

        return results

    def generate_report(self, check_results: List[CheckResult], materials: List[MaterialItem]) -> ReportSummary:
        status_counts = {
            CheckStatus.PASSED: 0,
            CheckStatus.WARNING: 0,
            CheckStatus.FAILED: 0,
            CheckStatus.BLOCKED: 0,
            CheckStatus.PENDING: 0
        }

        missing_materials = []
        old_versions = []
        amount_mismatches = []
        missing_seals = []
        corrections_needed = []

        for result in check_results:
            status_counts[result.status] += 1

            if result.rule_id == "REQUIRED_MISSING":
                missing_materials.append(result.details.get("requirement", ""))
                corrections_needed.append({
                    "type": "missing_material",
                    "item": result.details.get("requirement", ""),
                    "severity": "blocked",
                    "action": "补充缺失材料"
                })

            elif result.rule_id == "OLD_VERSION_DETECTED":
                old_versions.append(result.details.get("file", ""))
                corrections_needed.append({
                    "type": "old_version",
                    "item": result.details.get("file", ""),
                    "severity": "warning",
                    "action": "确认是否需要替换为新版本"
                })

            elif result.rule_id == "AMOUNT_MISMATCH":
                amount_mismatches.append(f"{result.details.get('file1')} vs {result.details.get('file2')}")
                corrections_needed.append({
                    "type": "amount_mismatch",
                    "item": f"{result.details.get('file1')} vs {result.details.get('file2')}",
                    "severity": "failed",
                    "action": "统一报价金额"
                })

            elif result.rule_id in ["SEAL_MISSING", "SEAL_NOT_VERIFIED"]:
                missing_seals.append(result.details.get("file", "盖章页"))
                corrections_needed.append({
                    "type": "missing_seal",
                    "item": result.details.get("file", "盖章页"),
                    "severity": "blocked" if result.rule_id == "SEAL_MISSING" else "warning",
                    "action": "补充或确认盖章"
                })

        can_submit = (
            status_counts[CheckStatus.BLOCKED] == 0 and
            status_counts[CheckStatus.FAILED] == 0
        )

        return ReportSummary(
            total_checks=len(check_results),
            passed=status_counts[CheckStatus.PASSED],
            warnings=status_counts[CheckStatus.WARNING],
            failed=status_counts[CheckStatus.FAILED],
            blocked=status_counts[CheckStatus.BLOCKED],
            can_submit=can_submit,
            missing_materials=missing_materials,
            old_versions=old_versions,
            amount_mismatches=amount_mismatches,
            missing_seals=missing_seals,
            corrections_needed=corrections_needed
        )
