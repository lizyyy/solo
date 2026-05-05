import pandas as pd
import os
import re
from datetime import datetime
from typing import Dict, Any, List, Optional
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from .exceptions import ReportGenerationError


class ReportGenerator:
    """报告生成器 - 生成 Excel、PDF、Markdown 三种格式的报告"""
    
    def __init__(self, template_content: str, metrics_config: Dict[str, Any]):
        self.template_content = template_content
        self.metrics_config = metrics_config
        self.generated_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    def _prepare_template_context(self, 
                                    metrics_result: Dict[str, Any],
                                    orders_df: pd.DataFrame) -> Dict[str, Any]:
        """准备模板渲染上下文"""
        context = {
            'report_title': self.metrics_config.get('report_title', '销售报告'),
            'report_period': self.metrics_config.get('report_period', {'start': '', 'end': ''}),
            'gmv': metrics_result.get('gmv', 0),
            'total_orders': metrics_result.get('total_orders', 0),
            'average_order_value': metrics_result.get('average_order_value', 0),
            'refund_rate': metrics_result.get('refund_rate', 0),
            'channel_conversion': [],
            'top_channels': '',
            'anomaly_count': 0,
            'anomaly_rate': 0.0,
            'anomalous_orders': [],
            'conclusion': '',
            'suggestion_1': '优化高退款渠道的用户体验',
            'suggestion_2': '增加高转化渠道的营销投入',
            'suggestion_3': '建立异常订单实时监控机制',
            'generated_at': self.generated_at
        }
        
        channel_conversion = metrics_result.get('channel_conversion', {})
        channel_display_names = {}
        for ch in self.metrics_config.get('channels', []):
            channel_display_names[ch['name']] = ch.get('display_name', ch['name'])
        
        channel_stats = []
        total_orders = metrics_result.get('total_orders', len(orders_df))
        for channel, rate in channel_conversion.items():
            order_count = len(orders_df[orders_df['channel'] == channel])
            channel_stats.append({
                'name': channel,
                'display_name': channel_display_names.get(channel, channel),
                'order_count': order_count,
                'rate': rate
            })
        
        context['channel_conversion'] = channel_stats
        
        top_channels = metrics_result.get('top_channels', {})
        if top_channels:
            top_channel_names = []
            for ch, count in top_channels.items():
                display_name = channel_display_names.get(ch, ch)
                top_channel_names.append(f"{display_name} ({count} 单)")
            context['top_channels'] = '、'.join(top_channel_names)
        
        anomalous_orders = metrics_result.get('anomalous_orders', pd.DataFrame())
        if not anomalous_orders.empty:
            context['anomaly_count'] = len(anomalous_orders)
            context['anomaly_rate'] = round((len(anomalous_orders) / total_orders) * 100, 2) if total_orders > 0 else 0
            
            anomalies_list = []
            for idx, row in anomalous_orders.iterrows():
                reasons = row.get('reasons', [])
                anomalies_list.append({
                    'order_id': row.get('order_id', ''),
                    'amount': row.get('amount', 0),
                    'status': row.get('status', ''),
                    'channel': channel_display_names.get(row.get('channel', ''), row.get('channel', '')),
                    'reasons': '；'.join(reasons) if reasons else ''
                })
            context['anomalous_orders'] = anomalies_list
        
        gmv = metrics_result.get('gmv', 0)
        refund_rate = metrics_result.get('refund_rate', 0)
        top_channel = '未知'
        if channel_stats:
            top_channel = channel_stats[0]['display_name']
        
        performance = "良好" if gmv > 5000 else "一般"
        refund_status = "处于健康水平" if refund_rate < 5 else "需要关注"
        
        context['conclusion'] = f"""本月销售表现{performance}，GMV 达到 {gmv} 元。
退款率为 {refund_rate}%，{refund_status}。
主要销售渠道为 {top_channel}。"""
        
        return context
    
    def _render_template(self, context: Dict[str, Any]) -> str:
        """渲染模板"""
        content = self.template_content
        
        def replace_simple_variables(match):
            var_name = match.group(1)
            value = context.get(var_name, '')
            return str(value)
        
        content = re.sub(r'\{\{(\w+)\}\}', replace_simple_variables, content)
        
        def replace_nested_variables(match):
            var_path = match.group(1)
            parts = var_path.split('.')
            value = context
            for part in parts:
                if isinstance(value, dict):
                    value = value.get(part, '')
                else:
                    value = ''
                    break
            return str(value)
        
        content = re.sub(r'\{\{(\w+\.[\w.]+)\}\}', replace_nested_variables, content)
        
        def replace_list_sections(match):
            section_name = match.group(1)
            section_template = match.group(2)
            items = context.get(section_name, [])
            
            if not items:
                return ''
            
            result_lines = []
            for item in items:
                if isinstance(item, dict):
                    item_content = section_template
                    for key, value in item.items():
                        item_content = item_content.replace(f'{{{{{key}}}}}', str(value))
                    result_lines.append(item_content)
            
            return '\n'.join(result_lines)
        
        content = re.sub(r'\{\{#(\w+)\}\}(.*?)\{\{\/\1\}\}', replace_list_sections, content, flags=re.DOTALL)
        
        return content
    
    def generate_markdown(self, 
                          metrics_result: Dict[str, Any],
                          orders_df: pd.DataFrame,
                          output_path: str) -> str:
        """
        生成 Markdown 格式报告
        
        Args:
            metrics_result: 指标计算结果
            orders_df: 订单数据 DataFrame
            output_path: 输出文件路径
            
        Returns:
            生成的 Markdown 内容
        """
        try:
            context = self._prepare_template_context(metrics_result, orders_df)
            markdown_content = self._render_template(context)
            
            os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(markdown_content)
            
            return markdown_content
            
        except Exception as e:
            raise ReportGenerationError(
                message=str(e),
                output_format='markdown',
                suggestions=['检查模板语法', '确保数据格式正确']
            ) from e
    
    def generate_excel(self,
                       metrics_result: Dict[str, Any],
                       orders_df: pd.DataFrame,
                       output_path: str) -> str:
        """
        生成 Excel 格式报告
        
        Args:
            metrics_result: 指标计算结果
            orders_df: 订单数据 DataFrame
            output_path: 输出文件路径
            
        Returns:
            输出文件路径
        """
        try:
            context = self._prepare_template_context(metrics_result, orders_df)
            
            wb = Workbook()
            
            ws_summary = wb.active
            ws_summary.title = "报告摘要"
            
            header_font = Font(bold=True, size=14, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            center_align = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            ws_summary['A1'] = context['report_title']
            ws_summary['A1'].font = Font(bold=True, size=16)
            ws_summary.merge_cells('A1:B1')
            
            ws_summary['A2'] = f"报告周期: {context['report_period'].get('start', '')} - {context['report_period'].get('end', '')}"
            ws_summary.merge_cells('A2:B2')
            
            ws_summary['A3'] = f"生成时间: {self.generated_at}"
            ws_summary.merge_cells('A3:B3')
            
            ws_summary['A5'] = "核心指标"
            ws_summary['A5'].font = Font(bold=True, size=12)
            ws_summary.merge_cells('A5:B5')
            
            metrics_data = [
                ("指标", "数值"),
                ("GMV", f"{context['gmv']} 元"),
                ("总订单数", f"{context['total_orders']} 笔"),
                ("平均客单价", f"{context['average_order_value']} 元"),
                ("退款率", f"{context['refund_rate']}%"),
                ("异常订单数", f"{context['anomaly_count']} 笔"),
                ("异常订单占比", f"{context['anomaly_rate']}%")
            ]
            
            for row_idx, (metric, value) in enumerate(metrics_data, start=6):
                ws_summary[f'A{row_idx}'] = metric
                ws_summary[f'B{row_idx}'] = value
                
                if row_idx == 6:
                    ws_summary[f'A{row_idx}'].font = header_font
                    ws_summary[f'B{row_idx}'].font = header_font
                    ws_summary[f'A{row_idx}'].fill = header_fill
                    ws_summary[f'B{row_idx}'].fill = header_fill
                
                ws_summary[f'A{row_idx}'].border = thin_border
                ws_summary[f'B{row_idx}'].border = thin_border
            
            ws_summary.column_dimensions['A'].width = 20
            ws_summary.column_dimensions['B'].width = 30
            
            ws_channels = wb.create_sheet("渠道分析")
            
            ws_channels['A1'] = "渠道分布分析"
            ws_channels['A1'].font = Font(bold=True, size=14)
            ws_channels.merge_cells('A1:D1')
            
            channels_data = [
                ("渠道", "订单数", "占比", "排名")
            ]
            
            channel_stats = context['channel_conversion']
            for idx, ch in enumerate(channel_stats, start=1):
                channels_data.append((
                    ch['display_name'],
                    ch['order_count'],
                    f"{ch['rate']}%",
                    idx
                ))
            
            for row_idx, row_data in enumerate(channels_data, start=2):
                for col_idx, value in enumerate(row_data, start=1):
                    cell = ws_channels.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    
                    if row_idx == 2:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = center_align
            
            ws_channels.column_dimensions['A'].width = 15
            ws_channels.column_dimensions['B'].width = 12
            ws_channels.column_dimensions['C'].width = 10
            ws_channels.column_dimensions['D'].width = 8
            
            ws_anomalies = wb.create_sheet("异常订单")
            
            ws_anomalies['A1'] = "异常订单详情"
            ws_anomalies['A1'].font = Font(bold=True, size=14)
            ws_anomalies.merge_cells('A1:E1')
            
            anomalies_headers = ["订单ID", "金额", "状态", "渠道", "异常原因"]
            for col_idx, header in enumerate(anomalies_headers, start=1):
                cell = ws_anomalies.cell(row=2, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
            
            anomalies_list = context['anomalous_orders']
            for row_idx, anomaly in enumerate(anomalies_list, start=3):
                ws_anomalies.cell(row=row_idx, column=1, value=anomaly.get('order_id', '')).border = thin_border
                ws_anomalies.cell(row=row_idx, column=2, value=anomaly.get('amount', 0)).border = thin_border
                ws_anomalies.cell(row=row_idx, column=3, value=anomaly.get('status', '')).border = thin_border
                ws_anomalies.cell(row=row_idx, column=4, value=anomaly.get('channel', '')).border = thin_border
                ws_anomalies.cell(row=row_idx, column=5, value=anomaly.get('reasons', '')).border = thin_border
            
            ws_anomalies.column_dimensions['A'].width = 15
            ws_anomalies.column_dimensions['B'].width = 12
            ws_anomalies.column_dimensions['C'].width = 10
            ws_anomalies.column_dimensions['D'].width = 12
            ws_anomalies.column_dimensions['E'].width = 50
            
            ws_data = wb.create_sheet("原始数据")
            
            for col_idx, col_name in enumerate(orders_df.columns, start=1):
                cell = ws_data.cell(row=1, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
            
            for row_idx, row_data in enumerate(dataframe_to_rows(orders_df, index=False, header=False), start=2):
                for col_idx, value in enumerate(row_data, start=1):
                    cell = ws_data.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
            
            for col_idx in range(1, len(orders_df.columns) + 1):
                ws_data.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)].width = 18
            
            ws_conclusion = wb.create_sheet("结论与建议")
            
            ws_conclusion['A1'] = "结论"
            ws_conclusion['A1'].font = Font(bold=True, size=14)
            
            ws_conclusion['A2'] = context['conclusion']
            ws_conclusion.merge_cells('A2:C5')
            ws_conclusion['A2'].alignment = Alignment(wrap_text=True, vertical='top')
            
            ws_conclusion['A7'] = "建议"
            ws_conclusion['A7'].font = Font(bold=True, size=14)
            
            suggestions = [
                context['suggestion_1'],
                context['suggestion_2'],
                context['suggestion_3']
            ]
            
            for idx, suggestion in enumerate(suggestions, start=1):
                ws_conclusion[f'A{7 + idx}'] = f"{idx}. {suggestion}"
            
            ws_conclusion.column_dimensions['A'].width = 60
            
            os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
            wb.save(output_path)
            
            return output_path
            
        except Exception as e:
            raise ReportGenerationError(
                message=str(e),
                output_format='excel',
                suggestions=['检查数据格式', '确保 Excel 写入权限']
            ) from e
    
    def generate_pdf(self,
                     metrics_result: Dict[str, Any],
                     orders_df: pd.DataFrame,
                     output_path: str) -> str:
        """
        生成 PDF 格式报告
        
        Args:
            metrics_result: 指标计算结果
            orders_df: 订单数据 DataFrame
            output_path: 输出文件路径
            
        Returns:
            输出文件路径
        """
        try:
            context = self._prepare_template_context(metrics_result, orders_df)
            
            pdf = FPDF()
            pdf.add_page()
            
            pdf.set_font("Arial", 'B', 16)
            pdf.cell(0, 10, txt=context['report_title'], ln=True, align='C')
            
            pdf.set_font("Arial", size=10)
            pdf.cell(0, 8, txt=f"报告周期: {context['report_period'].get('start', '')} - {context['report_period'].get('end', '')}", ln=True, align='C')
            pdf.cell(0, 8, txt=f"生成时间: {self.generated_at}", ln=True, align='C')
            
            pdf.ln(10)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, txt="1. 核心指标概览", ln=True)
            
            pdf.set_font("Arial", size=10)
            pdf.cell(40, 8, txt="指标", border=1, align='C')
            pdf.cell(60, 8, txt="数值", border=1, align='C')
            pdf.ln()
            
            metrics_list = [
                ("GMV", f"{context['gmv']} 元"),
                ("总订单数", f"{context['total_orders']} 笔"),
                ("平均客单价", f"{context['average_order_value']} 元"),
                ("退款率", f"{context['refund_rate']}%"),
                ("异常订单数", f"{context['anomaly_count']} 笔"),
                ("异常订单占比", f"{context['anomaly_rate']}%")
            ]
            
            for metric, value in metrics_list:
                pdf.cell(40, 8, txt=metric, border=1)
                pdf.cell(60, 8, txt=str(value), border=1)
                pdf.ln()
            
            pdf.ln(10)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, txt="2. 渠道分布", ln=True)
            
            pdf.set_font("Arial", size=10)
            pdf.cell(40, 8, txt="渠道", border=1, align='C')
            pdf.cell(30, 8, txt="订单数", border=1, align='C')
            pdf.cell(30, 8, txt="占比", border=1, align='C')
            pdf.ln()
            
            for ch in context['channel_conversion']:
                pdf.cell(40, 8, txt=str(ch['display_name']), border=1)
                pdf.cell(30, 8, txt=str(ch['order_count']), border=1)
                pdf.cell(30, 8, txt=f"{ch['rate']}%", border=1)
                pdf.ln()
            
            pdf.ln(10)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, txt="3. 异常订单", ln=True)
            
            pdf.set_font("Arial", size=10)
            pdf.cell(0, 8, txt=f"异常订单总数: {context['anomaly_count']} 笔", ln=True)
            pdf.cell(0, 8, txt=f"异常订单占比: {context['anomaly_rate']}%", ln=True)
            
            if context['anomalous_orders']:
                pdf.ln(5)
                pdf.cell(30, 8, txt="订单ID", border=1, align='C')
                pdf.cell(25, 8, txt="金额", border=1, align='C')
                pdf.cell(80, 8, txt="异常原因", border=1, align='C')
                pdf.ln()
                
                for anomaly in context['anomalous_orders']:
                    pdf.cell(30, 8, txt=str(anomaly.get('order_id', '')), border=1)
                    pdf.cell(25, 8, txt=str(anomaly.get('amount', 0)), border=1)
                    pdf.cell(80, 8, txt=str(anomaly.get('reasons', ''))[:40], border=1)
                    pdf.ln()
            
            pdf.ln(10)
            
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(0, 10, txt="4. 结论与建议", ln=True)
            
            pdf.set_font("Arial", size=10)
            pdf.multi_cell(0, 8, txt="结论:")
            pdf.multi_cell(0, 8, txt=context['conclusion'])
            
            pdf.ln(5)
            pdf.multi_cell(0, 8, txt="建议:")
            suggestions = [
                context['suggestion_1'],
                context['suggestion_2'],
                context['suggestion_3']
            ]
            for idx, suggestion in enumerate(suggestions, start=1):
                pdf.multi_cell(0, 8, txt=f"{idx}. {suggestion}")
            
            os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
            pdf.output(output_path)
            
            return output_path
            
        except Exception as e:
            raise ReportGenerationError(
                message=str(e),
                output_format='pdf',
                suggestions=['检查 PDF 生成库', '确保写入权限']
            ) from e
