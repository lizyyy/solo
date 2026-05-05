import pandas as pd
import os
import re
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime
from openpyxl import load_workbook

from .exceptions import ReportVerificationError, ConsistencyError


class ReportVerifier:
    """报告验证器 - 验证多份报告内容的一致性"""
    
    KEY_SECTIONS = ['执行摘要', '核心指标', '渠道分析', '异常订单', '结论与建议', '结论', '建议']
    
    def __init__(self):
        self.extraction_errors: List[Dict[str, Any]] = []
    
    def extract_report_content(self, file_path: str, format_type: str) -> Dict[str, Any]:
        """
        从报告文件中提取关键内容
        
        Args:
            file_path: 报告文件路径
            format_type: 格式类型 ('excel', 'pdf', 'markdown')
            
        Returns:
            字典包含提取的关键内容
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"报告文件不存在: {file_path}")
        
        self.extraction_errors = []
        
        try:
            if format_type == 'markdown':
                return self._extract_from_markdown(file_path)
            elif format_type == 'excel':
                return self._extract_from_excel(file_path)
            elif format_type == 'pdf':
                return self._extract_from_pdf(file_path)
            else:
                raise ReportVerificationError(
                    message=f"不支持的格式类型: {format_type}",
                    report_type=format_type
                )
        except Exception as e:
            raise ReportVerificationError(
                message=f"提取报告内容失败: {str(e)}",
                report_type=format_type
            ) from e
    
    def _extract_from_markdown(self, file_path: str) -> Dict[str, Any]:
        """从 Markdown 文件提取内容"""
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        result = {
            'title': self._extract_title_from_markdown(content),
            'sections': self._extract_sections_from_markdown(content),
            'metrics': self._extract_metrics_from_markdown(content),
            'tables': self._extract_tables_from_markdown(content),
            'anomalies': self._extract_anomalies_from_markdown(content),
            'conclusion': self._extract_conclusion_from_markdown(content),
            'raw_content': content
        }
        
        return result
    
    def _extract_title_from_markdown(self, content: str) -> str:
        """从 Markdown 提取标题"""
        lines = content.strip().split('\n')
        for line in lines:
            line = line.strip()
            if line.startswith('# ') and len(line) > 2:
                return line[2:].strip()
        return '未找到标题'
    
    def _extract_sections_from_markdown(self, content: str) -> List[str]:
        """从 Markdown 提取章节标题"""
        sections = []
        pattern = r'^#{1,3}\s+(.+)$'
        for match in re.finditer(pattern, content, re.MULTILINE):
            sections.append(match.group(1).strip())
        return sections
    
    def _extract_metrics_from_markdown(self, content: str) -> Dict[str, Any]:
        """从 Markdown 提取关键指标"""
        metrics = {}
        
        gmv_match = re.search(r'GMV[：:]\s*([\d,\.]+', content)
        if gmv_match:
            gmv_str = gmv_match.group(1).replace(',', '')
            try:
                metrics['gmv'] = float(gmv_str)
            except ValueError:
                pass
        
        refund_match = re.search(r'退款率[：:]\s*([\d\.]+)\s*%', content)
        if refund_match:
            try:
                metrics['refund_rate'] = float(refund_match.group(1))
            except ValueError:
                pass
        
        orders_match = re.search(r'总订单数[：:]\s*(\d+)', content)
        if orders_match:
            try:
                metrics['total_orders'] = int(orders_match.group(1))
            except ValueError:
                pass
        
        avg_match = re.search(r'平均客单价[：:]\s*([\d,\.]+)', content)
        if avg_match:
            avg_str = avg_match.group(1).replace(',', '')
            try:
                metrics['average_order_value'] = float(avg_str)
            except ValueError:
                pass
        
        anomaly_match = re.search(r'异常订单总数[：:]\s*(\d+)', content)
        if anomaly_match:
            try:
                metrics['anomaly_count'] = int(anomaly_match.group(1))
            except ValueError:
                pass
        
        return metrics
    
    def _extract_tables_from_markdown(self, content: str) -> List[Dict[str, Any]]:
        """从 Markdown 提取表格信息"""
        tables = []
        
        table_pattern = r'\|.*\|[\s\S]*?\|.*\|'
        for match in re.finditer(table_pattern, content):
            table_text = match.group(0)
            lines = [line for line in table_text.split('\n') if '|' in line]
            
            if len(lines) >= 2:
                row_count = len(lines) - 2
                
                header_line = lines[0]
                headers = [h.strip() for h in header_line.split('|') if h.strip()]
                
                tables.append({
                    'row_count': row_count,
                    'headers': headers,
                    'full_text': table_text
                })
        
        return tables
    
    def _extract_anomalies_from_markdown(self, content: str) -> List[str]:
        """从 Markdown 提取异常订单相关信息"""
        anomalies = []
        
        anomaly_section_pattern = r'##.*异常订单.*\s*([\s\S]*?)(?=##|$)'
        match = re.search(anomaly_section_pattern, content)
        
        if match:
            section_content = match.group(1)
            row_pattern = r'\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]+)\|'
            for row_match in re.finditer(row_pattern, section_content):
                order_id = row_match.group(1).strip()
                if order_id and order_id != '订单ID' and not order_id.startswith('-'):
                    anomalies.append(order_id)
        
        return anomalies
    
    def _extract_conclusion_from_markdown(self, content: str) -> str:
        """从 Markdown 提取结论"""
        conclusion_pattern = r'##.*结论.*\s*([\s\S]*?)(?=##|$)'
        match = re.search(conclusion_pattern, content)
        
        if match:
            conclusion = match.group(1).strip()
            lines = [line.strip() for line in conclusion.split('\n') if line.strip()]
            return '\n'.join(lines)
        
        return ''
    
    def _extract_from_excel(self, file_path: str) -> Dict[str, Any]:
        """从 Excel 文件提取内容"""
        wb = load_workbook(file_path, data_only=True)
        
        result = {
            'title': '未找到标题',
            'sections': [],
            'metrics': {},
            'tables': [],
            'anomalies': [],
            'conclusion': '',
            'raw_content': ''
        }
        
        if '报告摘要' in wb.sheetnames:
            ws_summary = wb['报告摘要']
            if ws_summary['A1'].value:
                result['title'] = str(ws_summary['A1'].value)
            
            metrics = {}
            for row in range(6, 15):
                metric_name = ws_summary[f'A{row}'].value
                metric_value = ws_summary[f'B{row}'].value
                
                if metric_name and metric_value:
                    name_str = str(metric_name)
                    value_str = str(metric_value)
                    
                    if 'GMV' in name_str:
                        try:
                            gmv_match = re.search(r'([\d,\.]+)', value_str)
                            if gmv_match:
                                metrics['gmv'] = float(gmv_match.group(1).replace(',', ''))
                        except (ValueError, AttributeError):
                            pass
                    elif '退款率' in name_str:
                        try:
                            refund_match = re.search(r'([\d\.]+)', value_str)
                            if refund_match:
                                metrics['refund_rate'] = float(refund_match.group(1))
                        except (ValueError, AttributeError):
                            pass
                    elif '总订单数' in name_str:
                        try:
                            orders_match = re.search(r'(\d+)', value_str)
                            if orders_match:
                                metrics['total_orders'] = int(orders_match.group(1))
                        except (ValueError, AttributeError):
                            pass
                    elif '平均客单价' in name_str:
                        try:
                            avg_match = re.search(r'([\d,\.]+)', value_str)
                            if avg_match:
                                metrics['average_order_value'] = float(avg_match.group(1).replace(',', ''))
                        except (ValueError, AttributeError):
                            pass
                    elif '异常订单数' in name_str:
                        try:
                            anomaly_match = re.search(r'(\d+)', value_str)
                            if anomaly_match:
                                metrics['anomaly_count'] = int(anomaly_match.group(1))
                        except (ValueError, AttributeError):
                            pass
            
            result['metrics'] = metrics
        
        if '渠道分析' in wb.sheetnames:
            result['sections'].append('渠道分析')
            ws_channels = wb['渠道分析']
            row_count = 0
            for row in range(3, 20):
                if ws_channels[f'A{row}'].value:
                    row_count += 1
            result['tables'].append({
                'row_count': row_count,
                'headers': ['渠道', '订单数', '占比', '排名'],
                'full_text': ''
            })
        
        if '异常订单' in wb.sheetnames:
            result['sections'].append('异常订单')
            ws_anomalies = wb['异常订单']
            anomalies = []
            for row in range(3, 100):
                order_id = ws_anomalies[f'A{row}'].value
                if order_id:
                    anomalies.append(str(order_id))
            result['anomalies'] = anomalies
            
            result['tables'].append({
                'row_count': len(anomalies),
                'headers': ['订单ID', '金额', '状态', '渠道', '异常原因'],
                'full_text': ''
            })
        
        if '结论与建议' in wb.sheetnames:
            result['sections'].append('结论与建议')
            ws_conclusion = wb['结论与建议']
            conclusion_parts = []
            for row in range(2, 20):
                cell_value = ws_conclusion[f'A{row}'].value
                if cell_value:
                    conclusion_parts.append(str(cell_value))
            result['conclusion'] = '\n'.join(conclusion_parts)
        
        result['sections'] = list(set(result['sections']))
        
        return result
    
    def _extract_from_pdf(self, file_path: str) -> Dict[str, Any]:
        """从 PDF 文件提取内容（简化实现）"""
        try:
            import fitz
            doc = fitz.open(file_path)
            content = ''
            for page in doc:
                content += page.get_text()
            
            result = {
                'title': '未找到标题',
                'sections': [],
                'metrics': {},
                'tables': [],
                'anomalies': [],
                'conclusion': '',
                'raw_content': content
            }
            
            lines = content.split('\n')
            if lines:
                result['title'] = lines[0].strip()
            
            for section in self.KEY_SECTIONS:
                if section in content:
                    result['sections'].append(section)
            
            metrics = {}
            gmv_match = re.search(r'GMV\s*[:：]?\s*([\d,\.]+)', content)
            if gmv_match:
                try:
                    metrics['gmv'] = float(gmv_match.group(1).replace(',', ''))
                except ValueError:
                    pass
            
            refund_match = re.search(r'退款率\s*[:：]?\s*([\d\.]+)\s*%', content)
            if refund_match:
                try:
                    metrics['refund_rate'] = float(refund_match.group(1))
                except ValueError:
                    pass
            
            orders_match = re.search(r'总订单数\s*[:：]?\s*(\d+)', content)
            if orders_match:
                try:
                    metrics['total_orders'] = int(orders_match.group(1))
                except ValueError:
                    pass
            
            result['metrics'] = metrics
            
            anomaly_match = re.search(r'异常订单总数\s*[:：]?\s*(\d+)', content)
            if anomaly_match:
                try:
                    result['metrics']['anomaly_count'] = int(anomaly_match.group(1))
                except ValueError:
                    pass
            
            conclusion_pattern = r'结论[：:]\s*([\s\S]*?)(?=建议|$)'
            conclusion_match = re.search(conclusion_pattern, content)
            if conclusion_match:
                result['conclusion'] = conclusion_match.group(1).strip()
            
            return result
            
        except ImportError:
            return {
                'title': 'PDF 提取需要安装 PyMuPDF',
                'sections': [],
                'metrics': {},
                'tables': [],
                'anomalies': [],
                'conclusion': '',
                'raw_content': ''
            }
        except Exception as e:
            self.extraction_errors.append({
                'type': 'pdf_extraction_error',
                'message': str(e)
            })
            return {
                'title': 'PDF 提取失败',
                'sections': [],
                'metrics': {},
                'tables': [],
                'anomalies': [],
                'conclusion': '',
                'raw_content': ''
            }
    
    def verify_consistency(self, reports: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
        """
        验证多份报告的一致性
        
        Args:
            reports: 字典，key 为格式类型，value 为提取的报告内容
            
        Returns:
            验证结果字典
        """
        if len(reports) < 2:
            return {
                'consistent': True,
                'differences': [],
                'message': '只有一份报告，无需比较',
                'report_count': len(reports)
            }
        
        differences = []
        
        report_list = list(reports.items())
        base_format, base_content = report_list[0]
        
        for other_format, other_content in report_list[1:]:
            diffs = self._compare_two_reports(
                base_content, other_content,
                base_format, other_format
            )
            differences.extend(diffs)
        
        consistent = len(differences) == 0
        
        return {
            'consistent': consistent,
            'differences': differences,
            'report_count': len(reports),
            'formats_compared': list(reports.keys())
        }
    
    def _compare_two_reports(self, 
                              content1: Dict[str, Any],
                              content2: Dict[str, Any],
                              format1: str,
                              format2: str) -> List[Dict[str, Any]]:
        """比较两份报告的内容"""
        differences = []
        
        if content1.get('title') != content2.get('title'):
            differences.append({
                'type': 'title_mismatch',
                'location': '报告标题',
                'expected': content1.get('title'),
                'actual': content2.get('title'),
                'formats': [format1, format2],
                'suggestion': f'确保 {format1.upper()} 和 {format2.upper()} 报告的标题应该一致'
            })
        
        sections1 = set(content1.get('sections', []))
        sections2 = set(content2.get('sections', []))
        
        if sections1 != sections2:
            missing_in_2 = sections1 - sections2
            missing_in_1 = sections2 - sections1
            
            if missing_in_2:
                differences.append({
                    'type': 'section_missing',
                    'location': f'{format2.upper()} 报告',
                    'expected': list(missing_in_2),
                    'actual': None,
                    'formats': [format1, format2],
                    'suggestion': f'请在 {format2.upper()} 报告中添加以下章节: {", ".join(missing_in_2)}'
                })
            
            if missing_in_1:
                differences.append({
                    'type': 'section_extra',
                    'location': f'{format1.upper()} 报告',
                    'expected': None,
                    'actual': list(missing_in_1),
                    'formats': [format1, format2],
                    'suggestion': f'请检查 {format1.upper()} 报告中多余的章节: {", ".join(missing_in_1)}'
                })
        
        metrics1 = content1.get('metrics', {})
        metrics2 = content2.get('metrics', {})
        
        for key in ['gmv', 'refund_rate', 'total_orders', 'average_order_value', 'anomaly_count']:
            val1 = metrics1.get(key)
            val2 = metrics2.get(key)
            
            if val1 is not None and val2 is not None:
                if isinstance(val1, float) and isinstance(val2, float):
                    if abs(val1 - val2) > 0.01:
                        differences.append({
                            'type': 'metric_mismatch',
                            'location': f'指标: {key}',
                            'expected': val1,
                            'actual': val2,
                            'formats': [format1, format2],
                            'suggestion': f'检查 {format1.upper()} 和 {format2.upper()} 报告中 {key} 指标值应该一致'
                        })
                elif val1 != val2:
                    differences.append({
                        'type': 'metric_mismatch',
                        'location': f'指标: {key}',
                        'expected': val1,
                        'actual': val2,
                        'formats': [format1, format2],
                        'suggestion': f'检查 {format1.upper()} 和 {format2.upper()} 报告中 {key} 指标值应该一致'
                    })
        
        tables1 = content1.get('tables', [])
        tables2 = content2.get('tables', [])
        
        if len(tables1) != len(tables2):
            differences.append({
                'type': 'table_count_mismatch',
                'location': '表格数量',
                'expected': len(tables1),
                'actual': len(tables2),
                'formats': [format1, format2],
                'suggestion': f'{format1.upper()} 报告有 {len(tables1)} 个表格，{format2.upper()} 报告有 {len(tables2)} 个表格'
            })
        
        for i, (table1, table2) in enumerate(zip(tables1, tables2)):
            if table1.get('row_count') != table2.get('row_count'):
                differences.append({
                    'type': 'table_row_count_mismatch',
                    'location': f'表格 #{i+1}',
                    'expected': table1.get('row_count'),
                    'actual': table2.get('row_count'),
                    'formats': [format1, format2],
                    'suggestion': f'表格 #{i+1} 的行数不一致'
                })
        
        anomalies1 = set(content1.get('anomalies', []))
        anomalies2 = set(content2.get('anomalies', []))
        
        if anomalies1 != anomalies2:
            missing_in_2 = anomalies1 - anomalies2
            missing_in_1 = anomalies2 - anomalies1
            
            if missing_in_2:
                differences.append({
                    'type': 'anomaly_missing',
                    'location': f'{format2.upper()} 报告',
                    'expected': list(missing_in_2),
                    'formats': [format1, format2],
                    'suggestion': f'{format2.upper()} 报告缺少异常订单: {", ".join(missing_in_2)}'
                })
            
            if missing_in_1:
                differences.append({
                    'type': 'anomaly_extra',
                    'location': f'{format1.upper()} 报告',
                    'actual': list(missing_in_1),
                    'formats': [format1, format2],
                    'suggestion': f'{format1.upper()} 报告有多余的异常订单: {", ".join(missing_in_1)}'
                })
        
        conclusion1 = content1.get('conclusion', '').strip()
        conclusion2 = content2.get('conclusion', '').strip()
        
        if conclusion1 and conclusion2:
            def normalize_text(text: str) -> str:
                return ' '.join(text.lower().split())
            
            if normalize_text(conclusion1) != normalize_text(conclusion2):
                differences.append({
                    'type': 'conclusion_mismatch',
                    'location': '结论部分',
                    'expected': conclusion1[:200] + '...' if len(conclusion1) > 200 else conclusion1,
                    'actual': conclusion2[:200] + '...' if len(conclusion2) > 200 else conclusion2,
                    'formats': [format1, format2],
                    'suggestion': f'确保 {format1.upper()} 和 {format2.upper()} 报告的结论应该一致'
                })
        
        return differences
    
    def get_extraction_errors(self) -> List[Dict[str, Any]]:
        """获取提取错误列表"""
        return self.extraction_errors
