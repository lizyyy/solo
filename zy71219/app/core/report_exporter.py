from datetime import datetime
from typing import List, Dict, Any
from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from app.models import LetterOfCredit, Document, Discrepancy, LcClause, VersionRecord
from app.core.comparer import DISCREPANCY_CATEGORIES


SEVERITY_COLORS = {
    "CRITICAL": {"hex": "#FF4444", "excel": "FF4444"},
    "HIGH": {"hex": "#FF8800", "excel": "FF8800"},
    "MEDIUM": {"hex": "#FFCC00", "excel": "FFCC00"},
    "LOW": {"hex": "#88CC88", "excel": "88CC88"},
}

STATUS_COLORS = {
    "OPEN": {"hex": "#FF4444", "excel": "FF4444"},
    "IN_PROGRESS": {"hex": "#FF8800", "excel": "FF8800"},
    "RESOLVED": {"hex": "#44AA44", "excel": "44AA44"},
    "ACCEPTED": {"hex": "#8888FF", "excel": "8888FF"},
    "CLOSED": {"hex": "#AAAAAA", "excel": "AAAAAA"},
}


class ReportExporter:
    def __init__(self, db: Session):
        self.db = db

    def export_discrepancy_report_excel(self, lc_id: int) -> BytesIO:
        lc = self.db.query(LetterOfCredit).filter(LetterOfCredit.id == lc_id).first()
        if not lc:
            raise ValueError(f"Letter of Credit {lc_id} not found")

        discrepancies = self.db.query(Discrepancy).filter(Discrepancy.lc_id == lc_id).all()
        documents = self.db.query(Document).filter(Document.lc_id == lc_id, Document.is_active == True).all()

        output = BytesIO()
        wb = Workbook()

        self._create_summary_sheet(wb, lc, discrepancies, documents)
        self._create_discrepancies_sheet(wb, lc, discrepancies)
        self._create_documents_sheet(wb, documents)
        self._create_version_history_sheet(wb, lc_id)

        wb.save(output)
        output.seek(0)

        return output

    def export_discrepancy_report_pdf(self, lc_id: int) -> BytesIO:
        lc = self.db.query(LetterOfCredit).filter(LetterOfCredit.id == lc_id).first()
        if not lc:
            raise ValueError(f"Letter of Credit {lc_id} not found")

        discrepancies = self.db.query(Discrepancy).filter(Discrepancy.lc_id == lc_id).all()
        documents = self.db.query(Document).filter(Document.lc_id == lc_id, Document.is_active == True).all()

        output = BytesIO()
        pdf_doc = SimpleDocTemplate(output, pagesize=A4, title=f"不符点报告 - {lc.lc_number}")

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=20,
            alignment=1,
        )
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=12,
            textColor=colors.HexColor('#333366'),
        )
        normal_style = styles['Normal']

        elements = []

        elements.append(Paragraph(f"信用证单据不符点报告", title_style))
        elements.append(Paragraph(f"信用证编号: {lc.lc_number}", normal_style))
        elements.append(Spacer(1, 0.2 * inch))

        summary_data = [
            ["项目", "内容"],
            ["信用证编号", lc.lc_number],
            ["开证行", lc.issuing_bank or "-"],
            ["申请人", lc.applicant or "-"],
            ["受益人", lc.beneficiary or "-"],
            ["金额", f"{lc.currency or ''} {lc.amount:,.2f}" if lc.amount else "-"],
            ["状态", lc.status],
            ["总不符点数", str(len(discrepancies))],
            ["待处理不符点", str(len([d for d in discrepancies if d.status == "OPEN"]))],
            ["已解决不符点", str(len([d for d in discrepancies if d.status in ["RESOLVED", "CLOSED"]]))],
        ]

        summary_table = Table(summary_data, colWidths=[2 * inch, 4 * inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333366')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))

        elements.append(Paragraph("一、基本信息", heading_style))
        elements.append(summary_table)
        elements.append(Spacer(1, 0.3 * inch))

        if discrepancies:
            elements.append(Paragraph("二、不符点明细", heading_style))

            for idx, disc in enumerate(discrepancies, 1):
                cat_info = DISCREPANCY_CATEGORIES.get(disc.discrepancy_type, {"name": disc.discrepancy_type})

                disc_data = [
                    ["字段", "内容"],
                    ["不符点类型", cat_info.get("name", disc.discrepancy_type)],
                    ["严重程度", self._get_severity_text(disc.severity)],
                    ["状态", self._get_status_text(disc.status)],
                    ["描述", disc.description or "-"],
                    ["原因", disc.reason or "-"],
                    ["影响范围", disc.impact_scope or "-"],
                    ["下一步动作", disc.next_action or "-"],
                ]

                disc_table = Table(disc_data, colWidths=[1.5 * inch, 4.5 * inch])
                severity_color = SEVERITY_COLORS.get(disc.severity, {}).get("hex", "#FFCC00")
                disc_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(severity_color)),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                    ('PADDING', (0, 0), (-1, -1), 6),
                ]))

                elements.append(Paragraph(f"{idx}. {cat_info.get('name', disc.discrepancy_type)}", styles['Heading3']))
                elements.append(disc_table)
                elements.append(Spacer(1, 0.2 * inch))

        if documents:
            elements.append(Paragraph("三、单据清单", heading_style))

            doc_data = [["单据类型", "单据编号", "版本", "状态", "提交人"]]
            for doc in documents:
                doc_data.append([
                    self._get_doc_type_text(doc.document_type),
                    doc.document_number,
                    f"v{doc.version}",
                    "有效" if doc.is_active else "已作废",
                    doc.submitted_by or "-",
                ])

            doc_table = Table(doc_data, colWidths=[1.2 * inch, 2 * inch, 0.8 * inch, 0.8 * inch, 1.2 * inch])
            doc_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333366')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('PADDING', (0, 0), (-1, -1), 6),
            ]))

            elements.append(doc_table)

        pdf_doc.build(elements)
        output.seek(0)

        return output

    def _create_summary_sheet(self, wb: Workbook, lc: LetterOfCredit,
                              discrepancies: List[Discrepancy], documents: List[Document]):
        ws = wb.active
        ws.title = "汇总"

        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="333366", end_color="333366", fill_type="solid")
        normal_font = Font(size=11)

        data = [
            ["信用证不符点检查报告", ""],
            ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["", ""],
            ["信用证编号", lc.lc_number],
            ["开证行", lc.issuing_bank or "-"],
            ["申请人", lc.applicant or "-"],
            ["受益人", lc.beneficiary or "-"],
            ["金额", f"{lc.currency or ''} {lc.amount:,.2f}" if lc.amount else "-"],
            ["最迟装运期", lc.latest_shipment_date.strftime("%Y-%m-%d") if lc.latest_shipment_date else "-"],
            ["有效期", lc.expiry_date.strftime("%Y-%m-%d") if lc.expiry_date else "-"],
            ["当前状态", self._get_status_text(lc.status)],
            ["", ""],
            ["不符点统计", ""],
            ["总不符点数", len(discrepancies)],
            ["严重(Critical)", len([d for d in discrepancies if d.severity == "CRITICAL"])],
            ["高(High)", len([d for d in discrepancies if d.severity == "HIGH"])],
            ["中(Medium)", len([d for d in discrepancies if d.severity == "MEDIUM"])],
            ["低(Low)", len([d for d in discrepancies if d.severity == "LOW"])],
            ["", ""],
            ["状态统计", ""],
            ["待处理(Open)", len([d for d in discrepancies if d.status == "OPEN"])],
            ["处理中(In Progress)", len([d for d in discrepancies if d.status == "IN_PROGRESS"])],
            ["已解决(Resolved)", len([d for d in discrepancies if d.status == "RESOLVED"])],
            ["已接受(Accepted)", len([d for d in discrepancies if d.status == "ACCEPTED"])],
            ["已关闭(Closed)", len([d for d in discrepancies if d.status == "CLOSED"])],
            ["", ""],
            ["单据统计", ""],
            ["有效单据数", len(documents)],
            ["提单", len([d for d in documents if d.document_type.upper() in ["BL", "BILL_OF_LADING"]])],
            ["商业发票", len([d for d in documents if d.document_type.upper() in ["INVOICE", "COMMERCIAL_INVOICE"]])],
            ["装箱单", len([d for d in documents if d.document_type.upper() in ["PL", "PACKING_LIST"]])],
        ]

        for row_idx, row in enumerate(data, 1):
            for col_idx, value in enumerate(row, 1):
                if col_idx == 2 and row_idx in [1, 12, 19, 27]:
                    continue

                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 1:
                    cell.font = Font(bold=True, size=14)
                    cell.alignment = Alignment(horizontal="center")
                elif row_idx in [12, 19, 27]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="333366", end_color="333366", fill_type="solid")
                else:
                    cell.font = normal_font

                if col_idx == 1 and row_idx > 1:
                    cell.font = Font(bold=True, size=11)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ws.merge_cells(start_row=12, start_column=1, end_row=12, end_column=2)
        ws.merge_cells(start_row=19, start_column=1, end_row=19, end_column=2)
        ws.merge_cells(start_row=27, start_column=1, end_row=27, end_column=2)

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 60

    def _create_discrepancies_sheet(self, wb: Workbook, lc: LetterOfCredit, discrepancies: List[Discrepancy]):
        ws = wb.create_sheet("不符点明细")

        headers = ["序号", "类型", "严重程度", "状态", "描述", "原因", "影响范围", "下一步动作",
                   "涉及单据", "涉及条款", "创建时间", "更新时间"]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="333366", end_color="333366", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx, disc in enumerate(discrepancies, 2):
            cat_info = DISCREPANCY_CATEGORIES.get(disc.discrepancy_type, {"name": disc.discrepancy_type})

            values = [
                row_idx - 1,
                cat_info.get("name", disc.discrepancy_type),
                self._get_severity_text(disc.severity),
                self._get_status_text(disc.status),
                disc.description or "",
                disc.reason or "",
                disc.impact_scope or "",
                disc.next_action or "",
                disc.document_ref or "",
                disc.clause_ref or "",
                disc.created_at.strftime("%Y-%m-%d %H:%M:%S") if disc.created_at else "",
                disc.updated_at.strftime("%Y-%m-%d %H:%M:%S") if disc.updated_at else "",
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)

                if col_idx == 3:
                    color = SEVERITY_COLORS.get(disc.severity, {}).get("excel", "FFCC00")
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                    cell.font = Font(bold=True)
                elif col_idx == 4:
                    color = STATUS_COLORS.get(disc.status, {}).get("excel", "AAAAAA")
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                    cell.font = Font(bold=True)

        widths = [6, 18, 10, 12, 40, 30, 30, 30, 20, 15, 18, 18]
        for col_idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col_idx)].width = width

    def _create_documents_sheet(self, wb: Workbook, documents: List[Document]):
        ws = wb.create_sheet("单据清单")

        headers = ["序号", "单据类型", "单据编号", "版本", "状态", "提交人", "提交时间", "备注"]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="333366", end_color="333366", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for row_idx, doc in enumerate(documents, 2):
            values = [
                row_idx - 1,
                self._get_doc_type_text(doc.document_type),
                doc.document_number,
                f"v{doc.version}",
                "有效" if doc.is_active else "已作废",
                doc.submitted_by or "-",
                doc.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if doc.submitted_at else "",
                doc.remarks or "",
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)

                if col_idx == 5 and not doc.is_active:
                    cell.font = Font(color="888888", strike=True)

        widths = [6, 15, 25, 8, 10, 15, 20, 40]
        for col_idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col_idx)].width = width

    def _create_version_history_sheet(self, wb: Workbook, lc_id: int):
        ws = wb.create_sheet("版本历史")

        headers = ["序号", "关联类型", "关联ID", "操作类型", "版本", "修改原因", "操作人", "操作时间"]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="333366", end_color="333366", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        version_records = self.db.query(VersionRecord).filter(
            (VersionRecord.related_type == "LC") & (VersionRecord.related_id == lc_id) |
            (VersionRecord.related_type == "DOCUMENT") |
            (VersionRecord.related_type == "DISCREPANCY")
        ).order_by(VersionRecord.created_at.desc()).all()

        for row_idx, record in enumerate(version_records, 2):
            values = [
                row_idx - 1,
                self._get_related_type_text(record.related_type),
                record.related_id,
                self._get_action_text(record.action),
                record.version,
                record.change_reason or "-",
                record.operator or "-",
                record.created_at.strftime("%Y-%m-%d %H:%M:%S") if record.created_at else "",
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        widths = [6, 12, 10, 12, 8, 30, 15, 20]
        for col_idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col_idx)].width = width

    def _get_severity_text(self, severity: str) -> str:
        mapping = {
            "CRITICAL": "严重",
            "HIGH": "高",
            "MEDIUM": "中",
            "LOW": "低",
        }
        return mapping.get(severity, severity)

    def _get_status_text(self, status: str) -> str:
        mapping = {
            "DRAFT": "草稿",
            "CHECKING": "检查中",
            "DISCREPANCY_FOUND": "发现不符点",
            "RESOLVING": "整改中",
            "READY_FOR_SUBMISSION": "可提交",
            "SUBMITTED": "已提交",
            "OPEN": "待处理",
            "IN_PROGRESS": "处理中",
            "RESOLVED": "已解决",
            "ACCEPTED": "已接受",
            "CLOSED": "已关闭",
        }
        return mapping.get(status, status)

    def _get_doc_type_text(self, doc_type: str) -> str:
        mapping = {
            "BL": "提单",
            "BILL_OF_LADING": "提单",
            "INVOICE": "商业发票",
            "COMMERCIAL_INVOICE": "商业发票",
            "PL": "装箱单",
            "PACKING_LIST": "装箱单",
            "CERTIFICATE": "证明",
            "OTHER": "其他",
        }
        return mapping.get(doc_type.upper(), doc_type)

    def _get_related_type_text(self, related_type: str) -> str:
        mapping = {
            "LC": "信用证",
            "DOCUMENT": "单据",
            "DISCREPANCY": "不符点",
        }
        return mapping.get(related_type, related_type)

    def _get_action_text(self, action: str) -> str:
        mapping = {
            "CREATE": "创建",
            "UPDATE": "更新",
            "DEACTIVATE": "作废",
            "STATUS_CHANGE": "状态变更",
            "CORRECTION": "修正",
        }
        return mapping.get(action, action)
