from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.responses import StreamingResponse, JSONResponse
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from database import get_db, init_db
from schemas import (
    CustomerCreate, CustomerUpdate, CustomerResponse,
    SalesOrderCreate, SalesOrderResponse,
    ReturnRecordCreate, ReturnRecordResponse,
    PaymentRecordCreate, PaymentRecordResponse,
    RecalculateRequest, RecalculateResponse,
    DebtReportRequest, DebtReportResponse,
    ErrorResponse, ErrorCode
)
from services import (
    BusinessException,
    create_customer, get_customers, update_customer, get_customer_by_id,
    create_sales_order, get_sales_orders, get_sales_order_by_id,
    create_return_record, get_return_records,
    create_payment_record, get_payment_records,
    recalculate_debt, generate_debt_report, get_debt_reports
)

app = FastAPI(
    title="乡镇农资店赊销管理系统",
    description="赊销回款退货抵扣欠款重算后端API",
    version="1.0.0"
)


@app.on_event("startup")
def startup_event():
    init_db()


@app.exception_handler(BusinessException)
async def business_exception_handler(request, exc: BusinessException):
    status_code_map = {
        ErrorCode.MISSING_FIELD: status.HTTP_400_BAD_REQUEST,
        ErrorCode.INVALID_STATUS: status.HTTP_400_BAD_REQUEST,
        ErrorCode.NEED_REVIEW: status.HTTP_409_CONFLICT,
        ErrorCode.ALREADY_PROCESSED: status.HTTP_409_CONFLICT,
        ErrorCode.NOT_FOUND: status.HTTP_404_NOT_FOUND,
        ErrorCode.BUSINESS_ERROR: status.HTTP_400_BAD_REQUEST,
    }
    return JSONResponse(
        status_code=status_code_map.get(exc.code, status.HTTP_400_BAD_REQUEST),
        content={
            "code": exc.code.value,
            "message": exc.message,
            "details": exc.details
        }
    )


@app.post("/api/customers/", response_model=CustomerResponse, tags=["客户管理"])
def create_customer_api(customer: CustomerCreate, db: Session = Depends(get_db)):
    return create_customer(db, customer)


@app.get("/api/customers/", response_model=List[CustomerResponse], tags=["客户管理"])
def list_customers(
    skip: int = 0, limit: int = 100, name: str = None, db: Session = Depends(get_db)
):
    return get_customers(db, skip, limit, name)


@app.get("/api/customers/{customer_id}", response_model=CustomerResponse, tags=["客户管理"])
def get_customer(customer_id: int, db: Session = Depends(get_db)):
    customer = get_customer_by_id(db, customer_id)
    if not customer:
        raise BusinessException(
            code=ErrorCode.NOT_FOUND,
            message=f"客户ID {customer_id} 不存在"
        )
    return customer


@app.put("/api/customers/{customer_id}", response_model=CustomerResponse, tags=["客户管理"])
def update_customer_api(
    customer_id: int, customer_update: CustomerUpdate, db: Session = Depends(get_db)
):
    return update_customer(db, customer_id, customer_update)


@app.post("/api/sales-orders/", response_model=SalesOrderResponse, tags=["赊销订单"])
def create_sales_order_api(order: SalesOrderCreate, db: Session = Depends(get_db)):
    return create_sales_order(db, order)


@app.get("/api/sales-orders/", response_model=List[SalesOrderResponse], tags=["赊销订单"])
def list_sales_orders(
    skip: int = 0, limit: int = 100,
    customer_id: int = None, status: str = None,
    db: Session = Depends(get_db)
):
    return get_sales_orders(db, skip, limit, customer_id, status)


@app.get("/api/sales-orders/{order_id}", response_model=SalesOrderResponse, tags=["赊销订单"])
def get_sales_order(order_id: int, db: Session = Depends(get_db)):
    order = get_sales_order_by_id(db, order_id)
    if not order:
        raise BusinessException(
            code=ErrorCode.NOT_FOUND,
            message=f"订单ID {order_id} 不存在"
        )
    return order


@app.post("/api/returns/", response_model=ReturnRecordResponse, tags=["退货管理"])
def create_return_api(return_data: ReturnRecordCreate, db: Session = Depends(get_db)):
    return create_return_record(db, return_data)


@app.get("/api/returns/", response_model=List[ReturnRecordResponse], tags=["退货管理"])
def list_returns(
    skip: int = 0, limit: int = 100, order_id: int = None, db: Session = Depends(get_db)
):
    return get_return_records(db, skip, limit, order_id)


@app.post("/api/payments/", response_model=PaymentRecordResponse, tags=["回款管理"])
def create_payment_api(payment: PaymentRecordCreate, db: Session = Depends(get_db)):
    return create_payment_record(db, payment)


@app.get("/api/payments/", response_model=List[PaymentRecordResponse], tags=["回款管理"])
def list_payments(
    skip: int = 0, limit: int = 100, order_id: int = None, db: Session = Depends(get_db)
):
    return get_payment_records(db, skip, limit, order_id)


@app.post("/api/recalculate/", response_model=RecalculateResponse, tags=["欠款重算"])
def recalculate_debt_api(request: RecalculateRequest, db: Session = Depends(get_db)):
    db_order, previous_debt = recalculate_debt(db, request.order_id)
    return RecalculateResponse(
        order_id=db_order.id,
        order_no=db_order.order_no,
        previous_debt=previous_debt,
        new_debt=db_order.debt_amount,
        total_paid=db_order.paid_amount,
        total_returned=db_order.returned_amount,
        difference=db_order.debt_amount - previous_debt
    )


@app.post("/api/debt-reports/", response_model=DebtReportResponse, tags=["欠款报告"])
def create_debt_report_api(request: DebtReportRequest, db: Session = Depends(get_db)):
    report = generate_debt_report(db, request)
    customer_name = None
    if report.customer_id:
        customer = get_customer_by_id(db, report.customer_id)
        customer_name = customer.name if customer else None
    
    return DebtReportResponse(
        id=report.id,
        report_no=report.report_no,
        report_date=report.report_date,
        customer_id=report.customer_id,
        customer_name=customer_name,
        period_start=report.period_start,
        period_end=report.period_end,
        total_debt=report.total_debt,
        total_paid=report.total_paid,
        total_returned=report.total_returned,
        final_debt=report.final_debt,
        status=report.status,
        need_review=report.need_review,
        review_remarks=report.review_remarks
    )


@app.get("/api/debt-reports/", response_model=List[DebtReportResponse], tags=["欠款报告"])
def list_debt_reports(
    skip: int = 0, limit: int = 100, customer_id: int = None, db: Session = Depends(get_db)
):
    reports = get_debt_reports(db, skip, limit, customer_id)
    result = []
    for report in reports:
        customer_name = None
        if report.customer_id:
            customer = get_customer_by_id(db, report.customer_id)
            customer_name = customer.name if customer else None
        result.append(DebtReportResponse(
            id=report.id,
            report_no=report.report_no,
            report_date=report.report_date,
            customer_id=report.customer_id,
            customer_name=customer_name,
            period_start=report.period_start,
            period_end=report.period_end,
            total_debt=report.total_debt,
            total_paid=report.total_paid,
            total_returned=report.total_returned,
            final_debt=report.final_debt,
            status=report.status,
            need_review=report.need_review,
            review_remarks=report.review_remarks
        ))
    return result


@app.get("/api/debt-reports/{report_id}/export", tags=["欠款报告"])
def export_debt_report(report_id: int, db: Session = Depends(get_db)):
    reports = get_debt_reports(db, 0, 1, None)
    report = next((r for r in reports if r.id == report_id), None)
    if not report:
        raise BusinessException(
            code=ErrorCode.NOT_FOUND,
            message=f"报告ID {report_id} 不存在"
        )
    
    customer_name = "所有客户"
    if report.customer_id:
        customer = get_customer_by_id(db, report.customer_id)
        customer_name = customer.name if customer else "未知客户"

    wb = Workbook()
    ws = wb.active
    ws.title = "欠款报告"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    headers = ["项目", "金额", "说明"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    data = [
        ["报告编号", report.report_no, ""],
        ["客户", customer_name, ""],
        ["报告日期", report.report_date.strftime("%Y-%m-%d %H:%M:%S"), ""],
        ["", "", ""],
        ["初始欠款总额", report.total_debt, "所有订单实际金额总和"],
        ["已回款总额", report.total_paid, "已确认回款金额"],
        ["退货抵扣总额", report.total_returned, "已确认退货抵扣金额"],
        ["", "", ""],
        ["最终欠款", report.final_debt, "= 初始欠款 - 已回款 - 退货抵扣"],
        ["", "", ""],
        ["是否需要复核", "是" if report.need_review else "否", ""],
    ]

    for row, row_data in enumerate(data, 2):
        for col, value in enumerate(row_data, 1):
            ws.cell(row=row, column=col, value=value)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=debt_report_{report.report_no}.xlsx"}
    )


@app.get("/api/health", tags=["系统"])
def health_check():
    return {"status": "ok", "timestamp": datetime.now().isoformat()}
