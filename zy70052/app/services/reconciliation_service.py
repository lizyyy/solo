from typing import Optional, Dict, Any, List
from datetime import datetime, date
from decimal import Decimal
import json
from io import BytesIO

from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.models import (
    ExtensionApplication, LoanAccount, RepaymentPlan,
    PenaltySnapshot, CreditLimitRecord, ApprovalHistory,
    AuditLog
)
from app.utils import (
    ApplicationStatus, AccountStatus, IdGenerator, DateTimeUtils,
    BusinessType, OperationType
)


class ReconciliationService:
    def __init__(self, db: Session):
        self.db = db
    
    def get_extension_reconciliation_data(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
        status: Optional[ApplicationStatus] = None
    ) -> List[Dict[str, Any]]:
        query = self.db.query(ExtensionApplication).order_by(
            ExtensionApplication.requested_at.desc()
        )
        
        if start_date:
            query = query.filter(ExtensionApplication.requested_at >= start_date)
        if end_date:
            query = query.filter(ExtensionApplication.requested_at <= end_date)
        if status:
            query = query.filter(ExtensionApplication.status == status.value)
        
        applications = query.all()
        
        result = []
        for app in applications:
            loan_account = self.db.query(LoanAccount).get(app.loan_account_id)
            
            approval_histories = self.db.query(ApprovalHistory).filter(
                ApprovalHistory.application_id == app.id
            ).order_by(ApprovalHistory.operation_time.asc()).all()
            
            penalty_snapshots = self.db.query(PenaltySnapshot).filter(
                PenaltySnapshot.extension_application_id == app.id
            ).all()
            
            result.append({
                "application_no": app.application_no,
                "account_no": loan_account.account_no if loan_account else None,
                "customer_name": loan_account.customer_name if loan_account else None,
                "extension_months": app.extension_months,
                "extension_reason": app.extension_reason,
                "status": app.status,
                "requested_at": app.requested_at.isoformat() if app.requested_at else None,
                "applicant": app.applicant_name,
                "first_approved_at": app.first_approval_at.isoformat() if app.first_approval_at else None,
                "final_approved_at": app.final_approval_at.isoformat() if app.final_approval_at else None,
                "executed_at": app.executed_at.isoformat() if app.executed_at else None,
                "executed_status": app.executed_status,
                "execution_error": app.execution_error,
                "approval_count": len(approval_histories),
                "penalty_amount": sum(s.penalty_amount for s in penalty_snapshots) if penalty_snapshots else Decimal("0"),
                "rejected_at": app.rejected_at.isoformat() if app.rejected_at else None,
                "reject_reason": app.reject_reason,
                "cancelled_at": app.cancelled_at.isoformat() if app.cancelled_at else None
            })
        
        return result
    
    def get_loan_account_reconciliation_data(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None
    ) -> List[Dict[str, Any]]:
        query = self.db.query(LoanAccount).order_by(LoanAccount.created_at.desc())
        
        if start_date:
            query = query.filter(LoanAccount.created_at >= start_date)
        if end_date:
            query = query.filter(LoanAccount.created_at <= end_date)
        
        accounts = query.all()
        
        result = []
        for acc in accounts:
            current_plan = self.db.query(RepaymentPlan).filter(
                RepaymentPlan.loan_account_id == acc.id,
                RepaymentPlan.is_current == True
            ).first()
            
            extension_apps = self.db.query(ExtensionApplication).filter(
                ExtensionApplication.loan_account_id == acc.id
            ).all()
            
            executed_count = sum(
                1 for app in extension_apps 
                if app.status == ApplicationStatus.EXECUTED.value
            )
            
            result.append({
                "account_no": acc.account_no,
                "customer_id": acc.customer_id,
                "customer_name": acc.customer_name,
                "status": acc.status,
                "loan_amount": str(acc.loan_amount),
                "remaining_principal": str(acc.remaining_principal),
                "total_interest": str(acc.total_interest),
                "paid_interest": str(acc.paid_interest),
                "annual_interest_rate": str(acc.annual_interest_rate),
                "loan_term": acc.loan_term,
                "original_maturity_date": acc.original_maturity_date.isoformat() if acc.original_maturity_date else None,
                "current_maturity_date": acc.current_maturity_date.isoformat() if acc.current_maturity_date else None,
                "credit_limit_used": str(acc.credit_limit_used),
                "extension_count": acc.extension_count,
                "max_extension_count": acc.max_extension_count,
                "current_plan_no": current_plan.plan_no if current_plan else None,
                "executed_extensions": executed_count,
                "version": acc.version
            })
        
        return result
    
    def get_limit_reconciliation_data(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None
    ) -> List[Dict[str, Any]]:
        query = self.db.query(CreditLimitRecord).order_by(
            CreditLimitRecord.operation_time.desc()
        )
        
        if start_date:
            query = query.filter(CreditLimitRecord.operation_time >= start_date)
        if end_date:
            query = query.filter(CreditLimitRecord.operation_time <= end_date)
        
        records = query.all()
        
        return [
            {
                "record_no": r.record_no,
                "customer_id": r.customer_id,
                "operation_type": r.operation_type,
                "operation_reason": r.operation_reason,
                "before_limit": str(r.before_limit),
                "before_used": str(r.before_used),
                "after_limit": str(r.after_limit),
                "after_used": str(r.after_used),
                "change_amount": str(r.change_amount),
                "operation_time": r.operation_time.isoformat() if r.operation_time else None,
                "related_application_id": r.related_application_id,
                "related_business_no": r.related_business_no
            }
            for r in records
        ]
    
    def get_audit_reconciliation_data(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
        business_type: Optional[BusinessType] = None
    ) -> List[Dict[str, Any]]:
        query = self.db.query(AuditLog).order_by(AuditLog.audit_time.desc())
        
        if start_date:
            query = query.filter(AuditLog.audit_time >= start_date)
        if end_date:
            query = query.filter(AuditLog.audit_time <= end_date)
        if business_type:
            query = query.filter(AuditLog.business_type == business_type.value)
        
        logs = query.limit(1000).all()
        
        return [
            {
                "log_no": l.log_no,
                "audit_time": l.audit_time.isoformat() if l.audit_time else None,
                "business_type": l.business_type,
                "business_no": l.business_no,
                "operation_type": l.operation_type,
                "operation_desc": l.operation_desc,
                "operator_id": l.operator_id,
                "operator_name": l.operator_name,
                "operator_type": l.operator_type,
                "request_path": l.request_path,
                "request_method": l.request_method,
                "request_ip": l.request_ip,
                "operation_result": l.operation_result,
                "error_message": l.error_message
            }
            for l in logs
        ]
    
    def export_to_excel(
        self,
        data: List[Dict[str, Any]],
        sheet_name: str = "Sheet1"
    ) -> bytes:
        if not data:
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(["暂无数据"])
            
            output = BytesIO()
            wb.save(output)
            return output.getvalue()
        
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = list(data[0].keys())
        ws.append(headers)
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, key in enumerate(headers, start=1):
                value = row_data.get(key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else "")
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='left', vertical='center')
        
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 20
        
        ws.freeze_panes = "A2"
        
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
    
    def export_extension_reconciliation(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
        status: Optional[ApplicationStatus] = None
    ) -> bytes:
        data = self.get_extension_reconciliation_data(start_date, end_date, status)
        return self.export_to_excel(data, "展期申请对账")
    
    def export_loan_account_reconciliation(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None
    ) -> bytes:
        data = self.get_loan_account_reconciliation_data(start_date, end_date)
        return self.export_to_excel(data, "贷款账户对账")
    
    def export_limit_reconciliation(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None
    ) -> bytes:
        data = self.get_limit_reconciliation_data(start_date, end_date)
        return self.export_to_excel(data, "额度变更对账")
    
    def export_audit_reconciliation(
        self,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None,
        business_type: Optional[BusinessType] = None
    ) -> bytes:
        data = self.get_audit_reconciliation_data(start_date, end_date, business_type)
        return self.export_to_excel(data, "审计日志对账")
    
    def check_data_consistency(
        self,
        application_no: str
    ) -> Dict[str, Any]:
        from app.services.extension_service import ExtensionApplicationException
        
        application = self.db.query(ExtensionApplication).filter(
            ExtensionApplication.application_no == application_no
        ).first()
        
        if not application:
            raise ExtensionApplicationException(f"申请不存在: {application_no}")
        
        loan_account = self.db.query(LoanAccount).get(application.loan_account_id)
        
        consistency_checks = []
        
        if application.status == ApplicationStatus.EXECUTED.value:
            if loan_account:
                expected_extension_count = loan_account.extension_count
                actual_extensions = self.db.query(ExtensionApplication).filter(
                    ExtensionApplication.loan_account_id == loan_account.id,
                    ExtensionApplication.status == ApplicationStatus.EXECUTED.value
                ).count()
                
                consistency_checks.append({
                    "check": "展期次数一致性",
                    "expected": expected_extension_count,
                    "actual": actual_extensions,
                    "passed": expected_extension_count == actual_extensions
                })
            
            if application.new_repayment_plan_id:
                new_plan = self.db.query(RepaymentPlan).get(application.new_repayment_plan_id)
                if new_plan:
                    consistency_checks.append({
                        "check": "新还款计划状态一致性",
                        "expected": True,
                        "actual": new_plan.is_current,
                        "passed": new_plan.is_current
                    })
                    
                    if loan_account:
                        consistency_checks.append({
                            "check": "账户当前计划关联一致性",
                            "expected": new_plan.id,
                            "actual": loan_account.current_repayment_plan_id,
                            "passed": loan_account.current_repayment_plan_id == new_plan.id
                        })
        
        if application.status in [
            ApplicationStatus.FINAL_APPROVED.value,
            ApplicationStatus.EXECUTED.value,
            ApplicationStatus.EXECUTE_FAILED.value
        ]:
            penalty_count = self.db.query(PenaltySnapshot).filter(
                PenaltySnapshot.extension_application_id == application.id
            ).count()
            
            consistency_checks.append({
                "check": "罚息快照存在性",
                "expected": ">=1",
                "actual": penalty_count,
                "passed": penalty_count >= 1
            })
        
        approval_count = self.db.query(ApprovalHistory).filter(
            ApprovalHistory.application_id == application.id
        ).count()
        
        expected_approvals = 1
        if application.status in [
            ApplicationStatus.FIRST_APPROVED.value,
            ApplicationStatus.FINAL_APPROVED.value,
            ApplicationStatus.EXECUTED.value,
            ApplicationStatus.EXECUTE_FAILED.value
        ]:
            expected_approvals = 2
        if application.status in [
            ApplicationStatus.EXECUTED.value,
            ApplicationStatus.EXECUTE_FAILED.value
        ]:
            expected_approvals = 3
        
        consistency_checks.append({
            "check": "审批历史完整性",
            "expected": f">={expected_approvals}",
            "actual": approval_count,
            "passed": approval_count >= expected_approvals
        })
        
        all_passed = all(check["passed"] for check in consistency_checks)
        
        return {
            "application_no": application_no,
            "status": application.status,
            "consistency_passed": all_passed,
            "checks": consistency_checks
        }
