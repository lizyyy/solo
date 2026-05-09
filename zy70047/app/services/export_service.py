import os
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import Config
from app import db
from app.models import SparePart, Reservation, StockAlarm, OperationLog, WorkOrder


def ensure_export_dir():
    if not os.path.exists(Config.EXPORT_DIR):
        os.makedirs(Config.EXPORT_DIR)
    return Config.EXPORT_DIR


def generate_filename(prefix, ext='xlsx'):
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'{prefix}_{timestamp}.{ext}'


def export_stock_alarm_review():
    ensure_export_dir()
    filename = generate_filename('stock_alarm_review')
    filepath = os.path.join(Config.EXPORT_DIR, filename)
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = '待复核告警'
    
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    
    headers = ['告警ID', '备件编码', '备件名称', '告警类型', '告警级别', 
               '当前库存', '已预占', '可用库存', '最小库存', '安全库存',
               '关联工单', '关联预占ID', '告警内容', '是否需要人工复核',
               '创建时间']
    
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    alarms = StockAlarm.query.filter(
        StockAlarm.reviewed == False
    ).order_by(StockAlarm.needs_review.desc(), StockAlarm.created_at.desc()).all()
    
    alarm_type_map = {
        'STOCK_OUT': '库存为0',
        'PARTIAL_RESERVE': '部分预占',
        'SAFETY_STOCK_BREACH': '安全库存不足',
        'MIN_STOCK_BREACH': '最小库存不足'
    }
    
    for row_idx, alarm in enumerate(alarms, 2):
        ws1.cell(row=row_idx, column=1, value=alarm.alarm_id)
        ws1.cell(row=row_idx, column=2, value=alarm.part_code)
        ws1.cell(row=row_idx, column=3, value=alarm.part_name)
        ws1.cell(row=row_idx, column=4, value=alarm_type_map.get(alarm.alarm_type, alarm.alarm_type))
        ws1.cell(row=row_idx, column=5, value=alarm.alarm_level)
        ws1.cell(row=row_idx, column=6, value=alarm.current_stock)
        ws1.cell(row=row_idx, column=7, value=alarm.reserved_qty)
        ws1.cell(row=row_idx, column=8, value=alarm.available_qty)
        ws1.cell(row=row_idx, column=9, value=alarm.min_stock)
        ws1.cell(row=row_idx, column=10, value=alarm.safety_stock)
        ws1.cell(row=row_idx, column=11, value=alarm.related_order_no or '')
        ws1.cell(row=row_idx, column=12, value=alarm.related_reservation_id or '')
        ws1.cell(row=row_idx, column=13, value=alarm.message)
        ws1.cell(row=row_idx, column=14, value='是' if alarm.needs_review else '否')
        ws1.cell(row=row_idx, column=15, value=alarm.created_at.strftime('%Y-%m-%d %H:%M:%S') if alarm.created_at else '')
    
    ws2 = wb.create_sheet('备件池概览')
    
    headers2 = ['备件编码', '备件名称', '分类', '当前库存', '已预占', '可用库存',
                '最小库存', '安全库存', '库存状态', '预占工单数', '单价']
    
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    parts = SparePart.query.all()
    
    for row_idx, part in enumerate(parts, 2):
        if part.available_qty < part.min_stock:
            status = '严重不足'
        elif part.available_qty < part.safety_stock:
            status = '预警'
        else:
            status = '正常'
        
        reservation_count = Reservation.query.filter(
            Reservation.part_code == part.part_code,
            Reservation.status.in_(['RESERVED', 'PARTIAL'])
        ).count()
        
        ws2.cell(row=row_idx, column=1, value=part.part_code)
        ws2.cell(row=row_idx, column=2, value=part.part_name)
        ws2.cell(row=row_idx, column=3, value=part.category or '')
        ws2.cell(row=row_idx, column=4, value=part.total_stock)
        ws2.cell(row=row_idx, column=5, value=part.reserved_qty)
        ws2.cell(row=row_idx, column=6, value=part.available_qty)
        ws2.cell(row=row_idx, column=7, value=part.min_stock)
        ws2.cell(row=row_idx, column=8, value=part.safety_stock)
        ws2.cell(row=row_idx, column=9, value=status)
        ws2.cell(row=row_idx, column=10, value=reservation_count)
        ws2.cell(row=row_idx, column=11, value=part.unit_price)
    
    ws3 = wb.create_sheet('活跃预占记录')
    
    headers3 = ['预占ID', '工单编号', '备件编码', '备件名称', '申请数量',
                '已预占', '已领用', '已释放', '状态', '使用替代件',
                '预占人', '预占时间']
    
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    reservations = Reservation.query.filter(
        Reservation.status.in_(['RESERVED', 'PARTIAL'])
    ).order_by(Reservation.reserved_at.desc()).all()
    
    status_map = {
        'PENDING': '待处理',
        'RESERVED': '已预占',
        'PARTIAL': '部分预占',
        'COMPLETED': '已完成',
        'RELEASED': '已释放'
    }
    
    for row_idx, r in enumerate(reservations, 2):
        ws3.cell(row=row_idx, column=1, value=r.reservation_id)
        ws3.cell(row=row_idx, column=2, value=r.order_no)
        ws3.cell(row=row_idx, column=3, value=r.part_code)
        ws3.cell(row=row_idx, column=4, value=r.part_name)
        ws3.cell(row=row_idx, column=5, value=r.requested_qty)
        ws3.cell(row=row_idx, column=6, value=r.reserved_qty)
        ws3.cell(row=row_idx, column=7, value=r.used_qty)
        ws3.cell(row=row_idx, column=8, value=r.released_qty)
        ws3.cell(row=row_idx, column=9, value=status_map.get(r.status, r.status))
        ws3.cell(row=row_idx, column=10, value='是' if r.used_substitute else '否')
        ws3.cell(row=row_idx, column=11, value=r.reserved_by or '')
        ws3.cell(row=row_idx, column=12, value=r.reserved_at.strftime('%Y-%m-%d %H:%M:%S') if r.reserved_at else '')
    
    ws4 = wb.create_sheet('复核说明')
    ws4.cell(row=1, column=1, value='【库存告警业务复核说明】').font = Font(bold=True, size=14)
    ws4.cell(row=2, column=1, value='导出时间: ' + datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    
    review_guide = [
        '',
        '一、需要人工复核的情况 (needs_review=True):',
        '   1. 库存为0 (STOCK_OUT) - 需确认紧急采购或转用替代件',
        '   2. 部分预占 (PARTIAL_RESERVE) - 需确认是否接受部分供应或调整工单',
        '   3. 可用库存低于最小库存 (MIN_STOCK_BREACH) - 需启动紧急采购流程',
        '   4. 使用无效替代件 - 需确认替代关系配置',
        '',
        '二、可自动处理的情况:',
        '   1. 可用库存低于安全库存但高于最小库存 (WARNING级别)',
        '   2. 有可用替代件时的库存不足',
        '   3. 重复请求的幂等返回',
        '',
        '三、复核要点:',
        '   1. 检查关联工单优先级，优先保障紧急工单',
        '   2. 确认替代件的可用性和兼容性',
        '   3. 核实库存数据准确性，排除数据异常',
        '   4. 记录复核结果和处理措施'
    ]
    
    for i, line in enumerate(review_guide):
        ws4.cell(row=i + 4, column=1, value=line)
    
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column].width = adjusted_width
    
    wb.save(filepath)
    return filepath


def export_operation_logs(start_date=None, end_date=None, operation_type=None):
    ensure_export_dir()
    filename = generate_filename('operation_logs')
    filepath = os.path.join(Config.EXPORT_DIR, filename)
    
    wb = Workbook()
    ws = wb.active
    ws.title = '操作日志'
    
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
    
    headers = ['日志ID', '操作类型', '工单编号', '备件编码', '变更数量', '变更金额',
               '操作人', '操作时间', '备注']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    query = OperationLog.query
    if start_date:
        query = query.filter(OperationLog.operation_time >= start_date)
    if end_date:
        query = query.filter(OperationLog.operation_time <= end_date)
    if operation_type:
        query = query.filter(OperationLog.operation_type == operation_type)
    
    logs = query.order_by(OperationLog.operation_time.desc()).all()
    
    op_type_map = {
        'WORK_ORDER_CREATE': '创建工单',
        'RESERVE_SUCCESS': '预占成功',
        'RESERVE_FAIL': '预占失败',
        'USE_SPARE': '领用备件',
        'RELEASE_RESERVATION': '释放预占',
        'STOCK_IN': '入库',
        'STOCK_ADJUST': '库存调整'
    }
    
    for row_idx, log in enumerate(logs, 2):
        ws.cell(row=row_idx, column=1, value=log.log_id)
        ws.cell(row=row_idx, column=2, value=op_type_map.get(log.operation_type, log.operation_type))
        ws.cell(row=row_idx, column=3, value=log.order_no or '')
        ws.cell(row=row_idx, column=4, value=log.part_code or '')
        ws.cell(row=row_idx, column=5, value=log.change_qty)
        ws.cell(row=row_idx, column=6, value=log.change_amount)
        ws.cell(row=row_idx, column=7, value=log.operator or '')
        ws.cell(row=row_idx, column=8, value=log.operation_time.strftime('%Y-%m-%d %H:%M:%S') if log.operation_time else '')
        ws.cell(row=row_idx, column=9, value=log.note or '')
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(filepath)
    return filepath


def get_export_list():
    ensure_export_dir()
    files = []
    for f in os.listdir(Config.EXPORT_DIR):
        if f.endswith('.xlsx'):
            filepath = os.path.join(Config.EXPORT_DIR, f)
            stat = os.stat(filepath)
            files.append({
                'filename': f,
                'filepath': filepath,
                'size': stat.st_size,
                'created_at': datetime.fromtimestamp(stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S')
            })
    files.sort(key=lambda x: x['created_at'], reverse=True)
    return files
