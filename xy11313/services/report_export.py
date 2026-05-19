import pandas as pd
from datetime import datetime
from typing import List, Dict, Optional
from sqlalchemy.orm import Session
from sqlalchemy import and_, or_, func
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models import (
    Incident, Bus, Route, RouteStop, Driver,
    StudentAppeal, GPSRecord, IncidentEvidence,
    AnomalyType, Responsibility
)

class IncidentQuery:
    def __init__(self, db_session: Session):
        self.db = db_session
    
    def query_incidents(self, 
                        responsibility: Optional[str] = None,
                        status: Optional[str] = None,
                        anomaly_type: Optional[str] = None,
                        start_date: Optional[datetime] = None,
                        end_date: Optional[datetime] = None,
                        bus_id: Optional[int] = None,
                        route_id: Optional[int] = None,
                        driver_id: Optional[int] = None,
                        page: int = 1,
                        page_size: int = 50) -> Dict:
        
        query = self.db.query(Incident)
        
        if responsibility:
            query = query.filter_by(responsibility=responsibility)
        
        if status:
            query = query.filter_by(status=status)
        
        if anomaly_type:
            query = query.filter_by(anomaly_type=anomaly_type)
        
        if start_date:
            query = query.filter(Incident.incident_date >= start_date)
        
        if end_date:
            query = query.filter(Incident.incident_date <= end_date)
        
        if bus_id:
            query = query.filter_by(bus_id=bus_id)
        
        if route_id:
            query = query.filter_by(route_id=route_id)
        
        if driver_id:
            bus_ids = [b.id for b in self.db.query(Bus).filter_by(driver_id=driver_id).all()]
            if bus_ids:
                query = query.filter(Incident.bus_id.in_(bus_ids))
        
        total_count = query.count()
        
        incidents = query.order_by(Incident.incident_date.desc()) \
                         .offset((page - 1) * page_size) \
                         .limit(page_size) \
                         .all()
        
        result_list = []
        for inc in incidents:
            bus = self.db.query(Bus).get(inc.bus_id)
            route = self.db.query(Route).get(inc.route_id)
            stop = self.db.query(RouteStop).get(inc.stop_id)
            driver = None
            if bus:
                driver = self.db.query(Driver).filter_by(bus_id=bus.id).first()
            
            result_list.append({
                'id': inc.id,
                'incident_number': inc.incident_number,
                'incident_date': inc.incident_date.strftime('%Y-%m-%d %H:%M:%S'),
                'anomaly_type': inc.anomaly_type,
                'anomaly_type_name': self._get_anomaly_type_name(inc.anomaly_type),
                'bus_number': bus.bus_number if bus else '',
                'route_name': route.route_name if route else '',
                'stop_name': stop.stop_name if stop else '',
                'driver_name': driver.name if driver else '',
                'scheduled_time': inc.scheduled_time,
                'actual_time': inc.actual_time,
                'time_difference_minutes': inc.time_difference_minutes,
                'description': inc.description,
                'responsibility': inc.responsibility,
                'responsibility_name': self._get_responsibility_name(inc.responsibility),
                'status': inc.status,
                'reviewed_by': inc.reviewed_by,
                'reviewed_at': inc.reviewed_at.strftime('%Y-%m-%d %H:%M:%S') if inc.reviewed_at else '',
                'created_at': inc.created_at.strftime('%Y-%m-%d %H:%M:%S')
            })
        
        return {
            'total_count': total_count,
            'page': page,
            'page_size': page_size,
            'total_pages': (total_count + page_size - 1) // page_size,
            'data': result_list
        }
    
    def _get_anomaly_type_name(self, anomaly_type: str) -> str:
        names = {
            'late_arrival': '晚点到达',
            'early_departure': '提前发车',
            'missing_stop': '漏站',
            'speeding': '超速',
            'route_deviation': '路线偏离',
            'driver_absent': '司机缺席',
            'parent_appeal': '家长申诉',
            'gps_mismatch': 'GPS数据不符'
        }
        return names.get(anomaly_type, anomaly_type)
    
    def _get_responsibility_name(self, responsibility: str) -> str:
        names = {
            'unassigned': '未分配',
            'driver': '司机责任',
            'traffic': '交通原因',
            'school': '学校原因',
            'parent': '家长原因',
            'weather': '天气原因',
            'other': '其他原因'
        }
        return names.get(responsibility, responsibility)
    
    def get_statistics(self, start_date: Optional[datetime] = None, 
                       end_date: Optional[datetime] = None) -> Dict:
        query = self.db.query(Incident)
        
        if start_date:
            query = query.filter(Incident.incident_date >= start_date)
        if end_date:
            query = query.filter(Incident.incident_date <= end_date)
        
        total_incidents = query.count()
        
        by_type = {}
        for anomaly_type in [a.value for a in AnomalyType]:
            count = query.filter_by(anomaly_type=anomaly_type).count()
            if count > 0:
                by_type[self._get_anomaly_type_name(anomaly_type)] = count
        
        by_responsibility = {}
        for resp in [r.value for r in Responsibility]:
            count = query.filter_by(responsibility=resp).count()
            if count > 0:
                by_responsibility[self._get_responsibility_name(resp)] = count
        
        by_status = {
            'pending': query.filter_by(status='pending').count(),
            'reviewed': query.filter_by(status='reviewed').count(),
            'resolved': query.filter_by(status='resolved').count()
        }
        
        return {
            'total_incidents': total_incidents,
            'by_anomaly_type': by_type,
            'by_responsibility': by_responsibility,
            'by_status': by_status
        }

class ReportExporter:
    def __init__(self, db_session: Session):
        self.db = db_session
    
    def export_incidents_to_excel(self, 
                                   output_file: str,
                                   responsibility: Optional[str] = None,
                                   status: Optional[str] = None,
                                   anomaly_type: Optional[str] = None,
                                   start_date: Optional[datetime] = None,
                                   end_date: Optional[datetime] = None) -> Dict:
        query_service = IncidentQuery(self.db)
        result = query_service.query_incidents(
            responsibility=responsibility,
            status=status,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            page=1,
            page_size=10000
        )
        
        wb = Workbook()
        ws = wb.active
        ws.title = "事件报告"
        
        headers = [
            '事件编号', '事件日期', '异常类型', '车牌号', 
            '线路名称', '站点名称', '司机姓名', '计划时间',
            '实际时间', '延误(分钟)', '事件描述', '责任方',
            '状态', '复核人', '复核时间'
        ]
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        for row, inc in enumerate(result['data'], 2):
            data = [
                inc['incident_number'],
                inc['incident_date'],
                inc['anomaly_type_name'],
                inc['bus_number'],
                inc['route_name'],
                inc['stop_name'],
                inc['driver_name'],
                inc['scheduled_time'],
                inc['actual_time'],
                inc['time_difference_minutes'],
                inc['description'],
                inc['responsibility_name'],
                inc['status'],
                inc['reviewed_by'],
                inc['reviewed_at']
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = thin_border
                
                if inc['responsibility'] == 'driver':
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                elif inc['responsibility'] == 'traffic':
                    cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        
        column_widths = [18, 20, 15, 12, 20, 20, 12, 10, 10, 12, 30, 15, 10, 12, 20]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'
        
        self._add_summary_sheet(wb, result['data'])
        
        wb.save(output_file)
        
        return {
            'success': True,
            'output_file': output_file,
            'total_records': result['total_count']
        }
    
    def _add_summary_sheet(self, wb: Workbook, incidents: List[Dict]):
        ws = wb.create_sheet(title="统计摘要")
        
        header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        
        ws.cell(row=1, column=1, value='校车调度异常事件统计报告').font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value=f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        
        ws.cell(row=4, column=1, value='按异常类型统计').font = Font(bold=True, size=12)
        ws.cell(row=5, column=1, value='异常类型').fill = header_fill
        ws.cell(row=5, column=1).font = header_font
        ws.cell(row=5, column=2, value='数量').fill = header_fill
        ws.cell(row=5, column=2).font = header_font
        
        type_counts = {}
        for inc in incidents:
            at = inc['anomaly_type_name']
            type_counts[at] = type_counts.get(at, 0) + 1
        
        row = 6
        for atype, count in type_counts.items():
            ws.cell(row=row, column=1, value=atype)
            ws.cell(row=row, column=2, value=count)
            row += 1
        
        row += 2
        ws.cell(row=row, column=1, value='按责任方统计').font = Font(bold=True, size=12)
        row += 1
        ws.cell(row=row, column=1, value='责任方').fill = header_fill
        ws.cell(row=row, column=1).font = header_font
        ws.cell(row=row, column=2, value='数量').fill = header_fill
        ws.cell(row=row, column=2).font = header_font
        
        resp_counts = {}
        for inc in incidents:
            resp = inc['responsibility_name']
            resp_counts[resp] = resp_counts.get(resp, 0) + 1
        
        row += 1
        for resp, count in resp_counts.items():
            ws.cell(row=row, column=1, value=resp)
            ws.cell(row=row, column=2, value=count)
            row += 1
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 10
    
    def export_incidents_to_csv(self, 
                                 output_file: str,
                                 responsibility: Optional[str] = None,
                                 status: Optional[str] = None,
                                 anomaly_type: Optional[str] = None,
                                 start_date: Optional[datetime] = None,
                                 end_date: Optional[datetime] = None) -> Dict:
        query_service = IncidentQuery(self.db)
        result = query_service.query_incidents(
            responsibility=responsibility,
            status=status,
            anomaly_type=anomaly_type,
            start_date=start_date,
            end_date=end_date,
            page=1,
            page_size=10000
        )
        
        df = pd.DataFrame(result['data'])
        df.to_csv(output_file, index=False, encoding='utf-8-sig')
        
        return {
            'success': True,
            'output_file': output_file,
            'total_records': result['total_count']
        }
