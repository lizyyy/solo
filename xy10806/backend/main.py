from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from sqlalchemy.orm import Session
from pydantic import BaseModel
from datetime import datetime, timedelta
import time
import json
from typing import Optional, List
from io import BytesIO
from openpyxl import Workbook

from database import (
    Base, engine, get_db,
    ClientApplication, UserConsent, AccessToken, RefreshToken,
    RevocationEvent, BackgroundTask, RequestLog,
    AuthorizationStatus, TokenStatus, TaskStatus, RevocationStatus
)
from services import (
    generate_id, RevocationOrchestrator, CompensationService,
    RequestLogger, TaskInterceptorService, DataGenerator
)

Base.metadata.create_all(bind=engine)

app = FastAPI(title="OAuth 授权撤回中心 API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup_event():
    db = next(get_db())
    DataGenerator.create_sample_data(db)


class RevocationRequest(BaseModel):
    consent_id: str
    user_id: str
    reason: str
    initiated_by: str = "user"


class CompensationRequest(BaseModel):
    consent_id: str
    reason: str


class TaskExecutionRequest(BaseModel):
    task_id: str


@app.get("/api/health")
def health_check():
    return {"status": "healthy", "timestamp": datetime.utcnow().isoformat()}


@app.get("/api/applications")
def list_applications(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    apps = db.query(ClientApplication).offset(skip).limit(limit).all()
    return {
        "total": db.query(ClientApplication).count(),
        "items": [
            {
                "id": app.id,
                "name": app.name,
                "client_id": app.client_id,
                "description": app.description,
                "is_active": app.is_active,
                "created_at": app.created_at.isoformat()
            }
            for app in apps
        ]
    }


@app.get("/api/applications/{app_id}")
def get_application(app_id: str, db: Session = Depends(get_db)):
    app = db.query(ClientApplication).filter(ClientApplication.id == app_id).first()
    if not app:
        raise HTTPException(status_code=404, detail="Application not found")
    return {
        "id": app.id,
        "name": app.name,
        "client_id": app.client_id,
        "description": app.description,
        "is_active": app.is_active,
        "created_at": app.created_at.isoformat(),
        "consent_count": len(app.consents),
        "active_token_count": len([t for t in app.tokens if t.status == TokenStatus.VALID])
    }


@app.get("/api/consents")
def list_consents(
    user_id: Optional[str] = None,
    status: Optional[AuthorizationStatus] = None,
    application_id: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(UserConsent)
    if user_id:
        query = query.filter(UserConsent.user_id == user_id)
    if status:
        query = query.filter(UserConsent.status == status)
    if application_id:
        query = query.filter(UserConsent.application_id == application_id)

    consents = query.offset(skip).limit(limit).all()
    return {
        "total": query.count(),
        "items": [
            {
                "id": c.id,
                "user_id": c.user_id,
                "application_id": c.application_id,
                "application_name": c.application.name,
                "scope": c.scope,
                "status": c.status,
                "granted_at": c.granted_at.isoformat(),
                "revoked_at": c.revoked_at.isoformat() if c.revoked_at else None,
                "expires_at": c.expires_at.isoformat() if c.expires_at else None
            }
            for c in consents
        ]
    }


@app.get("/api/consents/{consent_id}")
def get_consent(consent_id: str, db: Session = Depends(get_db)):
    consent = db.query(UserConsent).filter(UserConsent.id == consent_id).first()
    if not consent:
        raise HTTPException(status_code=404, detail="Consent not found")
    return {
        "id": consent.id,
        "user_id": consent.user_id,
        "application_id": consent.application_id,
        "application_name": consent.application.name,
        "scope": consent.scope,
        "status": consent.status,
        "granted_at": consent.granted_at.isoformat(),
        "revoked_at": consent.revoked_at.isoformat() if consent.revoked_at else None,
        "expires_at": consent.expires_at.isoformat() if consent.expires_at else None,
        "tokens": [
            {
                "id": t.id,
                "status": t.status,
                "issued_at": t.issued_at.isoformat(),
                "expires_at": t.expires_at.isoformat()
            }
            for t in consent.access_tokens
        ],
        "revocation_events": [
            {
                "id": e.id,
                "status": e.status,
                "initiated_at": e.initiated_at.isoformat()
            }
            for e in consent.revocation_events
        ]
    }


@app.get("/api/tokens")
def list_tokens(
    user_id: Optional[str] = None,
    status: Optional[TokenStatus] = None,
    consent_id: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(AccessToken)
    if user_id:
        query = query.filter(AccessToken.user_id == user_id)
    if status:
        query = query.filter(AccessToken.status == status)
    if consent_id:
        query = query.filter(AccessToken.consent_id == consent_id)

    tokens = query.offset(skip).limit(limit).all()
    return {
        "total": query.count(),
        "items": [
            {
                "id": t.id,
                "token_hash": t.token_hash[:16] + "...",
                "user_id": t.user_id,
                "application_id": t.application_id,
                "application_name": t.application.name,
                "consent_id": t.consent_id,
                "status": t.status,
                "issued_at": t.issued_at.isoformat(),
                "expires_at": t.expires_at.isoformat(),
                "revoked_at": t.revoked_at.isoformat() if t.revoked_at else None
            }
            for t in tokens
        ]
    }


@app.get("/api/revocations")
def list_revocations(
    user_id: Optional[str] = None,
    status: Optional[RevocationStatus] = None,
    consent_id: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(RevocationEvent)
    if user_id:
        query = query.filter(RevocationEvent.user_id == user_id)
    if status:
        query = query.filter(RevocationEvent.status == status)
    if consent_id:
        query = query.filter(RevocationEvent.consent_id == consent_id)

    events = query.order_by(RevocationEvent.initiated_at.desc()).offset(skip).limit(limit).all()
    return {
        "total": query.count(),
        "items": [
            {
                "id": e.id,
                "request_id": e.request_id,
                "consent_id": e.consent_id,
                "user_id": e.user_id,
                "reason": e.reason,
                "status": e.status,
                "initiated_by": e.initiated_by,
                "initiated_at": e.initiated_at.isoformat(),
                "completed_at": e.completed_at.isoformat() if e.completed_at else None,
                "tokens_revoked": e.tokens_revoked,
                "tasks_intercepted": e.tasks_intercepted
            }
            for e in events
        ]
    }


@app.get("/api/revocations/{event_id}")
def get_revocation(event_id: str, db: Session = Depends(get_db)):
    event = db.query(RevocationEvent).filter(RevocationEvent.id == event_id).first()
    if not event:
        raise HTTPException(status_code=404, detail="Revocation event not found")
    return {
        "id": event.id,
        "request_id": event.request_id,
        "consent_id": event.consent_id,
        "user_id": event.user_id,
        "reason": event.reason,
        "status": event.status,
        "initiated_by": event.initiated_by,
        "initiated_at": event.initiated_at.isoformat(),
        "completed_at": event.completed_at.isoformat() if event.completed_at else None,
        "tokens_revoked": event.tokens_revoked,
        "tasks_intercepted": event.tasks_intercepted,
        "error_message": event.error_message
    }


@app.post("/api/revocations")
def initiate_revocation(
    request: RevocationRequest,
    db: Session = Depends(get_db)
):
    start_time = time.time()
    request_id = generate_id()

    try:
        orchestrator = RevocationOrchestrator(db)
        result = orchestrator.initiate_revocation(
            consent_id=request.consent_id,
            user_id=request.user_id,
            reason=request.reason,
            initiated_by=request.initiated_by
        )

        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=result.get("request_id", request_id),
            endpoint="/api/revocations",
            method="POST",
            request_input=request.dict(),
            result=result,
            responsible_node="revocation-service",
            status_code=200,
            duration_ms=duration,
            revocation_event_id=result.get("event_id")
        )

        return result
    except Exception as e:
        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=request_id,
            endpoint="/api/revocations",
            method="POST",
            request_input=request.dict(),
            error=str(e),
            responsible_node="revocation-service",
            status_code=500,
            duration_ms=duration
        )
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/revocations/{event_id}/propagate")
def propagate_revocation(
    event_id: str,
    db: Session = Depends(get_db)
):
    start_time = time.time()
    request_id = generate_id()

    try:
        orchestrator = RevocationOrchestrator(db)
        result = orchestrator.propagate_revocation(event_id)

        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=request_id,
            endpoint=f"/api/revocations/{event_id}/propagate",
            method="POST",
            request_input={"event_id": event_id},
            result=result,
            responsible_node="propagation-service",
            status_code=200,
            duration_ms=duration,
            revocation_event_id=event_id
        )

        return result
    except Exception as e:
        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=request_id,
            endpoint=f"/api/revocations/{event_id}/propagate",
            method="POST",
            request_input={"event_id": event_id},
            error=str(e),
            responsible_node="propagation-service",
            status_code=500,
            duration_ms=duration
        )
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/compensations")
def compensate_revocation(
    request: CompensationRequest,
    db: Session = Depends(get_db)
):
    start_time = time.time()
    request_id = generate_id()

    try:
        result = CompensationService.compensate_revocation(
            db,
            consent_id=request.consent_id,
            reason=request.reason
        )

        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=request_id,
            endpoint="/api/compensations",
            method="POST",
            request_input=request.dict(),
            result=result,
            responsible_node="compensation-service",
            status_code=200,
            duration_ms=duration
        )

        if not result["success"]:
            raise HTTPException(status_code=400, detail=result["message"])
        return result
    except HTTPException:
        raise
    except Exception as e:
        duration = (time.time() - start_time) * 1000
        RequestLogger.log_request(
            db,
            request_id=request_id,
            endpoint="/api/compensations",
            method="POST",
            request_input=request.dict(),
            error=str(e),
            responsible_node="compensation-service",
            status_code=500,
            duration_ms=duration
        )
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/tasks")
def list_tasks(
    user_id: Optional[str] = None,
    status: Optional[TaskStatus] = None,
    consent_id: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(BackgroundTask)
    if user_id:
        query = query.filter(BackgroundTask.user_id == user_id)
    if status:
        query = query.filter(BackgroundTask.status == status)
    if consent_id:
        query = query.filter(BackgroundTask.consent_id == consent_id)

    tasks = query.order_by(BackgroundTask.created_at.desc()).offset(skip).limit(limit).all()
    return {
        "total": query.count(),
        "items": [
            {
                "id": t.id,
                "task_type": t.task_type,
                "user_id": t.user_id,
                "token_id": t.token_id,
                "consent_id": t.consent_id,
                "status": t.status,
                "priority": t.priority,
                "created_at": t.created_at.isoformat(),
                "started_at": t.started_at.isoformat() if t.started_at else None,
                "completed_at": t.completed_at.isoformat() if t.completed_at else None,
                "intercepted_at": t.intercepted_at.isoformat() if t.intercepted_at else None,
                "retry_count": t.retry_count
            }
            for t in tasks
        ]
    }


@app.post("/api/tasks/{task_id}/execute")
def execute_task(
    task_id: str,
    db: Session = Depends(get_db)
):
    task = db.query(BackgroundTask).filter(BackgroundTask.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    if TaskInterceptorService.should_intercept_task(db, task):
        TaskInterceptorService.intercept_task(db, task)
        return {
            "success": False,
            "task_id": task_id,
            "intercepted": True,
            "message": "Task was intercepted due to token revocation"
        }

    if task.status != TaskStatus.PENDING:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot execute task in status: {task.status}"
        )

    task.status = TaskStatus.RUNNING
    task.started_at = datetime.utcnow()
    db.commit()

    task.status = TaskStatus.COMPLETED
    task.completed_at = datetime.utcnow()
    db.commit()

    return {
        "success": True,
        "task_id": task_id,
        "intercepted": False,
        "message": "Task executed successfully"
    }


@app.get("/api/audit-logs")
def list_audit_logs(
    request_id: Optional[str] = None,
    endpoint: Optional[str] = None,
    responsible_node: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(RequestLog)
    if request_id:
        query = query.filter(RequestLog.request_id == request_id)
    if endpoint:
        query = query.filter(RequestLog.endpoint == endpoint)
    if responsible_node:
        query = query.filter(RequestLog.responsible_node == responsible_node)

    logs = query.order_by(RequestLog.created_at.desc()).offset(skip).limit(limit).all()
    return {
        "total": query.count(),
        "items": [
            {
                "id": l.id,
                "request_id": l.request_id,
                "revocation_event_id": l.revocation_event_id,
                "endpoint": l.endpoint,
                "method": l.method,
                "request_input": json.loads(l.request_input),
                "result": json.loads(l.result) if l.result else None,
                "error": l.error,
                "responsible_node": l.responsible_node,
                "status_code": l.status_code,
                "duration_ms": l.duration_ms,
                "created_at": l.created_at.isoformat()
            }
            for l in logs
        ]
    }


@app.get("/api/dashboard/stats")
def get_dashboard_stats(db: Session = Depends(get_db)):
    return {
        "applications": db.query(ClientApplication).count(),
        "active_consents": db.query(UserConsent).filter(
            UserConsent.status == AuthorizationStatus.ACTIVE
        ).count(),
        "revoked_consents": db.query(UserConsent).filter(
            UserConsent.status.in_([
                AuthorizationStatus.REVOKED,
                AuthorizationStatus.PENDING_REVOCATION
            ])
        ).count(),
        "valid_tokens": db.query(AccessToken).filter(
            AccessToken.status == TokenStatus.VALID
        ).count(),
        "revoked_tokens": db.query(AccessToken).filter(
            AccessToken.status == TokenStatus.REVOKED
        ).count(),
        "pending_tasks": db.query(BackgroundTask).filter(
            BackgroundTask.status == TaskStatus.PENDING
        ).count(),
        "intercepted_tasks": db.query(BackgroundTask).filter(
            BackgroundTask.status == TaskStatus.INTERCEPTED
        ).count(),
        "total_revocations": db.query(RevocationEvent).count(),
        "completed_revocations": db.query(RevocationEvent).filter(
            RevocationEvent.status == RevocationStatus.COMPLETED
        ).count()
    }


@app.get("/api/export/revocations")
def export_revocations(
    status: Optional[RevocationStatus] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(RevocationEvent)
    if status:
        query = query.filter(RevocationEvent.status == status)
    if start_date:
        query = query.filter(RevocationEvent.initiated_at >= datetime.fromisoformat(start_date))
    if end_date:
        query = query.filter(RevocationEvent.initiated_at <= datetime.fromisoformat(end_date))

    events = query.order_by(RevocationEvent.initiated_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "撤回记录"

    headers = [
        "事件ID", "请求ID", "用户ID", "授权ID", "撤回原因",
        "状态", "发起者", "发起时间", "完成时间",
        "撤回令牌数", "拦截任务数", "错误信息"
    ]
    ws.append(headers)

    for e in events:
        ws.append([
            e.id,
            e.request_id,
            e.user_id,
            e.consent_id,
            e.reason,
            e.status,
            e.initiated_by,
            e.initiated_at.isoformat(),
            e.completed_at.isoformat() if e.completed_at else "",
            e.tokens_revoked,
            e.tasks_intercepted,
            e.error_message or ""
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=revocation_records.xlsx"}
    )


@app.get("/api/export/impact/{consent_id}")
def export_impact_analysis(consent_id: str, db: Session = Depends(get_db)):
    consent = db.query(UserConsent).filter(UserConsent.id == consent_id).first()
    if not consent:
        raise HTTPException(status_code=404, detail="Consent not found")

    tokens = db.query(AccessToken).filter(AccessToken.consent_id == consent_id).all()
    tasks = db.query(BackgroundTask).filter(BackgroundTask.consent_id == consent_id).all()
    revocations = db.query(RevocationEvent).filter(RevocationEvent.consent_id == consent_id).all()

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "授权信息"
    ws1.append(["授权ID", consent.id])
    ws1.append(["用户ID", consent.user_id])
    ws1.append(["应用名称", consent.application.name])
    ws1.append(["授权范围", consent.scope])
    ws1.append(["当前状态", consent.status])
    ws1.append(["授权时间", consent.granted_at.isoformat()])
    ws1.append(["过期时间", consent.expires_at.isoformat() if consent.expires_at else ""])

    ws2 = wb.create_sheet("关联令牌")
    ws2.append(["令牌ID", "状态", "签发时间", "过期时间", "撤回时间"])
    for t in tokens:
        ws2.append([
            t.id, t.status,
            t.issued_at.isoformat(),
            t.expires_at.isoformat(),
            t.revoked_at.isoformat() if t.revoked_at else ""
        ])

    ws3 = wb.create_sheet("关联任务")
    ws3.append(["任务ID", "任务类型", "状态", "创建时间", "拦截时间"])
    for t in tasks:
        ws3.append([
            t.id, t.task_type, t.status,
            t.created_at.isoformat(),
            t.intercepted_at.isoformat() if t.intercepted_at else ""
        ])

    ws4 = wb.create_sheet("撤回历史")
    ws4.append(["事件ID", "状态", "发起时间", "完成时间", "撤回令牌数", "拦截任务数"])
    for r in revocations:
        ws4.append([
            r.id, r.status,
            r.initiated_at.isoformat(),
            r.completed_at.isoformat() if r.completed_at else "",
            r.tokens_revoked, r.tasks_intercepted
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=impact_analysis_{consent_id[:8]}.xlsx"}
    )


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
