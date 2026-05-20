import json
import csv
from io import StringIO, BytesIO
from typing import Optional, Dict, Any, List
from datetime import datetime
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from app.models.models import UserPreference, ChangeHistory, SendInterception, AnomalyQueue
from app.schemas.preference import ExportRequest


class ExportService:
    def __init__(self, db: Session):
        self.db = db

    def export_preferences(
        self, request: ExportRequest
    ) -> tuple[bytes, str, str]:
        query = self.db.query(UserPreference)
        
        if request.user_id:
            query = query.filter(UserPreference.user_id == request.user_id)
        if request.channel:
            query = query.filter(UserPreference.channel == request.channel)
        if request.business_scene:
            query = query.filter(UserPreference.business_scene == request.business_scene)
        if request.start_date:
            query = query.filter(UserPreference.created_at >= request.start_date)
        if request.end_date:
            query = query.filter(UserPreference.created_at <= request.end_date)

        preferences = query.all()
        
        data = [
            {
                "id": p.id,
                "user_id": p.user_id,
                "channel": p.channel.value,
                "business_scene": p.business_scene.value,
                "enabled": p.enabled,
                "source": p.source.value,
                "source_priority": p.source_priority,
                "status": p.status.value,
                "created_at": p.created_at.isoformat() if p.created_at else None,
                "updated_at": p.updated_at.isoformat() if p.updated_at else None,
                "expires_at": p.expires_at.isoformat() if p.expires_at else None,
            }
            for p in preferences
        ]

        filename = f"preferences_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if request.export_format == "json":
            return self._to_json(data), filename + ".json", "application/json"
        elif request.export_format == "csv":
            return self._to_csv(data), filename + ".csv", "text/csv"
        elif request.export_format == "excel":
            return self._to_excel(data, "preferences"), filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            return self._to_json(data), filename + ".json", "application/json"

    def export_change_history(
        self, request: ExportRequest
    ) -> tuple[bytes, str, str]:
        query = self.db.query(ChangeHistory)
        
        if request.user_id:
            query = query.filter(ChangeHistory.user_id == request.user_id)
        if request.start_date:
            query = query.filter(ChangeHistory.created_at >= request.start_date)
        if request.end_date:
            query = query.filter(ChangeHistory.created_at <= request.end_date)

        histories = query.all()
        
        data = [
            {
                "id": h.id,
                "preference_id": h.preference_id,
                "user_id": h.user_id,
                "channel": h.channel.value if h.channel else None,
                "business_scene": h.business_scene.value if h.business_scene else None,
                "source": h.source.value if h.source else None,
                "operator": h.operator,
                "change_type": h.change_type,
                "created_at": h.created_at.isoformat() if h.created_at else None,
            }
            for h in histories
        ]

        filename = f"change_history_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if request.export_format == "json":
            return self._to_json(data), filename + ".json", "application/json"
        elif request.export_format == "csv":
            return self._to_csv(data), filename + ".csv", "text/csv"
        elif request.export_format == "excel":
            return self._to_excel(data, "change_history"), filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            return self._to_json(data), filename + ".json", "application/json"

    def export_interceptions(
        self, request: ExportRequest
    ) -> tuple[bytes, str, str]:
        query = self.db.query(SendInterception)
        
        if request.user_id:
            query = query.filter(SendInterception.user_id == request.user_id)
        if request.channel:
            query = query.filter(SendInterception.channel == request.channel)
        if request.business_scene:
            query = query.filter(SendInterception.business_scene == request.business_scene)
        if request.start_date:
            query = query.filter(SendInterception.created_at >= request.start_date)
        if request.end_date:
            query = query.filter(SendInterception.created_at <= request.end_date)

        interceptions = query.all()
        
        data = [
            {
                "id": i.id,
                "preference_id": i.preference_id,
                "user_id": i.user_id,
                "channel": i.channel.value,
                "business_scene": i.business_scene.value,
                "status": i.status.value,
                "reason": i.reason,
                "interception_rule": i.interception_rule,
                "message_id": i.message_id,
                "created_at": i.created_at.isoformat() if i.created_at else None,
            }
            for i in interceptions
        ]

        filename = f"interceptions_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if request.export_format == "json":
            return self._to_json(data), filename + ".json", "application/json"
        elif request.export_format == "csv":
            return self._to_csv(data), filename + ".csv", "text/csv"
        elif request.export_format == "excel":
            return self._to_excel(data, "interceptions"), filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            return self._to_json(data), filename + ".json", "application/json"

    def export_anomaly_queue(
        self, request: ExportRequest
    ) -> tuple[bytes, str, str]:
        query = self.db.query(AnomalyQueue)
        
        if request.user_id:
            query = query.filter(AnomalyQueue.user_id == request.user_id)
        if request.channel:
            query = query.filter(AnomalyQueue.channel == request.channel)
        if request.business_scene:
            query = query.filter(AnomalyQueue.business_scene == request.business_scene)
        if request.start_date:
            query = query.filter(AnomalyQueue.created_at >= request.start_date)
        if request.end_date:
            query = query.filter(AnomalyQueue.created_at <= request.end_date)

        anomalies = query.all()
        
        data = [
            {
                "id": a.id,
                "user_id": a.user_id,
                "channel": a.channel.value,
                "business_scene": a.business_scene.value,
                "anomaly_type": a.anomaly_type,
                "description": a.description,
                "source": a.source.value,
                "status": a.status,
                "retry_count": a.retry_count,
                "max_retries": a.max_retries,
                "resolver": a.resolver,
                "resolution_note": a.resolution_note,
                "created_at": a.created_at.isoformat() if a.created_at else None,
                "resolved_at": a.resolved_at.isoformat() if a.resolved_at else None,
            }
            for a in anomalies
        ]

        filename = f"anomaly_queue_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if request.export_format == "json":
            return self._to_json(data), filename + ".json", "application/json"
        elif request.export_format == "csv":
            return self._to_csv(data), filename + ".csv", "text/csv"
        elif request.export_format == "excel":
            return self._to_excel(data, "anomaly_queue"), filename + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            return self._to_json(data), filename + ".json", "application/json"

    def _to_json(self, data: List[Dict[str, Any]]) -> bytes:
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")

    def _to_csv(self, data: List[Dict[str, Any]]) -> bytes:
        if not data:
            return b""
        
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)
        
        return output.getvalue().encode("utf-8-sig")

    def _to_excel(self, data: List[Dict[str, Any]], sheet_name: str) -> bytes:
        output = BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        if data:
            headers = list(data[0].keys())
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for row_idx, row_data in enumerate(data, 2):
                for col_idx, key in enumerate(headers, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(key))
                    cell.alignment = Alignment(wrap_text=True)

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width

        wb.save(output)
        return output.getvalue()
