import csv
import json
import os
from typing import List, Dict, Any, Tuple
from .weight_curve import WeightCurveManager
from .models import ProcessingStatus


class DataLoader:
    def __init__(self, curve_manager: WeightCurveManager):
        self.curve_manager = curve_manager

    @staticmethod
    def _read_csv(path: str) -> List[Dict]:
        rows = []
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for r in reader:
                rows.append(dict(r))
        return rows

    @staticmethod
    def _read_json(path: str) -> List[Dict]:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
        if isinstance(data, dict) and "records" in data:
            return data["records"]
        if isinstance(data, dict) and "data" in data:
            return data["data"]
        return [data]

    @staticmethod
    def _read_xlsx(path: str) -> List[Dict]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            alt = path.replace(".xlsx", ".csv").replace(".xls", ".csv")
            if os.path.exists(alt):
                return DataLoader._read_csv(alt)
            return []
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {}
            for i, h in enumerate(headers):
                row[h] = r[i] if i < len(r) else None
            rows.append(row)
        return rows

    def load_file(self, path: str) -> Tuple[int, int]:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            rows = self._read_csv(path)
        elif ext in (".json",):
            rows = self._read_json(path)
        elif ext in (".xlsx", ".xls"):
            rows = self._read_xlsx(path)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")
        ok = 0
        total = len(rows)
        source = os.path.basename(path)
        for i, row in enumerate(rows, start=1):
            try:
                rec, extra = self.curve_manager.import_row(row, source_file=source, row_num=i)
                self.curve_manager.add_record(rec)
                ok += 1
            except Exception as e:
                continue
        return ok, total

    def load_directory(self, dir_path: str) -> Dict[str, Tuple[int, int]]:
        results = {}
        for fn in sorted(os.listdir(dir_path)):
            fp = os.path.join(dir_path, fn)
            if not os.path.isfile(fp):
                continue
            ext = os.path.splitext(fn)[1].lower()
            if ext not in (".csv", ".json", ".xlsx", ".xls"):
                continue
            results[fn] = self.load_file(fp)
        return results
