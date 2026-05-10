from sqlalchemy.orm import Session
from typing import Optional, List, Dict, Any
from datetime import datetime
import csv
import io
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from ..models.models import (
    ExportRecord,
    TermRule,
    RuleStatus,
    RuleType,
    AuditLog,
    ReviewRequest,
    GrayRelease,
    GrayEffectCheck
)
from ..schemas.schemas import ExportRequest
from .audit_service import AuditService
import json


class ExportService:
    def __init__(self, export_dir: str = "./exports"):
        self.export_dir = export_dir
        os.makedirs(export_dir, exist_ok=True)

    def create_export(self, db: Session, request: ExportRequest) -> ExportRecord:
        export_record = ExportRecord(
            export_type=request.export_type,
            export_params=json.dumps({
                "rule_ids": request.rule_ids,
                "start_date": request.start_date.isoformat() if request.start_date else None,
                "end_date": request.end_date.isoformat() if request.end_date else None,
                "statuses": [s.value for s in request.statuses] if request.statuses else None
            }, ensure_ascii=False),
            status="processing",
            created_by=request.created_by,
            created_at=datetime.utcnow()
        )
        db.add(export_record)
        db.flush()

        try:
            if request.export_type == "rules":
                file_path, file_name, record_count = self._export_rules(
                    db,
                    request.rule_ids,
                    request.statuses,
                    request.start_date,
                    request.end_date
                )
            elif request.export_type == "history":
                file_path, file_name, record_count = self._export_history(
                    db,
                    request.rule_ids,
                    request.start_date,
                    request.end_date
                )
            elif request.export_type == "full_report":
                file_path, file_name, record_count = self._export_full_report(
                    db,
                    request.rule_ids,
                    request.start_date,
                    request.end_date
                )
            else:
                raise ValueError(f"不支持的导出类型: {request.export_type}")

            export_record.file_path = file_path
            export_record.file_name = file_name
            export_record.record_count = record_count
            export_record.status = "completed"
            export_record.completed_at = datetime.utcnow()

            if request.rule_ids and len(request.rule_ids) == 1:
                AuditService.log_action(
                    db=db,
                    rule_id=request.rule_ids[0],
                    action="EXPORT",
                    actor=request.created_by,
                    reason=f"导出 {request.export_type} 数据",
                    details={"export_id": export_record.id}
                )

        except Exception as e:
            export_record.status = "failed"
            db.flush()
            raise e

        return export_record

    def _export_rules(
        self,
        db: Session,
        rule_ids: Optional[List[int]],
        statuses: Optional[List[RuleStatus]],
        start_date: Optional[datetime],
        end_date: Optional[datetime]
    ) -> tuple:
        query = db.query(TermRule)
        
        if rule_ids:
            query = query.filter(TermRule.id.in_(rule_ids))
        if statuses:
            query = query.filter(TermRule.status.in_(statuses))
        if start_date:
            query = query.filter(TermRule.updated_at >= start_date)
        if end_date:
            query = query.filter(TermRule.updated_at <= end_date)

        rules = query.order_by(TermRule.updated_at.desc()).all()

        file_name = f"规则导出_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(self.export_dir, file_name)

        wb = Workbook()
        ws = wb.active
        ws.title = "规则列表"

        headers = [
            "规则ID", "搜索词", "规则类型", "匹配方式", "优先级",
            "动作", "状态", "词库ID", "创建人", "创建时间",
            "最后更新", "原因"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        status_display = {
            RuleStatus.DRAFT: "草稿",
            RuleStatus.PENDING_REVIEW: "待审核",
            RuleStatus.REVIEW_REJECTED: "审核驳回",
            RuleStatus.PENDING_GRAY: "待灰度",
            RuleStatus.IN_GRAY: "灰度中",
            RuleStatus.GRAY_REJECTED: "灰度不通过",
            RuleStatus.PRODUCTION: "生产生效",
            RuleStatus.ROLLED_BACK: "已回滚",
            RuleStatus.DEPRECATED: "已废弃"
        }
        
        type_display = {
            RuleType.BLACKLIST: "黑名单",
            RuleType.WHITELIST: "白名单"
        }

        for row_idx, rule in enumerate(rules, 2):
            ws.cell(row=row_idx, column=1, value=rule.id)
            ws.cell(row=row_idx, column=2, value=rule.term)
            ws.cell(row=row_idx, column=3, value=type_display.get(rule.rule_type, rule.rule_type.value))
            ws.cell(row=row_idx, column=4, value=rule.match_type)
            ws.cell(row=row_idx, column=5, value=rule.priority)
            ws.cell(row=row_idx, column=6, value=rule.action)
            ws.cell(row=row_idx, column=7, value=status_display.get(rule.status, rule.status.value))
            ws.cell(row=row_idx, column=8, value=rule.library_id)
            ws.cell(row=row_idx, column=9, value=rule.created_by)
            ws.cell(row=row_idx, column=10, value=rule.created_at.strftime("%Y-%m-%d %H:%M:%S") if rule.created_at else "")
            ws.cell(row=row_idx, column=11, value=rule.updated_at.strftime("%Y-%m-%d %H:%M:%S") if rule.updated_at else "")
            ws.cell(row=row_idx, column=12, value=rule.reason or "")

        for col in range(1, 13):
            ws.column_dimensions[chr(64 + col)].width = 20

        wb.save(file_path)
        return file_path, file_name, len(rules)

    def _export_history(
        self,
        db: Session,
        rule_ids: Optional[List[int]],
        start_date: Optional[datetime],
        end_date: Optional[datetime]
    ) -> tuple:
        query = db.query(AuditLog)
        
        if rule_ids:
            query = query.filter(AuditLog.rule_id.in_(rule_ids))
        if start_date:
            query = query.filter(AuditLog.timestamp >= start_date)
        if end_date:
            query = query.filter(AuditLog.timestamp <= end_date)

        logs = query.order_by(AuditLog.timestamp.desc()).all()

        file_name = f"历史记录导出_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(self.export_dir, file_name)

        wb = Workbook()
        ws = wb.active
        ws.title = "操作历史"

        headers = [
            "日志ID", "规则ID", "操作类型", "原状态", "新状态",
            "操作人", "操作时间", "原因", "详情摘要"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        action_display = {
            "CREATE": "创建规则",
            "UPDATE": "更新规则",
            "SUBMIT_REVIEW": "提交审核",
            "APPROVE_REVIEW": "审核通过",
            "REJECT_REVIEW": "审核驳回",
            "START_GRAY": "开始灰度",
            "APPROVE_GRAY": "灰度通过",
            "REJECT_GRAY": "灰度不通过",
            "ROLLBACK": "回滚",
            "DEPRECATE": "废弃",
            "EXPORT": "导出"
        }
        
        status_display = {
            RuleStatus.DRAFT: "草稿",
            RuleStatus.PENDING_REVIEW: "待审核",
            RuleStatus.REVIEW_REJECTED: "审核驳回",
            RuleStatus.PENDING_GRAY: "待灰度",
            RuleStatus.IN_GRAY: "灰度中",
            RuleStatus.GRAY_REJECTED: "灰度不通过",
            RuleStatus.PRODUCTION: "生产生效",
            RuleStatus.ROLLED_BACK: "已回滚",
            RuleStatus.DEPRECATED: "已废弃"
        }

        for row_idx, log in enumerate(logs, 2):
            ws.cell(row=row_idx, column=1, value=log.id)
            ws.cell(row=row_idx, column=2, value=log.rule_id)
            ws.cell(row=row_idx, column=3, value=action_display.get(log.action, log.action))
            ws.cell(row=row_idx, column=4, value=status_display.get(log.from_status, log.from_status.value if log.from_status else ""))
            ws.cell(row=row_idx, column=5, value=status_display.get(log.to_status, log.to_status.value if log.to_status else ""))
            ws.cell(row=row_idx, column=6, value=log.actor)
            ws.cell(row=row_idx, column=7, value=log.timestamp.strftime("%Y-%m-%d %H:%M:%S") if log.timestamp else "")
            ws.cell(row=row_idx, column=8, value=log.reason or "")
            
            details_summmary = ""
            if log.details:
                try:
                    details = json.loads(log.details)
                    details_summmary = self._summarize_details(details)
                except:
                    details_summmary = str(log.details)[:100]
            ws.cell(row=row_idx, column=9, value=details_summmary)

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 20

        wb.save(file_path)
        return file_path, file_name, len(logs)

    def _export_full_report(
        self,
        db: Session,
        rule_ids: Optional[List[int]],
        start_date: Optional[datetime],
        end_date: Optional[datetime]
    ) -> tuple:
        file_name = f"完整报告_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path = os.path.join(self.export_dir, file_name)

        wb = Workbook()

        ws_summary = wb.active
        ws_summary.title = "汇总概览"
        
        summary_headers = ["统计项", "数量/说明"]
        for col, header in enumerate(summary_headers, 1):
            cell = ws_summary.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")

        total_rules = db.query(TermRule).count()
        production_rules = db.query(TermRule).filter(TermRule.status == RuleStatus.PRODUCTION).count()
        pending_review = db.query(TermRule).filter(TermRule.status == RuleStatus.PENDING_REVIEW).count()
        in_gray = db.query(TermRule).filter(TermRule.status == RuleStatus.IN_GRAY).count()
        blacklist = db.query(TermRule).filter(TermRule.rule_type == RuleType.BLACKLIST).count()
        whitelist = db.query(TermRule).filter(TermRule.rule_type == RuleType.WHITELIST).count()

        summary_data = [
            ("规则总数", total_rules),
            ("生产生效规则", production_rules),
            ("待审核规则", pending_review),
            ("灰度中规则", in_gray),
            ("黑名单规则", blacklist),
            ("白名单规则", whitelist),
            ("导出时间", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"))
        ]

        for row_idx, (key, value) in enumerate(summary_data, 2):
            ws_summary.cell(row=row_idx, column=1, value=key)
            ws_summary.cell(row=row_idx, column=2, value=value)

        ws_summary.column_dimensions['A'].width = 20
        ws_summary.column_dimensions['B'].width = 30

        self._add_rules_sheet(wb, db, rule_ids, start_date, end_date)
        self._add_history_sheet(wb, db, rule_ids, start_date, end_date)
        self._add_review_sheet(wb, db, rule_ids)
        self._add_gray_sheet(wb, db, rule_ids)

        wb.save(file_path)
        
        total_records = (
            db.query(TermRule).count() +
            db.query(AuditLog).count() +
            db.query(ReviewRequest).count() +
            db.query(GrayRelease).count()
        )
        
        return file_path, file_name, total_records

    def _add_rules_sheet(self, wb: Workbook, db: Session, rule_ids: Optional[List[int]], start_date, end_date):
        ws = wb.create_sheet("规则详情")
        
        headers = [
            "规则ID", "搜索词", "类型", "匹配方式", "优先级",
            "动作", "状态", "词库ID", "创建人", "创建时间",
            "最后更新", "原因"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        query = db.query(TermRule)
        if rule_ids:
            query = query.filter(TermRule.id.in_(rule_ids))
        rules = query.order_by(TermRule.updated_at.desc()).all()

        type_display = {"blacklist": "黑名单", "whitelist": "白名单"}
        status_display = {
            "draft": "草稿", "pending_review": "待审核", "review_rejected": "审核驳回",
            "pending_gray": "待灰度", "in_gray": "灰度中", "gray_rejected": "灰度不通过",
            "production": "生产生效", "rolled_back": "已回滚", "deprecated": "已废弃"
        }

        for row_idx, rule in enumerate(rules, 2):
            ws.cell(row=row_idx, column=1, value=rule.id)
            ws.cell(row=row_idx, column=2, value=rule.term)
            ws.cell(row=row_idx, column=3, value=type_display.get(rule.rule_type.value, rule.rule_type.value))
            ws.cell(row=row_idx, column=4, value=rule.match_type)
            ws.cell(row=row_idx, column=5, value=rule.priority)
            ws.cell(row=row_idx, column=6, value=rule.action)
            ws.cell(row=row_idx, column=7, value=status_display.get(rule.status.value, rule.status.value))
            ws.cell(row=row_idx, column=8, value=rule.library_id)
            ws.cell(row=row_idx, column=9, value=rule.created_by)
            ws.cell(row=row_idx, column=10, value=rule.created_at.strftime("%Y-%m-%d %H:%M:%S") if rule.created_at else "")
            ws.cell(row=row_idx, column=11, value=rule.updated_at.strftime("%Y-%m-%d %H:%M:%S") if rule.updated_at else "")
            ws.cell(row=row_idx, column=12, value=rule.reason or "")

        for col in range(1, 13):
            ws.column_dimensions[chr(64 + col)].width = 18

    def _add_history_sheet(self, wb: Workbook, db: Session, rule_ids: Optional[List[int]], start_date, end_date):
        ws = wb.create_sheet("操作历史")
        
        headers = [
            "日志ID", "规则ID", "操作", "原状态", "新状态",
            "操作人", "时间", "原因", "详情"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

        query = db.query(AuditLog)
        if rule_ids:
            query = query.filter(AuditLog.rule_id.in_(rule_ids))
        logs = query.order_by(AuditLog.timestamp.desc()).all()

        action_display = {
            "CREATE": "创建", "UPDATE": "更新", "SUBMIT_REVIEW": "提交审核",
            "APPROVE_REVIEW": "审核通过", "REJECT_REVIEW": "审核驳回",
            "START_GRAY": "开始灰度", "APPROVE_GRAY": "灰度通过",
            "REJECT_GRAY": "灰度驳回", "ROLLBACK": "回滚",
            "DEPRECATE": "废弃", "EXPORT": "导出"
        }
        status_display = {
            "draft": "草稿", "pending_review": "待审核", "review_rejected": "审核驳回",
            "pending_gray": "待灰度", "in_gray": "灰度中", "gray_rejected": "灰度不通过",
            "production": "生产生效", "rolled_back": "已回滚", "deprecated": "已废弃"
        }

        for row_idx, log in enumerate(logs, 2):
            ws.cell(row=row_idx, column=1, value=log.id)
            ws.cell(row=row_idx, column=2, value=log.rule_id)
            ws.cell(row=row_idx, column=3, value=action_display.get(log.action, log.action))
            ws.cell(row=row_idx, column=4, value=status_display.get(log.from_status.value, log.from_status.value) if log.from_status else "")
            ws.cell(row=row_idx, column=5, value=status_display.get(log.to_status.value, log.to_status.value) if log.to_status else "")
            ws.cell(row=row_idx, column=6, value=log.actor)
            ws.cell(row=row_idx, column=7, value=log.timestamp.strftime("%Y-%m-%d %H:%M:%S") if log.timestamp else "")
            ws.cell(row=row_idx, column=8, value=log.reason or "")
            ws.cell(row=row_idx, column=9, value=log.details or "")

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 18

    def _add_review_sheet(self, wb: Workbook, db: Session, rule_ids: Optional[List[int]]):
        ws = wb.create_sheet("审核记录")
        
        headers = [
            "审核ID", "规则ID", "状态", "提交人", "审核人",
            "提交时间", "审核时间", "提交备注", "审核意见"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

        query = db.query(ReviewRequest)
        if rule_ids:
            query = query.filter(ReviewRequest.rule_id.in_(rule_ids))
        reviews = query.order_by(ReviewRequest.requested_at.desc()).all()

        status_display = {"pending": "待审核", "approved": "已通过", "rejected": "已驳回"}

        for row_idx, review in enumerate(reviews, 2):
            ws.cell(row=row_idx, column=1, value=review.id)
            ws.cell(row=row_idx, column=2, value=review.rule_id)
            ws.cell(row=row_idx, column=3, value=status_display.get(review.status.value, review.status.value))
            ws.cell(row=row_idx, column=4, value=review.requested_by)
            ws.cell(row=row_idx, column=5, value=review.reviewer or "")
            ws.cell(row=row_idx, column=6, value=review.requested_at.strftime("%Y-%m-%d %H:%M:%S") if review.requested_at else "")
            ws.cell(row=row_idx, column=7, value=review.reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if review.reviewed_at else "")
            ws.cell(row=row_idx, column=8, value=review.comments or "")
            ws.cell(row=row_idx, column=9, value=review.review_comment or "")

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 18

    def _add_gray_sheet(self, wb: Workbook, db: Session, rule_ids: Optional[List[int]]):
        ws = wb.create_sheet("灰度发布")
        
        headers = [
            "灰度ID", "规则ID", "流量比例", "是否活跃", "效果状态",
            "创建人", "开始时间", "结束时间", "效果说明"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")

        query = db.query(GrayRelease)
        if rule_ids:
            query = query.filter(GrayRelease.rule_id.in_(rule_ids))
        grays = query.order_by(GrayRelease.created_at.desc()).all()

        effect_display = {"pending": "待验证", "passed": "通过", "failed": "未通过"}

        for row_idx, gray in enumerate(grays, 2):
            ws.cell(row=row_idx, column=1, value=gray.id)
            ws.cell(row=row_idx, column=2, value=gray.rule_id)
            ws.cell(row=row_idx, column=3, value=f"{gray.traffic_percentage}%")
            ws.cell(row=row_idx, column=4, value="是" if gray.is_active else "否")
            ws.cell(row=row_idx, column=5, value=effect_display.get(gray.effect_status.value, gray.effect_status.value))
            ws.cell(row=row_idx, column=6, value=gray.created_by)
            ws.cell(row=row_idx, column=7, value=gray.start_time.strftime("%Y-%m-%d %H:%M:%S") if gray.start_time else "")
            ws.cell(row=row_idx, column=8, value=gray.end_time.strftime("%Y-%m-%d %H:%M:%S") if gray.end_time else "")
            ws.cell(row=row_idx, column=9, value=gray.effect_comment or "")

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 18

    def _summarize_details(self, details: Dict[str, Any]) -> str:
        summaries = []
        
        if "review_id" in details:
            summaries.append(f"审核ID: {details['review_id']}")
        if "gray_release_id" in details:
            summaries.append(f"灰度ID: {details['gray_release_id']}")
        if "traffic_percentage" in details:
            summaries.append(f"流量: {details['traffic_percentage']}%")
        if "transition_description" in details:
            summaries.append(details["transition_description"])
        if "check_count" in details:
            summaries.append(f"回查数: {details.get('check_count', 0)}")
        if "review_comment" in details:
            summaries.append(f"备注: {details['review_comment'][:50]}")
            
        return "; ".join(summaries)

    def get_export_record(self, db: Session, export_id: int) -> Optional[ExportRecord]:
        return db.query(ExportRecord).filter(ExportRecord.id == export_id).first()

    def list_export_records(self, db: Session, limit: int = 100) -> List[ExportRecord]:
        return (
            db.query(ExportRecord)
            .order_by(ExportRecord.created_at.desc())
            .limit(limit)
            .all()
        )
