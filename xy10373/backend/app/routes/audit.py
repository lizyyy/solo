from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import and_, or_
from datetime import date, datetime, timedelta
from typing import List, Optional
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from ..database import get_db
from ..models import Patient, Caregiver, CareCertificate, ReplacementRequest, OperationLog, WardRule
from ..schemas import OperationLogResponse
from ..utils.validators import calculate_days_remaining, update_expired_certificates

router = APIRouter(prefix="/api/audit", tags=["audit"])


@router.get("/logs", response_model=List[OperationLogResponse])
def get_operation_logs(
    operation_type: Optional[str] = None,
    patient_id: Optional[str] = None,
    certificate_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    result: Optional[str] = None,
    skip: int = Query(0, ge=0),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    query = db.query(OperationLog)
    
    if operation_type:
        query = query.filter(OperationLog.operation_type == operation_type)
    
    if patient_id:
        query = query.filter(OperationLog.patient_id == patient_id)
    
    if certificate_id:
        query = query.filter(OperationLog.certificate_id == certificate_id)
    
    if start_date:
        query = query.filter(OperationLog.created_at >= start_date)
    
    if end_date:
        end_datetime = datetime.combine(end_date, datetime.max.time())
        query = query.filter(OperationLog.created_at <= end_datetime)
    
    if result:
        query = query.filter(OperationLog.result == result)
    
    logs = query.order_by(OperationLog.created_at.desc()).offset(skip).limit(limit).all()
    
    results = []
    for log in logs:
        patient = None
        certificate = None
        
        if log.patient_id:
            patient = db.query(Patient).filter(Patient.patient_id == log.patient_id).first()
        
        if log.certificate_id:
            certificate = db.query(CareCertificate).filter(CareCertificate.id == log.certificate_id).first()
        
        results.append({
            "id": log.id,
            "operation_type": log.operation_type,
            "patient_id": log.patient_id,
            "patient_name": patient.name if patient else None,
            "certificate_id": log.certificate_id,
            "certificate_no": certificate.certificate_no if certificate else None,
            "operator": log.operator,
            "action": log.action,
            "old_value": log.old_value,
            "new_value": log.new_value,
            "source_file": log.source_file,
            "result": log.result,
            "notes": log.notes,
            "created_at": log.created_at
        })
    
    return results


@router.get("/export")
def export_audit_report(
    ward: Optional[str] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db)
):
    update_expired_certificates(db)
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "陪护证汇总"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    headers1 = ["证件号", "患者姓名", "患者ID", "病区", "床号", "陪护人姓名", "身份证号", 
                "与患者关系", "发证日期", "有效期至", "剩余天数", "状态", "来源文件", "备注"]
    
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    cert_query = db.query(CareCertificate)
    if ward:
        cert_query = cert_query.join(Patient, CareCertificate.patient_id == Patient.patient_id).filter(Patient.ward == ward)
    
    certificates = cert_query.order_by(CareCertificate.created_at.desc()).all()
    
    status_map = {
        "active": "有效",
        "expired": "已过期",
        "expiring_soon": "即将过期",
        "cancelled": "已注销",
        "replaced": "已换人"
    }
    
    for row_idx, cert in enumerate(certificates, 2):
        patient = db.query(Patient).filter(Patient.patient_id == cert.patient_id).first()
        caregiver = db.query(Caregiver).filter(Caregiver.caregiver_id == cert.caregiver_id).first()
        
        days_remaining = calculate_days_remaining(cert.expiry_date)
        status = cert.status
        if cert.status == "active":
            if days_remaining < 0:
                status = "expired"
            elif days_remaining <= 3:
                status = "expiring_soon"
        
        row_data = [
            cert.certificate_no,
            patient.name if patient else "",
            cert.patient_id,
            patient.ward if patient else "",
            patient.bed_no if patient else "",
            caregiver.name if caregiver else "",
            caregiver.id_card if caregiver else "",
            caregiver.relation_to_patient if caregiver else "",
            cert.issue_date.strftime("%Y-%m-%d") if cert.issue_date else "",
            cert.expiry_date.strftime("%Y-%m-%d") if cert.expiry_date else "",
            days_remaining if days_remaining >= 0 else 0,
            status_map.get(status, status),
            cert.source_file or "",
            cert.notes or ""
        ]
        
        for col, value in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = center_align
            
            if status == "expired":
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif status == "expiring_soon":
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    for col in range(1, len(headers1) + 1):
        ws1.column_dimensions[ws1.cell(row=1, column=col).column_letter].width = 18
    
    ws2 = wb.create_sheet("换人申请")
    headers2 = ["申请单号", "患者姓名", "病区", "原陪护人", "新陪护人", "关系", "申请原因", 
                "申请人", "申请时间", "审批状态", "审批人", "审批时间", "审批备注"]
    
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    repl_query = db.query(ReplacementRequest)
    if ward:
        repl_query = repl_query.join(Patient, ReplacementRequest.patient_id == Patient.patient_id).filter(Patient.ward == ward)
    if start_date:
        repl_query = repl_query.filter(ReplacementRequest.request_date >= start_date)
    if end_date:
        repl_query = repl_query.filter(ReplacementRequest.request_date <= end_date)
    
    replacements = repl_query.order_by(ReplacementRequest.created_at.desc()).all()
    
    status_map2 = {
        "pending": "待审批",
        "approved": "已通过",
        "rejected": "已驳回"
    }
    
    for row_idx, req in enumerate(replacements, 2):
        patient = db.query(Patient).filter(Patient.patient_id == req.patient_id).first()
        old_caregiver = db.query(Caregiver).filter(Caregiver.caregiver_id == req.old_caregiver_id).first()
        
        row_data = [
            req.request_no,
            patient.name if patient else "",
            patient.ward if patient else "",
            old_caregiver.name if old_caregiver else "",
            req.new_caregiver_name,
            req.new_caregiver_relation,
            req.reason,
            req.requested_by,
            req.request_date.strftime("%Y-%m-%d %H:%M:%S") if req.request_date else "",
            status_map2.get(req.status, req.status),
            req.approved_by or "",
            req.approval_date.strftime("%Y-%m-%d %H:%M:%S") if req.approval_date else "",
            req.approval_notes or ""
        ]
        
        for col, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = center_align
    
    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = 18
    
    ws3 = wb.create_sheet("操作日志")
    headers3 = ["操作时间", "操作类型", "操作员", "操作内容", "患者", "证件号", 
                "旧值", "新值", "来源文件", "结果", "备注"]
    
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    log_query = db.query(OperationLog)
    if ward:
        log_query = log_query.join(Patient, OperationLog.patient_id == Patient.patient_id).filter(Patient.ward == ward)
    if start_date:
        log_query = log_query.filter(OperationLog.created_at >= start_date)
    if end_date:
        end_datetime = datetime.combine(end_date, datetime.max.time())
        log_query = log_query.filter(OperationLog.created_at <= end_datetime)
    
    logs = log_query.order_by(OperationLog.created_at.desc()).all()
    
    op_type_map = {
        "PATIENT_REGISTER": "患者登记",
        "PATIENT_UPDATE": "患者更新",
        "CERTIFICATE_ISSUE": "证件办理",
        "CERTIFICATE_RENEW": "证件续期",
        "CERTIFICATE_CANCEL": "证件注销",
        "REPLACEMENT_REQUEST": "换人申请",
        "REPLACEMENT_APPROVE": "换人审批通过",
        "REPLACEMENT_REJECT": "换人审批驳回"
    }
    
    result_map = {
        "success": "成功",
        "failed": "失败"
    }
    
    for row_idx, log in enumerate(logs, 2):
        patient = None
        certificate = None
        
        if log.patient_id:
            patient = db.query(Patient).filter(Patient.patient_id == log.patient_id).first()
        if log.certificate_id:
            certificate = db.query(CareCertificate).filter(CareCertificate.id == log.certificate_id).first()
        
        row_data = [
            log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "",
            op_type_map.get(log.operation_type, log.operation_type),
            log.operator,
            log.action,
            f"{patient.name}({log.patient_id})" if patient else (log.patient_id or ""),
            certificate.certificate_no if certificate else "",
            log.old_value or "",
            log.new_value or "",
            log.source_file or "",
            result_map.get(log.result, log.result),
            log.notes or ""
        ]
        
        for col, value in enumerate(row_data, 1):
            cell = ws3.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = center_align
            
            if log.result == "failed":
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    for col in range(1, len(headers3) + 1):
        ws3.column_dimensions[ws3.cell(row=1, column=col).column_letter].width = 20
    
    ws4 = wb.create_sheet("过期提醒")
    headers4 = ["证件号", "患者姓名", "病区", "陪护人", "过期日期", "剩余天数", "状态"]
    
    for col, header in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    today = date.today()
    three_days_later = today + timedelta(days=3)
    
    expiring_query = db.query(CareCertificate).filter(
        CareCertificate.status == "active",
        CareCertificate.expiry_date <= three_days_later
    )
    if ward:
        expiring_query = expiring_query.join(Patient, CareCertificate.patient_id == Patient.patient_id).filter(Patient.ward == ward)
    
    expiring_certs = expiring_query.order_by(CareCertificate.expiry_date).all()
    
    for row_idx, cert in enumerate(expiring_certs, 2):
        patient = db.query(Patient).filter(Patient.patient_id == cert.patient_id).first()
        caregiver = db.query(Caregiver).filter(Caregiver.caregiver_id == cert.caregiver_id).first()
        days_remaining = calculate_days_remaining(cert.expiry_date)
        
        status = "即将过期" if days_remaining >= 0 else "已过期"
        
        row_data = [
            cert.certificate_no,
            patient.name if patient else "",
            patient.ward if patient else "",
            caregiver.name if caregiver else "",
            cert.expiry_date.strftime("%Y-%m-%d") if cert.expiry_date else "",
            max(days_remaining, 0),
            status
        ]
        
        for col, value in enumerate(row_data, 1):
            cell = ws4.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = center_align
            
            if days_remaining < 0:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif days_remaining <= 3:
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    for col in range(1, len(headers4) + 1):
        ws4.column_dimensions[ws4.cell(row=1, column=col).column_letter].width = 18
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"陪护证审计报告_{date.today().strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/statistics")
def get_statistics(db: Session = Depends(get_db)):
    update_expired_certificates(db)
    
    today = date.today()
    three_days_later = today + timedelta(days=3)
    
    total_patients = db.query(Patient).count()
    active_patients = db.query(Patient).filter(Patient.is_discharged == False).count()
    discharged_patients = total_patients - active_patients
    
    total_certificates = db.query(CareCertificate).count()
    active_certificates = db.query(CareCertificate).filter(
        CareCertificate.status == "active",
        CareCertificate.expiry_date >= today
    ).count()
    expired_certificates = db.query(CareCertificate).filter(
        CareCertificate.status == "active",
        CareCertificate.expiry_date < today
    ).count()
    cancelled_certificates = db.query(CareCertificate).filter(CareCertificate.status == "cancelled").count()
    
    expiring_soon = db.query(CareCertificate).filter(
        CareCertificate.status == "active",
        CareCertificate.expiry_date >= today,
        CareCertificate.expiry_date <= three_days_later
    ).count()
    
    pending_replacements = db.query(ReplacementRequest).filter(ReplacementRequest.status == "pending").count()
    approved_replacements = db.query(ReplacementRequest).filter(ReplacementRequest.status == "approved").count()
    
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())
    
    today_operations = db.query(OperationLog).filter(
        OperationLog.created_at >= today_start,
        OperationLog.created_at <= today_end
    ).count()
    
    failed_operations = db.query(OperationLog).filter(OperationLog.result == "failed").count()
    
    return {
        "patients": {
            "total": total_patients,
            "active": active_patients,
            "discharged": discharged_patients
        },
        "certificates": {
            "total": total_certificates,
            "active": active_certificates,
            "expired": expired_certificates,
            "cancelled": cancelled_certificates,
            "expiring_soon": expiring_soon
        },
        "replacements": {
            "pending": pending_replacements,
            "approved": approved_replacements
        },
        "operations": {
            "today": today_operations,
            "failed": failed_operations
        }
    }
