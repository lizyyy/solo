import io
from datetime import datetime

from typing import Optional

from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.models import (
    Appeal, AppealStatus, MedicalRecord, Deduction, Material, MaterialStatus,
    Flag, FlagType,
)
from app.schemas import AppealCreate, AppealOut, AppealUpdate, CleanResult
from app.services.cleaner import (
    clean_dataframe, check_duplicate_appeal, fix_typos,
)

router = APIRouter(prefix="/entry", tags=["录入"])


def _generate_appeal_no(db: Session) -> str:
    today = datetime.utcnow().strftime("%Y%m%d")
    prefix = f"AP{today}"
    count = db.query(Appeal).filter(Appeal.appeal_no.like(f"{prefix}%")).count()
    return f"{prefix}-{count + 1:04d}"


@router.post("/appeals", response_model=AppealOut)
def create_appeal(data: AppealCreate, db: Session = Depends(get_db)):
    existing = db.query(Appeal).filter(
        Appeal.admission_no == data.admission_no,
        Appeal.patient_name == data.patient_name,
        Appeal.status.notin_(["rejected", "returned"]),
    ).all()

    appeal = Appeal(
        appeal_no=_generate_appeal_no(db),
        patient_name=data.patient_name,
        admission_no=data.admission_no,
        insurance_no=data.insurance_no,
        operator=data.operator,
        status=AppealStatus.DRAFT,
    )
    db.add(appeal)
    db.flush()

    if check_duplicate_appeal(
        [{"admission_no": e.admission_no, "patient_name": e.patient_name, "status": e.status.value} for e in existing],
        {"admission_no": data.admission_no, "patient_name": data.patient_name},
    ):
        flag = Flag(
            appeal_id=appeal.id,
            flag_type=FlagType.DUPLICATE_APPEAL,
            detail=f"住院号{data.admission_no}患者{data.patient_name}存在未关闭申诉",
        )
        db.add(flag)

    if data.medical_record:
        mr = MedicalRecord(appeal_id=appeal.id, **data.medical_record.model_dump())
        db.add(mr)

    if data.deductions:
        for d in data.deductions:
            db.add(Deduction(appeal_id=appeal.id, **d.model_dump()))

    if data.materials:
        for m in data.materials:
            db.add(Material(appeal_id=appeal.id, **m.model_dump()))

    required_categories = ["病案摘要", "扣款清单", "申诉材料"]
    uploaded_cats = {m.category for m in (data.materials or [])}
    missing = [c for c in required_categories if c not in uploaded_cats]
    if missing:
        flag = Flag(
            appeal_id=appeal.id,
            flag_type=FlagType.MATERIAL_MISSING,
            detail=f"缺少必要材料分类: {', '.join(missing)}",
        )
        db.add(flag)

    db.commit()
    db.refresh(appeal)
    return appeal


@router.put("/appeals/{appeal_id}", response_model=AppealOut)
def update_appeal(appeal_id: int, data: AppealUpdate, db: Session = Depends(get_db)):
    appeal = db.query(Appeal).filter(Appeal.id == appeal_id).first()
    if not appeal:
        raise HTTPException(status_code=404, detail="申诉记录不存在")

    update_data = data.model_dump(exclude_unset=True)
    for field, new_val in update_data.items():
        old_val = getattr(appeal, field)
        if old_val != new_val:
            from app.models.models import ChangeHistory
            change = ChangeHistory(
                appeal_id=appeal.id,
                field_name=field,
                old_value=str(old_val) if old_val is not None else None,
                new_value=str(new_val) if new_val is not None else None,
                changed_by=data.operator,
                reason="人工修改",
            )
            db.add(change)
            setattr(appeal, field, new_val)

    db.commit()
    db.refresh(appeal)
    return appeal


@router.get("/appeals", response_model=list[AppealOut])
def list_appeals(status: Optional[AppealStatus] = None, db: Session = Depends(get_db)):
    q = db.query(Appeal)
    if status:
        q = q.filter(Appeal.status == status)
    return q.order_by(Appeal.created_at.desc()).all()


@router.get("/appeals/{appeal_id}", response_model=AppealOut)
def get_appeal(appeal_id: int, db: Session = Depends(get_db)):
    appeal = db.query(Appeal).filter(Appeal.id == appeal_id).first()
    if not appeal:
        raise HTTPException(status_code=404, detail="申诉记录不存在")
    return appeal


@router.post("/appeals/import", response_model=CleanResult)
async def import_appeals_from_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析Excel文件")

    ws = wb.active
    rows = []
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip() if h else f"col_{j}" for j, h in enumerate(row)]
            continue
        rows.append({headers[j]: cell for j, cell in enumerate(row) if j < len(headers)})

    cleaned = clean_dataframe(rows, key_fields=["住院号", "患者姓名"])

    for row in cleaned["rows"]:
        patient_name = row.get("患者姓名", "")
        admission_no = row.get("住院号", "")
        if not patient_name:
            continue

        appeal = Appeal(
            appeal_no=_generate_appeal_no(db),
            patient_name=patient_name,
            admission_no=admission_no,
            insurance_no=row.get("医保号"),
            operator=row.get("经办人"),
            status=AppealStatus.DRAFT,
        )
        db.add(appeal)
        db.flush()

        mr = MedicalRecord(
            appeal_id=appeal.id,
            diagnosis_code=row.get("诊断编码"),
            diagnosis_name=row.get("诊断名称"),
            admission_date=row.get("入院日期"),
            discharge_date=row.get("出院日期"),
            department=row.get("科室"),
            attending_doctor=row.get("主治医师"),
            summary=row.get("病案摘要"),
        )
        db.add(mr)

        if row.get("扣款项目编码") or row.get("扣款金额"):
            ded = Deduction(
                appeal_id=appeal.id,
                item_code=row.get("扣款项目编码"),
                item_name=row.get("扣款项目名称"),
                deduction_amount=float(row.get("扣款金额") or 0),
                reason_code=row.get("扣款原因编码"),
                reason_text=row.get("扣款原因"),
                rule_version=row.get("规则版本"),
            )
            db.add(ded)

    db.commit()
    wb.close()
    return cleaned
