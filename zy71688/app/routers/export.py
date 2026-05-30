import os
from datetime import datetime

from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

from app.config import EXPORT_DIR
from app.database import get_db
from app.models.models import Appeal, ChangeHistory

router = APIRouter(prefix="/export", tags=["导出"])

os.makedirs(EXPORT_DIR, exist_ok=True)

HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
WARN_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
OVERRIDE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, row_num, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_data(ws, row_num, max_col, fill=None):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.border = THIN_BORDER
        cell.alignment = WRAP_ALIGN
        if fill:
            cell.fill = fill


@router.get("/appeals/{appeal_id}/report")
def export_appeal_report(appeal_id: int, db: Session = Depends(get_db)):
    appeal = db.query(Appeal).filter(Appeal.id == appeal_id).first()
    if not appeal:
        raise HTTPException(status_code=404, detail="申诉记录不存在")

    changes = db.query(ChangeHistory).filter(ChangeHistory.appeal_id == appeal_id).all()
    changed_fields = {(c.field_name, c.changed_by, c.changed_at, c.old_value, c.new_value) for c in changes}

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "申诉概要"
    headers = ["申诉编号", "患者姓名", "住院号", "医保号", "状态", "经办人", "复核人", "创建时间", "更新时间"]
    ws1.append(headers)
    _style_header(ws1, 1, len(headers))

    row_data = [
        appeal.appeal_no, appeal.patient_name, appeal.admission_no,
        appeal.insurance_no, appeal.status.value, appeal.operator,
        appeal.reviewer,
        appeal.created_at.strftime("%Y-%m-%d %H:%M") if appeal.created_at else "",
        appeal.updated_at.strftime("%Y-%m-%d %H:%M") if appeal.updated_at else "",
    ]
    ws1.append(row_data)
    _style_data(ws1, 2, len(headers))

    for col in range(1, len(headers) + 1):
        ws1.column_dimensions[ws1.cell(row=1, column=col).column_letter].width = 18

    if appeal.medical_record:
        ws2 = wb.create_sheet("病案摘要")
        mr = appeal.medical_record
        mr_headers = ["诊断编码", "诊断名称", "入院日期", "出院日期", "科室", "主治医师", "摘要"]
        ws2.append(mr_headers)
        _style_header(ws2, 1, len(mr_headers))
        ws2.append([
            mr.diagnosis_code, mr.diagnosis_name, mr.admission_date,
            mr.discharge_date, mr.department, mr.attending_doctor, mr.summary,
        ])
        _style_data(ws2, 2, len(mr_headers))
        for col in range(1, len(mr_headers) + 1):
            ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = 20

    if appeal.deductions:
        ws3 = wb.create_sheet("扣款明细")
        ded_headers = ["项目编码", "项目名称", "扣款金额", "原因编码", "原因", "规则版本", "是否争议"]
        ws3.append(ded_headers)
        _style_header(ws3, 1, len(ded_headers))
        for i, d in enumerate(appeal.deductions, start=2):
            ws3.append([
                d.item_code, d.item_name, d.deduction_amount,
                d.reason_code, d.reason_text, d.rule_version,
                "是" if d.is_disputed else "否",
            ])
            is_overridden = any(f[0].startswith(f"deduction.{d.id}") for f in changed_fields)
            _style_data(ws3, i, len(ded_headers), fill=OVERRIDE_FILL if is_overridden else None)
        for col in range(1, len(ded_headers) + 1):
            ws3.column_dimensions[ws3.cell(row=1, column=col).column_letter].width = 18

    if appeal.rule_matches:
        ws4 = wb.create_sheet("规则匹配")
        rm_headers = ["规则编码", "规则名称", "规则版本", "版本是否最新", "匹配结果", "说明"]
        ws4.append(rm_headers)
        _style_header(ws4, 1, len(rm_headers))
        for i, rm in enumerate(appeal.rule_matches, start=2):
            ws4.append([
                rm.rule_code, rm.rule_name, rm.rule_version,
                "是" if rm.is_version_latest else "否",
                rm.match_result, rm.explanation,
            ])
            fill = WARN_FILL if not rm.is_version_latest else None
            _style_data(ws4, i, len(rm_headers), fill=fill)
        for col in range(1, len(rm_headers) + 1):
            ws4.column_dimensions[ws4.cell(row=1, column=col).column_letter].width = 20

    if appeal.processing:
        ws5 = wb.create_sheet("处理记录")
        proc = appeal.processing
        proc_headers = ["检查项", "结果", "备注"]
        ws5.append(proc_headers)
        _style_header(ws5, 1, len(proc_headers))
        proc_data = [
            ["材料检查", "通过" if proc.material_check_passed else "未通过", proc.material_check_note or ""],
            ["规则检查", "通过" if proc.rule_check_passed else "未通过", proc.rule_check_note or ""],
            ["进度说明", "", proc.progress_note or ""],
            ["差异解释", "", proc.difference_explanation or ""],
        ]
        for i, row in enumerate(proc_data, start=2):
            ws5.append(row)
            _style_data(ws5, i, len(proc_headers))
        for col in range(1, len(proc_headers) + 1):
            ws5.column_dimensions[ws5.cell(row=1, column=col).column_letter].width = 30

    if appeal.materials:
        ws6 = wb.create_sheet("材料清单")
        mat_headers = ["分类", "文件名", "描述", "状态", "上传时间", "验证时间"]
        ws6.append(mat_headers)
        _style_header(ws6, 1, len(mat_headers))
        for i, m in enumerate(appeal.materials, start=2):
            ws6.append([
                m.category, m.file_name or "", m.description or "",
                m.status.value,
                m.uploaded_at.strftime("%Y-%m-%d %H:%M") if m.uploaded_at else "",
                m.verified_at.strftime("%Y-%m-%d %H:%M") if m.verified_at else "",
            ])
            fill = WARN_FILL if m.status.value == "missing" else None
            _style_data(ws6, i, len(mat_headers), fill=fill)
        for col in range(1, len(mat_headers) + 1):
            ws6.column_dimensions[ws6.cell(row=1, column=col).column_letter].width = 20

    if appeal.flags:
        ws7 = wb.create_sheet("标记汇总")
        flag_headers = ["标记类型", "详情", "是否解决", "创建时间"]
        ws7.append(flag_headers)
        _style_header(ws7, 1, len(flag_headers))
        for i, f in enumerate(appeal.flags, start=2):
            ws7.append([
                f.flag_type.value, f.detail or "",
                "已解决" if f.is_resolved else "未解决",
                f.created_at.strftime("%Y-%m-%d %H:%M") if f.created_at else "",
            ])
            fill = None if f.is_resolved else WARN_FILL
            _style_data(ws7, i, len(flag_headers), fill=fill)
        for col in range(1, len(flag_headers) + 1):
            ws7.column_dimensions[ws7.cell(row=1, column=col).column_letter].width = 25

    if changes:
        ws8 = wb.create_sheet("变更历史")
        ch_headers = ["字段", "旧值", "新值", "修改人", "修改时间", "原因"]
        ws8.append(ch_headers)
        _style_header(ws8, 1, len(ch_headers))
        for i, c in enumerate(changes, start=2):
            ws8.append([
                c.field_name, c.old_value or "", c.new_value or "",
                c.changed_by or "",
                c.changed_at.strftime("%Y-%m-%d %H:%M") if c.changed_at else "",
                c.reason or "",
            ])
            _style_data(ws8, i, len(ch_headers), fill=OVERRIDE_FILL)
        for col in range(1, len(ch_headers) + 1):
            ws8.column_dimensions[ws8.cell(row=1, column=col).column_letter].width = 22

    ts = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    filename = f"appeal_report_{appeal.appeal_no}_{ts}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    return FileResponse(
        path=filepath,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )


@router.get("/appeals/batch-report")
def export_batch_report(
    status: Optional[str] = None, db: Session = Depends(get_db)
):
    q = db.query(Appeal)
    if status:
        from app.models.models import AppealStatus
        try:
            q = q.filter(Appeal.status == AppealStatus(status))
        except ValueError:
            raise HTTPException(status_code=400, detail=f"无效状态: {status}")

    appeals = q.order_by(Appeal.created_at.desc()).all()
    if not appeals:
        raise HTTPException(status_code=404, detail="无符合条件的申诉记录")

    wb = Workbook()
    ws = wb.active
    ws.title = "申诉汇总"

    headers = [
        "申诉编号", "患者姓名", "住院号", "医保号", "状态",
        "经办人", "复核人", "创建时间", "扣款总额", "标记数", "未解决标记",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for i, a in enumerate(appeals, start=2):
        total_deduction = sum(d.deduction_amount or 0 for d in a.deductions)
        flag_count = len(a.flags)
        unresolved = sum(1 for f in a.flags if not f.is_resolved)
        ws.append([
            a.appeal_no, a.patient_name, a.admission_no,
            a.insurance_no, a.status.value, a.operator,
            a.reviewer,
            a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else "",
            total_deduction, flag_count, unresolved,
        ])
        fill = WARN_FILL if unresolved > 0 else None
        _style_data(ws, i, len(headers), fill=fill)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    ts = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    filename = f"appeal_batch_{ts}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    return FileResponse(
        path=filepath,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )
