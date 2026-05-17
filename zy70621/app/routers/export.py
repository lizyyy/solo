from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from app import crud, models
from app.database import get_db
from app.models import RepairStatus

router = APIRouter(
    prefix="/export",
    tags=["export"],
)


def create_excel_response(rows, headers, filename):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/repairs")
def export_repair_orders(
    status: Optional[RepairStatus] = None,
    is_overdue: Optional[bool] = None,
    is_outsourced: Optional[bool] = None,
    building_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    orders = crud.get_repair_orders(
        db, skip=0, limit=10000,
        status=status, is_overdue=is_overdue,
        is_outsourced=is_outsourced, building_id=building_id
    )
    
    headers = [
        "工单编号", "楼栋房间", "报修人", "联系电话", "报修类型",
        "描述", "紧急程度", "状态", "处理人", "是否超时",
        "是否重复工单", "报修时间", "预计完成时间", "实际完成时间"
    ]
    
    rows = []
    for order in orders:
        building = f"{order.building.building_name} {order.building.unit_number}-{order.building.room_number}" if order.building else ""
        handler_name = order.handler.name if order.handler else ""
        
        rows.append([
            order.order_no,
            building,
            order.reporter_name or "",
            order.reporter_phone or "",
            order.repair_type or "",
            order.description or "",
            order.urgency.value,
            order.status.value,
            handler_name,
            "是" if order.is_overdue else "否",
            "是" if order.is_duplicate else "否",
            order.reported_at.strftime("%Y-%m-%d %H:%M") if order.reported_at else "",
            order.expected_completion_time.strftime("%Y-%m-%d %H:%M") if order.expected_completion_time else "",
            order.actual_completion_time.strftime("%Y-%m-%d %H:%M") if order.actual_completion_time else ""
        ])
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"repair_orders_{timestamp}.xlsx"
    
    return create_excel_response(rows, headers, filename)


@router.get("/overdue")
def export_overdue_orders(db: Session = Depends(get_db)):
    orders = crud.get_overdue_orders(db, skip=0, limit=10000)
    
    headers = [
        "工单编号", "楼栋房间", "报修人", "报修类型",
        "紧急程度", "状态", "处理人", "报修时间", "超时时长(小时)"
    ]
    
    rows = []
    for order in orders:
        building = f"{order.building.building_name} {order.building.unit_number}-{order.building.room_number}" if order.building else ""
        handler_name = order.handler.name if order.handler else ""
        
        overdue_hours = 0
        if order.reported_at:
            overdue_hours = (datetime.now() - order.reported_at).total_seconds() / 3600
            overdue_hours = round(overdue_hours, 1)
        
        rows.append([
            order.order_no,
            building,
            order.reporter_name or "",
            order.repair_type or "",
            order.urgency.value,
            order.status.value,
            handler_name,
            order.reported_at.strftime("%Y-%m-%d %H:%M") if order.reported_at else "",
            overdue_hours
        ])
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"overdue_orders_{timestamp}.xlsx"
    
    return create_excel_response(rows, headers, filename)


@router.get("/outsourced")
def export_outsourced_orders(db: Session = Depends(get_db)):
    orders = crud.get_outsourced_orders(db, skip=0, limit=10000)
    
    headers = [
        "工单编号", "楼栋房间", "报修类型", "外包公司",
        "外包人员", "外包时间", "预计完成时间", "费用预估"
    ]
    
    rows = []
    for order in orders:
        building = f"{order.building.building_name} {order.building.unit_number}-{order.building.room_number}" if order.building else ""
        outsourcing_record = order.outsourcing_records[-1] if order.outsourcing_records else None
        
        outsourcer_name = ""
        company_name = ""
        outsourcing_time = ""
        expected_completion = ""
        cost_estimate = ""
        
        if outsourcing_record and outsourcing_record.outsourcer:
            outsourcer_name = outsourcing_record.outsourcer.name
            company_name = outsourcing_record.outsourcer.company_name or ""
            outsourcing_time = outsourcing_record.outsourcing_time.strftime("%Y-%m-%d %H:%M") if outsourcing_record.outsourcing_time else ""
            expected_completion = outsourcing_record.expected_completion.strftime("%Y-%m-%d %H:%M") if outsourcing_record.expected_completion else ""
            cost_estimate = outsourcing_record.cost_estimate or ""
        
        rows.append([
            order.order_no,
            building,
            order.repair_type or "",
            company_name,
            outsourcer_name,
            outsourcing_time,
            expected_completion,
            cost_estimate
        ])
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"outsourced_orders_{timestamp}.xlsx"
    
    return create_excel_response(rows, headers, filename)


@router.get("/audit-logs")
def export_audit_logs(
    order_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    logs = crud.get_audit_logs(db, order_id=order_id, skip=0, limit=10000)
    
    headers = [
        "工单编号", "操作类型", "操作人", "旧状态", "新状态",
        "操作结论", "操作原因", "原始输入", "操作时间"
    ]
    
    rows = []
    for log in logs:
        order_no = log.repair_order.order_no if log.repair_order else ""
        
        rows.append([
            order_no,
            log.action or "",
            log.operator or "",
            log.old_status or "",
            log.new_status or "",
            log.conclusion or "",
            log.reason or "",
            log.original_input or "",
            log.created_at.strftime("%Y-%m-%d %H:%M") if log.created_at else ""
        ])
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"audit_logs_{timestamp}.xlsx"
    
    return create_excel_response(rows, headers, filename)


@router.get("/statistics")
def export_statistics(db: Session = Depends(get_db)):
    all_orders = crud.get_repair_orders(db, skip=0, limit=10000)
    
    total_count = len(all_orders)
    overdue_count = sum(1 for o in all_orders if o.is_overdue)
    outsourced_count = sum(1 for o in all_orders if o.status == RepairStatus.OUTSOURCED)
    completed_count = sum(1 for o in all_orders if o.status in [RepairStatus.COMPLETED, RepairStatus.VERIFIED, RepairStatus.CLOSED])
    pending_count = sum(1 for o in all_orders if o.status == RepairStatus.PENDING)
    
    status_stats = {}
    for status in RepairStatus:
        count = sum(1 for o in all_orders if o.status == status)
        status_stats[status.value] = count
    
    urgency_stats = {}
    for urgency in models.UrgencyLevel:
        count = sum(1 for o in all_orders if o.urgency == urgency)
        urgency_stats[urgency.value] = count
    
    headers = ["统计类别", "统计项", "数值", "占比(%)"]
    rows = []
    
    rows.append(["总体统计", "总工单数量", total_count, "100"])
    rows.append(["总体统计", "已完成工单", completed_count, round(completed_count/total_count*100, 2) if total_count else 0])
    rows.append(["总体统计", "待处理工单", pending_count, round(pending_count/total_count*100, 2) if total_count else 0])
    rows.append(["总体统计", "超时工单", overdue_count, round(overdue_count/total_count*100, 2) if total_count else 0])
    rows.append(["总体统计", "外包工单", outsourced_count, round(outsourced_count/total_count*100, 2) if total_count else 0])
    rows.append(["", "", "", ""])
    
    rows.append(["状态统计", "", "", ""])
    for status, count in status_stats.items():
        rows.append(["状态统计", status, count, round(count/total_count*100, 2) if total_count else 0])
    
    rows.append(["", "", "", ""])
    rows.append(["紧急程度统计", "", "", ""])
    for urgency, count in urgency_stats.items():
        rows.append(["紧急程度统计", urgency, count, round(count/total_count*100, 2) if total_count else 0])
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"statistics_report_{timestamp}.xlsx"
    
    return create_excel_response(rows, headers, filename)
