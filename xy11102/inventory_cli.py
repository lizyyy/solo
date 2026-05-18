#!/usr/bin/env python3
import click
import pandas as pd
import hashlib
from pathlib import Path
from datetime import datetime
import re


class InventoryProcessor:
    def __init__(self, input_file, output_dir):
        self.input_file = Path(input_file)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        self.STANDARD_BARCODE_LENGTH = 13
        self.CRITICAL_COLUMNS = [
            '条码',
            '书名',
            '作者',
            '出版社',
            '定价',
            '寄卖人',
            '寄卖日期',
            '库存状态',
            '货架位置',
            '盘点数量'
        ]
        
    def load_data(self):
        if self.input_file.suffix == '.xlsx':
            df = pd.read_excel(self.input_file, dtype=str)
        elif self.input_file.suffix == '.csv':
            df = pd.read_csv(self.input_file, dtype=str)
        else:
            raise ValueError(f"不支持的文件格式: {self.input_file.suffix}")
        return df
    
    def detect_missing_leading_zeros(self, df):
        issues = []
        for idx, row in df.iterrows():
            barcode = str(row.get('条码', ''))
            if barcode and len(barcode) < self.STANDARD_BARCODE_LENGTH:
                if re.match(r'^\d+$', barcode):
                    expected_length = self.STANDARD_BARCODE_LENGTH
                    if len(barcode) < expected_length:
                        issues.append({
                            '行号': idx + 2,
                            '原始条码': barcode,
                            '问题类型': '条码前导零丢失',
                            '建议修复': barcode.zfill(expected_length),
                            '缺失零个数': expected_length - len(barcode)
                        })
        return pd.DataFrame(issues)
    
    def detect_duplicate_scans(self, df):
        barcode_counts = df['条码'].value_counts()
        duplicates = barcode_counts[barcode_counts > 1]
        issues = []
        
        for barcode, count in duplicates.items():
            duplicate_rows = df[df['条码'] == barcode]
            for idx, row in duplicate_rows.iterrows():
                issues.append({
                    '行号': idx + 2,
                    '条码': barcode,
                    '书名': row.get('书名', ''),
                    '问题类型': '重复扫码',
                    '出现次数': count,
                    '重复行': ', '.join([str(i + 2) for i in duplicate_rows.index.tolist()])
                })
        return pd.DataFrame(issues)
    
    def fix_barcode_leading_zeros(self, df):
        df = df.copy()
        for idx, row in df.iterrows():
            barcode = str(row.get('条码', ''))
            if barcode and len(barcode) < self.STANDARD_BARCODE_LENGTH:
                if re.match(r'^\d+$', barcode):
                    df.at[idx, '条码'] = barcode.zfill(self.STANDARD_BARCODE_LENGTH)
                    df.at[idx, '条码修复标记'] = '已补前导零'
                else:
                    df.at[idx, '条码修复标记'] = '非数字条码'
            else:
                df.at[idx, '条码修复标记'] = ''
        return df
    
    def stable_sort(self, df):
        df_sorted = df.sort_values(
            by=['条码', '寄卖日期', '书名'],
            ascending=[True, True, True]
        ).reset_index(drop=True)
        return df_sorted
    
    def generate_file_hash(self, df):
        content = df.to_csv(index=False, lineterminator='\n')
        return hashlib.sha256(content.encode('utf-8')).hexdigest()[:16]
    
    def process(self):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        
        df = self.load_data()
        
        leading_zero_issues = self.detect_missing_leading_zeros(df)
        duplicate_issues = self.detect_duplicate_scans(df)
        
        df_fixed = self.fix_barcode_leading_zeros(df)
        df_sorted = self.stable_sort(df_fixed)
        
        file_hash = self.generate_file_hash(df_sorted)
        
        output_columns = [col for col in self.CRITICAL_COLUMNS if col in df_sorted.columns]
        output_columns.append('条码修复标记')
        
        df_output = df_sorted[output_columns]
        
        main_output = self.output_dir / f'inventory_result_{timestamp}_{file_hash}.xlsx'
        issues_output = self.output_dir / f'inventory_issues_{timestamp}_{file_hash}.xlsx'
        
        from openpyxl.styles import numbers
        
        with pd.ExcelWriter(main_output, engine='openpyxl') as writer:
            df_output.to_excel(writer, sheet_name='盘点结果', index=False)
            worksheet = writer.sheets['盘点结果']
            barcode_col_idx = df_output.columns.get_loc('条码') + 1
            for row in range(2, len(df_output) + 2):
                cell = worksheet.cell(row=row, column=barcode_col_idx)
                cell.number_format = numbers.FORMAT_TEXT
            
            summary_data = {
                '统计项': [
                    '总记录数',
                    '条码前导零丢失数量',
                    '重复扫码记录数',
                    '文件哈希',
                    '生成时间'
                ],
                '数值': [
                    len(df),
                    len(leading_zero_issues),
                    len(duplicate_issues),
                    file_hash,
                    datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                ]
            }
            pd.DataFrame(summary_data).to_excel(writer, sheet_name='统计摘要', index=False)
        
        with pd.ExcelWriter(issues_output, engine='openpyxl') as writer:
            if not leading_zero_issues.empty:
                leading_zero_issues.to_excel(writer, sheet_name='条码前导零问题', index=False)
                worksheet = writer.sheets['条码前导零问题']
                original_col_idx = leading_zero_issues.columns.get_loc('原始条码') + 1
                suggest_col_idx = leading_zero_issues.columns.get_loc('建议修复') + 1
                for row in range(2, len(leading_zero_issues) + 2):
                    worksheet.cell(row=row, column=original_col_idx).number_format = numbers.FORMAT_TEXT
                    worksheet.cell(row=row, column=suggest_col_idx).number_format = numbers.FORMAT_TEXT
            if not duplicate_issues.empty:
                duplicate_issues.to_excel(writer, sheet_name='重复扫码问题', index=False)
                worksheet = writer.sheets['重复扫码问题']
                barcode_col_idx = duplicate_issues.columns.get_loc('条码') + 1
                for row in range(2, len(duplicate_issues) + 2):
                    worksheet.cell(row=row, column=barcode_col_idx).number_format = numbers.FORMAT_TEXT
            
            if leading_zero_issues.empty and duplicate_issues.empty:
                pd.DataFrame({'消息': ['未发现异常问题']}).to_excel(
                    writer, sheet_name='检查结果', index=False
                )
        
        return {
            'main_output': main_output,
            'issues_output': issues_output,
            'total_records': len(df),
            'leading_zero_issues': len(leading_zero_issues),
            'duplicate_issues': len(duplicate_issues),
            'file_hash': file_hash
        }


@click.group()
def cli():
    """二手书店寄卖库存盘点 CLI - 处理条码前导零丢失、重复扫码和可复跑输出"""
    pass


@cli.command()
@click.argument('input_file', type=click.Path(exists=True))
@click.option('--output-dir', '-o', default='./output', help='输出目录')
def check(input_file, output_dir):
    """检查库存数据中的异常问题"""
    click.echo(f"\n=== 二手书店寄卖库存盘点检查 ===")
    click.echo(f"输入文件: {input_file}")
    click.echo(f"输出目录: {output_dir}\n")
    
    processor = InventoryProcessor(input_file, output_dir)
    result = processor.process()
    
    click.echo(click.style("✓ 处理完成！", fg='green'))
    click.echo(f"\n统计摘要:")
    click.echo(f"  总记录数: {result['total_records']}")
    click.echo(f"  条码前导零丢失: {click.style(str(result['leading_zero_issues']), fg='yellow')} 条")
    click.echo(f"  重复扫码记录: {click.style(str(result['duplicate_issues']), fg='red')} 条")
    click.echo(f"  文件哈希: {result['file_hash']}")
    click.echo(f"\n输出文件:")
    click.echo(f"  盘点结果: {result['main_output']}")
    click.echo(f"  异常报告: {result['issues_output']}")


@cli.command()
@click.argument('sample_dir', type=click.Path(), default='./samples')
def generate_samples(sample_dir):
    """生成二手书店寄卖库存盘点样例数据"""
    import samples as samples_module
    samples_module.generate_sample_files(sample_dir)
    click.echo(click.style(f"✓ 样例数据已生成到: {sample_dir}", fg='green'))


if __name__ == '__main__':
    cli()
