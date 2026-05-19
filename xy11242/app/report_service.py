import io
import csv
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session

from app.models import Book, BookStatus, BookException, ImportLog, BadRecord
from app.schemas import BookQueryParams
from app import crud


class ReportExporter:
    @staticmethod
    def _format_book_data(book: Book, db: Session) -> dict:
        exceptions = crud.get_book_exceptions(db, book.id)
        exception_types = ", ".join([exc.exception_type.value for exc in exceptions if not exc.resolved])

        return {
            "ID": book.id,
            "ISBN": book.isbn or "",
            "书名": book.title or "",
            "作者": book.author or "",
            "出版社": book.publisher or "",
            "品相": book.condition.value if book.condition else "",
            "年级标签": book.grade_level.value if book.grade_level else "",
            "状态": book.status.value if book.status else "",
            "负责人": book.volunteer or "",
            "架位": book.shelf_location or "",
            "当前异常": exception_types,
            "导入来源": book.import_source.value if book.import_source else "",
            "导入批次": book.import_batch_id or "",
            "创建时间": book.created_at.strftime("%Y-%m-%d %H:%M:%S") if book.created_at else "",
            "更新时间": book.updated_at.strftime("%Y-%m-%d %H:%M:%S") if book.updated_at else "",
            "备注": book.remarks or ""
        }

    @classmethod
    def export_to_excel(cls, db: Session, books: List[Book], query_params: BookQueryParams = None) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "书籍清单"

        headers = [
            "ID", "ISBN", "书名", "作者", "出版社", "品相", "年级标签",
            "状态", "负责人", "架位", "当前异常", "导入来源", "导入批次",
            "创建时间", "更新时间", "备注"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        for row_num, book in enumerate(books, 2):
            book_data = cls._format_book_data(book, db)
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num, value=book_data[header])
                if book.status == BookStatus.EXCEPTION and header == "状态":
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif book.status == BookStatus.APPROVED and header == "状态":
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column].width = adjusted_width

        if query_params:
            summary_ws = wb.create_sheet("查询条件")
            summary_ws.append(["查询条件", "值"])
            summary_ws.append(["负责人", query_params.volunteer or "全部"])
            summary_ws.append(["状态", query_params.status.value if query_params.status else "全部"])
            summary_ws.append(["异常类型", query_params.exception_type.value if query_params.exception_type else "全部"])
            summary_ws.append(["年级", query_params.grade_level.value if query_params.grade_level else "全部"])
            summary_ws.append(["品相", query_params.condition.value if query_params.condition else "全部"])
            summary_ws.append(["开始日期", query_params.start_date.strftime("%Y-%m-%d") if query_params.start_date else ""])
            summary_ws.append(["结束日期", query_params.end_date.strftime("%Y-%m-%d") if query_params.end_date else ""])
            summary_ws.append(["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
            summary_ws.append(["记录总数", len(books)])

            for col in summary_ws.columns:
                summary_ws.column_dimensions[col[0].column_letter].width = 20

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @classmethod
    def export_to_csv(cls, db: Session, books: List[Book]) -> bytes:
        output = io.StringIO()
        headers = [
            "ID", "ISBN", "书名", "作者", "出版社", "品相", "年级标签",
            "状态", "负责人", "架位", "当前异常", "导入来源", "导入批次",
            "创建时间", "更新时间", "备注"
        ]

        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()

        for book in books:
            writer.writerow(cls._format_book_data(book, db))

        return output.getvalue().encode('utf-8-sig')

    @staticmethod
    def generate_monthly_summary(db: Session, year: int, month: int) -> dict:
        from sqlalchemy import func, extract

        start_date = datetime(year, month, 1)
        if month == 12:
            end_date = datetime(year + 1, 1, 1)
        else:
            end_date = datetime(year, month + 1, 1)

        total_books = db.query(Book).filter(
            Book.created_at >= start_date,
            Book.created_at < end_date
        ).count()

        status_stats = db.query(
            Book.status,
            func.count(Book.id)
        ).filter(
            Book.created_at >= start_date,
            Book.created_at < end_date
        ).group_by(Book.status).all()

        volunteer_stats = db.query(
            Book.volunteer,
            func.count(Book.id)
        ).filter(
            Book.created_at >= start_date,
            Book.created_at < end_date,
            Book.volunteer.isnot(None)
        ).group_by(Book.volunteer).order_by(func.count(Book.id).desc()).all()

        exception_stats = db.query(
            BookException.exception_type,
            func.count(BookException.id)
        ).filter(
            BookException.created_at >= start_date,
            BookException.created_at < end_date
        ).group_by(BookException.exception_type).all()

        import_stats = db.query(
            ImportLog.source,
            func.sum(ImportLog.total_records)
        ).filter(
            ImportLog.imported_at >= start_date,
            ImportLog.imported_at < end_date
        ).group_by(ImportLog.source).all()

        total_exceptions = db.query(BookException).filter(
            BookException.created_at >= start_date,
            BookException.created_at < end_date
        ).count()

        resolved_exceptions = db.query(BookException).filter(
            BookException.created_at >= start_date,
            BookException.created_at < end_date,
            BookException.resolved == True
        ).count()

        return {
            "period": f"{year}年{month}月",
            "total_books": total_books,
            "status_breakdown": {status.value if status else "未知": count for status, count in status_stats},
            "volunteer_breakdown": {volunteer: count for volunteer, count in volunteer_stats},
            "exception_breakdown": {exc_type.value: count for exc_type, count in exception_stats},
            "import_breakdown": {source.value: int(total) if total else 0 for source, total in import_stats},
            "exception_rate": round(total_exceptions / total_books * 100, 2) if total_books > 0 else 0,
            "exception_resolution_rate": round(resolved_exceptions / total_exceptions * 100, 2) if total_exceptions > 0 else 0,
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

    @classmethod
    def export_monthly_report(cls, db: Session, year: int, month: int) -> bytes:
        summary = cls.generate_monthly_summary(db, year, month)

        wb = Workbook()
        ws = wb.active
        ws.title = "月度复盘"

        ws.append(["公益书库月度复盘报告", ""])
        ws.append(["统计期间", summary["period"]])
        ws.append(["生成时间", summary["generated_at"]])
        ws.append([])
        ws.append(["一、总体统计", ""])
        ws.append(["入库书籍总数", summary["total_books"]])
        ws.append(["异常率", f"{summary['exception_rate']}%"])
        ws.append(["异常解决率", f"{summary['exception_resolution_rate']}%"])
        ws.append([])

        ws.append(["二、状态分布", ""])
        for status, count in summary["status_breakdown"].items():
            ws.append([status, count])
        ws.append([])

        ws.append(["三、志愿者贡献", ""])
        for volunteer, count in summary["volunteer_breakdown"].items():
            ws.append([volunteer, count])
        ws.append([])

        ws.append(["四、异常类型分布", ""])
        for exc_type, count in summary["exception_breakdown"].items():
            ws.append([exc_type, count])
        ws.append([])

        ws.append(["五、导入来源分布", ""])
        for source, count in summary["import_breakdown"].items():
            ws.append([source, count])

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 25

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def export_bad_records(db: Session, records: List[BadRecord]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "导入失败记录"

        headers = ["ID", "原始位置", "原始数据", "失败原因", "修改建议", "是否已重试", "重试时间", "是否解决"]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill

        for row_num, record in enumerate(records, 2):
            ws.cell(row=row_num, column=1, value=record.id)
            ws.cell(row=row_num, column=2, value=record.original_position)
            ws.cell(row=row_num, column=3, value=record.raw_data)
            ws.cell(row=row_num, column=4, value=record.failure_reason)
            ws.cell(row=row_num, column=5, value=record.suggestion)
            ws.cell(row=row_num, column=6, value="是" if record.is_retried else "否")
            ws.cell(row=row_num, column=7, value=record.retried_at.strftime("%Y-%m-%d %H:%M:%S") if record.retried_at else "")
            ws.cell(row=row_num, column=8, value="是" if record.resolved else "否")

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 20

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
