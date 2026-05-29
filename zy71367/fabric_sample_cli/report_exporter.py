import os
import json
from datetime import datetime, date
from typing import List, Dict, Any, Optional
from tabulate import tabulate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import DataStore, FabricSample, Garment, MatchingRecord, ArrivalReminder, InspectionStatus, ArrivalStatus
from .storage import StorageManager


class ReportExporter:
    def __init__(self, storage: StorageManager):
        self.storage = storage
        self.export_dir = os.path.join(os.path.dirname(__file__), 'exports')

    def export_matching_report(self, store: DataStore, include_calculations: bool = True) -> Dict[str, Any]:
        report_data = {
            'report_type': '样卡匹配报告',
            'generated_at': datetime.now().isoformat(),
            'summary': {
                'total_samples': len(store.samples),
                'missing_samples': sum(1 for s in store.samples.values() if s.is_missing),
                'total_garments': len(store.garments),
                'matched_pairs': len(store.matches),
                'unmatched_samples': sum(1 for s in store.samples.values() if not any(m.sample_id == s.sample_id for m in store.matches.values()) and not s.is_missing),
                'unmatched_garments': sum(1 for g in store.garments.values() if g.sample_id is None),
                'auto_matches': sum(1 for m in store.matches.values() if not m.is_manual),
                'manual_matches': sum(1 for m in store.matches.values() if m.is_manual),
            },
            'matches': [],
            'calculation_details': []
        }

        for match in store.matches.values():
            sample = store.samples.get(match.sample_id)
            garment = store.garments.get(match.garment_id)
            supplier = store.suppliers.get(match.supplier_id)

            match_data = {
                'match_id': match.match_id,
                'match_score': match.match_score,
                'color_match_score': match.color_match_score,
                'supplier_match_score': match.supplier_match_score,
                'is_manual': match.is_manual,
                'matched_by': match.matched_by,
                'matched_at': match.matched_at.isoformat(),
                'sample': {
                    'sample_id': sample.sample_id if sample else None,
                    'fabric_name': sample.fabric_name if sample else None,
                    'color_code': sample.color_spec.color_code if sample else None,
                    'normalized_hex': sample.color_spec.normalized_hex if sample else None,
                },
                'garment': {
                    'garment_id': garment.garment_id if garment else None,
                    'style_name': garment.style_name if garment else None,
                    'color_code': garment.color_spec.color_code if garment and garment.color_spec else None,
                },
                'supplier': {
                    'supplier_id': supplier.supplier_id if supplier else None,
                    'supplier_name': supplier.supplier_name if supplier else None,
                }
            }
            report_data['matches'].append(match_data)

        if include_calculations:
            report_data['calculation_details'].append({
                'title': '匹配置信度计算方法',
                'formula': '总分 = 颜色匹配度 * 0.6 + 供应商匹配度 * 0.3 + 编号关联度 * 0.1',
                'weights': {
                    'color': 0.6,
                    'supplier': 0.3,
                    'id_correlation': 0.1
                },
                'threshold': 0.5,
                'description': '总分 >= 0.5 判定为匹配成功'
            })

            report_data['calculation_details'].append({
                'title': '颜色匹配度计算方法',
                'formula': '匹配度 = 1 - 色差 / 255',
                'color_distance_formula': 'Redmean公式: sqrt((2+rmean/256)*r² + 4*g² + (2+(255-rmean)/256)*b²)',
                'thresholds': {
                    'exact_match': '< 10',
                    'similar': '10 - 40',
                    'acceptable': '40 - 100',
                    'different': '> 100'
                }
            })

        return report_data

    def export_arrival_report(self, store: DataStore, include_calculations: bool = True) -> Dict[str, Any]:
        report_data = {
            'report_type': '到货与检验状态报告',
            'generated_at': datetime.now().isoformat(),
            'reference_date': date.today().isoformat(),
            'summary': {
                'total_samples': len(store.samples),
                'on_time': 0,
                'early': 0,
                'delayed': 0,
                'not_arrived': 0,
                'missing_samples': 0,
                'by_status': {},
            },
            'reminders': [],
            'calculation_details': []
        }

        for status in InspectionStatus:
            count = sum(1 for s in store.samples.values() if s.inspection_status == status)
            report_data['summary']['by_status'][status.value] = count

        for reminder in store.reminders.values():
            sample = store.samples.get(reminder.sample_id)

            if reminder.arrival_status == ArrivalStatus.ON_TIME:
                report_data['summary']['on_time'] += 1
            elif reminder.arrival_status == ArrivalStatus.EARLY:
                report_data['summary']['early'] += 1
            elif reminder.arrival_status == ArrivalStatus.DELAYED:
                report_data['summary']['delayed'] += 1
            elif reminder.arrival_status == ArrivalStatus.NOT_ARRIVED:
                report_data['summary']['not_arrived'] += 1

            if sample and sample.is_missing:
                report_data['summary']['missing_samples'] += 1

            reminder_data = {
                'reminder_id': reminder.reminder_id,
                'sample_id': reminder.sample_id,
                'garment_id': reminder.garment_id,
                'arrival_status': reminder.arrival_status.value,
                'expected_arrival': reminder.expected_arrival.isoformat(),
                'actual_arrival': reminder.actual_arrival.isoformat() if reminder.actual_arrival else None,
                'inspection_deadline': reminder.inspection_deadline.isoformat() if reminder.inspection_deadline else None,
                'days_overdue': reminder.days_overdue,
                'days_before_deadline': reminder.days_before_deadline,
                'status_transition_chain': reminder.status_transition_chain,
                'fabric_name': sample.fabric_name if sample else None,
                'inspection_status': sample.inspection_status.value if sample else None,
                'calculation_process': reminder.calculation_process if include_calculations else None
            }
            report_data['reminders'].append(reminder_data)

        if include_calculations:
            report_data['calculation_details'].append({
                'title': '到货状态判定规则',
                'rules': [
                    '实际到货日期 = 预计到货日期 → 准时到货',
                    '实际到货日期 < 预计到货日期 → 提前到货',
                    '实际到货日期 > 预计到货日期 → 到货逾期 (逾期天数=实际-预计)',
                    '当前日期 > 预计到货日期 且 未到货 → 到货逾期 (逾期天数=当前-预计)',
                    '当前日期 <= 预计到货日期 且 未到货 → 未到货',
                ],
                'thresholds': {
                    'delay_warning': '>= 3 天 (警告)',
                    'delay_error': '> 7 天 (错误)',
                    'deadline_warning': '<= 7 天 (检验截止提醒)'
                }
            })

            report_data['calculation_details'].append({
                'title': '状态流转可靠性校验',
                'checks': [
                    '到货日期 > 检验截止日期 且 检验状态 ≠ 待检验/已撤回 → 状态流转异常',
                    '状态变更必须遵循流转规则: PENDING → IN_PROGRESS → PASSED/FAILED → REINSPECT → ...',
                    '撤回操作必须解除所有关联的成衣匹配'
                ]
            })

        return report_data

    def export_to_excel(self, store: DataStore, filename: Optional[str] = None) -> str:
        if not filename:
            filename = f"面料样卡报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        elif not filename.endswith('.xlsx'):
            filename = f"{filename}.xlsx"
        filepath = os.path.join(self.export_dir, filename)

        wb = Workbook()

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        ws1 = wb.active
        ws1.title = "样卡清单"
        headers = ["样卡编号", "面料名称", "面料成分", "色号", "标准化HEX", "供应商", "预计到货", "实际到货", "检验截止", "检验状态", "样卡缺失", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, sample in enumerate(store.samples.values(), 2):
            supplier = store.suppliers.get(sample.supplier_id)
            data = [
                sample.sample_id,
                sample.fabric_name,
                sample.fabric_type or "",
                sample.color_spec.color_code,
                sample.color_spec.normalized_hex or "",
                supplier.supplier_name if supplier else sample.supplier_id,
                sample.expected_arrival.isoformat() if sample.expected_arrival else "",
                sample.arrival_date.isoformat() if sample.arrival_date else "",
                sample.inspection_deadline.isoformat() if sample.inspection_deadline else "",
                sample.inspection_status.value,
                "是" if sample.is_missing else "否",
                sample.remark or ""
            ]
            for col, value in enumerate(data, 1):
                cell = ws1.cell(row=row, column=col, value=value)
                cell.border = thin_border
                if sample.is_missing:
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        ws2 = wb.create_sheet("成衣清单")
        headers = ["成衣编号", "款式名称", "款式代码", "关联样卡", "供应商", "颜色", "试装日期", "走秀顺序", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, garment in enumerate(store.garments.values(), 2):
            supplier = store.suppliers.get(garment.supplier_id) if garment.supplier_id else None
            data = [
                garment.garment_id,
                garment.style_name,
                garment.style_code or "",
                garment.sample_id or "",
                supplier.supplier_name if supplier else (garment.supplier_id or ""),
                garment.color_spec.color_code if garment.color_spec else "",
                garment.fitting_date.isoformat() if garment.fitting_date else "",
                garment.show_order or "",
                garment.remark or ""
            ]
            for col, value in enumerate(data, 1):
                cell = ws2.cell(row=row, column=col, value=value)
                cell.border = thin_border

        ws3 = wb.create_sheet("匹配结果")
        headers = ["匹配ID", "样卡编号", "成衣编号", "供应商", "匹配总分", "颜色匹配分", "供应商匹配分", "匹配类型", "匹配人", "匹配时间"]
        for col, header in enumerate(headers, 1):
            cell = ws3.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, match in enumerate(store.matches.values(), 2):
            supplier = store.suppliers.get(match.supplier_id)
            data = [
                match.match_id,
                match.sample_id,
                match.garment_id,
                supplier.supplier_name if supplier else match.supplier_id,
                match.match_score,
                match.color_match_score,
                match.supplier_match_score,
                "人工" if match.is_manual else "自动",
                match.matched_by or "",
                match.matched_at.strftime('%Y-%m-%d %H:%M:%S')
            ]
            for col, value in enumerate(data, 1):
                cell = ws3.cell(row=row, column=col, value=value)
                cell.border = thin_border
                if match.match_score < 0.7:
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        ws4 = wb.create_sheet("到货提醒")
        headers = ["提醒ID", "样卡编号", "成衣编号", "到货状态", "预计到货", "实际到货", "检验截止", "逾期天数", "距截止天数", "状态流转链"]
        for col, header in enumerate(headers, 1):
            cell = ws4.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, reminder in enumerate(store.reminders.values(), 2):
            data = [
                reminder.reminder_id,
                reminder.sample_id,
                reminder.garment_id or "",
                reminder.arrival_status.value,
                reminder.expected_arrival.isoformat(),
                reminder.actual_arrival.isoformat() if reminder.actual_arrival else "",
                reminder.inspection_deadline.isoformat() if reminder.inspection_deadline else "",
                reminder.days_overdue,
                reminder.days_before_deadline if reminder.days_before_deadline is not None else "",
                " → ".join(reminder.status_transition_chain)
            ]
            for col, value in enumerate(data, 1):
                cell = ws4.cell(row=row, column=col, value=value)
                cell.border = thin_border
                if reminder.arrival_status == ArrivalStatus.DELAYED:
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        ws5 = wb.create_sheet("错误清单")
        headers = ["错误ID", "错误类型", "错误代码", "严重程度", "错误描述", "关联样卡", "关联成衣", "创建时间", "是否解决"]
        for col, header in enumerate(headers, 1):
            cell = ws5.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row, error in enumerate(store.errors.values(), 2):
            data = [
                error.error_id,
                error.error_type,
                error.error_code,
                error.severity,
                error.message,
                error.related_sample_id or "",
                error.related_garment_id or "",
                error.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                "是" if error.resolved else "否"
            ]
            for col, value in enumerate(data, 1):
                cell = ws5.cell(row=row, column=col, value=value)
                cell.border = thin_border
                if error.severity == "error":
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif error.severity == "warning":
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        for ws in [ws1, ws2, ws3, ws4, ws5]:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filepath)
        return filepath

    def print_console_report(self, store: DataStore, report_type: str = "summary") -> str:
        output = []

        if report_type in ["summary", "all"]:
            output.append("")
            output.append("=" * 80)
            output.append("【输出区】面料样卡管理汇总报告")
            output.append("=" * 80)

            summary_data = [
                ["总样卡数", len(store.samples)],
                ["样卡缺失", sum(1 for s in store.samples.values() if s.is_missing)],
                ["总成衣数", len(store.garments)],
                ["供应商数", len(store.suppliers)],
                ["匹配成功", len(store.matches)],
                ["自动匹配", sum(1 for m in store.matches.values() if not m.is_manual)],
                ["人工匹配", sum(1 for m in store.matches.values() if m.is_manual)],
                ["待检验", sum(1 for s in store.samples.values() if s.inspection_status.value == "待检验")],
                ["检验中", sum(1 for s in store.samples.values() if s.inspection_status.value == "检验中")],
                ["检验通过", sum(1 for s in store.samples.values() if s.inspection_status.value == "检验通过")],
                ["检验不合格", sum(1 for s in store.samples.values() if s.inspection_status.value == "检验不合格")],
                ["已撤回", sum(1 for s in store.samples.values() if s.inspection_status.value == "已撤回")],
                ["到货提醒", len(store.reminders)],
                ["错误记录", len(store.errors)],
            ]
            output.append(tabulate(summary_data, headers=["指标", "数值"], tablefmt="grid"))
            output.append("")

        if report_type in ["matches", "all"]:
            output.append("")
            output.append("=" * 80)
            output.append("【输出区】匹配结果明细")
            output.append("=" * 80)

            if store.matches:
                match_data = []
                for m in store.matches.values():
                    sample = store.samples.get(m.sample_id)
                    garment = store.garments.get(m.garment_id)
                    match_data.append([
                        m.match_id[-8:],
                        m.sample_id,
                        sample.fabric_name if sample else "",
                        m.garment_id,
                        garment.style_name if garment else "",
                        f"{m.match_score:.2%}",
                        "人工" if m.is_manual else "自动"
                    ])
                output.append(tabulate(match_data, headers=["匹配ID", "样卡编号", "面料", "成衣编号", "款式", "匹配度", "类型"], tablefmt="grid"))
            else:
                output.append("暂无匹配记录")
            output.append("")

        if report_type in ["errors", "all"]:
            output.append("")
            output.append("=" * 80)
            output.append("【错误清单】")
            output.append("=" * 80)

            unresolved_errors = [e for e in store.errors.values() if not e.resolved]
            if unresolved_errors:
                error_data = []
                for e in unresolved_errors:
                    error_data.append([
                        e.error_id[-8:],
                        e.error_type,
                        e.error_code,
                        e.severity,
                        e.message[:50] + "..." if len(e.message) > 50 else e.message,
                        e.related_sample_id or "",
                        e.created_at.strftime('%m-%d %H:%M')
                    ])
                output.append(tabulate(error_data, headers=["错误ID", "类型", "代码", "严重度", "描述", "关联样卡", "时间"], tablefmt="grid"))
            else:
                output.append("暂无未解决错误")
            output.append("")

        if report_type in ["history", "all"]:
            output.append("")
            output.append("=" * 80)
            output.append("【历史目录】最近10条操作记录")
            output.append("=" * 80)

            history_list = sorted(store.history.values(), key=lambda h: h.operation_time, reverse=True)[:10]
            if history_list:
                hist_data = []
                for h in history_list:
                    hist_data.append([
                        h.record_id[-8:],
                        h.operation_type.value,
                        h.entity_type,
                        h.entity_id,
                        h.operator or "",
                        h.change_reason[:40] if h.change_reason else "",
                        h.operation_time.strftime('%m-%d %H:%M')
                    ])
                output.append(tabulate(hist_data, headers=["记录ID", "操作", "实体", "编号", "操作人", "原因", "时间"], tablefmt="grid"))
            else:
                output.append("暂无历史记录")
            output.append("")

        return "\n".join(output)
