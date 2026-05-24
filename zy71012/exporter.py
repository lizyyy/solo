from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from datetime import datetime
import models
import crud


def export_to_excel(
    db: Session,
    status: str = None,
    elevator_no: str = None,
    is_timeout: bool = None,
    start_time: datetime = None,
    end_time: datetime = None,
    source: str = None,
    maintenance_person: str = None,
    has_merged: bool = None,
    has_resubmit: bool = None,
    include_merged: bool = False,
) -> Workbook:
    alarms, total = crud.get_alarms(
        db, 0, 10000, status, elevator_no, is_timeout,
        start_time, end_time, source, maintenance_person,
        has_merged, has_resubmit, include_merged
    )
    
    wb = Workbook()
    
    ws_main = wb.active
    ws_main.title = "报警汇总"
    
    headers = [
        "报警编号", "电梯编号", "报警时间", "乘客人数", "报警来源",
        "位置", "状态", "是否超时", "超时原因", "维保人员",
        "派单时间", "到场时间", "解决时间", "处理结果",
        "合并到主记录ID", "合并次数", "重新提交次数", "创建人", "创建时间"
    ]
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws_main.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    status_colors = {
        "pending": "FFFFEB",
        "maintenance_dispatched": "FFF2CC",
        "on_site": "DDEBF7",
        "rescuing": "DAEEF3",
        "resolved": "E2EFDA",
        "reviewing": "DDEBF7",
        "rejected": "FCE4D6",
        "closed": "E2EFDA",
        "cancelled": "D9D9D9"
    }
    
    for row, alarm in enumerate(alarms, 2):
        ws_main.cell(row=row, column=1, value=alarm.alarm_no)
        ws_main.cell(row=row, column=2, value=alarm.elevator_no)
        ws_main.cell(row=row, column=3, value=alarm.alarm_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.alarm_time else "")
        ws_main.cell(row=row, column=4, value=alarm.passenger_count)
        ws_main.cell(row=row, column=5, value=alarm.source)
        ws_main.cell(row=row, column=6, value=alarm.location or "")
        ws_main.cell(row=row, column=7, value=alarm.status)
        ws_main.cell(row=row, column=8, value="是" if alarm.is_timeout else "否")
        ws_main.cell(row=row, column=9, value=alarm.timeout_reason or "")
        ws_main.cell(row=row, column=10, value=alarm.maintenance_person or "")
        ws_main.cell(row=row, column=11, value=alarm.dispatched_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.dispatched_time else "")
        ws_main.cell(row=row, column=12, value=alarm.arrived_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.arrived_time else "")
        ws_main.cell(row=row, column=13, value=alarm.resolved_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.resolved_time else "")
        ws_main.cell(row=row, column=14, value=alarm.resolution or "")
        ws_main.cell(row=row, column=15, value=alarm.parent_id or "")
        ws_main.cell(row=row, column=16, value=alarm.merge_count)
        ws_main.cell(row=row, column=17, value=alarm.resubmit_count)
        ws_main.cell(row=row, column=18, value=alarm.created_by)
        ws_main.cell(row=row, column=19, value=alarm.created_at.strftime("%Y-%m-%d %H:%M:%S") if alarm.created_at else "")
        
        status_color = status_colors.get(alarm.status, "FFFFFF")
        fill = PatternFill(start_color=status_color, end_color=status_color, fill_type="solid")
        for col in range(1, len(headers) + 1):
            ws_main.cell(row=row, column=col).fill = fill
    
    column_widths = [18, 12, 20, 10, 10, 20, 18, 10, 20, 12, 20, 20, 20, 30, 18, 10, 12, 12, 20]
    for col, width in enumerate(column_widths, 1):
        ws_main.column_dimensions[get_column_letter(col)].width = width
    
    ws_calls = wb.create_sheet("通话记录")
    call_headers = [
        "报警编号", "通话类型", "通话人", "联系电话",
        "通话时间", "时长(秒)", "通话内容", "操作员", "记录时间"
    ]
    
    for col, header in enumerate(call_headers, 1):
        cell = ws_calls.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    call_row = 2
    for alarm in alarms:
        for call in alarm.call_records:
            ws_calls.cell(row=call_row, column=1, value=alarm.alarm_no)
            ws_calls.cell(row=call_row, column=2, value=call.call_type)
            ws_calls.cell(row=call_row, column=3, value=call.caller)
            ws_calls.cell(row=call_row, column=4, value=call.caller_phone or "")
            ws_calls.cell(row=call_row, column=5, value=call.call_time.strftime("%Y-%m-%d %H:%M:%S") if call.call_time else "")
            ws_calls.cell(row=call_row, column=6, value=call.duration or "")
            ws_calls.cell(row=call_row, column=7, value=call.content)
            ws_calls.cell(row=call_row, column=8, value=call.operator)
            ws_calls.cell(row=call_row, column=9, value=call.created_at.strftime("%Y-%m-%d %H:%M:%S") if call.created_at else "")
            call_row += 1
    
    call_widths = [18, 12, 12, 15, 20, 10, 40, 12, 20]
    for col, width in enumerate(call_widths, 1):
        ws_calls.column_dimensions[get_column_letter(col)].width = width
    
    ws_status = wb.create_sheet("状态流转")
    status_headers = [
        "报警编号", "原状态", "新状态", "操作员", "备注", "变更时间"
    ]
    
    for col, header in enumerate(status_headers, 1):
        cell = ws_status.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    status_row = 2
    for alarm in alarms:
        for history in alarm.status_histories:
            ws_status.cell(row=status_row, column=1, value=alarm.alarm_no)
            ws_status.cell(row=status_row, column=2, value=history.from_status or "无")
            ws_status.cell(row=status_row, column=3, value=history.to_status)
            ws_status.cell(row=status_row, column=4, value=history.operator)
            ws_status.cell(row=status_row, column=5, value=history.remark or "")
            ws_status.cell(row=status_row, column=6, value=history.change_time.strftime("%Y-%m-%d %H:%M:%S") if history.change_time else "")
            status_row += 1
    
    status_widths = [18, 15, 15, 12, 30, 20]
    for col, width in enumerate(status_widths, 1):
        ws_status.column_dimensions[get_column_letter(col)].width = width
    
    ws_audit = wb.create_sheet("修改日志")
    audit_headers = [
        "报警编号", "修改字段", "原值", "新值",
        "操作员", "修改原因", "修改时间"
    ]
    
    for col, header in enumerate(audit_headers, 1):
        cell = ws_audit.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    audit_row = 2
    for alarm in alarms:
        for audit in alarm.audit_logs:
            ws_audit.cell(row=audit_row, column=1, value=alarm.alarm_no)
            ws_audit.cell(row=audit_row, column=2, value=audit.field_name)
            ws_audit.cell(row=audit_row, column=3, value=audit.old_value or "")
            ws_audit.cell(row=audit_row, column=4, value=audit.new_value or "")
            ws_audit.cell(row=audit_row, column=5, value=audit.operator)
            ws_audit.cell(row=audit_row, column=6, value=audit.change_reason or "")
            ws_audit.cell(row=audit_row, column=7, value=audit.change_time.strftime("%Y-%m-%d %H:%M:%S") if audit.change_time else "")
            audit_row += 1
    
    audit_widths = [18, 15, 20, 20, 12, 20, 20]
    for col, width in enumerate(audit_widths, 1):
        ws_audit.column_dimensions[get_column_letter(col)].width = width
    
    return wb


def export_single_to_excel(db: Session, alarm_id: int) -> Workbook:
    alarm = crud.get_alarm(db, alarm_id)
    if not alarm:
        raise ValueError("报警记录不存在")
    
    wb = Workbook()
    
    ws_info = wb.active
    ws_info.title = "报警详情"
    
    info_data = [
        ["报警编号", alarm.alarm_no],
        ["电梯编号", alarm.elevator_no],
        ["报警时间", alarm.alarm_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.alarm_time else ""],
        ["乘客人数", alarm.passenger_count],
        ["报警来源", alarm.source],
        ["位置", alarm.location or ""],
        ["描述", alarm.description or ""],
        ["当前状态", alarm.status],
        ["是否超时", "是" if alarm.is_timeout else "否"],
        ["超时原因", alarm.timeout_reason or ""],
        ["维保人员", alarm.maintenance_person or ""],
        ["维保电话", alarm.maintenance_phone or ""],
        ["派单时间", alarm.dispatched_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.dispatched_time else ""],
        ["到场时间", alarm.arrived_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.arrived_time else ""],
        ["解决时间", alarm.resolved_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.resolved_time else ""],
        ["处理结果", alarm.resolution or ""],
    ]
    
    if alarm.parent_id:
        info_data.append(["已合并到主记录ID", alarm.parent_id])
    
    info_data.extend([
        ["合并记录数", alarm.merge_count],
        ["重新提交次数", alarm.resubmit_count],
        ["审核人", alarm.reviewer or ""],
        ["审核意见", alarm.review_comment or ""],
        ["审核时间", alarm.review_time.strftime("%Y-%m-%d %H:%M:%S") if alarm.review_time else ""],
        ["创建人", alarm.created_by],
        ["创建时间", alarm.created_at.strftime("%Y-%m-%d %H:%M:%S") if alarm.created_at else ""],
    ])
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for row, (key, value) in enumerate(info_data, 1):
        key_cell = ws_info.cell(row=row, column=1, value=key)
        key_cell.fill = header_fill
        key_cell.font = header_font
        ws_info.cell(row=row, column=2, value=value)
    
    ws_info.column_dimensions['A'].width = 15
    ws_info.column_dimensions['B'].width = 50
    
    if alarm.previous_rejection:
        ws_reject = wb.create_sheet("驳回历史")
        reject_headers = ["驳回时间", "审核人", "驳回意见"]
        for col, header in enumerate(reject_headers, 1):
            cell = ws_reject.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        history = alarm.previous_rejection.get("history", [])
        for row, item in enumerate(history, 2):
            ws_reject.cell(row=row, column=1, value=item.get("reject_time", ""))
            ws_reject.cell(row=row, column=2, value=item.get("reviewer", ""))
            ws_reject.cell(row=row, column=3, value=item.get("comment", ""))
        
        ws_reject.column_dimensions['A'].width = 25
        ws_reject.column_dimensions['B'].width = 12
        ws_reject.column_dimensions['C'].width = 50
    
    if alarm.merged_alarms:
        ws_merged = wb.create_sheet("合并记录")
        merged_headers = ["合并的报警编号", "电梯编号", "报警时间", "创建人", "合并备注"]
        for col, header in enumerate(merged_headers, 1):
            cell = ws_merged.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        for row, merged in enumerate(alarm.merged_alarms, 2):
            ws_merged.cell(row=row, column=1, value=merged.alarm_no)
            ws_merged.cell(row=row, column=2, value=merged.elevator_no)
            ws_merged.cell(row=row, column=3, value=merged.alarm_time.strftime("%Y-%m-%d %H:%M:%S") if merged.alarm_time else "")
            ws_merged.cell(row=row, column=4, value=merged.created_by)
        
        ws_merged.column_dimensions['A'].width = 18
        ws_merged.column_dimensions['B'].width = 12
        ws_merged.column_dimensions['C'].width = 20
        ws_merged.column_dimensions['D'].width = 12
        ws_merged.column_dimensions['E'].width = 30
    
    ws_calls = wb.create_sheet("通话记录")
    call_headers = ["通话类型", "通话人", "联系电话", "通话时间", "时长(秒)", "通话内容", "操作员"]
    for col, header in enumerate(call_headers, 1):
        cell = ws_calls.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    for row, call in enumerate(alarm.call_records, 2):
        ws_calls.cell(row=row, column=1, value=call.call_type)
        ws_calls.cell(row=row, column=2, value=call.caller)
        ws_calls.cell(row=row, column=3, value=call.caller_phone or "")
        ws_calls.cell(row=row, column=4, value=call.call_time.strftime("%Y-%m-%d %H:%M:%S") if call.call_time else "")
        ws_calls.cell(row=row, column=5, value=call.duration or "")
        ws_calls.cell(row=row, column=6, value=call.content)
        ws_calls.cell(row=row, column=7, value=call.operator)
    
    call_widths = [12, 12, 15, 20, 10, 40, 12]
    for col, width in enumerate(call_widths, 1):
        ws_calls.column_dimensions[get_column_letter(col)].width = width
    
    ws_status = wb.create_sheet("状态流转")
    status_headers = ["原状态", "新状态", "操作员", "备注", "变更时间"]
    for col, header in enumerate(status_headers, 1):
        cell = ws_status.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    for row, history in enumerate(alarm.status_histories, 2):
        ws_status.cell(row=row, column=1, value=history.from_status or "无")
        ws_status.cell(row=row, column=2, value=history.to_status)
        ws_status.cell(row=row, column=3, value=history.operator)
        ws_status.cell(row=row, column=4, value=history.remark or "")
        ws_status.cell(row=row, column=5, value=history.change_time.strftime("%Y-%m-%d %H:%M:%S") if history.change_time else "")
    
    status_widths = [15, 15, 12, 30, 20]
    for col, width in enumerate(status_widths, 1):
        ws_status.column_dimensions[get_column_letter(col)].width = width
    
    ws_audit = wb.create_sheet("修改日志")
    audit_headers = ["修改字段", "原值", "新值", "操作员", "修改原因", "修改时间"]
    for col, header in enumerate(audit_headers, 1):
        cell = ws_audit.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    for row, audit in enumerate(alarm.audit_logs, 2):
        ws_audit.cell(row=row, column=1, value=audit.field_name)
        ws_audit.cell(row=row, column=2, value=audit.old_value or "")
        ws_audit.cell(row=row, column=3, value=audit.new_value or "")
        ws_audit.cell(row=row, column=4, value=audit.operator)
        ws_audit.cell(row=row, column=5, value=audit.change_reason or "")
        ws_audit.cell(row=row, column=6, value=audit.change_time.strftime("%Y-%m-%d %H:%M:%S") if audit.change_time else "")
    
    audit_widths = [15, 20, 20, 12, 20, 20]
    for col, width in enumerate(audit_widths, 1):
        ws_audit.column_dimensions[get_column_letter(col)].width = width
    
    return wb
