import csv
import io
import json
from datetime import datetime, date
from typing import Optional, List, Dict, Any
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from models import ShortageRecord, CompensationRecord, SettlementRecord
from schemas import ExportRequest, OperatorContext
from data_masking import DataMasking
from idempotent import IdempotentManager
from audit_service import AuditService


class ExportService:
    def __init__(self, db: Session):
        self.db = db
        self.audit_service = AuditService(db)

    def export_shortages(
        self,
        request: ExportRequest,
        operator_context: OperatorContext
    ) -> Dict[str, Any]:
        idempotent_key = request.idempotent_key or IdempotentManager.generate_key("export_shortage", request.dict())
        existing_hash = IdempotentManager.get_existing_result(self.db, idempotent_key)
        if existing_hash:
            return {"success": True, "message": "重复请求，导出已处理", "idempotent": True}

        start_dt = datetime.combine(request.start_date, datetime.min.time())
        end_dt = datetime.combine(request.end_date, datetime.max.time())

        shortages = self.db.query(ShortageRecord).filter(
            ShortageRecord.identified_at >= start_dt,
            ShortageRecord.identified_at <= end_dt
        ).all()

        data = []
        for s in shortages:
            row = {
                "缺货单号": s.shortage_no,
                "订单号": s.order.order_no if s.order else "",
                "商品ID": s.product_id,
                "商品名称": s.product_name,
                "缺货数量": s.shortage_quantity,
                "缺货金额": s.shortage_amount,
                "状态": s.status.value,
                "识别人": s.identified_by,
                "识别时间": s.identified_at.strftime("%Y-%m-%d %H:%M:%S") if s.identified_at else "",
                "确认人": s.confirmed_by or "",
                "确认时间": s.confirmed_at.strftime("%Y-%m-%d %H:%M:%S") if s.confirmed_at else "",
                "客户ID": s.order.customer_id if s.order else "",
                "客户名称": s.order.customer_name if s.order else "",
                "客户电话": s.order.customer_phone if s.order else "",
            }

            if not request.include_sensitive:
                row = DataMasking.mask_dict(row, operator_context.operator_role, False)

            data.append(row)

        filename = f"shortages_{request.start_date.strftime('%Y%m%d')}_{request.end_date.strftime('%Y%m%d')}"

        IdempotentManager.check_and_set(self.db, idempotent_key, "export_shortage")

        self.audit_service.log_operation(
            operation_type="export",
            operator_context=operator_context,
            result="allowed",
            reason=f"导出缺货记录{len(data)}条",
            reference_type="shortage",
            request_data=request.dict()
        )

        return {
            "success": True,
            "message": f"成功导出{len(data)}条记录",
            "data": data,
            "filename": filename,
            "count": len(data)
        }

    def export_compensations(
        self,
        request: ExportRequest,
        operator_context: OperatorContext
    ) -> Dict[str, Any]:
        idempotent_key = request.idempotent_key or IdempotentManager.generate_key("export_compensation", request.dict())
        existing_hash = IdempotentManager.get_existing_result(self.db, idempotent_key)
        if existing_hash:
            return {"success": True, "message": "重复请求，导出已处理", "idempotent": True}

        start_dt = datetime.combine(request.start_date, datetime.min.time())
        end_dt = datetime.combine(request.end_date, datetime.max.time())

        compensations = self.db.query(CompensationRecord).filter(
            CompensationRecord.created_at >= start_dt,
            CompensationRecord.created_at <= end_dt
        ).all()

        data = []
        for c in compensations:
            row = {
                "补偿单号": c.compensation_no,
                "缺货单号": c.shortage.shortage_no if c.shortage else "",
                "补偿类型": c.compensation_type.value,
                "退款金额": c.amount,
                "券面值": c.coupon_value,
                "券号": c.coupon_id or "",
                "换货商品ID": c.exchange_product_id or "",
                "换货商品名称": c.exchange_product_name or "",
                "状态": c.status.value,
                "操作人": c.operator_name,
                "处理时间": c.processed_at.strftime("%Y-%m-%d %H:%M:%S") if c.processed_at else "",
                "回滚人": c.rolled_back_by or "",
                "回滚时间": c.rolled_back_at.strftime("%Y-%m-%d %H:%M:%S") if c.rolled_back_at else "",
            }

            if not request.include_sensitive:
                row = DataMasking.mask_dict(row, operator_context.operator_role, False)

            data.append(row)

        filename = f"compensations_{request.start_date.strftime('%Y%m%d')}_{request.end_date.strftime('%Y%m%d')}"

        IdempotentManager.check_and_set(self.db, idempotent_key, "export_compensation")

        self.audit_service.log_operation(
            operation_type="export",
            operator_context=operator_context,
            result="allowed",
            reason=f"导出补偿记录{len(data)}条",
            reference_type="compensation",
            request_data=request.dict()
        )

        return {
            "success": True,
            "message": f"成功导出{len(data)}条记录",
            "data": data,
            "filename": filename,
            "count": len(data)
        }

    def export_settlements(
        self,
        request: ExportRequest,
        operator_context: OperatorContext
    ) -> Dict[str, Any]:
        idempotent_key = request.idempotent_key or IdempotentManager.generate_key("export_settlement", request.dict())
        existing_hash = IdempotentManager.get_existing_result(self.db, idempotent_key)
        if existing_hash:
            return {"success": True, "message": "重复请求，导出已处理", "idempotent": True}

        start_dt = datetime.combine(request.start_date, datetime.min.time())
        end_dt = datetime.combine(request.end_date, datetime.max.time())

        settlements = self.db.query(SettlementRecord).filter(
            SettlementRecord.processed_at >= start_dt,
            SettlementRecord.processed_at <= end_dt
        ).all()

        data = []
        for s in settlements:
            row = {
                "结算单号": s.settlement_no,
                "结算日期": s.settlement_date.strftime("%Y-%m-%d"),
                "缺货总数": s.total_shortage_count,
                "缺货总金额": s.total_shortage_amount,
                "退款总金额": s.total_refund_amount,
                "券总面值": s.total_coupon_value,
                "状态": s.status.value,
                "操作人": s.operator_name,
                "处理时间": s.processed_at.strftime("%Y-%m-%d %H:%M:%S") if s.processed_at else "",
            }

            if not request.include_sensitive:
                row = DataMasking.mask_dict(row, operator_context.operator_role, False)

            data.append(row)

        filename = f"settlements_{request.start_date.strftime('%Y%m%d')}_{request.end_date.strftime('%Y%m%d')}"

        IdempotentManager.check_and_set(self.db, idempotent_key, "export_settlement")

        self.audit_service.log_operation(
            operation_type="export",
            operator_context=operator_context,
            result="allowed",
            reason=f"导出结算记录{len(data)}条",
            reference_type="settlement",
            request_data=request.dict()
        )

        return {
            "success": True,
            "message": f"成功导出{len(data)}条记录",
            "data": data,
            "filename": filename,
            "count": len(data)
        }

    def to_csv(self, data: List[Dict], filename: str) -> str:
        if not data:
            return ""

        filepath = f"/tmp/{filename}.csv"
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=data[0].keys())
            writer.writeheader()
            writer.writerows(data)

        return filepath

    def to_excel(self, data: List[Dict], filename: str) -> str:
        if not data:
            return ""

        filepath = f"/tmp/{filename}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        headers = list(data[0].keys())
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_num, row_data in enumerate(data, 2):
            for col_num, key in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num, value=row_data.get(key, ""))
                cell.alignment = Alignment(horizontal="left")

        for col_num in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20

        wb.save(filepath)
        return filepath

    def export_to_file(
        self,
        request: ExportRequest,
        operator_context: OperatorContext,
        export_format: str = "xlsx"
    ) -> Dict[str, Any]:
        if request.export_type == "shortage":
            result = self.export_shortages(request, operator_context)
        elif request.export_type == "compensation":
            result = self.export_compensations(request, operator_context)
        elif request.export_type == "settlement":
            result = self.export_settlements(request, operator_context)
        else:
            return {"success": False, "message": f"不支持的导出类型: {request.export_type}"}

        if not result.get("success") or not result.get("data"):
            return result

        if export_format == "csv":
            filepath = self.to_csv(result["data"], result["filename"])
        else:
            filepath = self.to_excel(result["data"], result["filename"])

        result["filepath"] = filepath
        return result
