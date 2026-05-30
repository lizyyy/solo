import io
from typing import Optional
from datetime import date
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import BondLedger, CouponSchedule, RedemptionNotice, CustodyReceipt, FundCalendar, WarningSheet
from services.cashflow import build_cashflow_schedule
from services.receipt_match import match_receipts
from services.gap_warning import detect_gaps
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

router = APIRouter()

_header_font = Font(bold=True, size=11)
_header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_header_font_white = Font(bold=True, size=11, color="FFFFFF")
_suspect_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
_error_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
_error_font = Font(color="FFFFFF")
_thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _write_headers(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _header_font_white
        cell.fill = _header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _thin_border


def _write_row(ws, row, values, quality_col=None):
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.border = _thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if quality_col is not None:
        status = values[quality_col - 1] if quality_col <= len(values) else None
        if status == "error":
            for col in range(1, len(values) + 1):
                ws.cell(row=row, column=col).fill = _error_fill
                ws.cell(row=row, column=col).font = _error_font
        elif status == "suspect":
            for col in range(1, len(values) + 1):
                ws.cell(row=row, column=col).fill = _suspect_fill


@router.get("/cashflow-schedule")
def export_cashflow_schedule(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    schedule = build_cashflow_schedule(db, start_date, end_date)

    wb = Workbook()
    ws = wb.active
    ws.title = "现金流排程"

    headers = ["日期", "债券代码", "债券名称", "类型", "金额(万)", "来源表", "来源ID", "质量状态", "异常说明"]
    _write_headers(ws, headers)

    for i, item in enumerate(schedule.items, 2):
        notes = "; ".join(item.anomaly_notes) if item.anomaly_notes else ""
        _write_row(ws, i, [
            str(item.date), item.bond_code, item.bond_name,
            item.flow_type, item.amount, item.source_table,
            item.source_id, item.quality_status, notes,
        ], quality_col=8)

    ws2 = wb.create_sheet("缺口与预警")
    gap_headers = ["日期", "债券代码", "类型", "金额(万)", "预警类型", "预警级别", "备注"]
    _write_headers(ws2, gap_headers)
    for i, w in enumerate(schedule.warnings, 2):
        notes = "; ".join(w.get("notes", [])) if w.get("notes") else ""
        _write_row(ws2, i, [
            w.get("date", ""), w.get("bond_code", ""), w.get("flow_type", ""),
            w.get("amount", ""), w.get("warning_type", ""),
            w.get("warning_level", ""), notes,
        ])

    for col in range(1, 10):
        ws.column_dimensions[chr(64 + col)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=cashflow_schedule.xlsx"},
    )


@router.get("/warning-detail")
def export_warning_detail(db: Session = Depends(get_db)):
    warnings = db.query(WarningSheet).order_by(WarningSheet.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "预警明细"

    headers = ["预警ID", "关联债券ID", "预警类型", "预警级别", "描述", "受影响记录", "状态", "处理意见", "处理人", "处理时间", "来源类型", "来源详情"]
    _write_headers(ws, headers)

    for i, w in enumerate(warnings, 2):
        affected = str(w.affected_records) if w.affected_records else ""
        _write_row(ws, i, [
            w.id, w.bond_id or "", w.warning_type, w.warning_level,
            w.description, affected, w.status, w.resolution or "",
            w.resolved_by or "", str(w.resolved_at) if w.resolved_at else "",
            w.source_type, w.source_detail or "",
        ])

    for col in range(1, 13):
        ws.column_dimensions[chr(64 + col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=warning_detail.xlsx"},
    )


@router.get("/receipt-match")
def export_receipt_match(db: Session = Depends(get_db)):
    result = match_receipts(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "回执匹配"

    headers = ["匹配状态", "债券代码", "回执日期", "回执金额(万)", "匹配对象", "对方金额(万)", "回执编号"]
    _write_headers(ws, headers)

    row = 2
    for m in result.matched:
        _write_row(ws, row, [
            "已匹配", m.get("bond_code", ""), m.get("receipt_date", ""),
            m.get("receipt_amount", ""), m.get("matched_to", ""),
            m.get("coupon_amount", m.get("redemption_price", "")), "",
        ])
        row += 1

    for mr in result.missing_receipt:
        _write_row(ws, row, [
            "缺少回执", mr.get("bond_code", ""), mr.get("payment_date", ""),
            "", mr.get("flow_type", ""), mr.get("expected_amount", ""), "",
        ], quality_col=1)
        ws.cell(row=row, column=1).fill = _suspect_fill
        row += 1

    for ur in result.unmatched_receipt:
        _write_row(ws, row, [
            "回执无匹配", ur.get("bond_code", ""), ur.get("receipt_date", ""),
            ur.get("receipt_amount", ""), "", "", ur.get("receipt_no", ""),
        ], quality_col=1)
        ws.cell(row=row, column=1).fill = _suspect_fill
        row += 1

    for col in range(1, 8):
        ws.column_dimensions[chr(64 + col)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=receipt_match.xlsx"},
    )


@router.get("/full-report")
def export_full_report(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "总览"

    bonds = db.query(BondLedger).all()
    headers = ["债券代码", "债券名称", "发行人", "面值(万)", "票面利率", "到期日", "赎回方式", "状态", "质量状态", "来源"]
    _write_headers(ws, headers)
    for i, b in enumerate(bonds, 2):
        _write_row(ws, i, [
            b.bond_code, b.bond_name, b.issuer or "", b.face_value,
            b.coupon_rate, str(b.maturity_date), b.redemption_type,
            b.status, b.quality_status,
            f"{b.source_type}: {b.source_detail}" if b.source_detail else b.source_type,
        ], quality_col=9)

    for col in range(1, 11):
        ws.column_dimensions[chr(64 + col)].width = 18

    ws2 = wb.create_sheet("现金流排程")
    schedule = build_cashflow_schedule(db, start_date, end_date)
    cf_headers = ["日期", "债券代码", "债券名称", "类型", "金额(万)", "来源", "质量", "异常"]
    _write_headers(ws2, cf_headers)
    for i, item in enumerate(schedule.items, 2):
        notes = "; ".join(item.anomaly_notes) if item.anomaly_notes else ""
        _write_row(ws2, i, [
            str(item.date), item.bond_code, item.bond_name,
            item.flow_type, item.amount, f"{item.source_table}:{item.source_id}",
            item.quality_status, notes,
        ], quality_col=7)

    ws3 = wb.create_sheet("预警明细")
    warnings = db.query(WarningSheet).order_by(WarningSheet.created_at.desc()).all()
    w_headers = ["预警ID", "债券ID", "类型", "级别", "描述", "状态", "处理人"]
    _write_headers(ws3, w_headers)
    for i, w in enumerate(warnings, 2):
        _write_row(ws3, i, [
            w.id, w.bond_id or "", w.warning_type, w.warning_level,
            w.description, w.status, w.resolved_by or "",
        ])

    ws4 = wb.create_sheet("回执匹配")
    receipt_result = match_receipts(db)
    r_headers = ["状态", "债券代码", "日期", "金额(万)", "匹配信息", "回执编号"]
    _write_headers(ws4, r_headers)
    row = 2
    for m in receipt_result.matched:
        _write_row(ws4, row, ["已匹配", m.get("bond_code", ""), m.get("receipt_date", ""),
                               m.get("receipt_amount", ""), m.get("matched_to", ""), ""])
        row += 1
    for mr in receipt_result.missing_receipt:
        _write_row(ws4, row, ["缺少回执", mr.get("bond_code", ""), mr.get("payment_date", ""),
                               "", mr.get("flow_type", ""), ""], quality_col=1)
        row += 1
    for ur in receipt_result.unmatched_receipt:
        _write_row(ws4, row, ["无匹配", ur.get("bond_code", ""), ur.get("receipt_date", ""),
                               ur.get("receipt_amount", ""), "", ur.get("receipt_no", "")], quality_col=1)
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=full_report.xlsx"},
    )
