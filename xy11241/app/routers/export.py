from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from ..database import get_db
from .. import crud, models
from ..utils import mask_sensitive_fields, logger

router = APIRouter()


@router.get("/books/csv", summary="导出书籍列表为CSV")
def export_books_csv(
    grade: Optional[str] = None,
    condition: Optional[str] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db)
):
    books = crud.get_books(db, grade=grade, condition=condition, status=status)
    
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    
    headers = ["ID", "ISBN", "书名", "作者", "出版社", "品相", "年级", "数量", "捐赠人", "状态", "创建时间"]
    writer.writerow(headers)
    
    for book in books:
        row = [
            book.id,
            book.isbn or "",
            book.title,
            book.author or "",
            book.publisher or "",
            book.condition,
            book.grade,
            book.book_count,
            mask_sensitive_fields(book.donor_name) if book.donor_name else "",
            book.status,
            book.created_at.strftime("%Y-%m-%d %H:%M:%S") if book.created_at else ""
        ]
        writer.writerow(row)
    
    output.seek(0)
    logger.info("导出CSV文件成功")
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=books_export.csv"}
    )


@router.get("/books/excel", summary="导出书籍列表为Excel")
def export_books_excel(
    grade: Optional[str] = None,
    condition: Optional[str] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db)
):
    books = crud.get_books(db, grade=grade, condition=condition, status=status)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "书籍列表"
    
    headers = ["ID", "ISBN", "书名", "作者", "出版社", "品相", "年级", "数量", "捐赠人", "状态", "创建时间"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    for row_num, book in enumerate(books, 2):
        ws.cell(row=row_num, column=1, value=book.id)
        ws.cell(row=row_num, column=2, value=book.isbn or "")
        ws.cell(row=row_num, column=3, value=book.title)
        ws.cell(row=row_num, column=4, value=book.author or "")
        ws.cell(row=row_num, column=5, value=book.publisher or "")
        ws.cell(row=row_num, column=6, value=book.condition)
        ws.cell(row=row_num, column=7, value=book.grade)
        ws.cell(row=row_num, column=8, value=book.book_count)
        ws.cell(row=row_num, column=9, value=mask_sensitive_fields(book.donor_name) if book.donor_name else "")
        ws.cell(row=row_num, column=10, value=book.status)
        ws.cell(row=row_num, column=11, value=book.created_at.strftime("%Y-%m-%d %H:%M:%S") if book.created_at else "")
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    logger.info("导出Excel文件成功")
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=books_export.xlsx"}
    )


@router.get("/shelf-list/excel", summary="导出版清单为Excel")
def export_shelf_list_excel(
    grade: Optional[str] = None,
    condition: Optional[str] = None,
    status: str = "待上架",
    db: Session = Depends(get_db)
):
    books = crud.get_books(db, grade=grade, condition=condition, status=status)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "上架清单"
    
    ws.merge_cells('A1:I1')
    title_cell = ws.cell(row=1, column=1, value="公益书库上架清单")
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal='center')
    
    headers = ["序号", "ISBN", "书名", "作者", "品相", "适用年级", "数量", "状态", "备注"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    for idx, book in enumerate(books, 1):
        row_num = idx + 3
        ws.cell(row=row_num, column=1, value=idx)
        ws.cell(row=row_num, column=2, value=book.isbn or "无ISBN")
        ws.cell(row=row_num, column=3, value=book.title)
        ws.cell(row=row_num, column=4, value=book.author or "")
        ws.cell(row=row_num, column=5, value=book.condition)
        ws.cell(row=row_num, column=6, value=book.grade)
        ws.cell(row=row_num, column=7, value=book.book_count)
        ws.cell(row=row_num, column=8, value=book.status)
        ws.cell(row=row_num, column=9, value=book.remarks or "")
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    logger.info("导出版清单成功")
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=shelf_list.xlsx"}
    )


@router.get("/import-records/excel/{batch_no}", summary="导出导入记录为Excel")
def export_import_records_excel(batch_no: str, db: Session = Depends(get_db)):
    batch = crud.get_import_batch(db, batch_no)
    if not batch:
        return {"success": False, "message": "批次不存在"}
    
    records = crud.get_batch_records(db, batch.id)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "导入记录"
    
    ws.merge_cells('A1:H1')
    title_cell = ws.cell(row=1, column=1, value=f"批量导入记录 - {batch_no}")
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal='center')
    
    ws.cell(row=2, column=1, value=f"总计: {batch.total_count} 条")
    ws.cell(row=2, column=3, value=f"成功: {batch.success_count} 条")
    ws.cell(row=2, column=5, value=f"失败: {batch.failed_count} 条")
    ws.cell(row=2, column=7, value=f"状态: {batch.status}")
    
    headers = ["行号", "ISBN", "书名", "品相", "年级", "处理结果", "操作", "原因"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    for idx, record in enumerate(records, 1):
        row_num = idx + 4
        ws.cell(row=row_num, column=1, value=record.row_number)
        ws.cell(row=row_num, column=2, value=record.isbn or "无ISBN")
        ws.cell(row=row_num, column=3, value=record.title)
        ws.cell(row=row_num, column=4, value=record.condition or "")
        ws.cell(row=row_num, column=5, value=record.grade or "")
        ws.cell(row=row_num, column=6, value="成功" if record.is_success else "失败")
        ws.cell(row=row_num, column=7, value=record.action_taken)
        ws.cell(row=row_num, column=8, value=record.reason)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[column].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    logger.info(f"导出批次 {batch_no} 的导入记录成功")
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=import_records_{batch_no}.xlsx"}
    )
