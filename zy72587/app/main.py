from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, FileResponse
from typing import List, Optional
from datetime import datetime, timedelta
import io
import pandas as pd

from .models import (
    FeatureRecord,
    CheckParameters,
    ConflictResolution,
    CheckStep,
    RecordStatus,
)
from .workflow import workflow_manager

app = FastAPI(title="时间窗特征穿越检查系统")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/", response_class=HTMLResponse)
async def root():
    with open("app/static/index.html", "r", encoding="utf-8") as f:
        return f.read()


@app.post("/api/sessions")
async def create_session(created_by: str = "xiaomeng"):
    session = workflow_manager.create_session(created_by)
    return {"session_id": session.session_id, "message": "会话创建成功"}


@app.get("/api/sessions/{session_id}")
async def get_session(session_id: str):
    session = workflow_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    store = workflow_manager.get_result_store(session_id)
    return store.get_summary_data()


@app.post("/api/sessions/{session_id}/step1/import")
async def step1_import(
    session_id: str,
    records: List[FeatureRecord],
    parameters: Optional[CheckParameters] = None,
):
    try:
        session = workflow_manager.step1_import_bucket(session_id, records, parameters)
        store = workflow_manager.get_result_store(session_id)
        return {
            "message": "步骤1完成：线上实验桶导入成功",
            "session": store.get_summary_data(),
            "records": store.get_all_records_for_api(),
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/sessions/{session_id}/step2/negatives")
async def step2_negatives(
    session_id: str,
    records: List[FeatureRecord],
):
    try:
        session = workflow_manager.step2_review_negatives(session_id, records)
        store = workflow_manager.get_result_store(session_id)
        return {
            "message": "步骤2完成：负样本列表导入并检测完成",
            "session": store.get_summary_data(),
            "records": store.get_all_records_for_api(),
            "conflicts": store.get_conflicts_for_api(),
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/api/sessions/{session_id}/conflicts")
async def get_conflicts(session_id: str):
    session = workflow_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    store = workflow_manager.get_result_store(session_id)
    return {"conflicts": store.get_conflicts_for_api()}


@app.post("/api/sessions/{session_id}/conflicts/{record_id}/resolve")
async def resolve_conflict(
    session_id: str,
    record_id: str,
    resolution: ConflictResolution,
    resolved_by: str,
):
    try:
        session = workflow_manager.resolve_conflict(
            session_id, record_id, resolution, resolved_by
        )
        store = workflow_manager.get_result_store(session_id)
        return {
            "message": "冲突处理完成",
            "conflicts": store.get_conflicts_for_api(),
            "records": store.get_all_records_for_api(),
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/sessions/{session_id}/step3/summary")
async def step3_summary(
    session_id: str,
    summary: str,
    reviewer: Optional[str] = None,
):
    try:
        session = workflow_manager.step3_update_summary(session_id, summary, reviewer)
        store = workflow_manager.get_result_store(session_id)
        return {
            "message": "步骤3完成：摘要更新，检查流程结束",
            "session": store.get_summary_data(),
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/sessions/{session_id}/resupplement")
async def resupplement_records(
    session_id: str,
    records: List[FeatureRecord],
):
    try:
        session = workflow_manager.resupplement_and_recalculate(session_id, records)
        store = workflow_manager.get_result_store(session_id)
        return {
            "message": "补录并重算完成",
            "session": store.get_summary_data(),
            "records": store.get_all_records_for_api(),
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/api/sessions/{session_id}/records")
async def get_records(session_id: str):
    session = workflow_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    store = workflow_manager.get_result_store(session_id)
    return {"records": store.get_all_records_for_api()}


@app.get("/api/sessions/{session_id}/export")
async def export_records(session_id: str):
    session = workflow_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    store = workflow_manager.get_result_store(session_id)

    records = store.get_all_records_for_export()
    conflicts = store.get_conflicts_for_export()
    summary = store.get_summary_data()

    from .self_check import check_page_api_export_consistency
    page_data = store.get_all_records_for_page()
    api_data = store.get_all_records_for_api()
    consistency_result = check_page_api_export_consistency(page_data, api_data, records)

    df_records = pd.DataFrame(records)
    df_conflicts = pd.DataFrame(conflicts) if conflicts else pd.DataFrame(
        columns=["样本ID", "特征ID", "特征名称", "冲突类型", "记录唯一标识", "线上实验桶取值", "负样本列表取值", "冲突描述", "处理状态", "处理人(评测运营)", "处理时间", "参数版本", "冲突处理历史"]
    )

    params_df = pd.DataFrame([
        {"检查参数项": key, "取值": value, "说明": "所有sheet使用此参数版本"}
        for key, value in summary["parameters"].items()
    ])

    self_check_rows = []
    for sc in summary["self_check_results"]:
        self_check_rows.append({
            "自检项": sc["check_type"],
            "是否通过": "是" if sc["passed"] else "否",
            "发现问题数": sc["found_issues"],
            "详细说明": sc["details"],
            "自检时间": sc.get("checked_at", ""),
        })
    self_check_rows.append({
        "自检项": "页面/接口/导出三端一致性",
        "是否通过": "是" if consistency_result.passed else "否",
        "发现问题数": consistency_result.found_issues,
        "详细说明": consistency_result.details + "。（此检查在导出时实时执行，确保三端数据完全对齐，尤其特征缺失默认分记录）",
        "自检时间": datetime.now().isoformat(),
    })
    df_self_check = pd.DataFrame(self_check_rows)

    op_rows = []
    for line in summary["operation_log"]:
        op_rows.append({"操作流水": line})
    df_operation = pd.DataFrame(op_rows)

    missing_rows = []
    for idx, line in enumerate(summary.get("pending_review_details", []), 1):
        missing_rows.append({"序号": idx, "【待推荐负责人复核】特征缺失给默认分记录": line, "当前状态": "待复核，结论不能直接发"})
    for idx, line in enumerate(summary.get("feature_missing_reviewed_details", []), 1):
        missing_rows.append({"序号": len(missing_rows) + idx, "【已复核】特征缺失给默认分记录": line, "当前状态": f"已复核，复核人: {summary.get('reviewer', '')}"})
    if not missing_rows:
        missing_rows.append({"序号": 1, "特征缺失给默认分记录": "无", "当前状态": "无异常"})
    df_missing = pd.DataFrame(missing_rows)

    cover_data = [
        {"项": "文件说明", "内容": "时间窗特征穿越检查导出明细"},
        {"项": "会话ID", "内容": session_id},
        {"项": "创建人(评测运营)", "内容": session.created_by},
        {"项": "当前步骤", "内容": summary.get("current_step_text", "")},
        {"项": "推荐负责人", "内容": session.reviewer or "（待填写）特征缺失默认分记录需此角色复核"},
        {"项": "流程是否锁定", "内容": "是（完成步骤3后锁定）" if session.is_locked else "否"},
        {"项": "检查摘要", "内容": session.summary or "（步骤3后填写）"},
        {"项": "", "内容": ""},
        {"项": "【重点关注】数据一致性保证", "内容": summary["_data_consistency_note"]},
        {"项": "", "内容": ""},
        {"项": "自检总览", "内容": f"共{len(summary['self_check_results'])+1}项自检，未通过{sum(1 for r in summary['self_check_results'] if not r['passed']) + (0 if consistency_result.passed else 1)}项"},
        {"项": "待推荐负责人复核记录数", "内容": summary.get("pending_review_count", 0)},
        {"项": "已复核通过的特征缺失记录数", "内容": summary.get("feature_missing_reviewed_count", 0)},
        {"项": "冲突记录数", "内容": summary.get("conflict_count", 0)},
    ]
    df_cover = pd.DataFrame(cover_data)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_cover.to_excel(writer, index=False, sheet_name="0-文件说明与自检概览")
        df_records.to_excel(writer, index=False, sheet_name="1-检查明细(与页面/接口同源)")
        df_missing.to_excel(writer, index=False, sheet_name="2-特征缺失默认分重点")
        df_conflicts.to_excel(writer, index=False, sheet_name="3-冲突证据清单")
        df_self_check.to_excel(writer, index=False, sheet_name="4-自检结果")
        params_df.to_excel(writer, index=False, sheet_name="5-参数版本与取舍理由")
        df_operation.to_excel(writer, index=False, sheet_name="6-完整操作流水")

        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        red_font = Font(color="C00000", bold=True)
        highlight_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        for sheet in writer.sheets.values():
            for col in sheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = min(len(str(cell.value)), 80)
                    except:
                        pass
                adjusted_width = max_length + 4
                sheet.column_dimensions[column].width = adjusted_width

        ws1 = writer.sheets["1-检查明细(与页面/接口同源)"]
        for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
            default_cell = row[9]
            status_cell = row[14]
            if default_cell.value and "默认分" in str(default_cell.value):
                for cell in row:
                    cell.fill = highlight_fill

    output.seek(0)
    filename = f"时间窗特征穿越检查_{session_id}_{'步骤1完成' if session.current_step.value == 'step1_import' else '步骤2完成' if session.current_step.value == 'step2_review_negatives' else '已锁定'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{filename.encode('utf-8').decode('latin-1')}"}
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.post("/api/sessions/{session_id}/self-check/consistency")
async def check_consistency(session_id: str):
    session = workflow_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    store = workflow_manager.get_result_store(session_id)

    from .self_check import check_page_api_export_consistency

    page_data = store.get_all_records_for_page()
    api_data = store.get_all_records_for_api()
    detail_data = store.get_all_records_for_export()

    result = check_page_api_export_consistency(page_data, api_data, detail_data)
    return {"result": result.dict()}
