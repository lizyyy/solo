from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime, timedelta
from typing import List, Optional
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from database import get_db, init_db, ServiceFee, Refund, Order, Settlement
import schemas
from services import (
    LeaderService, CommissionRuleService, OrderService,
    RefundService, SettlementService, AdjustmentService,
    AuditLogService
)
import json
from datetime import datetime as dt_datetime


class DateTimeEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, dt_datetime):
            return obj.isoformat()
        return super().default(obj)


def safe_json_dumps(obj, **kwargs):
    return json.dumps(obj, cls=DateTimeEncoder, **kwargs)

app = FastAPI(title="社群分账退款冲抵佣金阶梯API", version="1.0.0")


@app.on_event("startup")
def startup_event():
    init_db()


@app.post("/leaders/", response_model=schemas.LeaderResponse, tags=["团长"])
def create_leader(leader: schemas.LeaderCreate, db: Session = Depends(get_db)):
    return LeaderService.create_leader(db, leader)


@app.get("/leaders/", response_model=List[schemas.LeaderResponse], tags=["团长"])
def get_leaders(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return LeaderService.get_all_leaders(db, skip, limit)


@app.get("/leaders/{leader_id}", response_model=schemas.LeaderResponse, tags=["团长"])
def get_leader(leader_id: int, db: Session = Depends(get_db)):
    leader = LeaderService.get_leader(db, leader_id)
    if not leader:
        raise HTTPException(status_code=404, detail="团长不存在")
    return leader


@app.post("/commission-rules/", response_model=schemas.CommissionRuleResponse, tags=["佣金规则"])
def create_commission_rule(rule: schemas.CommissionRuleCreate, db: Session = Depends(get_db)):
    return CommissionRuleService.create_rule(db, rule)


@app.post("/orders/", response_model=schemas.OrderResponse, tags=["订单"])
def create_order(order: schemas.OrderCreateRequest, db: Session = Depends(get_db)):
    original_input = safe_json_dumps(order.model_dump(), ensure_ascii=False)
    
    existing = db.query(Order).filter(Order.order_no == order.order_no).first()
    if existing:
        AuditLogService.create_log(
            db,
            action="create_order_failed",
            original_input=original_input,
            processed_by=order.processed_by,
            conclusion="失败：订单号已存在",
            settlement_id=None
        )
        raise HTTPException(status_code=400, detail="订单号已存在")
    
    leader = LeaderService.get_leader(db, order.leader_id)
    if not leader:
        AuditLogService.create_log(
            db,
            action="create_order_failed",
            original_input=original_input,
            processed_by=order.processed_by,
            conclusion="失败：团长不存在",
            settlement_id=None
        )
        raise HTTPException(status_code=400, detail="团长不存在")
    
    db_order, is_duplicate = OrderService.create_order(db, schemas.OrderCreate(**order.model_dump()))
    
    if is_duplicate or db_order.is_duplicate:
        AuditLogService.create_log(
            db,
            action="create_order_duplicate",
            original_input=original_input,
            processed_by=order.processed_by,
            conclusion=f"检测到重复订单，原始订单ID: {db_order.duplicate_of}",
            settlement_id=None
        )
    else:
        AuditLogService.create_log(
            db,
            action="create_order_success",
            original_input=original_input,
            processed_by=order.processed_by,
            conclusion=f"订单创建成功，订单ID: {db_order.id}",
            settlement_id=None
        )
    
    return db_order


@app.get("/orders/{order_id}", response_model=schemas.OrderResponse, tags=["订单"])
def get_order(order_id: int, db: Session = Depends(get_db)):
    order = OrderService.get_order(db, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    return order


@app.post("/refunds/", response_model=schemas.RefundResponse, tags=["退款"])
def create_refund(refund: schemas.RefundCreateRequest, db: Session = Depends(get_db)):
    original_input = safe_json_dumps(refund.model_dump(), ensure_ascii=False)
    
    existing = db.query(Refund).filter(Refund.refund_no == refund.refund_no).first()
    if existing:
        AuditLogService.create_log(
            db,
            action="create_refund_failed",
            original_input=original_input,
            processed_by=refund.processed_by,
            conclusion="失败：退款单号已存在",
            settlement_id=None
        )
        raise HTTPException(status_code=400, detail="退款单号已存在")
    
    order = OrderService.get_order(db, refund.order_id)
    if not order:
        AuditLogService.create_log(
            db,
            action="create_refund_failed",
            original_input=original_input,
            processed_by=refund.processed_by,
            conclusion="失败：订单不存在",
            settlement_id=None
        )
        raise HTTPException(status_code=404, detail="订单不存在")
    
    db_refund = RefundService.create_refund(db, schemas.RefundCreate(**refund.model_dump()))
    
    AuditLogService.create_log(
        db,
        action="create_refund_success",
        original_input=original_input,
        processed_by=refund.processed_by,
        conclusion=f"退款创建成功，退款ID: {db_refund.id}",
        settlement_id=None
    )
    
    return db_refund


@app.put("/refunds/{refund_id}/process", response_model=schemas.RefundResponse, tags=["退款"])
def process_refund(refund_id: int, processed_by: str, db: Session = Depends(get_db)):
    original_input = safe_json_dumps({"refund_id": refund_id, "processed_by": processed_by}, ensure_ascii=False)
    
    refund = db.query(Refund).filter(Refund.id == refund_id).first()
    if not refund:
        AuditLogService.create_log(
            db,
            action="process_refund_failed",
            original_input=original_input,
            processed_by=processed_by,
            conclusion="失败：退款记录不存在",
            settlement_id=None
        )
        raise HTTPException(status_code=404, detail="退款记录不存在")
    
    if refund.status == "processed":
        AuditLogService.create_log(
            db,
            action="process_refund_failed",
            original_input=original_input,
            processed_by=processed_by,
            conclusion="失败：退款已处理",
            settlement_id=None
        )
        raise HTTPException(status_code=400, detail="退款已处理")
    
    refund.status = "processed"
    refund.processed_at = datetime.utcnow()
    refund.processed_by = processed_by
    db.commit()
    db.refresh(refund)
    
    AuditLogService.create_log(
        db,
        action="process_refund_success",
        original_input=original_input,
        processed_by=processed_by,
        conclusion=f"退款处理成功，退款ID: {refund_id}",
        settlement_id=None
    )
    
    return refund


@app.post("/settlements/", response_model=schemas.SettlementResponse, tags=["分账结算"])
def create_settlement(settlement: schemas.SettlementCreateRequest, db: Session = Depends(get_db)):
    original_input = safe_json_dumps(settlement.model_dump(), ensure_ascii=False)
    
    leader = LeaderService.get_leader(db, settlement.leader_id)
    if not leader:
        AuditLogService.create_log(
            db,
            action="create_settlement_failed",
            original_input=original_input,
            processed_by=settlement.processed_by,
            conclusion="失败：团长不存在",
            settlement_id=None
        )
        raise HTTPException(status_code=404, detail="团长不存在")
    
    db_settlement = SettlementService.create_settlement(db, schemas.SettlementCreate(**settlement.model_dump()))
    
    AuditLogService.create_log(
        db,
        action="create_settlement_success",
        original_input=original_input,
        processed_by=settlement.processed_by,
        conclusion=f"结算单创建成功，结算ID: {db_settlement.id}",
        settlement_id=db_settlement.id
    )
    
    return db_settlement


@app.get("/settlements/{settlement_id}", response_model=schemas.SettlementResponse, tags=["分账结算"])
def get_settlement(settlement_id: int, db: Session = Depends(get_db)):
    settlement = SettlementService.get_settlement(db, settlement_id)
    if not settlement:
        raise HTTPException(status_code=404, detail="结算单不存在")
    return settlement


@app.post("/settlements/calculate", response_model=schemas.SettlementResponse, tags=["分账结算"])
def calculate_settlement(
    request: schemas.SettlementCalculateRequest,
    db: Session = Depends(get_db)
):
    original_input = safe_json_dumps(request.model_dump(), ensure_ascii=False)
    
    settlement = db.query(Settlement).filter(Settlement.id == request.settlement_id).first()
    if not settlement:
        AuditLogService.create_log(
            db,
            action="calculate_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion="失败：结算单不存在",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail="计算失败：结算单不存在")
    
    if settlement.status != "draft":
        AuditLogService.create_log(
            db,
            action="calculate_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion=f"失败：结算单状态不正确，当前状态: {settlement.status}",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail=f"计算失败：结算单状态不正确，当前状态: {settlement.status}")
    
    settlement = SettlementService.calculate_settlement(
        db, request.settlement_id, request.processed_by
    )
    
    return settlement


@app.post("/settlements/process", response_model=schemas.SettlementResponse, tags=["分账结算"])
def process_settlement(
    request: schemas.SettlementProcessRequest,
    db: Session = Depends(get_db)
):
    original_input = safe_json_dumps(request.model_dump(), ensure_ascii=False)
    
    settlement = db.query(Settlement).filter(Settlement.id == request.settlement_id).first()
    if not settlement:
        AuditLogService.create_log(
            db,
            action="process_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion="失败：结算单不存在",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail="处理失败：结算单不存在")
    
    if settlement.status != "calculated":
        AuditLogService.create_log(
            db,
            action="process_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion=f"失败：结算单状态不正确，当前状态: {settlement.status}",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail=f"处理失败：结算单状态不正确，当前状态: {settlement.status}")
    
    settlement = SettlementService.process_settlement(
        db, request.settlement_id, request.processed_by
    )
    
    return settlement


@app.post("/settlements/close", response_model=schemas.SettlementResponse, tags=["分账结算"])
def close_settlement(
    request: schemas.SettlementCloseRequest,
    db: Session = Depends(get_db)
):
    original_input = safe_json_dumps(request.model_dump(), ensure_ascii=False)
    
    settlement = db.query(Settlement).filter(Settlement.id == request.settlement_id).first()
    if not settlement:
        AuditLogService.create_log(
            db,
            action="close_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion="失败：结算单不存在",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail="关闭失败：结算单不存在")
    
    if settlement.status == "closed":
        AuditLogService.create_log(
            db,
            action="close_settlement_failed",
            original_input=original_input,
            processed_by=request.processed_by,
            conclusion="失败：结算单已关闭",
            settlement_id=request.settlement_id
        )
        raise HTTPException(status_code=400, detail="关闭失败：结算单已关闭")
    
    settlement = SettlementService.close_settlement(
        db, request.settlement_id, request.processed_by, request.close_reason
    )
    
    return settlement


@app.post("/adjustments/", response_model=schemas.AdjustmentResponse, tags=["人工修正"])
def create_adjustment(
    adjustment: schemas.AdjustmentCreate,
    db: Session = Depends(get_db)
):
    original_input = safe_json_dumps(adjustment.model_dump(), ensure_ascii=False)
    
    settlement = db.query(Settlement).filter(Settlement.id == adjustment.settlement_id).first()
    if not settlement:
        AuditLogService.create_log(
            db,
            action="create_adjustment_failed",
            original_input=original_input,
            processed_by=adjustment.processed_by,
            conclusion="失败：结算单不存在",
            settlement_id=adjustment.settlement_id
        )
        raise HTTPException(status_code=400, detail="修正失败：结算单不存在")
    
    if settlement.status not in ["calculated", "processed"]:
        AuditLogService.create_log(
            db,
            action="create_adjustment_failed",
            original_input=original_input,
            processed_by=adjustment.processed_by,
            conclusion=f"失败：结算单状态不正确，当前状态: {settlement.status}",
            settlement_id=adjustment.settlement_id
        )
        raise HTTPException(status_code=400, detail=f"修正失败：结算单状态不正确，当前状态: {settlement.status}")
    
    db_adjustment = AdjustmentService.create_adjustment(db, adjustment)
    
    return db_adjustment


@app.get("/settlements/{settlement_id}/audit-logs", response_model=List[schemas.AuditLogResponse], tags=["审计日志"])
def get_audit_logs(settlement_id: int, db: Session = Depends(get_db)):
    return AuditLogService.get_logs_by_settlement(db, settlement_id)


@app.get("/audit-logs/", response_model=List[schemas.AuditLogResponse], tags=["审计日志"])
def get_all_audit_logs(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return AuditLogService.get_all_logs(db, skip, limit)


@app.get("/settlements/{settlement_id}/export", tags=["导出报告"])
def export_settlement(settlement_id: int, db: Session = Depends(get_db)):
    settlement = SettlementService.get_settlement(db, settlement_id)
    if not settlement:
        raise HTTPException(status_code=404, detail="结算单不存在")
    
    output = io.BytesIO()
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "结算汇总"
    
    headers = ["项目", "金额", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    data = [
        ["结算单号", settlement.settlement_no, ""],
        ["团长ID", settlement.leader_id, ""],
        ["结算周期", f"{settlement.start_date} ~ {settlement.end_date}", ""],
        ["订单总金额", settlement.total_order_amount, ""],
        ["退款总金额", settlement.total_refund_amount, ""],
        ["净订单金额", settlement.net_order_amount, ""],
        ["平台服务费", settlement.service_fee, ""],
        ["应发佣金", settlement.commission_amount, ""],
        ["实发佣金", settlement.final_leader_amount, "退款冲抵后"],
        ["状态", settlement.status, ""],
    ]
    
    for row, row_data in enumerate(data, 2):
        for col, value in enumerate(row_data, 1):
            ws1.cell(row=row, column=col, value=value)
    
    ws2 = wb.create_sheet("订单明细")
    order_headers = ["订单号", "用户", "手机", "金额", "创建时间"]
    for col, header in enumerate(order_headers, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    for row, order in enumerate(settlement.orders, 2):
        ws2.cell(row=row, column=1, value=order.order_no)
        ws2.cell(row=row, column=2, value=order.user_name or "")
        ws2.cell(row=row, column=3, value=order.user_phone or "")
        ws2.cell(row=row, column=4, value=order.total_amount)
        ws2.cell(row=row, column=5, value=str(order.created_at))
    
    ws3 = wb.create_sheet("退款明细")
    refund_headers = ["退款号", "订单号", "金额", "原因", "创建时间"]
    for col, header in enumerate(refund_headers, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    for row, refund in enumerate(settlement.refunds, 2):
        ws3.cell(row=row, column=1, value=refund.refund_no)
        ws3.cell(row=row, column=2, value=refund.order.order_no)
        ws3.cell(row=row, column=3, value=refund.refund_amount)
        ws3.cell(row=row, column=4, value=refund.refund_reason or "")
        ws3.cell(row=row, column=5, value=str(refund.created_at))
    
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=settlement_{settlement.settlement_no}.xlsx"}
    )


@app.post("/service-fees/", tags=["服务费"])
def create_service_fee(fee_rate: float, min_fee: float = 0, max_fee: Optional[float] = None, db: Session = Depends(get_db)):
    db.query(ServiceFee).update({ServiceFee.is_active: False})
    db_fee = ServiceFee(fee_rate=fee_rate, min_fee=min_fee, max_fee=max_fee)
    db.add(db_fee)
    db.commit()
    db.refresh(db_fee)
    return db_fee


@app.post("/init-sample-data/", tags=["系统"])
def init_sample_data(db: Session = Depends(get_db)):
    leader1 = LeaderService.create_leader(db, schemas.LeaderCreate(
        leader_code="LEADER001",
        name="张团长",
        phone="13800138001",
        email="zhang@example.com"
    ))
    
    leader2 = LeaderService.create_leader(db, schemas.LeaderCreate(
        leader_code="LEADER002",
        name="李团长",
        phone="13800138002",
        email="li@example.com"
    ))
    
    CommissionRuleService.create_rule(db, schemas.CommissionRuleCreate(
        tier_min=0,
        tier_max=1000,
        commission_rate=0.08
    ))
    CommissionRuleService.create_rule(db, schemas.CommissionRuleCreate(
        tier_min=1000,
        tier_max=5000,
        commission_rate=0.1
    ))
    CommissionRuleService.create_rule(db, schemas.CommissionRuleCreate(
        tier_min=5000,
        tier_max=None,
        commission_rate=0.12
    ))
    
    db.query(ServiceFee).update({ServiceFee.is_active: False})
    db.add(ServiceFee(fee_rate=0.05, min_fee=10, max_fee=500))
    db.commit()
    
    now = datetime.utcnow()
    for i in range(5):
        OrderService.create_order(db, schemas.OrderCreate(
            order_no=f"ORD{now.strftime('%Y%m%d')}{i:03d}",
            leader_id=leader1.id,
            user_name=f"用户{i+1}",
            user_phone=f"1390000000{i}",
            total_amount=199.0 + i * 50,
            product_count=1 + i
        ))
    
    refund = RefundService.create_refund(db, schemas.RefundCreate(
        refund_no=f"REF{now.strftime('%Y%m%d')}001",
        order_id=1,
        refund_amount=99.0,
        refund_reason="商品质量问题"
    ))
    refund.status = "processed"
    refund.processed_at = now
    refund.processed_by = "system"
    db.commit()
    
    return {"message": "示例数据初始化完成", "leader_count": 2, "order_count": 5}
