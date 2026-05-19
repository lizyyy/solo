import csv
import json
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
from io import StringIO

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from app.models import (
    InspectionRecord,
    InspectionSummary,
    IssueType,
    ReviewStatus,
    ExportFormat,
)
from app.utils import get_storage, mask_dict, mask_object


class DataExporter:
    def __init__(self):
        self._storage = get_storage()
    
    def _mask_record(self, record: InspectionRecord) -> Dict[str, Any]:
        record_dict = record.dict()
        return mask_dict(record_dict)
    
    def export_to_csv(self, records: List[InspectionRecord], include_masked: bool = True) -> str:
        output = StringIO()
        writer = csv.writer(output)
        
        writer.writerow([
            '记录ID', '通话ID', '坐席ID', '坐席姓名', '客户电话', '客户姓名',
            '通话时间', '通话时长(秒)', '问题类型', '问题描述', '严重程度',
            '审核状态', '审核人', '审核时间', '扫描时间',
        ])
        
        for record in records:
            record_dict = self._mask_record(record) if include_masked else record.dict()
            
            if not record.issues:
                writer.writerow([
                    record.id,
                    record.call_id,
                    record_dict['metadata']['agent_id'],
                    record_dict['metadata']['agent_name'],
                    record_dict['metadata']['customer_phone'],
                    record_dict['metadata']['customer_name'],
                    record_dict['metadata']['call_start_time'],
                    record_dict['metadata']['call_duration'],
                    '', '', '', '', '', '',
                    record_dict['scanned_at'],
                ])
            else:
                for issue in record.issues:
                    writer.writerow([
                        record.id,
                        record.call_id,
                        record_dict['metadata']['agent_id'],
                        record_dict['metadata']['agent_name'],
                        record_dict['metadata']['customer_phone'],
                        record_dict['metadata']['customer_name'],
                        record_dict['metadata']['call_start_time'],
                        record_dict['metadata']['call_duration'],
                        issue.issue_type.value,
                        issue.description,
                        issue.severity,
                        issue.review_status.value,
                        issue.reviewed_by or '',
                        issue.reviewed_at or '',
                        record_dict['scanned_at'],
                    ])
        
        return output.getvalue()
    
    def export_to_excel(self, records: List[InspectionRecord], include_masked: bool = True) -> bytes:
        wb = openpyxl.Workbook()
        
        ws = wb.active
        ws.title = "质检记录"
        
        headers = [
            '记录ID', '通话ID', '坐席ID', '坐席姓名', '客户电话', '客户姓名',
            '通话时间', '通话时长(秒)', '问题类型', '问题描述', '严重程度',
            '审核状态', '审核人', '审核时间', '扫描时间',
        ]
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        row = 2
        for record in records:
            record_dict = self._mask_record(record) if include_masked else record.dict()
            
            if not record.issues:
                ws.append([
                    record.id,
                    record.call_id,
                    record_dict['metadata']['agent_id'],
                    record_dict['metadata']['agent_name'],
                    record_dict['metadata']['customer_phone'],
                    record_dict['metadata']['customer_name'],
                    str(record_dict['metadata']['call_start_time']),
                    record_dict['metadata']['call_duration'],
                    '', '', '', '', '', '',
                    str(record_dict['scanned_at']),
                ])
                row += 1
            else:
                for issue in record.issues:
                    ws.append([
                        record.id,
                        record.call_id,
                        record_dict['metadata']['agent_id'],
                        record_dict['metadata']['agent_name'],
                        record_dict['metadata']['customer_phone'],
                        record_dict['metadata']['customer_name'],
                        str(record_dict['metadata']['call_start_time']),
                        record_dict['metadata']['call_duration'],
                        issue.issue_type.value,
                        issue.description,
                        issue.severity,
                        issue.review_status.value,
                        issue.reviewed_by or '',
                        str(issue.reviewed_at) if issue.reviewed_at else '',
                        str(record_dict['scanned_at']),
                    ])
                    row += 1
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15
        
        ws2 = wb.create_sheet("坏记录")
        bad_records = self._storage.get_all_bad_records()
        
        bad_headers = ['ID', '源文件', '原始位置', '错误类型', '错误信息', '原始内容', '建议修复', '创建时间']
        for col, header in enumerate(bad_headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for br in bad_records:
            br_dict = mask_dict(br.dict()) if include_masked else br.dict()
            ws2.append([
                br_dict['id'],
                br_dict['source_file'],
                br_dict['original_position'],
                br_dict['error_type'],
                br_dict['error_message'],
                br_dict['raw_content'],
                br_dict['suggested_fix'] or '',
                str(br_dict['created_at']),
            ])
        
        for col in range(1, len(bad_headers) + 1):
            ws2.column_dimensions[chr(64 + col)].width = 18
        
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
    
    def export_to_json(self, records: List[InspectionRecord], include_masked: bool = True) -> str:
        records_data = []
        for record in records:
            record_dict = self._mask_record(record) if include_masked else record.dict()
            records_data.append(record_dict)
        
        return json.dumps(records_data, ensure_ascii=False, indent=2, default=str)
    
    def export_records(self, record_ids: List[str] = None, format: ExportFormat = ExportFormat.CSV, include_masked: bool = True):
        if record_ids:
            records = []
            for rid in record_ids:
                record = self._storage.get_record(rid)
                if record:
                    records.append(record)
        else:
            records = self._storage.get_all_records()
        
        if format == ExportFormat.CSV:
            return self.export_to_csv(records, include_masked)
        elif format == ExportFormat.EXCEL:
            return self.export_to_excel(records, include_masked)
        elif format == ExportFormat.JSON:
            return self.export_to_json(records, include_masked)
        else:
            raise ValueError(f"不支持的导出格式: {format}")
    
    def get_summary(self) -> InspectionSummary:
        records = self._storage.get_all_records()
        bad_records = self._storage.get_all_bad_records()
        
        total_issues = 0
        issues_by_type = {}
        issues_by_severity = {}
        pending_review = 0
        confirmed_issues = 0
        rejected_issues = 0
        
        for record in records:
            for issue in record.issues:
                total_issues += 1
                issue_type_key = issue.issue_type.value
                issues_by_type[issue_type_key] = issues_by_type.get(issue_type_key, 0) + 1
                issues_by_severity[issue.severity] = issues_by_severity.get(issue.severity, 0) + 1
                
                if issue.review_status == ReviewStatus.PENDING:
                    pending_review += 1
                elif issue.review_status == ReviewStatus.CONFIRMED:
                    confirmed_issues += 1
                elif issue.review_status == ReviewStatus.REJECTED:
                    rejected_issues += 1
        
        from app.services.scanner import get_scanner
        scanner = get_scanner()
        top_sensitive_words = scanner.get_sensitive_word_stats(records)
        
        scan_dates = [r.scanned_at for r in records if r.scanned_at]
        scan_date_range = None
        if scan_dates:
            scan_date_range = {
                "min": min(scan_dates),
                "max": max(scan_dates),
            }
        
        return InspectionSummary(
            total_records=len(records),
            total_issues=total_issues,
            issues_by_type=issues_by_type,
            issues_by_severity=issues_by_severity,
            pending_review=pending_review,
            confirmed_issues=confirmed_issues,
            rejected_issues=rejected_issues,
            bad_records_count=len(bad_records),
            top_sensitive_words=top_sensitive_words,
            scan_date_range=scan_date_range,
        )
    
    def save_to_file(self, content, file_path: str, format: ExportFormat):
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        
        if format == ExportFormat.EXCEL:
            with open(path, 'wb') as f:
                f.write(content)
        else:
            with open(path, 'w', encoding='utf-8') as f:
                f.write(content)
        
        return str(path.absolute())


_exporter_instance = None


def get_exporter() -> DataExporter:
    global _exporter_instance
    if _exporter_instance is None:
        _exporter_instance = DataExporter()
    return _exporter_instance
