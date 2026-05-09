import io
from datetime import datetime
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session
from sqlalchemy import func
from database import LogEntry
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


class ReportService:
    
    @staticmethod
    def generate_summary(db: Session, params: Dict[str, Any]) -> Dict[str, Any]:
        query = db.query(LogEntry)
        
        if params.get("source_id"):
            query = query.filter(LogEntry.source_id == params["source_id"])
        if params.get("log_level"):
            query = query.filter(LogEntry.log_level == params["log_level"].value)
        if params.get("start_time"):
            query = query.filter(LogEntry.log_time >= params["start_time"])
        if params.get("end_time"):
            query = query.filter(LogEntry.log_time <= params["end_time"])
        
        total = query.count()
        
        level_stats = db.query(
            LogEntry.log_level,
            func.count(LogEntry.id)
        )
        if params.get("source_id"):
            level_stats = level_stats.filter(LogEntry.source_id == params["source_id"])
        if params.get("start_time"):
            level_stats = level_stats.filter(LogEntry.log_time >= params["start_time"])
        if params.get("end_time"):
            level_stats = level_stats.filter(LogEntry.log_time <= params["end_time"])
        level_stats = level_stats.group_by(LogEntry.log_level).all()
        
        module_stats = db.query(
            LogEntry.module,
            func.count(LogEntry.id)
        )
        if params.get("source_id"):
            module_stats = module_stats.filter(LogEntry.source_id == params["source_id"])
        if params.get("log_level"):
            module_stats = module_stats.filter(LogEntry.log_level == params["log_level"].value)
        if params.get("start_time"):
            module_stats = module_stats.filter(LogEntry.log_time >= params["start_time"])
        if params.get("end_time"):
            module_stats = module_stats.filter(LogEntry.log_time <= params["end_time"])
        module_stats = module_stats.group_by(LogEntry.module).order_by(func.count(LogEntry.id).desc()).limit(10).all()
        
        error_entries = query.filter(
            LogEntry.log_level.in_(["ERROR", "CRITICAL"])
        ).order_by(LogEntry.log_time.desc()).limit(50).all()
        
        return {
            "total_logs": total,
            "level_distribution": {level: count for level, count in level_stats},
            "top_modules": {module: count for module, count in module_stats if module},
            "recent_errors": [
                {
                    "id": entry.id,
                    "log_time": entry.log_time,
                    "log_level": entry.log_level,
                    "message": entry.message[:200] if entry.message else ""
                }
                for entry in error_entries
            ]
        }
    
    @staticmethod
    def export_to_excel(db: Session, params: Dict[str, Any]) -> bytes:
        query = db.query(LogEntry)
        
        if params.get("source_id"):
            query = query.filter(LogEntry.source_id == params["source_id"])
        if params.get("log_level"):
            query = query.filter(LogEntry.log_level == params["log_level"].value)
        if params.get("start_time"):
            query = query.filter(LogEntry.log_time >= params["start_time"])
        if params.get("end_time"):
            query = query.filter(LogEntry.log_time <= params["end_time"])
        
        logs = query.order_by(LogEntry.log_time.desc()).all()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "日志明细"
        
        headers = ["ID", "日志时间", "级别", "模块", "消息", "Trace ID", "额外数据"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        for row_idx, log in enumerate(logs, 2):
            ws.cell(row=row_idx, column=1, value=log.id)
            ws.cell(row=row_idx, column=2, value=log.log_time.strftime("%Y-%m-%d %H:%M:%S") if log.log_time else "")
            ws.cell(row=row_idx, column=3, value=log.log_level)
            ws.cell(row=row_idx, column=4, value=log.module or "")
            ws.cell(row=row_idx, column=5, value=log.message or "")
            ws.cell(row=row_idx, column=6, value=log.trace_id or "")
            ws.cell(row=row_idx, column=7, value=str(log.extra_data) if log.extra_data else "")
        
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 50
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 30
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer.getvalue()
