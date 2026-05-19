from datetime import datetime
from typing import Optional
from io import BytesIO
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import models
from services import ExemptionService


class ExportService:
    @staticmethod
    def export_compatibility_report(
        db: Session,
        start_date: Optional[datetime] = None,
        end_date: Optional[datetime] = None
    ) -> bytes:
        wb = Workbook()
        
        ws1 = wb.active
        ws1.title = "失败样例汇总"
        
        headers = [
            "ID", "页面路径", "浏览器矩阵", "失败用例数", "报告人",
            "结论", "结论说明", "豁免状态", "创建时间", "更新时间"
        ]
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        query = db.query(models.FailureCase)
        if start_date:
            query = query.filter(models.FailureCase.created_at >= start_date)
        if end_date:
            query = query.filter(models.FailureCase.created_at <= end_date)
        
        failure_cases = query.all()
        
        for row_idx, failure in enumerate(failure_cases, 2):
            exemption_status = ExemptionService.get_exemption_status(db, failure.id)
            
            ws1.cell(row=row_idx, column=1, value=failure.id)
            ws1.cell(row=row_idx, column=2, value=failure.page_path)
            ws1.cell(row=row_idx, column=3, value=str(failure.browser_matrix))
            ws1.cell(row=row_idx, column=4, value=len(failure.failure_cases) if failure.failure_cases else 0)
            ws1.cell(row=row_idx, column=5, value=failure.reporter)
            ws1.cell(row=row_idx, column=6, value=failure.conclusion)
            ws1.cell(row=row_idx, column=7, value=failure.conclusion_note or "")
            ws1.cell(row=row_idx, column=8, value=exemption_status)
            ws1.cell(row=row_idx, column=9, value=failure.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws1.cell(row=row_idx, column=10, value=failure.updated_at.strftime("%Y-%m-%d %H:%M:%S") if failure.updated_at else "")
        
        for col in range(1, 11):
            ws1.column_dimensions[chr(64 + col)].width = 18
        
        ws2 = wb.create_sheet(title="豁免申请明细")
        
        exemption_headers = [
            "ID", "失败样例ID", "页面路径", "豁免原因", "豁免浏览器",
            "到期时间", "申请人", "状态", "审核结果", "审核人", "审核时间"
        ]
        
        for col, header in enumerate(exemption_headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        exemptions_query = db.query(models.Exemption)
        if start_date:
            exemptions_query = exemptions_query.filter(models.Exemption.created_at >= start_date)
        if end_date:
            exemptions_query = exemptions_query.filter(models.Exemption.created_at <= end_date)
        
        exemptions = exemptions_query.all()
        
        for row_idx, exemption in enumerate(exemptions, 2):
            page_path = ""
            if exemption.failure_case:
                page_path = exemption.failure_case.page_path
            
            is_expired = exemption.expire_at < datetime.now()
            status_display = exemption.status
            if exemption.status == "approved" and is_expired:
                status_display = "已过期"
            
            ws2.cell(row=row_idx, column=1, value=exemption.id)
            ws2.cell(row=row_idx, column=2, value=exemption.failure_id)
            ws2.cell(row=row_idx, column=3, value=page_path)
            ws2.cell(row=row_idx, column=4, value=exemption.exemption_reason)
            ws2.cell(row=row_idx, column=5, value=", ".join(exemption.exempt_browsers) if exemption.exempt_browsers else "")
            ws2.cell(row=row_idx, column=6, value=exemption.expire_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws2.cell(row=row_idx, column=7, value=exemption.applicant)
            ws2.cell(row=row_idx, column=8, value=status_display)
            ws2.cell(row=row_idx, column=9, value=exemption.review_result or "")
            ws2.cell(row=row_idx, column=10, value=exemption.reviewer or "")
            ws2.cell(row=row_idx, column=11, value=exemption.reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if exemption.reviewed_at else "")
        
        for col in range(1, 12):
            ws2.column_dimensions[chr(64 + col)].width = 18
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
