import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
import re
from models import Student, Equipment, Rental, DamageRecord, ImportLog
from app import db

class ExcelImporter:
    COLUMN_MAPPINGS = {
        'student': [
            ['姓名', '学生姓名', 'name', 'Name', '学生'],
            ['学号', '学生编号', 'student_id', 'ID', '编号'],
            ['电话', '手机号', 'phone', '联系电话'],
            ['邮箱', 'email', '电子邮箱']
        ],
        'equipment': [
            ['设备编号', '序列号', 'serial_number', '设备ID', '串号'],
            ['设备名称', '乐器名称', 'name', '名称'],
            ['类型', '设备类型', 'type', '乐器类型'],
            ['品牌', 'brand'],
            ['型号', 'model'],
            ['押金金额', '押金', 'deposit_amount', '押金标准'],
            ['日租金', 'daily_rate', '租金']
        ],
        'rental': [
            ['租赁单号', 'rental_number', '单据号', '编号'],
            ['租赁日期', '借出日期', 'rent_date', '开始日期'],
            ['应还日期', '到期日期', 'due_date', '归还日期'],
            ['实际归还日期', 'return_date', '归还日期'],
            ['已交押金', 'deposit_paid', '实交押金'],
            ['已退押金', 'deposit_refunded', '退款金额'],
            ['租金', 'rental_fee', '租赁费用'],
            ['损坏赔偿', 'damage_fee', '赔偿金额'],
            ['状态', 'status', '租赁状态']
        ],
        'damage': [
            ['损坏描述', 'description', '损坏情况'],
            ['损坏程度', 'severity', '严重程度'],
            ['维修费用', 'repair_cost'],
            ['赔偿费用', 'fee_charged', '扣费金额'],
            ['报告人', 'reported_by'],
            ['是否解决', 'resolved']
        ]
    }

    def __init__(self):
        self.warnings = []
        self.errors = []

    def detect_column_type(self, header_name):
        header_clean = str(header_name).strip().lower()
        for entity_type, column_groups in self.COLUMN_MAPPINGS.items():
            for column_group in column_groups:
                for alias in column_group:
                    if alias.lower() in header_clean or header_clean in alias.lower():
                        return entity_type, column_group[0]
        return None, None

    def clean_value(self, value, field_type=None):
        if pd.isna(value) or value is None or str(value).strip() == '':
            return None
        
        value_str = str(value).strip()
        
        if field_type == 'date':
            date_patterns = [
                '%Y-%m-%d', '%Y/%m/%d', '%Y年%m月%d日',
                '%m-%d-%Y', '%m/%d/%Y',
                '%d-%m-%Y', '%d/%m/%Y',
                '%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M:%S'
            ]
            for pattern in date_patterns:
                try:
                    return datetime.strptime(value_str, pattern)
                except ValueError:
                    continue
            try:
                return pd.to_datetime(value).to_pydatetime()
            except:
                return None
        
        if field_type == 'number':
            value_str = re.sub(r'[^\d.-]', '', value_str)
            try:
                return float(value_str)
            except ValueError:
                return None
        
        if field_type == 'boolean':
            return value_str.lower() in ['是', 'yes', 'true', '1', '有', '已解决']
        
        return value_str

    def handle_merged_cells(self, ws):
        merged_data = {}
        for merged_range in ws.merged_cells.ranges:
            top_left_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            value = top_left_cell.value
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_data[(row, col)] = value
        return merged_data

    def import_excel(self, file_path, sheet_name=None, imported_by='system'):
        self.warnings = []
        self.errors = []
        
        wb = load_workbook(filename=file_path, data_only=True)
        
        if sheet_name is None:
            sheet_name = wb.sheetnames[0]
        
        ws = wb[sheet_name]
        merged_data = self.handle_merged_cells(ws)
        
        data = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=False), 1):
            row_data = []
            for col_idx, cell in enumerate(row, 1):
                if (row_idx, col_idx) in merged_data:
                    row_data.append(merged_data[(row_idx, col_idx)])
                else:
                    row_data.append(cell.value)
            data.append(row_data)
        
        header_row = 0
        for i, row in enumerate(data[:10]):
            non_empty = sum(1 for c in row if c is not None and str(c).strip())
            if non_empty >= 3:
                header_row = i
                break
        
        headers = data[header_row]
        column_map = {}
        for col_idx, header in enumerate(headers):
            if header is None:
                continue
            entity_type, field_name = self.detect_column_type(header)
            if entity_type and field_name:
                column_map[col_idx] = (entity_type, field_name)
        
        if not column_map:
            self.errors.append("未能识别任何有效列，请检查Excel格式")
            return {'success': False, 'errors': self.errors}
        
        records_processed = 0
        records_created = 0
        records_updated = 0
        
        entity_types_present = set(et for et, _ in column_map.values())
        
        for row_idx in range(header_row + 1, len(data)):
            row = data[row_idx]
            records_processed += 1
            
            row_data = {}
            for col_idx, (entity_type, field_name) in column_map.items():
                if col_idx < len(row):
                    row_data.setdefault(entity_type, {})[field_name] = row[col_idx]
            
            if not any(row_data.values()):
                continue
            
            try:
                student = None
                if 'student' in row_data and row_data['student'].get('姓名'):
                    student_name = self.clean_value(row_data['student'].get('姓名'))
                    if student_name:
                        student_id_val = self.clean_value(row_data['student'].get('学号'))
                        student = Student.query.filter(
                            (Student.name == student_name) | 
                            (Student.student_id == student_id_val)
                        ).first()
                        
                        if not student:
                            student = Student(
                                name=student_name,
                                student_id=student_id_val,
                                phone=self.clean_value(row_data['student'].get('电话')),
                                email=self.clean_value(row_data['student'].get('邮箱'))
                            )
                            db.session.add(student)
                            records_created += 1
                        else:
                            if student_id_val and student.student_id != student_id_val:
                                student.student_id = student_id_val
                            student.phone = self.clean_value(row_data['student'].get('电话')) or student.phone
                            student.email = self.clean_value(row_data['student'].get('邮箱')) or student.email
                            records_updated += 1
                
                equipment = None
                if 'equipment' in row_data and row_data['equipment'].get('设备编号'):
                    serial_number = self.clean_value(row_data['equipment'].get('设备编号'))
                    if serial_number:
                        equipment = Equipment.query.filter_by(serial_number=serial_number).first()
                        
                        if not equipment:
                            equipment = Equipment(
                                serial_number=serial_number,
                                name=self.clean_value(row_data['equipment'].get('设备名称')) or '未知设备',
                                type=self.clean_value(row_data['equipment'].get('类型')),
                                brand=self.clean_value(row_data['equipment'].get('品牌')),
                                model=self.clean_value(row_data['equipment'].get('型号')),
                                deposit_amount=self.clean_value(row_data['equipment'].get('押金金额'), 'number') or 0,
                                daily_rate=self.clean_value(row_data['equipment'].get('日租金'), 'number') or 0
                            )
                            db.session.add(equipment)
                            records_created += 1
                        else:
                            equipment.name = self.clean_value(row_data['equipment'].get('设备名称')) or equipment.name
                            equipment.type = self.clean_value(row_data['equipment'].get('类型')) or equipment.type
                            equipment.deposit_amount = self.clean_value(row_data['equipment'].get('押金金额'), 'number') or equipment.deposit_amount
                            equipment.daily_rate = self.clean_value(row_data['equipment'].get('日租金'), 'number') or equipment.daily_rate
                            records_updated += 1
                
                if 'rental' in row_data and student and equipment:
                    rental_number = self.clean_value(row_data['rental'].get('租赁单号'))
                    rental = None
                    if rental_number:
                        rental = Rental.query.filter_by(rental_number=rental_number).first()
                    
                    if not rental:
                        rental = Rental(
                            student_id=student.id,
                            equipment_id=equipment.id,
                            rental_number=rental_number,
                            rent_date=self.clean_value(row_data['rental'].get('租赁日期'), 'date') or datetime.now(),
                            due_date=self.clean_value(row_data['rental'].get('应还日期'), 'date'),
                            return_date=self.clean_value(row_data['rental'].get('实际归还日期'), 'date'),
                            deposit_paid=self.clean_value(row_data['rental'].get('已交押金'), 'number') or 0,
                            deposit_refunded=self.clean_value(row_data['rental'].get('已退押金'), 'number') or 0,
                            rental_fee=self.clean_value(row_data['rental'].get('租金'), 'number') or 0,
                            damage_fee=self.clean_value(row_data['rental'].get('损坏赔偿'), 'number') or 0,
                            status=self.clean_value(row_data['rental'].get('状态')) or 'active'
                        )
                        db.session.add(rental)
                        records_created += 1
                    else:
                        rental.rent_date = self.clean_value(row_data['rental'].get('租赁日期'), 'date') or rental.rent_date
                        rental.due_date = self.clean_value(row_data['rental'].get('应还日期'), 'date') or rental.due_date
                        rental.return_date = self.clean_value(row_data['rental'].get('实际归还日期'), 'date') or rental.return_date
                        rental.deposit_paid = self.clean_value(row_data['rental'].get('已交押金'), 'number') or rental.deposit_paid
                        rental.deposit_refunded = self.clean_value(row_data['rental'].get('已退押金'), 'number') or rental.deposit_refunded
                        records_updated += 1
                    
                    if 'damage' in row_data and row_data['damage'].get('损坏描述'):
                        damage_desc = self.clean_value(row_data['damage'].get('损坏描述'))
                        if damage_desc:
                            damage_record = DamageRecord(
                                rental_id=rental.id,
                                equipment_id=equipment.id,
                                description=damage_desc,
                                severity=self.clean_value(row_data['damage'].get('损坏程度')),
                                repair_cost=self.clean_value(row_data['damage'].get('维修费用'), 'number') or 0,
                                fee_charged=self.clean_value(row_data['damage'].get('赔偿费用'), 'number') or 0,
                                reported_by=self.clean_value(row_data['damage'].get('报告人')),
                                resolved=self.clean_value(row_data['damage'].get('是否解决'), 'boolean') or False
                            )
                            db.session.add(damage_record)
                            records_created += 1
                
                db.session.commit()
                
            except Exception as e:
                db.session.rollback()
                self.errors.append(f"第 {row_idx + 1} 行处理失败: {str(e)}")
        
        import_log = ImportLog(
            filename=file_path.split('/')[-1],
            sheet_name=sheet_name,
            imported_by=imported_by,
            records_processed=records_processed,
            records_created=records_created,
            records_updated=records_updated,
            errors='\n'.join(self.errors) if self.errors else None,
            notes='\n'.join(self.warnings) if self.warnings else None
        )
        db.session.add(import_log)
        db.session.commit()
        
        return {
            'success': len(self.errors) < records_processed,
            'records_processed': records_processed,
            'records_created': records_created,
            'records_updated': records_updated,
            'errors': self.errors,
            'warnings': self.warnings,
            'import_log_id': import_log.id
        }
