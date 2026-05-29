import os
from datetime import datetime
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models import (
    ComparisonTask, DenoiseParams, AudioSegment,
    ProcessingResult, ListeningRecord, AnomalyRecord,
    ComparisonReport, TaskStatus
)
from ..schemas import ComparisonReportExport, TaskMetricsResponse, MetricsComparison


EXPORT_DIR = "./exports"
os.makedirs(EXPORT_DIR, exist_ok=True)


class ReportExporter:
    def __init__(self, db: Session):
        self.db = db

    def _generate_batch_no(self) -> str:
        now = datetime.now()
        return f"RPT{now.strftime('%Y%m%d')}{now.strftime('%H%M%S')}{now.strftime('%f')[:3]}"

    def _generate_report_name(
        self,
        task: ComparisonTask,
        report_type: str
    ) -> str:
        now = datetime.now()
        date_str = now.strftime("%Y%m%d")
        time_str = now.strftime("%H%M%S")
        return f"{date_str}_{task.batch_no}_{report_type}_{time_str}"

    def calculate_metrics(
        self,
        task_id: int
    ) -> TaskMetricsResponse:
        task = self.db.query(ComparisonTask).filter(
            ComparisonTask.id == task_id
        ).first()

        if task is None:
            raise ValueError(f"Task {task_id} not found")

        active_params = [p for p in task.params if p.is_active]
        voice_segments = [
            a for a in task.audio_segments
            if a.audio_type in ("original", "voice_segment")
        ]

        comparisons = []
        for params in active_params:
            results = [
                r for r in task.processing_results
                if r.params_id == params.id
            ]
            listenings = [
                l for l in task.listening_records
                if l.params_id == params.id
            ]
            anomalies = [
                a for a in task.anomalies
                if a.params_id == params.id and not a.resolved
            ]

            if not results:
                continue

            avg_snr = sum(r.snr_improvement or 0 for r in results) / len(results)
            avg_voice = sum(r.voice_preservation_rate or 0 for r in results) / len(results)
            avg_noise = sum(r.noise_reduction_rate or 0 for r in results) / len(results)

            avg_naturalness = 0
            avg_overall = 0
            if listenings:
                avg_naturalness = sum(
                    l.naturalness_score or 0 for l in listenings
                ) / len(listenings)
                avg_overall = sum(
                    l.overall_score or 0 for l in listenings
                ) / len(listenings)

            comparisons.append(MetricsComparison(
                params_id=params.id,
                params_name=params.name,
                version=params.version,
                avg_snr_improvement=round(avg_snr, 2),
                avg_voice_preservation=round(avg_voice, 1),
                avg_noise_reduction=round(avg_noise, 1),
                avg_naturalness_score=round(avg_naturalness, 1),
                avg_overall_score=round(avg_overall, 1),
                anomaly_count=len(anomalies)
            ))

        return TaskMetricsResponse(
            task_id=task.id,
            batch_no=task.batch_no,
            total_segments=len(voice_segments),
            total_params=len(active_params),
            comparisons=comparisons
        )

    def export_report(
        self,
        task_id: int,
        export_request: ComparisonReportExport
    ) -> ComparisonReport:
        task = self.db.query(ComparisonTask).filter(
            ComparisonTask.id == task_id
        ).first()

        if task is None:
            raise ValueError(f"Task {task_id} not found")

        report_name = self._generate_report_name(task, export_request.report_type)
        batch_no = self._generate_batch_no()
        file_ext = export_request.format.lower()
        file_path = os.path.join(EXPORT_DIR, f"{report_name}.{file_ext}")

        if file_ext == "xlsx":
            self._export_to_excel(task, export_request.report_type, file_path)
        else:
            raise ValueError(f"Unsupported format: {export_request.format}")

        metrics = self.calculate_metrics(task_id)
        summary = {
            "task_id": task.id,
            "task_name": task.name,
            "batch_no": task.batch_no,
            "report_batch_no": batch_no,
            "report_type": export_request.report_type,
            "total_params": metrics.total_params,
            "total_segments": metrics.total_segments,
            "total_comparisons": len(metrics.comparisons),
            "exported_by": export_request.exported_by,
            "exported_at": datetime.utcnow().isoformat(),
            "best_params": self._find_best_params(metrics)
        }

        report = ComparisonReport(
            task_id=task.id,
            batch_no=batch_no,
            report_type=export_request.report_type,
            name=report_name,
            file_path=file_path,
            format=file_ext,
            summary=summary,
            exported_by=export_request.exported_by
        )

        self.db.add(report)
        self.db.flush()

        if task.status != TaskStatus.COMPLETED:
            task.status = TaskStatus.COMPLETED
            self.db.flush()

        return report

    def _find_best_params(self, metrics: TaskMetricsResponse) -> Optional[Dict[str, Any]]:
        if not metrics.comparisons:
            return None

        def score(c: MetricsComparison) -> float:
            return (
                c.avg_voice_preservation * 0.4 +
                c.avg_naturalness_score * 0.3 +
                c.avg_noise_reduction * 0.2 +
                c.avg_overall_score * 0.1 -
                c.anomaly_count * 10
            )

        best = max(metrics.comparisons, key=score)
        return {
            "params_id": best.params_id,
            "params_name": best.params_name,
            "version": best.version,
            "composite_score": round(score(best), 2),
            "avg_voice_preservation": best.avg_voice_preservation,
            "avg_naturalness_score": best.avg_naturalness_score
        }

    def _export_to_excel(
        self,
        task: ComparisonTask,
        report_type: str,
        file_path: str
    ):
        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        info_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        def style_header(ws, row: int, max_col: int):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        def auto_width(ws, max_col: int):
            for col in range(1, max_col + 1):
                max_length = 0
                column_letter = get_column_letter(col)
                for row in ws.iter_rows(min_col=col, max_col=col):
                    for cell in row:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

        ws_summary = wb.active
        ws_summary.title = "概述"
        ws_summary.append(["音频降噪对比报告"])
        ws_summary["A1"].font = Font(bold=True, size=16)
        ws_summary.merge_cells("A1:D1")

        ws_summary.append([])
        ws_summary.append(["任务信息"])
        ws_summary.append(["任务ID", task.id])
        ws_summary.append(["批次号", task.batch_no])
        ws_summary.append(["任务名称", task.name])
        ws_summary.append(["状态", task.status])
        ws_summary.append(["创建时间", task.created_at.strftime("%Y-%m-%d %H:%M:%S")])
        ws_summary.append(["创建人", task.created_by or ""])
        ws_summary.append(["备注", task.manual_notes or ""])
        ws_summary.append([])

        metrics = self.calculate_metrics(task.id)
        ws_summary.append(["统计摘要"])
        ws_summary.append(["参数组数", metrics.total_params])
        ws_summary.append(["音频片段数", metrics.total_segments])
        ws_summary.append(["对比结果数", len(task.processing_results)])
        ws_summary.append(["试听记录数", len(task.listening_records)])
        ws_summary.append(["异常记录数", len([a for a in task.anomalies if not a.resolved])])
        ws_summary.append([])

        if metrics.comparisons:
            best = self._find_best_params(metrics)
            if best:
                ws_summary.append(["推荐参数"])
                ws_summary.append(["参数名称", best["params_name"]])
                ws_summary.append(["版本", best["version"]])
                ws_summary.append(["综合评分", best["composite_score"]])
                ws_summary.append(["人声保留率", f"{best['avg_voice_preservation']:.1f}%"])
                ws_summary.append(["自然度评分", f"{best['avg_naturalness_score']:.1f}/100"])

        auto_width(ws_summary, 4)

        if report_type in ("full", "metrics"):
            ws_metrics = wb.create_sheet("指标对比")
            headers = [
                "参数ID", "参数名称", "版本",
                "平均SNR提升(dB)", "人声保留率(%)", "降噪率(%)",
                "自然度评分", "综合评分", "异常数"
            ]
            ws_metrics.append(headers)
            style_header(ws_metrics, 1, len(headers))

            for cmp in metrics.comparisons:
                composite = (
                    cmp.avg_voice_preservation * 0.4 +
                    cmp.avg_naturalness_score * 0.3 +
                    cmp.avg_noise_reduction * 0.2 +
                    cmp.avg_overall_score * 0.1 -
                    cmp.anomaly_count * 10
                )
                row = [
                    cmp.params_id, cmp.params_name, cmp.version,
                    cmp.avg_snr_improvement, cmp.avg_voice_preservation,
                    cmp.avg_noise_reduction, cmp.avg_naturalness_score,
                    round(composite, 2), cmp.anomaly_count
                ]
                ws_metrics.append(row)
                if cmp.anomaly_count > 0:
                    for col in range(1, len(headers) + 1):
                        ws_metrics.cell(row=ws_metrics.max_row, column=col).fill = warning_fill

            auto_width(ws_metrics, len(headers))

        if report_type in ("full", "anomalies") and task.anomalies:
            ws_anomalies = wb.create_sheet("异常记录")
            headers = [
                "异常ID", "类型", "严重程度", "消息",
                "建议", "关联参数", "关联音频",
                "检测时间", "是否解决"
            ]
            ws_anomalies.append(headers)
            style_header(ws_anomalies, 1, len(headers))

            for anomaly in task.anomalies:
                params_name = next(
                    (p.name for p in task.params if p.id == anomaly.params_id), ""
                )
                audio_name = next(
                    (a.name for a in task.audio_segments if a.id == anomaly.audio_segment_id), ""
                )
                row = [
                    anomaly.id, anomaly.anomaly_type, anomaly.severity,
                    anomaly.message, anomaly.suggestion or "",
                    params_name, audio_name,
                    anomaly.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
                    "是" if anomaly.resolved else "否"
                ]
                ws_anomalies.append(row)
                if not anomaly.resolved:
                    for col in range(1, len(headers) + 1):
                        ws_anomalies.cell(row=ws_anomalies.max_row, column=col).fill = warning_fill

            auto_width(ws_anomalies, len(headers))

        if report_type in ("full", "listening") and task.listening_records:
            ws_listening = wb.create_sheet("试听记录")
            headers = [
                "记录ID", "试听人", "音频片段", "参数",
                "自然度", "降噪效果", "综合",
                "伪影", "回声", "发闷", "备注"
            ]
            ws_listening.append(headers)
            style_header(ws_listening, 1, len(headers))

            for record in task.listening_records:
                audio_name = next(
                    (a.name for a in task.audio_segments if a.id == record.audio_segment_id), ""
                )
                params_name = next(
                    (p.name for p in task.params if p.id == record.params_id), ""
                )
                row = [
                    record.id, record.listener or "", audio_name, params_name,
                    record.naturalness_score or "", record.noise_reduction_score or "",
                    record.overall_score or "",
                    "是" if record.has_artifacts else "否",
                    "是" if record.has_echo else "否",
                    "是" if record.has_muffled else "否",
                    record.comments or ""
                ]
                ws_listening.append(row)
                has_issues = any([record.has_artifacts, record.has_echo, record.has_muffled])
                if has_issues or (record.naturalness_score and record.naturalness_score < 50):
                    for col in range(1, len(headers) + 1):
                        ws_listening.cell(row=ws_listening.max_row, column=col).fill = warning_fill

            auto_width(ws_listening, len(headers))

        if report_type == "full":
            ws_results = wb.create_sheet("详细结果")
            headers = [
                "结果ID", "音频片段", "参数", "参数版本",
                "SNR前(dB)", "SNR后(dB)", "SNR提升(dB)",
                "人声保留率(%)", "降噪率(%)",
                "处理时长(s)", "输出文件"
            ]
            ws_results.append(headers)
            style_header(ws_results, 1, len(headers))

            for result in task.processing_results:
                audio_name = next(
                    (a.name for a in task.audio_segments if a.id == result.audio_segment_id), ""
                )
                params = next(
                    (p for p in task.params if p.id == result.params_id), None
                )
                params_name = params.name if params else ""
                params_version = params.version if params else ""
                row = [
                    result.id, audio_name, params_name, params_version,
                    result.snr_before, result.snr_after, result.snr_improvement,
                    result.voice_preservation_rate, result.noise_reduction_rate,
                    result.processing_time, result.output_file_path or ""
                ]
                ws_results.append(row)
                has_anomaly = any(
                    a.audio_segment_id == result.audio_segment_id and
                    a.params_id == result.params_id and not a.resolved
                    for a in task.anomalies
                )
                if has_anomaly:
                    for col in range(1, len(headers) + 1):
                        ws_results.cell(row=ws_results.max_row, column=col).fill = info_fill

            auto_width(ws_results, len(headers))

            ws_params = wb.create_sheet("参数详情")
            headers = [
                "参数ID", "名称", "版本", "是否激活",
                "降噪强度", "人声保护", "攻击程度",
                "来源", "备注"
            ]
            ws_params.append(headers)
            style_header(ws_params, 1, len(headers))

            for params in task.params:
                pj = params.params_json or {}
                row = [
                    params.id, params.name, params.version,
                    "是" if params.is_active else "否",
                    pj.get("noise_reduction", ""),
                    pj.get("voice_protection", ""),
                    pj.get("aggressiveness", ""),
                    params.source or "",
                    params.manual_notes or ""
                ]
                ws_params.append(row)

            auto_width(ws_params, len(headers))

        wb.save(file_path)
