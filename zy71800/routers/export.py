import io
from typing import Optional
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import OccupationWriteoff, WriteoffChangeLog, CreditLedger, TransactionFlow, CreditAlert
from services.detection_service import run_all_detections

router = APIRouter(prefix="/api/export", tags=["导出与检测"])


@router.get("/writeoff/excel")
def export_writeoff_excel(
    customer_id: Optional[str] = None,
    status: Optional[str] = None,
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook

    query = db.query(OccupationWriteoff)
    if customer_id:
        query = query.filter(OccupationWriteoff.customer_id == customer_id)
    if status:
        query = query.filter(OccupationWriteoff.status == status)
    results = query.order_by(OccupationWriteoff.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "备用金占用冲销"
    headers = [
        "ID", "客户ID", "客户名称", "占用金额", "冲销金额", "占用类型",
        "状态", "风险标记", "结论", "关联授信ID", "关联交易ID", "依据引用",
        "备注", "复核人", "复核日期", "版本", "创建时间", "更新时间"
    ]
    ws.append(headers)

    for w in results:
        ws.append([
            w.id, w.customer_id, w.customer_name, w.occupation_amount,
            w.writeoff_amount, w.occupation_type, w.status,
            ",".join(w.get_risk_tags()), w.conclusion, w.credit_id,
            ",".join(str(t) for t in w.get_transaction_ids()),
            str(w.get_evidence_refs()), w.remarks, w.reviewer,
            str(w.review_date or ""), w.version,
            str(w.created_at or ""), str(w.updated_at or "")
        ])

    ws2 = wb.create_sheet("变更日志")
    ws2.append(["冲销ID", "变更类型", "旧值", "新值", "变更描述", "告警级别", "操作人", "时间"])
    all_logs = []
    for w in results:
        logs = db.query(WriteoffChangeLog).filter(WriteoffChangeLog.writeoff_id == w.id).all()
        all_logs.extend(logs)
    for log in all_logs:
        ws2.append([
            log.writeoff_id, log.change_type, log.old_value, log.new_value,
            log.change_description, log.alert_level, log.operator, str(log.created_at or "")
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=writeoff_export.xlsx"},
    )


@router.get("/credit/excel")
def export_credit_excel(db: Session = Depends(get_db)):
    from openpyxl import Workbook

    results = db.query(CreditLedger).order_by(CreditLedger.created_at.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "授信台账"
    ws.append([
        "ID", "客户ID", "客户名称", "授信类型", "授信额度", "已用额度",
        "冻结额度", "可用额度", "状态", "生效日", "到期日", "来源", "备注", "创建时间"
    ])
    for c in results:
        ws.append([
            c.id, c.customer_id, c.customer_name, c.credit_type,
            c.credit_amount, c.used_amount, c.frozen_amount, c.available_amount,
            c.status, c.effective_date, c.expiry_date, c.source, c.remarks,
            str(c.created_at or "")
        ])

    ws2 = wb.create_sheet("授信告警")
    ws2.append(["ID", "授信ID", "告警类型", "告警详情", "状态", "创建时间"])
    alerts = db.query(CreditAlert).order_by(CreditAlert.created_at.desc()).all()
    for a in alerts:
        ws2.append([a.id, a.credit_id, a.alert_type, a.alert_detail, a.status, str(a.created_at or "")])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=credit_export.xlsx"},
    )


@router.get("/transaction/excel")
def export_transaction_excel(
    customer_id: Optional[str] = None,
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook

    query = db.query(TransactionFlow)
    if customer_id:
        query = query.filter(TransactionFlow.customer_id == customer_id)
    results = query.order_by(TransactionFlow.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "交易流水"
    ws.append([
        "ID", "交易编号", "客户ID", "客户名称", "交易类型", "金额",
        "占用类型", "交易日期", "关联授信ID", "版本", "是否补传", "批次号", "创建时间"
    ])
    for t in results:
        ws.append([
            t.id, t.transaction_no, t.customer_id, t.customer_name,
            t.transaction_type, t.amount, t.occupation_type, t.transaction_date,
            t.credit_id, t.version, "是" if t.is_supplementary else "否",
            t.batch_no, str(t.created_at or "")
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=transaction_export.xlsx"},
    )


@router.post("/detect/all")
def run_detections(db: Session = Depends(get_db)):
    results = run_all_detections(db)
    return {"total_alerts": len(results), "alerts": results}
