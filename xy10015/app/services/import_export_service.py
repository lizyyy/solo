import io
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional, Callable
from decimal import Decimal

from sqlalchemy.orm import Session
from fastapi import UploadFile, HTTPException, status

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.models.store import Store, Product
from app.models.inventory import Inventory
from app.models.audit import ImportExportLog
from app.services.audit_service import AuditService
from app.config import get_settings

settings = get_settings()


REQUIRED_INVENTORY_FIELDS = ["store_code", "product_sku", "quantity"]
REQUIRED_PRODUCT_FIELDS = ["sku", "name", "unit"]
REQUIRED_STORE_FIELDS = ["code", "name"]


class ImportExportService:
    @staticmethod
    def _read_excel(file: UploadFile) -> List[Dict[str, Any]]:
        try:
            content = file.file.read()
            workbook = load_workbook(io.BytesIO(content), data_only=True)
            sheet = workbook.active
            headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]

            rows = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if all(cell is None or str(cell).strip() == "" for cell in row):
                    continue
                row_data = {}
                for idx, header in enumerate(headers):
                    if header:
                        value = row[idx] if idx < len(row) else None
                        if isinstance(value, datetime):
                            value = value.strftime("%Y-%m-%d %H:%M:%S")
                        row_data[header] = value
                rows.append(row_data)
            return rows, headers
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"读取Excel文件失败: {str(e)}"
            )

    @staticmethod
    def _read_csv(file: UploadFile) -> List[Dict[str, Any]]:
        try:
            content = file.file.read().decode("utf-8")
            reader = csv.DictReader(io.StringIO(content))
            return list(reader), reader.fieldnames or []
        except UnicodeDecodeError:
            try:
                file.file.seek(0)
                content = file.file.read().decode("gbk")
                reader = csv.DictReader(io.StringIO(content))
                return list(reader), reader.fieldnames or []
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"读取CSV文件失败: {str(e)}"
                )
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"读取CSV文件失败: {str(e)}"
            )

    @staticmethod
    def _validate_headers(headers: List[str], required_fields: List[str]) -> List[str]:
        missing = [field for field in required_fields if field not in headers]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"缺少必填字段: {', '.join(missing)}"
            )
        return headers

    @staticmethod
    def import_inventory(
        db: Session,
        file: UploadFile,
        user_id: int
    ) -> Dict[str, Any]:
        file_ext = file.filename.split(".")[-1].lower() if file.filename else "xlsx"

        log = AuditService.log_import(
            db=db,
            module="inventory",
            file_name=file.filename or "unknown",
            status="processing",
            started_at=datetime.utcnow()
        )

        try:
            if file_ext in ["xlsx", "xls"]:
                rows, headers = ImportExportService._read_excel(file)
            elif file_ext == "csv":
                rows, headers = ImportExportService._read_csv(file)
            else:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="不支持的文件格式，仅支持xlsx和csv"
                )

            ImportExportService._validate_headers(headers, REQUIRED_INVENTORY_FIELDS)

            if not rows:
                raise HTTPException(status_code=400, detail="文件中没有数据")

            if len(rows) > settings.EXPORT_MAX_ROWS:
                raise HTTPException(
                    status_code=400,
                    detail=f"超出最大导入行数 {settings.EXPORT_MAX_ROWS}"
                )

            success_count = 0
            failed_count = 0
            skipped_count = 0
            errors = []

            for idx, row in enumerate(rows, start=2):
                try:
                    store_code = str(row.get("store_code", "")).strip()
                    product_sku = str(row.get("product_sku", "")).strip()
                    quantity_str = str(row.get("quantity", "")).strip()
                    cost_price_str = str(row.get("cost_price", "")).strip() or "0"
                    sale_price_str = str(row.get("sale_price", "")).strip() or "0"

                    if not store_code or not product_sku:
                        skipped_count += 1
                        errors.append({"row": idx, "error": "门店编码或商品SKU为空"})
                        continue

                    try:
                        quantity = int(quantity_str)
                        if quantity < 0:
                            raise ValueError("数量不能为负数")
                    except ValueError:
                        failed_count += 1
                        errors.append({"row": idx, "error": f"数量格式错误: {quantity_str}"})
                        continue

                    try:
                        cost_price = Decimal(cost_price_str)
                        sale_price = Decimal(sale_price_str)
                    except ValueError:
                        failed_count += 1
                        errors.append({"row": idx, "error": "价格格式错误"})
                        continue

                    store = db.query(Store).filter(
                        Store.code == store_code,
                        Store.is_deleted == False
                    ).first()
                    if not store:
                        failed_count += 1
                        errors.append({"row": idx, "error": f"门店不存在: {store_code}"})
                        continue

                    product = db.query(Product).filter(
                        Product.sku == product_sku,
                        Product.is_deleted == False
                    ).first()
                    if not product:
                        failed_count += 1
                        errors.append({"row": idx, "error": f"商品不存在: {product_sku}"})
                        continue

                    inventory = db.query(Inventory).filter(
                        Inventory.store_id == store.id,
                        Inventory.product_id == product.id,
                        Inventory.is_deleted == False
                    ).first()

                    if inventory:
                        inventory.quantity = quantity
                        inventory.available_quantity = quantity - inventory.reserved_quantity
                        if cost_price:
                            inventory.cost_price = cost_price
                        if sale_price:
                            inventory.sale_price = sale_price
                        inventory.updated_by = user_id
                    else:
                        inventory = Inventory(
                            store_id=store.id,
                            product_id=product.id,
                            quantity=quantity,
                            reserved_quantity=0,
                            available_quantity=quantity,
                            cost_price=cost_price or product.default_cost,
                            sale_price=sale_price or product.default_sale_price,
                            created_by=user_id
                        )
                        db.add(inventory)

                    success_count += 1
                except Exception as e:
                    failed_count += 1
                    errors.append({"row": idx, "error": str(e)})

            db.commit()

            final_status = "completed" if failed_count == 0 else "partial"
            AuditService.update_import_export_log(
                db=db,
                log_id=log.id,
                success_count=success_count,
                failed_count=failed_count,
                skipped_count=skipped_count,
                status=final_status,
                error_details=errors,
                completed_at=datetime.utcnow()
            )

            return {
                "success": True,
                "total": len(rows),
                "success_count": success_count,
                "failed_count": failed_count,
                "skipped_count": skipped_count,
                "errors": errors[:100]
            }

        except HTTPException:
            db.rollback()
            AuditService.update_import_export_log(
                db=db,
                log_id=log.id,
                status="failed",
                completed_at=datetime.utcnow()
            )
            raise
        except Exception as e:
            db.rollback()
            AuditService.update_import_export_log(
                db=db,
                log_id=log.id,
                status="failed",
                completed_at=datetime.utcnow()
            )
            raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")

    @staticmethod
    def export_inventory(
        db: Session,
        store_ids: Optional[List[int]] = None,
        module: str = "inventory"
    ) -> bytes:
        from app.models.store import Store, Product

        query = db.query(Inventory).filter(Inventory.is_deleted == False)
        if store_ids:
            query = query.filter(Inventory.store_id.in_(store_ids))

        inventories = query.all()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "库存数据"

        headers = ["门店编码", "门店名称", "商品SKU", "商品名称", "商品类别",
                   "单位", "库存数量", "可用数量", "成本价", "售价"]
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row_idx, inv in enumerate(inventories, start=2):
            store = db.query(Store).filter(Store.id == inv.store_id).first()
            product = db.query(Product).filter(Product.id == inv.product_id).first()

            row_data = [
                store.code if store else "",
                store.name if store else "",
                product.sku if product else "",
                product.name if product else "",
                product.category if product else "",
                product.unit if product else "",
                inv.quantity,
                inv.available_quantity,
                float(inv.cost_price) if inv.cost_price else 0,
                float(inv.sale_price) if inv.sale_price else 0
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border

        for col in sheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            sheet.column_dimensions[column].width = adjusted_width

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        return output.getvalue()

    @staticmethod
    def import_products(
        db: Session,
        file: UploadFile,
        user_id: int
    ) -> Dict[str, Any]:
        file_ext = file.filename.split(".")[-1].lower() if file.filename else "xlsx"

        log = AuditService.log_import(
            db=db,
            module="product",
            file_name=file.filename or "unknown",
            status="processing",
            started_at=datetime.utcnow()
        )

        try:
            if file_ext in ["xlsx", "xls"]:
                rows, headers = ImportExportService._read_excel(file)
            elif file_ext == "csv":
                rows, headers = ImportExportService._read_csv(file)
            else:
                raise HTTPException(status_code=400, detail="不支持的文件格式")

            ImportExportService._validate_headers(headers, REQUIRED_PRODUCT_FIELDS)

            success_count = 0
            failed_count = 0
            errors = []

            for idx, row in enumerate(rows, start=2):
                try:
                    sku = str(row.get("sku", "")).strip()
                    name = str(row.get("name", "")).strip()
                    barcode = str(row.get("barcode", "")).strip() or None
                    category = str(row.get("category", "")).strip() or None
                    unit = str(row.get("unit", "")).strip() or "件"
                    default_cost = Decimal(str(row.get("default_cost", "0")) or "0")
                    default_sale_price = Decimal(str(row.get("default_sale_price", "0")) or "0")

                    if not sku or not name:
                        failed_count += 1
                        errors.append({"row": idx, "error": "SKU或名称为空"})
                        continue

                    existing = db.query(Product).filter(Product.sku == sku, Product.is_deleted == False).first()

                    if existing:
                        existing.name = name
                        existing.barcode = barcode
                        existing.category = category
                        existing.unit = unit
                        existing.default_cost = default_cost
                        existing.default_sale_price = default_sale_price
                        existing.updated_by = user_id
                    else:
                        product = Product(
                            sku=sku,
                            name=name,
                            barcode=barcode,
                            category=category,
                            unit=unit,
                            default_cost=default_cost,
                            default_sale_price=default_sale_price,
                            created_by=user_id
                        )
                        db.add(product)

                    success_count += 1
                except Exception as e:
                    failed_count += 1
                    errors.append({"row": idx, "error": str(e)})

            db.commit()

            final_status = "completed" if failed_count == 0 else "partial"
            AuditService.update_import_export_log(
                db=db,
                log_id=log.id,
                success_count=success_count,
                failed_count=failed_count,
                status=final_status,
                error_details=errors,
                completed_at=datetime.utcnow()
            )

            return {
                "success": True,
                "success_count": success_count,
                "failed_count": failed_count,
                "errors": errors[:100]
            }

        except Exception as e:
            db.rollback()
            AuditService.update_import_export_log(
                db=db,
                log_id=log.id,
                status="failed",
                completed_at=datetime.utcnow()
            )
            raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")

    @staticmethod
    def get_import_export_template(module: str) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "导入模板"

        templates = {
            "inventory": {
                "headers": ["store_code", "product_sku", "quantity", "cost_price", "sale_price"],
                "sample": [
                    ["ST001", "SKU001", 100, 50.00, 89.00],
                    ["ST001", "SKU002", 50, 30.00, 59.00],
                    ["ST002", "SKU001", 200, 50.00, 89.00]
                ]
            },
            "product": {
                "headers": ["sku", "name", "barcode", "category", "unit", "default_cost", "default_sale_price"],
                "sample": [
                    ["SKU001", "示例商品1", "6901234567890", "食品", "件", 50.00, 89.00],
                    ["SKU002", "示例商品2", "6901234567891", "饮料", "瓶", 30.00, 59.00]
                ]
            },
            "store": {
                "headers": ["code", "name", "address", "phone", "description"],
                "sample": [
                    ["ST001", "中心店", "北京市朝阳区XXX路100号", "010-12345678", "示例门店"],
                    ["ST002", "分店", "上海市浦东新区XXX路50号", "021-87654321", "示例门店"]
                ]
            }
        }

        if module not in templates:
            raise HTTPException(status_code=400, detail="不支持的模块")

        template = templates[module]
        headers = template["headers"]
        samples = template["sample"]

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, sample_row in enumerate(samples, start=2):
            for col_idx, value in enumerate(sample_row, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=value)

        for col in sheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            sheet.column_dimensions[column].width = max_length + 5

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output.getvalue()
