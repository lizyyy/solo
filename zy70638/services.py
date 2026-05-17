from sqlalchemy.orm import Session
from datetime import datetime, timedelta
from typing import List, Optional
import database
import schemas
from exceptions import (
    ValidationException,
    InvalidStatusException,
    ManualReviewRequiredException,
    AlreadyProcessedException,
    ResourceNotFoundException,
    DuplicateOperationException
)


def get_today_str() -> str:
    return datetime.now().strftime("%Y-%m-%d")


class MemberService:
    @staticmethod
    def create_member(db: Session, member: schemas.MemberCreate) -> database.Member:
        existing_member = db.query(database.Member).filter(
            database.Member.phone == member.phone
        ).first()
        if existing_member:
            raise DuplicateOperationException(
                message="该手机号已注册会员",
                operation_type="member_registration",
                existing_id=existing_member.id
            )

        if not member.member_number:
            member.member_number = f"M{datetime.now().strftime('%Y%m%d%H%M%S')}"

        db_member = database.Member(**member.model_dump())
        db.add(db_member)
        db.commit()
        db.refresh(db_member)
        return db_member

    @staticmethod
    def get_member(db: Session, member_id: int) -> Optional[database.Member]:
        return db.query(database.Member).filter(database.Member.id == member_id).first()

    @staticmethod
    def get_member_by_phone(db: Session, phone: str) -> Optional[database.Member]:
        return db.query(database.Member).filter(database.Member.phone == phone).first()

    @staticmethod
    def update_member(db: Session, member_id: int, member_update: schemas.MemberUpdate) -> database.Member:
        db_member = MemberService.get_member(db, member_id)
        if not db_member:
            raise ResourceNotFoundException("Member", member_id)

        for field, value in member_update.model_dump(exclude_unset=True).items():
            setattr(db_member, field, value)

        db.commit()
        db.refresh(db_member)
        return db_member


class StationService:
    @staticmethod
    def create_station(db: Session, station: schemas.StationCreate) -> database.Station:
        existing = db.query(database.Station).filter(database.Station.name == station.name).first()
        if existing:
            raise DuplicateOperationException(
                message="该工位名称已存在",
                operation_type="station_creation",
                existing_id=existing.id
            )

        db_station = database.Station(**station.model_dump())
        db.add(db_station)
        db.commit()
        db.refresh(db_station)
        return db_station

    @staticmethod
    def get_station(db: Session, station_id: int) -> Optional[database.Station]:
        return db.query(database.Station).filter(database.Station.id == station_id).first()

    @staticmethod
    def get_all_stations(db: Session, only_active: bool = True) -> List[database.Station]:
        query = db.query(database.Station)
        if only_active:
            query = query.filter(database.Station.is_active == True)
        return query.order_by(database.Station.id).all()

    @staticmethod
    def get_available_stations(db: Session, station_type: str = None) -> List[database.Station]:
        query = db.query(database.Station).filter(
            database.Station.is_active == True,
            database.Station.status == "空闲"
        )
        if station_type:
            query = query.filter(database.Station.station_type == station_type)
        return query.all()


class QueueService:
    @staticmethod
    def generate_queue_number(db: Session, queue_date: str) -> tuple:
        last_queue = db.query(database.QueueNumber).filter(
            database.QueueNumber.queue_date == queue_date
        ).order_by(database.QueueNumber.sequence_number.desc()).first()

        sequence_number = last_queue.sequence_number + 1 if last_queue else 1
        display_number = f"A{sequence_number:03d}"
        return sequence_number, display_number

    @staticmethod
    def check_duplicate_queue(db: Session, member_id: int, queue_date: str) -> bool:
        existing = db.query(database.QueueNumber).filter(
            database.QueueNumber.member_id == member_id,
            database.QueueNumber.queue_date == queue_date,
            database.QueueNumber.status.in_(["等待中", "已叫号", "服务中"])
        ).first()
        return existing is not None

    @staticmethod
    def create_queue_number(
        db: Session,
        queue_create: schemas.QueueNumberCreate
    ) -> database.QueueNumber:
        member = MemberService.get_member(db, queue_create.member_id)
        if not member:
            raise ResourceNotFoundException("Member", queue_create.member_id)

        if not member.is_active:
            raise InvalidStatusException(
                message="会员已被禁用，无法取号",
                current_status="inactive"
            )

        queue_date = get_today_str()

        if QueueService.check_duplicate_queue(db, queue_create.member_id, queue_date):
            raise DuplicateOperationException(
                message="该会员今日已有有效排队号，请勿重复取号",
                operation_type="duplicate_queue"
            )

        sequence_number, display_number = QueueService.generate_queue_number(db, queue_date)

        priority = 0
        if member.membership_level == "VIP":
            priority = 2
        elif member.membership_level == "高级":
            priority = 1

        db_queue = database.QueueNumber(
            **queue_create.model_dump(),
            queue_date=queue_date,
            sequence_number=sequence_number,
            display_number=display_number,
            priority=priority
        )
        db.add(db_queue)
        db.commit()
        db.refresh(db_queue)
        return db_queue

    @staticmethod
    def get_queue_number(db: Session, queue_id: int) -> Optional[database.QueueNumber]:
        return db.query(database.QueueNumber).filter(database.QueueNumber.id == queue_id).first()

    @staticmethod
    def get_waiting_queues(db: Session, queue_date: str = None) -> List[database.QueueNumber]:
        if queue_date is None:
            queue_date = get_today_str()

        return db.query(database.QueueNumber).filter(
            database.QueueNumber.queue_date == queue_date,
            database.QueueNumber.status == "等待中"
        ).order_by(
            database.QueueNumber.priority.desc(),
            database.QueueNumber.sequence_number
        ).all()

    @staticmethod
    def call_next_queue(db: Session, station_id: int, operator: str = None) -> database.QueueNumber:
        station = StationService.get_station(db, station_id)
        if not station:
            raise ResourceNotFoundException("Station", station_id)

        if station.status != "空闲":
            raise InvalidStatusException(
                message="该工车位当前不可用",
                current_status=station.status,
                expected_status="空闲"
            )

        waiting_queues = QueueService.get_waiting_queues(db)
        if not waiting_queues:
            raise ResourceNotFoundException("当前没有等待的排队号")

        next_queue = waiting_queues[0]

        next_queue.status = "已叫号"
        next_queue.assigned_station_id = station_id
        next_queue.called_at = datetime.utcnow()

        station.status = "服务中"
        station.current_queue_id = next_queue.id

        db.commit()
        db.refresh(next_queue)
        return next_queue

    @staticmethod
    def start_service(db: Session, queue_id: int) -> database.QueueNumber:
        queue = QueueService.get_queue_number(db, queue_id)
        if not queue:
            raise ResourceNotFoundException("QueueNumber", queue_id)

        if queue.status != "已叫号":
            raise InvalidStatusException(
                message="当前排队状态不允许开始服务",
                current_status=queue.status,
                expected_status="已叫号"
            )

        queue.status = "服务中"
        queue.started_at = datetime.utcnow()

        db.commit()
        db.refresh(queue)
        return queue

    @staticmethod
    def complete_service(db: Session, queue_id: int) -> database.QueueNumber:
        queue = QueueService.get_queue_number(db, queue_id)
        if not queue:
            raise ResourceNotFoundException("QueueNumber", queue_id)

        if queue.status == "已完成":
            raise AlreadyProcessedException(
                message="该排队号已完成服务",
                processed_at=queue.completed_at.isoformat() if queue.completed_at else None
            )

        if queue.status != "服务中":
            raise InvalidStatusException(
                message="当前排队状态不允许完成服务",
                current_status=queue.status,
                expected_status="服务中"
            )

        queue.status = "已完成"
        queue.completed_at = datetime.utcnow()

        if queue.assigned_station_id:
            station = StationService.get_station(db, queue.assigned_station_id)
            if station:
                station.status = "空闲"
                station.current_queue_id = None

        db.commit()
        db.refresh(queue)
        return queue

    @staticmethod
    def mark_as_passed(
        db: Session,
        queue_id: int,
        reason: str = None,
        notes: str = None,
        operator: str = None
    ) -> database.PassedRecord:
        queue = QueueService.get_queue_number(db, queue_id)
        if not queue:
            raise ResourceNotFoundException("QueueNumber", queue_id)

        if queue.status == "过号":
            raise AlreadyProcessedException(
                message="该排队号已标记为过号"
            )

        if queue.status not in ["已叫号", "服务中"]:
            raise InvalidStatusException(
                message="当前排队状态不允许标记过号",
                current_status=queue.status
            )

        queue.status = "过号"

        passed_record = database.PassedRecord(
            queue_number_id=queue_id,
            reason=reason,
            notes=notes,
            operator=operator
        )
        db.add(passed_record)

        if queue.assigned_station_id:
            station = StationService.get_station(db, queue.assigned_station_id)
            if station:
                station.status = "空闲"
                station.current_queue_id = None

        db.commit()
        db.refresh(passed_record)
        return passed_record

    @staticmethod
    def requeue_passed(
        db: Session,
        passed_record_id: int,
        operator: str = None
    ) -> database.QueueNumber:
        passed_record = db.query(database.PassedRecord).filter(
            database.PassedRecord.id == passed_record_id
        ).first()

        if not passed_record:
            raise ResourceNotFoundException("PassedRecord", passed_record_id)

        if passed_record.new_queue_number_id:
            raise AlreadyProcessedException(
                message="该过号记录已重新排队",
                processed_at=passed_record.requeued_at.isoformat() if passed_record.requeued_at else None
            )

        original_queue = QueueService.get_queue_number(db, passed_record.queue_number_id)
        if not original_queue:
            raise ResourceNotFoundException("Original QueueNumber", passed_record.queue_number_id)

        queue_date = get_today_str()
        sequence_number, display_number = QueueService.generate_queue_number(db, queue_date)

        new_queue = database.QueueNumber(
            member_id=original_queue.member_id,
            service_type=original_queue.service_type,
            queue_date=queue_date,
            sequence_number=sequence_number,
            display_number=display_number,
            priority=original_queue.priority + 3,
            is_appointment=original_queue.is_appointment,
            appointment_id=original_queue.appointment_id,
            status="等待中"
        )
        db.add(new_queue)
        db.flush()

        passed_record.requeued_at = datetime.utcnow()
        passed_record.new_queue_number_id = new_queue.id
        if operator:
            passed_record.operator = operator

        db.commit()
        db.refresh(new_queue)
        return new_queue


class AppointmentService:
    @staticmethod
    def create_appointment(
        db: Session,
        appointment: schemas.AppointmentCreate
    ) -> database.Appointment:
        member = MemberService.get_member(db, appointment.member_id)
        if not member:
            raise ResourceNotFoundException("Member", appointment.member_id)

        appointment_time = appointment.appointment_time
        appointment_date = appointment_time.date().isoformat()

        existing = db.query(database.Appointment).filter(
            database.Appointment.member_id == appointment.member_id,
            database.Appointment.status.in_(["待确认", "已确认"]),
            database.Appointment.appointment_time >= appointment_time - timedelta(hours=1),
            database.Appointment.appointment_time <= appointment_time + timedelta(hours=1)
        ).first()

        if existing:
            raise DuplicateOperationException(
                message="该会员在同一时段已有预约",
                operation_type="appointment_booking",
                existing_id=existing.id
            )

        available_stations = StationService.get_available_stations(db)
        booked_station_id = None

        for station in available_stations:
            conflict = db.query(database.Appointment).filter(
                database.Appointment.booked_station_id == station.id,
                database.Appointment.status.in_(["待确认", "已确认"]),
                database.Appointment.appointment_time >= appointment_time - timedelta(minutes=30),
                database.Appointment.appointment_time <= appointment_time + timedelta(minutes=30)
            ).first()

            if not conflict:
                booked_station_id = station.id
                break

        db_appointment = database.Appointment(
            **appointment.model_dump(),
            booked_station_id=booked_station_id,
            status="已确认" if booked_station_id else "待确认"
        )
        db.add(db_appointment)
        db.commit()
        db.refresh(db_appointment)
        return db_appointment

    @staticmethod
    def get_appointment(db: Session, appointment_id: int) -> Optional[database.Appointment]:
        return db.query(database.Appointment).filter(database.Appointment.id == appointment_id).first()

    @staticmethod
    def cancel_appointment(db: Session, appointment_id: int) -> database.Appointment:
        appointment = AppointmentService.get_appointment(db, appointment_id)
        if not appointment:
            raise ResourceNotFoundException("Appointment", appointment_id)

        if appointment.status == "已取消":
            raise AlreadyProcessedException(
                message="该预约已取消"
            )

        if appointment.status == "已完成":
            raise InvalidStatusException(
                message="已完成的预约无法取消",
                current_status=appointment.status
            )

        appointment.status = "已取消"
        db.commit()
        db.refresh(appointment)
        return appointment


class ReportService:
    @staticmethod
    def generate_daily_report(db: Session, report_date: str) -> database.QueueReport:
        queues = db.query(database.QueueNumber).filter(
            database.QueueNumber.queue_date == report_date
        ).all()

        if not queues:
            raise ValidationException(
                message="该日期没有排队数据",
                field="report_date"
            )

        total_queues = len(queues)
        total_completed = sum(1 for q in queues if q.status == "已完成")
        total_passed = sum(1 for q in queues if q.status == "过号")

        wait_times = []
        service_times = []
        hour_counts = {}

        for q in queues:
            if q.called_at and q.created_at:
                wait_seconds = (q.called_at - q.created_at).total_seconds()
                wait_times.append(wait_seconds)

            if q.completed_at and q.started_at:
                service_seconds = (q.completed_at - q.started_at).total_seconds()
                service_times.append(service_seconds)

            if q.created_at:
                hour = q.created_at.hour
                hour_counts[hour] = hour_counts.get(hour, 0) + 1

        avg_wait_time = sum(wait_times) / len(wait_times) / 60 if wait_times else 0
        avg_service_time = sum(service_times) / len(service_times) / 60 if service_times else 0

        peak_hour = None
        if hour_counts:
            peak_hour_num = max(hour_counts.keys(), key=lambda h: hour_counts[h])
            peak_hour = f"{peak_hour_num}:00-{peak_hour_num + 1}:00"

        existing_report = db.query(database.QueueReport).filter(
            database.QueueReport.report_date == report_date
        ).first()

        if existing_report:
            existing_report.total_queues = total_queues
            existing_report.total_completed = total_completed
            existing_report.total_passed = total_passed
            existing_report.avg_wait_time = round(avg_wait_time, 2)
            existing_report.avg_service_time = round(avg_service_time, 2)
            existing_report.peak_hour = peak_hour
            db.commit()
            db.refresh(existing_report)
            return existing_report

        report = database.QueueReport(
            report_date=report_date,
            total_queues=total_queues,
            total_completed=total_completed,
            total_passed=total_passed,
            avg_wait_time=round(avg_wait_time, 2),
            avg_service_time=round(avg_service_time, 2),
            peak_hour=peak_hour
        )
        db.add(report)
        db.commit()
        db.refresh(report)
        return report

    @staticmethod
    def export_to_excel(db: Session, start_date: str, end_date: str, file_path: str):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "排队记录"

        headers = ["日期", "排队号", "会员", "服务类型", "状态", "优先级", "创建时间", "叫号时间", "完成时间", "等待时间(分钟)", "服务时间(分钟)"]

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        queues = db.query(database.QueueNumber).filter(
            database.QueueNumber.queue_date >= start_date,
            database.QueueNumber.queue_date <= end_date
        ).order_by(database.QueueNumber.queue_date, database.QueueNumber.sequence_number).all()

        for row, queue in enumerate(queues, 2):
            wait_time = None
            if queue.called_at and queue.created_at:
                wait_time = round((queue.called_at - queue.created_at).total_seconds() / 60, 1)

            service_time = None
            if queue.completed_at and queue.started_at:
                service_time = round((queue.completed_at - queue.started_at).total_seconds() / 60, 1)

            ws.cell(row=row, column=1, value=queue.queue_date)
            ws.cell(row=row, column=2, value=queue.display_number)
            ws.cell(row=row, column=3, value=queue.member.name if queue.member else "")
            ws.cell(row=row, column=4, value=queue.service_type)
            ws.cell(row=row, column=5, value=queue.status)
            ws.cell(row=row, column=6, value=queue.priority)
            ws.cell(row=row, column=7, value=queue.created_at.strftime("%Y-%m-%d %H:%M:%S") if queue.created_at else "")
            ws.cell(row=row, column=8, value=queue.called_at.strftime("%Y-%m-%d %H:%M:%S") if queue.called_at else "")
            ws.cell(row=row, column=9, value=queue.completed_at.strftime("%Y-%m-%d %H:%M:%S") if queue.completed_at else "")
            ws.cell(row=row, column=10, value=wait_time)
            ws.cell(row=row, column=11, value=service_time)

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 18

        wb.save(file_path)
        return file_path
