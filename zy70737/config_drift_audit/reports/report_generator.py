import csv
import json
from pathlib import Path
from typing import Optional
from datetime import datetime

from ..models import AuditResult, ReviewStatus, SourceLocation

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ReportGenerator:
    def __init__(self, output_dir: Optional[str] = None):
        if output_dir is None:
            output_dir = Path.cwd() / "audit_reports"
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_all(self, audit_result: AuditResult, prefix: str = "drift_audit") -> str:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = f"{prefix}_{timestamp}"

        csv_path = self.generate_csv(audit_result, f"{base_name}.csv")
        json_path = self.generate_json(audit_result, f"{base_name}.json")
        summary_path = self.generate_summary_txt(audit_result, f"{base_name}_summary.txt")

        if HAS_OPENPYXL:
            excel_path = self.generate_excel(audit_result, f"{base_name}.xlsx")
            return f"Reports generated: {csv_path}, {json_path}, {summary_path}, {excel_path}"

        return f"Reports generated: {csv_path}, {json_path}, {summary_path}"

    def generate_csv(self, audit_result: AuditResult, filename: str) -> str:
        output_path = self.output_dir / filename

        with open(output_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                'service_name',
                'config_key',
                'expected_value',
                'actual_value',
                'has_exemption',
                'exemption_reason',
                'exemption_expire_date',
                'review_status',
                'is_exemption_expired',
                'source_location',
            ])

            for record in audit_result.drift_records:
                location = audit_result.source_tracker.get_location(record.record_id)
                writer.writerow([
                    record.config_item.service_name,
                    record.config_item.config_key,
                    record.config_item.expected_value,
                    record.config_item.actual_value,
                    'Yes' if record.exemption else 'No',
                    record.exemption.reason if record.exemption else '',
                    record.exemption.expire_date if record.exemption else '',
                    record.review_status.value,
                    'Yes' if record.is_exemption_expired else 'No',
                    str(location) if location else '',
                ])

            if audit_result.source_tracker.bad_rows:
                writer.writerow([])
                writer.writerow(['BAD_ROWS', '', '', '', '', '', '', '', '', ''])
                writer.writerow([
                    'file_path',
                    'sheet_name',
                    'row_number',
                    'raw_content',
                ])
                for bad_row in audit_result.source_tracker.bad_rows:
                    writer.writerow([
                        bad_row.file_path,
                        bad_row.sheet_name or '',
                        bad_row.row_number,
                        bad_row.raw_content,
                    ])

        return str(output_path)

    def generate_json(self, audit_result: AuditResult, filename: str) -> str:
        output_path = self.output_dir / filename

        data = {
            'run_id': audit_result.run_id,
            'audit_time': audit_result.summary.audit_time,
            'summary': {
                'total_records': audit_result.summary.total_records,
                'drifted_records': audit_result.summary.drifted_records,
                'exempted_records': audit_result.summary.exempted_records,
                'expired_exemptions': audit_result.summary.expired_exemptions,
                'pending_review': audit_result.summary.pending_review,
                'approved_exemptions': audit_result.summary.approved_exemptions,
                'rejected_exemptions': audit_result.summary.rejected_exemptions,
                'no_exemption': audit_result.summary.no_exemption,
                'bad_rows_count': audit_result.summary.bad_rows_count,
            },
            'drift_records': [
                {
                    'service_name': r.config_item.service_name,
                    'config_key': r.config_item.config_key,
                    'expected_value': r.config_item.expected_value,
                    'actual_value': r.config_item.actual_value,
                    'exemption': {
                        'reason': r.exemption.reason,
                        'expire_date': r.exemption.expire_date,
                        'reviewer': r.exemption.reviewer,
                    } if r.exemption else None,
                    'review_status': r.review_status.value,
                    'is_exemption_expired': r.is_exemption_expired,
                    'source_location': str(audit_result.source_tracker.get_location(r.record_id)),
                }
                for r in audit_result.drift_records
            ],
            'bad_rows': [
                {
                    'file_path': br.file_path,
                    'sheet_name': br.sheet_name,
                    'row_number': br.row_number,
                    'raw_content': br.raw_content,
                }
                for br in audit_result.source_tracker.bad_rows
            ],
        }

        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

        return str(output_path)

    def generate_summary_txt(self, audit_result: AuditResult, filename: str) -> str:
        output_path = self.output_dir / filename

        lines = [
            "=" * 60,
            "配置漂移豁免到期复核排查报告",
            "=" * 60,
            "",
            f"运行ID: {audit_result.run_id}",
            f"审计时间: {audit_result.summary.audit_time}",
            "",
            "-" * 60,
            "统计摘要",
            "-" * 60,
            f"总记录数: {audit_result.summary.total_records}",
            f"漂移记录数: {audit_result.summary.drifted_records}",
            f"已豁免记录数: {audit_result.summary.exempted_records}",
            f"豁免已过期数: {audit_result.summary.expired_exemptions}",
            f"待复核数: {audit_result.summary.pending_review}",
            f"已批准数: {audit_result.summary.approved_exemptions}",
            f"已拒绝数: {audit_result.summary.rejected_exemptions}",
            f"无豁免数: {audit_result.summary.no_exemption}",
            f"坏行数: {audit_result.summary.bad_rows_count}",
            "",
        ]

        expired = audit_result.get_expired_exemptions()
        if expired:
            lines.extend([
                "-" * 60,
                "豁免已过期记录 (需重点关注)",
                "-" * 60,
                "",
            ])
            for r in expired:
                loc = audit_result.source_tracker.get_location(r.record_id)
                lines.extend([
                    f"服务: {r.config_item.service_name}",
                    f"配置项: {r.config_item.config_key}",
                    f"期望值: {r.config_item.expected_value}",
                    f"实际值: {r.config_item.actual_value}",
                    f"豁免理由: {r.exemption.reason if r.exemption else 'N/A'}",
                    f"豁免到期日: {r.exemption.expire_date if r.exemption else 'N/A'}",
                    f"来源: {str(loc) if loc else 'N/A'}",
                    "",
                ])

        pending = audit_result.get_records_by_status(ReviewStatus.PENDING)
        if pending:
            lines.extend([
                "-" * 60,
                "待复核记录",
                "-" * 60,
                "",
            ])
            for r in pending:
                loc = audit_result.source_tracker.get_location(r.record_id)
                lines.extend([
                    f"服务: {r.config_item.service_name}",
                    f"配置项: {r.config_item.config_key}",
                    f"期望值: {r.config_item.expected_value}",
                    f"实际值: {r.config_item.actual_value}",
                    f"豁免理由: {r.exemption.reason if r.exemption else 'N/A'}",
                    f"来源: {str(loc) if loc else 'N/A'}",
                    "",
                ])

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))

        return str(output_path)

    def generate_excel(self, audit_result: AuditResult, filename: str) -> str:
        if not HAS_OPENPYXL:
            raise ImportError(
                "openpyxl is required for Excel reports. "
                "Install it with: pip install openpyxl"
            )

        output_path = self.output_dir / filename
        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = "摘要"
        self._fill_summary_sheet(ws_summary, audit_result)

        ws_drift = wb.create_sheet("漂移详情")
        self._fill_drift_sheet(ws_drift, audit_result)

        ws_expired = wb.create_sheet("已过期豁免")
        self._fill_expired_sheet(ws_expired, audit_result)

        if audit_result.source_tracker.bad_rows:
            ws_bad = wb.create_sheet("坏行记录")
            self._fill_bad_rows_sheet(ws_bad, audit_result)

        wb.save(output_path)
        return str(output_path)

    def _fill_summary_sheet(self, ws, audit_result: AuditResult) -> None:
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        ws['A1'] = "配置漂移豁免到期复核排查报告"
        ws['A1'].font = Font(bold=True, size=14)

        ws['A3'] = "运行ID"
        ws['B3'] = audit_result.run_id
        ws['A4'] = "审计时间"
        ws['B4'] = audit_result.summary.audit_time

        ws['A6'] = "统计项"
        ws['B6'] = "数值"
        ws['A6'].fill = header_fill
        ws['B6'].fill = header_fill
        ws['A6'].font = header_font
        ws['B6'].font = header_font

        summary_items = [
            ("总记录数", audit_result.summary.total_records),
            ("漂移记录数", audit_result.summary.drifted_records),
            ("已豁免记录数", audit_result.summary.exempted_records),
            ("豁免已过期数", audit_result.summary.expired_exemptions),
            ("待复核数", audit_result.summary.pending_review),
            ("已批准数", audit_result.summary.approved_exemptions),
            ("已拒绝数", audit_result.summary.rejected_exemptions),
            ("无豁免数", audit_result.summary.no_exemption),
            ("坏行数", audit_result.summary.bad_rows_count),
        ]

        for i, (label, value) in enumerate(summary_items, start=7):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 20

    def _fill_drift_sheet(self, ws, audit_result: AuditResult) -> None:
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        expired_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        headers = [
            '服务名称', '配置项', '期望值', '实际值',
            '是否豁免', '豁免理由', '豁免到期日', '复核状态',
            '是否过期', '来源位置'
        ]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, record in enumerate(audit_result.drift_records, start=2):
            location = audit_result.source_tracker.get_location(record.record_id)

            row_data = [
                record.config_item.service_name,
                record.config_item.config_key,
                record.config_item.expected_value,
                record.config_item.actual_value,
                '是' if record.exemption else '否',
                record.exemption.reason if record.exemption else '',
                record.exemption.expire_date if record.exemption else '',
                self._translate_status(record.review_status.value),
                '是' if record.is_exemption_expired else '否',
                str(location) if location else '',
            ]

            for col, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if record.is_exemption_expired:
                    cell.fill = expired_fill

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 18

    def _fill_expired_sheet(self, ws, audit_result: AuditResult) -> None:
        expired = audit_result.get_expired_exemptions()
        if not expired:
            ws['A1'] = "无已过期豁免记录"
            return

        header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        headers = [
            '服务名称', '配置项', '期望值', '实际值',
            '豁免理由', '豁免到期日', '来源位置'
        ]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, record in enumerate(expired, start=2):
            location = audit_result.source_tracker.get_location(record.record_id)
            row_data = [
                record.config_item.service_name,
                record.config_item.config_key,
                record.config_item.expected_value,
                record.config_item.actual_value,
                record.exemption.reason if record.exemption else '',
                record.exemption.expire_date if record.exemption else '',
                str(location) if location else '',
            ]
            for col, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col, value=value)

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 20

    def _fill_bad_rows_sheet(self, ws, audit_result: AuditResult) -> None:
        header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        header_font = Font(bold=True, color="000000")

        headers = ['文件路径', '工作表', '行号', '原始内容']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, bad_row in enumerate(audit_result.source_tracker.bad_rows, start=2):
            row_data = [
                bad_row.file_path,
                bad_row.sheet_name or '',
                bad_row.row_number,
                bad_row.raw_content,
            ]
            for col, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col, value=value)

        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 50

    def _translate_status(self, status: str) -> str:
        translations = {
            'pending': '待复核',
            'approved': '已批准',
            'rejected': '已拒绝',
            'expired': '已过期',
            'no_exemption': '无豁免',
        }
        return translations.get(status, status)
