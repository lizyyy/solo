from typing import List
from pathlib import Path

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from .base_parser import BaseParser, ParseResult
from ..models import ConfigItem, ExemptionRecord, SourceLocation, SourceTracker, ReviewStatus


class ExcelParser(BaseParser):
    REQUIRED_CONFIG_COLUMNS = ['service_name', 'config_key', 'expected_value', 'actual_value']
    REQUIRED_EXEMPTION_COLUMNS = ['service_name', 'config_key', 'reason', 'expire_date']

    def __init__(self):
        if not HAS_OPENPYXL:
            raise ImportError(
                "openpyxl is required for Excel parsing. "
                "Install it with: pip install openpyxl"
            )

    def _parse_status(self, status_str: str) -> ReviewStatus:
        status_map = {
            'pending': ReviewStatus.PENDING,
            'approved': ReviewStatus.APPROVED,
            'rejected': ReviewStatus.REJECTED,
            'expired': ReviewStatus.EXPIRED,
            'no_exemption': ReviewStatus.NO_EXEMPTION,
        }
        return status_map.get(status_str.lower(), ReviewStatus.PENDING)

    def parse_config_items(self, file_path: str, sheet_name: str = None) -> ParseResult[ConfigItem]:
        path = self._validate_file(file_path)
        result = ParseResult[ConfigItem](file_path=file_path)
        source_tracker = SourceTracker()

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        actual_sheet = ws.title

        headers = self._read_headers(ws)
        self._validate_columns(headers, self.REQUIRED_CONFIG_COLUMNS, file_path)
        header_map = {h: i for i, h in enumerate(headers)}

        for row_num in range(2, ws.max_row + 1):
            row_data = self._read_row(ws, row_num, len(headers))
            try:
                item = ConfigItem(
                    service_name=self._clean_value(row_data[header_map.get('service_name', -1)]),
                    config_key=self._clean_value(row_data[header_map.get('config_key', -1)]),
                    expected_value=self._clean_value(row_data[header_map.get('expected_value', -1)]),
                    actual_value=self._clean_value(row_data[header_map.get('actual_value', -1)]),
                    source_id=f"{actual_sheet}:{row_num}",
                )
                result.items.append(item)
                location = SourceLocation(
                    file_path=file_path,
                    sheet_name=actual_sheet,
                    row_number=row_num,
                    raw_content=str(row_data),
                )
                source_tracker.add_location(item.row_hash, location)
            except Exception as e:
                location = SourceLocation(
                    file_path=file_path,
                    sheet_name=actual_sheet,
                    row_number=row_num,
                    raw_content=str(row_data),
                )
                source_tracker.add_bad_row(location)

        wb.close()
        result.source_tracker = source_tracker
        result.sort_items()
        return result

    def parse_exemptions(self, file_path: str, sheet_name: str = None) -> ParseResult[ExemptionRecord]:
        path = self._validate_file(file_path)
        result = ParseResult[ExemptionRecord](file_path=file_path)
        source_tracker = SourceTracker()

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        actual_sheet = ws.title

        headers = self._read_headers(ws)
        self._validate_columns(headers, self.REQUIRED_EXEMPTION_COLUMNS, file_path)
        header_map = {h: i for i, h in enumerate(headers)}

        for row_num in range(2, ws.max_row + 1):
            row_data = self._read_row(ws, row_num, len(headers))
            try:
                status_str = self._clean_value(row_data[header_map.get('status', -1)] if 'status' in header_map else '')
                status = self._parse_status(status_str) if status_str else ReviewStatus.PENDING
                record = ExemptionRecord(
                    service_name=self._clean_value(row_data[header_map.get('service_name', -1)]),
                    config_key=self._clean_value(row_data[header_map.get('config_key', -1)]),
                    reason=self._clean_value(row_data[header_map.get('reason', -1)]),
                    expire_date=self._clean_value(row_data[header_map.get('expire_date', -1)]),
                    reviewer=self._clean_value(row_data[header_map.get('reviewer', -1)] if 'reviewer' in header_map else ''),
                    status=status,
                )
                result.items.append(record)
                location = SourceLocation(
                    file_path=file_path,
                    sheet_name=actual_sheet,
                    row_number=row_num,
                    raw_content=str(row_data),
                )
                source_tracker.add_location(record.exemption_id, location)
            except Exception as e:
                location = SourceLocation(
                    file_path=file_path,
                    sheet_name=actual_sheet,
                    row_number=row_num,
                    raw_content=str(row_data),
                )
                source_tracker.add_bad_row(location)

        wb.close()
        result.source_tracker = source_tracker
        result.sort_items()
        return result

    def _read_headers(self, ws) -> List[str]:
        return [str(cell.value).strip() if cell.value else '' for cell in ws[1]]

    def _read_row(self, ws, row_num: int, col_count: int) -> List[str]:
        row = ws[row_num]
        return [str(cell.value).strip() if cell.value else '' for cell in row[:col_count]]

    def _clean_value(self, value: str) -> str:
        return value if value else ''

    def _validate_columns(self, actual_columns: List[str], required_columns: List[str], file_path: str):
        missing = set(required_columns) - set(actual_columns or [])
        if missing:
            raise ValueError(
                f"Missing required columns in {file_path}: {', '.join(missing)}. "
                f"Required columns: {', '.join(required_columns)}"
            )
