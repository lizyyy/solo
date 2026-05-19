import csv
import os
from typing import List
from models import HiddenDanger

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

class ReportExporter:
    def export_report(self, hazards: List[HiddenDanger], output_path: str):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        if output_path.endswith('.xlsx'):
            self._export_excel(hazards, output_path)
        else:
            self._export_csv(hazards, output_path)
    
    def _export_csv(self, hazards: List[HiddenDanger], output_path: str):
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                '隐患编号', '隐患描述', '位置', '责任人', '状态', 
                '异常类型', '发现日期', '整改期限', '创建时间', '更新时间'
            ])
            
            for h in hazards:
                writer.writerow([
                    h.hazard_id, h.description, h.location,
                    h.person_in_charge, h.status, h.exception_type,
                    h.discovered_date, h.deadline, h.created_at, h.updated_at
                ])
    
    def _export_excel(self, hazards: List[HiddenDanger], output_path: str):
        if not EXCEL_AVAILABLE:
            raise ImportError("请安装 openpyxl 以支持 Excel 导出: pip install openpyxl")
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '隐患记录'
        
        headers = [
            '隐患编号', '隐患描述', '位置', '责任人', '状态', 
            '异常类型', '发现日期', '整改期限', '创建时间', '更新时间'
        ]
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        for row, h in enumerate(hazards, 2):
            ws.cell(row=row, column=1, value=h.hazard_id)
            ws.cell(row=row, column=2, value=h.description)
            ws.cell(row=row, column=3, value=h.location)
            ws.cell(row=row, column=4, value=h.person_in_charge)
            ws.cell(row=row, column=5, value=h.status)
            ws.cell(row=row, column=6, value=h.exception_type)
            ws.cell(row=row, column=7, value=h.discovered_date)
            ws.cell(row=row, column=8, value=h.deadline)
            ws.cell(row=row, column=9, value=h.created_at)
            ws.cell(row=row, column=10, value=h.updated_at)
            
            status_cell = ws.cell(row=row, column=5)
            if h.status == 'closed':
                status_cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
                status_cell.font = Font(color='FFFFFF')
            elif h.status == 'open':
                status_cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
                status_cell.font = Font(color='FFFFFF')
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(output_path)
