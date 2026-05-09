from datetime import datetime, date
from typing import List, Optional, Dict, Any
from sqlalchemy.orm import Session
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.models.instrument import Instrument, InstrumentStatus
from app.models.borrow import Borrow, BorrowStatus
from app.models.calibration import Calibration, CalibrationStatus


STYLE_HEADER_FONT = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
STYLE_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
STYLE_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
STYLE_BORDER = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)


class ExportService:
    STATUS_MAP_CN = {
        "in_stock": "在库",
        "borrowed": "借出",
        "calibrating": "校准中",
        "sealed": "封存",
        "discarded": "报废",
    }
    
    BORROW_STATUS_MAP_CN = {
        "pending": "待确认",
        "borrowed": "借用中",
        "returned": "已归还",
        "overdue": "逾期",
        "cancelled": "已取消",
    }
    
    CALIBRATION_STATUS_MAP_CN = {
        "scheduled": "已排期",
        "in_progress": "进行中",
        "passed": "合格",
        "failed": "不合格",
        "cancelled": "已取消",
    }

    @staticmethod
    def _fmt_dt(dt) -> str:
        if not dt:
            return ""
        if isinstance(dt, datetime):
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(dt, date):
            return dt.strftime("%Y-%m-%d")
        return str(dt)

    @staticmethod
    def _apply_header_style(ws, headers):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = STYLE_HEADER_FONT
            cell.fill = STYLE_HEADER_FILL
            cell.alignment = STYLE_CENTER
            cell.border = STYLE_BORDER

    @staticmethod
    def _apply_data_border(ws, row, col_count):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = STYLE_CENTER
            cell.border = STYLE_BORDER

    @staticmethod
    def _auto_width(ws):
        for col in ws.columns:
            max_length = 0
            column_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 4, 60)
            ws.column_dimensions[column_letter].width = adjusted_width

    @staticmethod
    def export_instruments(db: Session, status: Optional[InstrumentStatus] = None) -> bytes:
        instruments = db.query(Instrument).all()
        if status:
            instruments = [i for i in instruments if i.status == status]
        
        headers = [
            "器具编号", "器具名称", "规格型号", "出厂编号", "制造商",
            "准确度等级", "测量范围", "存放位置", "当前状态",
            "校准周期(月)", "上次校准日期", "下次校准日期",
            "创建时间", "更新时间",
        ]
        
        wb = Workbook()
        ws = wb.active
        ws.title = "器具档案"
        
        ExportService._apply_header_style(ws, headers)
        
        for row_idx, instrument in enumerate(instruments, start=2):
            data = [
                instrument.code or "",
                instrument.name or "",
                instrument.specification or "",
                instrument.serial_number or "",
                instrument.manufacturer or "",
                instrument.accuracy or "",
                instrument.measurement_range or "",
                instrument.location or "",
                ExportService.STATUS_MAP_CN.get(instrument.status.value, instrument.status.value) if instrument.status else "",
                instrument.calibration_period_months or 0,
                ExportService._fmt_dt(instrument.last_calibration_date),
                ExportService._fmt_dt(instrument.next_calibration_date),
                ExportService._fmt_dt(instrument.created_at),
                ExportService._fmt_dt(instrument.updated_at),
            ]
            for col_idx, value in enumerate(data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            ExportService._apply_data_border(ws, row_idx, len(headers))
        
        ExportService._auto_width(ws)
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_borrows(db: Session, status: Optional[BorrowStatus] = None) -> bytes:
        query = db.query(Borrow)
        if status:
            query = query.filter(Borrow.status == status)
        borrows = query.order_by(Borrow.created_at.desc()).all()
        
        headers = [
            "借用单号", "器具编号", "借用人", "部门", "车间",
            "用途", "工单号", "借用日期", "预计归还日期",
            "实际归还日期", "状态", "是否逾期", "催还次数",
            "归还状况", "备注", "创建时间",
        ]
        
        wb = Workbook()
        ws = wb.active
        ws.title = "借用记录"
        
        ExportService._apply_header_style(ws, headers)
        
        from app.models.instrument import Instrument
        from app.models.user import User
        
        for row_idx, borrow in enumerate(borrows, start=2):
            instrument = db.query(Instrument).filter(Instrument.id == borrow.instrument_id).first()
            borrower = db.query(User).filter(User.id == borrow.borrower_id).first()
            
            data = [
                borrow.id,
                instrument.code if instrument else "",
                borrower.full_name if borrower else "",
                borrow.department or "",
                borrow.workshop or "",
                borrow.purpose or "",
                borrow.work_order or "",
                ExportService._fmt_dt(borrow.borrow_date),
                ExportService._fmt_dt(borrow.expected_return_date),
                ExportService._fmt_dt(borrow.actual_return_date),
                ExportService.BORROW_STATUS_MAP_CN.get(borrow.status.value, borrow.status.value) if borrow.status else "",
                "是" if borrow.is_overdue else "否",
                borrow.overdue_notice_count or 0,
                borrow.return_condition or "",
                borrow.remarks or "",
                ExportService._fmt_dt(borrow.created_at),
            ]
            for col_idx, value in enumerate(data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            ExportService._apply_data_border(ws, row_idx, len(headers))
        
        ExportService._auto_width(ws)
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_calibrations(db: Session, status: Optional[CalibrationStatus] = None) -> bytes:
        query = db.query(Calibration)
        if status:
            query = query.filter(Calibration.status == status)
        calibrations = query.order_by(Calibration.created_at.desc()).all()
        
        headers = [
            "校准单号", "器具编号", "校准类型", "计划校准日期",
            "实际校准日期", "校准机构", "证书编号",
            "测量不确定度", "环境条件", "结果", "状态",
            "备注", "创建时间",
        ]
        
        wb = Workbook()
        ws = wb.active
        ws.title = "校准记录"
        
        ExportService._apply_header_style(ws, headers)
        
        from app.models.instrument import Instrument
        
        for row_idx, cal in enumerate(calibrations, start=2):
            instrument = db.query(Instrument).filter(Instrument.id == cal.instrument_id).first()
            
            result_str = ""
            if cal.result_pass is True:
                result_str = "合格"
            elif cal.result_pass is False:
                result_str = "不合格"
            
            data = [
                cal.id,
                instrument.code if instrument else "",
                cal.calibration_type or "",
                ExportService._fmt_dt(cal.scheduled_date),
                ExportService._fmt_dt(cal.calibration_date),
                cal.calibration_agency or "",
                cal.certificate_number or "",
                cal.measurement_uncertainty or "",
                cal.environmental_conditions or "",
                result_str,
                ExportService.CALIBRATION_STATUS_MAP_CN.get(cal.status.value, cal.status.value) if cal.status else "",
                cal.remarks or "",
                ExportService._fmt_dt(cal.created_at),
            ]
            for col_idx, value in enumerate(data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            ExportService._apply_data_border(ws, row_idx, len(headers))
        
        ExportService._auto_width(ws)
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_full_ledger(db: Session) -> bytes:
        instruments = db.query(Instrument).order_by(Instrument.code.asc()).all()
        
        headers = [
            "器具编号", "器具名称", "规格型号", "当前状态",
            "校准周期(月)", "上次校准日期", "下次校准日期",
            "借用状态", "借用人", "部门", "预计归还日期",
            "最近校准结果", "历史借用次数", "是否逾期",
            "备注",
        ]
        
        wb = Workbook()
        ws = wb.active
        ws.title = "完整台账"
        
        ExportService._apply_header_style(ws, headers)
        
        from app.models.instrument import Instrument
        from app.models.borrow import Borrow
        from app.models.calibration import Calibration
        from app.models.user import User
        
        for row_idx, instrument in enumerate(instruments, start=2):
            active_borrow = (
                db.query(Borrow)
                .filter(
                    Borrow.instrument_id == instrument.id,
                    Borrow.status.in_([BorrowStatus.BORROWED, BorrowStatus.OVERDUE]),
                )
                .first()
            )
            
            borrower = None
            if active_borrow:
                borrower = db.query(User).filter(User.id == active_borrow.borrower_id).first()
            
            borrow_count = (
                db.query(Borrow)
                .filter(
                    Borrow.instrument_id == instrument.id,
                    Borrow.status == BorrowStatus.RETURNED,
                )
                .count()
            )
            
            last_cal = (
                db.query(Calibration)
                .filter(
                    Calibration.instrument_id == instrument.id,
                    Calibration.status.in_([CalibrationStatus.PASSED, CalibrationStatus.FAILED]),
                )
                .order_by(Calibration.calibration_date.desc())
                .first()
            )
            
            cal_result_str = ""
            if last_cal:
                if last_cal.result_pass is True:
                    cal_result_str = "合格"
                elif last_cal.result_pass is False:
                    cal_result_str = "不合格"
            
            data = [
                instrument.code or "",
                instrument.name or "",
                instrument.specification or "",
                ExportService.STATUS_MAP_CN.get(instrument.status.value, instrument.status.value) if instrument.status else "",
                instrument.calibration_period_months or 0,
                ExportService._fmt_dt(instrument.last_calibration_date),
                ExportService._fmt_dt(instrument.next_calibration_date),
                ExportService.BORROW_STATUS_MAP_CN.get(active_borrow.status.value, active_borrow.status.value) if active_borrow and active_borrow.status else "-",
                borrower.full_name if borrower else "-",
                active_borrow.department if active_borrow else "-",
                ExportService._fmt_dt(active_borrow.expected_return_date) if active_borrow else "-",
                cal_result_str or "-",
                borrow_count,
                "是" if (active_borrow and active_borrow.is_overdue) else "否",
                "",
            ]
            for col_idx, value in enumerate(data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            ExportService._apply_data_border(ws, row_idx, len(headers))
        
        ExportService._auto_width(ws)
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
