import csv
import json
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
try:
    import openpyxl
    from openpyxl import Workbook
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from .models import ParentAppeal, MatchedCase, RulingDecision, ReviewDecision
from .storage import IdempotentStore


class DataExporter:
    def __init__(self, store: IdempotentStore):
        self.store = store
    
    def export_appeals_to_json(self, output_path: str, status: Optional[str] = None) -> str:
        appeals = self.store.get_all_appeals()
        if status:
            appeals = [a for a in appeals if a.status.value == status]
        
        data = [a.model_dump() for a in appeals]
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
        
        return output_path
    
    def export_rulings_to_json(self, output_path: str) -> str:
        rulings = self.store.get_all_rulings()
        data = [r.model_dump() for r in rulings]
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
        
        return output_path
    
    def export_appeals_to_csv(self, output_path: str, status: Optional[str] = None) -> str:
        appeals = self.store.get_all_appeals()
        if status:
            appeals = [a for a in appeals if a.status.value == status]
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        if not appeals:
            with open(output_path, 'w', encoding='utf-8-sig') as f:
                pass
            return output_path
        
        fieldnames = [
            'appeal_id', 'parent_name', 'student_name', 'bus_id', 'route_id',
            'stop_name', 'scheduled_time', 'actual_arrival_time', 'appeal_time',
            'description', 'status', 'created_at', 'created_by'
        ]
        
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            
            for appeal in appeals:
                row = appeal.model_dump()
                for key in row:
                    if isinstance(row[key], datetime):
                        row[key] = row[key].isoformat()
                writer.writerow({k: row.get(k, '') for k in fieldnames})
        
        return output_path
    
    def export_rulings_to_csv(self, output_path: str) -> str:
        rulings = self.store.get_all_rulings()
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        if not rulings:
            with open(output_path, 'w', encoding='utf-8-sig') as f:
                pass
            return output_path
        
        fieldnames = [
            'ruling_id', 'case_id', 'appeal_id', 'result', 'reason',
            'ruled_at', 'ruled_by'
        ]
        
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            
            for ruling in rulings:
                row = ruling.model_dump()
                for key in row:
                    if isinstance(row[key], datetime):
                        row[key] = row[key].isoformat()
                writer.writerow({k: row.get(k, '') for k in fieldnames})
        
        return output_path
    
    def export_case_details_to_excel(self, output_path: str) -> str:
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export")
        
        wb = Workbook()
        
        ws1 = wb.active
        ws1.title = "申诉记录"
        
        appeals = self.store.get_all_appeals()
        if appeals:
            headers1 = ['申诉ID', '家长姓名', '学生姓名', '车牌号', '线路ID',
                       '站点名称', '计划时间', '实际到达时间', '申诉时间',
                       '描述', '状态', '创建人']
            ws1.append(headers1)
            
            for a in appeals:
                row = [
                    a.appeal_id, a.parent_name, a.student_name, a.bus_id,
                    a.route_id, a.stop_name,
                    a.scheduled_time.isoformat() if a.scheduled_time else '',
                    a.actual_arrival_time.isoformat() if a.actual_arrival_time else '',
                    a.appeal_time.isoformat() if a.appeal_time else '',
                    a.description, a.status.value, a.created_by
                ]
                ws1.append(row)
        
        ws2 = wb.create_sheet("裁定结果")
        rulings = self.store.get_all_rulings()
        if rulings:
            headers2 = ['裁定ID', '案例ID', '申诉ID', '裁定结果', '原因',
                       '裁定时间', '裁定人']
            ws2.append(headers2)
            
            for r in rulings:
                row = [
                    r.ruling_id, r.case_id, r.appeal_id, r.result.value,
                    r.reason, r.ruled_at.isoformat() if r.ruled_at else '',
                    r.ruled_by
                ]
                ws2.append(row)
        
        ws3 = wb.create_sheet("复核记录")
        reviews = self.store.get_all_reviews()
        if reviews:
            headers3 = ['复核ID', '案例ID', '申诉ID', '原裁定ID', '是否维持',
                       '新结果', '原因', '复核时间', '复核人']
            ws3.append(headers3)
            
            for r in reviews:
                row = [
                    r.review_id, r.case_id, r.appeal_id, r.original_ruling_id,
                    '是' if r.uphold else '否',
                    r.new_result.value if r.new_result else '',
                    r.reason,
                    r.reviewed_at.isoformat() if r.reviewed_at else '',
                    r.reviewed_by
                ]
                ws3.append(row)
        
        ws4 = wb.create_sheet("审计日志")
        audit_logs = self.store.get_audit_logs()
        if audit_logs:
            headers4 = ['日志ID', '实体类型', '实体ID', '操作', '操作人',
                       '操作人角色', '时间', '详情']
            ws4.append(headers4)
            
            for log in audit_logs:
                row = [
                    log.id, log.entity_type, log.entity_id, log.action,
                    log.operator, log.operator_role.value,
                    log.timestamp.isoformat() if log.timestamp else '',
                    json.dumps(log.details, ensure_ascii=False)
                ]
                ws4.append(row)
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        
        return output_path
    
    def export_summary_report(self, output_path: str) -> str:
        appeals = self.store.get_all_appeals()
        rulings = self.store.get_all_rulings()
        
        status_counts = {}
        for a in appeals:
            status = a.status.value
            status_counts[status] = status_counts.get(status, 0) + 1
        
        result_counts = {}
        for r in rulings:
            result = r.result.value
            result_counts[result] = result_counts.get(result, 0) + 1
        
        report = {
            "generated_at": datetime.now().isoformat(),
            "summary": {
                "total_appeals": len(appeals),
                "total_rulings": len(rulings),
                "appeal_status_counts": status_counts,
                "ruling_result_counts": result_counts
            }
        }
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        
        return output_path
