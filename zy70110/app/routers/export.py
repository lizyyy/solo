from datetime import date, datetime
from typing import Optional
from io import BytesIO
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import and_
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.database import get_db
from app.models import (
    Appointment, InspectionQueue, StorageLock, SkipRecord,
    LoadingReceipt, FailedTask, AppointmentStatus, InspectionStatus
)

router = APIRouter(prefix="/export", tags=["数据导出"])


def get_status_display(status):
    status_mapping = {
        "pending": "待确认",
        "confirmed": "已确认",
        "checked_in": "已签到",
        "in_inspection": "检测中",
        "inspection_passed": "检测合格",
        "inspection_failed": "检测不合格",
        "in_loading": "装卸中",
        "completed": "已完成",
        "cancelled": "已取消",
        "skipped": "已过号",
        "failed": "失败",
        "waiting": "等待中",
        "in_progress": "进行中",
        "passed": "合格",
        "available": "可用",
        "locked": "已锁定",
        "occupied": "已占用"
    }
    return status_mapping.get(status, status)


def format_datetime(dt):
    if not dt:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def format_date(d):
    if not d:
        return ""
    return d.strftime("%Y-%m-%d")


def apply_header_style(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )


def apply_data_style(cell):
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )


@router.get("/daily-report")
def export_daily_report(
    start_date: Optional[date] = Query(None, description="开始日期"),
    end_date: Optional[date] = Query(None, description="结束日期"),
    db: Session = Depends(get_db)
):
    if not start_date:
        start_date = date.today()
    if not end_date:
        end_date = date.today()
    
    appointments = db.query(Appointment).filter(
        and_(
            Appointment.scheduled_date >= start_date,
            Appointment.scheduled_date <= end_date
        )
    ).order_by(Appointment.scheduled_date, Appointment.created_at).all()
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "入库预约汇总"
    headers1 = ["预约编号", "客户名称", "联系电话", "车牌号", "司机姓名",
                "产品类型", "产品名称", "数量", "单位", "体积(立方米)",
                "温区要求", "预约日期", "时段", "当前状态", "仓位编码",
                "创建时间", "最后更新"]
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        apply_header_style(cell)
    
    for row, appt in enumerate(appointments, 2):
        data = [
            appt.appointment_no, appt.customer_name, appt.contact_phone,
            appt.license_plate, appt.driver_name, appt.product_type,
            appt.product_name, appt.quantity, appt.unit, appt.volume_cubic_meters,
            appt.temperature_requirement, format_date(appt.scheduled_date),
            appt.scheduled_time_slot, get_status_display(appt.status),
            appt.storage_bay.bay_code if appt.storage_bay else "",
            format_datetime(appt.created_at), format_datetime(appt.updated_at)
        ]
        for col, value in enumerate(data, 1):
            cell = ws1.cell(row=row, column=col, value=value)
            apply_data_style(cell)
    
    for col in range(1, len(headers1) + 1):
        ws1.column_dimensions[chr(64 + col)].width = 18
    ws1.row_dimensions[1].height = 25
    
    ws2 = wb.create_sheet("检测排队明细")
    headers2 = ["预约编号", "客户名称", "车牌号", "排队日期", "排队号",
                "检测窗口", "检测状态", "签到时间", "检测开始时间",
                "检测结束时间", "等待时长(分钟)", "检测员", "检测备注"]
    
    queue_records = db.query(InspectionQueue).filter(
        and_(
            InspectionQueue.queue_date >= start_date,
            InspectionQueue.queue_date <= end_date
        )
    ).order_by(InspectionQueue.queue_date, InspectionQueue.queue_number).all()
    
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        apply_header_style(cell)
    
    for row, record in enumerate(queue_records, 2):
        appt = record.appointment
        data = [
            appt.appointment_no if appt else "",
            appt.customer_name if appt else "",
            appt.license_plate if appt else "",
            format_date(record.queue_date),
            record.queue_number,
            record.inspection_window.window_code if record.inspection_window else "",
            get_status_display(record.status),
            format_datetime(record.checked_in_time),
            format_datetime(record.inspection_start_time),
            format_datetime(record.inspection_end_time),
            record.actual_wait_minutes or "",
            record.inspector_name or "",
            record.inspection_notes or ""
        ]
        for col, value in enumerate(data, 1):
            cell = ws2.cell(row=row, column=col, value=value)
            apply_data_style(cell)
    
    for col in range(1, len(headers2) + 1):
        ws2.column_dimensions[chr(64 + col)].width = 18
    ws2.row_dimensions[1].height = 25
    
    ws3 = wb.create_sheet("过号记录")
    headers3 = ["预约编号", "客户名称", "车牌号", "过号类型", "过号时间",
                "过号原因", "原排队号", "新排队号", "是否已重新分配",
                "重新分配时间", "操作人"]
    
    skip_records = db.query(SkipRecord).filter(
        SkipRecord.skip_time >= datetime.combine(start_date, datetime.min.time()),
        SkipRecord.skip_time <= datetime.combine(end_date, datetime.max.time())
    ).order_by(SkipRecord.skip_time.desc()).all()
    
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        apply_header_style(cell)
    
    for row, record in enumerate(skip_records, 2):
        appt = record.appointment
        data = [
            appt.appointment_no if appt else "",
            appt.customer_name if appt else "",
            appt.license_plate if appt else "",
            get_status_display(record.skip_type),
            format_datetime(record.skip_time),
            record.skip_reason,
            record.original_queue_number,
            record.new_queue_number or "",
            "是" if record.is_reassigned else "否",
            format_datetime(record.reassigned_at),
            record.operator_name or ""
        ]
        for col, value in enumerate(data, 1):
            cell = ws3.cell(row=row, column=col, value=value)
            apply_data_style(cell)
    
    for col in range(1, len(headers3) + 1):
        ws3.column_dimensions[chr(64 + col)].width = 18
    ws3.row_dimensions[1].height = 25
    
    ws4 = wb.create_sheet("装卸回执")
    headers4 = ["回执编号", "预约编号", "客户名称", "仓位编码", "装卸开始时间",
                "装卸结束时间", "实际数量", "实际体积", "库温读数",
                "装卸员", "验收状态", "差异说明", "客户确认", "确认人", "确认时间"]
    
    receipts = db.query(LoadingReceipt).filter(
        LoadingReceipt.created_at >= datetime.combine(start_date, datetime.min.time()),
        LoadingReceipt.created_at <= datetime.combine(end_date, datetime.max.time())
    ).order_by(LoadingReceipt.created_at.desc()).all()
    
    for col, header in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=col, value=header)
        apply_header_style(cell)
    
    for row, receipt in enumerate(receipts, 2):
        appt = receipt.appointment
        data = [
            receipt.receipt_no,
            appt.appointment_no if appt else "",
            appt.customer_name if appt else "",
            appt.storage_bay.bay_code if appt and appt.storage_bay else "",
            format_datetime(receipt.loading_start_time),
            format_datetime(receipt.loading_end_time),
            receipt.actual_quantity,
            receipt.actual_volume,
            receipt.temperature_reading or "",
            receipt.handler_name or "",
            "合格" if receipt.acceptance_status == "accepted" else "有差异",
            receipt.discrepancy_notes or "",
            "已确认" if receipt.customer_confirmation else "未确认",
            receipt.confirmed_by or "",
            format_datetime(receipt.confirmed_at)
        ]
        for col, value in enumerate(data, 1):
            cell = ws4.cell(row=row, column=col, value=value)
            apply_data_style(cell)
    
    for col in range(1, len(headers4) + 1):
        ws4.column_dimensions[chr(64 + col)].width = 18
    ws4.row_dimensions[1].height = 25
    
    ws5 = wb.create_sheet("异常处理记录")
    headers5 = ["预约编号", "任务类型", "任务描述", "错误信息",
                "错误详情", "发生时间", "重试次数", "最大重试次数",
                "最后重试时间", "是否已解决", "解决时间", "解决说明"]
    
    failed_tasks = db.query(FailedTask).filter(
        FailedTask.occurred_at >= datetime.combine(start_date, datetime.min.time()),
        FailedTask.occurred_at <= datetime.combine(end_date, datetime.max.time())
    ).order_by(FailedTask.occurred_at.desc()).all()
    
    for col, header in enumerate(headers5, 1):
        cell = ws5.cell(row=1, column=col, value=header)
        apply_header_style(cell)
    
    for row, task in enumerate(failed_tasks, 2):
        appt = task.appointment
        data = [
            appt.appointment_no if appt else "",
            task.task_type,
            task.task_description,
            task.error_message,
            task.error_details or "",
            format_datetime(task.occurred_at),
            task.retry_count,
            task.max_retries,
            format_datetime(task.last_retry_at),
            "已解决" if task.is_resolved else "待处理",
            format_datetime(task.resolved_at),
            task.resolution_notes or ""
        ]
        for col, value in enumerate(data, 1):
            cell = ws5.cell(row=row, column=col, value=value)
            apply_data_style(cell)
    
    for col in range(1, len(headers5) + 1):
        ws5.column_dimensions[chr(64 + col)].width = 18
    ws5.row_dimensions[1].height = 25
    
    ws6 = wb.create_sheet("业务统计")
    ws6.cell(row=1, column=1, value=f"入库业务复核报告 - {format_date(start_date)} 至 {format_date(end_date)}")
    ws6.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws6.merge_cells("A1:F1")
    
    stats_data = [
        ["统计指标", "数值"],
        ["总预约数", len(appointments)],
        ["已完成", len([a for a in appointments if a.status == AppointmentStatus.COMPLETED])],
        ["检测合格", len([a for a in appointments if a.status in [AppointmentStatus.INSPECTION_PASSED, AppointmentStatus.IN_LOADING, AppointmentStatus.COMPLETED]])],
        ["检测不合格", len([a for a in appointments if a.status == AppointmentStatus.INSPECTION_FAILED])],
        ["过号数量", len(skip_records)],
        ["过号后重新分配", len([s for s in skip_records if s.is_reassigned])],
        ["异常任务数", len(failed_tasks)],
        ["已解决异常", len([t for t in failed_tasks if t.is_resolved])],
        ["待解决异常", len([t for t in failed_tasks if not t.is_resolved])]
    ]
    
    for row, data in enumerate(stats_data, 3):
        for col, value in enumerate(data, 1):
            cell = ws6.cell(row=row, column=col, value=value)
            if row == 3:
                apply_header_style(cell)
            else:
                apply_data_style(cell)
    
    ws6.column_dimensions["A"].width = 20
    ws6.column_dimensions["B"].width = 15
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"冷库入库业务复核报告_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/appointment/{appointment_id}")
def export_single_appointment(
    appointment_id: int,
    db: Session = Depends(get_db)
):
    from app.models import Appointment
    
    appt = db.query(Appointment).filter(Appointment.id == appointment_id).first()
    if not appt:
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="预约不存在")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "入库单详情"
    
    ws.cell(row=1, column=1, value=f"冷库入库预约单详情 - {appt.appointment_no}")
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    
    basic_info = [
        ["基本信息", ""],
        ["预约编号", appt.appointment_no],
        ["客户名称", appt.customer_name],
        ["联系电话", appt.contact_phone],
        ["车牌号", appt.license_plate],
        ["司机姓名", appt.driver_name],
        ["预约日期", format_date(appt.scheduled_date)],
        ["预约时段", appt.scheduled_time_slot],
        ["当前状态", get_status_display(appt.status)],
        ["", ""],
        ["货物信息", ""],
        ["产品类型", appt.product_type],
        ["产品名称", appt.product_name],
        ["数量", f"{appt.quantity} {appt.unit}"],
        ["体积", f"{appt.volume_cubic_meters} 立方米"],
        ["温区要求", appt.temperature_requirement],
        ["", ""],
        ["仓位信息", ""],
        ["仓位编码", appt.storage_bay.bay_code if appt.storage_bay else "待分配"],
        ["仓位名称", appt.storage_bay.bay_name if appt.storage_bay else ""],
        ["所在区域", appt.storage_bay.zone if appt.storage_bay else ""],
        ["仓位容量", f"{appt.storage_bay.capacity_cubic_meters} 立方米" if appt.storage_bay else ""]
    ]
    
    for row, data in enumerate(basic_info, 3):
        cell1 = ws.cell(row=row, column=1, value=data[0])
        cell2 = ws.cell(row=row, column=2, value=data[1])
        if data[0] and not data[1]:
            cell1.font = Font(bold=True)
            cell1.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            ws.merge_cells(f"A{row}:H{row}")
    
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 30
    
    if appt.queue_record:
        ws2 = wb.create_sheet("检测记录")
        q = appt.queue_record
        inspection_data = [
            ["检测排队信息", ""],
            ["排队日期", format_date(q.queue_date)],
            ["排队号", q.queue_number],
            ["检测窗口", q.inspection_window.window_code if q.inspection_window else ""],
            ["检测状态", get_status_display(q.status)],
            ["签到时间", format_datetime(q.checked_in_time)],
            ["检测开始时间", format_datetime(q.inspection_start_time)],
            ["检测结束时间", format_datetime(q.inspection_end_time)],
            ["实际等待时长", f"{q.actual_wait_minutes} 分钟" if q.actual_wait_minutes else ""],
            ["检测员", q.inspector_name or ""],
            ["检测备注", q.inspection_notes or ""]
        ]
        for row, data in enumerate(inspection_data, 1):
            cell1 = ws2.cell(row=row, column=1, value=data[0])
            cell2 = ws2.cell(row=row, column=2, value=data[1])
            if not data[1]:
                cell1.font = Font(bold=True)
        ws2.column_dimensions["A"].width = 15
        ws2.column_dimensions["B"].width = 30
    
    if appt.loading_receipt:
        ws3 = wb.create_sheet("装卸回执")
        r = appt.loading_receipt
        receipt_data = [
            ["装卸回执信息", ""],
            ["回执编号", r.receipt_no],
            ["装卸开始时间", format_datetime(r.loading_start_time)],
            ["装卸结束时间", format_datetime(r.loading_end_time)],
            ["实际数量", f"{r.actual_quantity} {appt.unit}"],
            ["实际体积", f"{r.actual_volume} 立方米"],
            ["库温读数", f"{r.temperature_reading} °C" if r.temperature_reading else ""],
            ["装卸员", r.handler_name or ""],
            ["验收状态", "合格" if r.acceptance_status == "accepted" else "有差异"],
            ["差异说明", r.discrepancy_notes or "无"],
            ["客户确认", "已确认" if r.customer_confirmation else "未确认"],
            ["确认人", r.confirmed_by or ""],
            ["确认时间", format_datetime(r.confirmed_at)]
        ]
        for row, data in enumerate(receipt_data, 1):
            cell1 = ws3.cell(row=row, column=1, value=data[0])
            cell2 = ws3.cell(row=row, column=2, value=data[1])
            if not data[1]:
                cell1.font = Font(bold=True)
        ws3.column_dimensions["A"].width = 15
        ws3.column_dimensions["B"].width = 30
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"入库预约单_{appt.appointment_no}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
