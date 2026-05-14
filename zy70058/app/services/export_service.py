from datetime import datetime
import os
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from app import db
from app.models import (
    Account, InstallmentPlan, RevocationRecord, ExceptionRecord, PendingTask,
    BillRecord, FeeAmortization
)
from config import Config


class ExportService:
    
    @staticmethod
    def _get_export_filename(prefix):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"{prefix}_{timestamp}.xlsx"
    
    @staticmethod
    def _setup_excel_workbook(filename):
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        return wb
    
    @staticmethod
    def _apply_header_style(sheet, headers):
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        for col_idx, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
    
    @staticmethod
    def _auto_width(sheet):
        for col in sheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    cell_value = str(cell.value)
                    if cell_value and len(cell_value) > max_length:
                        max_length = len(cell_value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column].width = adjusted_width
    
    @staticmethod
    def export_revocation_report(start_date=None, end_date=None):
        wb = ExportService._setup_excel_workbook('revocation')
        
        filename = ExportService._get_export_filename('撤销记录汇总')
        file_path = os.path.join(Config.EXPORT_DIR, filename)
        
        summary_sheet = wb.create_sheet('撤销汇总')
        details_sheet = wb.create_sheet('明细详情')
        
        query = RevocationRecord.query
        if start_date:
            start_dt = datetime.strptime(start_date, '%Y-%m-%d')
            query = query.filter(RevocationRecord.created_at >= start_dt)
        if end_date:
            end_dt = datetime.strptime(end_date, '%Y-%m-%d')
            query = query.filter(RevocationRecord.created_at <= end_dt)
        
        revocations = query.order_by(RevocationRecord.created_at.desc()).all()
        
        plan_account_map = {}
        if revocations:
            plan_ids = [r.plan_id for r in revocations]
            plans = InstallmentPlan.query.filter(InstallmentPlan.id.in_(plan_ids)).all()
            for p in plans:
                plan_account_map[p.id] = p.account_id
        
        summary_headers = ['撤销类型', '数量', '总退款', '总违约金', '总待收', '总额度恢复']
        ExportService._apply_header_style(summary_sheet, summary_headers)
        
        by_type = {}
        total_refund, total_penalty, total_collect, total_credit = 0, 0, 0, 0
        
        for rev in revocations:
            rev_type = rev.revocation_type
            if rev_type not in by_type:
                by_type[rev_type] = {'count': 0, 'refund': 0, 'penalty': 0, 'collect': 0, 'credit': 0}
            by_type[rev_type]['count'] += 1
            by_type[rev_type]['refund'] += float(rev.total_refund)
            by_type[rev_type]['penalty'] += float(rev.penalty_fee)
            by_type[rev_type]['collect'] += float(rev.amount_to_collect)
            by_type[rev_type]['credit'] += float(rev.credit_restored)
            total_refund += float(rev.total_refund)
            total_penalty += float(rev.penalty_fee)
            total_collect += float(rev.amount_to_collect)
            total_credit += float(rev.credit_restored)
        
        row_idx = 2
        for rev_type, stats in by_type.items():
            summary_sheet.cell(row=row_idx, column=1, value=rev_type)
            summary_sheet.cell(row=row_idx, column=2, value=stats['count'])
            summary_sheet.cell(row=row_idx, column=3, value=stats['refund'])
            summary_sheet.cell(row=row_idx, column=4, value=stats['penalty'])
            summary_sheet.cell(row=row_idx, column=5, value=stats['collect'])
            summary_sheet.cell(row=row_idx, column=6, value=stats['credit'])
            row_idx += 1
        
        row_idx += 1
        summary_sheet.cell(row=row_idx, column=1, value='总计')
        summary_sheet.cell(row=row_idx, column=2, value=len(revocations))
        summary_sheet.cell(row=row_idx, column=3, value=total_refund)
        summary_sheet.cell(row=row_idx, column=4, value=total_penalty)
        summary_sheet.cell(row=row_idx, column=5, value=total_collect)
        summary_sheet.cell(row=row_idx, column=6, value=total_credit)
        
        ExportService._auto_width(summary_sheet)
        
        detail_headers = [
            'ID', '计划ID', '账户ID', '撤销类型', '交易ID', '已付本金', '已付手续费',
            '退款本金', '退款手续费', '总退款', '违约金', '待收金额',
            '撤销前剩余本金', '撤销前剩余手续费', '恢复额度', '撤销原因', '撤销时间'
        ]
        ExportService._apply_header_style(details_sheet, detail_headers)
        
        for row_idx, rev in enumerate(revocations, 2):
            details_sheet.cell(row=row_idx, column=1, value=rev.id)
            details_sheet.cell(row=row_idx, column=2, value=rev.plan_id)
            details_sheet.cell(row=row_idx, column=3, value=plan_account_map.get(rev.plan_id, ''))
            details_sheet.cell(row=row_idx, column=4, value=rev.revocation_type)
            details_sheet.cell(row=row_idx, column=5, value=rev.transaction_id)
            details_sheet.cell(row=row_idx, column=6, value=float(rev.original_paid_principal))
            details_sheet.cell(row=row_idx, column=7, value=float(rev.original_paid_fee))
            details_sheet.cell(row=row_idx, column=8, value=float(rev.refund_principal))
            details_sheet.cell(row=row_idx, column=9, value=float(rev.refund_fee))
            details_sheet.cell(row=row_idx, column=10, value=float(rev.total_refund))
            details_sheet.cell(row=row_idx, column=11, value=float(rev.penalty_fee))
            details_sheet.cell(row=row_idx, column=12, value=float(rev.amount_to_collect))
            details_sheet.cell(row=row_idx, column=13, value=float(rev.remaining_principal_before))
            details_sheet.cell(row=row_idx, column=14, value=float(rev.remaining_fee_before))
            details_sheet.cell(row=row_idx, column=15, value=float(rev.credit_restored))
            details_sheet.cell(row=row_idx, column=16, value=rev.reason or '')
            details_sheet.cell(row=row_idx, column=17, value=rev.created_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        ExportService._auto_width(details_sheet)
        
        wb.save(file_path)
        return file_path
    
    @staticmethod
    def export_exception_report(severity=None, status=None):
        wb = ExportService._setup_excel_workbook('exception')
        
        filename = ExportService._get_export_filename('异常记录')
        file_path = os.path.join(Config.EXPORT_DIR, filename)
        
        summary_sheet = wb.create_sheet('异常汇总')
        details_sheet = wb.create_sheet('明细详情')
        
        query = ExceptionRecord.query
        if severity:
            query = query.filter_by(severity=severity)
        if status:
            query = query.filter_by(status=status)
        
        exceptions = query.order_by(ExceptionRecord.detected_at.desc()).all()
        
        summary_headers = ['严重级别', '数量', '占比', '状态分布', '未解决数量']
        ExportService._apply_header_style(summary_sheet, summary_headers)
        
        by_severity = {}
        for exc in exceptions:
            if exc.severity not in by_severity:
                by_severity[exc.severity] = {'total': 0, 'unresolved': 0}
            by_severity[exc.severity]['total'] += 1
            if exc.status in [ExceptionRecord.STATUS_OPEN, ExceptionRecord.STATUS_INVESTIGATING]:
                by_severity[exc.severity]['unresolved'] += 1
        
        row_idx = 2
        total = len(exceptions)
        for sev, stats in by_severity.items():
            summary_sheet.cell(row=row_idx, column=1, value=sev)
            summary_sheet.cell(row=row_idx, column=2, value=stats['total'])
            summary_sheet.cell(row=row_idx, column=3, value=f"{stats['total']/total*100:.2f}%")
            summary_sheet.cell(row=row_idx, column=4, value=stats['unresolved'])
            row_idx += 1
        
        ExportService._auto_width(summary_sheet)
        
        detail_headers = [
            'ID', '账户ID', '计划ID', '异常类型', '严重级别', '标题', '描述',
            '状态', '检测时间', '解决时间'
        ]
        ExportService._apply_header_style(details_sheet, detail_headers)
        
        for row_idx, exc in enumerate(exceptions, 2):
            details_sheet.cell(row=row_idx, column=1, value=exc.id)
            details_sheet.cell(row=row_idx, column=2, value=exc.account_id or '')
            details_sheet.cell(row=row_idx, column=3, value=exc.plan_id or '')
            details_sheet.cell(row=row_idx, column=4, value=exc.exception_type)
            details_sheet.cell(row=row_idx, column=5, value=exc.severity)
            details_sheet.cell(row=row_idx, column=6, value=exc.title)
            details_sheet.cell(row=row_idx, column=7, value=exc.description or '')
            details_sheet.cell(row=row_idx, column=8, value=exc.status)
            details_sheet.cell(row=row_idx, column=9, value=exc.detected_at.strftime('%Y-%m-%d %H:%M:%S'))
            details_sheet.cell(row=row_idx, column=10, value=exc.resolved_at.strftime('%Y-%m-%d %H:%M:%S') if exc.resolved_at else '')
        
        ExportService._auto_width(details_sheet)
        
        wb.save(file_path)
        return file_path
    
    @staticmethod
    def export_pending_tasks_report(task_type=None, status=None):
        wb = ExportService._setup_excel_workbook('pending_tasks')
        
        filename = ExportService._get_export_filename('待处理任务')
        file_path = os.path.join(Config.EXPORT_DIR, filename)
        
        summary_sheet = wb.create_sheet('任务汇总')
        details_sheet = wb.create_sheet('明细详情')
        
        query = PendingTask.query
        if task_type:
            query = query.filter_by(task_type=task_type)
        if status:
            query = query.filter_by(status=status)
        
        tasks = query.order_by(PendingTask.scheduled_at.asc()).all()
        
        summary_headers = ['任务类型', '待处理', '处理中', '已完成', '失败', '总计']
        ExportService._apply_header_style(summary_sheet, summary_headers)
        
        by_type = {}
        for task in tasks:
            if task.task_type not in by_type:
                by_type[task.task_type] = {
                    'pending': 0, 'processing': 0, 'completed': 0, 'failed': 0, 'total': 0
                }
            by_type[task.task_type]['total'] += 1
            if task.status == PendingTask.STATUS_PENDING:
                by_type[task.task_type]['pending'] += 1
            elif task.status == PendingTask.STATUS_PROCESSING:
                by_type[task.task_type]['processing'] += 1
            elif task.status == PendingTask.STATUS_COMPLETED:
                by_type[task.task_type]['completed'] += 1
            elif task.status == PendingTask.STATUS_FAILED:
                by_type[task.task_type]['failed'] += 1
        
        row_idx = 2
        for task_type_name, stats in by_type.items():
            summary_sheet.cell(row=row_idx, column=1, value=task_type_name)
            summary_sheet.cell(row=row_idx, column=2, value=stats['pending'])
            summary_sheet.cell(row=row_idx, column=3, value=stats['processing'])
            summary_sheet.cell(row=row_idx, column=4, value=stats['completed'])
            summary_sheet.cell(row=row_idx, column=5, value=stats['failed'])
            summary_sheet.cell(row=row_idx, column=6, value=stats['total'])
            row_idx += 1
        
        ExportService._auto_width(summary_sheet)
        
        detail_headers = [
            'ID', '任务类型', '账户ID', '计划ID', '交易ID', '标题', '描述',
            '状态', '重试次数', '最后错误', '计划时间', '开始时间', '完成时间'
        ]
        ExportService._apply_header_style(details_sheet, detail_headers)
        
        for row_idx, task in enumerate(tasks, 2):
            details_sheet.cell(row=row_idx, column=1, value=task.id)
            details_sheet.cell(row=row_idx, column=2, value=task.task_type)
            details_sheet.cell(row=row_idx, column=3, value=task.account_id or '')
            details_sheet.cell(row=row_idx, column=4, value=task.plan_id or '')
            details_sheet.cell(row=row_idx, column=5, value=task.transaction_id or '')
            details_sheet.cell(row=row_idx, column=6, value=task.title)
            details_sheet.cell(row=row_idx, column=7, value=task.description or '')
            details_sheet.cell(row=row_idx, column=8, value=task.status)
            details_sheet.cell(row=row_idx, column=9, value=task.retry_count)
            details_sheet.cell(row=row_idx, column=10, value=task.last_error or '')
            details_sheet.cell(row=row_idx, column=11, value=task.scheduled_at.strftime('%Y-%m-%d %H:%M:%S'))
            details_sheet.cell(row=row_idx, column=12, value=task.started_at.strftime('%Y-%m-%d %H:%M:%S') if task.started_at else '')
            details_sheet.cell(row=row_idx, column=13, value=task.completed_at.strftime('%Y-%m-%d %H:%M:%S') if task.completed_at else '')
        
        ExportService._auto_width(details_sheet)
        
        wb.save(file_path)
        return file_path
