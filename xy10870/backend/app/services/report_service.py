from datetime import datetime, timedelta
from sqlalchemy.orm import Session
from sqlalchemy import func, desc
from typing import List, Dict
import io
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from ..models.models import RunRequest, Student, LanguageEnvironment, RequestStatus, ResultSummary


class ReportService:
    @staticmethod
    def generate_report_data(db: Session, days: int = 7) -> Dict:
        end_date = datetime.utcnow()
        start_date = end_date - timedelta(days=days)

        total_requests = db.query(RunRequest).filter(
            RunRequest.created_at >= start_date
        ).count()

        successful = db.query(RunRequest).filter(
            RunRequest.created_at >= start_date,
            RunRequest.status == RequestStatus.SUCCESS
        ).count()

        success_rate = (successful / total_requests * 100) if total_requests > 0 else 0

        avg_exec_time = db.query(func.avg(RunRequest.execution_time_ms)).filter(
            RunRequest.created_at >= start_date,
            RunRequest.execution_time_ms.isnot(None)
        ).scalar() or 0

        student_stats = db.query(
            Student.name,
            func.count(RunRequest.id).label('count')
        ).join(RunRequest).filter(
            RunRequest.created_at >= start_date
        ).group_by(Student.id).order_by(desc('count')).limit(5).all()

        top_students = [{"name": s[0], "count": s[1]} for s in student_stats]

        lang_stats = db.query(
            LanguageEnvironment.name,
            func.count(RunRequest.id).label('count')
        ).join(RunRequest).filter(
            RunRequest.created_at >= start_date
        ).group_by(LanguageEnvironment.id).all()

        language_distribution = [{"language": l[0], "count": l[1]} for l in lang_stats]

        daily_stats = []
        for i in range(days):
            day_start = (end_date - timedelta(days=days - 1 - i)).replace(hour=0, minute=0, second=0, microsecond=0)
            day_end = day_start + timedelta(days=1)

            day_total = db.query(RunRequest).filter(
                RunRequest.created_at >= day_start,
                RunRequest.created_at < day_end
            ).count()

            day_success = db.query(RunRequest).filter(
                RunRequest.created_at >= day_start,
                RunRequest.created_at < day_end,
                RunRequest.status == RequestStatus.SUCCESS
            ).count()

            daily_stats.append({
                "date": day_start.strftime("%Y-%m-%d"),
                "total": day_total,
                "success": day_success,
                "success_rate": (day_success / day_total * 100) if day_total > 0 else 0
            })

        return {
            "date_range": f"{start_date.strftime('%Y-%m-%d')} ~ {end_date.strftime('%Y-%m-%d')}",
            "total_requests": total_requests,
            "success_rate": round(success_rate, 2),
            "avg_execution_time": round(avg_exec_time, 2),
            "top_students": top_students,
            "language_distribution": language_distribution,
            "daily_stats": daily_stats
        }

    @staticmethod
    def export_to_csv(db: Session) -> bytes:
        requests = db.query(RunRequest).order_by(desc(RunRequest.created_at)).all()

        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow([
            '请求ID', '学生ID', '学生姓名', '语言', '状态',
            '创建时间', '完成时间', '执行时间(ms)', '退出码', '错误信息'
        ])

        for req in requests:
            writer.writerow([
                req.request_id,
                req.student.student_id if req.student else '',
                req.student.name if req.student else '',
                req.language.name if req.language else '',
                req.status,
                req.created_at.strftime('%Y-%m-%d %H:%M:%S') if req.created_at else '',
                req.completed_at.strftime('%Y-%m-%d %H:%M:%S') if req.completed_at else '',
                req.execution_time_ms or '',
                req.exit_code or '',
                req.error_message or ''
            ])

        return output.getvalue().encode('utf-8-sig')

    @staticmethod
    def export_to_excel(db: Session) -> bytes:
        wb = Workbook()

        ws1 = wb.active
        ws1.title = "运行请求"

        headers = [
            '请求ID', '学生ID', '学生姓名', '语言', '状态',
            '创建时间', '完成时间', '执行时间(ms)', '退出码', '错误信息'
        ]

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        requests = db.query(RunRequest).order_by(desc(RunRequest.created_at)).all()

        status_colors = {
            RequestStatus.SUCCESS: "C6EFCE",
            RequestStatus.FAILED: "FFC7CE",
            RequestStatus.TIMEOUT: "FFEB9C",
            RequestStatus.MERGED: "E2EFDA",
            RequestStatus.CANCELLED: "D9D9D9"
        }

        for row, req in enumerate(requests, 2):
            ws1.cell(row=row, column=1, value=req.request_id)
            ws1.cell(row=row, column=2, value=req.student.student_id if req.student else '')
            ws1.cell(row=row, column=3, value=req.student.name if req.student else '')
            ws1.cell(row=row, column=4, value=req.language.name if req.language else '')
            status_cell = ws1.cell(row=row, column=5, value=req.status)
            if req.status in status_colors:
                status_cell.fill = PatternFill(start_color=status_colors[req.status], fill_type="solid")
            ws1.cell(row=row, column=6, value=req.created_at.strftime('%Y-%m-%d %H:%M:%S') if req.created_at else '')
            ws1.cell(row=row, column=7, value=req.completed_at.strftime('%Y-%m-%d %H:%M:%S') if req.completed_at else '')
            ws1.cell(row=row, column=8, value=req.execution_time_ms or '')
            ws1.cell(row=row, column=9, value=req.exit_code or '')
            ws1.cell(row=row, column=10, value=req.error_message or '')

        for col in range(1, 11):
            ws1.column_dimensions[chr(64 + col)].width = 18

        ws2 = wb.create_sheet("统计概览")

        report_data = ReportService.generate_report_data(db, days=7)

        ws2.cell(row=1, column=1, value="代码运行平台统计报告").font = Font(bold=True, size=14)
        ws2.merge_cells('A1:D1')

        stats_data = [
            ["统计周期", report_data["date_range"]],
            ["总请求数", report_data["total_requests"]],
            ["成功率(%)", report_data["success_rate"]],
            ["平均执行时间(ms)", report_data["avg_execution_time"]]
        ]

        for row, (key, value) in enumerate(stats_data, 3):
            ws2.cell(row=row, column=1, value=key).font = Font(bold=True)
            ws2.cell(row=row, column=2, value=value)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def update_daily_summary(db: Session):
        today = datetime.utcnow().date()
        start_time = datetime.combine(today, datetime.min.time())
        end_time = start_time + timedelta(days=1)

        total = db.query(RunRequest).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time
        ).count()

        successful = db.query(RunRequest).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.status == RequestStatus.SUCCESS
        ).count()

        failed = db.query(RunRequest).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.status == RequestStatus.FAILED
        ).count()

        timed_out = db.query(RunRequest).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.status == RequestStatus.TIMEOUT
        ).count()

        merged = db.query(RunRequest).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.status == RequestStatus.MERGED
        ).count()

        avg_exec_time = db.query(func.avg(RunRequest.execution_time_ms)).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.execution_time_ms.isnot(None)
        ).scalar() or 0

        total_exec_time = db.query(func.sum(RunRequest.execution_time_ms)).filter(
            RunRequest.created_at >= start_time,
            RunRequest.created_at < end_time,
            RunRequest.execution_time_ms.isnot(None)
        ).scalar() or 0

        summary = db.query(ResultSummary).filter(ResultSummary.date == start_time).first()

        if not summary:
            summary = ResultSummary(
                date=start_time,
                total_requests=total,
                successful=successful,
                failed=failed,
                timed_out=timed_out,
                merged=merged,
                avg_execution_time_ms=avg_exec_time,
                total_execution_time_ms=total_exec_time,
                peak_concurrent_containers=0,
                quota_exceeded_count=0
            )
            db.add(summary)
        else:
            summary.total_requests = total
            summary.successful = successful
            summary.failed = failed
            summary.timed_out = timed_out
            summary.merged = merged
            summary.avg_execution_time_ms = avg_exec_time
            summary.total_execution_time_ms = total_exec_time

        db.commit()
