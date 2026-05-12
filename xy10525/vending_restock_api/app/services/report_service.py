from sqlalchemy.orm import Session
from datetime import datetime, timedelta, date
from typing import List
from sqlalchemy import and_

from app.models.models import (
    Route, RestockTask, TaskItem, Inventory, Fault, Product, Machine,
    RouteStatus, TaskStatus, FaultStatus
)
from app.schemas.schemas import (
    DailyReport, StockoutRisk, RecoveryItem, ExecutionDiff, FaultResponse
)
from app.config import settings


class ReportService:
    def __init__(self, db: Session):
        self.db = db

    def generate_daily_report(self, report_date: date) -> DailyReport:
        dt = datetime.combine(report_date, datetime.min.time())
        
        routes = self.db.query(Route).filter(
            and_(
                Route.scheduled_date >= dt,
                Route.scheduled_date <= dt.replace(hour=23, minute=59, second=59)
            )
        ).all()
        
        total_routes = len(routes)
        completed_routes = sum(1 for r in routes if r.status == RouteStatus.COMPLETED)
        
        all_tasks = []
        for route in routes:
            all_tasks.extend(route.tasks)
        
        total_tasks = len(all_tasks)
        completed_tasks = sum(1 for t in all_tasks if t.status == TaskStatus.COMPLETED)
        
        total_restocked = 0
        total_recovered = 0
        execution_diffs: List[ExecutionDiff] = []
        
        for task in all_tasks:
            for item in task.items:
                actual = item.actual_quantity or 0
                if item.item_type == "restock":
                    total_restocked += actual
                elif item.item_type == "recovery":
                    total_recovered += actual
                
                if item.actual_quantity is not None:
                    diff = item.actual_quantity - item.requested_quantity
                    if diff != 0:
                        execution_diffs.append(ExecutionDiff(
                            task_id=task.id,
                            machine_id=task.machine_id,
                            item_type=item.item_type,
                            product_id=item.product_id,
                            requested_quantity=item.requested_quantity,
                            actual_quantity=item.actual_quantity,
                            difference=diff
                        ))
        
        stockout_risks = self._analyze_stockout_risks()
        recovery_items = self._analyze_recovery_items()
        unresolved_faults = self._get_unresolved_faults()
        
        return DailyReport(
            report_date=report_date,
            total_routes=total_routes,
            completed_routes=completed_routes,
            total_tasks=total_tasks,
            completed_tasks=completed_tasks,
            total_quantity_restocked=total_restocked,
            total_quantity_recovered=total_recovered,
            stockout_risks=stockout_risks,
            recovery_items=recovery_items,
            execution_diffs=execution_diffs,
            unresolved_faults=unresolved_faults
        )

    def _analyze_stockout_risks(self) -> List[StockoutRisk]:
        risks: List[StockoutRisk] = []
        
        inventories = self.db.query(Inventory).all()
        for inv in inventories:
            if inv.quantity < inv.min_level:
                machine = self.db.query(Machine).filter(Machine.id == inv.machine_id).first()
                product = self.db.query(Product).filter(Product.id == inv.product_id).first()
                
                shortage = inv.min_level - inv.quantity
                
                if shortage < 5:
                    risk_level = "LOW"
                elif shortage < 15:
                    risk_level = "MEDIUM"
                else:
                    risk_level = "HIGH"
                
                risks.append(StockoutRisk(
                    machine_id=inv.machine_id,
                    machine_name=machine.name if machine else inv.machine_id,
                    product_id=inv.product_id,
                    product_name=product.name if product else inv.product_id,
                    current_quantity=inv.quantity,
                    min_level=inv.min_level,
                    shortage=shortage,
                    risk_level=risk_level
                ))
        
        risks.sort(key=lambda x: {"LOW": 0, "MEDIUM": 1, "HIGH": 2}[x.risk_level], reverse=True)
        return risks

    def _analyze_recovery_items(self) -> List[RecoveryItem]:
        items: List[RecoveryItem] = []
        expiry_threshold = datetime.utcnow() + timedelta(days=settings.EXPIRE_WARNING_DAYS)
        
        inventories = self.db.query(Inventory).filter(
            Inventory.expiry_date.isnot(None),
            Inventory.expiry_date <= expiry_threshold,
            Inventory.quantity > 0
        ).all()
        
        for inv in inventories:
            machine = self.db.query(Machine).filter(Machine.id == inv.machine_id).first()
            product = self.db.query(Product).filter(Product.id == inv.product_id).first()
            
            days_until_expiry = (inv.expiry_date - datetime.utcnow()).days
            
            items.append(RecoveryItem(
                machine_id=inv.machine_id,
                machine_name=machine.name if machine else inv.machine_id,
                product_id=inv.product_id,
                product_name=product.name if product else inv.product_id,
                quantity=inv.quantity,
                expiry_date=inv.expiry_date,
                days_until_expiry=days_until_expiry
            ))
        
        items.sort(key=lambda x: x.days_until_expiry)
        return items

    def _get_unresolved_faults(self) -> List[FaultResponse]:
        faults = self.db.query(Fault).filter(
            Fault.status.in_([FaultStatus.OPEN, FaultStatus.IN_PROGRESS])
        ).order_by(Fault.priority.desc()).all()
        return [FaultResponse.model_validate(f) for f in faults]

    def export_report_excel(self, report: DailyReport, filepath: str):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        wb = Workbook()
        
        ws_summary = wb.active
        ws_summary.title = "Summary"
        self._write_summary(ws_summary, report)
        
        ws_risks = wb.create_sheet("Stockout Risks")
        self._write_risks(ws_risks, report.stockout_risks)
        
        ws_recovery = wb.create_sheet("Recovery Items")
        self._write_recovery(ws_recovery, report.recovery_items)
        
        ws_diffs = wb.create_sheet("Execution Diffs")
        self._write_diffs(ws_diffs, report.execution_diffs)
        
        ws_faults = wb.create_sheet("Unresolved Faults")
        self._write_faults(ws_faults, report.unresolved_faults)
        
        wb.save(filepath)
        return filepath

    def _write_summary(self, ws, report: DailyReport):
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        ws["A1"] = "Daily Restock Report"
        ws["A1"].font = Font(size=16, bold=True)
        ws["A2"] = f"Date: {report.report_date}"
        
        data = [
            ["Metric", "Value"],
            ["Total Routes", report.total_routes],
            ["Completed Routes", report.completed_routes],
            ["Route Completion Rate", f"{(report.completed_routes/report.total_routes*100):.1f}%" if report.total_routes else "0%"],
            ["", ""],
            ["Total Tasks", report.total_tasks],
            ["Completed Tasks", report.completed_tasks],
            ["Task Completion Rate", f"{(report.completed_tasks/report.total_tasks*100):.1f}%" if report.total_tasks else "0%"],
            ["", ""],
            ["Total Quantity Restocked", report.total_quantity_restocked],
            ["Total Quantity Recovered", report.total_quantity_recovered],
        ]
        
        for i, row in enumerate(data, start=4):
            for j, cell_value in enumerate(row, start=1):
                cell = ws.cell(row=i, column=j, value=cell_value)
                if i == 4:
                    cell.font = header_font
                    cell.fill = header_fill

    def _write_risks(self, ws, risks: List[StockoutRisk]):
        headers = ["Machine ID", "Machine Name", "Product ID", "Product Name", 
                   "Current Qty", "Min Level", "Shortage", "Risk Level"]
        self._write_headers(ws, headers)
        
        for i, risk in enumerate(risks, start=2):
            ws.cell(row=i, column=1, value=risk.machine_id)
            ws.cell(row=i, column=2, value=risk.machine_name)
            ws.cell(row=i, column=3, value=risk.product_id)
            ws.cell(row=i, column=4, value=risk.product_name)
            ws.cell(row=i, column=5, value=risk.current_quantity)
            ws.cell(row=i, column=6, value=risk.min_level)
            ws.cell(row=i, column=7, value=risk.shortage)
            ws.cell(row=i, column=8, value=risk.risk_level)

    def _write_recovery(self, ws, items: List[RecoveryItem]):
        headers = ["Machine ID", "Machine Name", "Product ID", "Product Name",
                   "Quantity", "Expiry Date", "Days Until Expiry"]
        self._write_headers(ws, headers)
        
        for i, item in enumerate(items, start=2):
            ws.cell(row=i, column=1, value=item.machine_id)
            ws.cell(row=i, column=2, value=item.machine_name)
            ws.cell(row=i, column=3, value=item.product_id)
            ws.cell(row=i, column=4, value=item.product_name)
            ws.cell(row=i, column=5, value=item.quantity)
            ws.cell(row=i, column=6, value=item.expiry_date.strftime("%Y-%m-%d"))
            ws.cell(row=i, column=7, value=item.days_until_expiry)

    def _write_diffs(self, ws, diffs: List[ExecutionDiff]):
        headers = ["Task ID", "Machine ID", "Item Type", "Product ID",
                   "Requested Qty", "Actual Qty", "Difference"]
        self._write_headers(ws, headers)
        
        for i, diff in enumerate(diffs, start=2):
            ws.cell(row=i, column=1, value=diff.task_id)
            ws.cell(row=i, column=2, value=diff.machine_id)
            ws.cell(row=i, column=3, value=diff.item_type)
            ws.cell(row=i, column=4, value=diff.product_id)
            ws.cell(row=i, column=5, value=diff.requested_quantity)
            ws.cell(row=i, column=6, value=diff.actual_quantity)
            ws.cell(row=i, column=7, value=diff.difference)

    def _write_faults(self, ws, faults: List[FaultResponse]):
        headers = ["Fault ID", "Machine ID", "Fault Type", "Status", 
                   "Priority", "Description", "Reported At"]
        self._write_headers(ws, headers)
        
        for i, fault in enumerate(faults, start=2):
            ws.cell(row=i, column=1, value=fault.id)
            ws.cell(row=i, column=2, value=fault.machine_id)
            ws.cell(row=i, column=3, value=fault.fault_type)
            ws.cell(row=i, column=4, value=fault.status.value)
            ws.cell(row=i, column=5, value=fault.priority)
            ws.cell(row=i, column=6, value=fault.description)
            ws.cell(row=i, column=7, value=fault.reported_at.strftime("%Y-%m-%d %H:%M"))

    def _write_headers(self, ws, headers: List[str]):
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        for i, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=i, value=header)
            cell.font = header_font
            cell.fill = header_fill
