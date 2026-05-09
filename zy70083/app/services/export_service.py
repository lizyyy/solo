from typing import List, Optional, Dict, Tuple
from sqlalchemy.orm import Session
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from ..models import (
    ApplicationRecord, FamilyMember, VerificationRecord,
    ProcessingHistory, LotteryPool, ApplicationStatus
)


class ExportService:
    @staticmethod
    def generate_application_excel(
        db: Session,
        pool_id: Optional[str] = None,
        status_filter: Optional[str] = None
    ) -> bytes:
        applications = db.query(ApplicationRecord)
        if pool_id:
            applications = applications.filter(ApplicationRecord.lottery_pool_id == pool_id)
        if status_filter:
            applications = applications.filter(ApplicationRecord.current_status == status_filter)
        applications = applications.all()

        wb = Workbook()
        ws = wb.active
        ws.title = "公租房摇号资格结果"

        headers = [
            "申请编号", "申请人姓名", "身份证号", "联系电话",
            "家庭住址", "家庭人数", "当前状态", "资格规则",
            "收入核验结果", "社保核验结果", "住房核验结果",
            "摇号池", "最终结果", "申请时间", "最后更新时间"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        status_map = {
            ApplicationStatus.CREATED.value: "已创建",
            ApplicationStatus.UNDER_REVIEW.value: "审核中",
            ApplicationStatus.INCOME_VERIFICATION.value: "收入核验中",
            ApplicationStatus.SOCIAL_INSURANCE_VERIFICATION.value: "社保核验中",
            ApplicationStatus.HOUSING_STATUS_VERIFICATION.value: "住房核验中",
            ApplicationStatus.VERIFIED.value: "核验通过",
            ApplicationStatus.REJECTED.value: "已拒绝",
            ApplicationStatus.LOTTERY_LOCKED.value: "已锁定",
            ApplicationStatus.PUBLIC_ANNOUNCEMENT.value: "公示中",
            ApplicationStatus.OBJECTION_RAISED.value: "有异议待处理",
            ApplicationStatus.OBJECTION_RESOLVED.value: "异议已处理",
            ApplicationStatus.FINAL_RESULT.value: "最终结果"
        }

        for row_idx, app in enumerate(applications, 2):
            family_size = len(app.family_members)
            rule_name = app.applied_rule.rule_name if app.applied_rule else "默认规则"

            income_result = ExportService._get_verification_result(app, "income")
            insurance_result = ExportService._get_verification_result(app, "social_insurance")
            housing_result = ExportService._get_verification_result(app, "housing_status")

            final_result_map = {"qualified": "符合资格", "not_qualified": "不符合资格"}

            data = [
                app.application_number,
                app.applicant_name,
                app.applicant_id_card,
                app.contact_phone,
                app.household_address,
                family_size,
                status_map.get(app.current_status, app.current_status),
                rule_name,
                income_result,
                insurance_result,
                housing_result,
                app.lottery_pool_id or "-",
                final_result_map.get(app.final_result, app.final_result or "-"),
                app.created_at.strftime("%Y-%m-%d %H:%M:%S") if app.created_at else "-",
                app.updated_at.strftime("%Y-%m-%d %H:%M:%S") if app.updated_at else "-"
            ]

            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = center_align

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col) if col <= 26 else 'A' + chr(64 + col - 26)].width = 18

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def _get_verification_result(application: ApplicationRecord, verification_type: str) -> str:
        record = next(
            (v for v in application.verification_records if v.verification_type == verification_type),
            None
        )
        if not record:
            return "未核验"
        result_map = {"passed": "通过", "failed": "未通过", "pending": "核验中"}
        return result_map.get(record.verification_result, record.verification_result or "-")

    @staticmethod
    def generate_family_detail_excel(
        db: Session,
        application_id: int
    ) -> Optional[bytes]:
        application = db.query(ApplicationRecord).filter(
            ApplicationRecord.id == application_id
        ).first()
        if not application:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "家庭成员详情"

        headers = [
            "序号", "姓名", "身份证号", "与申请人关系", "是否主申请人",
            "月收入(元)", "社保缴纳月数", "住房贡献面积(㎡)"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for idx, member in enumerate(application.family_members, 2):
            data = [
                idx - 1,
                member.name,
                member.id_card,
                member.relation,
                "是" if member.is_main_applicant else "否",
                member.monthly_income or 0,
                member.social_insurance_months or 0,
                member.housing_area_contribution or 0
            ]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx, column=col, value=value)
                cell.alignment = center_align

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col) if col <= 26 else 'A' + chr(64 + col - 26)].width = 18

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_processing_history_excel(
        db: Session,
        application_id: int
    ) -> Optional[bytes]:
        application = db.query(ApplicationRecord).filter(
            ApplicationRecord.id == application_id
        ).first()
        if not application:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "处理历史记录"

        headers = [
            "序号", "时间", "操作", "原状态", "新状态", "卡点", "操作人", "备注"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for idx, history in enumerate(reversed(application.processing_history), 2):
            data = [
                idx - 1,
                history.created_at.strftime("%Y-%m-%d %H:%M:%S") if history.created_at else "-",
                history.action,
                history.previous_status or "-",
                history.new_status,
                history.checkpoint or "-",
                history.operator or "-",
                history.remark or "-"
            ]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx, column=col, value=value)
                cell.alignment = center_align

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col) if col <= 26 else 'A' + chr(64 + col - 26)].width = 20

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_pool_summary_excel(
        db: Session,
        pool_id: str
    ) -> Optional[bytes]:
        pool = db.query(LotteryPool).filter(LotteryPool.pool_id == pool_id).first()
        if not pool:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "摇号池汇总"

        ws.cell(row=1, column=1, value="摇号池信息").font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value="摇号池编号:").font = Font(bold=True)
        ws.cell(row=2, column=2, value=pool.pool_id)
        ws.cell(row=3, column=1, value="摇号池名称:").font = Font(bold=True)
        ws.cell(row=3, column=2, value=pool.pool_name)
        ws.cell(row=4, column=1, value="年份/批次:").font = Font(bold=True)
        ws.cell(row=4, column=2, value=f"{pool.lottery_year}年第{pool.lottery_batch}批")
        ws.cell(row=5, column=1, value="总配额:").font = Font(bold=True)
        ws.cell(row=5, column=2, value=pool.total_quota)
        ws.cell(row=6, column=1, value="已锁定人数:").font = Font(bold=True)
        ws.cell(row=6, column=2, value=pool.locked_count)
        ws.cell(row=7, column=1, value="已选中人数:").font = Font(bold=True)
        ws.cell(row=7, column=2, value=pool.selected_count)
        ws.cell(row=8, column=1, value="状态:").font = Font(bold=True)
        ws.cell(row=8, column=2, value="活跃" if pool.is_active else "已关闭")

        applications = db.query(ApplicationRecord).filter(
            ApplicationRecord.lottery_pool_id == pool_id
        ).all()

        status_counts = {}
        for app in applications:
            status = app.current_status
            status_counts[status] = status_counts.get(status, 0) + 1

        ws.cell(row=10, column=1, value="各状态统计:").font = Font(bold=True)
        row_idx = 11
        for status, count in status_counts.items():
            ws.cell(row=row_idx, column=1, value=status)
            ws.cell(row=row_idx, column=2, value=count)
            row_idx += 1

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
