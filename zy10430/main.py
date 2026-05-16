from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import JSONResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Any
from datetime import datetime
from sqlalchemy.orm import Session
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from models import (
    get_db, init_db, Application, WhitelistLabel, BlockRecord, Status
)

app = FastAPI(title="指标基数护栏API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_CARDINALITY = 10000
CARDINALITY_WARNING = 5000

class ApplicationCreate(BaseModel):
    metric_name: str = Field(..., description="指标名称")
    labels: Dict[str, str] = Field(..., description="标签键值对")
    reason: str = Field(..., description="申请理由")
    estimated_values: Optional[Dict[str, int]] = Field(None, description="各标签可能值数量")

class ApplicationResponse(BaseModel):
    id: int
    metric_name: str
    labels: Dict[str, str]
    estimated_cardinality: int
    reason: str
    status: str
    created_at: str
    block_reason: Optional[str]
    process_result: Optional[Dict[str, Any]]
    raw_input: Optional[Dict[str, Any]]

class ReviewRequest(BaseModel):
    reviewer: str = Field(..., description="审核人")
    comment: Optional[str] = Field(None, description="审核意见")

class WhitelistAddRequest(BaseModel):
    label_key: str = Field(..., description="标签键")
    description: Optional[str] = Field(None, description="描述")
    created_by: str = Field("admin", description="创建人")

def estimate_cardinality(labels: Dict[str, str], estimated_values: Optional[Dict[str, int]] = None) -> int:
    if estimated_values:
        cardinality = 1
        for key in labels.keys():
            cardinality *= estimated_values.get(key, 10)
        return cardinality
    else:
        return 10 ** len(labels)

def validate_labels(labels: Dict[str, str], db: Session) -> tuple:
    whitelist = db.query(WhitelistLabel).all()
    whitelist_keys = {w.label_key for w in whitelist}
    invalid_labels = []
    for key in labels.keys():
        if key not in whitelist_keys:
            invalid_labels.append(key)
    return len(invalid_labels) == 0, invalid_labels

def check_duplicate_application(metric_name: str, labels: Dict[str, str], db: Session) -> Optional[Application]:
    labels_json = json.dumps(labels, sort_keys=True)
    applications = db.query(Application).filter(Application.metric_name == metric_name).all()
    for app in applications:
        if json.dumps(app.labels, sort_keys=True) == labels_json:
            return app
    return None

def process_application(application: Application, db: Session):
    process_result = {
        "checks": [],
        "passed": True,
        "final_status": Status.PENDING
    }
    
    labels_valid, invalid_labels = validate_labels(application.labels, db)
    labels_check = {
        "check": "label_whitelist",
        "passed": labels_valid,
        "message": "所有标签在白名单内" if labels_valid else f"存在未授权标签: {invalid_labels}"
    }
    process_result["checks"].append(labels_check)
    if not labels_valid:
        process_result["passed"] = False
    
    cardinality = application.estimated_cardinality
    cardinality_level = "normal"
    if cardinality >= MAX_CARDINALITY:
        cardinality_level = "critical"
        cardinality_check = {
            "check": "cardinality_estimation",
            "passed": False,
            "message": f"基数估算过高 ({cardinality})，超过最大值 {MAX_CARDINALITY}",
            "cardinality": cardinality
        }
        process_result["passed"] = False
    elif cardinality >= CARDINALITY_WARNING:
        cardinality_level = "warning"
        cardinality_check = {
            "check": "cardinality_estimation",
            "passed": True,
            "message": f"基数估算较高 ({cardinality})，建议确认",
            "cardinality": cardinality
        }
    else:
        cardinality_check = {
            "check": "cardinality_estimation",
            "passed": True,
            "message": f"基数估算正常 ({cardinality})",
            "cardinality": cardinality
        }
    process_result["checks"].append(cardinality_check)
    
    if not process_result["passed"]:
        process_result["final_status"] = Status.BLOCKED
        application.status = Status.BLOCKED
        application.block_reason = "; ".join([c["message"] for c in process_result["checks"] if not c["passed"]])
        
        block_record = BlockRecord(
            application_id=application.id,
            metric_name=application.metric_name,
            block_type="auto_block",
            block_detail=process_result
        )
        db.add(block_record)
    else:
        process_result["final_status"] = Status.APPROVED
        application.status = Status.APPROVED
    
    application.process_result = process_result
    db.commit()
    db.refresh(application)
    return application

@app.on_event("startup")
async def startup_event():
    init_db()

@app.post("/api/v1/applications", response_model=ApplicationResponse)
async def create_application(
    request: ApplicationCreate,
    db: Session = Depends(get_db)
):
    raw_input = request.dict()
    
    existing = check_duplicate_application(request.metric_name, request.labels, db)
    if existing:
        return ApplicationResponse(
            id=existing.id,
            metric_name=existing.metric_name,
            labels=existing.labels,
            estimated_cardinality=existing.estimated_cardinality,
            reason=existing.reason,
            status=existing.status,
            created_at=existing.created_at.isoformat(),
            block_reason=existing.block_reason,
            process_result=existing.process_result,
            raw_input=existing.raw_input
        )
    
    estimated_cardinality = estimate_cardinality(request.labels, request.estimated_values)
    
    application = Application(
        metric_name=request.metric_name,
        labels=request.labels,
        estimated_cardinality=estimated_cardinality,
        reason=request.reason,
        raw_input=raw_input
    )
    
    db.add(application)
    db.commit()
    db.refresh(application)
    
    application = process_application(application, db)
    
    return ApplicationResponse(
        id=application.id,
        metric_name=application.metric_name,
        labels=application.labels,
        estimated_cardinality=application.estimated_cardinality,
        reason=application.reason,
        status=application.status,
        created_at=application.created_at.isoformat(),
        block_reason=application.block_reason,
        process_result=application.process_result,
        raw_input=application.raw_input
    )

@app.get("/api/v1/applications")
async def list_applications(
    status: Optional[str] = None,
    metric_name: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    query = db.query(Application)
    if status:
        query = query.filter(Application.status == status)
    if metric_name:
        query = query.filter(Application.metric_name.contains(metric_name))
    
    total = query.count()
    applications = query.order_by(Application.created_at.desc()).offset(skip).limit(limit).all()
    
    return {
        "total": total,
        "items": [app.to_dict() for app in applications]
    }

@app.get("/api/v1/applications/{app_id}")
async def get_application(
    app_id: int,
    db: Session = Depends(get_db)
):
    application = db.query(Application).filter(Application.id == app_id).first()
    if not application:
        raise HTTPException(status_code=404, detail="申请不存在")
    return application.to_dict()

@app.post("/api/v1/applications/{app_id}/approve")
async def approve_application(
    app_id: int,
    request: ReviewRequest,
    db: Session = Depends(get_db)
):
    application = db.query(Application).filter(Application.id == app_id).first()
    if not application:
        raise HTTPException(status_code=404, detail="申请不存在")
    
    if application.status not in [Status.PENDING, Status.BLOCKED]:
        raise HTTPException(
            status_code=400, 
            detail=f"当前状态 {application.status} 不允许人工批准，仅 PENDING 或 BLOCKED 状态可操作"
        )
    
    application.status = Status.APPROVED
    application.reviewed_by = request.reviewer
    application.review_comment = request.comment or "人工放行"
    application.reviewed_at = datetime.utcnow()
    application.block_reason = None
    
    if application.process_result:
        application.process_result["final_status"] = Status.APPROVED
        application.process_result["manual_review"] = {
            "action": "approve",
            "reviewer": request.reviewer,
            "comment": request.comment
        }
    
    db.commit()
    db.refresh(application)
    return application.to_dict()

@app.post("/api/v1/applications/{app_id}/reject")
async def reject_application(
    app_id: int,
    request: ReviewRequest,
    db: Session = Depends(get_db)
):
    application = db.query(Application).filter(Application.id == app_id).first()
    if not application:
        raise HTTPException(status_code=404, detail="申请不存在")
    
    if application.status not in [Status.PENDING, Status.APPROVED]:
        raise HTTPException(
            status_code=400, 
            detail=f"当前状态 {application.status} 不允许人工拒绝"
        )
    
    application.status = Status.REJECTED
    application.reviewed_by = request.reviewer
    application.review_comment = request.comment or "人工拒绝"
    application.reviewed_at = datetime.utcnow()
    
    db.commit()
    db.refresh(application)
    return application.to_dict()

@app.post("/api/v1/applications/{app_id}/apply")
async def mark_applied(
    app_id: int,
    db: Session = Depends(get_db)
):
    application = db.query(Application).filter(Application.id == app_id).first()
    if not application:
        raise HTTPException(status_code=404, detail="申请不存在")
    
    if application.status != Status.APPROVED:
        raise HTTPException(
            status_code=400, 
            detail=f"仅 APPROVED 状态可标记为已应用，当前状态: {application.status}"
        )
    
    application.status = Status.APPLIED
    application.applied_at = datetime.utcnow()
    
    db.commit()
    db.refresh(application)
    return application.to_dict()

@app.get("/api/v1/whitelist")
async def list_whitelist(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    total = db.query(WhitelistLabel).count()
    labels = db.query(WhitelistLabel).order_by(WhitelistLabel.created_at.desc()).offset(skip).limit(limit).all()
    return {
        "total": total,
        "items": [label.to_dict() for label in labels]
    }

@app.post("/api/v1/whitelist")
async def add_to_whitelist(
    request: WhitelistAddRequest,
    db: Session = Depends(get_db)
):
    existing = db.query(WhitelistLabel).filter(WhitelistLabel.label_key == request.label_key).first()
    if existing:
        return existing.to_dict()
    
    label = WhitelistLabel(
        label_key=request.label_key,
        description=request.description,
        created_by=request.created_by
    )
    db.add(label)
    db.commit()
    db.refresh(label)
    return label.to_dict()

@app.get("/api/v1/report")
async def export_report(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    format: str = Query("json", enum=["json", "xlsx"]),
    db: Session = Depends(get_db)
):
    query = db.query(Application)
    
    if start_date:
        query = query.filter(Application.created_at >= datetime.fromisoformat(start_date))
    if end_date:
        query = query.filter(Application.created_at <= datetime.fromisoformat(end_date))
    
    applications = query.order_by(Application.created_at.desc()).all()
    block_records = db.query(BlockRecord).all()
    
    report_data = {
        "summary": {
            "total_applications": len(applications),
            "blocked_count": len([a for a in applications if a.status == Status.BLOCKED]),
            "approved_count": len([a for a in applications if a.status == Status.APPROVED]),
            "rejected_count": len([a for a in applications if a.status == Status.REJECTED]),
            "applied_count": len([a for a in applications if a.status == Status.APPLIED]),
            "pending_count": len([a for a in applications if a.status == Status.PENDING]),
            "generated_at": datetime.utcnow().isoformat()
        },
        "applications": [app.to_dict() for app in applications],
        "block_records": [br.to_dict() for br in block_records]
    }
    
    if format == "xlsx":
        wb = Workbook()
        
        ws_summary = wb.active
        ws_summary.title = "汇总"
        headers = ["统计项", "数量"]
        for col, header in enumerate(headers, 1):
            cell = ws_summary.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        summary_items = [
            ("总申请数", report_data["summary"]["total_applications"]),
            ("拦截数", report_data["summary"]["blocked_count"]),
            ("批准数", report_data["summary"]["approved_count"]),
            ("拒绝数", report_data["summary"]["rejected_count"]),
            ("已应用数", report_data["summary"]["applied_count"]),
            ("待处理数", report_data["summary"]["pending_count"])
        ]
        for row, (key, value) in enumerate(summary_items, 2):
            ws_summary.cell(row=row, column=1, value=key)
            ws_summary.cell(row=row, column=2, value=value)
        
        ws_apps = wb.create_sheet("申请明细")
        headers = ["ID", "指标名", "标签", "估算基数", "状态", "创建时间", "拦截原因", "审核人", "审核意见"]
        for col, header in enumerate(headers, 1):
            cell = ws_apps.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        for row, app in enumerate(applications, 2):
            ws_apps.cell(row=row, column=1, value=app.id)
            ws_apps.cell(row=row, column=2, value=app.metric_name)
            ws_apps.cell(row=row, column=3, value=json.dumps(app.labels, ensure_ascii=False))
            ws_apps.cell(row=row, column=4, value=app.estimated_cardinality)
            ws_apps.cell(row=row, column=5, value=app.status)
            ws_apps.cell(row=row, column=6, value=app.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws_apps.cell(row=row, column=7, value=app.block_reason or "")
            ws_apps.cell(row=row, column=8, value=app.reviewed_by or "")
            ws_apps.cell(row=row, column=9, value=app.review_comment or "")
            
            if app.status == Status.BLOCKED:
                for col in range(1, 10):
                    ws_apps.cell(row=row, column=col).fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        
        ws_blocks = wb.create_sheet("拦截记录")
        headers = ["ID", "申请ID", "指标名", "拦截类型", "拦截时间"]
        for col, header in enumerate(headers, 1):
            cell = ws_blocks.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        for row, br in enumerate(block_records, 2):
            ws_blocks.cell(row=row, column=1, value=br.id)
            ws_blocks.cell(row=row, column=2, value=br.application_id)
            ws_blocks.cell(row=row, column=3, value=br.metric_name)
            ws_blocks.cell(row=row, column=4, value=br.block_type)
            ws_blocks.cell(row=row, column=5, value=br.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        
        filename = f"guardrail_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = f"/tmp/{filename}"
        wb.save(filepath)
        return FileResponse(filepath, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    
    return report_data

@app.get("/api/v1/stats")
async def get_statistics(db: Session = Depends(get_db)):
    applications = db.query(Application).all()
    block_records = db.query(BlockRecord).all()
    
    status_stats = {}
    for status in [Status.PENDING, Status.BLOCKED, Status.APPROVED, Status.REJECTED, Status.APPLIED]:
        status_stats[status] = len([a for a in applications if a.status == status])
    
    metric_stats = {}
    for app in applications:
        if app.metric_name not in metric_stats:
            metric_stats[app.metric_name] = {"count": 0, "blocked": 0}
        metric_stats[app.metric_name]["count"] += 1
        if app.status == Status.BLOCKED:
            metric_stats[app.metric_name]["blocked"] += 1
    
    return {
        "status_statistics": status_stats,
        "metric_statistics": metric_stats,
        "total_block_records": len(block_records),
        "average_cardinality": sum(a.estimated_cardinality for a in applications) / len(applications) if applications else 0
    }

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
