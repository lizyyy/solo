import random
import uuid
from datetime import datetime, timedelta
from typing import List, Dict, Any, Tuple
from sqlalchemy.orm import Session
from database import DutyRecord, SamplingRule, ProcessingBatch, OperationLog, CleanupCandidate
from schemas import DutyRecordCreate, SamplingRuleCreate, ProcessingBatchCreate, CleanupCandidateCreate


def generate_batch_id(department: str, date_str: str) -> str:
    date_part = date_str.replace("-", "")
    random_str = str(uuid.uuid4())[:8].upper()
    return f"{department.upper()}-{date_part}-{random_str}"


def create_sampling_rule(db: Session, rule_data: SamplingRuleCreate) -> SamplingRule:
    db.query(SamplingRule).filter(
        SamplingRule.department == rule_data.department,
        SamplingRule.is_active == True
    ).update({"is_active": False})
    
    rule = SamplingRule(
        version=rule_data.version,
        name=rule_data.name,
        description=rule_data.description,
        department=rule_data.department,
        sampling_rate=rule_data.sampling_rate,
        min_incidents=rule_data.min_incidents,
        include_departments=rule_data.include_departments,
        exclude_departments=rule_data.exclude_departments,
        receipt_timeout_hours=rule_data.receipt_timeout_hours,
        created_by=rule_data.created_by
    )
    db.add(rule)
    db.commit()
    db.refresh(rule)
    
    log_operation(db, "create_rule", None, None, rule_data.created_by, {
        "rule_version": rule_data.version,
        "rule_name": rule_data.name
    })
    return rule


def get_active_rule(db: Session, department: str) -> SamplingRule:
    rule = db.query(SamplingRule).filter(
        SamplingRule.department == department,
        SamplingRule.is_active == True
    ).first()
    if not rule:
        rule = db.query(SamplingRule).filter(
            SamplingRule.department == "总部",
            SamplingRule.is_active == True
        ).first()
    return rule


def get_rule_by_version(db: Session, version: str) -> SamplingRule:
    return db.query(SamplingRule).filter(SamplingRule.version == version).first()


def create_duty_record(db: Session, record_data: DutyRecordCreate) -> DutyRecord:
    record = DutyRecord(
        batch_id=record_data.batch_id,
        duty_date=record_data.duty_date,
        department=record_data.department,
        duty_person=record_data.duty_person,
        phone=record_data.phone,
        incident_count=record_data.incident_count,
        incidents=record_data.incidents,
        external_receipt_status=record_data.external_receipt_status,
        raw_input=record_data.raw_input
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def batch_create_duty_records(db: Session, records: List[DutyRecordCreate]) -> List[DutyRecord]:
    created_records = []
    for record_data in records:
        record = DutyRecord(
            batch_id=record_data.batch_id,
            duty_date=record_data.duty_date,
            department=record_data.department,
            duty_person=record_data.duty_person,
            phone=record_data.phone,
            incident_count=record_data.incident_count,
            incidents=record_data.incidents,
            external_receipt_status=record_data.external_receipt_status,
            raw_input=record_data.raw_input
        )
        db.add(record)
        created_records.append(record)
    db.commit()
    for record in created_records:
        db.refresh(record)
    return created_records


def sample_call_chain(db: Session, batch_id: str, rule_version: str = None) -> Tuple[List[DutyRecord], Dict[str, Any]]:
    batch = db.query(ProcessingBatch).filter(ProcessingBatch.batch_id == batch_id).first()
    if not batch:
        raise ValueError(f"Batch {batch_id} not found")
    
    if rule_version:
        rule = get_rule_by_version(db, rule_version)
    else:
        rule = get_active_rule(db, batch.department)
    
    if not rule:
        raise ValueError(f"No sampling rule found for department {batch.department}")
    
    all_records = db.query(DutyRecord).filter(DutyRecord.batch_id == batch_id).all()
    
    sampled_records = []
    late_receipt_count = 0
    
    for record in all_records:
        if record.department in rule.exclude_departments:
            continue
        
        if record.department not in rule.include_departments and "*" not in rule.include_departments:
            continue
        
        if record.external_receipt_status == "late":
            late_receipt_count += 1
        
        if record.incident_count >= rule.min_incidents:
            sampled_records.append(record)
            continue
        
        if random.random() < rule.sampling_rate:
            sampled_records.append(record)
    
    summary = {
        "total_records": len(all_records),
        "sampled_count": len(sampled_records),
        "sampling_rate_used": rule.sampling_rate,
        "min_incidents_threshold": rule.min_incidents,
        "late_receipt_count": late_receipt_count,
        "receipt_timeout_hours": rule.receipt_timeout_hours,
        "rule_version": rule.version,
        "rule_name": rule.name
    }
    
    batch.rule_snapshot = {
        "version": rule.version,
        "name": rule.name,
        "description": rule.description,
        "sampling_rate": rule.sampling_rate,
        "min_incidents": rule.min_incidents,
        "include_departments": rule.include_departments,
        "exclude_departments": rule.exclude_departments,
        "receipt_timeout_hours": rule.receipt_timeout_hours
    }
    batch.total_records = len(all_records)
    batch.sampled_count = len(sampled_records)
    batch.summary = summary
    batch.status = "completed"
    batch.completed_at = datetime.utcnow()
    db.commit()
    
    log_operation(db, "sampling", batch_id, None, "system", summary)
    
    return sampled_records, summary


def log_operation(db: Session, operation_type: str, batch_id: str = None, 
                  record_id: int = None, operator: str = "system", 
                  details: Dict[str, Any] = None, status: str = "success",
                  error_message: str = None):
    log = OperationLog(
        operation_type=operation_type,
        batch_id=batch_id,
        record_id=record_id,
        operator=operator,
        details=details or {},
        status=status,
        error_message=error_message
    )
    db.add(log)
    db.commit()


def create_cleanup_candidate(db: Session, candidate_data: CleanupCandidateCreate, 
                             operator: str) -> CleanupCandidate:
    candidate_id = f"CLEAN-{str(uuid.uuid4())[:8].upper()}"
    
    summary = {
        "affected_batches": len(candidate_data.batch_ids),
        "affected_records": len(candidate_data.record_ids),
        "operation_type": candidate_data.operation_type
    }
    
    if candidate_data.batch_ids:
        batch_details = []
        for batch_id in candidate_data.batch_ids:
            batch = db.query(ProcessingBatch).filter(ProcessingBatch.batch_id == batch_id).first()
            if batch:
                batch_details.append({
                    "batch_id": batch_id,
                    "department": batch.department,
                    "date_range": f"{batch.start_date} to {batch.end_date}",
                    "total_records": batch.total_records
                })
        summary["batch_details"] = batch_details
    
    candidate = CleanupCandidate(
        candidate_id=candidate_id,
        operation_type=candidate_data.operation_type,
        batch_ids=candidate_data.batch_ids,
        record_ids=candidate_data.record_ids,
        reason=candidate_data.reason,
        summary=summary,
        status="pending"
    )
    db.add(candidate)
    db.commit()
    db.refresh(candidate)
    
    log_operation(db, "create_cleanup_candidate", None, None, operator, {
        "candidate_id": candidate_id,
        "operation_type": candidate_data.operation_type,
        "reason": candidate_data.reason
    })
    
    return candidate


def approve_cleanup_candidate(db: Session, candidate_id: str, approver: str) -> CleanupCandidate:
    candidate = db.query(CleanupCandidate).filter(CleanupCandidate.candidate_id == candidate_id).first()
    if not candidate:
        raise ValueError(f"Candidate {candidate_id} not found")
    
    candidate.status = "approved"
    candidate.approved_by = approver
    candidate.approved_at = datetime.utcnow()
    db.commit()
    
    log_operation(db, "approve_cleanup", None, None, approver, {
        "candidate_id": candidate_id
    })
    
    return candidate


def execute_cleanup(db: Session, candidate_id: str, executor: str) -> Dict[str, Any]:
    candidate = db.query(CleanupCandidate).filter(CleanupCandidate.candidate_id == candidate_id).first()
    if not candidate:
        raise ValueError(f"Candidate {candidate_id} not found")
    
    if candidate.status != "approved":
        raise ValueError(f"Candidate {candidate_id} is not approved")
    
    results = {
        "deleted_batches": 0,
        "deleted_records": 0,
        "deleted_logs": 0
    }
    
    if candidate.batch_ids:
        for batch_id in candidate.batch_ids:
            db.query(DutyRecord).filter(DutyRecord.batch_id == batch_id).delete()
            db.query(ProcessingBatch).filter(ProcessingBatch.batch_id == batch_id).delete()
            results["deleted_batches"] += 1
    
    if candidate.record_ids:
        results["deleted_records"] = db.query(DutyRecord).filter(
            DutyRecord.id.in_(candidate.record_ids)
        ).delete()
    
    candidate.status = "executed"
    candidate.executed_by = executor
    candidate.executed_at = datetime.utcnow()
    db.commit()
    
    log_operation(db, "execute_cleanup", None, None, executor, {
        "candidate_id": candidate_id,
        "results": results
    })
    
    return results


def get_operation_logs_by_time(db: Session, start_time: datetime, end_time: datetime, 
                               operation_type: str = None) -> List[OperationLog]:
    query = db.query(OperationLog).filter(
        OperationLog.operation_time >= start_time,
        OperationLog.operation_time <= end_time
    )
    if operation_type:
        query = query.filter(OperationLog.operation_type == operation_type)
    return query.order_by(OperationLog.operation_time.desc()).all()


def get_batch_raw_records(db: Session, batch_id: str) -> List[DutyRecord]:
    return db.query(DutyRecord).filter(DutyRecord.batch_id == batch_id).all()


def generate_sampling_export_csv(db: Session, batch_id: str, sampled_only: bool = True) -> Tuple[str, Dict[str, Any]]:
    batch = db.query(ProcessingBatch).filter(ProcessingBatch.batch_id == batch_id).first()
    if not batch:
        raise ValueError(f"Batch {batch_id} not found")
    
    if sampled_only:
        records, summary = sample_call_chain(db, batch_id)
    else:
        records = get_batch_raw_records(db, batch_id)
        summary = {
            "total_records": len(records),
            "sampled_count": len(records),
            "export_all": True
        }
    
    import csv
    from io import StringIO
    
    output = StringIO()
    writer = csv.writer(output)
    
    headers = [
        "记录ID", "批次ID", "值班日期", "部门", "值班人员", 
        "联系电话", "事件数量", "外部回执状态",
        "回执时间", "创建时间", "更新时间"
    ]
    writer.writerow(headers)
    
    for record in records:
        writer.writerow([
            record.id,
            record.batch_id,
            record.duty_date,
            record.department,
            record.duty_person,
            record.phone,
            record.incident_count,
            record.external_receipt_status,
            record.external_receipt_time.strftime("%Y-%m-%d %H:%M:%S") if record.external_receipt_time else "",
            record.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            record.updated_at.strftime("%Y-%m-%d %H:%M:%S")
        ])
    
    csv_content = output.getvalue()
    output.close()
    
    export_summary = {
        **summary,
        "export_format": "CSV",
        "export_time": datetime.utcnow().isoformat(),
        "batch_id": batch_id,
        "department": batch.department,
        "date_range": f"{batch.start_date} ~ {batch.end_date}"
    }
    
    return csv_content, export_summary


def generate_sampling_export_excel(db: Session, batch_id: str, sampled_only: bool = True) -> Tuple[bytes, Dict[str, Any]]:
    batch = db.query(ProcessingBatch).filter(ProcessingBatch.batch_id == batch_id).first()
    if not batch:
        raise ValueError(f"Batch {batch_id} not found")
    
    if sampled_only:
        records, summary = sample_call_chain(db, batch_id)
    else:
        records = get_batch_raw_records(db, batch_id)
        summary = {
            "total_records": len(records),
            "sampled_count": len(records),
            "export_all": True
        }
    
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from io import BytesIO
    
    output = BytesIO()
    wb = Workbook()
    
    ws = wb.active
    ws.title = "采样结果"
    
    headers = [
        "记录ID", "批次ID", "值班日期", "部门", "值班人员",
        "联系电话", "事件数量", "外部回执状态",
        "回执时间", "创建时间", "更新时间"
    ]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    for row, record in enumerate(records, 2):
        ws.cell(row=row, column=1, value=record.id)
        ws.cell(row=row, column=2, value=record.batch_id)
        ws.cell(row=row, column=3, value=record.duty_date)
        ws.cell(row=row, column=4, value=record.department)
        ws.cell(row=row, column=5, value=record.duty_person)
        ws.cell(row=row, column=6, value=record.phone)
        ws.cell(row=row, column=7, value=record.incident_count)
        ws.cell(row=row, column=8, value=record.external_receipt_status)
        ws.cell(row=row, column=9, value=record.external_receipt_time.strftime("%Y-%m-%d %H:%M:%S") if record.external_receipt_time else "")
        ws.cell(row=row, column=10, value=record.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row, column=11, value=record.updated_at.strftime("%Y-%m-%d %H:%M:%S"))
        
        if record.external_receipt_status == "late":
            late_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            for col in range(1, 12):
                ws.cell(row=row, column=col).fill = late_fill
    
    for col in range(1, 12):
        ws.column_dimensions[chr(64 + col)].width = 18
    
    ws_rules = wb.create_sheet(title="采样规则")
    if batch.rule_snapshot:
        rule_headers = ["规则版本", "规则名称", "采样率", "最小事件数阈值", "回执超时小时"]
        for col, header in enumerate(rule_headers, 1):
            cell = ws_rules.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        ws_rules.cell(row=2, column=1, value=batch.rule_snapshot.get("version", ""))
        ws_rules.cell(row=2, column=2, value=batch.rule_snapshot.get("name", ""))
        ws_rules.cell(row=2, column=3, value=batch.rule_snapshot.get("sampling_rate", 0))
        ws_rules.cell(row=2, column=4, value=batch.rule_snapshot.get("min_incidents", 0))
        ws_rules.cell(row=2, column=5, value=batch.rule_snapshot.get("receipt_timeout_hours", 0))
    
    ws_incidents = wb.create_sheet(title="事件详情")
    incident_headers = ["记录ID", "事件ID", "事件类型", "级别", "开始时间", "处理人", "描述"]
    for col, header in enumerate(incident_headers, 1):
        cell = ws_incidents.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    incident_row = 2
    for record in records:
        if record.incidents:
            for incident in record.incidents:
                ws_incidents.cell(row=incident_row, column=1, value=record.id)
                ws_incidents.cell(row=incident_row, column=2, value=incident.get("incident_id", ""))
                ws_incidents.cell(row=incident_row, column=3, value=incident.get("type", ""))
                ws_incidents.cell(row=incident_row, column=4, value=incident.get("level", ""))
                ws_incidents.cell(row=incident_row, column=5, value=incident.get("start_time", ""))
                ws_incidents.cell(row=incident_row, column=6, value=incident.get("handler", ""))
                ws_incidents.cell(row=incident_row, column=7, value=incident.get("description", ""))
                incident_row += 1
    
    wb.save(output)
    excel_content = output.getvalue()
    output.close()
    
    export_summary = {
        **summary,
        "export_format": "Excel",
        "export_time": datetime.utcnow().isoformat(),
        "batch_id": batch_id,
        "department": batch.department,
        "date_range": f"{batch.start_date} ~ {batch.end_date}",
        "sheets": ["采样结果", "采样规则", "事件详情"]
    }
    
    return excel_content, export_summary
