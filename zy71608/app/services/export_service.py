import os
from typing import Dict, Any, List, Optional
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import and_

from app.config import settings
from app.models import (
    ReductionApplication,
    Contract,
    AnomalyFlag,
    AnomalyStatus,
    AuditLog,
    OperationType,
)
from app.services.audit_service import AuditService


class ExportService:
    def __init__(self, db: Session):
        self.db = db
        self.audit_service = AuditService(db)

        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.anomaly_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.manual_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        self.center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def export_reduction_ledger(
        self,
        operator: str,
        status: Optional[str] = None,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        tenant_name: Optional[str] = None,
        has_anomaly: Optional[bool] = None,
    ) -> Dict[str, Any]:
        query = self.db.query(ReductionApplication).filter(ReductionApplication.is_active == True)

        if status:
            query = query.filter(ReductionApplication.status == status)
        if start_date:
            query = query.filter(ReductionApplication.closure_start_date >= start_date)
        if end_date:
            query = query.filter(ReductionApplication.closure_end_date <= end_date)
        if tenant_name:
            query = query.filter(ReductionApplication.tenant_name.like(f"%{tenant_name}%"))
        if has_anomaly is not None:
            query = query.filter(ReductionApplication.has_anomaly == has_anomaly)

        applications = query.order_by(ReductionApplication.created_at.desc()).all()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"租金减免台账_{timestamp}.xlsx"
        file_path = settings.EXPORT_DIR / file_name

        wb = Workbook()

        self._create_ledger_sheet(wb, applications)
        self._create_details_sheet(wb, applications)
        self._create_anomaly_sheet(wb, applications)
        self._create_audit_sheet(wb, applications)
        self._create_readme_sheet(wb)

        wb.save(str(file_path))

        self.audit_service.log_operation(
            operation_type=OperationType.EXPORT.value,
            operator=operator,
            table_name="reduction_applications",
            new_values={
                "record_count": len(applications),
                "file_name": file_name,
                "filters": {
                    "status": status,
                    "start_date": str(start_date) if start_date else None,
                    "end_date": str(end_date) if end_date else None,
                    "tenant_name": tenant_name,
                    "has_anomaly": has_anomaly,
                },
            },
            change_reason="导出租金减免台账",
        )

        self.db.commit()

        return {
            "file_name": file_name,
            "file_path": str(file_path),
            "record_count": len(applications),
            "total_reduction": sum(
                app.total_reduction_amount for app in applications if app.total_reduction_amount
            ),
        }

    def _create_ledger_sheet(self, wb: Workbook, applications: List[ReductionApplication]):
        ws = wb.active
        ws.title = "减免台账"

        headers = [
            "序号", "申请编号", "租户名称", "铺位编号", "合同编号", "合同版本",
            "减免原因", "闭店开始日期", "闭店结束日期", "申请天数", "核定天数",
            "减免比例", "月租金标准", "月物业费标准", "租金减免", "物业费减免",
            "减免合计", "状态", "当前环节", "是否异常", "人工修改", "创建时间",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        for row, app in enumerate(applications, 2):
            contract = (
                self.db.query(Contract)
                .filter(Contract.id == app.contract_id)
                .first()
            )

            anomalies = (
                self.db.query(AnomalyFlag)
                .filter(
                    AnomalyFlag.application_id == app.id,
                    AnomalyFlag.status == AnomalyStatus.OPEN.value,
                )
                .all()
            )

            manual_logs = (
                self.db.query(AuditLog)
                .filter(
                    AuditLog.application_id == app.id,
                    AuditLog.operation_type == OperationType.MANUAL_OVERRIDE.value,
                )
                .order_by(AuditLog.created_at.desc())
                .all()
            )

            row_data = [
                row - 1,
                app.application_no,
                app.tenant_name,
                app.store_code or "",
                contract.contract_no if contract else "",
                f"V{contract.version}" if contract else "",
                app.reduction_reason,
                str(app.closure_start_date),
                str(app.closure_end_date),
                app.applied_days,
                app.approved_days or "",
                f"{float(app.reduction_ratio * 100):.2f}%",
                float(app.monthly_rent_standard or (contract.monthly_rent if contract else 0)),
                float(app.monthly_service_fee_standard or (contract.monthly_service_fee if contract else 0)),
                float(app.calculated_rent_reduction or 0),
                float(app.calculated_service_fee_reduction or 0),
                float(app.total_reduction_amount or 0),
                app.status,
                app.current_step,
                "是" if app.has_anomaly else "否",
                "是" if app.manual_override else "否",
                app.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            ]

            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = self.thin_border

                if anomalies and col <= len(headers):
                    cell.fill = self.anomaly_fill

                if app.manual_override and not anomalies:
                    cell.fill = self.manual_fill

                if col in [1, 4, 5, 6, 8, 9, 10, 11, 12, 18, 19, 20, 21, 22]:
                    cell.alignment = self.center_align
                else:
                    cell.alignment = self.left_align

        column_widths = [6, 18, 20, 12, 18, 10, 25, 14, 14, 10, 10, 10, 14, 14, 14, 14, 14, 14, 14, 10, 10, 20]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"

    def _create_details_sheet(self, wb: Workbook, applications: List[ReductionApplication]):
        ws = wb.create_sheet("减免明细")

        headers = [
            "申请编号", "租户名称", "合同编号", "合同版本",
            "试算版本", "日租金", "日物业费", "实际减免天数",
            "租金减免", "物业费减免", "减免合计", "计算方法",
            "是否人工修改", "修改人", "修改原因", "计算时间",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        row = 2
        for app in applications:
            from app.models import ReductionCalculation

            contract = (
                self.db.query(Contract)
                .filter(Contract.id == app.contract_id)
                .first()
            )

            calculations = (
                self.db.query(ReductionCalculation)
                .filter(ReductionCalculation.application_id == app.id)
                .order_by(ReductionCalculation.calculation_version.desc())
                .all()
            )

            for calc in calculations:
                row_data = [
                    app.application_no,
                    app.tenant_name,
                    contract.contract_no if contract else "",
                    f"V{contract.version}" if contract else "",
                    f"V{calc.calculation_version}{' (当前)' if calc.is_current else ''}",
                    float(calc.daily_rent or 0),
                    float(calc.daily_service_fee or 0),
                    calc.actual_reduction_days or 0,
                    float(calc.rent_reduction or 0),
                    float(calc.service_fee_reduction or 0),
                    float(calc.total_reduction or 0),
                    calc.calculation_method or "",
                    "是" if calc.is_manual_override else "否",
                    calc.override_by or "",
                    calc.override_reason or "",
                    calc.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                ]

                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = self.thin_border

                    if calc.is_manual_override:
                        cell.fill = self.manual_fill

                    if col in [1, 2, 3, 4, 5, 8, 13, 14, 16]:
                        cell.alignment = self.center_align
                    else:
                        cell.alignment = self.left_align

                row += 1

        column_widths = [18, 20, 18, 10, 12, 12, 12, 14, 14, 14, 14, 35, 12, 12, 30, 20]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"

    def _create_anomaly_sheet(self, wb: Workbook, applications: List[ReductionApplication]):
        ws = wb.create_sheet("异常记录")

        headers = [
            "申请编号", "租户名称", "异常类型", "严重程度", "状态",
            "异常字段", "原值", "期望值", "异常描述",
            "检测时间", "处理人", "处理时间", "处理方式", "处理说明",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        row = 2
        for app in applications:
            anomalies = (
                self.db.query(AnomalyFlag)
                .filter(AnomalyFlag.application_id == app.id)
                .order_by(AnomalyFlag.severity.desc(), AnomalyFlag.created_at.desc())
                .all()
            )

            for anomaly in anomalies:
                row_data = [
                    app.application_no,
                    app.tenant_name,
                    anomaly.anomaly_type,
                    anomaly.severity,
                    anomaly.status,
                    anomaly.field_name or "",
                    anomaly.old_value or "",
                    anomaly.expected_value or "",
                    anomaly.description,
                    str(anomaly.detected_at) if anomaly.detected_at else "",
                    anomaly.resolved_by or "",
                    str(anomaly.resolved_at) if anomaly.resolved_at else "",
                    anomaly.resolution or "",
                    anomaly.resolution_notes or "",
                ]

                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = self.thin_border

                    if anomaly.status == AnomalyStatus.OPEN.value:
                        cell.fill = self.anomaly_fill

                    if col in [1, 2, 3, 4, 5, 6, 10, 11, 12]:
                        cell.alignment = self.center_align
                    else:
                        cell.alignment = self.left_align

                row += 1

        column_widths = [18, 20, 18, 10, 10, 18, 18, 18, 40, 14, 12, 14, 30, 30]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"

    def _create_audit_sheet(self, wb: Workbook, applications: List[ReductionApplication]):
        ws = wb.create_sheet("操作历史")

        headers = [
            "申请编号", "操作类型", "操作人", "操作时间",
            "表名", "记录ID", "字段名", "原值", "新值",
            "变更原因", "IP地址",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border

        row = 2
        for app in applications:
            logs = (
                self.db.query(AuditLog)
                .filter(AuditLog.application_id == app.id)
                .order_by(AuditLog.created_at.desc())
                .all()
            )

            for log in logs:
                row_data = [
                    app.application_no,
                    log.operation_type,
                    log.operator,
                    log.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                    log.table_name or "",
                    log.record_id or "",
                    log.field_name or "",
                    log.old_value or "",
                    log.new_value or "",
                    log.change_reason or "",
                    log.ip_address or "",
                ]

                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = self.thin_border

                    if log.operation_type == OperationType.MANUAL_OVERRIDE.value:
                        cell.fill = self.manual_fill

                    if col in [1, 2, 3, 4, 5, 6, 7, 11]:
                        cell.alignment = self.center_align
                    else:
                        cell.alignment = self.left_align

                row += 1

        column_widths = [18, 16, 12, 20, 20, 10, 18, 25, 25, 40, 14]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"

    def _create_readme_sheet(self, wb: Workbook):
        ws = wb.create_sheet("说明", 0)

        ws.merge_cells("A1:F1")
        title_cell = ws.cell(row=1, column=1, value="租金减免台账导出说明")
        title_cell.font = Font(bold=True, size=16)
        title_cell.alignment = self.center_align

        instructions = [
            ["", ""],
            ["导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["", ""],
            ["工作表说明", ""],
            ["  1. 减免台账", "租金减免申请的汇总信息，包含主要字段和计算结果"],
            ["  2. 减免明细", "每次试算的详细记录，包含计算过程和版本信息"],
            ["  3. 异常记录", "所有检测到的异常及其处理状态"],
            ["  4. 操作历史", "所有操作记录，特别是人工修改的追踪"],
            ["", ""],
            ["颜色标记说明", ""],
            ["  红色背景", "存在未解决的异常，需要特别关注"],
            ["  黄色背景", "存在人工修改记录，修改内容已留痕"],
            ["", ""],
            ["人工修改追溯", ""],
            ["  - 所有人工修改均记录在'操作历史'表中", ""],
            ["  - 可查看修改人、修改时间、修改前后的值", ""],
            ["  - 修改原因必须填写，便于后续审计", ""],
            ["", ""],
            ["异常处理说明", ""],
            ["  - 减免天数超限：超过最大减免天数限制（默认180天）", ""],
            ["  - 合同版本错误：使用了非最新版本或非生效状态的合同", ""],
            ["  - 重复申请：同一合同在相同时间段内有多笔申请", ""],
            ["  - 数据不一致：导入数据与合同数据存在差异", ""],
            ["  - 计算异常：计算结果与预期值存在偏差", ""],
            ["", ""],
            ["数据来源", ""],
            ["  - 租赁合同：合同基本信息和租金标准", ""],
            ["  - 租金计划：各期租金明细", ""],
            ["  - 闭店证明：实际闭店时间和核实情况", ""],
            ["  - 减免申请：租户提交的减免申请", ""],
            ["  - 补充协议：双方签署的补充协议", ""],
            ["  - 审批报告：各级审批意见", ""],
        ]

        for row, (key, value) in enumerate(instructions, 3):
            key_cell = ws.cell(row=row, column=1, value=key)
            value_cell = ws.cell(row=row, column=2, value=value)

            if key and not key.startswith(" "):
                key_cell.font = Font(bold=True)

            key_cell.alignment = self.left_align
            value_cell.alignment = self.left_align

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 60

    def export_single_application(self, application_id: int, operator: str) -> Dict[str, Any]:
        application = (
            self.db.query(ReductionApplication)
            .filter(ReductionApplication.id == application_id)
            .first()
        )

        if not application:
            raise ValueError(f"减免申请 {application_id} 不存在")

        return self.export_reduction_ledger(
            operator=operator,
            status=None,
            start_date=None,
            end_date=None,
            tenant_name=None,
            has_anomaly=None,
        )

    def get_export_files(self) -> List[Dict[str, Any]]:
        files = []
        for file_path in sorted(settings.EXPORT_DIR.glob("*.xlsx"), reverse=True):
            stat = file_path.stat()
            files.append({
                "file_name": file_path.name,
                "file_path": str(file_path),
                "size": stat.st_size,
                "created_at": datetime.fromtimestamp(stat.st_ctime),
            })
        return files
