import json
import csv
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime
from dataclasses import asdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .models import (
    CheckResult,
    PlanComparison,
    ParseError,
    RiskLevel,
    RegressionConclusion
)


class SourceTracker:
    def __init__(self):
        self.error_locations: Dict[str, List[ParseError]] = {}
        self.plan_locations: Dict[str, List[str]] = {}

    def track_errors(self, errors: List[ParseError]):
        for error in errors:
            file_path = error.source_location.file_path
            if file_path not in self.error_locations:
                self.error_locations[file_path] = []
            self.error_locations[file_path].append(error)

    def track_plan_location(self, plan_id: str, file_path: str, line_number: int):
        location_key = f"{file_path}:{line_number}"
        if plan_id not in self.plan_locations:
            self.plan_locations[plan_id] = []
        self.plan_locations[plan_id].append(location_key)

    def get_error_summary(self) -> Dict[str, Any]:
        summary = {}
        for file_path, errors in self.error_locations.items():
            error_types = {}
            for error in errors:
                error_type = error.error_type
                if error_type not in error_types:
                    error_types[error_type] = 0
                error_types[error_type] += 1

            summary[file_path] = {
                'total_errors': len(errors),
                'error_types': error_types,
                'line_numbers': sorted([e.source_location.line_number for e in errors])
            }
        return summary


class ReportGenerator:
    RISK_COLORS = {
        RiskLevel.CRITICAL: 'FF0000',
        RiskLevel.HIGH: 'FF6600',
        RiskLevel.MEDIUM: 'FFCC00',
        RiskLevel.LOW: '99CC00',
        RiskLevel.NONE: '00CC00'
    }

    CONCLUSION_COLORS = {
        RegressionConclusion.REGRESSED: 'FF0000',
        RegressionConclusion.NOT_REGRESSED: '00CC00',
        RegressionConclusion.NEED_INVESTIGATION: 'FFCC00',
        RegressionConclusion.FALSE_POSITIVE: '999999'
    }

    def __init__(self):
        self.source_tracker = SourceTracker()

    def generate_full_report(
        self,
        check_result: CheckResult,
        output_dir: str,
        report_name: str = "plan_regression_report"
    ) -> Dict[str, str]:
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        generated_files = {}

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        json_path = output_path / f"{report_name}_{timestamp}.json"
        self._generate_json_report(check_result, json_path)
        generated_files['json'] = str(json_path)

        excel_path = output_path / f"{report_name}_{timestamp}.xlsx"
        self._generate_excel_report(check_result, excel_path)
        generated_files['excel'] = str(excel_path)

        csv_path = output_path / f"{report_name}_{timestamp}.csv"
        self._generate_csv_report(check_result, csv_path)
        generated_files['csv'] = str(csv_path)

        summary_path = output_path / f"{report_name}_{timestamp}_summary.txt"
        self._generate_text_summary(check_result, summary_path)
        generated_files['summary'] = str(summary_path)

        return generated_files

    def _generate_json_report(self, check_result: CheckResult, output_path: Path):
        report_data = {
            'generated_at': check_result.generated_at.isoformat(),
            'summary': check_result.summary,
            'comparisons': [],
            'parse_errors': []
        }

        for comp in check_result.comparisons:
            comp_dict = {
                'query_template': comp.query_template,
                'parameter_set_normalized': comp.parameter_set_normalized,
                'risk_level': comp.risk_level.value,
                'confirm_status': comp.confirm_status.value,
                'conclusion': comp.conclusion.value,
                'differences': {},
                'old_plan': {
                    'scan_type': comp.old_plan.summary.scan_type,
                    'join_type': comp.old_plan.summary.join_type,
                    'estimated_rows': comp.old_plan.summary.estimated_rows,
                    'estimated_cost': comp.old_plan.summary.estimated_cost,
                    'used_indexes': comp.old_plan.summary.used_indexes,
                    'source_location': asdict(comp.old_plan.source_location) if comp.old_plan.source_location else None
                },
                'new_plan': {
                    'scan_type': comp.new_plan.summary.scan_type,
                    'join_type': comp.new_plan.summary.join_type,
                    'estimated_rows': comp.new_plan.summary.estimated_rows,
                    'estimated_cost': comp.new_plan.summary.estimated_cost,
                    'used_indexes': comp.new_plan.summary.used_indexes,
                    'source_location': asdict(comp.new_plan.source_location) if comp.new_plan.source_location else None
                },
                'notes': comp.notes,
                'review_by': comp.review_by
            }

            for field, diff in comp.differences.items():
                comp_dict['differences'][field] = {
                    'old_value': diff.old_value,
                    'new_value': diff.new_value,
                    'change_percent': diff.change_percent
                }

            report_data['comparisons'].append(comp_dict)

        for error in check_result.parse_errors:
            report_data['parse_errors'].append({
                'error_type': error.error_type,
                'error_message': error.error_message,
                'source_location': asdict(error.source_location)
            })

        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(report_data, f, ensure_ascii=False, indent=2, sort_keys=True)

    def _generate_excel_report(self, check_result: CheckResult, output_path: Path):
        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = "Summary"
        self._fill_summary_sheet(ws_summary, check_result)

        ws_comparisons = wb.create_sheet("Comparisons")
        self._fill_comparisons_sheet(ws_comparisons, check_result)

        ws_errors = wb.create_sheet("Parse Errors")
        self._fill_errors_sheet(ws_errors, check_result)

        ws_details = wb.create_sheet("Details")
        self._fill_details_sheet(ws_details, check_result)

        wb.save(output_path)

    def _fill_summary_sheet(self, ws, check_result: CheckResult):
        headers = ['Metric', 'Value']
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')

        summary = check_result.summary
        for key, value in summary.items():
            if isinstance(value, dict):
                for sub_key, sub_value in value.items():
                    ws.append([f"{key} - {sub_key}", sub_value])
            else:
                ws.append([key, value])

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _fill_comparisons_sheet(self, ws, check_result: CheckResult):
        headers = [
            'Risk Level', 'Conclusion', 'Query Template', 'Param Signature',
            'Old Scan Type', 'New Scan Type',
            'Old Cost', 'New Cost', 'Cost Change %',
            'Old Rows', 'New Rows', 'Rows Change %',
            'Old Indexes', 'New Indexes',
            'Confirm Status', 'Old Source', 'New Source'
        ]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')

        for comp in check_result.comparisons:
            old_loc = f"{comp.old_plan.source_location.file_path}:{comp.old_plan.source_location.line_number}" if comp.old_plan.source_location else ""
            new_loc = f"{comp.new_plan.source_location.file_path}:{comp.new_plan.source_location.line_number}" if comp.new_plan.source_location else ""

            cost_diff = comp.differences.get('estimated_cost')
            rows_diff = comp.differences.get('estimated_rows')

            row = [
                comp.risk_level.value,
                comp.conclusion.value,
                comp.query_template[:100],
                comp.parameter_set_normalized,
                comp.old_plan.summary.scan_type,
                comp.new_plan.summary.scan_type,
                comp.old_plan.summary.estimated_cost,
                comp.new_plan.summary.estimated_cost,
                f"{cost_diff.change_percent:.2f}%" if cost_diff else "0%",
                comp.old_plan.summary.estimated_rows,
                comp.new_plan.summary.estimated_rows,
                f"{rows_diff.change_percent:.2f}%" if rows_diff else "0%",
                ', '.join(comp.old_plan.summary.used_indexes),
                ', '.join(comp.new_plan.summary.used_indexes),
                comp.confirm_status.value,
                old_loc,
                new_loc
            ]
            ws.append(row)

            risk_color = self.RISK_COLORS.get(comp.risk_level, 'FFFFFF')
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(start_color=risk_color, end_color=risk_color, fill_type='solid')

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _fill_errors_sheet(self, ws, check_result: CheckResult):
        headers = ['File', 'Line', 'Error Type', 'Error Message', 'Raw Content']
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')

        for error in check_result.parse_errors:
            row = [
                error.source_location.file_path,
                error.source_location.line_number,
                error.error_type,
                error.error_message,
                error.source_location.raw_content[:200]
            ]
            ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _fill_details_sheet(self, ws, check_result: CheckResult):
        headers = [
            'Param Signature', 'Field', 'Old Value', 'New Value', 'Change %'
        ]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')

        for comp in check_result.comparisons:
            if comp.differences:
                for field, diff in comp.differences.items():
                    row = [
                        comp.parameter_set_normalized,
                        field,
                        str(diff.old_value),
                        str(diff.new_value),
                        f"{diff.change_percent:.2f}%" if hasattr(diff, 'change_percent') else ""
                    ]
                    ws.append(row)

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def _generate_csv_report(self, check_result: CheckResult, output_path: Path):
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                'Risk Level', 'Conclusion', 'Query Template', 'Param Signature',
                'Old Scan Type', 'New Scan Type', 'Old Cost', 'New Cost',
                'Old Rows', 'New Rows', 'Confirm Status'
            ])

            for comp in check_result.comparisons:
                writer.writerow([
                    comp.risk_level.value,
                    comp.conclusion.value,
                    comp.query_template[:100],
                    comp.parameter_set_normalized,
                    comp.old_plan.summary.scan_type,
                    comp.new_plan.summary.scan_type,
                    comp.old_plan.summary.estimated_cost,
                    comp.new_plan.summary.estimated_cost,
                    comp.old_plan.summary.estimated_rows,
                    comp.new_plan.summary.estimated_rows,
                    comp.confirm_status.value
                ])

    def _generate_text_summary(self, check_result: CheckResult, output_path: Path):
        lines = []
        lines.append("=" * 60)
        lines.append("Query Plan Regression Report")
        lines.append(f"Generated: {check_result.generated_at.strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("=" * 60)
        lines.append("")

        lines.append("SUMMARY:")
        for key, value in check_result.summary.items():
            if isinstance(value, dict):
                lines.append(f"  {key}:")
                for sub_key, sub_value in value.items():
                    lines.append(f"    {sub_key}: {sub_value}")
            else:
                lines.append(f"  {key}: {value}")
        lines.append("")

        if check_result.comparisons:
            lines.append("REGRESSION DETAILS (By Risk Level):")
            lines.append("")
            for comp in check_result.comparisons:
                if comp.risk_level != RiskLevel.NONE:
                    lines.append(f"  [{comp.risk_level.value}] {comp.query_template[:60]}...")
                    lines.append(f"    Conclusion: {comp.conclusion.value}")
                    for field, diff in comp.differences.items():
                        lines.append(f"    {field}: {diff.old_value} -> {diff.new_value}")
                    lines.append("")

        if check_result.parse_errors:
            lines.append(f"PARSE ERRORS: ({len(check_result.parse_errors)} total)")
            for error in check_result.parse_errors:
                lines.append(f"  {error.source_location.file_path}:{error.source_location.line_number}")
                lines.append(f"    [{error.error_type}] {error.error_message}")
                lines.append("")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))

    def print_console_summary(self, check_result: CheckResult):
        print("\n" + "=" * 60)
        print("Query Plan Regression Check Summary")
        print("=" * 60)

        print(f"\nTotal comparisons: {check_result.summary.get('total_comparisons', 0)}")
        print(f"Parse errors: {check_result.summary.get('total_errors', 0)}")

        risk_counts = check_result.summary.get('risk_level_counts', {})
        print("\nRisk Level Distribution:")
        for level in [RiskLevel.CRITICAL, RiskLevel.HIGH, RiskLevel.MEDIUM, RiskLevel.LOW, RiskLevel.NONE]:
            count = risk_counts.get(level.value, 0)
            if count > 0:
                print(f"  {level.value:10s}: {count}")

        conclusion_counts = check_result.summary.get('conclusion_counts', {})
        print("\nConclusion Distribution:")
        for concl in [RegressionConclusion.REGRESSED, RegressionConclusion.NEED_INVESTIGATION,
                      RegressionConclusion.NOT_REGRESSED, RegressionConclusion.FALSE_POSITIVE]:
            count = conclusion_counts.get(concl.value, 0)
            if count > 0:
                print(f"  {concl.value:20s}: {count}")

        if check_result.summary.get('regressed_count', 0) > 0:
            print(f"\n⚠️  Found {check_result.summary['regressed_count']} regressed queries!")
        else:
            print("\n✅ No regression detected.")

        print("\n" + "=" * 60 + "\n")
