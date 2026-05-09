from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc
from typing import Optional, List
from io import BytesIO
from datetime import datetime
from database import get_db
from models import (
    Permission, DetectionReport, OverlimitRecord, RectificationTask,
    ReviewReceipt, SupervisionReport, OverlimitStatus, RectificationStatus, PollutantType
)
from services import TraceService
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import tempfile
import os

router = APIRouter(prefix="/api/v1", tags=["数据追溯与导出"])


@router.get("/trace/{permit_no}", summary="按许可证号查询完整追溯链")
def get_full_trace(permit_no: str, db: Session = Depends(get_db)):
    trace_data = TraceService.get_full_trace(db, permit_no)
    
    result = {
        "permission": None,
        "detection_reports": [],
        "overlimit_records": [],
        "rectification_tasks": [],
        "review_receipts": [],
        "supervision_reports": [],
        "timeline": trace_data["timeline"]
    }
    
    perm = trace_data["permission"]
    result["permission"] = {
        "id": perm.id,
        "permit_no": perm.permit_no,
        "enterprise_name": perm.enterprise_name,
        "pollutant_type": perm.pollutant_type.value,
        "pollutant_name": perm.pollutant_name,
        "limit_value": perm.limit_value,
        "limit_unit": perm.limit_unit,
        "effective_date": perm.effective_date.isoformat() if perm.effective_date else None,
        "expiry_date": perm.expiry_date.isoformat() if perm.expiry_date else None
    }
    
    for dr in trace_data["detection_reports"]:
        result["detection_reports"].append({
            "id": dr.id,
            "report_no": dr.report_no,
            "detection_date": dr.detection_date.isoformat(),
            "detection_value": dr.detection_value,
            "detection_unit": dr.detection_unit,
            "is_overlimit": dr.is_overlimit,
            "lab_name": dr.lab_name
        })
    
    for ol in trace_data["overlimit_records"]:
        result["overlimit_records"].append({
            "id": ol.id,
            "overlimit_value": ol.overlimit_value,
            "overlimit_ratio": ol.overlimit_ratio,
            "detection_date": ol.detection_date.isoformat(),
            "identification_date": ol.identification_date.isoformat(),
            "status": ol.status.value,
            "description": ol.description
        })
    
    for rt in trace_data["rectification_tasks"]:
        result["rectification_tasks"].append({
            "id": rt.id,
            "task_no": rt.task_no,
            "deadline": rt.deadline.isoformat() if rt.deadline else None,
            "actual_completion_date": rt.actual_completion_date.isoformat() if rt.actual_completion_date else None,
            "status": rt.status.value,
            "rectification_measures": rt.rectification_measures,
            "responsible_person": rt.responsible_person
        })
    
    for rr in trace_data["review_receipts"]:
        result["review_receipts"].append({
            "id": rr.id,
            "receipt_no": rr.receipt_no,
            "review_date": rr.review_date.isoformat(),
            "reviewer": rr.reviewer,
            "review_result": rr.review_result,
            "is_qualified": rr.is_qualified,
            "redetection_value": rr.redetection_value
        })
    
    for sr in trace_data["supervision_reports"]:
        result["supervision_reports"].append({
            "id": sr.id,
            "report_no": sr.report_no,
            "report_date": sr.report_date.isoformat(),
            "reporter": sr.reporter,
            "supervision_content": sr.supervision_content,
            "supervision_result": sr.supervision_result
        })
    
    return result


@router.get("/export/overlimit-summary", summary="导出超标汇总报告")
def export_overlimit_summary(
    status: Optional[OverlimitStatus] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db)
):
    query = db.query(
        OverlimitRecord, Permission, DetectionReport
    ).join(
        Permission, OverlimitRecord.permission_id == Permission.id
    ).join(
        DetectionReport, OverlimitRecord.detection_report_id == DetectionReport.id
    )
    
    if status:
        query = query.filter(OverlimitRecord.status == status)
    
    if start_date:
        query = query.filter(OverlimitRecord.detection_date >= start_date)
    
    if end_date:
        query = query.filter(OverlimitRecord.detection_date <= end_date)
    
    records = query.order_by(desc(OverlimitRecord.detection_date)).all()
    
    wb = openpyxl.Workbook()
    
    ws = wb.active
    ws.title = "超标汇总"
    
    headers = [
        "序号", "企业名称", "许可证编号", "污染物类型", "污染物名称",
        "许可限值", "单位", "检测报告编号", "检测日期", "检测值",
        "超标值", "超标率(%)", "超标判定时间", "状态"
    ]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    for row_idx, (ol, perm, dr) in enumerate(records, 2):
        data = [
            row_idx - 1,
            perm.enterprise_name,
            perm.permit_no,
            perm.pollutant_type.value,
            perm.pollutant_name,
            perm.limit_value,
            perm.limit_unit,
            dr.report_no,
            ol.detection_date.strftime("%Y-%m-%d %H:%M"),
            dr.detection_value,
            ol.overlimit_value,
            round(ol.overlimit_ratio, 2),
            ol.identification_date.strftime("%Y-%m-%d %H:%M"),
            ol.status.value
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = center_alignment
    
    ws2 = wb.create_sheet("统计信息")
    
    pollutant_stats = {}
    status_stats = {}
    
    for ol, perm, dr in records:
        p_type = perm.pollutant_type.value
        p_name = perm.pollutant_name
        status = ol.status.value
        
        key = f"{p_type}-{p_name}"
        if key not in pollutant_stats:
            pollutant_stats[key] = {"count": 0, "total_ratio": 0}
        pollutant_stats[key]["count"] += 1
        pollutant_stats[key]["total_ratio"] += ol.overlimit_ratio
        
        if status not in status_stats:
            status_stats[status] = 0
        status_stats[status] += 1
    
    ws2.cell(row=1, column=1, value="超标统计汇总").font = Font(bold=True, size=14)
    ws2.cell(row=3, column=1, value="总超标记录数:").font = Font(bold=True)
    ws2.cell(row=3, column=2, value=len(records))
    
    ws2.cell(row=5, column=1, value="按状态统计:").font = Font(bold=True)
    row = 6
    for status, count in status_stats.items():
        ws2.cell(row=row, column=1, value=status)
        ws2.cell(row=row, column=2, value=count)
        row += 1
    
    row += 2
    ws2.cell(row=row, column=1, value="按污染物统计:").font = Font(bold=True)
    row += 1
    
    ws2.cell(row=row, column=1, value="污染物类型").font = Font(bold=True)
    ws2.cell(row=row, column=2, value="污染物名称").font = Font(bold=True)
    ws2.cell(row=row, column=3, value="超标次数").font = Font(bold=True)
    ws2.cell(row=row, column=4, value="平均超标率(%)").font = Font(bold=True)
    row += 1
    
    for key, stats in pollutant_stats.items():
        p_type, p_name = key.split("-", 1)
        ws2.cell(row=row, column=1, value=p_type)
        ws2.cell(row=row, column=2, value=p_name)
        ws2.cell(row=row, column=3, value=stats["count"])
        ws2.cell(row=row, column=4, value=round(stats["total_ratio"] / stats["count"], 2))
        row += 1
    
    for col in range(1, 15):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    temp_file.close()
    
    filename = f"超标汇总_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return FileResponse(
        path=temp_file.name,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/export/rectification-tasks", summary="导出整改任务清单")
def export_rectification_tasks(
    status: Optional[RectificationStatus] = None,
    overdue: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    from sqlalchemy import and_
    
    query = db.query(
        RectificationTask, OverlimitRecord, Permission
    ).join(
        OverlimitRecord, RectificationTask.overlimit_record_id == OverlimitRecord.id
    ).join(
        Permission, OverlimitRecord.permission_id == Permission.id
    )
    
    if status:
        query = query.filter(RectificationTask.status == status)
    
    if overdue is not None:
        now = datetime.utcnow()
        if overdue:
            query = query.filter(
                RectificationTask.deadline < now,
                RectificationTask.status.in_([RectificationStatus.PENDING, RectificationStatus.IN_PROGRESS])
            )
        else:
            query = query.filter(
                RectificationTask.deadline >= now,
                RectificationTask.status.in_([RectificationStatus.PENDING, RectificationStatus.IN_PROGRESS])
            )
    
    tasks = query.order_by(desc(RectificationTask.created_at)).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "整改任务清单"
    
    headers = [
        "序号", "企业名称", "污染物名称", "整改任务编号", "超标值",
        "超标率(%)", "截止日期", "实际完成日期", "负责人",
        "联系方式", "整改措施", "状态", "是否超期"
    ]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
    
    now = datetime.utcnow()
    
    for row_idx, (task, ol, perm) in enumerate(tasks, 2):
        is_overdue = task.deadline < now and task.status in [RectificationStatus.PENDING, RectificationStatus.IN_PROGRESS]
        
        data = [
            row_idx - 1,
            perm.enterprise_name,
            perm.pollutant_name,
            task.task_no,
            ol.overlimit_value,
            round(ol.overlimit_ratio, 2),
            task.deadline.strftime("%Y-%m-%d %H:%M") if task.deadline else "",
            task.actual_completion_date.strftime("%Y-%m-%d %H:%M") if task.actual_completion_date else "",
            task.responsible_person or "",
            task.contact_info or "",
            task.rectification_measures or "",
            task.status.value,
            "是" if is_overdue else "否"
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col == 13 and is_overdue:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    for col in range(1, 14):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    temp_file.close()
    
    filename = f"整改任务清单_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return FileResponse(
        path=temp_file.name,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/export/trace/{permit_no}", summary="导出单许可证完整追溯报告")
def export_trace_report(permit_no: str, db: Session = Depends(get_db)):
    trace_data = TraceService.get_full_trace(db, permit_no)
    
    wb = openpyxl.Workbook()
    
    perm = trace_data["permission"]
    
    ws1 = wb.active
    ws1.title = "许可证信息"
    ws1.cell(row=1, column=1, value="排污许可指标信息").font = Font(bold=True, size=16)
    
    info_rows = [
        ["企业名称", perm.enterprise_name],
        ["许可证编号", perm.permit_no],
        ["污染物类型", perm.pollutant_type.value],
        ["污染物名称", perm.pollutant_name],
        ["许可限值", f"{perm.limit_value} {perm.limit_unit}"],
        ["生效日期", perm.effective_date.isoformat()],
        ["到期日期", perm.expiry_date.isoformat()]
    ]
    
    for row_idx, (key, value) in enumerate(info_rows, 3):
        ws1.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws1.cell(row=row_idx, column=2, value=value)
    
    ws2 = wb.create_sheet("检测报告记录")
    headers2 = ["序号", "报告编号", "检测日期", "检测值", "单位", "实验室", "是否超标"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for idx, dr in enumerate(trace_data["detection_reports"], 2):
        ws2.cell(row=idx, column=1, value=idx - 1)
        ws2.cell(row=idx, column=2, value=dr.report_no)
        ws2.cell(row=idx, column=3, value=dr.detection_date.strftime("%Y-%m-%d %H:%M"))
        ws2.cell(row=idx, column=4, value=dr.detection_value)
        ws2.cell(row=idx, column=5, value=dr.detection_unit)
        ws2.cell(row=idx, column=6, value=dr.lab_name or "")
        cell = ws2.cell(row=idx, column=7, value="是" if dr.is_overlimit else "否")
        if dr.is_overlimit:
            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    if trace_data["overlimit_records"]:
        ws3 = wb.create_sheet("超标记录")
        headers3 = ["序号", "超标值", "超标率(%)", "检测日期", "判定时间", "状态", "描述"]
        for col, h in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
        
        for idx, ol in enumerate(trace_data["overlimit_records"], 2):
            ws3.cell(row=idx, column=1, value=idx - 1)
            ws3.cell(row=idx, column=2, value=ol.overlimit_value)
            ws3.cell(row=idx, column=3, value=round(ol.overlimit_ratio, 2))
            ws3.cell(row=idx, column=4, value=ol.detection_date.strftime("%Y-%m-%d %H:%M"))
            ws3.cell(row=idx, column=5, value=ol.identification_date.strftime("%Y-%m-%d %H:%M"))
            ws3.cell(row=idx, column=6, value=ol.status.value)
            ws3.cell(row=idx, column=7, value=ol.description or "")
    
    if trace_data["rectification_tasks"]:
        ws4 = wb.create_sheet("整改任务")
        headers4 = ["序号", "任务编号", "截止日期", "实际完成", "负责人", "联系方式", "状态"]
        for col, h in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        
        for idx, rt in enumerate(trace_data["rectification_tasks"], 2):
            ws4.cell(row=idx, column=1, value=idx - 1)
            ws4.cell(row=idx, column=2, value=rt.task_no)
            ws4.cell(row=idx, column=3, value=rt.deadline.strftime("%Y-%m-%d") if rt.deadline else "")
            ws4.cell(row=idx, column=4, value=rt.actual_completion_date.strftime("%Y-%m-%d") if rt.actual_completion_date else "")
            ws4.cell(row=idx, column=5, value=rt.responsible_person or "")
            ws4.cell(row=idx, column=6, value=rt.contact_info or "")
            ws4.cell(row=idx, column=7, value=rt.status.value)
    
    if trace_data["review_receipts"]:
        ws5 = wb.create_sheet("复查记录")
        headers5 = ["序号", "回执编号", "复查日期", "复查人", "复查结果", "是否合格", "复检值"]
        for col, h in enumerate(headers5, 1):
            cell = ws5.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
        
        for idx, rr in enumerate(trace_data["review_receipts"], 2):
            ws5.cell(row=idx, column=1, value=idx - 1)
            ws5.cell(row=idx, column=2, value=rr.receipt_no)
            ws5.cell(row=idx, column=3, value=rr.review_date.strftime("%Y-%m-%d %H:%M"))
            ws5.cell(row=idx, column=4, value=rr.reviewer or "")
            ws5.cell(row=idx, column=5, value="通过" if rr.review_result else "未通过")
            ws5.cell(row=idx, column=6, value="是" if rr.is_qualified else "否")
            ws5.cell(row=idx, column=7, value=rr.redetection_value if rr.redetection_value else "")
    
    if trace_data["supervision_reports"]:
        ws6 = wb.create_sheet("监管报告")
        headers6 = ["序号", "报告编号", "报告日期", "报告人", "监管内容", "监管结果", "建议"]
        for col, h in enumerate(headers6, 1):
            cell = ws6.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="A5A5A5", end_color="A5A5A5", fill_type="solid")
        
        for idx, sr in enumerate(trace_data["supervision_reports"], 2):
            ws6.cell(row=idx, column=1, value=idx - 1)
            ws6.cell(row=idx, column=2, value=sr.report_no)
            ws6.cell(row=idx, column=3, value=sr.report_date.strftime("%Y-%m-%d %H:%M"))
            ws6.cell(row=idx, column=4, value=sr.reporter or "")
            ws6.cell(row=idx, column=5, value=sr.supervision_content or "")
            ws6.cell(row=idx, column=6, value=sr.supervision_result or "")
            ws6.cell(row=idx, column=7, value=sr.suggestion or "")
    
    for ws in wb.worksheets:
        for col in range(1, 15):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    temp_file.close()
    
    filename = f"追溯报告_{permit_no}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return FileResponse(
        path=temp_file.name,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/export/detection-template", summary="下载检测报告导入模板")
def download_import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "检测报告导入"
    
    headers = [
        "report_no", "permit_no", "detection_date",
        "detection_value", "detection_unit", "detection_method",
        "lab_name", "operator", "remark"
    ]
    
    header_names = [
        "检测报告编号*", "许可证编号*", "检测时间*",
        "检测值*", "单位*", "检测方法",
        "实验室名称", "操作人员", "备注"
    ]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for col, (code, name) in enumerate(zip(headers, header_names), 1):
        cell = ws.cell(row=1, column=col, value=code)
        cell.font = Font(bold=True, color="808080")
        
        cell2 = ws.cell(row=2, column=col, value=name)
        cell2.font = header_font
        cell2.fill = header_fill
    
    examples = [
        ["DET202601001", "PERMIT-2026-001", "2026-01-15 14:30", "45.5", "mg/L", "GB 11914-89", "市环境监测站", "张三", ""],
        ["DET202601002", "PERMIT-2026-001", "2026-02-20 10:00", "52.3", "mg/L", "GB 11914-89", "市环境监测站", "李四", ""]
    ]
    
    for row_idx, row_data in enumerate(examples, 3):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    ws2 = wb.create_sheet("填写说明")
    instructions = [
        "检测报告导入模板填写说明",
        "",
        "必填字段（带*）：",
        "1. report_no - 检测报告编号，必须唯一",
        "2. permit_no - 许可证编号，必须先在系统中存在",
        "3. detection_date - 检测时间，格式：YYYY-MM-DD HH:MM",
        "4. detection_value - 检测值，数值类型",
        "5. detection_unit - 单位，必须与许可证单位一致",
        "",
        "可选字段：",
        "6. detection_method - 检测方法标准",
        "7. lab_name - 实验室名称",
        "8. operator - 操作人员",
        "9. remark - 备注信息",
        "",
        "重要提示：",
        "- 黄色背景行为示例数据，导入前请删除",
        "- 检测日期必须在许可证有效期内",
        "- 单位不一致会导致导入失败",
        "- 超过限值会自动生成超标记录"
    ]
    
    for idx, text in enumerate(instructions, 1):
        cell = ws2.cell(row=idx, column=1, value=text)
        if idx == 1:
            cell.font = Font(bold=True, size=14)
        elif "必填字段" in text or "可选字段" in text or "重要提示" in text:
            cell.font = Font(bold=True)
    
    for col in range(1, 10):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    ws2.column_dimensions['A'].width = 80
    
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    temp_file.close()
    
    return FileResponse(
        path=temp_file.name,
        filename="检测报告导入模板.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
