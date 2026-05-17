import csv
import os
from typing import List, Dict, Any, Tuple, Optional
from pathlib import Path

from models import (
    Order, Material, OutboundRecord, ReturnRecord, CompensationRecord,
    BadRecord, SourceInfo, MaterialType, DamageLevel, CompensationStatus
)

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class DataParser:
    def __init__(self):
        self.bad_records: List[BadRecord] = []

    def _create_source_info(self, file_path: str, row_number: int, raw_content: str, sheet_name: Optional[str] = None) -> SourceInfo:
        return SourceInfo(
            file_path=file_path,
            sheet_name=sheet_name,
            row_number=row_number,
            raw_content=raw_content
        )

    def _add_bad_record(self, source: SourceInfo, error_message: str, error_type: str):
        self.bad_records.append(BadRecord(
            source=source,
            error_message=error_message,
            error_type=error_type
        ))

    def parse_csv(self, file_path: str) -> Tuple[List[Dict[str, Any]], List[BadRecord]]:
        records = []
        bad_records = []
        file_path = os.path.abspath(file_path)

        try:
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f)
                for row_num, row in enumerate(reader, start=2):
                    raw_content = ','.join([f"{k}={v}" for k, v in row.items()])
                    source = self._create_source_info(file_path, row_num, raw_content)
                    records.append({
                        'data': row,
                        'source': source
                    })
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='gbk') as f:
                reader = csv.DictReader(f)
                for row_num, row in enumerate(reader, start=2):
                    raw_content = ','.join([f"{k}={v}" for k, v in row.items()])
                    source = self._create_source_info(file_path, row_num, raw_content)
                    records.append({
                        'data': row,
                        'source': source
                    })
        except Exception as e:
            source = self._create_source_info(file_path, 0, str(e))
            bad_records.append(BadRecord(source, f"文件读取失败: {str(e)}", "FileError"))

        return records, bad_records

    def parse_excel(self, file_path: str, sheet_name: Optional[str] = None) -> Tuple[List[Dict[str, Any]], List[BadRecord]]:
        if not EXCEL_AVAILABLE:
            source = self._create_source_info(file_path, 0, "")
            return [], [BadRecord(source, "openpyxl未安装，无法解析Excel文件", "DependencyError")]

        records = []
        bad_records = []
        file_path = os.path.abspath(file_path)

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            if sheet_name:
                sheets = [sheet_name] if sheet_name in wb.sheetnames else []
            else:
                sheets = wb.sheetnames

            for sheet in sheets:
                ws = wb[sheet]
                headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []

                for row_num in range(2, ws.max_row + 1):
                    row_data = {}
                    raw_cells = []
                    for col_num, header in enumerate(headers):
                        cell_value = ws.cell(row=row_num, column=col_num + 1).value
                        row_data[header] = str(cell_value) if cell_value is not None else ""
                        raw_cells.append(str(cell_value) if cell_value is not None else "")

                    raw_content = ','.join(raw_cells)
                    source = self._create_source_info(file_path, row_num, raw_content, sheet)
                    records.append({
                        'data': row_data,
                        'source': source
                    })

            wb.close()
        except Exception as e:
            source = self._create_source_info(file_path, 0, str(e))
            bad_records.append(BadRecord(source, f"Excel文件读取失败: {str(e)}", "FileError"))

        return records, bad_records

    def parse_file(self, file_path: str, sheet_name: Optional[str] = None) -> Tuple[List[Dict[str, Any]], List[BadRecord]]:
        ext = Path(file_path).suffix.lower()
        if ext in ['.xlsx', '.xls']:
            return self.parse_excel(file_path, sheet_name)
        elif ext == '.csv':
            return self.parse_csv(file_path)
        else:
            source = self._create_source_info(file_path, 0, "")
            return [], [BadRecord(source, f"不支持的文件格式: {ext}", "FormatError")]

    def _safe_get(self, data: Dict[str, Any], key: str, default: Any = "") -> str:
        for k, v in data.items():
            if k and k.strip() == key:
                return str(v).strip() if v else default
        return default

    def parse_orders(self, file_path: str) -> Tuple[List[Order], List[BadRecord]]:
        records, bad_records = self.parse_file(file_path)
        orders: List[Order] = []

        for record in records:
            data = record['data']
            source = record['source']

            try:
                order_id = self._safe_get(data, '订单ID') or self._safe_get(data, 'order_id')
                event_name = self._safe_get(data, '活动名称') or self._safe_get(data, 'event_name')
                event_date = self._safe_get(data, '活动日期') or self._safe_get(data, 'event_date')
                customer_name = self._safe_get(data, '客户姓名') or self._safe_get(data, 'customer_name')
                customer_phone = self._safe_get(data, '客户电话') or self._safe_get(data, 'customer_phone')

                if not order_id:
                    self._add_bad_record(source, "订单ID不能为空", "MissingField")
                    continue

                orders.append(Order(
                    order_id=order_id,
                    event_name=event_name,
                    event_date=event_date,
                    customer_name=customer_name,
                    customer_phone=customer_phone,
                    source=source
                ))
            except Exception as e:
                self._add_bad_record(source, f"订单解析失败: {str(e)}", "ParseError")

        return orders, bad_records + self.bad_records

    def parse_materials(self, file_path: str) -> Tuple[List[Material], List[BadRecord]]:
        records, bad_records = self.parse_file(file_path)
        materials: List[Material] = []
        self.bad_records = []

        for record in records:
            data = record['data']
            source = record['source']

            try:
                material_id = self._safe_get(data, '物料ID') or self._safe_get(data, 'material_id')
                name = self._safe_get(data, '物料名称') or self._safe_get(data, 'name')
                type_str = self._safe_get(data, '物料类型') or self._safe_get(data, 'material_type')
                unit_price_str = self._safe_get(data, '单价') or self._safe_get(data, 'unit_price') or '0'
                quantity_str = self._safe_get(data, '数量') or self._safe_get(data, 'quantity') or '0'

                if not material_id:
                    self._add_bad_record(source, "物料ID不能为空", "MissingField")
                    continue

                material_type = MaterialType.OTHER
                for mt in MaterialType:
                    if type_str and mt.value in type_str:
                        material_type = mt
                        break

                unit_price = float(unit_price_str) if unit_price_str else 0.0
                quantity = int(quantity_str) if quantity_str else 0

                materials.append(Material(
                    material_id=material_id,
                    name=name,
                    material_type=material_type,
                    unit_price=unit_price,
                    quantity=quantity,
                    source=source
                ))
            except Exception as e:
                self._add_bad_record(source, f"物料解析失败: {str(e)}", "ParseError")

        return materials, bad_records + self.bad_records

    def parse_outbound_records(self, file_path: str) -> Tuple[List[OutboundRecord], List[BadRecord]]:
        records, bad_records = self.parse_file(file_path)
        outbound_records: List[OutboundRecord] = []
        self.bad_records = []

        for record in records:
            data = record['data']
            source = record['source']

            try:
                outbound_id = self._safe_get(data, '出库ID') or self._safe_get(data, 'outbound_id')
                order_id = self._safe_get(data, '订单ID') or self._safe_get(data, 'order_id')
                material_id = self._safe_get(data, '物料ID') or self._safe_get(data, 'material_id')
                quantity_str = self._safe_get(data, '出库数量') or self._safe_get(data, 'quantity') or '0'
                outbound_date = self._safe_get(data, '出库日期') or self._safe_get(data, 'outbound_date')
                handler = self._safe_get(data, '经办人') or self._safe_get(data, 'handler')

                if not outbound_id:
                    self._add_bad_record(source, "出库ID不能为空", "MissingField")
                    continue

                quantity = int(quantity_str) if quantity_str else 0

                outbound_records.append(OutboundRecord(
                    outbound_id=outbound_id,
                    order_id=order_id,
                    material_id=material_id,
                    quantity=quantity,
                    outbound_date=outbound_date,
                    handler=handler,
                    source=source
                ))
            except Exception as e:
                self._add_bad_record(source, f"出库记录解析失败: {str(e)}", "ParseError")

        return outbound_records, bad_records + self.bad_records

    def parse_return_records(self, file_path: str) -> Tuple[List[ReturnRecord], List[BadRecord]]:
        records, bad_records = self.parse_file(file_path)
        return_records: List[ReturnRecord] = []
        self.bad_records = []

        for record in records:
            data = record['data']
            source = record['source']

            try:
                return_id = self._safe_get(data, '归还ID') or self._safe_get(data, 'return_id')
                order_id = self._safe_get(data, '订单ID') or self._safe_get(data, 'order_id')
                material_id = self._safe_get(data, '物料ID') or self._safe_get(data, 'material_id')
                quantity_str = self._safe_get(data, '归还数量') or self._safe_get(data, 'quantity') or '0'
                return_date = self._safe_get(data, '归还日期') or self._safe_get(data, 'return_date')
                checker = self._safe_get(data, '检查人') or self._safe_get(data, 'checker')
                damage_str = self._safe_get(data, '损坏程度') or self._safe_get(data, 'damage_level') or ''

                if not return_id:
                    self._add_bad_record(source, "归还ID不能为空", "MissingField")
                    continue

                quantity = int(quantity_str) if quantity_str else 0

                damage_level = DamageLevel.NONE
                for dl in DamageLevel:
                    if damage_str and dl.value in damage_str:
                        damage_level = dl
                        break

                damage_description = self._safe_get(data, '损坏描述') or self._safe_get(data, 'damage_description')

                return_records.append(ReturnRecord(
                    return_id=return_id,
                    order_id=order_id,
                    material_id=material_id,
                    quantity=quantity,
                    return_date=return_date,
                    checker=checker,
                    damage_level=damage_level,
                    damage_description=damage_description,
                    source=source
                ))
            except Exception as e:
                self._add_bad_record(source, f"归还记录解析失败: {str(e)}", "ParseError")

        return return_records, bad_records + self.bad_records

    def parse_compensation_records(self, file_path: str) -> Tuple[List[CompensationRecord], List[BadRecord]]:
        records, bad_records = self.parse_file(file_path)
        compensation_records: List[CompensationRecord] = []
        self.bad_records = []

        for record in records:
            data = record['data']
            source = record['source']

            try:
                compensation_id = self._safe_get(data, '赔付ID') or self._safe_get(data, 'compensation_id')
                order_id = self._safe_get(data, '订单ID') or self._safe_get(data, 'order_id')
                material_id = self._safe_get(data, '物料ID') or self._safe_get(data, 'material_id')
                damage_str = self._safe_get(data, '损坏程度') or self._safe_get(data, 'damage_level') or ''
                amount_str = self._safe_get(data, '赔付金额') or self._safe_get(data, 'compensation_amount') or '0'
                status_str = self._safe_get(data, '赔付状态') or self._safe_get(data, 'status') or ''
                recorded_at = self._safe_get(data, '记录时间') or self._safe_get(data, 'recorded_at')
                notes = self._safe_get(data, '备注') or self._safe_get(data, 'notes')

                if not compensation_id:
                    self._add_bad_record(source, "赔付ID不能为空", "MissingField")
                    continue

                damage_level = DamageLevel.NONE
                for dl in DamageLevel:
                    if damage_str and dl.value in damage_str:
                        damage_level = dl
                        break

                compensation_amount = float(amount_str) if amount_str else 0.0

                status = CompensationStatus.PENDING
                for cs in CompensationStatus:
                    if status_str and cs.value in status_str:
                        status = cs
                        break

                compensation_records.append(CompensationRecord(
                    compensation_id=compensation_id,
                    order_id=order_id,
                    material_id=material_id,
                    damage_level=damage_level,
                    compensation_amount=compensation_amount,
                    status=status,
                    recorded_at=recorded_at,
                    notes=notes,
                    source=source
                ))
            except Exception as e:
                self._add_bad_record(source, f"赔付记录解析失败: {str(e)}", "ParseError")

        return compensation_records, bad_records + self.bad_records
