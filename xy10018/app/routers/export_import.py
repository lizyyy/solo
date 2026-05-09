import io
import csv
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook
from ..database import get_db
from ..security import get_current_user, RoleChecker
from ..models import User, ReissueStatus, Reissue
from ..services import log_operation, create_reissue, record_failed_task
from ..models import OperationType
from ..schemas import ReissueCreate

router = APIRouter(prefix="/api", tags=["导入导出"])

manager_checker = RoleChecker(["admin", "manager"])
cs_checker = RoleChecker(["admin", "manager", "cs"])

def reissue_to_row(reissue: Reissue) -> list:
    return [
        reissue.id,
        reissue.order_no,
        reissue.customer_name,
        reissue.customer_phone,
        reissue.address,
        reissue.product_name,
        reissue.product_sku,
        reissue.quantity,
        reissue.reason,
        reissue.description,
        reissue.status.value,
        reissue.tracking_number,
        reissue.shipping_company,
        reissue.shipping_cost,
        reissue.remarks,
        reissue.version,
        reissue.created_by,
        reissue.assigned_to,
        reissue.created_at.isoformat() if reissue.created_at else "",
        reissue.updated_at.isoformat() if reissue.updated_at else ""
    ]

EXCEL_HEADERS = [
    "ID", "订单号", "客户姓名", "客户电话", "地址",
    "产品名称", "产品SKU", "数量", "原因", "描述",
    "状态", "快递单号", "快递公司", "运费", "备注",
    "版本", "创建人", "指派人", "创建时间", "更新时间"
]

@router.get("/export/excel")
async def export_excel(
    status: Optional[ReissueStatus] = None,
    assigned_to: Optional[int] = None,
    current_user: User = Depends(manager_checker),
    db: Session = Depends(get_db)
):
    query = db.query(Reissue)
    if status:
        query = query.filter(Reissue.status == status)
    if assigned_to:
        query = query.filter(Reissue.assigned_to == assigned_to)
    
    reissues = query.all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "补发单列表"
    
    ws.append(EXCEL_HEADERS)
    
    for reissue in reissues:
        ws.append(reissue_to_row(reissue))
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    log_operation(db, current_user.id, OperationType.EXPORT, None, {
        "format": "excel",
        "count": len(reissues),
        "status_filter": status.value if status else None,
        "assigned_to": assigned_to
    })
    db.commit()
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=reissues.xlsx"}
    )

@router.get("/export/csv")
async def export_csv(
    status: Optional[ReissueStatus] = None,
    assigned_to: Optional[int] = None,
    current_user: User = Depends(manager_checker),
    db: Session = Depends(get_db)
):
    query = db.query(Reissue)
    if status:
        query = query.filter(Reissue.status == status)
    if assigned_to:
        query = query.filter(Reissue.assigned_to == assigned_to)
    
    reissues = query.all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(EXCEL_HEADERS)
    
    for reissue in reissues:
        writer.writerow(reissue_to_row(reissue))
    
    output.seek(0)
    
    log_operation(db, current_user.id, OperationType.EXPORT, None, {
        "format": "csv",
        "count": len(reissues),
        "status_filter": status.value if status else None,
        "assigned_to": assigned_to
    })
    db.commit()
    
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode('utf-8')),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=reissues.csv"}
    )

@router.post("/import")
async def import_data(
    file: UploadFile = File(...),
    current_user: User = Depends(cs_checker),
    db: Session = Depends(get_db)
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="没有文件名")
    
    if not (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        raise HTTPException(status_code=400, detail="只支持Excel文件")
    
    try:
        contents = await file.read()
        wb = load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1]]
        required_fields = ["订单号", "客户姓名", "产品名称"]
        
        for field in required_fields:
            if field not in headers:
                raise HTTPException(status_code=400, detail=f"缺少必需列: {field}")
        
        success_count = 0
        failed_rows = []
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row_dict = dict(zip(headers, row))
            
            try:
                reissue_data = ReissueCreate(
                    order_no=str(row_dict["订单号"]).strip(),
                    customer_name=str(row_dict["客户姓名"]).strip(),
                    customer_phone=str(row_dict.get("客户电话", "")) if row_dict.get("客户电话") else None,
                    address=str(row_dict.get("地址", "")) if row_dict.get("地址") else None,
                    product_name=str(row_dict["产品名称"]).strip(),
                    product_sku=str(row_dict.get("产品SKU", "")) if row_dict.get("产品SKU") else None,
                    quantity=int(row_dict.get("数量", 1)) if row_dict.get("数量") else 1,
                    reason=str(row_dict.get("原因", "")) if row_dict.get("原因") else None,
                    description=str(row_dict.get("描述", "")) if row_dict.get("描述") else None,
                    assigned_to=current_user.id
                )
                
                create_reissue(db, reissue_data, created_by=current_user.id)
                success_count += 1
            except Exception as e:
                failed_rows.append({
                    "row": row_idx,
                    "order_no": row_dict.get("订单号"),
                    "error": str(e)
                })
        
        log_operation(db, current_user.id, OperationType.IMPORT, None, {
            "filename": file.filename,
            "success_count": success_count,
            "failed_count": len(failed_rows)
        })
        db.commit()
        
        return {
            "success_count": success_count,
            "failed_count": len(failed_rows),
            "failed_rows": failed_rows[:10]
        }
        
    except HTTPException:
        raise
    except Exception as e:
        record_failed_task(
            db, task_name="import_data",
            error_message=str(e),
            parameters={"filename": file.filename}
        )
        raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")
