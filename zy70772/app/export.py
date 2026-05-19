from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from typing import List
from app.models import FeatureFlag, CodeReference
from io import BytesIO


class ExcelExporter:
    def export_flags(self, flags: List[FeatureFlag]) -> bytes:
        wb = Workbook()
        
        ws1 = wb.active
        ws1.title = "Feature Flags"
        self._write_flags_sheet(ws1, flags)
        
        ws2 = wb.create_sheet("Code References")
        self._write_references_sheet(ws2, flags)
        
        ws3 = wb.create_sheet("Risk Summary")
        self._write_risk_summary(ws3, flags)
        
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
    
    def _write_flags_sheet(self, ws, flags: List[FeatureFlag]):
        headers = [
            "ID", "Name", "Description", "Default Value",
            "Status", "Risk Level", "Deletion Suggestion",
            "Experiment Status", "Owner", "Reference Count",
            "Created At", "Updated At", "Notes"
        ]
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        for row, flag in enumerate(flags, 2):
            ws.cell(row=row, column=1, value=flag.id)
            ws.cell(row=row, column=2, value=flag.name)
            ws.cell(row=row, column=3, value=flag.description or "")
            ws.cell(row=row, column=4, value="Enabled" if flag.default_value else "Disabled")
            ws.cell(row=row, column=5, value=flag.status.value)
            ws.cell(row=row, column=6, value=flag.risk_level.value if flag.risk_level else "")
            ws.cell(row=row, column=7, value=flag.deletion_suggestion.value if flag.deletion_suggestion else "")
            ws.cell(row=row, column=8, value=flag.experiment_status.value if flag.experiment_status else "")
            ws.cell(row=row, column=9, value=flag.owner or "")
            ws.cell(row=row, column=10, value=len(flag.code_references))
            ws.cell(row=row, column=11, value=flag.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws.cell(row=row, column=12, value=flag.updated_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws.cell(row=row, column=13, value=flag.notes or "")
            
            risk_cell = ws.cell(row=row, column=6)
            if flag.risk_level:
                if flag.risk_level.value in ["high", "critical"]:
                    risk_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif flag.risk_level.value == "medium":
                    risk_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                elif flag.risk_level.value in ["safe", "low"]:
                    risk_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        
        for col in range(1, 14):
            ws.column_dimensions[chr(64 + col)].width = 18
    
    def _write_references_sheet(self, ws, flags: List[FeatureFlag]):
        headers = ["Flag ID", "Flag Name", "File Path", "Line Number", "Code Snippet", "Language", "Repository", "Found At"]
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        row = 2
        for flag in flags:
            for ref in flag.code_references:
                ws.cell(row=row, column=1, value=flag.id)
                ws.cell(row=row, column=2, value=flag.name)
                ws.cell(row=row, column=3, value=ref.file_path)
                ws.cell(row=row, column=4, value=ref.line_number)
                ws.cell(row=row, column=5, value=ref.code_snippet or "")
                ws.cell(row=row, column=6, value=ref.language or "")
                ws.cell(row=row, column=7, value=ref.repository or "")
                ws.cell(row=row, column=8, value=ref.found_at.strftime("%Y-%m-%d %H:%M:%S") if ref.found_at else "")
                row += 1
        
        for col in range(1, 9):
            ws.column_dimensions[chr(64 + col)].width = 25
    
    def _write_risk_summary(self, ws, flags: List[FeatureFlag]):
        from collections import defaultdict
        
        risk_counts = defaultdict(int)
        suggestion_counts = defaultdict(int)
        status_counts = defaultdict(int)
        
        for flag in flags:
            if flag.risk_level:
                risk_counts[flag.risk_level.value] += 1
            if flag.deletion_suggestion:
                suggestion_counts[flag.deletion_suggestion.value] += 1
            status_counts[flag.status.value] += 1
        
        ws.cell(row=1, column=1, value="Risk Level Summary").font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value="Risk Level").font = Font(bold=True)
        ws.cell(row=2, column=2, value="Count").font = Font(bold=True)
        
        row = 3
        for risk, count in sorted(risk_counts.items()):
            ws.cell(row=row, column=1, value=risk)
            ws.cell(row=row, column=2, value=count)
            row += 1
        
        row += 2
        ws.cell(row=row, column=1, value="Deletion Suggestion Summary").font = Font(bold=True, size=14)
        row += 1
        ws.cell(row=row, column=1, value="Suggestion").font = Font(bold=True)
        ws.cell(row=row, column=2, value="Count").font = Font(bold=True)
        
        row += 1
        for suggestion, count in sorted(suggestion_counts.items()):
            ws.cell(row=row, column=1, value=suggestion)
            ws.cell(row=row, column=2, value=count)
            row += 1
        
        row += 2
        ws.cell(row=row, column=1, value="Status Summary").font = Font(bold=True, size=14)
        row += 1
        ws.cell(row=row, column=1, value="Status").font = Font(bold=True)
        ws.cell(row=row, column=2, value="Count").font = Font(bold=True)
        
        row += 1
        for status, count in sorted(status_counts.items()):
            ws.cell(row=row, column=1, value=status)
            ws.cell(row=row, column=2, value=count)
            row += 1
        
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15
