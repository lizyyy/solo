from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Dict, Any

from pydantic import BaseModel


class FontUsageResult(BaseModel):
    font_name: str
    file_path: str
    file_type: str
    page: Optional[str] = None
    section: Optional[str] = None
    context: Optional[str] = None
    detected_from: str = "content"
    metadata: Dict[str, Any] = {}


class BaseFileParser(ABC):
    @property
    @abstractmethod
    def supported_extensions(self) -> List[str]:
        pass
    
    @abstractmethod
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        pass


class PPTXParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".pptx"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            
            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1
                fonts_in_slide = set()
                
                for shape in slide.shapes:
                    if not shape.has_text_frame:
                        continue
                    
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font and run.font.name:
                                font_name = run.font.name.strip()
                                if font_name and font_name not in fonts_in_slide:
                                    fonts_in_slide.add(font_name)
                                    results.append(FontUsageResult(
                                        font_name=font_name,
                                        file_path=str(file_path),
                                        file_type="pptx",
                                        page=f"Slide {slide_num}",
                                        detected_from="text_run",
                                    ))
                
                try:
                    from pptx.oxml.ns import nsmap
                    slide_xml = slide._element
                    font_elements = slide_xml.xpath('.//a:rFonts', namespaces=nsmap)
                    
                    for font_elem in font_elements:
                        latin = font_elem.get('typeface')
                        east_asian = font_elem.get('ea')
                        
                        if latin and latin not in fonts_in_slide:
                            fonts_in_slide.add(latin)
                            results.append(FontUsageResult(
                                font_name=latin,
                                file_path=str(file_path),
                                file_type="pptx",
                                page=f"Slide {slide_num}",
                                detected_from="xml_rfonts",
                            ))
                        
                        if east_asian and east_asian not in fonts_in_slide:
                            fonts_in_slide.add(east_asian)
                            results.append(FontUsageResult(
                                font_name=east_asian,
                                file_path=str(file_path),
                                file_type="pptx",
                                page=f"Slide {slide_num}",
                                detected_from="xml_rfonts_eastasian",
                            ))
                except Exception:
                    pass
                    
        except ImportError:
            pass
        except Exception as e:
            pass
        
        return results


class SVGParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".svg"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            from lxml import etree
            
            tree = etree.parse(str(file_path))
            root = tree.getroot()
            
            nsmap = {'svg': 'http://www.w3.org/2000/svg'}
            
            fonts_found = set()
            
            style_elems = root.xpath('//svg:style', namespaces=nsmap)
            for style_elem in style_elems:
                style_text = style_elem.text or ""
                fonts = self._extract_font_families_from_css(style_text)
                for font in fonts:
                    if font not in fonts_found:
                        fonts_found.add(font)
                        results.append(FontUsageResult(
                            font_name=font,
                            file_path=str(file_path),
                            file_type="svg",
                            section="style",
                            detected_from="inline_css",
                        ))
            
            text_elems = root.xpath('//svg:text[@font-family]', namespaces=nsmap)
            for elem in text_elems:
                font_family = elem.get('font-family', '')
                fonts = self._parse_font_family_value(font_family)
                for font in fonts:
                    if font not in fonts_found:
                        fonts_found.add(font)
                        results.append(FontUsageResult(
                            font_name=font,
                            file_path=str(file_path),
                            file_type="svg",
                            detected_from="attribute",
                        ))
            
            style_attr_elems = root.xpath('//*[@style]', namespaces=nsmap)
            for elem in style_attr_elems:
                style_attr = elem.get('style', '')
                fonts = self._extract_font_families_from_css(style_attr)
                for font in fonts:
                    if font not in fonts_found:
                        fonts_found.add(font)
                        results.append(FontUsageResult(
                            font_name=font,
                            file_path=str(file_path),
                            file_type="svg",
                            detected_from="style_attribute",
                        ))
                        
        except ImportError:
            pass
        except Exception:
            pass
        
        return results
    
    def _extract_font_families_from_css(self, css_text: str) -> List[str]:
        import re
        fonts = []
        patterns = [
            r'font-family\s*:\s*([^;}]+)',
            r'font\s*:[^;}]*?([a-zA-Z\u4e00-\u9fa5][^;},]+)',
        ]
        
        for pattern in patterns:
            matches = re.findall(pattern, css_text, re.IGNORECASE)
            for match in matches:
                parsed = self._parse_font_family_value(match)
                fonts.extend(parsed)
        
        return list(set(fonts))
    
    def _parse_font_family_value(self, value: str) -> List[str]:
        fonts = []
        parts = value.split(',')
        for part in parts:
            font = part.strip().strip("'\"")
            if font:
                fonts.append(font)
        return fonts


class CSSParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".css"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                css_content = f.read()
            
            fonts_found = set()
            
            import re
            
            patterns = [
                r'font-family\s*:\s*([^;}]+)',
                r'@font-face\s*\{[^}]*font-family\s*:\s*([^;}]+)',
            ]
            
            for pattern in patterns:
                matches = re.findall(pattern, css_content, re.IGNORECASE)
                for match in matches:
                    parts = self._parse_font_family_list(match)
                    for font in parts:
                        if font and font not in fonts_found:
                            fonts_found.add(font)
                            results.append(FontUsageResult(
                                font_name=font,
                                file_path=str(file_path),
                                file_type="css",
                                detected_from="regex",
                            ))
                            
        except Exception:
            pass
        
        return results
    
    def _parse_font_family_list(self, value: str) -> List[str]:
        fonts = []
        import re
        
        quoted = re.findall(r"['\"]([^'\"]+)['\"]", value)
        for q in quoted:
            if q:
                fonts.append(q.strip())
        
        remaining = re.sub(r"['\"][^'\"]+['\"]", "", value)
        parts = remaining.split(',')
        for part in parts:
            font = part.strip()
            if font and font not in fonts:
                fonts.append(font)
        
        return fonts


class HTMLParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".html", ".htm"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            from lxml import etree, html
            
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            
            tree = html.fromstring(content)
            
            fonts_found = set()
            
            style_elems = tree.xpath('//style')
            css_parser = CSSParser()
            for style_elem in style_elems:
                style_text = style_elem.text_content() or ""
                import re
                matches = re.findall(r'font-family\s*:\s*([^;}]+)', style_text, re.IGNORECASE)
                for match in matches:
                    parts = match.split(',')
                    for part in parts:
                        font = part.strip().strip("'\"")
                        if font and font not in fonts_found:
                            fonts_found.add(font)
                            results.append(FontUsageResult(
                                font_name=font,
                                file_path=str(file_path),
                                file_type="html",
                                section="style",
                                detected_from="inline_css",
                            ))
            
            style_attr_elems = tree.xpath('//*[@style]')
            for elem in style_attr_elems:
                style_attr = elem.get('style', '')
                import re
                matches = re.findall(r'font-family\s*:\s*([^;}]+)', style_attr, re.IGNORECASE)
                for match in matches:
                    parts = match.split(',')
                    for part in parts:
                        font = part.strip().strip("'\"")
                        if font and font not in fonts_found:
                            fonts_found.add(font)
                            results.append(FontUsageResult(
                                font_name=font,
                                file_path=str(file_path),
                                file_type="html",
                                detected_from="style_attribute",
                            ))
            
            font_elems = tree.xpath('//font[@face]')
            for elem in font_elems:
                face = elem.get('face', '')
                parts = face.split(',')
                for part in parts:
                    font = part.strip()
                    if font and font not in fonts_found:
                        fonts_found.add(font)
                        results.append(FontUsageResult(
                            font_name=font,
                            file_path=str(file_path),
                            file_type="html",
                            detected_from="font_face",
                        ))
                        
        except ImportError:
            pass
        except Exception:
            pass
        
        return results


class PDFParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".pdf"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            import pdfplumber
            
            fonts_found = {}
            
            with pdfplumber.open(str(file_path)) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_num = page_idx + 1
                    
                    try:
                        chars = page.chars
                        if chars:
                            for char in chars:
                                font_name = char.get('fontname', '')
                                if font_name:
                                    clean_name = font_name
                                    if '+' in clean_name:
                                        clean_name = clean_name.split('+', 1)[1]
                                    if '-' in clean_name and not clean_name.startswith('-'):
                                        clean_name = clean_name.rsplit('-', 1)[0]
                                    
                                    if clean_name not in fonts_found:
                                        fonts_found[clean_name] = set()
                                    fonts_found[clean_name].add(f"Page {page_num}")
                    except Exception:
                        pass
                    
                    try:
                        page_fonts = page.fonts
                        if page_fonts:
                            for font in page_fonts:
                                font_name = font.get('name', font.get('fontname', ''))
                                if font_name:
                                    clean_name = font_name
                                    if '+' in clean_name:
                                        clean_name = clean_name.split('+', 1)[1]
                                    
                                    if clean_name not in fonts_found:
                                        fonts_found[clean_name] = set()
                                    fonts_found[clean_name].add(f"Page {page_num}")
                    except Exception:
                        pass
            
            for font_name, pages in fonts_found.items():
                pages_str = ", ".join(sorted(pages, key=lambda x: int(x.split()[-1]) if x.split()[-1].isdigit() else 0))
                results.append(FontUsageResult(
                    font_name=font_name,
                    file_path=str(file_path),
                    file_type="pdf",
                    page=pages_str if len(pages) < 10 else f"Multiple pages ({len(pages)} pages)",
                    detected_from="pdf_metadata",
                ))
                
        except ImportError:
            pass
        except Exception:
            pass
        
        return results


class ManifestParser(BaseFileParser):
    @property
    def supported_extensions(self) -> List[str]:
        return [".json", ".yml", ".yaml"]
    
    def parse(self, file_path: Path) -> List[FontUsageResult]:
        results: List[FontUsageResult] = []
        
        try:
            if file_path.suffix.lower() == ".json":
                import json
                with open(file_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
            else:
                import yaml
                with open(file_path, "r", encoding="utf-8") as f:
                    data = yaml.safe_load(f)
            
            fonts_found = set()
            
            def extract_fonts_from_value(value, section: str = ""):
                if isinstance(value, str):
                    return [value]
                elif isinstance(value, list):
                    return [str(v) for v in value if isinstance(v, str)]
                elif isinstance(value, dict):
                    fonts = []
                    for k, v in value.items():
                        extracted = extract_fonts_from_value(v, f"{section}.{k}" if section else k)
                        fonts.extend(extracted)
                    return fonts
                return []
            
            def search_for_fonts(obj, section: str = ""):
                if isinstance(obj, dict):
                    for key, value in obj.items():
                        key_lower = key.lower()
                        if 'font' in key_lower:
                            fonts = extract_fonts_from_value(value, section)
                            
                            for font in fonts:
                                if font and font not in fonts_found:
                                    fonts_found.add(font)
                                    results.append(FontUsageResult(
                                        font_name=font,
                                        file_path=str(file_path),
                                        file_type=file_path.suffix.lstrip("."),
                                        section=f"{section}.{key}" if section else key,
                                        detected_from="manifest",
                                    ))
                        search_for_fonts(value, f"{section}.{key}" if section else key)
                elif isinstance(obj, list):
                    for idx, item in enumerate(obj):
                        search_for_fonts(item, f"{section}[{idx}]" if section else f"[{idx}]")
            
            search_for_fonts(data)
            
        except ImportError:
            pass
        except Exception:
            pass
        
        return results


class ParserRegistry:
    def __init__(self):
        self._parsers: List[BaseFileParser] = []
        self._extension_map: Dict[str, BaseFileParser] = {}
    
    def register(self, parser: BaseFileParser) -> None:
        self._parsers.append(parser)
        for ext in parser.supported_extensions:
            ext_lower = ext.lower()
            if ext_lower not in self._extension_map:
                self._extension_map[ext_lower] = parser
    
    def get_parser_for_extension(self, extension: str) -> Optional[BaseFileParser]:
        return self._extension_map.get(extension.lower())
    
    def get_supported_extensions(self) -> List[str]:
        return list(self._extension_map.keys())


def create_default_registry() -> ParserRegistry:
    registry = ParserRegistry()
    registry.register(PPTXParser())
    registry.register(SVGParser())
    registry.register(CSSParser())
    registry.register(HTMLParser())
    registry.register(PDFParser())
    registry.register(ManifestParser())
    return registry


def scan_directory(
    directory: Path,
    registry: Optional[ParserRegistry] = None,
) -> List[FontUsageResult]:
    if registry is None:
        registry = create_default_registry()
    
    all_results: List[FontUsageResult] = []
    supported_extensions = registry.get_supported_extensions()
    
    for ext in supported_extensions:
        pattern = f"**/*{ext}"
        try:
            for file_path in directory.rglob(pattern):
                if file_path.is_file():
                    parser = registry.get_parser_for_extension(file_path.suffix)
                    if parser:
                        results = parser.parse(file_path)
                        all_results.extend(results)
        except Exception:
            continue
    
    return all_results
