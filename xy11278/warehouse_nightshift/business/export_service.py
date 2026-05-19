import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from warehouse_nightshift.config import EXPORT_DIR
from warehouse_nightshift.models import Task, TaskStatus, ExceptionType, FailedRecord
from warehouse_nightshift.storage import UnitOfWork


class QueryFilter:
    def __init__(self,
                 operator: Optional[str] = None,
                 start_date: Optional[date] = None,
                 end_date: Optional[date] = None,
                 status: Optional[TaskStatus] = None,
                 exception_type: Optional[ExceptionType] = None,
                 shift: Optional[str] = None):
        self.operator = operator
        self.start_date = start_date
        self.end_date = end_date
        self.status = status
        self.exception_type = exception_type
        self.shift = shift


class ExportService:
    def __init__(self, uow: UnitOfWork):
        self.uow = uow

    def query_tasks(self, query_filter: QueryFilter) -> List[Task]:
        tasks = self.uow.tasks.query(
            operator=query_filter.operator,
            start_date=query_filter.start_date,
            end_date=query_filter.end_date,
            status=query_filter.status,
            exception_type=query_filter.exception_type
        )

        if query_filter.shift:
            tasks = [t for t in tasks if t.shift == query_filter.shift]

        return tasks

    def get_summary(self, query_filter: QueryFilter) -> Dict[str, Any]:
        tasks = self.query_tasks(query_filter)

        summary = {
            'total_tasks': len(tasks),
            'by_status': {},
            'by_exception_type': {},
            'by_operator': {},
            'by_shift': {},
            'total_duration': 0,
            'avg_duration': 0
        }

        for task in tasks:
            status_key = task.status.value
            summary['by_status'][status_key] = summary['by_status'].get(status_key, 0) + 1

            if task.exception_type:
                type_key = task.exception_type.value
                summary['by_exception_type'][type_key] = summary['by_exception_type'].get(type_key, 0) + 1

            if task.assigned_operator:
                summary['by_operator'][task.assigned_operator] = \
                    summary['by_operator'].get(task.assigned_operator, 0) + 1

            if task.shift:
                summary['by_shift'][task.shift] = summary['by_shift'].get(task.shift, 0) + 1

            summary['total_duration'] += task.estimated_duration

        if summary['total_tasks'] > 0:
            summary['avg_duration'] = round(summary['total_duration'] / summary['total_tasks'], 1)

        return summary

    def export_to_csv(self, query_filter: QueryFilter, filename: str) -> Path:
        tasks = self.query_tasks(query_filter)
        filepath = EXPORT_DIR / filename

        with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                '任务ID', '标题', '描述', '优先级', '负责人', '叉车',
                '状态', '日期', '班次', '预计时长(分钟)', '开始时间', '结束时间',
                '异常类型', '异常说明'
            ])

            for task in tasks:
                writer.writerow([
                    task.id,
                    task.title,
                    task.description,
                    task.priority,
                    task.assigned_operator or '',
                    task.assigned_forklift or '',
                    task.status.value,
                    task.scheduled_date.isoformat() if task.scheduled_date else '',
                    task.shift or '',
                    task.estimated_duration,
                    task.actual_start.isoformat() if task.actual_start else '',
                    task.actual_end.isoformat() if task.actual_end else '',
                    task.exception_type.value if task.exception_type else '',
                    task.exception_note or ''
                ])

        return filepath

    def export_to_excel(self, query_filter: QueryFilter, filename: str) -> Path:
        tasks = self.query_tasks(query_filter)
        summary = self.get_summary(query_filter)
        filepath = EXPORT_DIR / filename

        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = '汇总'

        headers = ['统计项', '数值']
        ws_summary.append(headers)

        ws_summary.append(['任务总数', summary['total_tasks']])
        ws_summary.append(['总预计时长(分钟)', summary['total_duration']])
        ws_summary.append(['平均预计时长(分钟)', summary['avg_duration']])
        ws_summary.append([])
        ws_summary.append(['按状态统计'])
        for status, count in summary['by_status'].items():
            ws_summary.append([status, count])
        ws_summary.append([])
        ws_summary.append(['按异常类型统计'])
        for ex_type, count in summary['by_exception_type'].items():
            ws_summary.append([ex_type, count])
        ws_summary.append([])
        ws_summary.append(['按负责人统计'])
        for operator, count in summary['by_operator'].items():
            ws_summary.append([operator, count])
        ws_summary.append([])
        ws_summary.append(['按班次统计'])
        for shift, count in summary['by_shift'].items():
            ws_summary.append([shift, count])

        self._style_worksheet(ws_summary)

        ws_tasks = wb.create_sheet('任务明细')
        ws_tasks.append([
            '任务ID', '标题', '描述', '优先级', '负责人', '叉车',
            '状态', '日期', '班次', '预计时长(分钟)', '开始时间', '结束时间',
            '异常类型', '异常说明'
        ])

        for task in tasks:
            ws_tasks.append([
                task.id,
                task.title,
                task.description,
                task.priority,
                task.assigned_operator or '',
                task.assigned_forklift or '',
                task.status.value,
                task.scheduled_date.isoformat() if task.scheduled_date else '',
                task.shift or '',
                task.estimated_duration,
                task.actual_start.isoformat() if task.actual_start else '',
                task.actual_end.isoformat() if task.actual_end else '',
                task.exception_type.value if task.exception_type else '',
                task.exception_note or ''
            ])

        self._style_worksheet(ws_tasks)

        wb.save(filepath)
        return filepath

    def _style_worksheet(self, ws) -> None:
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def export_failed_records(self, import_id: Optional[str] = None, only_unresolved: bool = True,
                              filename: Optional[str] = None) -> Path:
        if import_id:
            records = self.uow.failed_records.get_by_import_id(import_id)
        else:
            records = self.uow.failed_records.get_unresolved() if only_unresolved else self.uow.failed_records.get_all()

        if not filename:
            filename = f'失败记录_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'

        filepath = EXPORT_DIR / filename

        with open(filepath, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                '记录ID', '导入批次ID', '类型', '原始数据', '行号',
                '错误信息', '修改建议', '是否已处理', '创建时间'
            ])

            for record in records:
                writer.writerow([
                    record.id,
                    record.import_id,
                    record.record_type,
                    record.original_data,
                    record.row_number,
                    record.error_message,
                    record.suggestion,
                    '是' if record.resolved else '否',
                    record.created_at.isoformat()
                ])

        return filepath

    def export_monthly_report(self, year: int, month: int, filename: Optional[str] = None) -> Path:
        start_date = date(year, month, 1)
        if month == 12:
            end_date = date(year + 1, 1, 1)
        else:
            end_date = date(year, month + 1, 1)

        query_filter = QueryFilter(start_date=start_date, end_date=end_date)
        if not filename:
            filename = f'月度复盘报告_{year}年{month}月.xlsx'

        return self.export_to_excel(query_filter, filename)
