import click
from tabulate import tabulate
from datetime import datetime
from typing import List, Optional
from .storage import Storage
from .models import QualityRecord, Status, AbnormalType, RecordType, BatchResult


def get_storage():
    return Storage()


def safe_str(value, default=""):
    return str(value).strip() if value is not None else default


def validate_temperature(ctx, param, value):
    if value is None:
        return value
    try:
        return float(value)
    except ValueError:
        raise click.BadParameter("温度必须是数字")


@click.group()
def cli():
    """门店品控管理 CLI 工具"""
    pass


@cli.command()
def init():
    """初始化数据存储"""
    storage = get_storage()
    click.echo(f"数据存储已初始化，位置: {storage.data_file}")
    click.echo("系统已准备就绪！")


@cli.command()
@click.argument("csv_file", type=click.Path(exists=True))
@click.option("--skip-errors", is_flag=True, help="跳过错误行继续导入")
def import_csv(csv_file, skip_errors):
    """从 CSV 文件批量导入记录"""
    import csv

    storage = get_storage()
    success_ids = []
    failed_rows = []
    errors = []

    with open(csv_file, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row_num, row in enumerate(reader, start=2):
            try:
                record = create_record_from_row(storage, row)
                storage.add_record(record)
                success_ids.append(record.id)
            except Exception as e:
                failed_rows.append(row_num)
                errors.append(f"第 {row_num} 行: {str(e)}")
                if not skip_errors:
                    break

    result = BatchResult(success_ids, [str(r) for r in failed_rows], errors)
    click.echo(result.summary())
    if result.has_errors():
        click.echo("错误详情:")
        for error in errors:
            click.echo(f"  - {error}")
        click.echo("提示: 使用 --skip-errors 可以跳过错误行继续导入")


def create_record_from_row(storage: Storage, row: dict) -> QualityRecord:
    record_type_str = safe_str(row.get("类型"))
    if record_type_str == "留样":
        record_type = RecordType.SAMPLE
    elif record_type_str == "冰箱温度":
        record_type = RecordType.TEMPERATURE
    elif record_type_str == "废弃":
        record_type = RecordType.DISCARD
    else:
        raise ValueError(f"无效的记录类型: {record_type_str}")

    abnormal_type_str = safe_str(row.get("异常类型", "无异常"))
    abnormal_type_map = {
        "无异常": AbnormalType.NONE,
        "温度过高": AbnormalType.TEMPERATURE_HIGH,
        "温度过低": AbnormalType.TEMPERATURE_LOW,
        "超时未处理": AbnormalType.EXPIRED,
        "记录缺失": AbnormalType.MISSING,
        "格式错误": AbnormalType.FORMAT_ERROR,
    }
    abnormal_type = abnormal_type_map.get(abnormal_type_str, AbnormalType.NONE)

    temp_value = safe_str(row.get("温度"))
    temperature = float(temp_value) if temp_value else None

    record = QualityRecord(
        id=storage.generate_id(),
        record_type=record_type,
        date=safe_str(row.get("日期")),
        store_name=safe_str(row.get("门店")),
        responsible=safe_str(row.get("负责人")),
        item_name=safe_str(row.get("品项")),
        temperature=temperature,
        sample_time=safe_str(row.get("留样时间")) or None,
        discard_time=safe_str(row.get("废弃时间")) or None,
        status=Status.PENDING,
        abnormal_type=abnormal_type,
        remark=safe_str(row.get("备注")),
    )

    if not record.date:
        raise ValueError("日期不能为空")
    if not record.store_name:
        raise ValueError("门店不能为空")
    if not record.responsible:
        raise ValueError("负责人不能为空")
    if not record.item_name:
        raise ValueError("品项不能为空")

    return record


@cli.command()
@click.option("--responsible", "-r", help="按负责人筛选")
@click.option("--start-date", "-s", help="开始日期 (YYYY-MM-DD)")
@click.option("--end-date", "-e", help="结束日期 (YYYY-MM-DD)")
@click.option("--status", "-t", type=click.Choice(["待复核", "已通过", "已驳回"]), help="按状态筛选")
@click.option("--abnormal", "-a", type=click.Choice(["无异常", "温度过高", "温度过低", "超时未处理", "记录缺失", "格式错误"]), help="按异常类型筛选")
@click.option("--record-type", "-y", type=click.Choice(["留样", "冰箱温度", "废弃"]), help="按记录类型筛选")
@click.option("--summary", is_flag=True, help="显示摘要统计")
def list(responsible, start_date, end_date, status, abnormal, record_type, summary):
    """查询记录列表，支持多维度筛选"""
    storage = get_storage()

    status_enum = Status(status) if status else None
    abnormal_enum = AbnormalType(abnormal) if abnormal else None
    record_type_enum = RecordType(record_type) if record_type else None

    records = storage.list_records(
        responsible=responsible,
        start_date=start_date,
        end_date=end_date,
        status=status_enum,
        abnormal_type=abnormal_enum,
        record_type=record_type_enum,
    )

    if not records:
        click.echo("没有找到符合条件的记录")
        return

    if summary:
        show_summary(records)
    else:
        show_records_table(records)


def show_records_table(records: List[QualityRecord]):
    headers = ["ID", "日期", "门店", "负责人", "类型", "品项", "状态", "异常类型"]
    rows = []
    for r in records:
        rows.append([
            r.id,
            r.date,
            r.store_name,
            r.responsible,
            r.record_type.value,
            r.item_name,
            r.status.value,
            r.abnormal_type.value,
        ])

    click.echo(tabulate(rows, headers=headers, tablefmt="simple"))
    click.echo(f"\n共 {len(records)} 条记录")


def show_summary(records: List[QualityRecord]):
    click.echo("=== 记录摘要 ===")
    click.echo(f"总记录数: {len(records)}")

    status_counts = {}
    for r in records:
        status_counts[r.status.value] = status_counts.get(r.status.value, 0) + 1
    click.echo("\n状态分布:")
    for status, count in status_counts.items():
        click.echo(f"  {status}: {count}")

    abnormal_counts = {}
    for r in records:
        if r.abnormal_type != AbnormalType.NONE:
            abnormal_counts[r.abnormal_type.value] = abnormal_counts.get(r.abnormal_type.value, 0) + 1
    click.echo("\n异常分布:")
    if abnormal_counts:
        for abnormal, count in abnormal_counts.items():
            click.echo(f"  {abnormal}: {count}")
    else:
        click.echo("  无异常记录")

    type_counts = {}
    for r in records:
        type_counts[r.record_type.value] = type_counts.get(r.record_type.value, 0) + 1
    click.echo("\n类型分布:")
    for record_type, count in type_counts.items():
        click.echo(f"  {record_type}: {count}")


@cli.command()
@click.argument("record_id")
def show(record_id):
    """查看单条记录详情"""
    storage = get_storage()
    record = storage.get_record(record_id)
    if not record:
        click.echo(f"未找到记录: {record_id}")
        return

    click.echo("=== 记录详情 ===")
    click.echo(f"ID: {record.id}")
    click.echo(f"日期: {record.date}")
    click.echo(f"门店: {record.store_name}")
    click.echo(f"负责人: {record.responsible}")
    click.echo(f"类型: {record.record_type.value}")
    click.echo(f"品项: {record.item_name}")

    if record.temperature is not None:
        click.echo(f"温度: {record.temperature}°C")
    if record.sample_time:
        click.echo(f"留样时间: {record.sample_time}")
    if record.discard_time:
        click.echo(f"废弃时间: {record.discard_time}")

    click.echo(f"状态: {record.status.value}")
    click.echo(f"异常类型: {record.abnormal_type.value}")
    if record.remark:
        click.echo(f"备注: {record.remark}")
    if record.reviewed_by:
        click.echo(f"复核人: {record.reviewed_by}")
    if record.reviewed_at:
        click.echo(f"复核时间: {record.reviewed_at}")

    if record.abnormal_type != AbnormalType.NONE:
        click.echo("\n⚠️  异常说明:")
        if record.abnormal_type == AbnormalType.TEMPERATURE_HIGH:
            click.echo(f"  该记录温度 ({record.temperature}°C) 超出安全范围，已被标记为异常")
        elif record.abnormal_type == AbnormalType.TEMPERATURE_LOW:
            click.echo(f"  该记录温度 ({record.temperature}°C) 低于安全范围，已被标记为异常")
        elif record.abnormal_type == AbnormalType.EXPIRED:
            click.echo("  该记录超时未处理，已被标记为异常")
        elif record.abnormal_type == AbnormalType.MISSING:
            click.echo("  该记录关键信息缺失，已被标记为异常")
        elif record.abnormal_type == AbnormalType.FORMAT_ERROR:
            click.echo("  该记录格式错误，已被标记为异常")


@cli.command()
@click.argument("record_ids", nargs=-1)
@click.option("--reviewer", "-v", required=True, help="复核人姓名")
@click.option("--remark", "-m", help="复核备注")
def approve(record_ids, reviewer, remark):
    """批量通过记录复核"""
    storage = get_storage()
    success_ids = []
    failed_ids = []
    errors = []

    for record_id in record_ids:
        try:
            record = storage.get_record(record_id)
            if not record:
                failed_ids.append(record_id)
                errors.append(f"{record_id}: 记录不存在")
                continue

            record.status = Status.APPROVED
            record.reviewed_by = reviewer
            record.reviewed_at = datetime.now().isoformat()
            if remark:
                record.remark = remark

            storage.update_record(record)
            success_ids.append(record_id)
        except Exception as e:
            failed_ids.append(record_id)
            errors.append(f"{record_id}: {str(e)}")

    result = BatchResult(success_ids, failed_ids, errors)
    click.echo(result.summary())
    if result.has_errors():
        click.echo("失败记录:")
        for error in errors:
            click.echo(f"  - {error}")


@cli.command()
@click.argument("record_ids", nargs=-1)
@click.option("--reviewer", "-v", required=True, help="复核人姓名")
@click.option("--remark", "-m", required=True, help="驳回原因")
def reject(record_ids, reviewer, remark):
    """批量驳回记录"""
    storage = get_storage()
    success_ids = []
    failed_ids = []
    errors = []

    for record_id in record_ids:
        try:
            record = storage.get_record(record_id)
            if not record:
                failed_ids.append(record_id)
                errors.append(f"{record_id}: 记录不存在")
                continue

            record.status = Status.REJECTED
            record.reviewed_by = reviewer
            record.reviewed_at = datetime.now().isoformat()
            record.remark = remark

            storage.update_record(record)
            success_ids.append(record_id)
        except Exception as e:
            failed_ids.append(record_id)
            errors.append(f"{record_id}: {str(e)}")

    result = BatchResult(success_ids, failed_ids, errors)
    click.echo(result.summary())
    if result.has_errors():
        click.echo("失败记录:")
        for error in errors:
            click.echo(f"  - {error}")


@cli.command()
@click.argument("output_file", type=click.Path())
@click.option("--responsible", "-r", help="按负责人筛选")
@click.option("--start-date", "-s", help="开始日期 (YYYY-MM-DD)")
@click.option("--end-date", "-e", help="结束日期 (YYYY-MM-DD)")
@click.option("--status", "-t", type=click.Choice(["待复核", "已通过", "已驳回"]), help="按状态筛选")
@click.option("--abnormal", "-a", type=click.Choice(["无异常", "温度过高", "温度过低", "超时未处理", "记录缺失", "格式错误"]), help="按异常类型筛选")
@click.option("--record-type", "-y", type=click.Choice(["留样", "冰箱温度", "废弃"]), help="按记录类型筛选")
@click.option("--format", "-f", "fmt", type=click.Choice(["xlsx", "csv"]), default="xlsx", help="导出格式")
def export(output_file, responsible, start_date, end_date, status, abnormal, record_type, fmt):
    """导出查询结果为报告"""
    storage = get_storage()

    status_enum = Status(status) if status else None
    abnormal_enum = AbnormalType(abnormal) if abnormal else None
    record_type_enum = RecordType(record_type) if record_type else None

    records = storage.list_records(
        responsible=responsible,
        start_date=start_date,
        end_date=end_date,
        status=status_enum,
        abnormal_type=abnormal_enum,
        record_type=record_type_enum,
    )

    if not records:
        click.echo("没有找到符合条件的记录，无法导出")
        return

    if fmt == "xlsx":
        export_excel(output_file, records)
    else:
        export_csv(output_file, records)

    click.echo(f"报告已导出: {output_file}")
    click.echo(f"共导出 {len(records)} 条记录")


def export_excel(output_file: str, records: List[QualityRecord]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "品控记录"

    headers = [
        "ID", "日期", "门店", "负责人", "类型", "品项",
        "温度(°C)", "留样时间", "废弃时间", "状态", "异常类型",
        "备注", "复核人", "复核时间"
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    for row, record in enumerate(records, start=2):
        data = [
            record.id,
            record.date,
            record.store_name,
            record.responsible,
            record.record_type.value,
            record.item_name,
            record.temperature,
            record.sample_time,
            record.discard_time,
            record.status.value,
            record.abnormal_type.value,
            record.remark,
            record.reviewed_by,
            record.reviewed_at,
        ]

        for col, value in enumerate(data, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            if record.status == Status.REJECTED:
                cell.fill = red_fill
            elif record.abnormal_type != AbnormalType.NONE:
                cell.fill = yellow_fill

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15

    wb.save(output_file)


def export_csv(output_file: str, records: List[QualityRecord]):
    import csv

    headers = [
        "ID", "日期", "门店", "负责人", "类型", "品项",
        "温度(°C)", "留样时间", "废弃时间", "状态", "异常类型",
        "备注", "复核人", "复核时间"
    ]

    with open(output_file, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)

        for record in records:
            writer.writerow([
                record.id,
                record.date,
                record.store_name,
                record.responsible,
                record.record_type.value,
                record.item_name,
                record.temperature,
                record.sample_time,
                record.discard_time,
                record.status.value,
                record.abnormal_type.value,
                record.remark,
                record.reviewed_by,
                record.reviewed_at,
            ])


@cli.command()
@click.option("--yes", is_flag=True, help="确认清除所有数据")
def clear(yes):
    """清除所有数据（谨慎使用）"""
    if not yes:
        click.confirm("确定要清除所有数据吗？此操作不可恢复！", abort=True)

    storage = get_storage()
    count = storage.clear_all()
    click.echo(f"已清除 {count} 条记录")


@cli.command()
def data_path():
    """显示数据文件路径"""
    storage = get_storage()
    click.echo(f"数据文件位置: {storage.data_file}")


if __name__ == "__main__":
    cli()
