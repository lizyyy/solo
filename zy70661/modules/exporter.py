import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import Dict, List, Any
from datetime import datetime
from .tracker import SourceTracker


class ReportExporter:
    def __init__(self):
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        self.error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _format_worksheet(self, ws):
        for cell in ws[1]:
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for row in ws.iter_rows():
            for cell in row:
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def export(self, output_path: str, merged_result: Dict, summary_result: Dict,
               invalid_visit: List, invalid_channel: List, tracker: SourceTracker):
        wb = Workbook()
        wb.remove(wb.active)
        
        self._create_summary_sheet(wb, summary_result)
        self._create_merge_details_sheet(wb, merged_result)
        self._create_duplicate_customers_sheet(wb, merged_result)
        self._create_consultant_summary_sheet(wb, summary_result)
        self._create_channel_summary_sheet(wb, summary_result)
        self._create_invalid_rows_sheet(wb, invalid_visit, invalid_channel)
        self._create_merge_decisions_sheet(wb, tracker)
        
        wb.save(output_path)

    def _create_summary_sheet(self, wb: Workbook, summary_result: Dict):
        ws = wb.create_sheet("总览", 0)
        
        summary_data = [
            ['统计项', '数值'],
            ['总客户数', summary_result['total_customers']],
            ['重复认领客户数', summary_result['total_duplicate']],
            ['涉及顾问数', len(summary_result['consultants'])],
            ['涉及渠道数', len(summary_result['channel_summary'])],
            ['导出时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')]
        ]
        
        for row in summary_data:
            ws.append(row)
        
        self._format_worksheet(ws)

    def _create_merge_details_sheet(self, wb: Workbook, merged_result: Dict):
        ws = wb.create_sheet("合并详情", 1)
        
        if merged_result['merge_details']:
            df = pd.DataFrame(merged_result['merge_details'])
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                is_duplicate = ws.cell(row=row_idx, column=6).value == '是'
                if is_duplicate:
                    for cell in row:
                        cell.fill = self.warning_fill
        
        self._format_worksheet(ws)

    def _create_duplicate_customers_sheet(self, wb: Workbook, merged_result: Dict):
        ws = wb.create_sheet("重复认领明细", 2)
        
        if merged_result['duplicate_customers']:
            df = pd.DataFrame(merged_result['duplicate_customers'])
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                for cell in row:
                    cell.fill = self.warning_fill
        
        self._format_worksheet(ws)

    def _create_consultant_summary_sheet(self, wb: Workbook, summary_result: Dict):
        ws = wb.create_sheet("顾问汇总", 3)
        
        if summary_result['consultants']:
            df = pd.DataFrame(summary_result['consultants'])
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
        
        self._format_worksheet(ws)

    def _create_channel_summary_sheet(self, wb: Workbook, summary_result: Dict):
        ws = wb.create_sheet("渠道汇总", 4)
        
        if summary_result['channel_summary']:
            df = pd.DataFrame(summary_result['channel_summary'])
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
        
        self._format_worksheet(ws)

    def _create_invalid_rows_sheet(self, wb: Workbook, invalid_visit: List, invalid_channel: List):
        ws = wb.create_sheet("异常记录", 5)
        
        all_invalid = []
        
        for item in invalid_visit:
            all_invalid.append({
                '来源文件': item['source_file'],
                '来源类型': '来访表',
                '来源行号': item['source_row'],
                '错误原因': ' | '.join(item['errors']),
                '原始数据': str(item['original_data'])[:500]
            })
        
        for item in invalid_channel:
            all_invalid.append({
                '来源文件': item['source_file'],
                '来源类型': '渠道表',
                '来源行号': item['source_row'],
                '错误原因': ' | '.join(item['errors']),
                '原始数据': str(item['original_data'])[:500]
            })
        
        if all_invalid:
            df = pd.DataFrame(all_invalid)
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                for cell in row:
                    cell.fill = self.error_fill
        
        self._format_worksheet(ws)

    def _create_merge_decisions_sheet(self, wb: Workbook, tracker: SourceTracker):
        ws = wb.create_sheet("决策追踪", 6)
        
        decisions = tracker.get_all_decisions()
        
        if decisions:
            decision_data = []
            for d in decisions:
                decision_data.append({
                    '归一化电话': d['normalized_phone'],
                    '决策结果': d['decision'],
                    '决策原因': d['reason'],
                    '最终渠道': d['selected_channel'],
                    '最终顾问': d['selected_consultant'],
                    '所有渠道': ' | '.join(d['all_channels']),
                    '所有顾问': ' | '.join(d['all_consultants'])
                })
            
            df = pd.DataFrame(decision_data)
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
        
        self._format_worksheet(ws)
