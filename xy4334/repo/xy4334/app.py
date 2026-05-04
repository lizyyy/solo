from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from datetime import datetime, timedelta
import sqlite3
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
CORS(app)

DATABASE = 'sample_transfer.db'
TEMP_RANGE = {'min': 2, 'max': 8}
WARNING_HOURS = 24

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transfers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_number TEXT NOT NULL UNIQUE,
            transfer_person TEXT NOT NULL,
            temperature REAL NOT NULL,
            deadline TEXT NOT NULL,
            created_at TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            risk_level TEXT NOT NULL DEFAULT 'normal',
            issues TEXT,
            notes TEXT
        )
    ''')
    conn.commit()
    conn.close()

def calculate_risk(temperature, deadline_str, created_at_str):
    issues = []
    risk_level = 'normal'
    
    if temperature < TEMP_RANGE['min'] or temperature > TEMP_RANGE['max']:
        issues.append(f"温度越界: {temperature}°C (正常范围: {TEMP_RANGE['min']}-{TEMP_RANGE['max']}°C)")
        risk_level = 'high'
    
    try:
        deadline = datetime.fromisoformat(deadline_str)
        created_at = datetime.fromisoformat(created_at_str)
        time_diff = deadline - created_at
        
        if time_diff.total_seconds() < 0:
            issues.append("截止时间已过")
            risk_level = 'high'
        elif time_diff < timedelta(hours=WARNING_HOURS):
            issues.append(f"超时风险: 距截止时间不足{WARNING_HOURS}小时")
            if risk_level != 'high':
                risk_level = 'medium'
    except Exception as e:
        issues.append(f"时间解析错误: {str(e)}")
        risk_level = 'high'
    
    return {
        'risk_level': risk_level,
        'issues': issues
    }

@app.route('/api/transfers', methods=['GET'])
def get_transfers():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM transfers ORDER BY created_at DESC')
    rows = cursor.fetchall()
    conn.close()
    
    transfers = []
    for row in rows:
        transfer = dict(row)
        transfer['issues'] = json.loads(transfer['issues']) if transfer['issues'] else []
        transfers.append(transfer)
    
    return jsonify(transfers)

@app.route('/api/transfers', methods=['POST'])
def create_transfer():
    data = request.json
    batch_number = data.get('batch_number')
    transfer_person = data.get('transfer_person')
    temperature = data.get('temperature')
    deadline = data.get('deadline')
    notes = data.get('notes', '')
    
    if not all([batch_number, transfer_person, temperature, deadline]):
        return jsonify({'error': '缺少必要字段'}), 400
    
    try:
        temperature = float(temperature)
    except ValueError:
        return jsonify({'error': '温度必须是数字'}), 400
    
    created_at = datetime.now().isoformat()
    
    conn = get_db()
    cursor = conn.cursor()
    
    try:
        risk = calculate_risk(temperature, deadline, created_at)
        status = 'pending'
        if risk['risk_level'] == 'high' and len(risk['issues']) > 0:
            status = 'review'
        
        cursor.execute('''
            INSERT INTO transfers (batch_number, transfer_person, temperature, deadline, created_at, status, risk_level, issues, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            batch_number,
            transfer_person,
            temperature,
            deadline,
            created_at,
            status,
            risk['risk_level'],
            json.dumps(risk['issues']),
            notes
        ))
        conn.commit()
        transfer_id = cursor.lastrowid
        conn.close()
        
        return jsonify({
            'id': transfer_id,
            'message': '交接记录创建成功',
            'risk_level': risk['risk_level'],
            'issues': risk['issues'],
            'status': status
        }), 201
        
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'error': f'批次编号 "{batch_number}" 已存在'}), 400

@app.route('/api/transfers/<int:transfer_id>', methods=['PUT'])
def update_transfer(transfer_id):
    data = request.json
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM transfers WHERE id = ?', (transfer_id,))
    row = cursor.fetchone()
    
    if not row:
        conn.close()
        return jsonify({'error': '交接记录不存在'}), 404
    
    current_data = dict(row)
    
    batch_number = data.get('batch_number', current_data['batch_number'])
    transfer_person = data.get('transfer_person', current_data['transfer_person'])
    temperature = data.get('temperature', current_data['temperature'])
    deadline = data.get('deadline', current_data['deadline'])
    notes = data.get('notes', current_data['notes'])
    status = data.get('status', current_data['status'])
    
    try:
        temperature = float(temperature)
    except ValueError:
        conn.close()
        return jsonify({'error': '温度必须是数字'}), 400
    
    try:
        created_at = current_data['created_at']
        risk = calculate_risk(temperature, deadline, created_at)
        
        if status == 'pending':
            if risk['risk_level'] == 'high' and len(risk['issues']) > 0:
                status = 'review'
        
        cursor.execute('''
            UPDATE transfers 
            SET batch_number = ?, transfer_person = ?, temperature = ?, deadline = ?, 
                status = ?, risk_level = ?, issues = ?, notes = ?
            WHERE id = ?
        ''', (
            batch_number,
            transfer_person,
            temperature,
            deadline,
            status,
            risk['risk_level'],
            json.dumps(risk['issues']),
            notes,
            transfer_id
        ))
        conn.commit()
        conn.close()
        
        return jsonify({
            'message': '交接记录更新成功',
            'risk_level': risk['risk_level'],
            'issues': risk['issues'],
            'status': status
        })
        
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'error': f'批次编号 "{batch_number}" 已存在'}), 400

@app.route('/api/transfers/<int:transfer_id>/release', methods=['POST'])
def release_transfer(transfer_id):
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM transfers WHERE id = ?', (transfer_id,))
    row = cursor.fetchone()
    
    if not row:
        conn.close()
        return jsonify({'error': '交接记录不存在'}), 404
    
    cursor.execute('''
        UPDATE transfers SET status = 'released' WHERE id = ?
    ''', (transfer_id,))
    conn.commit()
    conn.close()
    
    return jsonify({'message': '样本已放行'})

@app.route('/api/transfers/<int:transfer_id>', methods=['DELETE'])
def delete_transfer(transfer_id):
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('DELETE FROM transfers WHERE id = ?', (transfer_id,))
    
    if cursor.rowcount == 0:
        conn.close()
        return jsonify({'error': '交接记录不存在'}), 404
    
    conn.commit()
    conn.close()
    
    return jsonify({'message': '交接记录已删除'})

@app.route('/api/import', methods=['POST'])
def import_transfers():
    if 'file' not in request.files:
        return jsonify({'error': '没有上传文件'}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400
    
    if not file.filename.endswith('.json'):
        return jsonify({'error': '只支持JSON格式文件'}), 400
    
    try:
        data = json.load(file)
    except Exception as e:
        return jsonify({'error': f'文件解析错误: {str(e)}'}), 400
    
    if not isinstance(data, list):
        return jsonify({'error': 'JSON格式错误，需要是数组格式'}), 400
    
    conn = get_db()
    cursor = conn.cursor()
    results = []
    
    for item in data:
        batch_number = item.get('batch_number')
        transfer_person = item.get('transfer_person')
        temperature = item.get('temperature')
        deadline = item.get('deadline')
        notes = item.get('notes', '')
        
        if not all([batch_number, transfer_person, temperature, deadline]):
            results.append({'batch_number': batch_number, 'error': '缺少必要字段'})
            continue
        
        try:
            temperature = float(temperature)
        except ValueError:
            results.append({'batch_number': batch_number, 'error': '温度必须是数字'})
            continue
        
        created_at = datetime.now().isoformat()
        
        try:
            risk = calculate_risk(temperature, deadline, created_at)
            status = 'pending'
            if risk['risk_level'] == 'high' and len(risk['issues']) > 0:
                status = 'review'
            
            cursor.execute('''
                INSERT INTO transfers (batch_number, transfer_person, temperature, deadline, created_at, status, risk_level, issues, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                batch_number,
                transfer_person,
                temperature,
                deadline,
                created_at,
                status,
                risk['risk_level'],
                json.dumps(risk['issues']),
                notes
            ))
            results.append({
                'batch_number': batch_number,
                'success': True,
                'risk_level': risk['risk_level'],
                'status': status
            })
        except sqlite3.IntegrityError:
            results.append({'batch_number': batch_number, 'error': '批次编号已存在'})
    
    conn.commit()
    conn.close()
    
    success_count = sum(1 for r in results if r.get('success'))
    fail_count = len(results) - success_count
    
    return jsonify({
        'message': f'导入完成: 成功{success_count}条, 失败{fail_count}条',
        'details': results
    })

@app.route('/api/export/today', methods=['GET'])
def export_today():
    today = datetime.now().date().isoformat()
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT * FROM transfers 
        WHERE date(created_at) = ?
        ORDER BY created_at DESC
    ''', (today,))
    
    rows = cursor.fetchall()
    conn.close()
    
    if not rows:
        return jsonify({'error': '今天没有交接记录'}), 404
    
    wb = Workbook()
    ws = wb.active
    ws.title = '样本交接报告'
    
    header_font = Font(bold=True, size=12)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    ws['A1'] = f'样本交接报告 - {today}'
    ws.merge_cells('A1:I1')
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    headers = ['批次编号', '交接人', '温度(°C)', '截止时间', '创建时间', '状态', '风险等级', '问题', '备注']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    status_map = {'pending': '待处理', 'released': '已放行', 'review': '需复核'}
    risk_map = {'normal': '正常', 'medium': '中等', 'high': '高'}
    
    for row_num, row in enumerate(rows, 4):
        transfer = dict(row)
        issues = json.loads(transfer['issues']) if transfer['issues'] else []
        
        ws.cell(row=row_num, column=1, value=transfer['batch_number']).border = thin_border
        ws.cell(row=row_num, column=2, value=transfer['transfer_person']).border = thin_border
        ws.cell(row=row_num, column=3, value=transfer['temperature']).border = thin_border
        ws.cell(row=row_num, column=4, value=transfer['deadline']).border = thin_border
        ws.cell(row=row_num, column=5, value=transfer['created_at']).border = thin_border
        ws.cell(row=row_num, column=6, value=status_map.get(transfer['status'], transfer['status'])).border = thin_border
        ws.cell(row=row_num, column=7, value=risk_map.get(transfer['risk_level'], transfer['risk_level'])).border = thin_border
        ws.cell(row=row_num, column=8, value='; '.join(issues) if issues else '无').border = thin_border
        ws.cell(row=row_num, column=9, value=transfer['notes'] or '').border = thin_border
    
    for col_num in range(1, 10):
        ws.column_dimensions[get_column_letter(col_num)].width = 20
    
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[3].height = 25
    
    filename = f'sample_transfer_report_{today}.xlsx'
    wb.save(filename)
    
    return send_file(
        filename,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/')
def index():
    return app.send_static_file('index.html')

if __name__ == '__main__':
    init_db()
    print(f"数据库已初始化: {os.path.abspath(DATABASE)}")
    app.run(debug=True, host='0.0.0.0', port=5001)
