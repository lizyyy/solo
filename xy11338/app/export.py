import io
import csv
from datetime import datetime
from typing import List, Dict, Any
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from fastapi.responses import StreamingResponse
from app.models import Issue, IssueItem, Claim, ClaimItem, PartReturn, Part
from app.utils import mask_sensitive_data


def create_excel_response(data: List[Dict[str, Any]], filename: str, sheet_name: str = "Sheet1") -> StreamingResponse:
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    if data:
        headers = list(data[0].keys())
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))
                cell.alignment = Alignment(horizontal='left')
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def create_csv_response(data: List[Dict[str, Any]], filename: str) -> StreamingResponse:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=data[0].keys() if data else [])
    writer.writeheader()
    writer.writerows(data)
    
    output_bytes = io.BytesIO(output.getvalue().encode('utf-8-sig'))
    
    return StreamingResponse(
        output_bytes,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def export_issues_report(db: Session, format: str = "xlsx"):
    issues = db.query(Issue).all()
    
    data = []
    for issue in issues:
        for item in issue.items:
            data.append({
                "领用单号": issue.issue_no,
                "工单编号": issue.work_order_no,
                "客户姓名": mask_sensitive_data(issue.customer_name),
                "客户电话": mask_sensitive_data(issue.customer_phone),
                "家电类型": issue.appliance_type,
                "故障描述": issue.fault_description,
                "零件编码": item.part.part_code if item.part else "",
                "零件名称": item.part.part_name if item.part else "",
                "领用数量": item.quantity,
                "状态": item.status.value if hasattr(item.status, 'value') else item.status,
                "需要返还旧件": "是" if item.old_part_expected else "否",
                "旧件已返还": "是" if item.old_part_returned else "否",
                "批次号": issue.batch_no,
                "创建时间": issue.created_at.strftime("%Y-%m-%d %H:%M:%S") if issue.created_at else ""
            })
    
    filename = f"领用报表_{datetime.now().strftime('%Y%m%d%H%M%S')}.{format}"
    if format == "csv":
        return create_csv_response(data, filename)
    return create_excel_response(data, filename, "领用明细")


def export_returns_report(db: Session, format: str = "xlsx"):
    returns = db.query(PartReturn).all()
    
    data = []
    for ret in returns:
        for item in ret.items:
            data.append({
                "返还单号": ret.return_no,
                "领用单号": ret.issue.issue_no if ret.issue else "",
                "工单编号": ret.issue.work_order_no if ret.issue else "",
                "零件编码": item.part.part_code if item.part else "",
                "零件名称": item.part.part_name if item.part else "",
                "返还数量": item.quantity,
                "是否缺陷": "是" if item.is_defective else "否",
                "缺陷描述": item.defect_description,
                "是否已验证": "是" if item.condition_verified else "否",
                "批次号": ret.batch_no,
                "返还日期": ret.return_date.strftime("%Y-%m-%d %H:%M:%S") if ret.return_date else ""
            })
    
    filename = f"返还报表_{datetime.now().strftime('%Y%m%d%H%M%S')}.{format}"
    if format == "csv":
        return create_csv_response(data, filename)
    return create_excel_response(data, filename, "返还明细")


def export_claims_report(db: Session, format: str = "xlsx"):
    claims = db.query(Claim).all()
    
    data = []
    for claim in claims:
        for item in claim.items:
            data.append({
                "索赔单号": claim.claim_no,
                "供应商": claim.vendor,
                "零件编码": item.part_code,
                "零件名称": item.part_name,
                "数量": item.quantity,
                "单价": item.unit_price,
                "金额": item.amount,
                "工单编号": item.work_order_no,
                "缺陷代码": item.defect_code,
                "缺陷描述": item.defect_description,
                "是否重复": "是" if item.is_duplicate else "否",
                "状态": claim.status.value if hasattr(claim.status, 'value') else claim.status,
                "批次号": claim.batch_no,
                "提交时间": claim.submitted_at.strftime("%Y-%m-%d %H:%M:%S") if claim.submitted_at else "",
                "批准时间": claim.approved_at.strftime("%Y-%m-%d %H:%M:%S") if claim.approved_at else ""
            })
    
    filename = f"索赔报表_{datetime.now().strftime('%Y%m%d%H%M%S')}.{format}"
    if format == "csv":
        return create_csv_response(data, filename)
    return create_excel_response(data, filename, "索赔明细")


def export_monthly_reconciliation(db: Session, year: int, month: int, format: str = "xlsx"):
    from datetime import datetime as dt
    
    start_date = dt(year, month, 1)
    if month == 12:
        end_date = dt(year + 1, 1, 1)
    else:
        end_date = dt(year, month + 1, 1)
    
    issues = db.query(Issue).filter(
        Issue.created_at >= start_date,
        Issue.created_at < end_date
    ).all()
    
    data = []
    for issue in issues:
        for item in issue.items:
            is_returned = item.old_part_returned
            
            return_qty = 0
            if issue.returns:
                for ret in issue.returns:
                    for ret_item in ret.items:
                        if ret_item.part_id == item.part_id:
                            return_qty += ret_item.quantity
            
            claim_qty = 0
            claim_amount = 0
            claim_items = db.query(ClaimItem).filter(
                ClaimItem.work_order_no == issue.work_order_no,
                ClaimItem.part_code == item.part.part_code if item.part else ""
            ).all()
            for ci in claim_items:
                if not ci.is_duplicate:
                    claim_qty += ci.quantity
                    claim_amount += ci.amount
            
            data.append({
                "月份": f"{year}年{month}月",
                "领用单号": issue.issue_no,
                "工单编号": issue.work_order_no,
                "零件编码": item.part.part_code if item.part else "",
                "零件名称": item.part.part_name if item.part else "",
                "领用数量": item.quantity,
                "需要返还": "是" if item.old_part_expected else "否",
                "已返还数量": return_qty,
                "未返还数量": item.quantity - return_qty if item.old_part_expected else 0,
                "已索赔数量": claim_qty,
                "索赔金额": claim_amount,
                "状态": "正常" if (not item.old_part_expected or return_qty >= item.quantity) and claim_qty == 0 else ("旧件缺失" if return_qty < item.quantity and item.old_part_expected else "已索赔"),
                "批次号": issue.batch_no,
                "领用日期": issue.created_at.strftime("%Y-%m-%d") if issue.created_at else ""
            })
    
    filename = f"月度核对报表_{year}年{month}月_{datetime.now().strftime('%Y%m%d%H%M%S')}.{format}"
    if format == "csv":
        return create_csv_response(data, filename)
    return create_excel_response(data, filename, "月度核对")


def get_claim_verification_details(db: Session, claim_id: int) -> List[Dict[str, Any]]:
    claim = db.query(Claim).filter(Claim.id == claim_id).first()
    if not claim:
        return []
    
    result = []
    for verification in claim.verifications:
        result.append({
            "校验规则": verification.rule_name,
            "校验结果": verification.result.value if hasattr(verification.result, 'value') else verification.result,
            "原因": verification.reason,
            "校验时间": verification.verified_at.strftime("%Y-%m-%d %H:%M:%S") if verification.verified_at else ""
        })
    
    return result
