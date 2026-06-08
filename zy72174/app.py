import os
import json
import csv
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, flash
from werkzeug.utils import secure_filename
import io

from database import get_db, init_db

app = Flask(__name__)
app.config['SECRET_KEY'] = 'demolition-resettlement-key'
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['DATABASE'] = 'resettlement.db'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs('templates', exist_ok=True)

from core_logic import (
    get_point_list, get_point_detail, update_point_field, add_review,
    get_statistics, export_to_excel, merge_and_match_points,
    STATUS_MAP, STATUS_COLORS
)

@app.route('/')
def index():
    stats = get_statistics()
    status_filter = request.args.get('status', 'all')
    keyword = request.args.get('keyword', '')
    points = get_point_list(status_filter=status_filter, keyword=keyword)
    
    return render_template('index.html',
                         stats=stats,
                         points=points,
                         status_filter=status_filter,
                         keyword=keyword,
                         STATUS_MAP=STATUS_MAP)

@app.route('/point/<int:point_id>')
def point_detail(point_id):
    data = get_point_detail(point_id)
    return render_template('detail.html',
                         point=data['point'],
                         versions=data['versions'],
                         feedback=data['feedback'],
                         photos=data['photos'],
                         notes=data['notes'],
                         reviews=data['reviews'],
                         STATUS_MAP=STATUS_MAP)

@app.route('/point/<int:point_id>/review', methods=['POST'])
def submit_review(point_id):
    reviewer = request.form.get('reviewer', '何工')
    review_result = request.form.get('review_result')
    review_note = request.form.get('review_note', '')
    
    add_review(point_id, reviewer, review_result, review_note)
    flash('复核意见已提交', 'success')
    return redirect(url_for('point_detail', point_id=point_id))

@app.route('/point/<int:point_id>/update', methods=['POST'])
def update_point(point_id):
    field_name = request.form.get('field_name')
    new_value = request.form.get('new_value')
    operator = request.form.get('operator', '何工')
    operation_note = request.form.get('operation_note', '')
    
    success, msg = update_point_field(point_id, field_name, new_value, operator, operation_note)
    flash(msg, 'success' if success else 'warning')
    return redirect(url_for('point_detail', point_id=point_id))

@app.route('/export')
def export_data():
    status_filter = request.args.get('status', 'all')
    wb = export_to_excel(status_filter=status_filter)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"城市更新拆迁安置清单_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output,
                     as_attachment=True,
                     download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/import', methods=['GET', 'POST'])
def import_data():
    if request.method == 'POST':
        source_type = request.form.get('source_type')
        source_name = request.form.get('source_name')
        operator = request.form.get('operator', '何工')
        
        file = request.files.get('file')
        if not file or not file.filename:
            flash('请选择要上传的文件', 'warning')
            return redirect(url_for('import_data'))
        
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        try:
            gis_data, feedback_data, photos_data, notes_data = _parse_uploaded_file(filepath, source_type)
        except Exception as e:
            flash(f'文件解析失败: {e}', 'danger')
            return redirect(url_for('import_data'))
        
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO data_sources (source_type, source_name, source_file, imported_by) VALUES (?, ?, ?, ?)",
            (source_type, source_name, filename, operator)
        )
        source_id = cursor.lastrowid
        conn.commit()
        conn.close()
        
        report = merge_and_match_points(
            gis_data, feedback_data, photos_data, notes_data,
            operator=operator, source_id=source_id
        )
        
        summary_parts = []
        if report['created'] > 0:
            summary_parts.append(f"新建{report['created']}个点位")
        if report['updated'] > 0:
            summary_parts.append(f"更新{report['updated']}个点位(面积差异)")
        if report['feedback_added'] > 0:
            summary_parts.append(f"导入{report['feedback_added']}条反馈")
        if report['photos_added'] > 0:
            summary_parts.append(f"导入{report['photos_added']}张照片")
        if report['notes_added'] > 0:
            summary_parts.append(f"导入{report['notes_added']}条备注")
        if report['conflicts']:
            conflict_nos = ', '.join(c['point_no'] for c in report['conflicts'])
            summary_parts.append(f"冲突待核实: {conflict_nos}")
        
        if summary_parts:
            flash(f'导入完成: {"; ".join(summary_parts)}', 'success')
        else:
            flash(f'文件 {filename} 已导入，但未发现可匹配的数据', 'info')
        
        return redirect(url_for('import_data'))
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM data_sources ORDER BY imported_at DESC')
    sources = [dict(row) for row in cursor.fetchall()]
    conn.close()
    
    return render_template('import.html', sources=sources)


def _parse_uploaded_file(filepath, source_type):
    gis_data = []
    feedback_data = []
    photos_data = []
    notes_data = []
    
    ext = os.path.splitext(filepath)[1].lower()
    raw_rows = []
    
    if ext == '.json':
        with open(filepath, 'r', encoding='utf-8') as f:
            content = json.load(f)
        if isinstance(content, list):
            raw_rows = content
        elif isinstance(content, dict):
            if 'points' in content:
                raw_rows = content['points']
            elif 'data' in content:
                raw_rows = content['data']
            else:
                raw_rows = [content]
    elif ext in ('.xlsx', '.xls'):
        from openpyxl import load_workbook
        wb = load_workbook(filepath, read_only=True)
        ws = wb.active
        headers = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h).strip() if h else '' for h in row]
                continue
            if headers:
                row_dict = {}
                for j, val in enumerate(row):
                    if j < len(headers) and headers[j]:
                        row_dict[headers[j]] = val
                if row_dict:
                    raw_rows.append(row_dict)
        wb.close()
    elif ext == '.csv':
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                raw_rows.append(dict(row))
    else:
        raise ValueError(f'不支持的文件格式: {ext}')
    
    if source_type in ('gis_system', 'gis_old'):
        for row in raw_rows:
            point = {
                'point_no': str(row.get('point_no', row.get('点位编号', ''))),
                'address': str(row.get('address', row.get('地址', ''))),
                'district': str(row.get('district', row.get('行政区', ''))) if row.get('district', row.get('行政区')) else None,
                'gis_lng': float(row.get('gis_lng', row.get('经度', 0)) or 0),
                'gis_lat': float(row.get('gis_lat', row.get('纬度', 0)) or 0),
                'property_type': str(row.get('property_type', row.get('性质', ''))),
                'area': float(row.get('area', row.get('面积', 0)) or 0),
                'households': int(float(row.get('households', row.get('户数', 0)) or 0)),
            }
            if point['point_no'] and point['address']:
                gis_data.append(point)
    
    elif source_type == 'feedback':
        for row in raw_rows:
            fb = {
                'point_no': str(row.get('point_no', row.get('点位编号', ''))),
                'feedback_type': str(row.get('feedback_type', row.get('反馈类型', ''))),
                'feedback_content': str(row.get('feedback_content', row.get('反馈内容', ''))),
                'feedback_source': str(row.get('feedback_source', row.get('反馈来源', ''))),
                'feedback_time': str(row.get('feedback_time', row.get('反馈时间', ''))) if row.get('feedback_time', row.get('反馈时间')) else None,
                'handler': str(row.get('handler', row.get('处理人', ''))) if row.get('handler', row.get('处理人')) else None,
                'handle_note': str(row.get('handle_note', row.get('处理意见', ''))) if row.get('handle_note', row.get('处理意见')) else None,
                'handled_at': str(row.get('handled_at', row.get('处理时间', ''))) if row.get('handled_at', row.get('处理时间')) else None,
                'is_resolved': int(float(row.get('is_resolved', row.get('是否已处理', 0)) or 0)),
            }
            if fb['point_no'] and fb['feedback_content']:
                feedback_data.append(fb)
    
    elif source_type == 'inspection':
        for row in raw_rows:
            photo = {
                'point_no': str(row.get('point_no', row.get('点位编号', ''))),
                'photo_path': str(row.get('photo_path', row.get('照片路径', ''))),
                'photo_desc': str(row.get('photo_desc', row.get('照片描述', ''))),
                'taken_at': str(row.get('taken_at', row.get('拍摄时间', ''))) if row.get('taken_at', row.get('拍摄时间')) else None,
                'taken_by': str(row.get('taken_by', row.get('拍摄人', ''))),
            }
            if photo['point_no']:
                photos_data.append(photo)
    
    elif source_type == 'manual':
        for row in raw_rows:
            note = {
                'point_no': str(row.get('point_no', row.get('点位编号', ''))),
                'street_name': str(row.get('street_name', row.get('街道名称', ''))),
                'note_content': str(row.get('note_content', row.get('备注内容', ''))),
                'operator': str(row.get('operator', row.get('操作人', ''))),
            }
            if note['point_no'] and note['note_content']:
                notes_data.append(note)
    
    return gis_data, feedback_data, photos_data, notes_data

@app.route('/api/stats')
def api_stats():
    return jsonify(get_statistics())

@app.route('/api/points')
def api_points():
    status_filter = request.args.get('status', 'all')
    keyword = request.args.get('keyword', '')
    points = get_point_list(status_filter=status_filter, keyword=keyword)
    return jsonify(points)

@app.route('/api/point/<int:point_id>')
def api_point_detail(point_id):
    return jsonify(get_point_detail(point_id))

if __name__ == '__main__':
    init_db()
    from data_import import import_sample_data
    import_sample_data()
    print('\n' + '='*60)
    print('城市更新拆迁安置清单管理系统 启动成功!')
    print('访问地址: http://localhost:5001')
    print('样例数据: 3个点位，包含3种状态')
    print('  - CQ-2026-001: 已处理（顺利记录）')
    print('  - CQ-2026-002: 待核实（需人工确认）')
    print('  - CQ-2026-003: 需现场复看（GIS旧口径）')
    print('='*60 + '\n')
    app.run(debug=False, host='0.0.0.0', port=5001)
