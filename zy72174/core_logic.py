import json
from datetime import datetime
from database import get_db

STATUS_MAP = {
    'pending': '待处理',
    'processed': '已处理',
    'need_verify': '待核实',
    'need_onsite': '需现场复看'
}

STATUS_COLORS = {
    'pending': 'secondary',
    'processed': 'success',
    'need_verify': 'warning',
    'need_onsite': 'danger'
}

SOURCE_TYPE_MAP = {
    'gis_system': 'GIS系统',
    'gis_old': 'GIS旧口径',
    'inspection': '巡检照片',
    'manual': '人工备注',
    'feedback': '居民反馈'
}

def get_point_list(status_filter=None, keyword=None):
    conn = get_db()
    cursor = conn.cursor()
    
    sql = '''
    SELECT p.*,
           (SELECT COUNT(*) FROM feedback f WHERE f.point_id = p.id AND f.is_resolved = 0) as unresolved_feedback,
           (SELECT COUNT(*) FROM inspection_photos ph WHERE ph.point_id = p.id) as photo_count,
           (SELECT COUNT(*) FROM manual_notes mn WHERE mn.point_id = p.id) as note_count
    FROM points p
    WHERE 1=1
    '''
    params = []
    
    if status_filter and status_filter != 'all':
        sql += ' AND p.status = ?'
        params.append(status_filter)
    
    if keyword:
        sql += ' AND (p.point_no LIKE ? OR p.address LIKE ? OR p.district LIKE ?)'
        kw = f'%{keyword}%'
        params.extend([kw, kw, kw])
    
    sql += ' ORDER BY p.point_no'
    
    cursor.execute(sql, params)
    points = [dict(row) for row in cursor.fetchall()]
    
    for pt in points:
        pt['status_text'] = STATUS_MAP.get(pt['status'], pt['status'])
        pt['status_color'] = STATUS_COLORS.get(pt['status'], 'secondary')
    
    conn.close()
    return points

def get_point_detail(point_id):
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM points WHERE id = ?', (point_id,))
    point = dict(cursor.fetchone())
    point['status_text'] = STATUS_MAP.get(point['status'], point['status'])
    point['status_color'] = STATUS_COLORS.get(point['status'], 'secondary')
    
    cursor.execute('''
    SELECT pv.*, ds.source_name, ds.source_type
    FROM point_versions pv
    LEFT JOIN data_sources ds ON pv.source_id = ds.id
    WHERE pv.point_id = ?
    ORDER BY pv.version, pv.created_at
    ''', (point_id,))
    versions = [dict(row) for row in cursor.fetchall()]
    for v in versions:
        v['source_type_text'] = SOURCE_TYPE_MAP.get(v['source_type'], v['source_type'])
    
    cursor.execute('SELECT * FROM feedback WHERE point_id = ? ORDER BY created_at DESC', (point_id,))
    feedback = [dict(row) for row in cursor.fetchall()]
    
    cursor.execute('SELECT * FROM inspection_photos WHERE point_id = ? ORDER BY created_at DESC', (point_id,))
    photos = [dict(row) for row in cursor.fetchall()]
    
    cursor.execute('SELECT * FROM manual_notes WHERE point_id = ? ORDER BY created_at DESC', (point_id,))
    notes = [dict(row) for row in cursor.fetchall()]
    
    cursor.execute('SELECT * FROM review_records WHERE point_id = ? ORDER BY review_time DESC', (point_id,))
    reviews = [dict(row) for row in cursor.fetchall()]
    
    conn.close()
    return {
        'point': point,
        'versions': versions,
        'feedback': feedback,
        'photos': photos,
        'notes': notes,
        'reviews': reviews
    }

def update_point_field(point_id, field_name, new_value, operator, operation_note, source_id=None):
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute(f'SELECT {field_name} FROM points WHERE id = ?', (point_id,))
    old_value = str(cursor.fetchone()[field_name])
    
    if str(old_value) == str(new_value):
        conn.close()
        return False, '值未变化，无需更新'
    
    cursor.execute('SELECT current_version FROM points WHERE id = ?', (point_id,))
    current_ver = cursor.fetchone()['current_version']
    new_ver = current_ver + 1
    
    cursor.execute(
        '''INSERT INTO point_versions 
           (point_id, version, field_name, old_value, new_value, source_id, operator, operation_note)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
        (point_id, new_ver, field_name, old_value, str(new_value), source_id, operator, operation_note)
    )
    
    cursor.execute(
        f'UPDATE points SET {field_name} = ?, current_version = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
        (new_value, new_ver, point_id)
    )
    
    conn.commit()
    conn.close()
    return True, '更新成功'

def add_review(point_id, reviewer, review_result, review_note):
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute(
        '''INSERT INTO review_records (point_id, reviewer, review_result, review_note)
           VALUES (?, ?, ?, ?)''',
        (point_id, reviewer, review_result, review_note)
    )
    
    if review_result == 'pass':
        update_point_field(point_id, 'status', 'processed', reviewer, 
                          '人工复核通过，已完成处理')
    elif review_result == 'need_verify':
        update_point_field(point_id, 'status', 'need_verify', reviewer, 
                          '复核发现问题，需进一步核实')
    elif review_result == 'need_onsite':
        update_point_field(point_id, 'status', 'need_onsite', reviewer, 
                          '需现场复看确认')
    
    conn.commit()
    conn.close()
    return True

def get_statistics():
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('''
    SELECT status, COUNT(*) as cnt 
    FROM points 
    GROUP BY status
    ''')
    status_stats = {row['status']: row['cnt'] for row in cursor.fetchall()}
    
    cursor.execute('SELECT COUNT(*) as cnt FROM points')
    total = cursor.fetchone()['cnt']
    
    cursor.execute('SELECT COUNT(*) as cnt FROM feedback WHERE is_resolved = 0')
    unresolved = cursor.fetchone()['cnt']
    
    cursor.execute('SELECT COUNT(*) as cnt FROM point_versions')
    version_count = cursor.fetchone()['cnt']
    
    conn.close()
    
    return {
        'total': total,
        'processed': status_stats.get('processed', 0),
        'need_verify': status_stats.get('need_verify', 0),
        'need_onsite': status_stats.get('need_onsite', 0),
        'pending': status_stats.get('pending', 0),
        'unresolved_feedback': unresolved,
        'version_count': version_count
    }

def export_to_excel(status_filter=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    wb = Workbook()
    
    statuses_to_export = [status_filter] if status_filter and status_filter != 'all' \
                         else ['processed', 'need_verify', 'need_onsite']
    
    first_sheet = True
    for status in statuses_to_export:
        if first_sheet:
            ws = wb.active
            first_sheet = False
        else:
            ws = wb.create_sheet()
        
        ws.title = STATUS_MAP.get(status, status)
        
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['编号', '地址', '行政区', '性质', '面积(㎡)', '户数', '状态', 
                   '待处理反馈', '版本号', '原始来源', '最后更新', '处理备注']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
        
        points = get_point_list(status_filter=status)
        
        for row_idx, pt in enumerate(points, 2):
            detail = get_point_detail(pt['id'])
            versions = detail['versions']
            first_source = versions[0]['source_name'] if versions else '未知'
            
            last_note = ''
            if versions:
                last_version = versions[-1]
                last_note = last_version.get('operation_note', '')
            
            row_data = [
                pt['point_no'],
                pt['address'],
                pt['district'],
                pt['property_type'],
                pt['area'],
                pt['households'],
                pt['status_text'],
                pt['unresolved_feedback'],
                pt['current_version'],
                first_source,
                pt['updated_at'],
                last_note
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                if col in [1, 3, 5, 6, 7, 8, 9]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
        
        col_widths = [18, 40, 12, 12, 12, 8, 16, 12, 10, 30, 22, 40]
        for col, width in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width
        
        for row in ws.iter_rows(min_row=2, max_row=len(points) + 1):
            for cell in row:
                if cell.column == 7:
                    status_val = cell.value
                    if status_val == '已处理':
                        cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                    elif status_val == '待核实':
                        cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                    elif status_val == '需现场复看':
                        cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    
    summary_ws = wb.create_sheet('汇总说明', 0)
    summary_ws['A1'] = '城市更新拆迁安置清单 - 汇总说明'
    summary_ws['A1'].font = Font(bold=True, size=16)
    summary_ws.merge_cells('A1:F1')
    
    stats = get_statistics()
    summary_data = [
        ['导出时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
        ['总点数', stats['total']],
        ['已处理', f"{stats['processed']} 点"],
        ['待核实', f"{stats['need_verify']} 点"],
        ['需现场复看', f"{stats['need_onsite']} 点"],
        ['待处理反馈', f"{stats['unresolved_feedback']} 条"],
        ['版本变更记录', f"{stats['version_count']} 条"],
        ['', ''],
        ['说明', '本清单包含城市更新拆迁安置点位的完整信息，所有变更均保留历史版本记录。'],
        ['', '状态说明：'],
        ['', '  - 已处理：资料齐全，复核通过，可进入下一流程'],
        ['', '  - 待核实：存在异议，需进一步核对材料'],
        ['', '  - 需现场复看：需赴现场确认相关情况'],
    ]
    
    for row_idx, (key, value) in enumerate(summary_data, 3):
        summary_ws.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        summary_ws.cell(row=row_idx, column=2, value=value)
    
    summary_ws.column_dimensions['A'].width = 15
    summary_ws.column_dimensions['B'].width = 60
    
    return wb

def merge_and_match_points(gis_data, feedback_data, photos_data, notes_data, operator='system'):
    conn = get_db()
    cursor = conn.cursor()
    
    report = {
        'created': 0,
        'updated': 0,
        'conflicts': []
    }
    
    for pt in gis_data:
        point_no = pt['point_no']
        cursor.execute('SELECT id, area, status FROM points WHERE point_no = ?', (point_no,))
        existing = cursor.fetchone()
        
        if existing is None:
            cursor.execute(
                '''INSERT INTO points 
                   (point_no, address, district, gis_lng, gis_lat, property_type, area, households)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
                (point_no, pt['address'], pt['district'], pt['gis_lng'], pt['gis_lat'],
                 pt['property_type'], pt['area'], pt['households'])
            )
            point_id = cursor.lastrowid
            
            for field in ['address', 'area', 'property_type']:
                cursor.execute(
                    '''INSERT INTO point_versions 
                       (point_id, version, field_name, old_value, new_value, operator, operation_note)
                       VALUES (?, ?, ?, ?, ?, ?, ?)''',
                    (point_id, 1, field, None, str(pt[field]), operator, '数据导入创建')
                )
            report['created'] += 1
        else:
            point_id = existing['id']
            old_area = existing['area']
            if abs(float(old_area) - float(pt['area'])) > 0.1:
                report['conflicts'].append({
                    'point_no': point_no,
                    'field': 'area',
                    'old': old_area,
                    'new': pt['area'],
                    'note': '面积差异需要复核'
                })
                
                cursor.execute('SELECT current_version FROM points WHERE id = ?', (point_id,))
                current_ver = cursor.fetchone()['current_version']
                new_ver = current_ver + 1
                
                cursor.execute(
                    '''INSERT INTO point_versions 
                       (point_id, version, field_name, old_value, new_value, operator, operation_note)
                       VALUES (?, ?, ?, ?, ?, ?, ?)''',
                    (point_id, new_ver, 'area', str(old_area), str(pt['area']), 
                     operator, 'GIS数据更新，面积有差异')
                )
                
                cursor.execute(
                    '''UPDATE points SET area = ?, current_version = ?, status = 'need_verify', 
                       updated_at = CURRENT_TIMESTAMP WHERE id = ?''',
                    (pt['area'], new_ver, point_id)
                )
                report['updated'] += 1
    
    conn.commit()
    conn.close()
    return report
