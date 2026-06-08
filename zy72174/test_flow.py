import os
import json
import io
from database import get_db, init_db
from core_logic import (
    merge_and_match_points, get_statistics, get_point_list,
    get_point_detail, export_to_excel, add_review
)
from data_import import (
    SAMPLE_GIS_POINTS, SAMPLE_RESIDENT_FEEDBACK,
    SAMPLE_INSPECTION_PHOTOS, SAMPLE_MANUAL_NOTES,
    OLD_VERSION_DATA
)
from openpyxl import load_workbook

if os.path.exists('resettlement.db'):
    os.remove('resettlement.db')

init_db()

UPLOAD_DIR = 'uploads'
os.makedirs(UPLOAD_DIR, exist_ok=True)

print("=" * 70)
print("城市更新拆迁安置清单 - 完整操作路测试")
print("=" * 70)

# ---- 第1步：通过文件导入走 GIS 点位 ----
gis_json_path = os.path.join(UPLOAD_DIR, 'gis_points.json')
with open(gis_json_path, 'w', encoding='utf-8') as f:
    json.dump(SAMPLE_GIS_POINTS, f, ensure_ascii=False, indent=2)

from app import _parse_uploaded_file

gis_data, fb1, ph1, nt1 = _parse_uploaded_file(gis_json_path, 'gis_system')
print(f"\n1. 解析GIS JSON文件: {len(gis_data)} 个点位")
assert len(gis_data) == 3, f"期望3个点位, 实际{len(gis_data)}"

conn = get_db()
cursor = conn.cursor()
cursor.execute(
    "INSERT INTO data_sources (source_type, source_name, source_file, imported_by) VALUES (?, ?, ?, ?)",
    ('gis_system', '2026年城市更新GIS点位普查_第3批', 'gis_points.json', '何工')
)
gis_source_id = cursor.lastrowid
conn.commit()
conn.close()

report1 = merge_and_match_points(gis_data, fb1, ph1, nt1, operator='何工', source_id=gis_source_id)
print(f"   归并结果: 新建{report1['created']}个, 更新{report1['updated']}个, 冲突{len(report1['conflicts'])}个")
assert report1['created'] == 3, f"期望新建3个, 实际{report1['created']}"

stats = get_statistics()
print(f"   当前统计: 总{stats['total']}个点位, 待处理{stats['pending']}个")

# ---- 第2步：导入居民反馈 ----
fb_json_path = os.path.join(UPLOAD_DIR, 'feedback.json')
with open(fb_json_path, 'w', encoding='utf-8') as f:
    json.dump(SAMPLE_RESIDENT_FEEDBACK, f, ensure_ascii=False, indent=2)

fb_empty, fb_data, ph_empty, nt_empty = _parse_uploaded_file(fb_json_path, 'feedback')
print(f"\n2. 解析反馈JSON文件: {len(fb_data)} 条反馈")
assert len(fb_data) == 3, f"期望3条反馈, 实际{len(fb_data)}"

conn = get_db()
cursor = conn.cursor()
cursor.execute(
    "INSERT INTO data_sources (source_type, source_name, source_file, imported_by) VALUES (?, ?, ?, ?)",
    ('feedback', '居民反馈汇总', 'feedback.json', '何工')
)
fb_source_id = cursor.lastrowid
conn.commit()
conn.close()

report2 = merge_and_match_points([], fb_data, [], [], operator='何工', source_id=fb_source_id)
print(f"   归并结果: 反馈导入{report2['feedback_added']}条")
assert report2['feedback_added'] == 3, f"期望导入3条反馈, 实际{report2['feedback_added']}"

# ---- 第3步：导入巡检照片 ----
ph_json_path = os.path.join(UPLOAD_DIR, 'photos.json')
with open(ph_json_path, 'w', encoding='utf-8') as f:
    json.dump(SAMPLE_INSPECTION_PHOTOS, f, ensure_ascii=False, indent=2)

ph_empty, fb_empty, ph_data, nt_empty = _parse_uploaded_file(ph_json_path, 'inspection')
print(f"\n3. 解析照片JSON文件: {len(ph_data)} 张照片")
assert len(ph_data) == 5, f"期望5张照片, 实际{len(ph_data)}"

conn = get_db()
cursor = conn.cursor()
cursor.execute(
    "INSERT INTO data_sources (source_type, source_name, source_file, imported_by) VALUES (?, ?, ?, ?)",
    ('inspection', '巡检照片', 'photos.json', '巡检组')
)
ph_source_id = cursor.lastrowid
conn.commit()
conn.close()

report3 = merge_and_match_points([], [], ph_data, [], operator='巡检组', source_id=ph_source_id)
print(f"   归并结果: 照片导入{report3['photos_added']}张")
assert report3['photos_added'] == 5, f"期望导入5张照片, 实际{report3['photos_added']}"

# ---- 第4步：导入街道手改备注 ----
nt_json_path = os.path.join(UPLOAD_DIR, 'notes.json')
with open(nt_json_path, 'w', encoding='utf-8') as f:
    json.dump(SAMPLE_MANUAL_NOTES, f, ensure_ascii=False, indent=2)

nt_empty_gis, fb_empty_gis, ph_empty_gis, nt_data = _parse_uploaded_file(nt_json_path, 'manual')
print(f"\n4. 解析备注JSON文件: {len(nt_data)} 条备注")
assert len(nt_data) == 3, f"期望3条备注, 实际{len(nt_data)}"

conn = get_db()
cursor = conn.cursor()
cursor.execute(
    "INSERT INTO data_sources (source_type, source_name, source_file, imported_by) VALUES (?, ?, ?, ?)",
    ('manual', '街道手改备注', 'notes.json', '街道办')
)
nt_source_id = cursor.lastrowid
conn.commit()
conn.close()

report4 = merge_and_match_points([], [], [], nt_data, operator='街道办', source_id=nt_source_id)
print(f"   归并结果: 备注导入{report4['notes_added']}条")
assert report4['notes_added'] == 3, f"期望导入3条备注, 实际{report4['notes_added']}"

# ---- 第5步：手动应用状态变更(模拟data_import中的OLD_VERSION_DATA) ----
print(f"\n5. 应用历史版本变更:")
for ov in OLD_VERSION_DATA:
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT id, current_version FROM points WHERE point_no = ?', (ov['point_no'],))
    row = cursor.fetchone()
    if row:
        pid = row['id']
        current_ver = row['current_version']
        new_ver = current_ver + 1
        cursor.execute(
            '''INSERT INTO point_versions 
               (point_id, version, field_name, old_value, new_value, source_id, operator, operation_note)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
            (pid, new_ver, ov['field_name'], ov['old_value'], ov['new_value'],
             nt_source_id, ov['operator'], ov['operation_note'])
        )
        if ov['field_name'] == 'status':
            cursor.execute(
                "UPDATE points SET status = ?, current_version = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (ov['new_value'], new_ver, pid)
            )
        else:
            cursor.execute(
                "UPDATE points SET current_version = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (new_ver, pid)
            )
        conn.commit()
        print(f"   {ov['point_no']}: {ov['field_name']} {ov['old_value']} -> {ov['new_value']}")
    conn.close()

# ---- 第6步：人工复核 ----
print(f"\n6. 人工复核:")
conn = get_db()
cursor = conn.cursor()
cursor.execute("SELECT id, point_no FROM points WHERE point_no = 'CQ-2026-001'")
row = cursor.fetchone()
conn.close()

add_review(row['id'], '何工', 'pass', '资料齐全，面积无异议，居民诉求已登记，可进入下一流程')
print(f"   CQ-2026-001: 复核通过 -> 已处理")

# ---- 第7步：核对页面状态 ----
print(f"\n7. 核对页面状态:")

points = get_point_list()
status_check = {}
for p in points:
    status_check[p['point_no']] = p['status_text']
    print(f"   {p['point_no']}: {p['status_text']} (v{p['current_version']})")

assert status_check.get('CQ-2026-001') == '已处理', f"CQ-2026-001应为已处理, 实际{status_check.get('CQ-2026-001')}"
assert status_check.get('CQ-2026-002') == '待核实', f"CQ-2026-002应为待核实, 实际{status_check.get('CQ-2026-002')}"
assert status_check.get('CQ-2026-003') == '需现场复看', f"CQ-2026-003应为需现场复看, 实际{status_check.get('CQ-2026-003')}"

# ---- 第8步：核对详情 ----
print(f"\n8. 核对详情数据:")

detail_001 = get_point_detail(1)
fb_count_001 = len(detail_001['feedback'])
ph_count_001 = len(detail_001['photos'])
nt_count_001 = len(detail_001['notes'])
rv_count_001 = len(detail_001['reviews'])
print(f"   CQ-2026-001: 反馈{fb_count_001}条, 照片{ph_count_001}张, 备注{nt_count_001}条, 复核{rv_count_001}条")
assert fb_count_001 == 1, f"001应有1条反馈, 实际{fb_count_001}"
assert ph_count_001 == 2, f"001应有2张照片, 实际{ph_count_001}"
assert nt_count_001 == 1, f"001应有1条备注, 实际{nt_count_001}"
assert rv_count_001 == 1, f"001应有1条复核, 实际{rv_count_001}"

detail_002 = get_point_detail(2)
fb_count_002 = len(detail_002['feedback'])
nt_count_002 = len(detail_002['notes'])
unresolved_002 = sum(1 for fb in detail_002['feedback'] if not fb['is_resolved'])
print(f"   CQ-2026-002: 反馈{fb_count_002}条(待处理{unresolved_002}), 备注{nt_count_002}条")
assert fb_count_002 == 1, f"002应有1条反馈, 实际{fb_count_002}"
assert unresolved_002 == 1, f"002应有1条待处理反馈"
assert nt_count_002 == 1, f"002应有1条备注, 实际{nt_count_002}"

detail_003 = get_point_detail(3)
ver_count_003 = len(detail_003['versions'])
area_versions = [v for v in detail_003['versions'] if v['field_name'] == 'area']
print(f"   CQ-2026-003: 版本{ver_count_003}条, 面积变更{len(area_versions)}次")
assert len(area_versions) >= 2, f"003面积应至少变更2次(初始+旧口径), 实际{len(area_versions)}次"

# ---- 第9步：导出核对 ----
print(f"\n9. 导出Excel核对:")
wb = export_to_excel()
output = io.BytesIO()
wb.save(output)
output.seek(0)
wb2 = load_workbook(output)

for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    if sheet_name == '汇总说明':
        print(f"   汇总说明: OK")
        continue
    row_count = ws.max_row - 1
    point_nos = []
    for r in range(2, ws.max_row + 1):
        point_nos.append(ws.cell(row=r, column=1).value)
    print(f"   {sheet_name}: {row_count}条记录 [{', '.join(str(p) for p in point_nos if p)}]")

stats = get_statistics()
assert stats['total'] == 3, f"总数应为3, 实际{stats['total']}"
assert stats['processed'] == 1, f"已处理应为1, 实际{stats['processed']}"
assert stats['need_verify'] == 1, f"待核实应为1, 实际{stats['need_verify']}"
assert stats['need_onsite'] == 1, f"需现场复看应为1, 实际{stats['need_onsite']}"

print("\n" + "=" * 70)
print("✅ 全部断言通过！完整操作路验证成功")
print("   导入(解析文件+归并) → 状态变更 → 人工复核 → 页面状态 → 导出内容")
print("=" * 70)
