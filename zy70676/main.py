from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import datetime
import json
import os

from database import get_db, init_db, SamplingPoint, Unit, Parameter, Threshold, FieldRecord, LabResult, AuditLog, ReviewReport
import schemas
from services import (
    UnitConversionService, PointMatchingService, ThresholdService,
    MissingSampleService, LabResultService, AuditService
)

app = FastAPI(title="水质采样实验室结果单位换算API", version="1.0.0")


@app.on_event("startup")
async def startup_event():
    init_db()


@app.post("/sampling-points/", response_model=schemas.SamplingPointResponse, tags=["采样点"])
def create_sampling_point(point: schemas.SamplingPointCreate, db: Session = Depends(get_db)):
    existing = db.query(SamplingPoint).filter(SamplingPoint.point_code == point.point_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="点位编号已存在")
    db_point = SamplingPoint(**point.model_dump())
    db.add(db_point)
    db.commit()
    db.refresh(db_point)
    return db_point


@app.get("/sampling-points/", response_model=List[schemas.SamplingPointResponse], tags=["采样点"])
def list_sampling_points(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return db.query(SamplingPoint).offset(skip).limit(limit).all()


@app.get("/sampling-points/{point_id}", response_model=schemas.SamplingPointResponse, tags=["采样点"])
def get_sampling_point(point_id: int, db: Session = Depends(get_db)):
    point = db.query(SamplingPoint).filter(SamplingPoint.id == point_id).first()
    if not point:
        raise HTTPException(status_code=404, detail="采样点不存在")
    return point


@app.patch("/sampling-points/{point_id}", response_model=schemas.SamplingPointResponse, tags=["采样点"])
def update_sampling_point(point_id: int, update: schemas.SamplingPointUpdate, db: Session = Depends(get_db)):
    point = db.query(SamplingPoint).filter(SamplingPoint.id == point_id).first()
    if not point:
        raise HTTPException(status_code=404, detail="采样点不存在")
    for key, value in update.model_dump(exclude_unset=True).items():
        setattr(point, key, value)
    db.commit()
    db.refresh(point)
    return point


@app.post("/units/", response_model=schemas.UnitResponse, tags=["单位"])
def create_unit(unit: schemas.UnitCreate, db: Session = Depends(get_db)):
    existing = db.query(Unit).filter(Unit.unit_code == unit.unit_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="单位编码已存在")
    db_unit = Unit(**unit.model_dump())
    db.add(db_unit)
    db.commit()
    db.refresh(db_unit)
    return db_unit


@app.get("/units/", response_model=List[schemas.UnitResponse], tags=["单位"])
def list_units(db: Session = Depends(get_db)):
    return db.query(Unit).all()


@app.post("/parameters/", response_model=schemas.ParameterResponse, tags=["参数"])
def create_parameter(param: schemas.ParameterCreate, db: Session = Depends(get_db)):
    existing = db.query(Parameter).filter(Parameter.param_code == param.param_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="参数编码已存在")
    db_param = Parameter(**param.model_dump())
    db.add(db_param)
    db.commit()
    db.refresh(db_param)
    return db_param


@app.get("/parameters/", response_model=List[schemas.ParameterResponse], tags=["参数"])
def list_parameters(db: Session = Depends(get_db)):
    return db.query(Parameter).all()


@app.post("/thresholds/", response_model=schemas.ThresholdResponse, tags=["阈值"])
def create_threshold(threshold: schemas.ThresholdCreate, db: Session = Depends(get_db)):
    db_threshold = Threshold(**threshold.model_dump())
    db.add(db_threshold)
    db.commit()
    db.refresh(db_threshold)
    return db_threshold


@app.get("/thresholds/", response_model=List[schemas.ThresholdResponse], tags=["阈值"])
def list_thresholds(db: Session = Depends(get_db)):
    return db.query(Threshold).all()


@app.post("/field-records/", response_model=schemas.FieldRecordResponse, tags=["现场记录"])
def create_field_record(record: schemas.FieldRecordCreate, db: Session = Depends(get_db)):
    existing = db.query(FieldRecord).filter(FieldRecord.record_code == record.record_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="记录编号已存在")
    if not PointMatchingService.validate_sampling_point(db, record.sampling_point_id):
        raise HTTPException(status_code=400, detail="采样点无效")
    db_record = FieldRecord(**record.model_dump())
    db.add(db_record)
    db.commit()
    db.refresh(db_record)
    return db_record


@app.get("/field-records/", response_model=List[schemas.FieldRecordResponse], tags=["现场记录"])
def list_field_records(skip: int = 0, limit: int = 100, status: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(FieldRecord)
    if status:
        query = query.filter(FieldRecord.status == status)
    return query.offset(skip).limit(limit).all()


@app.get("/field-records/{record_id}", response_model=schemas.FieldRecordResponse, tags=["现场记录"])
def get_field_record(record_id: int, db: Session = Depends(get_db)):
    record = db.query(FieldRecord).filter(FieldRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="现场记录不存在")
    return record


@app.patch("/field-records/{record_id}", response_model=schemas.FieldRecordResponse, tags=["现场记录"])
def update_field_record(record_id: int, update: schemas.FieldRecordUpdate, db: Session = Depends(get_db)):
    record = db.query(FieldRecord).filter(FieldRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="现场记录不存在")
    for key, value in update.model_dump(exclude_unset=True).items():
        setattr(record, key, value)
    db.commit()
    db.refresh(record)
    return record


@app.post("/lab-results/", response_model=schemas.LabResultResponse, tags=["实验室结果"])
def create_lab_result(lab_result: schemas.LabResultCreate, db: Session = Depends(get_db)):
    try:
        return LabResultService.create_lab_result(db, lab_result)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/lab-results/", response_model=List[schemas.LabResultResponse], tags=["实验室结果"])
def list_lab_results(
    field_record_id: Optional[int] = None,
    parameter_id: Optional[int] = None,
    is_approved: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    query = db.query(LabResult)
    if field_record_id:
        query = query.filter(LabResult.field_record_id == field_record_id)
    if parameter_id:
        query = query.filter(LabResult.parameter_id == parameter_id)
    if is_approved is not None:
        query = query.filter(LabResult.is_approved == is_approved)
    return query.all()


@app.get("/lab-results/{result_id}", response_model=schemas.LabResultResponse, tags=["实验室结果"])
def get_lab_result(result_id: int, db: Session = Depends(get_db)):
    result = db.query(LabResult).filter(LabResult.id == result_id).first()
    if not result:
        raise HTTPException(status_code=404, detail="实验室结果不存在")
    return result


@app.post("/lab-results/{result_id}/approve", response_model=schemas.LabResultResponse, tags=["实验室结果"])
def approve_lab_result(result_id: int, request: schemas.ApprovalRequest, db: Session = Depends(get_db)):
    try:
        return LabResultService.approve_lab_result(db, result_id, request.operator, request.conclusion)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/lab-results/{result_id}/correct", response_model=schemas.LabResultResponse, tags=["实验室结果"])
def correct_lab_result(result_id: int, correction: schemas.CorrectionRequest, db: Session = Depends(get_db)):
    try:
        return LabResultService.correct_lab_result(db, result_id, correction)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/lab-results/{result_id}/withdraw", response_model=schemas.LabResultResponse, tags=["实验室结果"])
def withdraw_lab_result(result_id: int, request: schemas.ApprovalRequest, db: Session = Depends(get_db)):
    try:
        return LabResultService.withdraw_lab_result(db, result_id, request.operator, request.conclusion or "撤回")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/lab-results/{result_id}/threshold-check", tags=["实验室结果"])
def check_lab_result_threshold(result_id: int, db: Session = Depends(get_db)):
    result = db.query(LabResult).filter(LabResult.id == result_id).first()
    if not result:
        raise HTTPException(status_code=404, detail="实验室结果不存在")
    return ThresholdService.check_threshold(
        db, result.parameter_id, result.raw_value, result.raw_unit_id
    )


@app.get("/audit-logs/", response_model=List[schemas.AuditLogResponse], tags=["审计日志"])
def list_audit_logs(
    entity_type: Optional[str] = None,
    operation_type: Optional[str] = None,
    operator: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(AuditLog)
    if entity_type:
        query = query.filter(AuditLog.entity_type == entity_type)
    if operation_type:
        query = query.filter(AuditLog.operation_type == operation_type)
    if operator:
        query = query.filter(AuditLog.operator == operator)
    return query.order_by(AuditLog.created_at.desc()).all()


@app.get("/audit-logs/{log_id}", tags=["审计日志"])
def get_audit_log(log_id: int, db: Session = Depends(get_db)):
    log = db.query(AuditLog).filter(AuditLog.id == log_id).first()
    if not log:
        raise HTTPException(status_code=404, detail="审计日志不存在")
    return {
        "id": log.id,
        "operation_type": log.operation_type,
        "entity_type": log.entity_type,
        "entity_id": log.entity_id,
        "original_data": json.loads(log.original_data) if log.original_data else None,
        "modified_data": json.loads(log.modified_data) if log.modified_data else None,
        "operator": log.operator,
        "conclusion": log.conclusion,
        "created_at": log.created_at
    }


@app.post("/review-reports/", response_model=schemas.ReviewReportResponse, tags=["复核报告"])
def create_review_report(report: schemas.ReviewReportCreate, db: Session = Depends(get_db)):
    existing = db.query(ReviewReport).filter(ReviewReport.report_code == report.report_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="报告编号已存在")

    missing_info = MissingSampleService.find_missing_samples(db, report.field_record_ids)

    issues = []
    for record_id in report.field_record_ids:
        results = db.query(LabResult).filter(LabResult.field_record_id == record_id).all()
        for result in results:
            threshold_result = ThresholdService.check_threshold(
                db, result.parameter_id, result.raw_value, result.raw_unit_id
            )
            if not threshold_result.get("passed"):
                field_record = db.query(FieldRecord).filter(FieldRecord.id == record_id).first()
                parameter = db.query(Parameter).filter(Parameter.id == result.parameter_id).first()
                issues.append({
                    "record_code": field_record.record_code if field_record else None,
                    "parameter": parameter.param_name if parameter else None,
                    "value": result.raw_value,
                    "message": threshold_result.get("message")
                })

    db_report = ReviewReport(
        report_code=report.report_code,
        field_record_ids=json.dumps(report.field_record_ids),
        reviewer=report.reviewer,
        issues=json.dumps(issues, ensure_ascii=False) if issues else None,
        missing_samples=json.dumps(missing_info.get("details", {}), ensure_ascii=False) if missing_info.get("has_missing") else None,
        conclusion="待复核"
    )
    db.add(db_report)
    db.commit()
    db.refresh(db_report)
    return db_report


@app.get("/review-reports/", response_model=List[schemas.ReviewReportResponse], tags=["复核报告"])
def list_review_reports(
    status: Optional[str] = None,
    reviewer: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(ReviewReport)
    if status:
        query = query.filter(ReviewReport.status == status)
    if reviewer:
        query = query.filter(ReviewReport.reviewer == reviewer)
    return query.all()


@app.get("/review-reports/{report_id}", tags=["复核报告"])
def get_review_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(ReviewReport).filter(ReviewReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="复核报告不存在")
    return {
        "id": report.id,
        "report_code": report.report_code,
        "field_record_ids": json.loads(report.field_record_ids) if report.field_record_ids else [],
        "reviewer": report.reviewer,
        "review_time": report.review_time,
        "status": report.status,
        "conclusion": report.conclusion,
        "issues": json.loads(report.issues) if report.issues else None,
        "missing_samples": json.loads(report.missing_samples) if report.missing_samples else None,
        "created_at": report.created_at
    }


@app.post("/review-reports/{report_id}/finalize", response_model=schemas.ReviewReportResponse, tags=["复核报告"])
def finalize_review_report(report_id: int, operator: str, conclusion: str, db: Session = Depends(get_db)):
    report = db.query(ReviewReport).filter(ReviewReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="复核报告不存在")
    if report.status == "finalized":
        raise HTTPException(status_code=400, detail="报告已完成")

    report.status = "finalized"
    report.conclusion = conclusion
    db.commit()
    db.refresh(report)

    AuditService.log_operation(
        db, "finalize", "ReviewReport", report_id,
        operator, conclusion
    )

    return report


@app.get("/review-reports/{report_id}/export", tags=["复核报告"])
def export_review_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(ReviewReport).filter(ReviewReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="复核报告不存在")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "复核报告"

        ws["A1"] = "水质采样实验室结果复核报告"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:F1")
        ws["A1"].alignment = Alignment(horizontal="center")

        ws["A3"] = f"报告编号: {report.report_code}"
        ws["A4"] = f"复核人: {report.reviewer}"
        ws["A5"] = f"复核时间: {report.review_time.strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A6"] = f"状态: {report.status}"
        ws["A7"] = f"结论: {report.conclusion or '-'}"

        ws["A9"] = "采样记录详情"
        ws["A9"].font = Font(bold=True)

        headers = ["记录编号", "采样点", "参数", "原值", "单位", "换算值", "标准单位", "审核状态"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=10, column=col, value=header)
            cell.font = Font(bold=True)

        field_record_ids = json.loads(report.field_record_ids) if report.field_record_ids else []
        row = 11
        for record_id in field_record_ids:
            field_record = db.query(FieldRecord).filter(FieldRecord.id == record_id).first()
            if not field_record:
                continue
            results = db.query(LabResult).filter(LabResult.field_record_id == record_id).all()
            for result in results:
                parameter = db.query(Parameter).filter(Parameter.id == result.parameter_id).first()
                raw_unit = db.query(Unit).filter(Unit.id == result.raw_unit_id).first()
                std_unit = db.query(Unit).filter(Unit.id == result.standard_unit_id).first() if result.standard_unit_id else None

                ws.cell(row=row, column=1, value=field_record.record_code)
                ws.cell(row=row, column=2, value=field_record.sampling_point.point_name)
                ws.cell(row=row, column=3, value=parameter.param_name if parameter else "-")
                ws.cell(row=row, column=4, value=result.raw_value)
                ws.cell(row=row, column=5, value=raw_unit.unit_name if raw_unit else "-")
                ws.cell(row=row, column=6, value=result.converted_value or "-")
                ws.cell(row=row, column=7, value=std_unit.unit_name if std_unit else "-")
                ws.cell(row=row, column=8, value="已审核" if result.is_approved else "未审核")
                row += 1

        if report.issues:
            row += 2
            ws.cell(row=row, column=1, value="问题记录").font = Font(bold=True)
            issues = json.loads(report.issues)
            row += 1
            for issue in issues:
                ws.cell(row=row, column=1, value=f"{issue.get('record_code', '-')}: {issue.get('parameter', '-')} - {issue.get('message', '-')}")
                row += 1

        if report.missing_samples:
            row += 2
            ws.cell(row=row, column=1, value="缺样记录").font = Font(bold=True)
            missing = json.loads(report.missing_samples)
            row += 1
            for record_code, info in missing.items():
                ws.cell(row=row, column=1, value=f"{record_code} ({info['point_name']}): 缺少参数 - {', '.join([p['param_name'] for p in info['missing_params']])}")
                row += 1

        os.makedirs("exports", exist_ok=True)
        filename = f"exports/{report.report_code}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        wb.save(filename)

        report.export_path = filename
        db.commit()

        return FileResponse(filename, filename=f"{report.report_code}.xlsx")
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装openpyxl，无法导出Excel")


@app.get("/utils/convert-unit", tags=["工具"])
def convert_unit(value: float, from_unit_id: int, to_unit_id: int, db: Session = Depends(get_db)):
    result = UnitConversionService.convert_value(db, value, from_unit_id, to_unit_id)
    if result is None:
        raise HTTPException(status_code=400, detail="单位换算失败")
    return {"value": value, "from_unit_id": from_unit_id, "to_unit_id": to_unit_id, "converted_value": result}


@app.get("/utils/missing-samples", tags=["工具"])
def check_missing_samples(field_record_ids: str = Query(..., description="现场记录ID列表，用逗号分隔"), db: Session = Depends(get_db)):
    ids = [int(id.strip()) for id in field_record_ids.split(",")]
    return MissingSampleService.find_missing_samples(db, ids)
