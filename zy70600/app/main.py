from fastapi import FastAPI, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from datetime import datetime
from typing import List, Optional
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.database import engine, get_db, Base
from app import models, schemas, services

Base.metadata.create_all(bind=engine)

app = FastAPI(
    title="冷链温控证据签收复核赔付结论后端API",
    description="处理冷链箱温度监控、门店签收、异常复核、赔付结论全流程管理",
    version="1.0.0"
)


@app.exception_handler(services.StatusTransitionError)
def handle_status_transition_error(request, exc):
    from fastapi.responses import JSONResponse
    return JSONResponse(
        status_code=400,
        content={"success": False, "message": str(exc)}
    )


@app.exception_handler(services.DuplicateUploadError)
def handle_duplicate_upload_error(request, exc):
    from fastapi.responses import JSONResponse
    return JSONResponse(
        status_code=409,
        content={"success": False, "message": str(exc)}
    )


@app.exception_handler(ValueError)
def handle_value_error(request, exc):
    from fastapi.responses import JSONResponse
    return JSONResponse(
        status_code=404,
        content={"success": False, "message": str(exc)}
    )


@app.post("/api/boxes/", response_model=schemas.ColdChainBox)
def create_box(box: schemas.ColdChainBoxCreate, db: Session = Depends(get_db)):
    try:
        return services.create_box(db, box)
    except services.DuplicateUploadError as e:
        raise HTTPException(status_code=409, detail=str(e))


@app.get("/api/boxes/", response_model=List[schemas.ColdChainBox])
def list_boxes(
    status: Optional[str] = None,
    batch_no: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    return services.list_boxes(db, status, batch_no, skip, limit)


@app.get("/api/boxes/{box_code}", response_model=schemas.BoxDetailResponse)
def get_box_detail(box_code: str, db: Session = Depends(get_db)):
    db_box = services.get_box_detail(db, box_code)
    if not db_box:
        raise HTTPException(status_code=404, detail=f"冷链箱 {box_code} 不存在")
    return db_box


@app.get("/api/boxes/{box_code}/audit-logs", response_model=List[schemas.AuditLog])
def get_box_audit_logs(box_code: str, db: Session = Depends(get_db)):
    return services.get_audit_logs_by_box_code(db, box_code)


@app.post("/api/boxes/{box_code}/status", response_model=schemas.ColdChainBox)
def transition_box_status(
    box_code: str,
    request: schemas.StatusTransitionRequest,
    db: Session = Depends(get_db)
):
    try:
        return services.transition_box_status(db, box_code, request)
    except services.StatusTransitionError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/boxes/{box_code}/correction", response_model=schemas.ColdChainBox)
def manual_correction(
    box_code: str,
    request: schemas.ManualCorrectionRequest,
    db: Session = Depends(get_db)
):
    try:
        return services.manual_correction(db, box_code, request)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/boxes/{box_code}/close", response_model=schemas.ColdChainBox)
def close_box(
    box_code: str,
    operator: str,
    db: Session = Depends(get_db)
):
    try:
        return services.close_box(db, box_code, operator)
    except services.StatusTransitionError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/temperature/", response_model=schemas.TemperatureSample)
def create_temperature_sample(
    sample: schemas.TemperatureSampleCreate,
    db: Session = Depends(get_db)
):
    try:
        return services.create_temperature_sample(db, sample)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@app.post("/api/signoffs/", response_model=schemas.StoreSignoff)
def create_signoff(signoff: schemas.StoreSignoffCreate, db: Session = Depends(get_db)):
    try:
        return services.create_signoff(db, signoff)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@app.post("/api/signoffs/{signoff_id}/status", response_model=schemas.StoreSignoff)
def transition_signoff_status(
    signoff_id: int,
    target_status: str,
    operator: str,
    db: Session = Depends(get_db)
):
    try:
        return services.transition_signoff_status(db, signoff_id, target_status, operator)
    except services.StatusTransitionError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/photos/", response_model=schemas.PhotoEvidence)
def create_photo_evidence(photo: schemas.PhotoEvidenceCreate, db: Session = Depends(get_db)):
    try:
        return services.create_photo_evidence(db, photo)
    except services.DuplicateUploadError as e:
        raise HTTPException(status_code=409, detail=str(e))


@app.post("/api/reviews/", response_model=schemas.ExceptionReview)
def create_exception_review(review: schemas.ExceptionReviewCreate, db: Session = Depends(get_db)):
    try:
        return services.create_exception_review(db, review)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@app.post("/api/compensations/", response_model=schemas.CompensationConclusion)
def create_compensation_conclusion(
    conclusion: schemas.CompensationConclusionCreate,
    db: Session = Depends(get_db)
):
    try:
        return services.create_compensation_conclusion(db, conclusion)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@app.post("/api/export/")
def export_data(export_request: schemas.ExportRequest, db: Session = Depends(get_db)):
    query = db.query(models.ColdChainBox)
    
    if export_request.box_codes:
        query = query.filter(models.ColdChainBox.box_code.in_(export_request.box_codes))
    if export_request.status:
        query = query.filter(models.ColdChainBox.status.in_(export_request.status))
    if export_request.start_date:
        query = query.filter(models.ColdChainBox.created_at >= export_request.start_date)
    if export_request.end_date:
        query = query.filter(models.ColdChainBox.created_at <= export_request.end_date)
    
    boxes = query.all()
    
    if export_request.export_format == "json":
        result = []
        for box in boxes:
            result.append({
                "box_code": box.box_code,
                "batch_no": box.batch_no,
                "product_name": box.product_name,
                "status": box.status,
                "temperature_min": box.temperature_min,
                "temperature_max": box.temperature_max,
                "created_at": box.created_at.isoformat(),
                "temperature_samples": [
                    {
                        "time": s.sample_time.isoformat(),
                        "temperature": s.temperature,
                        "is_anomaly": s.is_anomaly,
                        "probe_id": s.probe_id
                    }
                    for s in box.temperature_samples
                ],
                "signoffs": [
                    {
                        "store_code": s.store_code,
                        "store_name": s.store_name,
                        "signoff_person": s.signoff_person,
                        "signoff_time": s.signoff_time.isoformat() if s.signoff_time else None,
                        "temperature_arrival": s.temperature_arrival,
                        "has_exception": s.has_exception,
                        "exception_desc": s.exception_desc,
                        "status": s.status
                    }
                    for s in box.signoffs
                ],
                "reviews": [
                    {
                        "reviewer": r.reviewer,
                        "review_time": r.review_time.isoformat() if r.review_time else None,
                        "original_input": r.original_input,
                        "review_result": r.review_result,
                        "review_comment": r.review_comment,
                        "temperature_violation": r.temperature_violation,
                        "compensation_eligible": r.compensation_eligible,
                        "status": r.status
                    }
                    for r in box.reviews
                ],
                "compensations": [
                    {
                        "amount": c.compensation_amount,
                        "reason": c.compensation_reason,
                        "processor": c.processor,
                        "approved_by": c.approved_by,
                        "conclusion_time": c.conclusion_time.isoformat() if c.conclusion_time else None,
                        "status": c.status
                    }
                    for c in box.compensations
                ],
                "audit_logs": [
                    {
                        "action_type": a.action_type,
                        "operator": a.operator,
                        "old_status": a.old_status,
                        "new_status": a.new_status,
                        "comment": a.comment,
                        "change_details": a.change_details,
                        "created_at": a.created_at.isoformat()
                    }
                    for a in box.audit_logs
                ]
            })
        return {"success": True, "message": "导出成功", "data": result}
    
    wb = Workbook()
    
    ws1 = wb.active
    ws1.title = "冷链箱信息"
    headers1 = ["箱号", "批次号", "产品名称", "状态", "温度范围", "创建时间"]
    ws1.append(headers1)
    
    for box in boxes:
        ws1.append([
            box.box_code,
            box.batch_no,
            box.product_name,
            box.status,
            f"{box.temperature_min} ~ {box.temperature_max}",
            box.created_at.strftime("%Y-%m-%d %H:%M:%S")
        ])
    
    ws2 = wb.create_sheet("异常报告")
    headers2 = [
        "箱号", "批次号", "门店编码", "门店名称", "签收人", "签收时间",
        "到货温度", "是否异常", "异常描述",
        "复核人", "复核时间", "复核原始输入", "复核结论", "复核意见",
        "温度违规", "符合赔付",
        "赔付金额", "赔付原因", "赔付处理人", "审批人",
        "状态"
    ]
    ws2.append(headers2)
    
    for box in boxes:
        for review in box.reviews:
            signoff = db.query(models.StoreSignoff).filter(models.StoreSignoff.id == review.signoff_id).first()
            compensation = db.query(models.CompensationConclusion).filter(
                models.CompensationConclusion.review_id == review.id
            ).first()
            
            original_input_display = ""
            if review.original_input:
                try:
                    import json
                    original_data = json.loads(review.original_input)
                    original_input_display = original_data.get("caller_input", review.original_input)
                except:
                    original_input_display = review.original_input
            
            ws2.append([
                box.box_code,
                box.batch_no,
                signoff.store_code if signoff else "",
                signoff.store_name if signoff else "",
                signoff.signoff_person if signoff else "",
                signoff.signoff_time.strftime("%Y-%m-%d %H:%M:%S") if signoff and signoff.signoff_time else "",
                signoff.temperature_arrival if signoff else "",
                "是" if (signoff and signoff.has_exception) else "否",
                signoff.exception_desc if signoff else "",
                review.reviewer,
                review.review_time.strftime("%Y-%m-%d %H:%M:%S") if review.review_time else "",
                original_input_display,
                review.review_result,
                review.review_comment,
                "是" if review.temperature_violation else "否",
                "是" if review.compensation_eligible else "否",
                compensation.compensation_amount if compensation else 0,
                compensation.compensation_reason if compensation else "",
                compensation.processor if compensation else "",
                compensation.approved_by if compensation else "",
                review.status
            ])
    
    ws3 = wb.create_sheet("审计日志")
    headers3 = ["箱号", "操作类型", "操作人", "旧状态", "新状态", "备注", "变更详情", "操作时间"]
    ws3.append(headers3)
    
    for box in boxes:
        for audit in box.audit_logs:
            change_details_display = ""
            if audit.change_details:
                import json
                if isinstance(audit.change_details, dict):
                    change_details_display = json.dumps(audit.change_details, ensure_ascii=False)
                else:
                    change_details_display = str(audit.change_details)
            
            ws3.append([
                box.box_code,
                audit.action_type,
                audit.operator,
                audit.old_status,
                audit.new_status,
                audit.comment,
                change_details_display,
                audit.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])
    
    for ws in [ws1, ws2, ws3]:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", fill_type="solid")
    
    os.makedirs("exports", exist_ok=True)
    filename = f"exports/cold_chain_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    
    return FileResponse(
        filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=os.path.basename(filename)
    )


@app.get("/api/health")
def health_check():
    return {"success": True, "message": "服务运行正常", "timestamp": datetime.now().isoformat()}
