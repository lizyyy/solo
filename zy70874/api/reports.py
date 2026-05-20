from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from database import get_db
from models import ReconciliationRecord, ReconciliationSummary, SessionStatus

router = APIRouter()


def generate_csv_report(records, summary):
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(["院线对账报告"])
    writer.writerow([f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    writer.writerow([f"批次号: {summary.reconciliation_batch if summary else 'N/A'}"])
    writer.writerow([])
    
    writer.writerow(["汇总信息"])
    if summary:
        writer.writerow(["总场次", summary.total_sessions])
        writer.writerow(["已匹配", summary.matched_sessions])
        writer.writerow(["有差异待复核", summary.disputed_sessions])
        writer.writerow(["已核准", summary.approved_sessions])
        writer.writerow(["已拒绝", summary.rejected_sessions])
        writer.writerow([])
        writer.writerow(["总预期补贴", summary.total_expected_subsidy])
        writer.writerow(["总实际补贴", summary.total_actual_subsidy])
        writer.writerow(["总补贴差异", summary.total_subsidy_discrepancy])
        writer.writerow(["总最低票房缺口", summary.total_min_boxoffice_discrepancy])
        writer.writerow(["总退票扣减", summary.total_refund_deduction])
        writer.writerow(["总差异金额", summary.grand_total_discrepancy])
    writer.writerow([])
    
    writer.writerow([
        "场次编号", "影片名称", "影厅", "放映时间",
        "状态", "预期补贴", "实际补贴", "补贴差异",
        "最低票房承诺", "实际票房", "票房缺口",
        "退票扣减", "总差异", "差异类型", "差异说明", "复核备注"
    ])
    
    for record in records:
        session = record.session
        film_name = session.film_name if session else "N/A"
        hall_name = session.hall_name if session else "N/A"
        show_time = session.show_time.strftime('%Y-%m-%d %H:%M') if session else "N/A"
        session_code = session.session_code if session else "N/A"
        
        writer.writerow([
            session_code, film_name, hall_name, show_time,
            record.status.value,
            record.expected_subsidy, record.actual_subsidy, record.subsidy_discrepancy,
            record.expected_min_boxoffice, record.actual_boxoffice, record.min_boxoffice_discrepancy,
            record.refund_deduction, record.total_discrepancy,
            record.discrepancy_types or "",
            record.discrepancy_explanation or "",
            record.reviewer_notes or ""
        ])
    
    return output.getvalue()


def generate_excel_report(records, summary):
    output = io.BytesIO()
    wb = Workbook()
    
    ws_summary = wb.active
    ws_summary.title = "汇总"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    ws_summary["A1"] = "院线对账报告"
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary["A2"] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_summary["A3"] = f"批次号: {summary.reconciliation_batch if summary else 'N/A'}"
    
    if summary:
        summary_data = [
            ["统计项", "数值"],
            ["总场次", summary.total_sessions],
            ["已匹配", summary.matched_sessions],
            ["有差异待复核", summary.disputed_sessions],
            ["已核准", summary.approved_sessions],
            ["已拒绝", summary.rejected_sessions],
            [],
            ["财务统计项", "金额(元)"],
            ["总预期补贴", summary.total_expected_subsidy],
            ["总实际补贴", summary.total_actual_subsidy],
            ["总补贴差异", summary.total_subsidy_discrepancy],
            ["总最低票房缺口", summary.total_min_boxoffice_discrepancy],
            ["总退票扣减", summary.total_refund_deduction],
            ["总差异金额", summary.grand_total_discrepancy]
        ]
        
        for i, row in enumerate(summary_data, start=5):
            for j, val in enumerate(row):
                cell = ws_summary.cell(row=i, column=j+1, value=val)
                if i == 5 or i == 12:
                    cell.fill = header_fill
                    cell.font = header_font
    
    for col in ws_summary.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_summary.column_dimensions[column].width = adjusted_width
    
    ws_detail = wb.create_sheet(title="明细")
    headers = [
        "场次编号", "影片名称", "影厅", "放映时间",
        "状态", "预期补贴", "实际补贴", "补贴差异",
        "最低票房承诺", "实际票房", "票房缺口",
        "退票扣减", "总差异", "差异类型", "差异说明", "复核备注"
    ]
    
    for i, header in enumerate(headers, 1):
        cell = ws_detail.cell(row=1, column=i, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    for row_idx, record in enumerate(records, 2):
        session = record.session
        film_name = session.film_name if session else "N/A"
        hall_name = session.hall_name if session else "N/A"
        show_time = session.show_time.strftime('%Y-%m-%d %H:%M') if session else "N/A"
        session_code = session.session_code if session else "N/A"
        
        row_data = [
            session_code, film_name, hall_name, show_time,
            record.status.value,
            record.expected_subsidy, record.actual_subsidy, record.subsidy_discrepancy,
            record.expected_min_boxoffice, record.actual_boxoffice, record.min_boxoffice_discrepancy,
            record.refund_deduction, record.total_discrepancy,
            record.discrepancy_types or "",
            record.discrepancy_explanation or "",
            record.reviewer_notes or ""
        ]
        
        fill = None
        if record.status == SessionStatus.MATCHED:
            fill = green_fill
        elif record.status == SessionStatus.DISPUTED:
            fill = yellow_fill
        elif record.status == SessionStatus.APPROVED:
            fill = green_fill
        elif record.status == SessionStatus.REJECTED:
            fill = red_fill
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_detail.cell(row=row_idx, column=col_idx, value=val)
            if fill:
                cell.fill = fill
    
    for col in ws_detail.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws_detail.column_dimensions[column].width = adjusted_width
    
    wb.save(output)
    return output.getvalue()


@router.get("/export/{batch_id}")
def export_report(
    batch_id: str,
    format: str = "xlsx",
    db: Session = Depends(get_db)
):
    records = db.query(ReconciliationRecord).filter(
        ReconciliationRecord.reconciliation_batch == batch_id
    ).all()
    
    if not records:
        raise HTTPException(status_code=404, detail="批次不存在")
    
    summary = db.query(ReconciliationSummary).filter(
        ReconciliationSummary.reconciliation_batch == batch_id
    ).first()
    
    if format == "csv":
        content = generate_csv_report(records, summary)
        return StreamingResponse(
            io.StringIO(content),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=reconciliation_{batch_id}.csv"}
        )
    elif format == "xlsx":
        content = generate_excel_report(records, summary)
        return StreamingResponse(
            io.BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=reconciliation_{batch_id}.xlsx"}
        )
    else:
        raise HTTPException(status_code=400, detail="不支持的格式，请使用 csv 或 xlsx")
