import csv
import json
from io import StringIO, BytesIO
from datetime import datetime
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import (
    ReconciliationResult, ReconciliationBatch, Waybill,
    PenaltyHistory, ReviewRecord
)


class ReportService:
    def __init__(self, db: Session):
        self.db = db

    def generate_detailed_report(self, batch_id: str, format: str = 'xlsx') -> bytes:
        results = self.db.query(ReconciliationResult).filter(
            ReconciliationResult.batch_id == batch_id
        ).all()

        if format == 'xlsx':
            return self._generate_detailed_excel(results)
        elif format == 'csv':
            return self._generate_detailed_csv(results)
        else:
            return self._generate_detailed_json(results)

    def _generate_detailed_excel(self, results: List[ReconciliationResult]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "对账明细"

        headers = [
            '运单号', '状态', '审核状态', '是否晚点', '晚点时长(小时)', '晚点等级',
            '是否破损', '破损类型', '破损严重程度', '是否中转问题',
            '总扣罚金额', '晚点扣罚', '破损扣罚', '中转扣罚',
            '是否豁免', '豁免原因', '是否天气豁免', '是否重复扣罚',
            '差异说明', '审核人', '审核意见', '审核时间'
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row, result in enumerate(results, 2):
            waybill = result.waybill

            ws.cell(row=row, column=1, value=result.waybill_no)
            ws.cell(row=row, column=2, value=result.status)
            ws.cell(row=row, column=3, value=result.review_status)
            ws.cell(row=row, column=4, value='是' if result.is_delayed else '否')
            ws.cell(row=row, column=5, value=result.delay_hours)
            ws.cell(row=row, column=6, value=self._translate_level(result.delay_level))
            ws.cell(row=row, column=7, value='是' if result.is_damaged else '否')
            ws.cell(row=row, column=8, value=self._translate_damage_type(result.damage_type))
            ws.cell(row=row, column=9, value=self._translate_severity(result.damage_severity))
            ws.cell(row=row, column=10, value='是' if result.is_transfer_issue else '否')
            ws.cell(row=row, column=11, value=result.total_penalty)
            ws.cell(row=row, column=12, value=result.delay_penalty)
            ws.cell(row=row, column=13, value=result.damage_penalty)
            ws.cell(row=row, column=14, value=result.transfer_penalty)
            ws.cell(row=row, column=15, value='是' if result.is_exempt else '否')
            ws.cell(row=row, column=16, value=result.exempt_reason or '')
            ws.cell(row=row, column=17, value='是' if result.weather_exempt else '否')
            ws.cell(row=row, column=18, value='是' if result.is_duplicate_penalty else '否')
            ws.cell(row=row, column=19, value=result.discrepancy_explanation or '')
            ws.cell(row=row, column=20, value=result.reviewer or '')
            ws.cell(row=row, column=21, value=result.review_comment or '')
            ws.cell(row=row, column=22, value=str(result.reviewed_at) if result.reviewed_at else '')

            if result.is_exempt or result.weather_exempt:
                for col in range(1, 23):
                    ws.cell(row=row, column=col).fill = PatternFill(
                        start_color="92D050", end_color="92D050", fill_type="solid"
                    )
            elif result.total_penalty > 0:
                for col in range(1, 23):
                    ws.cell(row=row, column=col).fill = PatternFill(
                        start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
                    )

        for col in range(1, 23):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _generate_detailed_csv(self, results: List[ReconciliationResult]) -> bytes:
        output = StringIO()
        writer = csv.writer(output)

        writer.writerow([
            '运单号', '状态', '审核状态', '是否晚点', '晚点时长(小时)', '晚点等级',
            '是否破损', '破损类型', '破损严重程度', '是否中转问题',
            '总扣罚金额', '晚点扣罚', '破损扣罚', '中转扣罚',
            '是否豁免', '豁免原因', '是否天气豁免', '是否重复扣罚',
            '差异说明', '审核人', '审核意见', '审核时间'
        ])

        for result in results:
            writer.writerow([
                result.waybill_no, result.status, result.review_status,
                '是' if result.is_delayed else '否', result.delay_hours,
                self._translate_level(result.delay_level),
                '是' if result.is_damaged else '否',
                self._translate_damage_type(result.damage_type),
                self._translate_severity(result.damage_severity),
                '是' if result.is_transfer_issue else '否',
                result.total_penalty, result.delay_penalty, result.damage_penalty,
                result.transfer_penalty,
                '是' if result.is_exempt else '否', result.exempt_reason or '',
                '是' if result.weather_exempt else '否',
                '是' if result.is_duplicate_penalty else '否',
                result.discrepancy_explanation or '',
                result.reviewer or '', result.review_comment or '',
                str(result.reviewed_at) if result.reviewed_at else ''
            ])

        return output.getvalue().encode('utf-8-sig')

    def _generate_detailed_json(self, results: List[ReconciliationResult]) -> bytes:
        data = []
        for result in results:
            data.append({
                'waybill_no': result.waybill_no,
                'status': result.status,
                'review_status': result.review_status,
                'is_delayed': result.is_delayed,
                'delay_hours': result.delay_hours,
                'delay_level': result.delay_level,
                'is_damaged': result.is_damaged,
                'damage_type': result.damage_type,
                'damage_severity': result.damage_severity,
                'is_transfer_issue': result.is_transfer_issue,
                'total_penalty': result.total_penalty,
                'delay_penalty': result.delay_penalty,
                'damage_penalty': result.damage_penalty,
                'transfer_penalty': result.transfer_penalty,
                'is_exempt': result.is_exempt,
                'exempt_reason': result.exempt_reason,
                'weather_exempt': result.weather_exempt,
                'is_duplicate_penalty': result.is_duplicate_penalty,
                'discrepancy_explanation': result.discrepancy_explanation,
                'reviewer': result.reviewer,
                'review_comment': result.review_comment,
                'reviewed_at': str(result.reviewed_at) if result.reviewed_at else None,
                'penalty_details': result.penalty_details,
                'auto_calculation_details': result.auto_calculation_details,
                'manual_adjustment': result.manual_adjustment
            })
        return json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8')

    def generate_summary_report(self, batch_id: str) -> bytes:
        batch = self.db.query(ReconciliationBatch).filter(
            ReconciliationBatch.batch_id == batch_id
        ).first()

        results = self.db.query(ReconciliationResult).filter(
            ReconciliationResult.batch_id == batch_id
        ).all()

        wb = openpyxl.Workbook()

        ws1 = wb.active
        ws1.title = "汇总概览"
        self._fill_summary_overview(ws1, batch, results)

        ws2 = wb.create_sheet("扣罚统计")
        self._fill_penalty_statistics(ws2, results)

        ws3 = wb.create_sheet("异常分析")
        self._fill_exception_analysis(ws3, results)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _fill_summary_overview(self, ws, batch: ReconciliationBatch, results: List[ReconciliationResult]):
        headers = ['指标', '数值']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

        summary_data = [
            ['批次号', batch.batch_id],
            ['批次名称', batch.batch_name],
            ['生成时间', str(batch.generated_at)],
            ['运单总数', batch.total_waybills],
            ['已审核', batch.reconciled_count],
            ['待审核', batch.pending_count],
            ['已豁免', batch.exempt_count],
            ['扣罚总金额', batch.total_penalty],
        ]

        for row, (key, value) in enumerate(summary_data, 2):
            ws.cell(row=row, column=1, value=key).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 40

    def _fill_penalty_statistics(self, ws, results: List[ReconciliationResult]):
        headers = ['扣罚类型', '涉及运单数', '总金额', '平均金额', '最小金额', '最大金额']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

        delay_penalties = [r.delay_penalty for r in results if r.delay_penalty > 0]
        damage_penalties = [r.damage_penalty for r in results if r.damage_penalty > 0]
        transfer_penalties = [r.transfer_penalty for r in results if r.transfer_penalty > 0]
        all_penalties = [r.total_penalty for r in results if r.total_penalty > 0]

        stats_data = [
            ['晚点扣罚', len(delay_penalties), sum(delay_penalties),
             sum(delay_penalties) / len(delay_penalties) if delay_penalties else 0,
             min(delay_penalties) if delay_penalties else 0,
             max(delay_penalties) if delay_penalties else 0],
            ['破损扣罚', len(damage_penalties), sum(damage_penalties),
             sum(damage_penalties) / len(damage_penalties) if damage_penalties else 0,
             min(damage_penalties) if damage_penalties else 0,
             max(damage_penalties) if damage_penalties else 0],
            ['中转扣罚', len(transfer_penalties), sum(transfer_penalties),
             sum(transfer_penalties) / len(transfer_penalties) if transfer_penalties else 0,
             min(transfer_penalties) if transfer_penalties else 0,
             max(transfer_penalties) if transfer_penalties else 0],
            ['合计', len(all_penalties), sum(all_penalties),
             sum(all_penalties) / len(all_penalties) if all_penalties else 0,
             min(all_penalties) if all_penalties else 0,
             max(all_penalties) if all_penalties else 0],
        ]

        for row, data in enumerate(stats_data, 2):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=round(value, 2) if isinstance(value, float) else value)

        for col in range(1, 7):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    def _fill_exception_analysis(self, ws, results: List[ReconciliationResult]):
        headers = ['异常类型', '数量', '占比', '说明']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

        total = len(results)
        exception_data = [
            ['晚点异常', sum(1 for r in results if r.is_delayed),
             f"{sum(1 for r in results if r.is_delayed) / total * 100:.1f}%" if total else '0%',
             '实际到达时间晚于计划到达时间'],
            ['破损异常', sum(1 for r in results if r.is_damaged),
             f"{sum(1 for r in results if r.is_damaged) / total * 100:.1f}%" if total else '0%',
             '货物存在破损记录'],
            ['中转异常', sum(1 for r in results if r.is_transfer_issue),
             f"{sum(1 for r in results if r.is_transfer_issue) / total * 100:.1f}%" if total else '0%',
             '中转过程存在超时或其他问题'],
            ['天气豁免', sum(1 for r in results if r.weather_exempt),
             f"{sum(1 for r in results if r.weather_exempt) / total * 100:.1f}%" if total else '0%',
             '因天气原因豁免扣罚'],
            ['重复扣罚', sum(1 for r in results if r.is_duplicate_penalty),
             f"{sum(1 for r in results if r.is_duplicate_penalty) / total * 100:.1f}%" if total else '0%',
             '该运单已在其他批次扣罚'],
            ['人工豁免', sum(1 for r in results if r.is_exempt and not r.weather_exempt),
             f"{sum(1 for r in results if r.is_exempt and not r.weather_exempt) / total * 100:.1f}%" if total else '0%',
             '经审核人工豁免扣罚'],
        ]

        for row, data in enumerate(exception_data, 2):
            for col, value in enumerate(data, 1):
                ws.cell(row=row, column=col, value=value)

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 40

    def generate_discrepancy_report(self, waybill_no: str) -> bytes:
        result = self.db.query(ReconciliationResult).filter(
            ReconciliationResult.waybill_no == waybill_no
        ).first()

        if not result:
            raise ValueError(f"未找到运单 {waybill_no} 的对账结果")

        penalty_histories = self.db.query(PenaltyHistory).filter(
            PenaltyHistory.reconciliation_id == result.id
        ).all()

        review_records = self.db.query(ReviewRecord).filter(
            ReviewRecord.waybill_no == waybill_no
        ).order_by(ReviewRecord.created_at).all()

        wb = openpyxl.Workbook()

        ws1 = wb.active
        ws1.title = "差异说明"
        self._fill_discrepancy_explanation(ws1, result, penalty_histories, review_records)

        ws2 = wb.create_sheet("扣罚追溯")
        self._fill_penalty_traceability(ws2, penalty_histories)

        ws3 = wb.create_sheet("审核记录")
        self._fill_review_records(ws3, review_records)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _fill_discrepancy_explanation(self, ws, result: ReconciliationResult,
                                       penalty_histories: List[PenaltyHistory],
                                       review_records: List[ReviewRecord]):
        ws.merge_cells('A1:B1')
        title_cell = ws.cell(row=1, column=1, value=f"运单 {result.waybill_no} 差异说明报告")
        title_cell.font = Font(bold=True, size=16)
        title_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=3, column=1, value="基本信息").font = Font(bold=True, size=14)
        basic_info = [
            ['运单号', result.waybill_no],
            ['状态', result.status],
            ['审核状态', result.review_status],
            ['总扣罚金额', result.total_penalty],
        ]
        for row, (key, value) in enumerate(basic_info, 4):
            ws.cell(row=row, column=1, value=key).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)

        ws.cell(row=9, column=1, value="差异详情").font = Font(bold=True, size=14)

        discrepancies = []
        if result.is_delayed:
            discrepancies.append([
                '晚点',
                f"晚点 {result.delay_hours} 小时，等级: {self._translate_level(result.delay_level)}",
                f"晚点扣罚: {result.delay_penalty} 元",
                result.weather_exempt
            ])
        if result.is_damaged:
            discrepancies.append([
                '破损',
                f"类型: {self._translate_damage_type(result.damage_type)}, "
                f"严重程度: {self._translate_severity(result.damage_severity)}",
                f"破损扣罚: {result.damage_penalty} 元",
                False
            ])
        if result.is_transfer_issue:
            discrepancies.append([
                '中转问题',
                result.transfer_responsibility or '',
                f"中转扣罚: {result.transfer_penalty} 元",
                False
            ])
        if result.weather_exempt:
            discrepancies.append([
                '天气豁免',
                result.weather_info.get('description', '') if result.weather_info else '',
                '扣罚已豁免',
                True
            ])
        if result.is_duplicate_penalty:
            discrepancies.append([
                '重复扣罚',
                result.duplicate_source or '',
                '请注意核查',
                False
            ])
        if result.is_exempt:
            discrepancies.append([
                '人工豁免',
                result.exempt_reason or '',
                '扣罚已豁免',
                True
            ])

        if discrepancies:
            headers = ['异常类型', '详细描述', '扣罚/处理', '已豁免']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=10, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF")

            for row, disc in enumerate(discrepancies, 11):
                for col, value in enumerate(disc, 1):
                    cell = ws.cell(row=row, column=col, value=value)
                    if disc[3]:
                        cell.fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 10

    def _fill_penalty_traceability(self, ws, penalty_histories: List[PenaltyHistory]):
        ws.cell(row=1, column=1, value="扣罚追溯明细").font = Font(bold=True, size=14)

        headers = ['扣罚类型', '规则编码', '规则名称', '计算依据', '原始值', '扣罚金额', '来源批次', '计算路径']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")

        for row, history in enumerate(penalty_histories, 3):
            ws.cell(row=row, column=1, value=self._translate_penalty_type(history.penalty_type))
            ws.cell(row=row, column=2, value=history.rule_code)
            ws.cell(row=row, column=3, value=history.rule_name)
            ws.cell(row=row, column=4, value=str(history.calculation_basis))
            ws.cell(row=row, column=5, value=history.original_value)
            ws.cell(row=row, column=6, value=history.penalty_amount)
            ws.cell(row=row, column=7, value=history.source_batch or '')
            ws.cell(row=row, column=8, value=str(history.traceability_path or ''))

        for col in range(1, 9):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    def _fill_review_records(self, ws, review_records: List[ReviewRecord]):
        ws.cell(row=1, column=1, value="审核记录").font = Font(bold=True, size=14)

        headers = ['时间', '审核人', '操作', '原状态', '新状态', '原扣罚', '新扣罚', '调整原因', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            cell.font = Font(bold=True)

        for row, record in enumerate(review_records, 3):
            ws.cell(row=row, column=1, value=str(record.created_at))
            ws.cell(row=row, column=2, value=record.reviewer)
            ws.cell(row=row, column=3, value=record.action)
            ws.cell(row=row, column=4, value=record.old_status)
            ws.cell(row=row, column=5, value=record.new_status)
            ws.cell(row=row, column=6, value=record.old_total_penalty)
            ws.cell(row=row, column=7, value=record.new_total_penalty)
            ws.cell(row=row, column=8, value=record.adjustment_reason or '')
            ws.cell(row=row, column=9, value=record.comment or '')

        for col in range(1, 10):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    def _translate_level(self, level: str) -> str:
        translations = {
            'minor': '轻微',
            'moderate': '一般',
            'serious': '严重',
            'critical': '重大',
            'normal': '正常'
        }
        return translations.get(level, level or '')

    def _translate_damage_type(self, damage_type: str) -> str:
        translations = {
            'water_damage': '湿损',
            'physical_damage': '破损',
            'lost': '丢失',
            'other': '其他'
        }
        return translations.get(damage_type, damage_type or '')

    def _translate_severity(self, severity: str) -> str:
        translations = {
            'minor': '轻微',
            'moderate': '一般',
            'serious': '严重'
        }
        return translations.get(severity, severity or '')

    def _translate_penalty_type(self, penalty_type: str) -> str:
        translations = {
            'delay': '晚点扣罚',
            'damage': '破损扣罚',
            'transfer': '中转扣罚'
        }
        return translations.get(penalty_type, penalty_type or '')
