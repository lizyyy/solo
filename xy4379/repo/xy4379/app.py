import os
from datetime import datetime, timedelta
from flask import Flask, render_template, request, jsonify, send_file
from flask_cors import CORS
from models import db, Case, Document, Deadline, Todo, PendingReview
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///legal_assistant.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024  # 100MB

CORS(app)
db.init_app(app)

with app.app_context():
    db.create_all()

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt', 'jpg', 'jpeg', 'png', 'gif', 'bmp'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def get_file_type(filename):
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    type_map = {
        'pdf': 'pdf',
        'docx': 'word',
        'txt': 'text',
        'jpg': 'image',
        'jpeg': 'image',
        'png': 'image',
        'gif': 'image',
        'bmp': 'image'
    }
    return type_map.get(ext, 'other')

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/cases', methods=['GET'])
def get_cases():
    cases = Case.query.order_by(Case.updated_at.desc()).all()
    return jsonify([case.to_dict() for case in cases])

@app.route('/api/cases/<int:case_id>', methods=['GET'])
def get_case(case_id):
    case = Case.query.get_or_404(case_id)
    return jsonify(case.to_dict())

@app.route('/api/cases', methods=['POST'])
def create_case():
    data = request.json
    existing_case = Case.query.filter_by(case_number=data.get('case_number')).first()
    if existing_case:
        return jsonify({'error': '案件号已存在'}), 400
    
    case = Case(
        case_number=data.get('case_number'),
        client_name=data.get('client_name'),
        case_type=data.get('case_type'),
        status=data.get('status', 'active')
    )
    db.session.add(case)
    db.session.commit()
    return jsonify(case.to_dict()), 201

@app.route('/api/upload', methods=['POST'])
def upload_files():
    if 'files[]' not in request.files:
        return jsonify({'error': '没有文件上传'}), 400
    
    files = request.files.getlist('files[]')
    if not files or files[0].filename == '':
        return jsonify({'error': '没有选择文件'}), 400
    
    uploaded_files = []
    errors = []
    
    for file in files:
        if file and allowed_file(file.filename):
            try:
                filename = secure_filename(file.filename)
                timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
                unique_filename = f"{timestamp}_{filename}"
                
                case_number = request.form.get('case_number', '')
                client_name = request.form.get('client_name', '')
                
                if not case_number or not client_name:
                    extracted = extract_case_info_from_filename(filename)
                    if extracted.get('case_number'):
                        case_number = extracted['case_number']
                    if extracted.get('client_name'):
                        client_name = extracted['client_name']
                
                if not case_number or not client_name:
                    pending = PendingReview(
                        issue_type='missing_info',
                        description=f'文件 {filename} 无法识别案件号或当事人，请手动补充',
                        status='pending'
                    )
                    db.session.add(pending)
                    db.session.commit()
                    errors.append({'filename': filename, 'error': '无法识别案件信息，已加入待核对清单'})
                    continue
                
                case = Case.query.filter_by(case_number=case_number).first()
                if not case:
                    case = Case(
                        case_number=case_number,
                        client_name=client_name,
                        status='active'
                    )
                    db.session.add(case)
                    db.session.commit()
                
                existing_doc = Document.query.filter_by(
                    case_id=case.id,
                    file_name=filename
                ).first()
                
                if existing_doc:
                    pending = PendingReview(
                        document_id=existing_doc.id,
                        case_id=case.id,
                        issue_type='duplicate',
                        description=f'文件 {filename} 已存在于案件 {case_number} 中，是否需要合并？',
                        status='pending'
                    )
                    db.session.add(pending)
                    db.session.commit()
                    errors.append({'filename': filename, 'error': '文件已存在，已加入待核对清单'})
                    continue
                
                upload_dir = os.path.join(app.config['UPLOAD_FOLDER'], case_number)
                os.makedirs(upload_dir, exist_ok=True)
                
                file_path = os.path.join(upload_dir, unique_filename)
                file.save(file_path)
                
                file_size = os.path.getsize(file_path)
                file_type = get_file_type(filename)
                
                extracted_text = ''
                try:
                    if file_type == 'pdf':
                        extracted_text = extract_text_from_pdf(file_path)
                    elif file_type == 'word':
                        extracted_text = extract_text_from_docx(file_path)
                    elif file_type == 'text':
                        with open(file_path, 'r', encoding='utf-8') as f:
                            extracted_text = f.read()
                except Exception as e:
                    print(f"文本提取失败: {e}")
                
                document = Document(
                    case_id=case.id,
                    file_name=filename,
                    file_path=file_path,
                    file_type=file_type,
                    file_size=file_size,
                    extracted_text=extracted_text,
                    status='processed'
                )
                db.session.add(document)
                
                if extracted_text or filename:
                    dates = extract_dates(extracted_text + ' ' + filename)
                    for date_info in dates:
                        deadline = Deadline(
                            case_id=case.id,
                            deadline_type=date_info.get('type', '未知'),
                            deadline_date=date_info['date'],
                            description=date_info.get('description', ''),
                            source='自动提取',
                            confidence=date_info.get('confidence', 0.5),
                            is_confirmed=False
                        )
                        db.session.add(deadline)
                
                db.session.commit()
                uploaded_files.append({
                    'filename': filename,
                    'case_number': case_number,
                    'client_name': client_name
                })
                
            except Exception as e:
                db.session.rollback()
                errors.append({'filename': file.filename, 'error': str(e)})
        else:
            errors.append({'filename': file.filename, 'error': '不支持的文件类型'})
    
    return jsonify({
        'uploaded': uploaded_files,
        'errors': errors
    })

def extract_case_info_from_filename(filename):
    import re
    result = {'case_number': '', 'client_name': ''}
    
    case_patterns = [
        r'(\d{4})[民刑行商劳知执]初字第(\d+)号',
        r'(\d{4})[民刑行商劳知执]终字第(\d+)号',
        r'(\d{4})[民刑行商劳知执]再字第(\d+)号',
        r'案号[：:]\s*(\S+)',
        r'案件号[：:]\s*(\S+)',
    ]
    
    for pattern in case_patterns:
        match = re.search(pattern, filename)
        if match:
            if len(match.groups()) == 2:
                result['case_number'] = f"{match.group(1)}{match.group(2)}"
            else:
                result['case_number'] = match.group(1)
            break
    
    name_patterns = [
        r'当事人[：:]\s*(\S+)',
        r'原告[：:]\s*(\S+)',
        r'被告[：:]\s*(\S+)',
        r'申请人[：:]\s*(\S+)',
        r'被申请人[：:]\s*(\S+)',
    ]
    
    for pattern in name_patterns:
        match = re.search(pattern, filename)
        if match:
            result['client_name'] = match.group(1)
            break
    
    return result

def extract_text_from_pdf(file_path):
    try:
        import PyPDF2
        text = ''
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ''
        return text
    except Exception as e:
        return ''

def extract_text_from_docx(file_path):
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        text = ''
        for para in doc.paragraphs:
            text += para.text + '\n'
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + '\n'
        return text
    except Exception as e:
        return ''

def extract_dates(text):
    import re
    from datetime import datetime
    dates = []
    
    date_patterns = [
        (r'(\d{4})年(\d{1,2})月(\d{1,2})日', '日期'),
        (r'(\d{4})-(\d{1,2})-(\d{1,2})', '日期'),
        (r'(\d{4})/(\d{1,2})/(\d{1,2})', '日期'),
    ]
    
    deadline_keywords = {
        '举证': '举证期限',
        '举证期限': '举证期限',
        '补材料': '补充材料',
        '补充材料': '补充材料',
        '回访': '回访',
        '开庭': '开庭',
        '答辩': '答辩期限',
        '上诉': '上诉期限',
        '执行': '执行',
    }
    
    for pattern, date_type in date_patterns:
        matches = re.finditer(pattern, text)
        for match in matches:
            try:
                year = int(match.group(1))
                month = int(match.group(2))
                day = int(match.group(3))
                
                if 1900 <= year <= 2100 and 1 <= month <= 12 and 1 <= day <= 31:
                    date_obj = datetime(year, month, day)
                    
                    start = max(0, match.start() - 50)
                    end = min(len(text), match.end() + 50)
                    context = text[start:end]
                    
                    identified_type = date_type
                    confidence = 0.5
                    
                    for keyword, dtype in deadline_keywords.items():
                        if keyword in context:
                            identified_type = dtype
                            confidence = 0.9
                            break
                    
                    dates.append({
                        'date': date_obj,
                        'type': identified_type,
                        'description': context,
                        'confidence': confidence
                    })
            except ValueError:
                continue
    
    return dates

@app.route('/api/cases/<int:case_id>/documents', methods=['GET'])
def get_case_documents(case_id):
    documents = Document.query.filter_by(case_id=case_id).all()
    return jsonify([doc.to_dict() for doc in documents])

@app.route('/api/cases/<int:case_id>/deadlines', methods=['GET'])
def get_case_deadlines(case_id):
    deadlines = Deadline.query.filter_by(case_id=case_id).order_by(Deadline.deadline_date).all()
    return jsonify([d.to_dict() for d in deadlines])

@app.route('/api/deadlines/<int:deadline_id>', methods=['PUT'])
def update_deadline(deadline_id):
    deadline = Deadline.query.get_or_404(deadline_id)
    data = request.json
    
    if 'deadline_type' in data:
        deadline.deadline_type = data['deadline_type']
    if 'deadline_date' in data:
        deadline.deadline_date = datetime.fromisoformat(data['deadline_date'].replace('Z', '+00:00'))
    if 'description' in data:
        deadline.description = data['description']
    if 'is_confirmed' in data:
        deadline.is_confirmed = data['is_confirmed']
    if 'is_completed' in data:
        deadline.is_completed = data['is_completed']
        if data['is_completed']:
            deadline.completed_at = datetime.utcnow()
    
    db.session.commit()
    return jsonify(deadline.to_dict())

@app.route('/api/deadlines', methods=['POST'])
def create_deadline():
    data = request.json
    deadline = Deadline(
        case_id=data['case_id'],
        deadline_type=data['deadline_type'],
        deadline_date=datetime.fromisoformat(data['deadline_date'].replace('Z', '+00:00')),
        description=data.get('description', ''),
        source='手动录入',
        confidence=1.0,
        is_confirmed=True
    )
    db.session.add(deadline)
    db.session.commit()
    return jsonify(deadline.to_dict()), 201

@app.route('/api/todos', methods=['GET'])
def get_todos():
    status = request.args.get('status')
    query = Todo.query
    
    if status:
        query = query.filter_by(status=status)
    
    todos = query.order_by(Todo.due_date).all()
    return jsonify([todo.to_dict() for todo in todos])

@app.route('/api/todos', methods=['POST'])
def create_todo():
    data = request.json
    todo = Todo(
        case_id=data['case_id'],
        title=data['title'],
        description=data.get('description', ''),
        priority=data.get('priority', 'medium'),
        due_date=datetime.fromisoformat(data['due_date'].replace('Z', '+00:00')) if data.get('due_date') else None
    )
    db.session.add(todo)
    db.session.commit()
    return jsonify(todo.to_dict()), 201

@app.route('/api/todos/<int:todo_id>', methods=['PUT'])
def update_todo(todo_id):
    todo = Todo.query.get_or_404(todo_id)
    data = request.json
    
    if 'status' in data:
        todo.status = data['status']
        if data['status'] == 'completed':
            todo.completed_at = datetime.utcnow()
    
    db.session.commit()
    return jsonify(todo.to_dict())

@app.route('/api/pending-reviews', methods=['GET'])
def get_pending_reviews():
    status = request.args.get('status', 'pending')
    reviews = PendingReview.query.filter_by(status=status).order_by(PendingReview.created_at.desc()).all()
    return jsonify([r.to_dict() for r in reviews])

@app.route('/api/pending-reviews/<int:review_id>', methods=['PUT'])
def resolve_pending_review(review_id):
    review = PendingReview.query.get_or_404(review_id)
    data = request.json
    
    review.status = data.get('status', 'resolved')
    if review.status == 'resolved':
        review.resolved_at = datetime.utcnow()
    
    db.session.commit()
    return jsonify(review.to_dict())

@app.route('/api/export/markdown', methods=['GET'])
def export_markdown():
    cases = Case.query.filter_by(status='active').all()
    today = datetime.now().date()
    
    markdown = f"# 法律援助案件到期提醒报告\n\n"
    markdown += f"生成日期: {today.strftime('%Y年%m月%d日')}\n\n"
    
    for case in cases:
        deadlines = Deadline.query.filter_by(
            case_id=case.id,
            is_completed=False
        ).order_by(Deadline.deadline_date).all()
        
        if not deadlines:
            continue
        
        markdown += f"## 案件号: {case.case_number}\n"
        markdown += f"**当事人**: {case.client_name}\n"
        if case.case_type:
            markdown += f"**案件类型**: {case.case_type}\n"
        markdown += "\n"
        
        markdown += "### 待处理截止日期:\n\n"
        markdown += "| 类型 | 截止日期 | 剩余天数 | 状态 |\n"
        markdown += "|------|----------|----------|------|\n"
        
        for deadline in deadlines:
            deadline_date = deadline.deadline_date.date()
            days_remaining = (deadline_date - today).days
            
            if days_remaining < 0:
                status = '已过期'
            elif days_remaining <= 3:
                status = '紧急'
            elif days_remaining <= 7:
                status = '即将到期'
            else:
                status = '正常'
            
            markdown += f"| {deadline.deadline_type} | {deadline_date.strftime('%Y年%m月%d日')} | {days_remaining}天 | {status} |\n"
        
        markdown += "\n"
    
    output_path = 'case_report.md'
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(markdown)
    
    return send_file(output_path, as_attachment=True, download_name=f'case_report_{today.strftime("%Y%m%d")}.md')

@app.route('/api/export/csv', methods=['GET'])
def export_csv():
    import csv
    from io import StringIO
    
    cases = Case.query.filter_by(status='active').all()
    today = datetime.now().date()
    
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(['案件号', '当事人', '截止日期类型', '截止日期', '剩余天数', '状态', '已确认', '已完成'])
    
    for case in cases:
        deadlines = Deadline.query.filter_by(case_id=case.id).all()
        for deadline in deadlines:
            deadline_date = deadline.deadline_date.date()
            days_remaining = (deadline_date - today).days
            
            if days_remaining < 0:
                status = '已过期'
            elif days_remaining <= 3:
                status = '紧急'
            elif days_remaining <= 7:
                status = '即将到期'
            else:
                status = '正常'
            
            writer.writerow([
                case.case_number,
                case.client_name,
                deadline.deadline_type,
                deadline_date.strftime('%Y-%m-%d'),
                days_remaining,
                status,
                '是' if deadline.is_confirmed else '否',
                '是' if deadline.is_completed else '否'
            ])
    
    output.seek(0)
    csv_content = output.getvalue()
    
    output_path = 'case_report.csv'
    with open(output_path, 'w', encoding='utf-8-sig') as f:
        f.write(csv_content)
    
    return send_file(output_path, as_attachment=True, download_name=f'case_report_{today.strftime("%Y%m%d")}.csv')

@app.route('/api/stats', methods=['GET'])
def get_stats():
    today = datetime.now().date()
    tomorrow = today + timedelta(days=1)
    next_week = today + timedelta(days=7)
    
    total_cases = Case.query.count()
    active_cases = Case.query.filter_by(status='active').count()
    
    pending_deadlines = Deadline.query.filter(
        Deadline.is_completed == False,
        Deadline.deadline_date >= datetime.combine(today, datetime.min.time())
    ).count()
    
    urgent_deadlines = Deadline.query.filter(
        Deadline.is_completed == False,
        Deadline.deadline_date >= datetime.combine(today, datetime.min.time()),
        Deadline.deadline_date <= datetime.combine(tomorrow, datetime.max.time())
    ).count()
    
    this_week_deadlines = Deadline.query.filter(
        Deadline.is_completed == False,
        Deadline.deadline_date >= datetime.combine(today, datetime.min.time()),
        Deadline.deadline_date <= datetime.combine(next_week, datetime.max.time())
    ).count()
    
    pending_reviews = PendingReview.query.filter_by(status='pending').count()
    
    return jsonify({
        'total_cases': total_cases,
        'active_cases': active_cases,
        'pending_deadlines': pending_deadlines,
        'urgent_deadlines': urgent_deadlines,
        'this_week_deadlines': this_week_deadlines,
        'pending_reviews': pending_reviews
    })

if __name__ == '__main__':
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    app.run(debug=True, port=5000)
