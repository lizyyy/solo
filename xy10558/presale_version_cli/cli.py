from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import click
from rich.console import Console
from rich.table import Table
from rich.tree import Tree

from .checker import Checker
from .config import Config
from .models import (
    CHECK_RULES,
    CheckStatus,
    Document,
    DocumentType,
    ReviewItem,
    ReviewStatus,
)
from .reporter import Reporter
from .storage import Storage

console = Console()


def get_current_user() -> str:
    return os.environ.get("USER") or os.environ.get("USERNAME") or "unknown"


@click.group()
@click.option("--project-dir", "-p", default=".", help="项目目录路径")
@click.pass_context
def main(ctx, project_dir: str):
    ctx.ensure_object(dict)
    ctx.obj["project_dir"] = Path(project_dir).resolve()
    ctx.obj["storage"] = Storage(ctx.obj["project_dir"])


@main.command()
@click.argument("project_name")
@click.option("--version", "-v", default="v1.0.0", help="初始版本号")
@click.pass_context
def init(ctx, project_name: str, version: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if Config.is_initialized(project_dir):
        console.print("[yellow]项目已初始化，正在更新...[/yellow]")
        state = storage.load_state()
        if state.project_name != project_name:
            old_name = state.project_name
            state.project_name = project_name
            storage.add_audit_log(
                state,
                "RENAME",
                get_current_user(),
                f"重命名项目: {old_name} -> {project_name}",
                before={"name": old_name},
                after={"name": project_name},
            )
        storage.save_state(state)
        console.print(f"[green]✓ 项目已更新: {project_name}[/green]")
        return
    
    state = storage.initialize_project(project_name)
    state.current_version = version
    
    storage.add_audit_log(
        state,
        "INIT",
        get_current_user(),
        f"初始化项目: {project_name}, 版本: {version}",
    )
    storage.save_state(state)
    
    tree = Tree("[bold green]✓ 项目初始化完成[/bold green]")
    tree.add(f"项目名称: {project_name}")
    tree.add(f"项目目录: {project_dir}")
    tree.add(f"初始版本: {version}")
    tree.add(f"配置目录: {Config.get_config_path(project_dir)}")
    console.print(tree)
    
    console.print("\n[dim]下一步: 使用 [bold]presale import[/bold] 导入文档[/dim]")


@main.command()
@click.option("--requirements", "-r", help="需求清单文件路径")
@click.option("--quotation", "-q", help="报价表文件路径")
@click.option("--attachment", "-a", multiple=True, help="技术附件文件路径")
@click.option("--review", "-rv", help="评审意见文件路径")
@click.option("--version", "-v", default=None, help="文档版本号")
@click.option("--category", "-c", help="附件分类")
@click.pass_context
def import_doc(
    ctx, requirements, quotation, attachment, review, version, category):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        console.print("[yellow]请先运行: presale init <项目名称>[/yellow]")
        raise click.Abort()
    
    state = storage.load_state()
    doc_version = version or state.current_version
    
    imported = []
    
    if requirements:
        doc = _import_file(
            storage, state, requirements, DocumentType.REQUIREMENTS, doc_version
        )
        imported.append(("需求清单", doc))
    
    if quotation:
        doc = _import_file(
            storage, state, quotation, DocumentType.QUOTATION, doc_version
        )
        imported.append(("报价表", doc))
    
    for att_path in attachment:
        metadata = {}
        if category:
            metadata["category"] = category
        doc = _import_file(
            storage, state, att_path, DocumentType.ATTACHMENT, doc_version, metadata
        )
        imported.append(("附件", doc))
    
    if review:
        doc = _import_file(
            storage, state, review, DocumentType.REVIEW, doc_version
        )
        _parse_review_items(storage, state, review)
        imported.append(("评审意见", doc))
    
    if not imported:
        console.print("[yellow]未指定任何导入文件[/yellow]")
        console.print("[dim]使用 --requirements/--quotation/--attachment/--review 指定文件[/dim]")
        return
    
    table = Table(title="已导入文档")
    table.add_column("类型", style="cyan")
    table.add_column("名称", style="green")
    table.add_column("版本", style="yellow")
    table.add_column("校验和", style="dim")
    
    for doc_type, doc in imported:
        table.add_row(
            doc_type,
            doc.name,
            doc.version,
            doc.checksum[:12] + "...",
        )
    
    console.print(table)
    storage.save_state(state)
    console.print(f"\n[green]✓ 共导入 {len(imported)} 份文档[/green]")


def _import_file(
    storage: Storage, state, file_path: str, doc_type: DocumentType, version: str, metadata=None):
    path = Path(file_path)
    if not path.exists():
        console.print(f"[red]错误: 文件不存在 {file_path}[/red]")
        raise click.Abort()
    
    checksum = storage.calculate_checksum(path)
    existing = next(
        (d for d in state.documents
         if d.file_path == str(path) and d.is_active),
        None,
    )
    
    if existing:
        if existing.checksum == checksum:
            console.print(f"[dim]文件未变化，跳过: {path.name}[/dim]")
            return existing
        
        old_version = existing.version
        old_checksum = existing.checksum
        existing.is_active = False
        
        storage.add_audit_log(
            state,
            "DEACTIVATE",
            get_current_user(),
            f"停用旧版文档: {existing.name}",
            before={
                "version": old_version,
                "checksum": old_checksum,
            },
            after={"active": False},
        )
    
    now = datetime.now()
    doc = Document(
        id=str(uuid4()),
        type=doc_type,
        name=path.name,
        version=version,
        file_path=str(path),
        checksum=checksum,
        created_at=now,
        updated_at=now,
        metadata=metadata or {},
        is_active=True,
    )
    
    state.documents.append(doc)
    
    storage.add_audit_log(
        state,
        "IMPORT",
        get_current_user(),
        f"导入文档: {doc.name}, 类型: {doc_type.value}, 版本: {version}",
        after={
            "name": doc.name,
            "type": doc_type.value,
            "version": version,
            "checksum": checksum,
        },
    )
    
    return doc


def _parse_review_items(storage: Storage, state, review_file: str):
    import re
    
    content = ""
    try:
        path = Path(review_file)
        if path.suffix == ".xlsx":
            try:
                from openpyxl import load_workbook
                wb = load_workbook(review_file)
                ws = wb.active
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(" | ".join(str(cell) for cell in row if cell is not None))
                content = "\n".join(rows)
            except Exception:
                content = path.read_text(encoding="utf-8")
        else:
            content = path.read_text(encoding="utf-8")
    except Exception as e:
        console.print(f"[yellow]警告: 无法解析评审意见内容: {e}[/yellow]")
        return
    
    lines = content.split("\n")
    patterns = [
        r"^\s*[\d]+[.、]\s*(.*)",
        r"^\s*[-*]\s*(.*)",
    ]
    
    count = 0
    for line in lines:
        line = line.strip()
        if not line or len(line) < 5:
            continue
        
        matched = False
        for pattern in patterns:
            match = re.match(pattern, line)
            if match:
                _add_review_item(storage, state, match.group(1).strip())
                count += 1
                matched = True
                break
        
        if not matched and len(line) > 20:
            _add_review_item(storage, state, line)
            count += 1
    
    if count > 0:
        console.print(f"[green]解析到 {count} 条评审意见[/green]")


def _add_review_item(storage: Storage, state, content: str):
    existing = next(
        (r for r in state.review_items if r.content == content),
        None,
    )
    
    if existing:
        return
    
    now = datetime.now()
    item = ReviewItem(
        id=str(uuid4()),
        document_id="",
        content=content,
        status=ReviewStatus.OPEN,
        created_at=now,
        updated_at=now,
    )
    state.review_items.append(item)


@main.command()
@click.option("--verbose", "-v", is_flag=True, help="显示详细检查详情")
@click.pass_context
def check(ctx, verbose: bool):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    checker = Checker(state)
    results = checker.run_all_checks()
    
    table = Table(title="版本一致性检查结果")
    table.add_column("规则", style="cyan")
    table.add_column("状态", style="bold")
    table.add_column("消息")
    table.add_column("阻断", style="red")
    
    pass_count = 0
    fail_count = 0
    warning_count = 0
    overridden_count = 0
    blocking = []
    
    for result in sorted(results, key=lambda x: x.rule_id):
        rule = next(r for r in CHECK_RULES if r.id == result.rule_id)
        
        if result.is_overridden:
            status_style = "magenta"
            status_text = "OVERRIDDEN"
            overridden_count += 1
        elif result.status == CheckStatus.PASS:
            status_style = "green"
            status_text = "PASS"
            pass_count += 1
        elif result.status == CheckStatus.WARNING:
            status_style = "yellow"
            status_text = "WARN"
            warning_count += 1
        else:
            status_style = "red"
            status_text = "FAIL"
            fail_count += 1
        
        is_blocking = result.is_blocking and not result.is_overridden
        if is_blocking and result.status == CheckStatus.FAIL:
            blocking.append(result.message)
        
        table.add_row(
            rule.name,
            f"[{status_style}]{status_text}[/{status_style}]",
            result.message,
            "是" if result.is_blocking else "否",
        )
    
    console.print(table)
    
    if verbose:
        console.print("\n[bold]详细信息:[/bold]")
        for result in results:
            if result.details:
                console.print(f"\n[dim]{result.rule_id}: {result.message}[/dim]")
                for key, value in result.details.items():
                    console.print(f"  {key}: {value}")
    
    overridden_str = f", [magenta]~ {overridden_count} 已覆盖[/magenta]" if overridden_count > 0 else ""
    console.print(f"\n总计: [green]✓ {pass_count} 通过[/green], "
                  f"[yellow]! {warning_count} 警告[/yellow], "
                  f"[red]✗ {fail_count} 失败[/red]"
                  f"{overridden_str}")
    
    if blocking:
        console.print("\n[bold red]存在阻断项，方案暂不可发送:[/bold red]")
        for msg in blocking:
            console.print(f"  - {msg}")
    else:
        console.print("\n[bold green]✓ 无阻断项，方案可发送[/bold green]")
    
    storage.add_audit_log(
        state,
        "CHECK",
        get_current_user(),
        f"执行检查: pass={pass_count}, fail={fail_count}, warning={warning_count}",
    )
    storage.save_state(state)


@main.command()
@click.option("--history", "-h", is_flag=True, help="显示历史记录")
@click.option("--audit", "-a", is_flag=True, help="显示审计日志")
@click.option("--document", "-d", help="显示指定文档详情")
@click.pass_context
def detail(ctx, history: bool, audit: bool, document: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    
    if history:
        _show_history(state)
    elif audit:
        _show_audit(state)
    elif document:
        _show_document_detail(state, document)
    else:
        _show_overview(state)


def _show_overview(state):
    reporter = Reporter(state)
    summary = reporter.get_summary()
    
    tree = Tree(f"[bold blue]项目概览: {summary['project_name']}[/bold blue]")
    
    tree.add(f"[cyan]当前版本: {summary['current_version']}[/cyan]")
    
    docs_node = tree.add(f"[green]文档 ({summary['documents']['active']}/{summary['documents']['total']})[/green]")
    docs_by_type = reporter.get_documents_by_type()
    for doc_type, docs in docs_by_type.items():
        active_count = sum(1 for d in docs if d.is_active)
        docs_node.add(f"{doc_type}: {active_count} 活动")
    
    tree.add(f"[yellow]评审意见: {summary['reviews']['closed']}/{summary['reviews']['total']} 已关闭[/yellow]")
    
    check_node = tree.add("[bold]检查结果[/bold]")
    check_node.add(f"通过: {summary['pass']}")
    check_node.add(f"警告: {summary['warning']}")
    check_node.add(f"失败: {summary['fail']}")
    if summary['overridden'] > 0:
        check_node.add(f"已覆盖: {summary['overridden']}")
    
    can_send = "[bold green]可发送[/bold green]" if summary['can_send'] else "[bold red]不可发送[/bold red]"
    tree.add(f"发送状态: {can_send}")
    
    console.print(tree)
    
    if not summary['can_send']:
        console.print("\n[red]阻断项:[/red]")
        for msg in summary['blocking_failures']:
            console.print(f"  - {msg}")


def _show_history(state):
    table = Table(title="版本历史")
    table.add_column("版本", style="cyan")
    table.add_column("时间", style="yellow")
    table.add_column("说明")
    
    for entry in reversed(state.version_history):
        table.add_row(
            entry.get("version", "?"),
            entry.get("timestamp", "?"),
            entry.get("description", ""),
        )
    
    console.print(table)


def _show_audit(state):
    table = Table(title="审计日志 (最近20条)")
    table.add_column("时间", style="dim")
    table.add_column("操作", style="cyan")
    table.add_column("操作者", style="yellow")
    table.add_column("描述")
    
    logs = sorted(state.audit_logs, key=lambda x: x.timestamp, reverse=True)[:20]
    for log in logs:
        table.add_row(
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            log.action,
            log.actor,
            log.description,
        )
    
    console.print(table)


def _show_document_detail(state, doc_name: str):
    docs = [d for d in state.documents if doc_name in d.name]
    if not docs:
        console.print(f"[red]未找到文档: {doc_name}[/red]")
        return
    
    for doc in docs:
        tree = Tree("[bold]文档详情[/bold]")
        tree.add(f"名称: {doc.name}")
        tree.add(f"类型: {doc.type.value}")
        tree.add(f"版本: {doc.version}")
        tree.add(f"路径: {doc.file_path}")
        tree.add(f"校验和: {doc.checksum}")
        tree.add(f"创建时间: {doc.created_at}")
        tree.add(f"状态: {'[green]活动[/green]' if doc.is_active else '[red]已停用[/red]'}")
        if doc.metadata:
            meta_node = tree.add("元数据")
            for key, value in doc.metadata.items():
                meta_node.add(f"{key}: {value}")
        
        console.print(tree)


@main.command()
@click.option("--output", "-o", default="report.md", help="输出文件路径")
@click.option("--format", "-f", "fmt", default="markdown", 
              type=click.Choice(["markdown", "json"]),
              help="输出格式")
@click.pass_context
def report(ctx, output: str, fmt: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    reporter = Reporter(state)
    summary = reporter.get_summary()
    
    if fmt == "json":
        import json
        report_data = {
            "summary": summary,
            "documents": reporter.get_documents_by_type(),
            "version_history": reporter.get_version_history(),
            "check_results": [
                {
                    "rule_id": r.rule_id,
                    "status": r.status,
                    "message": r.message,
                    "details": r.details,
                    "is_blocking": r.is_blocking,
                    "is_overridden": r.is_overridden,
                }
                for r in state.check_results
            ],
        }
        
        output_path = Path(output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(report_data, f, ensure_ascii=False, indent=2, default=str)
        
        console.print(f"[green]✓ 报告已生成: {output_path}[/green]")
        return
    
    _generate_markdown_report(state, reporter, summary, output)


def _generate_markdown_report(state, reporter, summary, output: str):
    lines = []
    
    lines.append("# 售前方案版本报告")
    lines.append("")
    lines.append(f"**项目名称**: {summary['project_name']}")
    lines.append(f"**当前版本**: {summary['current_version']}")
    lines.append(f"**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("")
    
    lines.append("## 发送状态")
    lines.append("")
    if summary['can_send']:
        lines.append("✅ **可以发送给客户**")
    else:
        lines.append("❌ **暂不可发送**")
        lines.append("")
        lines.append("### 阻断项:")
        for msg in summary['blocking_failures']:
            lines.append(f"- ❌ {msg}")
    lines.append("")
    
    lines.append("## 检查结果汇总")
    lines.append("")
    lines.append(f"- ✅ 通过: {summary['pass']}/{summary['total_checks']}")
    lines.append(f"- ⚠️ 警告: {summary['warning']}")
    lines.append(f"- ❌ 失败: {summary['fail']}")
    if summary['overridden'] > 0:
        lines.append(f"- 🔵 人工覆盖: {summary['overridden']}")
    lines.append("")
    
    lines.append("## 文档清单")
    lines.append("")
    
    docs_by_type = reporter.get_documents_by_type()
    
    type_names = {
        "requirements": "需求清单",
        "quotation": "报价表",
        "attachment": "技术附件",
        "review": "评审意见",
    }
    
    for doc_type, docs in docs_by_type.items():
        lines.append(f"### {type_names.get(doc_type, doc_type)}")
        lines.append("")
        lines.append("| 名称 | 版本 | 状态 |")
        lines.append("|------|------|------|")
        for doc in sorted(docs, key=lambda x: x.created_at, reverse=True):
            status = "✅" if doc.is_active else "⛔"
            lines.append(f"| {doc.name} | {doc.version} | {status} |")
        lines.append("")
    
    lines.append("## 版本历史")
    lines.append("")
    for entry in reversed(state.version_history):
        lines.append(f"- **{entry.get('version', '?')}**: {entry.get('description', '')}")
    lines.append("")
    
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    
    console.print(f"[green]✓ 报告已生成: {output_path}[/green]")
    console.print(f"[dim]使用 'cat {output_path}' 查看报告[/dim]")


@main.command("override")
@click.argument("rule_id")
@click.option("--reason", "-r", required=True, help="覆盖原因")
@click.pass_context
def override_check(ctx, rule_id: str, reason: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    
    result = next(
        (r for r in state.check_results if r.rule_id == rule_id),
        None,
    )
    
    if not result:
        console.print(f"[red]未找到检查规则: {rule_id}[/red]")
        console.print("[dim]可用规则:[/dim]")
        for rule in CHECK_RULES:
            console.print(f"  - {rule.id}")
        raise click.Abort()
    
    before = {
        "is_overridden": result.is_overridden,
        "status": result.status,
    }
    
    result.is_overridden = True
    result.overridden_by = get_current_user()
    result.overridden_at = datetime.now()
    result.override_reason = reason
    
    after = {
        "is_overridden": True,
        "overridden_by": result.overridden_by,
        "overridden_at": result.overridden_at.isoformat(),
        "reason": reason,
    }
    
    storage.add_audit_log(
        state,
        "OVERRIDE",
        get_current_user(),
        f"人工覆盖检查规则 {rule_id}: {reason}",
        before=before,
        after=after,
    )
    
    storage.save_state(state)
    
    console.print(f"[green]✓ 已覆盖检查规则: {rule_id}[/green]")
    console.print(f"[dim]原因: {reason}[/dim]")
    console.print(f"[dim]操作者: {result.overridden_by}[/dim]")


@main.command("close-review")
@click.argument("review_id_or_index")
@click.option("--assignee", "-a", help="处理人")
@click.pass_context
def close_review(ctx, review_id_or_index: str, assignee: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    
    try:
        index = int(review_id_or_index) - 1
        if 0 <= index < len(state.review_items):
            item = state.review_items[index]
        else:
            item = None
    except ValueError:
        item = next(
            (r for r in state.review_items if r.id == review_id_or_index),
            None,
        )
    
    if not item:
        console.print(f"[red]未找到评审意见: {review_id_or_index}[/red]")
        raise click.Abort()
    
    before = {"status": item.status.value}
    
    item.status = ReviewStatus.CLOSED
    item.closed_at = datetime.now()
    item.assignee = assignee or get_current_user()
    item.updated_at = datetime.now()
    
    after = {
        "status": "closed",
        "assignee": item.assignee,
        "closed_at": item.closed_at.isoformat(),
    }
    
    storage.add_audit_log(
        state,
        "CLOSE_REVIEW",
        get_current_user(),
        f"关闭评审意见: {item.content[:50]}...",
        before=before,
        after=after,
    )
    
    storage.save_state(state)
    
    console.print(f"[green]✓ 已关闭评审意见[/green]")
    console.print(f"[dim]{item.content}[/dim]")


@main.command("list-review")
@click.pass_context
def list_review(ctx):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    
    if not state.review_items:
        console.print("[yellow]暂无评审意见[/yellow]")
        return
    
    table = Table(title="评审意见列表")
    table.add_column("#", style="dim")
    table.add_column("状态", style="bold")
    table.add_column("内容")
    table.add_column("处理人")
    
    for i, item in enumerate(state.review_items, 1):
        if item.status == ReviewStatus.CLOSED:
            status_style = "green"
            status_text = "已关闭"
        elif item.status == ReviewStatus.IN_PROGRESS:
            status_style = "yellow"
            status_text = "处理中"
        else:
            status_style = "red"
            status_text = "待处理"
        
        content = item.content[:60] + "..." if len(item.content) > 60 else item.content
        table.add_row(
            str(i),
            f"[{status_style}]{status_text}[/{status_style}]",
            content,
            item.assignee or "-",
        )
    
    console.print(table)


@main.command("bump-version")
@click.option("--major", is_flag=True, help="升级主版本")
@click.option("--minor", is_flag=True, help="升级次版本")
@click.option("--patch", is_flag=True, help="升级修订版本")
@click.option("--description", "-d", help="版本描述")
@click.pass_context
def bump_version(ctx, major: bool, minor: bool, patch: bool, description: str):
    storage: Storage = ctx.obj["storage"]
    project_dir: Path = ctx.obj["project_dir"]
    
    if not Config.is_initialized(project_dir):
        console.print("[red]错误: 项目未初始化[/red]")
        raise click.Abort()
    
    state = storage.load_state()
    old_version = state.current_version
    
    try:
        parts = old_version.lstrip("v").split(".")
        major_v, minor_v, patch_v = map(int, parts)
    except Exception:
        major_v, minor_v, patch_v = 1, 0, 0
    
    if major:
        major_v += 1
        minor_v = 0
        patch_v = 0
    elif minor:
        minor_v += 1
        patch_v = 0
    elif patch:
        patch_v += 1
    else:
        patch_v += 1
    
    new_version = f"v{major_v}.{minor_v}.{patch_v}"
    
    before = {"version": old_version}
    
    state.current_version = new_version
    state.version_history.append(
        {
            "version": new_version,
            "timestamp": datetime.now().isoformat(),
            "description": description or "版本升级",
        }
    )
    
    after = {"version": new_version}
    
    storage.add_audit_log(
        state,
        "BUMP_VERSION",
        get_current_user(),
        f"版本升级: {old_version} -> {new_version}",
        before=before,
        after=after,
    )
    
    storage.save_state(state)
    
    console.print(f"[green]✓ 版本已升级: {old_version} -> {new_version}[/green]")


if __name__ == "__main__":
    main()
