from datetime import datetime, date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from models import (
    db, RepairOrder, Instrument, Performance, SparePartOrder,
    Technician, RepairException, Reminder
)
from models import RepairOrderStatus, ExceptionStatus, SparePartOrderStatus


class ReportService:
    @staticmethod
    def generate_schedule_report(start_date=None, end_date=None):
        if not start_date:
            start_date = date.today()
        if not end_date:
            from datetime import timedelta
            end_date = start_date + timedelta(days=30)

        repair_orders = RepairOrder.query.filter(
            RepairOrder.created_at >= start_date,
            RepairOrder.created_at <= end_date
        ).order_by(RepairOrder.created_at.desc()).all()

        instruments_in_repair = Instrument.query.filter(
            Instrument.status.in_(["维修中", "待备件"])
        ).all()

        upcoming_performances = Performance.query.filter(
            Performance.performance_date >= start_date
        ).order_by(Performance.performance_date.asc()).all()

        spare_orders_pending = SparePartOrder.query.filter(
            SparePartOrder.status.in_(["已下单", "运输中", "已延误"])
        ).all()

        active_exceptions = RepairException.query.filter(
            RepairException.status.in_([ExceptionStatus.DETECTED, ExceptionStatus.ACKNOWLEDGED])
        ).all()

        unread_reminders = Reminder.query.filter_by(is_read=False).all()

        status_summary = {}
        for status in RepairOrderStatus:
            count = RepairOrder.query.filter_by(status=status).count()
            if count > 0:
                status_summary[status.value] = count

        exception_summary = {}
        for exc in active_exceptions:
            exc_type = exc.exception_type.value
            exception_summary[exc_type] = exception_summary.get(exc_type, 0) + 1

        at_risk_instruments = []
        for perf in upcoming_performances[:5]:
            for pi in perf.required_instruments:
                inst = pi.instrument
                if inst and inst.status in ["维修中", "待备件"]:
                    active_order = RepairOrder.query.filter(
                        RepairOrder.instrument_id == inst.id,
                        RepairOrder.status.notin_(["已归还", "已取消"])
                    ).first()
                    at_risk_instruments.append({
                        "instrument_id": inst.id,
                        "instrument_name": inst.name,
                        "instrument_status": inst.status.value if hasattr(inst.status, 'value') else str(inst.status),
                        "performance_id": perf.id,
                        "performance_name": perf.name,
                        "performance_date": perf.performance_date,
                        "repair_order_id": active_order.id if active_order else None,
                        "return_deadline": active_order.return_deadline if active_order else None,
                        "instrument": inst,
                        "performance": perf,
                        "repair_order": active_order
                    })

        return {
            "report_period": {
                "start_date": start_date,
                "end_date": end_date,
                "generated_at": datetime.now()
            },
            "summary": {
                "total_repair_orders": len(repair_orders),
                "instruments_in_repair": len(instruments_in_repair),
                "upcoming_performances": len(upcoming_performances),
                "pending_spare_orders": len(spare_orders_pending),
                "active_exceptions": len(active_exceptions),
                "unread_reminders": len(unread_reminders),
                "status_distribution": status_summary,
                "exception_distribution": exception_summary
            },
            "repair_orders": repair_orders,
            "instruments_in_repair": instruments_in_repair,
            "upcoming_performances": upcoming_performances[:5],
            "pending_spare_orders": spare_orders_pending,
            "active_exceptions": active_exceptions,
            "at_risk_instruments": at_risk_instruments,
            "unread_reminders": unread_reminders[:10]
        }

    @staticmethod
    def generate_repair_order_detail_report(repair_order_id):
        order = RepairOrder.query.get(repair_order_id)
        if not order:
            return None

        status_history = list(order.status_histories)
        exceptions = list(order.exceptions)
        spare_usage = list(order.spare_part_usages)
        reminders = list(order.reminders)
        confirmations = []

        from models import ConfirmationRecord
        confirmations = ConfirmationRecord.query.filter_by(
            repair_order_id=repair_order_id
        ).order_by(ConfirmationRecord.confirmed_at.desc()).all()

        return {
            "repair_order": order,
            "status_history": status_history,
            "exceptions": exceptions,
            "spare_part_usages": spare_usage,
            "reminders": reminders,
            "confirmations": confirmations,
            "generated_at": datetime.now()
        }

    @staticmethod
    def export_schedule_report_to_excel(start_date=None, end_date=None):
        report = ReportService.generate_schedule_report(start_date, end_date)

        output = BytesIO()
        wb = Workbook()

        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        def style_header(ws, row, cols):
            for col in range(1, cols + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

        ws1 = wb.active
        ws1.title = "总览"
        ws1.merge_cells('A1:E1')
        ws1['A1'] = f"民乐团乐器维修排程报告"
        ws1['A1'].font = Font(bold=True, size=16)
        ws1['A1'].alignment = center_align

        ws1['A3'] = "报告期间"
        ws1['B3'] = f"{report['report_period']['start_date']} 至 {report['report_period']['end_date']}"
        ws1['A4'] = "生成时间"
        ws1['B4'] = report['report_period']['generated_at'].strftime('%Y-%m-%d %H:%M:%S')

        headers = ["指标", "数量"]
        for i, h in enumerate(headers, 1):
            ws1.cell(row=6, column=i, value=h)
        style_header(ws1, 6, 2)

        summary_items = [
            ("维修工单总数", report['summary']['total_repair_orders']),
            ("维修中乐器数", report['summary']['instruments_in_repair']),
            ("即将到来演出", report['summary']['upcoming_performances']),
            ("待到货备件订单", report['summary']['pending_spare_orders']),
            ("待处理业务例外", report['summary']['active_exceptions']),
            ("未读提醒", report['summary']['unread_reminders']),
        ]
        for i, (k, v) in enumerate(summary_items, 7):
            ws1.cell(row=i, column=1, value=k)
            ws1.cell(row=i, column=2, value=v)

        ws1['A15'] = "工单状态分布"
        ws1['A15'].font = Font(bold=True, size=12)
        for i, (k, v) in enumerate(report['summary']['status_distribution'].items(), 16):
            ws1.cell(row=i, column=1, value=k)
            ws1.cell(row=i, column=2, value=v)

        ws2 = wb.create_sheet("维修工单")
        headers2 = ["工单号", "乐器", "状态", "优先级", "维修师", "创建时间", "预计完成", "归还期限"]
        for i, h in enumerate(headers2, 1):
            ws2.cell(row=1, column=i, value=h)
        style_header(ws2, 1, len(headers2))

        for i, order in enumerate(report['repair_orders'], 2):
            ws2.cell(row=i, column=1, value=order.order_no)
            ws2.cell(row=i, column=2, value=order.instrument.name if order.instrument else "")
            ws2.cell(row=i, column=3, value=order.status.value)
            ws2.cell(row=i, column=4, value=order.priority)
            ws2.cell(row=i, column=5, value=order.technician.name if order.technician else "未分配")
            ws2.cell(row=i, column=6, value=order.created_at.strftime('%Y-%m-%d') if order.created_at else "")
            ws2.cell(row=i, column=7, value=order.scheduled_complete_date.strftime('%Y-%m-%d') if order.scheduled_complete_date else "")
            ws2.cell(row=i, column=8, value=order.return_deadline.strftime('%Y-%m-%d') if order.return_deadline else "")

        ws3 = wb.create_sheet("业务例外")
        headers3 = ["例外类型", "状态", "关联工单", "关联演出", "发现时间", "描述"]
        for i, h in enumerate(headers3, 1):
            ws3.cell(row=1, column=i, value=h)
        style_header(ws3, 1, len(headers3))

        for i, exc in enumerate(report['active_exceptions'], 2):
            ws3.cell(row=i, column=1, value=exc.exception_type.value)
            ws3.cell(row=i, column=2, value=exc.status.value)
            ws3.cell(row=i, column=3, value=f"工单#{exc.repair_order_id}" if exc.repair_order_id else "")
            ws3.cell(row=i, column=4, value=f"演出#{exc.performance_id}" if exc.performance_id else "")
            ws3.cell(row=i, column=5, value=exc.detected_at.strftime('%Y-%m-%d %H:%M') if exc.detected_at else "")
            ws3.cell(row=i, column=6, value=exc.description)

        ws4 = wb.create_sheet("即将演出")
        headers4 = ["演出名称", "日期", "地点", "城市", "所需乐器数", "存在风险乐器"]
        for i, h in enumerate(headers4, 1):
            ws4.cell(row=1, column=i, value=h)
        style_header(ws4, 1, len(headers4))

        for i, perf in enumerate(report['upcoming_performances'], 2):
            ws4.cell(row=i, column=1, value=perf.name)
            ws4.cell(row=i, column=2, value=perf.performance_date.strftime('%Y-%m-%d') if perf.performance_date else "")
            ws4.cell(row=i, column=3, value=perf.venue)
            ws4.cell(row=i, column=4, value=perf.city or "")
            ws4.cell(row=i, column=5, value=perf.required_instruments.count())
            risk_count = sum(
                1 for pi in perf.required_instruments
                if pi.instrument and pi.instrument.status in ["维修中", "待备件"]
            )
            ws4.cell(row=i, column=6, value=risk_count)

        for ws in [ws1, ws2, ws3, ws4]:
            for col in ws.columns:
                max_length = 0
                column = None
                for cell in col:
                    if not hasattr(cell, 'column_letter'):
                        continue
                    if column is None:
                        column = cell.column_letter
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                if column and max_length > 0:
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column].width = adjusted_width

        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_repair_order_to_excel(repair_order_id):
        report = ReportService.generate_repair_order_detail_report(repair_order_id)
        if not report:
            return None

        output = BytesIO()
        wb = Workbook()

        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        def style_header(ws, row, cols):
            for col in range(1, cols + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

        order = report['repair_order']

        ws1 = wb.active
        ws1.title = "工单详情"
        ws1.merge_cells('A1:D1')
        ws1['A1'] = f"维修工单详情 - {order.order_no}"
        ws1['A1'].font = Font(bold=True, size=16)
        ws1['A1'].alignment = center_align

        details = [
            ("工单号", order.order_no),
            ("乐器", order.instrument.name if order.instrument else ""),
            ("当前状态", order.status.value),
            ("优先级", order.priority),
            ("维修师", order.technician.name if order.technician else "未分配"),
            ("故障描述", order.description),
            ("创建时间", order.created_at.strftime('%Y-%m-%d %H:%M') if order.created_at else ""),
            ("计划开始", order.scheduled_start_date.strftime('%Y-%m-%d') if order.scheduled_start_date else ""),
            ("计划完成", order.scheduled_complete_date.strftime('%Y-%m-%d') if order.scheduled_complete_date else ""),
            ("实际开始", order.actual_start_date.strftime('%Y-%m-%d') if order.actual_start_date else ""),
            ("实际完成", order.actual_complete_date.strftime('%Y-%m-%d') if order.actual_complete_date else ""),
            ("归还期限", order.return_deadline.strftime('%Y-%m-%d') if order.return_deadline else ""),
            ("实际归还", order.actual_return_date.strftime('%Y-%m-%d') if order.actual_return_date else ""),
            ("预计工时", str(order.estimated_hours) if order.estimated_hours else ""),
            ("实际工时", str(order.actual_hours) if order.actual_hours else ""),
            ("维修备注", order.repair_note or ""),
            ("质检备注", order.quality_check_note or ""),
        ]

        for i, (k, v) in enumerate(details, 3):
            ws1.cell(row=i, column=1, value=k)
            ws1.cell(row=i, column=1).font = Font(bold=True)
            ws1.cell(row=i, column=2, value=v)

        ws2 = wb.create_sheet("状态历史")
        headers2 = ["序号", "变更时间", "原状态", "新状态", "变更原因", "操作人"]
        for i, h in enumerate(headers2, 1):
            ws2.cell(row=1, column=i, value=h)
        style_header(ws2, 1, len(headers2))

        for i, hist in enumerate(report['status_history'], 2):
            ws2.cell(row=i, column=1, value=i)
            ws2.cell(row=i, column=2, value=hist.created_at.strftime('%Y-%m-%d %H:%M:%S') if hist.created_at else "")
            ws2.cell(row=i, column=3, value=hist.from_status or "-")
            ws2.cell(row=i, column=4, value=hist.to_status)
            ws2.cell(row=i, column=5, value=hist.change_reason or "")
            ws2.cell(row=i, column=6, value=hist.operated_by or "")

        ws3 = wb.create_sheet("备件使用")
        headers3 = ["序号", "备件名称", "数量", "单位", "使用日期", "使用人", "备注"]
        for i, h in enumerate(headers3, 1):
            ws3.cell(row=1, column=i, value=h)
        style_header(ws3, 1, len(headers3))

        for i, usage in enumerate(report['spare_part_usages'], 2):
            ws3.cell(row=i, column=1, value=i)
            ws3.cell(row=i, column=2, value=usage.spare_part.name if usage.spare_part else "")
            ws3.cell(row=i, column=3, value=usage.quantity)
            ws3.cell(row=i, column=4, value=usage.spare_part.unit if usage.spare_part else "")
            ws3.cell(row=i, column=5, value=usage.usage_date.strftime('%Y-%m-%d') if usage.usage_date else "")
            ws3.cell(row=i, column=6, value=usage.used_by or "")
            ws3.cell(row=i, column=7, value=usage.remarks or "")

        ws4 = wb.create_sheet("业务例外")
        headers4 = ["序号", "类型", "状态", "发现时间", "描述", "是否已确认"]
        for i, h in enumerate(headers4, 1):
            ws4.cell(row=1, column=i, value=h)
        style_header(ws4, 1, len(headers4))

        for i, exc in enumerate(report['exceptions'], 2):
            ws4.cell(row=i, column=1, value=i)
            ws4.cell(row=i, column=2, value=exc.exception_type.value)
            ws4.cell(row=i, column=3, value=exc.status.value)
            ws4.cell(row=i, column=4, value=exc.detected_at.strftime('%Y-%m-%d %H:%M') if exc.detected_at else "")
            ws4.cell(row=i, column=5, value=exc.description)
            ws4.cell(row=i, column=6, value="是" if exc.acknowledged_by else "否")

        ws5 = wb.create_sheet("人工确认记录")
        headers5 = ["序号", "确认类型", "确认时间", "确认人", "确认备注", "保存快照"]
        for i, h in enumerate(headers5, 1):
            ws5.cell(row=1, column=i, value=h)
        style_header(ws5, 1, len(headers5))

        for i, conf in enumerate(report['confirmations'], 2):
            ws5.cell(row=i, column=1, value=i)
            ws5.cell(row=i, column=2, value=conf.confirmation_type.value)
            ws5.cell(row=i, column=3, value=conf.confirmed_at.strftime('%Y-%m-%d %H:%M:%S') if conf.confirmed_at else "")
            ws5.cell(row=i, column=4, value=conf.confirmed_by)
            ws5.cell(row=i, column=5, value=conf.confirmation_note or "")
            ws5.cell(row=i, column=6, value="是" if conf.before_snapshot and conf.after_snapshot else "否")

        for ws in [ws1, ws2, ws3, ws4, ws5]:
            for col in ws.columns:
                max_length = 0
                column = None
                for cell in col:
                    if not hasattr(cell, 'column_letter'):
                        continue
                    if column is None:
                        column = cell.column_letter
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                if column and max_length > 0:
                    adjusted_width = min(max_length + 2, 60)
                    ws.column_dimensions[column].width = adjusted_width

        wb.save(output)
        output.seek(0)
        return output
