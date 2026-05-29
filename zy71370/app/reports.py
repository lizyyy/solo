import os
import json
from datetime import datetime
from typing import Optional
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from . import models, schemas
from .exceptions import ReportGenerationException, StudentNotFoundException

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def generate_substitution_report(
    db: Session,
    student_id: Optional[int] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None
) -> schemas.ReportResponse:
    query = db.query(models.SubstitutionRecord)

    if student_id:
        student = db.query(models.Student).filter(models.Student.id == student_id).first()
        if not student:
            raise StudentNotFoundException(student_id)
        query = query.filter(models.SubstitutionRecord.student_id == student_id)

    if start_date:
        query = query.filter(models.SubstitutionRecord.created_at >= start_date)
    if end_date:
        query = query.filter(models.SubstitutionRecord.created_at <= end_date)

    records = query.order_by(models.SubstitutionRecord.created_at.desc()).all()

    filename = f"substitution_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "颜料替代记录"

        headers = ["ID", "原颜料ID", "原颜料名称", "替代颜料ID", "替代颜料名称",
                   "色差(ΔE)", "学生ID", "状态", "审核备注", "创建时间"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        for row, record in enumerate(records, 2):
            ws.cell(row=row, column=1, value=record.id)
            ws.cell(row=row, column=2, value=record.original_paint_id)
            ws.cell(row=row, column=3, value=record.original_paint_name)
            ws.cell(row=row, column=4, value=record.substitute_paint_id)
            ws.cell(row=row, column=5, value=record.substitute_paint_name)
            cell = ws.cell(row=row, column=6, value=record.color_difference)
            if record.color_difference > 5:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            ws.cell(row=row, column=7, value=record.student_id or "-")
            ws.cell(row=row, column=8, value=record.status)
            ws.cell(row=row, column=9, value=record.review_notes or "-")
            ws.cell(row=row, column=10, value=record.created_at.strftime("%Y-%m-%d %H:%M:%S"))

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 18

        summary_ws = wb.create_sheet("统计汇总")
        summary_ws["A1"] = "颜料替代报告统计汇总"
        summary_ws["A1"].font = Font(bold=True, size=14)
        summary_ws["A3"] = "总替代次数"
        summary_ws["B3"] = len(records)

        if records:
            avg_delta = sum(r.color_difference for r in records) / len(records)
            max_delta = max(r.color_difference for r in records)
            summary_ws["A4"] = "平均色差"
            summary_ws["B4"] = round(avg_delta, 2)
            summary_ws["A5"] = "最大色差"
            summary_ws["B5"] = round(max_delta, 2)
            large_diff_count = sum(1 for r in records if r.color_difference > 5)
            summary_ws["A6"] = "色差超过5的次数"
            summary_ws["B6"] = large_diff_count

        wb.save(filepath)
    except Exception as e:
        raise ReportGenerationException("substitution", str(e))

    report = models.Report(
        report_type="substitution",
        file_path=filepath,
        generated_by="system",
        parameters=json.dumps({
            "student_id": student_id,
            "start_date": start_date.isoformat() if start_date else None,
            "end_date": end_date.isoformat() if end_date else None
        })
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return schemas.ReportResponse(
        id=report.id,
        report_type=report.report_type,
        file_path=filepath,
        download_url=f"/api/reports/download/{report.id}",
        created_at=report.created_at
    )


def generate_purchase_report(
    db: Session,
    student_id: Optional[int] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None
) -> schemas.ReportResponse:
    query = db.query(models.Purchase)

    if student_id:
        student = db.query(models.Student).filter(models.Student.id == student_id).first()
        if not student:
            raise StudentNotFoundException(student_id)
        query = query.filter(models.Purchase.student_id == student_id)

    if start_date:
        query = query.filter(models.Purchase.created_at >= start_date)
    if end_date:
        query = query.filter(models.Purchase.created_at <= end_date)

    purchases = query.order_by(models.Purchase.created_at.desc()).all()

    filename = f"purchase_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "采购记录"

        headers = ["采购单ID", "学生ID", "学生姓名", "状态", "总金额", "预算警告",
                   "是否含替代", "创建时间", "完成时间", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        total_amount_all = 0
        for row, purchase in enumerate(purchases, 2):
            student = db.query(models.Student).filter(
                models.Student.id == purchase.student_id
            ).first()
            has_substitute = any(item.is_substitute for item in purchase.items)

            ws.cell(row=row, column=1, value=purchase.id)
            ws.cell(row=row, column=2, value=purchase.student_id)
            ws.cell(row=row, column=3, value=student.name if student else "未知")
            status_cell = ws.cell(row=row, column=4, value=purchase.status)
            if purchase.status == "pending_review":
                status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            ws.cell(row=row, column=5, value=purchase.total_amount)
            cell = ws.cell(row=row, column=6, value="是" if purchase.budget_warning else "否")
            if purchase.budget_warning:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            ws.cell(row=row, column=7, value="是" if has_substitute else "否")
            ws.cell(row=row, column=8, value=purchase.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            ws.cell(row=row, column=9, value=purchase.completed_at.strftime("%Y-%m-%d %H:%M:%S") if purchase.completed_at else "-")
            ws.cell(row=row, column=10, value=purchase.notes or "-")

            total_amount_all += purchase.total_amount

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        detail_ws = wb.create_sheet("采购明细")
        detail_headers = ["采购单ID", "颜料ID", "颜料名称", "品牌", "数量", "单价", "小计",
                          "是否替代品", "原颜料ID", "色差", "替代原因"]
        for col, header in enumerate(detail_headers, 1):
            cell = detail_ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")

        detail_row = 2
        for purchase in purchases:
            for item in purchase.items:
                paint = db.query(models.PaintInventory).filter(
                    models.PaintInventory.id == item.paint_id
                ).first()
                detail_ws.cell(row=detail_row, column=1, value=purchase.id)
                detail_ws.cell(row=detail_row, column=2, value=item.paint_id)
                detail_ws.cell(row=detail_row, column=3, value=paint.name if paint else "未知")
                detail_ws.cell(row=detail_row, column=4, value=paint.brand if paint else "未知")
                detail_ws.cell(row=detail_row, column=5, value=item.quantity)
                detail_ws.cell(row=detail_row, column=6, value=item.unit_price)
                detail_ws.cell(row=detail_row, column=7, value=item.unit_price * item.quantity)
                cell = detail_ws.cell(row=detail_row, column=8, value="是" if item.is_substitute else "否")
                if item.is_substitute:
                    cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                detail_ws.cell(row=detail_row, column=9, value=item.original_paint_id or "-")
                detail_ws.cell(row=detail_row, column=10, value=round(item.color_difference, 2) if item.color_difference else "-")
                detail_ws.cell(row=detail_row, column=11, value=item.substitute_reason or "-")
                detail_row += 1

        for col in range(1, len(detail_headers) + 1):
            detail_ws.column_dimensions[chr(64 + col)].width = 15

        summary_ws = wb.create_sheet("统计汇总")
        summary_ws["A1"] = "采购报告统计汇总"
        summary_ws["A1"].font = Font(bold=True, size=14)
        summary_ws["A3"] = "采购单总数"
        summary_ws["B3"] = len(purchases)
        summary_ws["A4"] = "采购总金额"
        summary_ws["B4"] = round(total_amount_all, 2)
        budget_warning_count = sum(1 for p in purchases if p.budget_warning)
        summary_ws["A5"] = "预算警告次数"
        summary_ws["B5"] = budget_warning_count
        substitute_count = sum(1 for p in purchases if any(i.is_substitute for i in p.items))
        summary_ws["A6"] = "含替代品采购单数"
        summary_ws["B6"] = substitute_count

        wb.save(filepath)
    except Exception as e:
        raise ReportGenerationException("purchase", str(e))

    report = models.Report(
        report_type="purchase",
        file_path=filepath,
        generated_by="system",
        parameters=json.dumps({
            "student_id": student_id,
            "start_date": start_date.isoformat() if start_date else None,
            "end_date": end_date.isoformat() if end_date else None
        })
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return schemas.ReportResponse(
        id=report.id,
        report_type=report.report_type,
        file_path=filepath,
        download_url=f"/api/reports/download/{report.id}",
        created_at=report.created_at
    )


def generate_inventory_report(db: Session) -> schemas.ReportResponse:
    paints = db.query(models.PaintInventory).order_by(models.PaintInventory.id).all()

    filename = f"inventory_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "颜料库存"

        headers = ["ID", "颜料名称", "品牌", "L值", "a值", "b值", "HEX",
                   "库存", "单价", "采购链接", "是否停产", "数据质量", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        total_value = 0
        out_of_stock_count = 0
        for row, paint in enumerate(paints, 2):
            ws.cell(row=row, column=1, value=paint.id)
            ws.cell(row=row, column=2, value=paint.name)
            ws.cell(row=row, column=3, value=paint.brand)
            ws.cell(row=row, column=4, value=paint.l_value)
            ws.cell(row=row, column=5, value=paint.a_value)
            ws.cell(row=row, column=6, value=paint.b_value)
            ws.cell(row=row, column=7, value=paint.hex_code or "-")
            stock_cell = ws.cell(row=row, column=8, value=paint.stock)
            if paint.stock == 0:
                stock_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                out_of_stock_count += 1
            ws.cell(row=row, column=9, value=paint.price)
            ws.cell(row=row, column=10, value=paint.purchase_link or "-")
            cell = ws.cell(row=row, column=11, value="是" if paint.is_discontinued else "否")
            if paint.is_discontinued:
                cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            cell = ws.cell(row=row, column=12, value=paint.data_quality)
            if paint.data_quality == "dirty":
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            ws.cell(row=row, column=13, value=paint.notes or "-")

            total_value += paint.stock * paint.price

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        summary_ws = wb.create_sheet("统计汇总")
        summary_ws["A1"] = "库存报告统计汇总"
        summary_ws["A1"].font = Font(bold=True, size=14)
        summary_ws["A3"] = "颜料种类总数"
        summary_ws["B3"] = len(paints)
        summary_ws["A4"] = "库存总价值"
        summary_ws["B4"] = round(total_value, 2)
        summary_ws["A5"] = "断货种类数"
        summary_ws["B5"] = out_of_stock_count
        discontinued_count = sum(1 for p in paints if p.is_discontinued)
        summary_ws["A6"] = "已停产种类数"
        summary_ws["B6"] = discontinued_count
        dirty_data_count = sum(1 for p in paints if p.data_quality == "dirty")
        summary_ws["A7"] = "脏数据记录数"
        summary_ws["B7"] = dirty_data_count

        wb.save(filepath)
    except Exception as e:
        raise ReportGenerationException("inventory", str(e))

    report = models.Report(
        report_type="inventory",
        file_path=filepath,
        generated_by="system",
        parameters=json.dumps({})
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return schemas.ReportResponse(
        id=report.id,
        report_type=report.report_type,
        file_path=filepath,
        download_url=f"/api/reports/download/{report.id}",
        created_at=report.created_at
    )


def generate_budget_report(
    db: Session,
    student_id: Optional[int] = None
) -> schemas.ReportResponse:
    query = db.query(models.Student)
    if student_id:
        student = db.query(models.Student).filter(models.Student.id == student_id).first()
        if not student:
            raise StudentNotFoundException(student_id)
        query = query.filter(models.Student.id == student_id)

    students = query.order_by(models.Student.id).all()

    filename = f"budget_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "学生预算"

        headers = ["学生ID", "姓名", "学号", "年级", "总预算", "已消费", "剩余预算",
                   "使用率", "数据质量", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        total_budget = 0
        total_used = 0
        over_budget_count = 0
        for row, student in enumerate(students, 2):
            used = student.budget - student.remaining_budget
            usage_rate = (used / student.budget * 100) if student.budget > 0 else 0

            ws.cell(row=row, column=1, value=student.id)
            ws.cell(row=row, column=2, value=student.name)
            ws.cell(row=row, column=3, value=student.student_no)
            ws.cell(row=row, column=4, value=student.grade or "-")
            ws.cell(row=row, column=5, value=student.budget)
            ws.cell(row=row, column=6, value=round(used, 2))
            cell = ws.cell(row=row, column=7, value=student.remaining_budget)
            if student.remaining_budget < 0:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                over_budget_count += 1
            ws.cell(row=row, column=8, value=f"{usage_rate:.1f}%")
            cell = ws.cell(row=row, column=9, value=student.data_quality)
            if student.data_quality == "dirty":
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            ws.cell(row=row, column=10, value=student.notes or "-")

            total_budget += student.budget
            total_used += used

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        summary_ws = wb.create_sheet("统计汇总")
        summary_ws["A1"] = "预算报告统计汇总"
        summary_ws["A1"].font = Font(bold=True, size=14)
        summary_ws["A3"] = "学生总数"
        summary_ws["B3"] = len(students)
        summary_ws["A4"] = "总预算金额"
        summary_ws["B4"] = round(total_budget, 2)
        summary_ws["A5"] = "总已消费金额"
        summary_ws["B5"] = round(total_used, 2)
        summary_ws["A6"] = "整体预算使用率"
        summary_ws["B6"] = f"{(total_used / total_budget * 100):.1f}%" if total_budget > 0 else "0%"
        summary_ws["A7"] = "超预算学生数"
        summary_ws["B7"] = over_budget_count

        wb.save(filepath)
    except Exception as e:
        raise ReportGenerationException("budget", str(e))

    report = models.Report(
        report_type="budget",
        file_path=filepath,
        generated_by="system",
        parameters=json.dumps({"student_id": student_id})
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return schemas.ReportResponse(
        id=report.id,
        report_type=report.report_type,
        file_path=filepath,
        download_url=f"/api/reports/download/{report.id}",
        created_at=report.created_at
    )
