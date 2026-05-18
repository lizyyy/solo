from datetime import datetime
from typing import List, Optional
from fastapi import FastAPI, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill

from database import get_db, engine, Base
from models import ProcessingStatus, ChangeStatus, Frame as FrameModel
import schemas
from services import (
    CustomerService, PrescriptionService, LensOrderService,
    ProcessingStateMachine, DegreeChangeInterceptor, DegreeVersionService,
    PickupService
)

Base.metadata.create_all(bind=engine)

app = FastAPI(
    title="眼镜加工度数版本取件提醒后端API",
    description="管理顾客、验光单、镜片订单、镜架、改度记录、取件报告的完整系统",
    version="1.0.0"
)


@app.exception_handler(ValueError)
async def value_error_handler(request, exc):
    from fastapi.responses import JSONResponse
    return JSONResponse(
        status_code=400,
        content={
            "error_code": "VALIDATION_ERROR",
            "message": str(exc),
            "timestamp": datetime.now().isoformat()
        }
    )


@app.post("/customers/", response_model=schemas.Customer, tags=["顾客管理"])
def create_customer(customer: schemas.CustomerCreate, db: Session = Depends(get_db)):
    existing = CustomerService.get_by_phone(db, customer.phone)
    if existing:
        raise HTTPException(status_code=400, detail="该手机号已注册")
    return CustomerService.create(db, customer)


@app.get("/customers/{customer_id}", response_model=schemas.Customer, tags=["顾客管理"])
def get_customer(customer_id: int, db: Session = Depends(get_db)):
    customer = CustomerService.get_by_id(db, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="顾客不存在")
    return customer


@app.get("/customers/", response_model=List[schemas.Customer], tags=["顾客管理"])
def list_customers(phone: Optional[str] = None, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    from models import Customer
    query = db.query(Customer)
    if phone:
        query = query.filter(Customer.phone.contains(phone))
    return query.offset(skip).limit(limit).all()


@app.post("/prescriptions/", response_model=schemas.Prescription, tags=["验光单管理"])
def create_prescription(prescription: schemas.PrescriptionCreate, db: Session = Depends(get_db)):
    customer = CustomerService.get_by_id(db, prescription.customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="顾客不存在")
    return PrescriptionService.create(db, prescription)


@app.get("/prescriptions/customer/{customer_id}", response_model=List[schemas.Prescription], tags=["验光单管理"])
def get_customer_prescriptions(customer_id: int, db: Session = Depends(get_db)):
    return PrescriptionService.get_by_customer(db, customer_id)


@app.post("/frames/", response_model=schemas.Frame, tags=["镜架管理"])
def create_frame(frame: schemas.FrameCreate, db: Session = Depends(get_db)):
    existing = db.query(FrameModel).filter(FrameModel.sku == frame.sku).first()
    if existing:
        raise HTTPException(status_code=400, detail="该SKU已存在")
    db_frame = FrameModel(**frame.model_dump())
    db.add(db_frame)
    db.commit()
    db.refresh(db_frame)
    return db_frame


@app.get("/frames/", response_model=List[schemas.Frame], tags=["镜架管理"])
def list_frames(sku: Optional[str] = None, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    query = db.query(FrameModel)
    if sku:
        query = query.filter(FrameModel.sku.contains(sku))
    return query.offset(skip).limit(limit).all()


@app.post("/lens-orders/", response_model=schemas.LensOrder, tags=["镜片订单管理"])
def create_lens_order(order: schemas.LensOrderCreate, db: Session = Depends(get_db)):
    customer = CustomerService.get_by_id(db, order.customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="顾客不存在")

    from models import Prescription
    prescription = db.query(Prescription).filter(Prescription.id == order.prescription_id).first()
    if not prescription:
        raise HTTPException(status_code=404, detail="验光单不存在")

    if order.frame_id:
        frame = db.query(FrameModel).filter(FrameModel.id == order.frame_id).first()
        if not frame:
            raise HTTPException(status_code=404, detail="镜架不存在")

    return LensOrderService.create(db, order)


@app.get("/lens-orders/{order_id}", response_model=schemas.LensOrderDetail, tags=["镜片订单管理"])
def get_lens_order(order_id: int, db: Session = Depends(get_db)):
    order = LensOrderService.get_by_id(db, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")
    return order


@app.get("/lens-orders/", tags=["镜片订单管理"])
def list_lens_orders(
    status: Optional[ProcessingStatus] = None,
    customer_id: Optional[int] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db)
):
    skip = (page - 1) * page_size
    orders, total = LensOrderService.list_orders(db, status, customer_id, skip, page_size)
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": orders
    }


@app.patch("/lens-orders/{order_id}/status", tags=["加工状态管理"])
def update_order_status(
    order_id: int,
    status_update: schemas.LensOrderStatusUpdate,
    db: Session = Depends(get_db)
):
    order = LensOrderService.get_by_id(db, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")

    success, message = ProcessingStateMachine.transition(
        db, order, status_update.status, status_update.changed_by, status_update.notes
    )

    if not success:
        raise HTTPException(status_code=400, detail=message)

    db.commit()
    db.refresh(order)
    return {
        "success": True,
        "message": message,
        "current_status": order.status
    }


@app.get("/lens-orders/{order_id}/status-history", response_model=List[schemas.ProcessingStatusHistory], tags=["加工状态管理"])
def get_order_status_history(order_id: int, db: Session = Depends(get_db)):
    from models import ProcessingStatusHistory
    return db.query(ProcessingStatusHistory).filter(ProcessingStatusHistory.lens_order_id == order_id).order_by(ProcessingStatusHistory.changed_at.desc()).all()


@app.post("/degree-changes/", response_model=schemas.DegreeChangeRecord, tags=["改度管理"])
def request_degree_change(change_request: schemas.DegreeChangeRequest, db: Session = Depends(get_db)):
    order = LensOrderService.get_by_id(db, change_request.lens_order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")

    try:
        record = DegreeChangeInterceptor.check_and_record(db, change_request, order)
        db.commit()
        db.refresh(record)
        return record
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/degree-changes/{change_id}/review", response_model=schemas.DegreeChangeRecord, tags=["改度管理"])
def review_degree_change(
    change_id: int,
    review: schemas.DegreeChangeReview,
    db: Session = Depends(get_db)
):
    from models import DegreeChangeRecord
    record = db.query(DegreeChangeRecord).filter(DegreeChangeRecord.id == change_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="改度记录不存在")

    if record.status != ChangeStatus.PENDING_REVIEW.value:
        raise HTTPException(status_code=400, detail="该改度记录已处理")

    if not record.can_apply and review.status == ChangeStatus.APPROVED:
        raise HTTPException(status_code=400, detail="该改度已被拦截，无法批准")

    record.status = review.status.value
    record.reviewed_by = review.reviewed_by
    record.reviewed_at = datetime.now()
    record.review_notes = review.review_notes

    if review.status == ChangeStatus.APPROVED:
        import json
        changes = json.loads(record.new_value)
        new_prescription = DegreeVersionService.create_new_version(
            db, record.prescription_id, changes, review.reviewed_by
        )
        record.new_prescription_id = new_prescription.id

        from models import LensOrder
        order = db.query(LensOrder).filter(LensOrder.id == record.lens_order_id).first()
        if order:
            order.prescription_id = new_prescription.id

        record.applied_at = datetime.now()
        record.applied_by = review.reviewed_by
        record.processing_conclusion = "已批准并应用新度数"

    else:
        record.processing_conclusion = f"已拒绝: {review.review_notes}"

    db.commit()
    db.refresh(record)
    return record


@app.get("/degree-changes/", response_model=List[schemas.DegreeChangeRecord], tags=["改度管理"])
def list_degree_changes(
    order_id: Optional[int] = None,
    status: Optional[ChangeStatus] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    from models import DegreeChangeRecord
    query = db.query(DegreeChangeRecord)
    if order_id:
        query = query.filter(DegreeChangeRecord.lens_order_id == order_id)
    if status:
        query = query.filter(DegreeChangeRecord.status == status.value)
    return query.order_by(DegreeChangeRecord.requested_at.desc()).offset(skip).limit(limit).all()


@app.post("/pickup-reports/{order_id}/reminder", response_model=schemas.PickupReport, tags=["取件管理"])
def send_pickup_reminder(
    order_id: int,
    reminder: schemas.PickupReminder,
    db: Session = Depends(get_db)
):
    order = LensOrderService.get_by_id(db, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")

    if order.status != ProcessingStatus.READY_FOR_PICKUP.value:
        raise HTTPException(status_code=400, detail=f"订单状态不是待取件，当前状态: {order.status}")

    try:
        report = PickupService.send_reminder(db, order_id, reminder.reminder_type, reminder.sent_by, reminder.notes)
        db.commit()
        db.refresh(report)
        return report
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/pickup-reports/{order_id}/confirm", tags=["取件管理"])
def confirm_pickup(
    order_id: int,
    confirmation: schemas.PickupConfirmation,
    db: Session = Depends(get_db)
):
    try:
        report, order = PickupService.confirm_pickup(
            db, order_id, confirmation.picked_up_by, confirmation.pickup_notes
        )
        db.commit()
        db.refresh(report)
        return {
            "success": True,
            "message": "取件确认成功",
            "report": report,
            "order": order
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/pickup-reports/{order_id}", response_model=schemas.PickupReport, tags=["取件管理"])
def get_pickup_report(order_id: int, db: Session = Depends(get_db)):
    try:
        return PickupService.get_or_create_report(db, order_id)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/lens-orders/{order_id}/withdraw", response_model=schemas.LensOrder, tags=["订单管理"])
def withdraw_order(
    order_id: int,
    withdraw: schemas.OrderWithdraw,
    db: Session = Depends(get_db)
):
    try:
        return LensOrderService.withdraw_order(db, order_id, withdraw.reason, withdraw.changed_by)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/lens-orders/{order_id}/manual-correction", tags=["人工修正"])
def manual_correction(
    order_id: int,
    correction: schemas.ManualCorrection,
    db: Session = Depends(get_db)
):
    order = LensOrderService.get_by_id(db, order_id)
    if not order:
        raise HTTPException(status_code=404, detail="订单不存在")

    from models import ProcessingStatusHistory
    history = ProcessingStatusHistory(
        lens_order_id=order_id,
        from_status=f"手动修正: {correction.field_name}",
        to_status=f"{correction.old_value} -> {correction.new_value}",
        changed_by=correction.corrected_by,
        notes=f"人工修正: {correction.reason}"
    )
    db.add(history)

    if hasattr(order, correction.field_name):
        setattr(order, correction.field_name, correction.new_value)

    db.commit()
    return {
        "success": True,
        "message": "人工修正已记录",
        "correction": correction.model_dump()
    }


@app.post("/export/lens-orders", tags=["报告导出"])
def export_lens_orders(filter: schemas.ExportFilter, db: Session = Depends(get_db)):
    from models import LensOrder, Customer
    query = db.query(LensOrder).join(Customer)

    if filter.status:
        query = query.filter(LensOrder.status.in_([s.value for s in filter.status]))
    if filter.date_from:
        query = query.filter(LensOrder.created_at >= filter.date_from)
    if filter.date_to:
        query = query.filter(LensOrder.created_at <= filter.date_to)
    if filter.customer_id:
        query = query.filter(LensOrder.customer_id == filter.customer_id)

    orders = query.order_by(LensOrder.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "镜片加工订单"

    headers = ["订单号", "顾客姓名", "手机号", "状态", "镜片类型(右)", "镜片类型(左)", "创建时间", "状态更新时间", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    for row, order in enumerate(orders, 2):
        ws.cell(row=row, column=1, value=order.order_no)
        ws.cell(row=row, column=2, value=order.customer.name)
        ws.cell(row=row, column=3, value=order.customer.phone)
        ws.cell(row=row, column=4, value=order.status)
        ws.cell(row=row, column=5, value=order.lens_type_od)
        ws.cell(row=row, column=6, value=order.lens_type_os)
        ws.cell(row=row, column=7, value=order.created_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row, column=8, value=order.status_updated_at.strftime("%Y-%m-%d %H:%M:%S"))
        ws.cell(row=row, column=9, value=order.notes or "")

    for col in range(1, 10):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=lens_orders_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"}
    )


@app.get("/export/pickup-reports", tags=["报告导出"])
def export_pickup_reports(db: Session = Depends(get_db)):
    from models import PickupReport, LensOrder, Customer
    reports = db.query(PickupReport).join(LensOrder).join(Customer).order_by(PickupReport.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "取件报告"

    headers = ["报告编号", "订单号", "顾客姓名", "手机号", "质检状态", "待取件时间", "第一次提醒", "第二次提醒", "取件时间", "取件人"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    for row, report in enumerate(reports, 2):
        ws.cell(row=row, column=1, value=report.report_no)
        ws.cell(row=row, column=2, value=report.lens_order.order_no)
        ws.cell(row=row, column=3, value=report.lens_order.customer.name)
        ws.cell(row=row, column=4, value=report.lens_order.customer.phone)
        ws.cell(row=row, column=5, value="合格" if report.quality_pass else "不合格")
        ws.cell(row=row, column=6, value=report.pickup_ready_date.strftime("%Y-%m-%d %H:%M:%S") if report.pickup_ready_date else "")
        ws.cell(row=row, column=7, value=report.first_reminder_date.strftime("%Y-%m-%d %H:%M:%S") if report.first_reminder_date else "")
        ws.cell(row=row, column=8, value=report.second_reminder_date.strftime("%Y-%m-%d %H:%M:%S") if report.second_reminder_date else "")
        ws.cell(row=row, column=9, value=report.pickup_date.strftime("%Y-%m-%d %H:%M:%S") if report.pickup_date else "")
        ws.cell(row=row, column=10, value=report.picked_up_by or "")

    for col in range(1, 11):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=pickup_reports_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"}
    )


@app.get("/health", tags=["系统"])
def health_check():
    return {"status": "healthy", "timestamp": datetime.now().isoformat()}
